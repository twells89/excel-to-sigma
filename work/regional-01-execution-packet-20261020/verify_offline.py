"""Read-only 13-source scaling and actual-Excel conformance accounting."""
import argparse
from collections import Counter,defaultdict
import csv
import hashlib
import json
from pathlib import Path
import sys

parser=argparse.ArgumentParser(description=__doc__)
for name in ('repo','task-root','out'):parser.add_argument('--'+name,required=True)
a=parser.parse_args();root=Path(a.task_root);repo=Path(a.repo);out=Path(a.out);out.mkdir(parents=True,exist_ok=True)
sys.path[:0]=[str(repo/'skills/excel-to-sigma/scripts/lib'),str(repo/'skills/excel-to-sigma/tests')]
from native_exact_plan import build_exact_native_calculation_plan,lower_exact_native_calculations
from native_exact import compile_expression,evaluate_bound,Literal
from native_calculations import _ast
from workbook_expression import WorkbookContext
from native_formula_reference import evaluate_native
from openpyxl import load_workbook

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def write(name,value):(out/name).write_text(json.dumps(value,indent=2,ensure_ascii=False)+'\n')
def ast_counts(node,kinds,functions):
    if isinstance(node,dict):
        if 'kind'in node:kinds[node['kind']]+=1
        if node.get('kind')=='function':functions[node['name']]+=1
        for v in node.values():ast_counts(v,kinds,functions)
    elif isinstance(node,list):
        for v in node:ast_counts(v,kinds,functions)

sources={};status=Counter();identities=[]
for file in sorted((root/'work/complete-source-record-bindings-plans-20261020').glob('*.json')):
    inputs=json.loads(file.read_text());name=inputs['sourceName']
    source=root/'outputs'/name if name.startswith('fixture-') else root/'work/broader-eight-source-replay-20261020/source-originals'/name
    native=build_exact_native_calculation_plan(source,inputs)
    prior=json.loads((root/'work/complete-native-exact-replay-r1-20261020'/file.name).read_text())
    if native['nativePlanHash']!=prior['nativePlanHash'] or native['sourceSha256']!=prior['sourceSha256']:
        raise ValueError('source/plan identity changed: '+name)
    assigned={f:{'elementId':'TEST_ONLY_'+f,'name':s['name'],'rowIdColumn':'Row ID',
                  'rows':{str(i):f'00000000-0000-0000-0000-{i:012d}'for i in range(len(s['rows']))}}
              for f,s in inputs['sharedInputArchitecture']['inputs'].items()}
    lowered=lower_exact_native_calculations(native,assigned)
    categories=defaultdict(lambda:{'tables':0,'columns':0,'formulaCharacters':0,'widthHistogram':Counter()})
    for shard in lowered['shards']:
        kind=('formal-row-template'if shard['name'].startswith('Native Rows ') else
              'criteria-row-helper'if ' Criteria 'in shard['name'] else 'range-aggregate'if 'rangeDomain'in shard else
              'public-output'if shard['sourceFormulaCells']else 'scalar-helper')
        c=categories[kind];c['tables']+=1;c['columns']+=len(shard['columns']);c['widthHistogram'][len(shard['columns'])]+=1
        c['formulaCharacters']+=sum(len(column['formula'])for column in shard['columns'])
    kinds=Counter();functions=Counter()
    for record in native['formulas'].values():ast_counts(record['ast'],kinds,functions)
    counts=Counter(r['status']for r in native['formulas'].values());status.update(counts)
    sources[name]={'sourceSha256':native['sourceSha256'],'nativePlanHash':native['nativePlanHash'],
        'inputPlanHash':inputs['inputPlanHash'],'sourceGroup':'originalFive' if name.startswith('fixture-') else 'broaderEight','formulaCount':native['formulaCount'],'statusCounts':dict(counts),
        'inputRecordProfiles':{f:{'rows':len(s['rows']),'editableColumns':len(s['columns']),'naturalKeys':s['keyColumns']}
                               for f,s in inputs['sharedInputArchitecture']['inputs'].items()},
        'tableCount':len(lowered['shards']),'columnCount':sum(len(e['columns'])for e in lowered['shards']),
        'helperCount':lowered['helperCount'],'maxDependencyDepth':lowered['maxDependencyDepth'],
        'maxFormulaLength':lowered['maxFormulaLength'],'categories':dict(categories),'astNodes':dict(kinds),
        'functionNodes':dict(functions),'runtimeStatus':'UNVERIFIED_CANDIDATE',
        'constructionIdentities':'TEST_ONLY, never assigned target IDs or staging receipts'}
    print(name, len(lowered['shards']),sources[name]['columnCount'])
if sum(status.values())!=3651:raise ValueError('3651 source denominator changed')
write('scaling.json',{'formulaCount':3651,'denominators':{'originalFive':1865,'broaderEight':1786},'statusCounts':dict(status),'sources':sources})

packet=root/'work/six-excel-authoritative-oracles-20261020'
vector=packet/'native-semantics-vectors.xlsx';manifest=json.loads((packet/'native-semantics-manifest.json').read_text())
capture=[json.loads(line)for line in (packet/'native-semantics-excel-oracle.jsonl').read_text().splitlines()if line]
expected={(x['sheet'],x['address']):x for x in manifest['formulaInventory']}
if len(capture)!=30 or len(expected)!=30 or len({(x['sheet'],x['address'])for x in capture})!=30:
    raise ValueError('native oracle denominator or uniqueness changed')
if sha(vector)!=manifest['workbookSha256']:raise ValueError('native vector source hash changed')
book=load_workbook(vector,data_only=False);context=WorkbookContext.from_path(vector);rounds=[];diff=[]
for row in capture:
    identity=(row['sheet'],row['address']);record=expected[identity]
    if (row['excel']['sheet'],row['excel']['address'])!=identity:raise ValueError('native capture nested identity differs')
    if row['workbookSha256']!=sha(vector) or row['sourceFormula']!=record['sourceFormula'] or book[row['sheet']][row['address']].value!=record['sourceFormula']:
        raise ValueError('oracle source identity changed')
    if row['excel']['formula2']!=record['sourceFormula']:diff.append(identity)
    if record['sourceFormula'].startswith('=ROUNDUP('):
        bound=_ast(context.bind_formula(record['sourceFormula'],sheet=row['sheet'],address=row['address']))
        values={f'{s}!{cell.coordinate}':cell.value for s in book.sheetnames for cells in book[s]for cell in cells if cell.data_type!='f'}
        source=evaluate_bound(bound,lambda c:values.get(c));helpers={}
        def emit(formula):
            token=f'[oracle_helper_{len(helpers)}]';helpers[token]=evaluate_native(formula,helpers);return token
        emitted=compile_expression(bound,lambda c:Literal.value(values.get(c)),emit=emit)
        result=evaluate_native(emitted.number,helpers);error=evaluate_native(emitted.error,helpers)
        actual=float(row['excel']['rawValue'])
        if source!=actual or result!=actual or error!='':raise ValueError('ROUNDUP oracle disagreement')
        rounds.append({'address':row['address'],'number':book[row['sheet']]['A'+row['address'][1:]].value,
                       'digits':book[row['sheet']]['B'+row['address'][1:]].value,'actualExcel':actual,
                       'sourceEvaluator':source,'emittedScalarReference':result,'status':'PASS'})
if diff or len(rounds)!=6:raise ValueError('native formula2 or ROUNDUP conformance changed')
write('native-semantics-conformance.json',{'workbookSha256':sha(vector),
    'captureSha256':sha(packet/'native-semantics-excel-oracle.jsonl'),
    'captureIntegrity':{'passed':30,'formula2ExactTextDifferences':len(diff)},'roundupConformance':rounds,
    'booleanCriteria':{'capturedVectors':24,'compilerContract':'UNSUPPORTED_UNCHANGED',
                      'note':'Actual capture is retained; no unsupported type-domain contract is expanded.'},
    'sigmaRuntimeStatus':'UNVERIFIED_CANDIDATE'})

source=root/'work/broader-eight-source-replay-20261020/source-originals/01 Regional Sales Forecast.xlsx'
regional=build_exact_native_calculation_plan(source,json.loads((root/'work/complete-source-record-bindings-plans-20261020/01-regional-sales-forecast-xlsx.json').read_text()))
context=WorkbookContext.from_path(source);raw=[]
for file in sorted((packet/'raw').glob('01-*.csv')):
    for row in csv.DictReader(file.open()):raw.append({**row,'captureFile':str(file),'captureFileSha256':sha(file)})
if len(raw)!=214 or len({(r['sheet'],r['address'])for r in raw})!=214:raise ValueError('Regional capture incomplete')
errors=[];numeric_diff=[];formula_diff=[];overflow=[]
for row in raw:
    cell=row['sheet']+'!'+row['address'];record=regional['formulas'][cell]
    if row['formula2']!=record['originalFormula']:
        rebound=_ast(context.bind_formula(row['formula2'],sheet=row['sheet'],address=row['address']))
        formula_diff.append({'sourceCellId':cell,'sourceFormula':record['originalFormula'],'actualFormula2':row['formula2'],
                             'boundAstEqual':rebound==record['ast'],'status':'REVIEW_REQUIRED_STRING_DIFFERENCE'})
    if row['valueClass']=='missing' and row['display'].startswith('#'):
        errors.append({'sourceCellId':cell,'sourceFormula':record['originalFormula'],'displayedExcelError':row['display'],
                       'actualValueClass':row['valueClass'],'actualSerialClass':row['serialClass'],
                       'localCandidateValue':record.get('localExpected'),'status':'SOURCE_ERROR_RECONCILIATION_REQUIRED'})
    elif row['valueClass']=='real':
        actual=float(row['rawValue']);local=record.get('localExpected')
        if abs(actual-local)>1e-6:numeric_diff.append({'sourceCellId':cell,'actualExcel':actual,'localCandidate':local})
        if row['display'].startswith('#'):overflow.append({'sourceCellId':cell,'actualNumericValue':actual,'display':row['display'],'classification':'numeric-display-overflow'})
write('regional-excel-review.json',{'sourceSha256':sha(source),'captureCount':214,'valueClasses':dict(Counter(r['valueClass']for r in raw)),
    'sourceErrors':errors,'numericDifferences':numeric_diff,'displayOverflows':overflow,
    'formula2StringDifferences':formula_diff,'runtimeStatus':'SOURCE_REVIEW_PENDING_NATIVE_UNQUALIFIED'})
