"""Reported A1+1 counterexample must not escape literal-input staging scope.

Claimed target/evidence fields below are counterfeit unit data, deliberately
complete enough that legacy full-scope gates accept them. Scope must reject
before these claims can authorize native/full/production work.
"""
import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
from openpyxl import Workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'scripts/lib'),str(Path(__file__).resolve().parent)]
from source_inputs import build_source_input_plan
from migration_manifest import build_manifest,validate_manifest,validate_staging_approval,validate_completion
from migration_gate import gate_staging_manifest,gate_completion_manifest,gate_plan_staging_source
from input_orchestration import gate_input_only_staging,gate_input_staging,gate_input_completion
from input_sync import (build_artifact_records,emit_construction_program,emit_workbook_wiring,
                        emit_write_programs,plan_seed,plan_input_extension)
from native_calculations import build_native_calculation_plan,prepare_native_calculation_update,build_native_calculation_backup,_observed_bindings
from native_exact_plan import build_exact_native_calculation_plan,prepare_exact_native_scaffold,prepare_exact_native_update
import acceptance
import promotion
import test_promotion


def script(name):
    spec=importlib.util.spec_from_file_location(name.replace('-','_'),ROOT/'scripts'/name)
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);return module


class InputScopeEscapeTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.folder=Path(self.temp.name);self.source=self.folder/'reported.xlsx'
        book=Workbook();book.active['A1']=1;book.active['A2']='=A1+1';book.save(self.source)
        decisions={'sourceSha256':hashlib.sha256(self.source.read_bytes()).hexdigest(),'features':[{
            'featureId':'input','role':'input','reviewed':True,'keyColumns':['Value'],'literalColumns':['Value'],
            'columnTypes':{'Value':'number'},'recordBindings':[{'sourceCells':{'Value':'Sheet!A1'}}]}]}
        self.inputs=build_source_input_plan(self.source,decisions,connection_id='unit-connection',folder_id='unit-folder')

    def material(self,*,status='live',scope='input-tables-only',claimed=False):
        plan={'sourcePath':str(self.source),'sourceSha256':self.inputs['sourceSha256'],
              'inputPlanHash':self.inputs['inputPlanHash'],'sharedInputArchitecture':self.inputs['sharedInputArchitecture'],
              'formulaCells':[{'id':'Sheet!A2','mappingStatus':status}]}
        if scope is not None:plan['executionScope']=scope
        m=build_manifest(self.source,plan);m['target']={'documentId':'doc','version':1}
        target={**m['target'],'pageId':'page','elementId':'input-element','columnIds':{'Value':'value-column'},'rowIdColumnId':'row-id'}
        a={'status':'approved_for_staging','scope':'input-tables-only','sourceHash':m['sourceHash'],'planHash':m['planHash'],
           'owner':'unit owner','reviewer':'unit reviewer','reviewItems':['Sheet!A2'],'pendingFormulaIds':['Sheet!A2'],
           'mappedFeatures':['input'],'targetBindings':{'input':target}}
        _,audit=gate_plan_staging_source(plan,source_arg=str(self.source))
        receipt={'status':'APPROVED_STAGING','sourceHash':m['sourceHash'],'planHash':m['planHash'],
                 'targetBindings':a['targetBindings'],'mappedFeatures':['input']}
        receipt['artifacts']=build_artifact_records(m,self.inputs['sharedInputArchitecture'],receipt['targetBindings'])
        if claimed:
            m['formulaCells'][0]['targetBinding']={'documentId':'doc','version':1,'elementId':'el','columnId':'col'}
            m['formulaCells'][0]['verification']={'status':'passed','evidenceId':'claimed','sourceCellId':'Sheet!A2','targetId':'col',
                'sourceHash':m['sourceHash'],'planHash':m['planHash'],'targetVersion':1,'targetDocumentId':'doc'}
        return plan,m,a,audit,receipt

    def test_reported_one_cell_counterexample_rejects_full_staging_and_completion_entrypoints(self):
        _,m,a,audit,receipt=self.material(claimed=True)
        calls={'validate-staging':lambda:validate_staging_approval(m,a),
               'gate-staging':lambda:gate_staging_manifest(m,a),
               'input-full-staging':lambda:gate_input_staging(m,a,self.inputs,audit),
               'validate-completion':lambda:validate_completion(m),
               'gate-completion':lambda:gate_completion_manifest(m),
               'input-completion':lambda:gate_input_completion(m,receipt,{'status':'passed'})}
        for name,call in calls.items():
            with self.subTest(entrypoint=name),self.assertRaisesRegex(ValueError,'input-tables-only'):call()
        # Same one-cell declared live/evidence fixture, without the new scope,
        # preserves legacy behavior; these tests concern scope, not authenticity.
        _,full,approval,_,_=self.material(scope=None,claimed=True)
        approval.pop('scope')
        self.assertTrue(validate_staging_approval(full,approval));self.assertTrue(validate_completion(full))

    def test_input_gate_requires_all_unresolved_and_no_runtime_target_claims(self):
        for status in ('live','verification','cached'):
            _,m,a,audit,_=self.material(status=status)
            with self.subTest(status=status),self.assertRaisesRegex(ValueError,'unresolved'):
                gate_input_only_staging(m,a,self.inputs,audit,source_path=self.source)
        for field,value in [('targetMapping',{'logicalCheckId':'claimed'}),('evidence',[{'status':'passed'}]),('execution',{'status':'complete'}),('runtimeStatus','VERIFIED'),('targetBinding',{'elementId':'claimed'}),('actualTarget',{'elementId':'claimed'}),
                            ('runtimeEvidence',{'status':'passed'}),('verification',{'status':'passed'}),
                            ('runtime',{'status':'verified'}),('executionState',{'status':'complete'})]:
            _,m,a,audit,_=self.material(status='unresolved');m['formulaCells'][0][field]=value
            with self.subTest(field=field),self.assertRaisesRegex(ValueError,'runtime|target'):
                gate_input_only_staging(m,a,self.inputs,audit,source_path=self.source)
        _,m,a,audit,_=self.material(status='unresolved')
        receipt=gate_input_only_staging(m,a,self.inputs,audit,source_path=self.source)
        self.assertEqual(receipt['status'],'APPROVED_INPUT_STAGING')
        self.assertIn('upsert_rows',emit_construction_program(m,receipt,'input'))

    def test_native_and_full_input_wiring_cannot_accept_forged_full_receipt(self):
        _,m,a,audit,receipt=self.material(claimed=True)
        current={'workbookId':'doc','documentVersion':1,'contents':{'kind':'workbook','pages':[{'id':'page'}],
          'elements':[{'id':'input-element','kind':'input-table','name':'input','columns':[
            {'id':'value-column','name':'Value','type':'number'},{'id':'row-id','name':'Row ID','type':None}]}],
          'layout':'<Page id="page" />'}}
        rows={'input':[{'Value':1,'Row ID':'physical-row'}]};backup=build_native_calculation_backup(current,rows)
        legacy=build_native_calculation_plan(self.source,self.inputs);exact=build_exact_native_calculation_plan(self.source,self.inputs)
        calls={
            'legacy-native':lambda:prepare_native_calculation_update(legacy,self.inputs,m,receipt,current,rows,backup=backup),
            'exact-scaffold':lambda:prepare_exact_native_scaffold(exact,self.inputs,m,receipt,current,rows,backup=backup,seed_binding={},page_id='page'),
            'exact-update':lambda:prepare_exact_native_update(exact,self.inputs,m,receipt,current,rows,backup=backup,native_bindings={}),
            'native-input-binding':lambda:_observed_bindings(self.inputs,m,receipt,current['contents'],rows),
            'full-wiring':lambda:emit_workbook_wiring(m,receipt),
            'forged-full-receipt-input-create':lambda:emit_construction_program(m,receipt,'input')}
        for name,call in calls.items():
            with self.subTest(entrypoint=name),self.assertRaisesRegex(ValueError,'input-tables-only'):call()

    def test_production_acceptance_and_promotion_paths_reject_reported_manifest(self):
        _,m,a,audit,receipt=self.material(claimed=True)
        matrix=acceptance.seal_matrix({'version':1,'workflows':{'sales':{'sourceHash':m['sourceHash'],'planHash':m['planHash'],
            'documentId':'doc','version':1,'url':'unit://doc','sourceFormulaCount':1,
            'inputContract':{'manifest':m,'receipt':receipt}}},'cases':[]})
        p,current,rows,sdk,scope=test_promotion.PromotionTests().fixture()
        backup=p.build_promotion_backup(scope,current,rows,sdk,{})
        evidence={'matrix':matrix,'ledger':[],'finalReview':{}}
        fake_plan=p._sealed({'acceptance':evidence,'request':{'body':{'documentVersion':7},
            'path':'/v2/workbooks/original/contents'},'sdkActionsBefore':scope['sdkActionInventory']},'promotionHash')
        calls={'validate-matrix':lambda:acceptance.validate_matrix(matrix),
               'production-matrix':lambda:acceptance.validate_production_matrix(matrix),
               'reports':lambda:acceptance.build_workflow_reports(matrix,[]),
               'integration':lambda:acceptance.run_integration(matrix,None,{},enabled=False),
               'prepare-promotion':lambda:p.prepare_promotion(backup,current,current['contents'],{},evidence),
               'finalize-promotion':lambda:p.finalize_mutation(fake_plan,current,{}),
               'build-acceptance':lambda:acceptance.build_acceptance_matrix(
                   {w:m for w in ('forecast','research','sales','inventory','expense')},
                   {w:{} for w in ('forecast','research','sales','inventory','expense')},
                   receipts={w:receipt for w in ('forecast','research','sales','inventory','expense')})}
        for name,call in calls.items():
            with self.subTest(entrypoint=name),self.assertRaisesRegex(ValueError,'input-tables-only'):call()

    def test_paired_scope_removal_and_receipt_relabel_rejects_every_public_consumer(self):
        # This starts with a genuine gated receipt, then removes BOTH scope tags
        # while preserving the genuine, now-stale plan/artifact hashes.
        _,manifest,approval,audit,_=self.material(status='unresolved')
        receipt=gate_input_only_staging(manifest,approval,self.inputs,audit,source_path=self.source)
        forged=copy.deepcopy(manifest);forged['planProjection'].pop('executionScope')
        relabelled=copy.deepcopy(receipt);relabelled['status']='APPROVED_STAGING';relabelled.pop('scope')
        self.assertEqual(forged['planHash'],receipt['planHash'])
        self.assertEqual(relabelled['artifacts'],receipt['artifacts'])
        with self.assertRaisesRegex(ValueError,'plan hash'):validate_manifest(forged)
        seed=plan_seed([{'Value':1}],[],['Value'],'Row ID')
        matrix=acceptance.seal_matrix({'version':1,'workflows':{'sales':{
            'sourceHash':forged['sourceHash'],'planHash':forged['planHash'],
            'documentId':'doc','version':1,'url':'unit://doc','sourceFormulaCount':1,
            'inputContract':{'manifest':forged,'receipt':relabelled}}},'cases':[]})
        legacy=build_native_calculation_plan(self.source,self.inputs)
        exact=build_exact_native_calculation_plan(self.source,self.inputs)
        p,current,rows,sdk,scope=test_promotion.PromotionTests().fixture()
        backup=p.build_promotion_backup(scope,current,rows,sdk,{})
        evidence={'matrix':matrix,'ledger':[],'finalReview':{}}
        fake_plan=p._sealed({'acceptance':evidence,'request':{'body':{'documentVersion':7},
            'path':'/v2/workbooks/original/contents'},'sdkActionsBefore':scope['sdkActionInventory']},'promotionHash')
        calls={
            'construction':lambda:emit_construction_program(forged,relabelled,'input'),
            'seed-writes':lambda:emit_write_programs(seed,forged,relabelled,'input'),
            'extension':lambda:plan_input_extension(forged,relabelled,'input',[
                {'name':'Value','type':'number','id':'value-column'}]),
            'wiring':lambda:emit_workbook_wiring(forged,relabelled),
            'validate-matrix':lambda:acceptance.validate_matrix(matrix),
            'matrix-scope':lambda:acceptance.require_full_matrix_scope(matrix),
            'native-input-binding':lambda:_observed_bindings(self.inputs,forged,relabelled,{'elements':[]},{},input_binding_only=True),
            'legacy-native':lambda:prepare_native_calculation_update(legacy,self.inputs,forged,relabelled,{},{}),
            'exact-scaffold':lambda:prepare_exact_native_scaffold(exact,self.inputs,forged,relabelled,{}, {},backup={},seed_binding={},page_id='page'),
            'exact-update':lambda:prepare_exact_native_update(exact,self.inputs,forged,relabelled,{}, {},backup={},native_bindings={}),
            'prepare-promotion':lambda:p.prepare_promotion(backup,current,current['contents'],{},evidence),
            'finalize-promotion':lambda:p.finalize_mutation(fake_plan,current,{}),
            'build-acceptance':lambda:acceptance.build_acceptance_matrix(
                {w:forged for w in ('forecast','research','sales','inventory','expense')},
                {w:{} for w in ('forecast','research','sales','inventory','expense')},
                receipts={w:relabelled for w in ('forecast','research','sales','inventory','expense')}),
            'production-matrix':lambda:acceptance.validate_production_matrix(matrix),
            'ledger':lambda:acceptance.validate_ledger(matrix,[]),
            'append':lambda:acceptance.append_evidence([],matrix,{}),
            'reports':lambda:acceptance.build_workflow_reports(matrix,[]),
            'integration':lambda:acceptance.run_integration(matrix,None,{},enabled=False)}
        for name,call in calls.items():
            with self.subTest(entrypoint=name),self.assertRaisesRegex(ValueError,'plan hash'):call()
        cli=script('input-sync.py')
        for name,value in [('manifest',forged),('receipt',relabelled),('rows',[])]:
            (self.folder/(name+'.json')).write_text(json.dumps(value))
        for command,extra in [('emit-create',['--feature','input']),('emit-wiring',[]),
                              ('plan',['--feature','input','--current-rows',str(self.folder/'rows.json'),
                                       '--row-id','Row ID','--snapshot-id','unit-query','--snapshot-age-seconds','0'])]:
            output=self.folder/(command+'-forged.json')
            args=[command,'--manifest',str(self.folder/'manifest.json'),'--receipt',str(self.folder/'receipt.json'),
                  '--out',str(output),*extra]
            with self.subTest(cli=command),self.assertRaisesRegex(ValueError,'plan hash'):cli.main(args)
            with self.subTest(output=command):self.assertFalse(output.exists())

    def test_real_input_only_and_full_receipts_keep_their_allowed_program_paths(self):
        for scope in ('input-tables-only',None):
            _,m,a,audit,_=self.material(status='unresolved' if scope else 'live',scope=scope)
            if scope:
                receipt=gate_input_only_staging(m,a,self.inputs,audit,source_path=self.source)
            else:
                a.pop('scope')
                receipt=gate_input_staging(m,a,self.inputs,audit)
            seed=plan_seed([{'Value':1}],[],['Value'],'Row ID')
            with self.subTest(scope=scope):
                self.assertIn('upsert_rows',emit_construction_program(m,receipt,'input'))
                self.assertIn('upsert_rows',emit_write_programs(seed,m,receipt,'input')[0]['program'])
                matrix=acceptance.seal_matrix({'version':1,'workflows':{'sales':{
                    'inputContract':{'manifest':m,'receipt':receipt}}},'cases':[]})
                if scope:
                    with self.assertRaisesRegex(ValueError,'input-tables-only'):emit_workbook_wiring(m,receipt)
                    with self.assertRaisesRegex(ValueError,'input-tables-only'):acceptance.validate_matrix(matrix)
                else:
                    self.assertEqual(emit_workbook_wiring(m,receipt)['planHash'],m['planHash'])
                    self.assertTrue(acceptance.validate_matrix(matrix))

    def test_program_and_matrix_gates_validate_integrity_beyond_scope_labels(self):
        _,m,_,_,receipt=self.material(scope=None)
        original=self.source.read_bytes()
        for mutation in ('plan-seed','formula-record','receipt-plan','receipt-artifact','source-file'):
            manifest=copy.deepcopy(m);bound=copy.deepcopy(receipt)
            if mutation=='plan-seed':
                manifest['planProjection']['sharedInputArchitecture']['inputs']['input']['rows'][0]['Value']=2
                # Even recomputing the artifact approval cannot bless a stale plan.
                bound['artifacts']=build_artifact_records(manifest,
                    manifest['planProjection']['sharedInputArchitecture'],bound['targetBindings'])
            elif mutation=='formula-record':manifest['formulaCells'][0]['originalFormula']='=A1+2'
            elif mutation=='receipt-plan':bound['planHash']='f'*64
            elif mutation=='receipt-artifact':bound['artifacts']['input']['seedHash']='f'*64
            else:
                book=Workbook();book.active['A1']=2;book.active['A2']='=A1+1';book.save(self.source);book.close()
            matrix=acceptance.seal_matrix({'version':1,'workflows':{'sales':{
                'inputContract':{'manifest':manifest,'receipt':bound}}},'cases':[]})
            calls={'construction':lambda:emit_construction_program(manifest,bound,'input'),
                   'seed':lambda:emit_write_programs(plan_seed([{'Value':1}],[],['Value'],'Row ID'),manifest,bound,'input'),
                   'wiring':lambda:emit_workbook_wiring(manifest,bound),
                   'matrix':lambda:acceptance.validate_matrix(matrix)}
            try:
                for name,call in calls.items():
                    with self.subTest(mutation=mutation,entrypoint=name),self.assertRaises(ValueError):call()
            finally:self.source.write_bytes(original)

    def test_empty_wiring_cannot_skip_manifest_and_receipt_validation(self):
        m=build_manifest(self.source,{'sharedInputArchitecture':{'inputs':{}}})
        receipt={'status':'APPROVED_STAGING','sourceHash':m['sourceHash'],'planHash':m['planHash'],
                 'targetBindings':{},'artifacts':{}}
        m['planProjection']['unexpected-edit']=True
        for entrypoint in ('wiring','matrix'):
            matrix=acceptance.seal_matrix({'version':1,'workflows':{'sales':{
                'inputContract':{'manifest':m,'receipt':receipt}}},'cases':[]})
            with self.subTest(entrypoint=entrypoint),self.assertRaisesRegex(ValueError,'plan hash'):
                if entrypoint=='wiring':emit_workbook_wiring(m,receipt)
                else:acceptance.validate_matrix(matrix)
        m['planProjection'].pop('unexpected-edit');receipt['status']='UNAPPROVED'
        matrix=acceptance.seal_matrix({'version':1,'workflows':{'sales':{
            'inputContract':{'manifest':m,'receipt':receipt}}},'cases':[]})
        with self.assertRaisesRegex(ValueError,'APPROVED_STAGING'):emit_workbook_wiring(m,receipt)
        with self.assertRaisesRegex(ValueError,'APPROVED_STAGING'):acceptance.validate_matrix(matrix)

    def test_public_stage_cli_rejects_escape_and_allows_only_pending_input_phase(self):
        cli=script('input-sync.py')
        for status,inputs_only in [('live',False),('live',True),('unresolved',True)]:
            plan,m,a,audit,_=self.material(status=status)
            for name,value in [('plan',plan),('manifest',m),('approval',a),('inputs',self.inputs)]:
                (self.folder/(name+'.json')).write_text(json.dumps(value))
            args=['stage','--plan',str(self.folder/'plan.json'),'--manifest',str(self.folder/'manifest.json'),
                  '--approval',str(self.folder/'approval.json'),'--source',str(self.source),'--out',str(self.folder/'receipt.json')]
            if inputs_only:args+=['--inputs-only','--input-plan',str(self.folder/'inputs.json')]
            with self.subTest(status=status,inputs_only=inputs_only):
                if status=='live':
                    with self.assertRaises(ValueError):cli.main(args)
                else:
                    cli.main(args);result=json.loads((self.folder/'receipt.json').read_text())
                    self.assertEqual(result['status'],'APPROVED_INPUT_STAGING')
                    self.assertEqual(result['pendingFormulaIds'],['Sheet!A2'])

    def test_manifest_validation_opt_in_does_not_expand_full_scope_and_unknown_scopes_reject(self):
        _,m,_,_,_=self.material(status='unresolved')
        with self.assertRaisesRegex(ValueError,'input-tables-only'):validate_manifest(m)
        self.assertTrue(validate_manifest(m,allow_input_only=True))
        for scope in ('input-only','full','native-formulas','',None):
            plan={'executionScope':scope,'formulaCells':[{'id':'Sheet!A2','mappingStatus':'live'}]}
            bad=build_manifest(self.source,plan)
            with self.subTest(scope=scope),self.assertRaisesRegex(ValueError,'executionScope'):
                validate_manifest(bad,allow_input_only=True)

    def test_readonly_row_binding_cli_still_accepts_genuine_pending_input_receipt(self):
        _,m,a,audit,_=self.material(status='unresolved')
        receipt=gate_input_only_staging(m,a,self.inputs,audit,source_path=self.source)
        current={'workbookId':'doc','documentVersion':1,'contents':{'kind':'workbook','pages':[{'id':'page'}],
          'elements':[{'id':'input-element','kind':'input-table','name':'input','columns':[
            {'id':'value-column','name':'Value','type':'number'},{'id':'row-id','name':'Row ID','type':None}]}]}}
        rows={'input':[{'Value':1,'Row ID':'physical-row'}]}
        for name,value in [('inputs',self.inputs),('manifest',m),('receipt',receipt),('current',current),('rows',rows)]:
            (self.folder/(name+'.json')).write_text(json.dumps(value))
        cli=script('native-calculations.py')
        cli.main(['bind-input-rows','--input-plan',str(self.folder/'inputs.json'),'--manifest',str(self.folder/'manifest.json'),
            '--staging-receipt',str(self.folder/'receipt.json'),'--current',str(self.folder/'current.json'),
            '--observed-rows',str(self.folder/'rows.json'),'--out',str(self.folder/'assigned.json')])
        assigned=json.loads((self.folder/'assigned.json').read_text())
        self.assertEqual(assigned['input']['rows'],{'0':'physical-row'})
        with self.assertRaisesRegex(ValueError,'input-tables-only'):
            gate_staging_manifest(m,a)

    def test_audit_cells_are_pending_in_input_scope_and_legacy_defaults_stay_distinct(self):
        from openpyxl import load_workbook
        book=load_workbook(self.source);book.create_sheet('Audit')['A1']='=Sheet!A2';book.save(self.source)
        plan={'executionScope':'input-tables-only'}
        staged=build_manifest(self.source,plan)
        self.assertEqual({r['mappingStatus']for r in staged['formulaCells']},{'unresolved'})
        self.assertTrue(validate_manifest(staged,allow_input_only=True))
        legacy=build_manifest(self.source,{})
        self.assertEqual(next(r['mappingStatus']for r in legacy['formulaCells']if r['id']=='Audit!A1'),'verification')

    def test_input_only_approval_or_receipt_cannot_be_relabelled_as_full_scope(self):
        _,full,approval,_,receipt=self.material(scope=None,claimed=True)
        with self.assertRaisesRegex(ValueError,'input-tables-only'):
            validate_staging_approval(full,approval)
        full['stagingApproval']={'scope':'input-tables-only','status':'APPROVED_INPUT_STAGING'}
        with self.assertRaisesRegex(ValueError,'input-tables-only'):validate_completion(full)
        full['stagingApproval']=None
        receipt['scope']='input-tables-only'
        matrix=acceptance.seal_matrix({'version':1,'workflows':{'sales':{
            'inputContract':{'manifest':full,'receipt':receipt}}},'cases':[]})
        with self.assertRaisesRegex(ValueError,'input-tables-only'):acceptance.validate_matrix(matrix)

if __name__=='__main__':unittest.main()
