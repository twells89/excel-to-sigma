"""Optional, read-only, source-hash-bound Excel 16.114 comparison.

Run explicitly with the external workbook/oracle directory and an output JSON
path. This is excluded from normal generated-fixture CI.
"""
import argparse
import hashlib
import json
from pathlib import Path
import sys

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts" / "lib"))
from statistical_forecast import SOURCE_SHA256, calculate
from statistical_ets_candidate import CandidateFitError, forecast_with_diagnostics


def compare(actual, expected):
    return {"actual": actual, "excel": expected, "difference": actual - expected,
            "absoluteDifference": abs(actual - expected)}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("oracle_dir", type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    source_hash = hashlib.sha256(args.source.read_bytes()).hexdigest()
    if source_hash != SOURCE_SHA256:
        raise SystemExit("source workbook hash mismatch")
    oracle = json.loads((args.oracle_dir / "statistical-reference.json").read_text())
    vectors = json.loads((args.oracle_dir / "behavior-vectors-runtime.json").read_text())
    if oracle["sourceHash"] != source_hash or vectors["sourceHash"] != source_hash:
        raise SystemExit("Excel oracle source hash mismatch")
    wb = load_workbook(args.source, read_only=True, data_only=False)
    history = wb["History"]
    dates = [history.cell(i, 1).value for i in range(2, 38)]
    skus = ("SKU-A100", "SKU-B200", "SKU-C300")
    rows = [{"Month": dates[i], "SKU": sku, "Demand": history.cell(i+2, j+2).value}
            for i in range(36) for j, sku in enumerate(skus)]
    target_dates = [wb["Forecast"].cell(i, 1).value for i in range(4, 16)]
    params = wb["Params"]
    baseline_fit_diagnostics = {}

    def source_forecast(values, seasonality, horizon):
        outcome = forecast_with_diagnostics(values, seasonality, horizon)
        baseline_fit_diagnostics[skus[len(baseline_fit_diagnostics)]] = outcome["diagnostics"]
        return outcome["forecasts"]

    result = calculate(rows, params["B3"].value, params["B4"].value,
                       params["B5"].value, params["B6"].value, target_dates,
                       ets_forecaster=source_forecast)
    by_address = {}
    for i, period in enumerate(result["Periods"], 4):
        for j, sku in enumerate(skus):
            by_address[f"{chr(66+j)}{i}"] = period["ETS"][sku]
            by_address[f"{chr(69+j)}{i}"] = period["Linear"][sku]
        by_address[f"H{i}"] = period["SelectedSKUTrend"]
        by_address[f"I{i}"] = period["Selected3MMA"]
    for i, summary in enumerate(result["Summary"], 19):
        for col, key in (("B", "FY26ETSTotal"), ("C", "StdDev36M"),
                         ("D", "SafetyStock"), ("E", "ReorderPoint")):
            by_address[f"{col}{i}"] = summary[key]
    comparisons = {}
    for record in oracle["values"]:
        address = record["address"]
        if record["status"] == "excel_source_error":
            comparisons[address] = {"excel": record["excelDisplayedValue"],
                                    "candidate": result["IntendedSequenceDescending"],
                                    "status": "EXPLICIT_INTENDED_EXPRESSION_REPAIR"}
        else:
            comparisons[address] = compare(by_address[address], record["value"])
    vector_sheet = load_workbook(args.oracle_dir / "05 Statistical behavior vectors.xlsx",
                                 read_only=True, data_only=False)["OracleVectors"]
    vector_cols = {"baseline": "B", "last_plus_100": "C", "seasonality_6": "D",
                   "all_zero": "E", "constant_100": "F", "one_missing": "G"}
    vector_comparisons = {}
    for name, col in vector_cols.items():
        values = [vector_sheet[f"{col}{i}"].value for i in range(2, 38)]
        seasonality = 6 if name == "seasonality_6" else 12
        expected = [record["value"] for record in vectors["results"] if record["vector"] == name]
        if len(expected) != 12:
            raise SystemExit("incomplete Excel behavior vector: " + name)
        if any(v is None for v in values):
            from statistical_forecast import interpolate_missing
            values = interpolate_missing(values)
        try:
            outcome = forecast_with_diagnostics(values, seasonality, 12)
            predicted = outcome["forecasts"]
            vector_comparisons[name] = {"status": "NUMERIC_COMPARISON", "comparisonCount": 12,
                                        "diagnostics": outcome["diagnostics"],
                                        "comparisons": [compare(a, e) for a, e in zip(predicted, expected)]}
        except CandidateFitError as exc:
            vector_comparisons[name] = {"status": "CANDIDATE_FIT_FAILED",
                                        "code": exc.code, "diagnostics": exc.diagnostics,
                                        "comparisonCount": 0, "candidateForecasts": None,
                                        "excelValues": expected}
    by_group = {}
    for group, cols in {"ets": "BCD", "linear": "EFG", "selectedTrend": "H",
                        "trailingAverage": "I", "summary": "BCDE"}.items():
        cells = [comparisons[f"{c}{row}"] for row in (range(19, 22) if group == "summary" else range(4, 16)) for c in cols]
        by_group[group] = {"count": len(cells), "maxAbsoluteDifference": max(v["absoluteDifference"] for v in cells),
                           "meanAbsoluteDifference": sum(v["absoluteDifference"] for v in cells)/len(cells)}
    evidence_hashes = {name: hashlib.sha256((args.oracle_dir / name).read_bytes()).hexdigest()
                       for name in ("statistical-reference.json", "behavior-vectors-runtime.json",
                                    "05 Statistical behavior vectors.xlsx")}
    output = {"sourceHash": source_hash, "evidenceHashes": evidence_hashes,
              "excelVersion": oracle["excelVersion"],
              "candidate": "statsmodels 0.15.0 ETSModel(add,add,add; estimated initial states)",
              "readiness": "REVIEW_REQUIRED", "sourceFormulaCount": oracle["formulaCount"],
              "comparedNumeric": len(by_address), "sourceError": "Forecast!K4 #NAME?",
              "groups": by_group, "baseline": comparisons,
              "baselineFitDiagnostics": baseline_fit_diagnostics,
              "vectorExpectedCount": 72,
              "vectorNumericComparisons": sum(v["comparisonCount"] for v in vector_comparisons.values()),
              "vectors": vector_comparisons}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(output, indent=2, allow_nan=False) + "\n")
    print(json.dumps({"groups": by_group,
                      "vectors": {k: (max(vv["absoluteDifference"] for vv in v["comparisons"])
                                      if v["status"] == "NUMERIC_COMPARISON" else {"status": v["status"], "code": v["code"]})
                                  for k, v in vector_comparisons.items()}}, indent=2))


if __name__ == "__main__":
    main()
