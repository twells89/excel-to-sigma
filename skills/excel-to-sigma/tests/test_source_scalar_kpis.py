"""Exact source sums and guarded native append planning; no Sigma execution."""
import copy
import ast
import re
import json
import subprocess
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts' / 'lib'))
try:
    import source_scalar_kpis as scalars
except ModuleNotFoundError:
    scalars = None

SOURCE = ROOT / 'tests' / 'fixtures' / 'fixture-sales-pipeline.xlsx'


def evaluate_scalar(formula, rows):
    """Independent finite-array oracle for the small emitted expression surface."""
    tree = ast.parse(re.sub(r'\[([^]]+)\]', lambda m: 'Column(' + repr(m[1].split('/')[-1]) + ')', formula), mode='eval')
    def evaluate(node):
        if isinstance(node, ast.Constant):
            return Decimal(str(node.value)) if isinstance(node.value, (int, float)) else node.value
        if isinstance(node, ast.BinOp) and isinstance(node.op, ast.Mult):
            return [None if a is None or b is None else a*b
                    for a, b in zip(evaluate(node.left), evaluate(node.right))]
        if isinstance(node, ast.Call):
            args = [evaluate(arg) for arg in node.args]
            if node.func.id == 'Column':
                return [None if row.get(args[0]) is None else Decimal(str(row[args[0]]))
                        if args[0] != 'ID' else row['ID'] for row in rows]
            if node.func.id == 'In': return [value in args[1:] for value in args[0]]
            if node.func.id == 'SumIf':
                return sum((v for v, include in zip(*args) if include and v is not None), Decimal(0))
            if node.func.id == 'Coalesce': return next((value for value in args if value is not None), None)
        raise AssertionError('unexpected emitted expression')
    return evaluate(tree.body)


class TestSourceScalarKPIs(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(scalars, 'native scalar append support is missing')
        self.contract = scalars.extract_sales_overview_scalars(SOURCE)
        self.identity = {
            'sourceHash': self.contract['sourceHash'], 'documentId': 'doc',
            'inputElementId': 'input', 'rowIdColumn': 'ID',
            'evidenceId': 'original-assigned-row-readback', 'reviewer': 'reviewer',
            'rows': [{'sourceRow': row, 'originalKey': f'OPP-{row-1:04}',
                      'physicalId': f'physical-{row}'} for row in range(2, 50)]}
        self.snapshot = {
            'documentId': 'doc', 'version': 51, 'published': True, 'hasDraft': False,
            'input': {'elementId': 'input', 'pageId': 'sales-edit',
                      'name': 'Sales Business Inputs', 'rowIdColumn': 'ID',
                      'columnIds': {'ID': 'ID', 'ACV': 'acv', 'Probability': 'prob'},
                      'filters': [],
                      'rows': [{'ID': f'physical-{r}', 'Opportunity ID': f'OPP-{r-1:04}',
                                'ACV': r * 10, 'Probability': .1} for r in range(2, 50)]},
            'contents': {'pages': [{'id': 'sales-overview'}, {'id': 'sales-edit'}],
                         'elements': [{'id': 'input', 'kind': 'input-table',
                                       'name': 'Sales Business Inputs',
                                       'columns': [{'id': 'ID', 'name': 'ID'},
                                                   {'id': 'acv', 'name': 'ACV'},
                                                   {'id': 'prob', 'name': 'Probability'}]}],
                         'layout': '<Page id="sales-overview"></Page><Page id="sales-edit">'
                                   '<Element elementId="input" gridColumn="1 / 25" gridRow="4 / 28"/></Page>'},
            'nativeActions': {'hosts': [{'id': 'private-host', 'sequences': ['opaque-private-effect']}]}}
        workbook = load_workbook(SOURCE, data_only=False)
        for row, observed in enumerate(self.snapshot['input']['rows'], 2):
            observed.update(ACV=workbook['Pipeline'].cell(row, 10).value,
                            Probability=workbook['Pipeline'].cell(row, 9).value)
        workbook.close()

    def plan(self, ownership=None):
        return scalars.plan_native_scalar_append(self.contract, self.identity,
                                                self.snapshot, ownership=ownership)

    def readback(self, plan):
        after = copy.deepcopy(plan['before'])
        after['version'] += 1
        assigned = {}
        layout = []
        for n, item in enumerate(plan['owners']):
            if item['sourceId'] not in plan['emittedSourceIds']:
                continue
            eid, cid = f'actual-kpi-{n}', f'actual-value-{n}'
            assigned[item['sourceId']] = {'elementId': eid, 'columnId': cid}
            after['contents']['elements'].append({
                'id': eid, 'kind': 'kpi-chart', 'name': item['name'],
                'source': {'kind': 'table', 'elementId': 'input'},
                'columns': [{'id': cid, 'name': item['name'], 'formula': item['formula'],
                             'format': {'kind': 'number', 'formatString': '$,.0f'}}], 'value': {'columnId': cid}})
            x, y, w, h = item['frame']
            layout.append(f'<Element elementId="{eid}" gridColumn="{x} / {x+w}" gridRow="{y} / {y+h}"/>')
        after['contents']['layout'] = after['contents']['layout'].replace(
            '<Page id="sales-overview">', '<Page id="sales-overview">' + ''.join(layout))
        return after, assigned

    def test_exact_source_formula_hashes_and_no_cached_values(self):
        by_source = {x['sourceId']: x for x in self.contract['owners']}
        self.assertEqual(by_source['Overview!B6']['sourceFormulaHash'],
                         '4feef621e9713eece56dec8c65d3627d4191031faea88a4ea2f36ae13ac52959')
        self.assertEqual(by_source['Overview!B7']['sourceFormulaHash'],
                         '5a161c654c1a90a64209f743850c83d626842ebfbaa91a371e30643ba0ec9b05')
        self.assertEqual(by_source['Overview!B6']['sourceCells'], [f'Pipeline!J{r}' for r in range(2, 50)])
        self.assertEqual(by_source['Overview!B7']['numericDependencies'], ['Probability', 'ACV'])
        self.assertFalse(self.contract['formulaCachesUsed'])
        self.assertNotIn('expectedValue', str(self.contract))

    def test_architecture_exposes_scalar_contract_without_full_put_nominal_targets(self):
        import presentation
        plan = presentation.extract_presentation(SOURCE)
        architecture = presentation.build_business_input_architecture(plan, connection_id='conn', folder_id='folder')
        self.assertEqual(architecture.get('sourceScopedScalars'), self.contract)
        self.assertNotIn('Overview!B6', architecture['formulaTargets'])
        self.assertNotIn('Overview!B7', architecture['formulaTargets'])

    def test_other_sales_workbooks_keep_building_with_explicit_unresolved_scalar_scope(self):
        import presentation
        for mutation in ('no-overview', 'different-overview'):
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / 'other-sales.xlsx'
                workbook = load_workbook(SOURCE)
                if mutation == 'no-overview': del workbook['Overview']
                else: workbook['Overview']['B6'] = '=SUM(Pipeline!J2:J49)'
                workbook.save(path)
                workbook.close()
                plan = presentation.extract_presentation(path)
                architecture = presentation.build_business_input_architecture(plan, connection_id='conn', folder_id='folder')
                scalar = architecture['sourceScopedScalars']
                self.assertEqual(scalar['status'], 'UNRESOLVED_UNSUPPORTED_SOURCE_SHAPE')
                self.assertEqual(scalar['sourceHash'], plan['source']['sha256'])
                self.assertEqual(scalar['owners'], [])
                self.assertEqual(scalar['productionStatus'], 'BLOCKED')
                with self.assertRaises(ValueError):
                    scalars.plan_native_scalar_append(scalar, self.identity, self.snapshot)

    def test_prepared_exact_source_contract_cannot_be_reused_after_source_drift(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'sales.xlsx'
            path.write_bytes(SOURCE.read_bytes())
            contract = scalars.extract_sales_overview_scalars(path)
            workbook = load_workbook(path)
            workbook['Overview']['B6'] = '=SUM(Pipeline!J2:J49)'
            workbook.save(path)
            workbook.close()
            with self.assertRaises(ValueError):
                scalars.plan_native_scalar_append(contract, self.identity, self.snapshot)

    def test_native_id_projection_requires_explicit_formula_observation(self):
        self.snapshot['input']['nativeIdFormula'] = '[New Input Table/ID]'
        plan = self.plan()
        self.assertIn("source_id.formula() != '[New Input Table/ID]'", plan['program'])
        self.assertIn("k0.add_source_column('ID')", plan['program'])
        self.assertIn("id0.formula() != '[Sales Business Inputs/ID]'", plan['program'])
        self.snapshot['input']['nativeIdFormula'] = '1'
        with self.assertRaises(ValueError): self.plan()

    def test_partial_assigned_readback_allows_only_the_missing_append(self):
        plan = self.plan()
        after, assigned = self.readback(plan)
        ownership = scalars.bind_native_scalar_owners(plan, after, assigned)
        self.assertEqual(plan['missingSourceIds'], ['Overview!B6', 'Overview!B7'])
        self.assertEqual(plan['emittedSourceIds'], ['Overview!B6'])
        self.assertEqual(set(assigned), {'Overview!B6'})
        foreign_after, foreign_assignments = self.readback({**plan, 'emittedSourceIds': plan['missingSourceIds']})
        with self.assertRaisesRegex(ValueError, 'scope'):
            scalars.bind_native_scalar_owners(plan, foreign_after, foreign_assignments)
        self.snapshot = after
        retry = self.plan(ownership)
        self.assertEqual(retry['missingSourceIds'], ['Overview!B7'])
        self.assertEqual(retry['program'].count('.add_kpi('), 1)

    def test_source_scope_and_weighted_row_equation_drift_fail_closed(self):
        for cell, value in [('B6', "=SUM('Pipeline'!$J$2)"), ('L49', '=I49+J49'),
                            ('J2', '18000'), ('I2', True)]:
            with self.subTest(cell=cell), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / 'changed.xlsx'
                workbook = load_workbook(SOURCE)
                workbook['Overview' if cell == 'B6' else 'Pipeline'][cell] = value
                workbook.save(path)
                workbook.close()
                with self.assertRaisesRegex(ValueError, 'source|equation|scope'):
                    scalars.extract_sales_overview_scalars(path)

    def test_scope_survives_mutable_keys_and_excludes_inserted_duplicate_keys(self):
        baseline = self.plan()
        self.snapshot['input']['rows'][0]['Opportunity ID'] = 'user-renamed'
        self.snapshot['input']['rows'][0]['ACV'] = -1.25
        self.snapshot['input']['rows'].append({'ID': 'inserted', 'Opportunity ID': 'OPP-0001',
                                               'ACV': 1000000, 'Probability': 1})
        changed = self.plan()
        self.assertEqual(baseline['scopeHash'], changed['scopeHash'])
        self.assertEqual(baseline['owners'], changed['owners'])
        self.assertIn('"physical-49"', changed['owners'][0]['formula'])
        self.assertNotIn('inserted', changed['owners'][0]['formula'])
        self.assertNotIn('Opportunity ID', changed['owners'][0]['formula'])
        self.assertNotIn('Settings', changed['owners'][1]['formula'])

    def test_missing_or_duplicated_physical_rows_and_rebound_source_rows_block(self):
        for mutation in ('missing', 'duplicate', 'remapped', 'stale-source'):
            with self.subTest(mutation=mutation):
                identity, snapshot = copy.deepcopy(self.identity), copy.deepcopy(self.snapshot)
                if mutation == 'missing': snapshot['input']['rows'].pop()
                if mutation == 'duplicate': snapshot['input']['rows'].append(snapshot['input']['rows'][0])
                if mutation == 'remapped': identity['rows'][0]['sourceRow'] = 49
                if mutation == 'stale-source': identity['sourceHash'] = 'foreign'
                with self.assertRaises(ValueError):
                    scalars.plan_native_scalar_append(self.contract, identity, snapshot)

    def test_duplicate_or_malformed_input_column_readback_is_rejected(self):
        for extra in ({'id': 'acv', 'name': 'ACV'}, {'id': 'acv', 'name': 'Shadow'},
                      {'id': 'shadow'}, {'name': 'Shadow'}, None):
            with self.subTest(extra=extra):
                snapshot = copy.deepcopy(self.snapshot)
                snapshot['contents']['elements'][0]['columns'].append(extra)
                with self.assertRaises(ValueError):
                    scalars.plan_native_scalar_append(self.contract, self.identity, snapshot)

    def test_native_program_is_narrow_and_plan_replay_requires_inventory_guard(self):
        plan = self.plan()
        self.assertEqual(plan['mode'], 'append-native-sdk')
        self.assertEqual([o['frame'] for o in plan['owners']], [[1, 48, 12, 6], [13, 48, 12, 6]])
        compile(plan['program'], 'native-scalars.py', 'exec')
        self.assertEqual(plan['program'].count('.add_kpi('), 1)
        self.assertIn('host.list_elements()', plan['program'])
        self.assertIn('source.get_column(\'ID\')', plan['program'])
        self.assertIn('sigma.formula.sum_if(', plan['program'])
        self.assertNotIn('upsert_rows', plan['program'])
        self.assertNotIn('contents', plan['program'])
        self.assertEqual(plan['productionStatus'], 'BLOCKED')
        self.assertEqual(plan['verification'], 'UNRESOLVED')

    def test_emitted_formulas_preserve_decimal_vectors_and_empty_sum(self):
        formula_b6, formula_b7 = [owner['formula'] for owner in self.plan()['owners']]
        for acv, probability, expected_b6, expected_b7 in [
            (18000, .1, '4643400', '1741890'), (18001, .1, '4643401', '1741890.10'),
            (18000, .25, '4643400', '1744590'), (-1.25, .1, '4625398.75', '1740089.875'),
            (18000, 0, '4643400', '1740090'), (None, .1, '4625400', '1740090'),
            (18000, None, '4643400', '1740090')]:
            with self.subTest(acv=acv, probability=probability):
                rows = copy.deepcopy(self.snapshot['input']['rows'])
                rows[0].update(ACV=acv, Probability=probability)
                self.assertEqual(evaluate_scalar(formula_b6, rows), Decimal(expected_b6))
                self.assertEqual(evaluate_scalar(formula_b7, rows), Decimal(expected_b7))
        rows = copy.deepcopy(self.snapshot['input']['rows'])
        rows[-1]['ACV'] = 31104
        rows.append({'ID': 'unrelated', 'ACV': 999999999, 'Probability': 1})
        self.assertEqual(evaluate_scalar(formula_b6, rows), Decimal('4643404'))
        self.assertEqual(evaluate_scalar(formula_b7, rows), Decimal('1741891'))
        for row in rows: row['ACV'] = None
        self.assertEqual(evaluate_scalar(formula_b6, rows), 0)
        self.assertEqual(evaluate_scalar(formula_b7, rows), 0)

    def test_inventory_guard_stops_direct_program_replay_before_mutation(self):
        class Page:
            def list_elements(self): return [self]
            def id(self): return 'created-on-first-run'
        class Workbook:
            def get_page(self, name): return Page()
        class Sigma:
            workbook = Workbook()
        tree = ast.parse(self.plan()['program'])
        self.assertIsInstance(tree.body.pop(0), ast.Raise)
        with self.assertRaisesRegex(RuntimeError, 'inventory changed'):
            exec(compile(tree, 'admitted-native.py', 'exec'), {'sigma': Sigma(), 'Page': Page})

    def test_program_requires_external_admission_and_prints_ids_before_fallible_setters(self):
        program = self.plan()['program']
        with self.assertRaisesRegex(RuntimeError, 'external run receipt.*inverse'):
            exec(program, {})
        tree = ast.parse(program)
        for index, statement in enumerate(tree.body):
            if (isinstance(statement, ast.Assign) and isinstance(statement.value, ast.Call)
                    and isinstance(statement.value.func, ast.Attribute)
                    and statement.value.func.attr == 'add_kpi'):
                next_statement = tree.body[index + 1]
                self.assertIsInstance(next_statement, ast.Expr)
                self.assertEqual(next_statement.value.func.id, 'print')
                self.assertEqual(next_statement.value.args[0].right.func.attr, 'id')

    def test_assigned_ids_require_actual_readback_and_retry_is_noop(self):
        plan = self.plan()
        after, assigned = self.readback(plan)
        ownership = scalars.bind_native_scalar_owners(plan, after, assigned)
        self.assertEqual(ownership['targets']['Overview!B6']['elementId'], 'actual-kpi-0')
        self.assertEqual(ownership['targets']['Overview!B6']['version'], 52)
        self.snapshot = after
        retry = self.plan(ownership)
        self.assertEqual(retry['emittedSourceIds'], ['Overview!B7'])
        after, assigned = self.readback(retry)
        ownership = scalars.bind_native_scalar_owners(retry, after, assigned)
        self.assertEqual(ownership['targets']['Overview!B7']['columnId'], 'actual-value-1')
        self.assertEqual(ownership['targets']['Overview!B6']['version'], 53)
        self.snapshot = after
        final = self.plan(ownership)
        self.assertEqual(final['mode'], 'reconcile-existing')
        self.assertEqual(final['program'], '')
        with self.assertRaisesRegex(ValueError, 'collision|ownership'):
            self.plan()

    def test_wrong_formula_grouped_or_stale_owners_cannot_bind(self):
        plan = self.plan()
        for mutation in ('constant', 'grouped', 'foreign-id', 'private-action', 'scope-filter',
                         'unpublished-version', 'failed-format-setter'):
            with self.subTest(mutation=mutation):
                after, assigned = self.readback(plan)
                element = after['contents']['elements'][-1]
                if mutation == 'constant': element['columns'][0]['formula'] = '1741890'
                if mutation == 'grouped': element['groupings'] = [{'groupBy': ['stage']}]
                if mutation == 'foreign-id': assigned['Overview!B6']['elementId'] = 'declared-logical-id'
                if mutation == 'private-action': after['nativeActions'] = {'hosts': []}
                if mutation == 'scope-filter': element['filters'] = [{'formula': 'false'}]
                if mutation == 'unpublished-version': after['version'] = 51
                if mutation == 'failed-format-setter': element['columns'][0].pop('format')
                with self.assertRaises(ValueError):
                    scalars.bind_native_scalar_owners(plan, after, assigned)
        after, assigned = self.readback(plan)
        ownership = scalars.bind_native_scalar_owners(plan, after, assigned)
        after['contents']['elements'][-1]['columns'][0]['formula'] = '0'
        self.snapshot = after
        with self.assertRaisesRegex(ValueError, 'hash|formula'):
            self.plan(ownership)

    def test_cli_emits_plan_and_binds_actual_ids_without_sigma_calls(self):
        cli = ROOT / 'scripts' / 'source-scalar-kpis.py'
        with tempfile.TemporaryDirectory() as directory:
            tmp = Path(directory)
            for name, value in [('identity', self.identity), ('snapshot', self.snapshot)]:
                (tmp / (name + '.json')).write_text(json.dumps(value))
            subprocess.run([sys.executable, str(cli), 'prepare', str(SOURCE), '--out', str(tmp/'contract.json')],
                           check=True, capture_output=True)
            subprocess.run([sys.executable, str(cli), 'plan', '--contract', str(tmp/'contract.json'),
                            '--identity', str(tmp/'identity.json'), '--snapshot', str(tmp/'snapshot.json'),
                            '--out', str(tmp/'plan.json')], check=True, capture_output=True)
            plan = json.loads((tmp/'plan.json').read_text())
            self.assertEqual(plan['missingSourceIds'], ['Overview!B6', 'Overview!B7'])
            after, assigned = self.readback(plan)
            for name, value in [('after', after), ('assigned', assigned)]:
                (tmp/(name+'.json')).write_text(json.dumps(value))
            subprocess.run([sys.executable, str(cli), 'bind', '--plan', str(tmp/'plan.json'),
                            '--snapshot', str(tmp/'after.json'), '--assignments', str(tmp/'assigned.json'),
                            '--out', str(tmp/'owners.json')], check=True, capture_output=True)
            owners = json.loads((tmp/'owners.json').read_text())
            self.assertEqual(owners['targets']['Overview!B6']['elementId'], 'actual-kpi-0')
            self.assertEqual(owners['verification'], 'UNRESOLVED')


if __name__ == '__main__':
    unittest.main()
