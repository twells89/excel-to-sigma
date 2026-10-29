#!/usr/bin/env python3
"""Forecast archetype presentation: source number formats, the source title, chart de-duplication, flat chart
labels, peer KPI cards, the table-report page recipe, ``--look``, and the invariant that none of it changes a
value (the DM spec, parity targets and every emitted value formula equal main's).

Synthetic workbooks come from ``forecast_pnl_fixture`` (formats are set by ``mutate`` hooks); the 02 / 03
regression fixtures are the real ones. ``fixtures/forecast-regression/main-dfa51f4-forecast-values.json`` pins
what main (dfa51f4, before this port) emits for 02, 03 and two legacy plans.
"""
from __future__ import annotations

import contextlib
import copy
import hashlib
import importlib.util
import io
import json
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / "lib"))
sys.path.insert(0, str(ROOT / "tests"))

import forecast_pnl_fixture as fixture  # noqa: E402
import forecast_presentation as fp  # noqa: E402
import presentation_style as ps  # noqa: E402
import workbook_wire  # noqa: E402
from forecast_plan import PlanDraftError, draft_plan  # noqa: E402
from test_forecast_plan import add_metric_chart, find_cell  # noqa: E402

FIXTURES = ROOT / "tests" / "fixtures"
SAAS = FIXTURES / "broader-eight" / "02 SaaS Driver Forecast FY26.xlsx"
RETAIL = FIXTURES / "broader-eight" / "03 Retail Quarterly Budget.xlsx"
GOLDEN = json.loads((FIXTURES / "forecast-regression" / "main-dfa51f4-forecast-values.json").read_text())
ELS = {"glDetail": "glDetail", "spine": "spine", "assumptions": "assumptions"}
PAREN_RED = "#,##0;[Red](#,##0)"
PAREN_DASH = '#,##0;\\(#,##0\\);"-"'


def load_script(filename, name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


BUILD = load_script("build-forecast-model.py", "forecast_presentation_builder")
DRAFT = load_script("draft-forecast-plan.py", "forecast_presentation_draft_cli")

LOOK = {"version": 1,
        "theme": {"overrides": {"colors": {"text": "#121b39"}, "categoricalScheme": ["#4e79a7", "#f28e2b", "#e15759"],
                                "tableStyles": {"cellSpacing": "medium"}}},
        "styles": {"container": {"borderRadius": "pill"}, "kpi-chart": {"borderRadius": "pill"},
                   "chart": {"padding": "none"}, "table": {"borderRadius": "pill"}},
        "kpi": {"name": {"color": "#121b39", "fontSize": 13, "fontWeight": "bold"}, "value": {"fontSize": 30}},
        "text": {"title": {"size": 30, "color": "#121b39"}, "subline": {"size": 13, "color": "#667085"},
                 "color": "#121b39"},
        "tablePreset": "presentation"}


def set_formats(codes, default=None):
    """mutate hook: number formats per P&L row label (``default`` for every other numeric row)."""
    def mutate(pnl, gl, drivers):
        for row in range(fixture.HEADER_ROW + 1, pnl.max_row + 1):
            label = pnl.cell(row, 1).value
            code = codes.get(label, default)
            if code is None:
                continue
            for col in range(2, 2 + len(fixture.MONTHS) + 1):
                if pnl.cell(row, col).value is not None:
                    pnl.cell(row, col).number_format = code
    return mutate


def chain(*hooks):
    def mutate(pnl, gl, drivers):
        for hook in hooks:
            hook(pnl, gl, drivers)
    return mutate


def gl_format(code):
    def mutate(pnl, gl, drivers):
        for row in range(5, gl.max_row + 1):
            gl.cell(row, 6).number_format = code
    return mutate


def style_headers(fill, bold=True, color=None):
    from openpyxl.styles import Font, PatternFill

    def mutate(pnl, gl, drivers):
        for sheet, row in ((pnl, fixture.HEADER_ROW), (gl, 4), (drivers, 3)):
            for cell in sheet[row]:
                if cell.value is not None:
                    cell.fill = PatternFill("solid", fgColor=fill)
                    cell.font = Font(bold=bold, color=color)
    return mutate


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        self.path = self.dir / "pnl.xlsx"

    def draft(self, intent="what-if", look=None, currency_format="$,.0f", pct_format=".1%", **fixture_kwargs):
        fixture.build_pnl(self.path, **fixture_kwargs)
        plan, info = draft_plan(self.path, intent=intent, connection="conn", folder="fold", look=look,
                                currency_format=currency_format, pct_format=pct_format)
        return plan, info

    def spec(self, plan):
        return BUILD.build_wb_spec(plan, "<dm>", ELS)


def elements_of(spec):
    return {e["id"]: e for page in spec["pages"] for e in page["elements"]}


def label(card):
    name = card["name"]
    return name["text"] if isinstance(name, dict) else name


def layout_pages(layout):
    """``{page id: [(element id, c0, c1, r0, r1, child ids)]}`` of a layout string."""
    out = {}
    for match in re.finditer(r'<Page\b[^>]*\bid="([^"]+)"[^>]*>(.*?)</Page>', layout, re.S):
        page, body = match.groups()
        placed = []
        for found in re.finditer(r'<(Element|Container) elementId="([^"]+)"[^>]*gridColumn="(\d+) / (\d+)" '
                                 r'gridRow="(\d+) / (\d+)"[^>]*?(/>|>(.*?)</Container>)', body, re.S):
            kind, element_id, c0, c1, r0, r1, _, inner = found.groups()
            children = re.findall(r'elementId="([^"]+)"', inner or "")
            placed.append((element_id, int(c0), int(c1), int(r0), int(r1), children))
        out[page] = placed
    return out


# ================================================================== 1. number formats
class TestSourceFormats(Base):
    def test_formats_are_read_per_line_section_and_derived_row(self):
        plan, info = self.draft(mutate=set_formats({"EBITDA": PAREN_DASH, "Net Income": "0.0%"}, PAREN_RED))
        formats = plan["presentation"]["formats"]
        self.assertEqual(formats["money"]["format"], {"kind": "number", "formatString": "(,.0f"})
        self.assertTrue(formats["money"]["negativeRed"])
        self.assertEqual(formats["money"]["code"], PAREN_RED)
        self.assertEqual(formats["sections"]["Revenue"]["format"]["formatString"], "(,.0f")
        self.assertNotIn("negativeRed", formats["derived"]["EBITDA"])          # parentheses, no [Red]
        self.assertEqual(formats["derived"]["EBITDA"]["format"]["formatString"], "(,.0f")
        self.assertEqual(formats["derived"]["Net Income"]["format"]["formatString"], ".1%")
        self.assertEqual(formats["derived"]["Gross Profit"]["code"], PAREN_RED)
        self.assertEqual(formats["money"]["cell"], "P&L!B5")

    def test_a_line_whose_format_differs_keeps_the_money_format_and_is_listed(self):
        plan, _ = self.draft(mutate=set_formats({"Services": "0.0%"}, PAREN_RED))
        self.assertEqual(plan["presentation"]["formats"]["money"]["code"], PAREN_RED)
        self.assertEqual(plan["presentation"]["formats"]["lines"]["Services"]["format"]["formatString"], ".1%")
        issue = next(i for i in plan["draftReviewIssues"] if i["code"] == "NUMBER_FORMAT_APPROXIMATED")
        self.assertIn("Services: 0.0%", " ".join(issue["evidence"]))
        pivot = elements_of(self.spec(plan))["pnlPivot"]
        self.assertEqual(next(c for c in pivot["columns"] if c["id"] == "p_amt")["format"]["formatString"], "(,.0f")

    def test_every_numeric_column_takes_its_source_format(self):
        plan, _ = self.draft(mutate=chain(set_formats({"EBITDA": "#,##0.0", "Net Income": "0.0%"}, PAREN_RED),
                                          gl_format("#,##0.00")))
        elements = elements_of(self.spec(plan))
        fmt = lambda element, column: next(c for c in element["columns"] if c["id"] == column).get("format")
        self.assertEqual(fmt(elements["pnlPivot"], "p_amt")["formatString"], "(,.0f")
        self.assertEqual(fmt(elements["kpiRev"], "kpiRev_v")["formatString"], "(,.0f")
        self.assertEqual(fmt(elements["kpiEbitda"], "kpiEbitda_v")["formatString"], ",.1f")
        self.assertEqual(fmt(elements["kpiNI"], "kpiNI_v")["formatString"], ".1%")
        self.assertEqual(fmt(elements["glTable"], "g6")["formatString"], ",.2f")
        self.assertEqual(fmt(elements["glTable"], "g1"), plan["presentation"]["formats"]["glDate"])
        self.assertEqual(fmt(elements["glTable"], "g1")["kind"], "datetime")
        trend = {c["name"]: c.get("format", {}).get("formatString") for c in elements["monthlyTrend"]["columns"]}
        self.assertEqual((trend["Revenue"], trend["EBITDA"], trend["Net Income"]), ("(,.0f", ",.1f", ".1%"))
        self.assertEqual(fmt(elements["pnlCombined"], "c_amt")["formatString"], "(,.0f")

    def test_source_chart_series_take_the_formats_of_their_rows(self):
        plan, _ = self.draft(mutate=chain(set_formats({"Gross Profit": "#,##0.0"}, PAREN_RED),
                                          add_metric_chart(["Total Revenue", "Gross Profit"])))
        chart = elements_of(self.spec(plan))["sourceChart1"]
        formats = {c["name"]: c["format"]["formatString"] for c in chart["columns"] if c["id"].startswith("s1_y")}
        self.assertEqual(formats, {"Revenue": "(,.0f", "Gross Profit": ",.1f"})

    def test_general_cells_fall_back_to_the_cli_defaults_only(self):
        plan, _ = self.draft(currency_format="$,.2f", pct_format=".2%")
        formats = plan["presentation"]["formats"]
        self.assertNotIn("format", formats["money"])               # General: no source format recorded
        elements = elements_of(self.spec(plan))
        for column in (elements["pnlPivot"]["columns"][-1], elements["kpiRev"]["columns"][0],
                       elements["kpiEbitda"]["columns"][0], elements["glTable"]["columns"][-1]):
            self.assertEqual(column["format"], {"kind": "number", "formatString": "$,.2f"})
        self.assertEqual(elements["kpiGM"]["columns"][0]["format"]["formatString"], ".2%")
        rate = next(c for c in elements["assumpTable"]["columns"] if c["id"] == "d2")
        self.assertEqual(rate["format"]["formatString"], ".2%")

    def test_a_formatted_cell_never_falls_back(self):
        plan, _ = self.draft(currency_format="$,.2f", mutate=set_formats({}, "#,##0"))
        elements = elements_of(self.spec(plan))
        self.assertEqual(elements["kpiRev"]["columns"][0]["format"]["formatString"], ",.0f")
        self.assertEqual(elements["pnlPivot"]["columns"][-1]["format"]["formatString"], ",.0f")
        # a General GL amount shows like the P&L cells it feeds
        self.assertEqual(elements["glTable"]["columns"][-1]["format"]["formatString"], ",.0f")

    def test_red_and_parenthesised_negatives_get_the_rule(self):
        plan, _ = self.draft(mutate=chain(set_formats({}, PAREN_RED), gl_format(PAREN_RED)))
        elements = elements_of(self.spec(plan))
        rule = {"type": "single", "columnIds": ["p_amt"], "condition": "<", "value": 0,
                "style": {"color": "#EF4444", "backgroundColor": "none"}}
        pivot_rule, = elements["pnlPivot"]["conditionalFormats"]
        self.assertEqual({k: v for k, v in pivot_rule.items() if k.startswith(("type", "columnIds", "condition",
                                                                                   "value", "style"))}, rule)
        self.assertEqual(pivot_rule["includeSubtotals"], True)
        self.assertEqual(elements["glTable"]["conditionalFormats"], [dict(rule, columnIds=["g6"])])
        self.assertEqual(elements["pnlPivot"]["columns"][-1]["format"]["formatString"], "(,.0f")

    def test_parentheses_without_red_get_no_rule(self):
        plan, _ = self.draft(mutate=set_formats({}, "#,##0;\\(#,##0\\)"))
        elements = elements_of(self.spec(plan))
        self.assertNotIn("conditionalFormats", elements["pnlPivot"])
        self.assertEqual(elements["pnlPivot"]["columns"][-1]["format"]["formatString"], "(,.0f")

    def test_zero_as_a_dash_is_approximated_and_listed(self):
        plan, info = self.draft(mutate=set_formats({}, PAREN_DASH))
        self.assertEqual(plan["presentation"]["formats"]["money"]["format"]["formatString"], "(,.0f")
        issue = next(i for i in plan["draftReviewIssues"] if i["code"] == "NUMBER_FORMAT_APPROXIMATED")
        self.assertIn("zero section", " ".join(issue["evidence"]))
        self.assertIn("NUMBER_FORMAT_APPROXIMATED", info["reviewIssues"])

    def test_the_real_sources_formats(self):
        plan, _ = draft_plan(SAAS, intent="read-only", connection="c", folder="f")
        formats = plan["presentation"]["formats"]
        self.assertEqual((formats["money"]["format"]["formatString"], formats["money"]["negativeRed"]), ("(,.0f", True))
        self.assertEqual(formats["derived"]["Gross Margin %"]["format"]["formatString"], ".1%")
        self.assertEqual(formats["drivers"]["format"]["formatString"], ".1%")
        plan, _ = draft_plan(RETAIL, intent="read-only", connection="c", folder="f")
        self.assertEqual(plan["presentation"]["formats"]["money"]["code"], PAREN_RED)


# ================================================================== 2. title, header band, table headers
class TestTitleAndHeaders(Base):
    def test_the_title_is_the_source_title_cell_not_the_plan_name(self):
        plan, _ = self.draft()
        self.assertEqual(plan["presentation"]["title"], "Synthetic Co - rolling forecast")
        plan["name"] = "TEST renamed (delete me)"
        elements = elements_of(self.spec(plan))
        body = elements["pnlHeader"]["body"]
        self.assertIn("Synthetic Co - rolling forecast", body)
        self.assertIn("pnl.xlsx · P&amp;L sheet", body)
        self.assertNotIn("TEST renamed", body)
        self.assertNotIn("Monthly actuals and forecast performance", body)
        self.assertEqual(ps.disallowed_span_styles(body), [])
        self.assertEqual(elements["pnlHeaderBand"]["kind"], "container")
        self.assertEqual(elements["pnlHeaderBand"]["style"]["backgroundColor"], plan["presentation"]["accent"]["tint"])

    def test_the_real_titles(self):
        plan, _ = draft_plan(SAAS, intent="read-only", connection="c", folder="f")
        self.assertEqual(plan["presentation"]["title"], "FY26 Monthly P&L — Actuals Jan-Jun, Forecast Jul-Dec")
        plan, _ = draft_plan(RETAIL, intent="read-only", connection="c", folder="f")
        self.assertEqual(plan["presentation"]["title"], "Retail Budget — FY26 (Q1-Q2 actual, Q3-Q4 budget)")

    def test_a_plan_without_a_title_cell_uses_its_name(self):
        plan, _ = self.draft()
        plan["presentation"]["title"] = None
        self.assertIn(plan["name"], elements_of(self.spec(plan))["pnlHeader"]["body"])

    def test_source_header_fill_and_bold_style_the_pivot_and_tables(self):
        plan, _ = self.draft(mutate=style_headers("DDE7F5"))
        self.assertEqual(plan["presentation"]["headers"]["pivot"], {"fill": "#DDE7F5", "bold": True})
        elements = elements_of(self.spec(plan))
        header = {"backgroundColor": "#DDE7F5", "fontWeight": "bold"}
        self.assertEqual(elements["pnlPivot"]["tableStyle"],
                         {"preset": "presentation", "textStyles": {"header": header, "columnHeader": header}})
        self.assertEqual(elements["glTable"]["tableStyle"],
                         {"preset": "presentation", "textStyles": {"header": header}})
        self.assertEqual(elements["assumpTable"]["tableStyle"]["textStyles"]["header"], header)

    def test_a_dark_header_fill_keeps_its_light_text(self):
        plan, _ = self.draft(mutate=style_headers("1F3A5F", color="FFFFFF"))
        header = elements_of(self.spec(plan))["pnlPivot"]["tableStyle"]["textStyles"]["header"]
        self.assertEqual(header, {"backgroundColor": "#1F3A5F", "color": "#FFFFFF", "fontWeight": "bold"})

    def test_a_source_without_a_header_style_still_gets_the_presentation_preset(self):
        plan, _ = self.draft()
        self.assertEqual(elements_of(self.spec(plan))["pnlPivot"]["tableStyle"], {"preset": "presentation"})

    def test_the_real_header_fills(self):
        plan, _ = draft_plan(SAAS, intent="read-only", connection="c", folder="f")
        self.assertEqual(plan["presentation"]["headers"]["pivot"]["fill"], "#DDE7F5")
        plan, _ = draft_plan(RETAIL, intent="read-only", connection="c", folder="f")
        self.assertEqual(plan["presentation"]["headers"]["assumptions"]["fill"], "#DDE7F5")   # 03 Drivers

    def test_section_rows_are_bold_in_the_pivot(self):
        plan, _ = self.draft()
        self.assertEqual(elements_of(self.spec(plan))["pnlPivot"]["totals"], {"subTotalFontWeight": "bold"})


# ================================================================== 3. charts
EVERY_SERIES = ["Total Revenue", "EBITDA", "Net Income"]


class TestChartDeduplication(Base):
    def test_a_mapped_source_chart_that_covers_the_trend_replaces_it(self):
        plan, _ = self.draft(mutate=add_metric_chart(EVERY_SERIES, title="Revenue, EBITDA & Net Income"))
        elements = elements_of(self.spec(plan))
        self.assertNotIn("monthlyTrend", elements)
        self.assertEqual(elements["sourceChart1"]["name"], "Revenue, EBITDA & Net Income")
        self.assertEqual([c["name"] for c in elements["sourceChart1"]["columns"]][1:4],
                         ["Revenue", "EBITDA", "Net Income"])      # the source series order
        charts = [e for e in elements.values() if e["kind"] in ("line-chart", "bar-chart")]
        self.assertEqual(len(charts), 1)

    def test_without_a_source_chart_the_generated_trend_stays(self):
        plan, _ = self.draft()
        elements = elements_of(self.spec(plan))
        self.assertIn("monthlyTrend", elements)
        self.assertNotIn("sourceChart1", elements)

    def test_a_source_chart_with_fewer_series_keeps_both_charts(self):
        plan, _ = self.draft(mutate=add_metric_chart(["Total Revenue"], title="Revenue"))
        elements = elements_of(self.spec(plan))
        self.assertIn("monthlyTrend", elements)
        self.assertIn("sourceChart1", elements)

    def test_the_real_sources(self):
        saas, _ = draft_plan(SAAS, intent="read-only", connection="c", folder="f")
        ids = set(elements_of(self.spec(saas)))
        self.assertNotIn("monthlyTrend", ids)          # 02: "Revenue, EBITDA & Net Income" already shows it
        self.assertIn("sourceChart1", ids)
        retail, _ = draft_plan(RETAIL, intent="read-only", connection="c", folder="f")
        ids = set(elements_of(self.spec(retail)))
        self.assertTrue({"monthlyTrend", "sourceChart1"} <= ids)    # 03: the bar chart has one series

    def test_coverage_needs_the_series_the_grain_and_every_period(self):
        periods = [{"label": "Jan-26", "end": "2026-01-31"}, {"label": "Feb-26", "end": "2026-02-28"}]
        chart = {"title": "T", "grain": "month", "series": ["Revenue", "EBITDA"], "periodStart": "2026-01-01",
                 "periodEnd": "2026-02-28"}
        plan = {"presentation": {}, "periods": periods, "sourceCharts": [chart]}
        self.assertEqual(fp.trend_covered(plan, ["Revenue", "EBITDA"], "month"), "T")
        self.assertIsNone(fp.trend_covered(plan, ["Revenue", "EBITDA", "Net Income"], "month"))   # series missing
        self.assertIsNone(fp.trend_covered(plan, ["Revenue"], "quarter"))                           # grain differs
        self.assertIsNone(fp.trend_covered(dict(plan, sourceCharts=[dict(chart, grain="quarter")]),
                                           ["Revenue"], "month"))
        self.assertIsNone(fp.trend_covered(dict(plan, sourceCharts=[dict(chart, periodEnd="2026-01-31")]),
                                           ["Revenue"], "month"))                                   # one month only
        self.assertEqual(fp.trend_covered(dict(plan, sourceCharts=[dict(chart, grain="period", periodEnd=None)]),
                                          ["Revenue"], "quarter"), "T")                             # period plans
        self.assertIsNone(fp.trend_covered({"periods": periods, "sourceCharts": [chart]}, ["Revenue"], "month"),
                          "a legacy plan keeps both charts")


class TestChartStyle(Base):
    def test_category_labels_are_flat_for_twelve_categories_or_fewer(self):
        plan, _ = self.draft(mutate=add_metric_chart(["Total Revenue"], title="Revenue"))
        elements = elements_of(self.spec(plan))
        for key in ("monthlyTrend", "sourceChart1"):
            self.assertEqual(elements[key]["xAxis"]["format"], {"labels": {"labelAngle": 0}}, key)
        axis = ps.flat_labels({"columnId": "x"}, 13)
        self.assertNotIn("format", axis)
        self.assertEqual(ps.flat_labels({"columnId": "x"}, 12)["format"], {"labels": {"labelAngle": 0}})

    def test_charts_use_the_card_chrome_one_series_one_colour_and_the_scheme(self):
        plan, _ = self.draft(mutate=add_metric_chart(["Total Revenue"], title="Revenue"))
        spec = self.spec(plan)
        elements = elements_of(spec)
        accent = plan["presentation"]["accent"]["hex"]
        self.assertEqual(elements["sourceChart1"]["style"], ps.card_style())
        self.assertEqual(elements["sourceChart1"]["color"], {"by": "single", "value": accent})
        self.assertEqual(elements["sourceChart1"]["legend"], {"visibility": "hidden"})
        self.assertEqual(elements["monthlyTrend"]["style"], ps.card_style())
        self.assertNotIn("color", elements["monthlyTrend"])
        self.assertEqual(spec["settings"]["theme"]["overrides"]["categoricalScheme"][0], accent)

    def test_source_series_colours_are_kept(self):
        def colour(pnl, gl, drivers):
            add_metric_chart(["Total Revenue", "EBITDA"], title="Two")(pnl, gl, drivers)
            for series, rgb in zip(pnl._charts[-1].series, ("1F77B4", "D62728")):
                series.graphicalProperties.line.solidFill = rgb
        plan, _ = self.draft(mutate=colour)
        self.assertEqual(plan["sourceCharts"][0]["colors"], ["#1F77B4", "#D62728"])
        scheme = self.spec(plan)["settings"]["theme"]["overrides"]["categoricalScheme"]
        self.assertEqual(scheme[:2], ["#1F77B4", "#D62728"])

    def test_one_chart_is_full_width_and_two_sit_side_by_side(self):
        plan, _ = self.draft()
        page = layout_pages(self.spec(plan)["layout"])["pnl"]
        trend = next(p for p in page if p[0] == "monthlyTrend")
        self.assertEqual(trend[1:3], (1, 25))
        plan, _ = self.draft(mutate=add_metric_chart(["Total Revenue"], title="Revenue"))
        page = layout_pages(self.spec(plan)["layout"])["pnl"]
        a, b = (next(p for p in page if p[0] == key) for key in ("monthlyTrend", "sourceChart1"))
        self.assertEqual((a[1:3], b[1:3], a[3:5] == b[3:5]), ((1, 13), (13, 25), True))


# ================================================================== 4. KPI cards
class TestKpiCards(Base):
    def test_cards_are_equal_peers_with_the_same_formulas(self):
        plan, _ = self.draft(mutate=set_formats({}, PAREN_RED))
        peers = self.spec(plan)
        legacy = BUILD.build_wb_spec({k: v for k, v in plan.items() if k != "presentation"}, "<dm>", ELS)
        cards = [e for e in elements_of(peers).values() if e["kind"] == "kpi-chart"]
        self.assertEqual([c["id"] for c in cards], ["kpiRev", "kpiEbitda", "kpiNI", "kpiGM"])
        shape = lambda card: {k: (v if k != "value" else {x: y for x, y in v.items() if x != "columnId"})
                              for k, v in card.items() if k not in ("id", "name", "columns")}
        for card in cards:
            self.assertEqual(shape(card), shape(cards[0]))
            self.assertEqual(card["name"]["fontSize"], ps.KPI_LABEL_PX)
            self.assertEqual(card["value"]["fontSize"], ps.KPI_VALUE_PX)
            self.assertNotEqual(card["style"].get("borderColor"), "#2563EB")      # no accent on the first
            self.assertEqual(card["name"]["text"], card["name"]["text"].upper())
        old = {e["id"]: e["columns"][0]["formula"] for e in elements_of(legacy).values() if e["kind"] == "kpi-chart"}
        self.assertEqual({c["id"]: c["columns"][0]["formula"] for c in cards}, old)

    def test_cards_span_the_page_in_equal_widths(self):
        plan, _ = self.draft()
        page = layout_pages(self.spec(plan)["layout"])["pnl"]
        cards = [p for p in page if p[0].startswith("kpi")]
        self.assertEqual([c[1:3] for c in cards], [(1, 7), (7, 13), (13, 19), (19, 25)])
        self.assertEqual({c[3:5] for c in cards}, {(ps.HEADER_ROWS + 1, ps.HEADER_ROWS + 1 + ps.KPI_ROWS)})


# ================================================================== 5. page composition
class TestComposition(Base):
    def assert_covers_every_element_once(self, spec):
        placed = layout_pages(spec["layout"])
        for page in spec["pages"]:
            ids = [e["id"] for e in page["elements"]]
            laid = [p[0] for p in placed[page["id"]]] + [c for p in placed[page["id"]] for c in p[5]]
            self.assertEqual(sorted(laid), sorted(ids), page["id"])
            top = placed[page["id"]]
            for i, left in enumerate(top):
                for right in top[i + 1:]:
                    overlap = left[1] < right[2] and right[1] < left[2] and left[3] < right[4] and right[3] < left[4]
                    self.assertFalse(overlap, f"{page['id']}: {left[0]} overlaps {right[0]}")
        self.assertEqual([p["id"] for p in spec["pages"]], list(placed))

    def test_layout_covers_every_element_once_without_overlaps(self):
        for intent in ("read-only", "editable"):
            with self.subTest(intent=intent):
                plan, _ = self.draft(intent=intent, mutate=add_metric_chart(["Total Revenue"], title="Revenue"))
                spec = self.spec(plan)
                self.assert_covers_every_element_once(spec)
                for page in spec["pages"]:
                    self.assertTrue(page["elements"], f"{page['id']} would be empty")

    def test_the_real_sources_lay_out_cleanly(self):
        for source in (SAAS, RETAIL):
            for intent in ("read-only", "what-if"):
                with self.subTest(source=source.name, intent=intent):
                    plan, _ = draft_plan(source, intent=intent, connection="c", folder="f")
                    self.assert_covers_every_element_once(self.spec(plan))

    def test_page_order_is_header_kpis_charts_divider_detail_pivot(self):
        plan, _ = self.draft(mutate=add_metric_chart(["Total Revenue"], title="Revenue"))
        page = layout_pages(self.spec(plan)["layout"])["pnl"]
        order = [p[0] for p in sorted(page, key=lambda p: (p[3], p[1]))]
        self.assertEqual(order, ["pnlHeaderBand", "kpiRev", "kpiEbitda", "kpiNI", "kpiGM", "monthlyTrend",
                                 "sourceChart1", "pnlDivider", "pnlDetail", "pnlPivot"])
        divider = next(p for p in page if p[0] == "pnlDivider")
        pivot = next(p for p in page if p[0] == "pnlPivot")
        self.assertEqual(pivot[3], divider[4] + ps.LABEL_ROWS)         # nothing between the label and the pivot

    def test_the_pivot_is_sized_to_its_rows_with_the_visible_row_cap(self):
        plan, _ = self.draft()
        pivot = next(p for p in layout_pages(self.spec(plan)["layout"])["pnl"] if p[0] == "pnlPivot")
        rows = len(plan["lines"]) + len(plan["sections"])
        self.assertEqual(pivot[4] - pivot[3], ps.table_height(rows, extra_header_rows=fp.PIVOT_HEADER_ROWS, roomy=True))
        many = ps.table_height(200, extra_header_rows=fp.PIVOT_HEADER_ROWS, roomy=True)
        self.assertEqual(many, ps.table_height(ps.MAX_VISIBLE_ROWS, extra_header_rows=fp.PIVOT_HEADER_ROWS, roomy=True))
        self.assertGreater(ps.table_height(rows, extra_header_rows=3, roomy=True), ps.table_height(rows, 3))
        # live measurement (CSS px): body rows 42, pivot header rows ~25 (see TestRowHeightMaths)
        pixels = lambda n, header: (ps.table_height(n, extra_header_rows=header, roomy=True) - 1) * ps.GRID_ROW_PX
        self.assertGreaterEqual(pixels(14, 3), ps.CHROME_PX + 14 * 42 + 3 * 25)
        self.assertLess(pixels(14, 3), ps.CHROME_PX + 14 * 42 + 3 * 25 + 2 * ps.GRID_ROW_PX)   # at most a row of slack
        self.assertEqual(ps.table_height(14, extra_header_rows=3, roomy=False), 22)    # the spreadsheet preset is unchanged
        gl = next(p for p in layout_pages(self.spec(plan)["layout"])["detail"] if p[0] == "glTable")
        self.assertEqual(gl[4] - gl[3], ps.table_height(len(plan["actuals"]), roomy=True))

    def test_tables_use_the_presentation_preset_and_the_card_chrome(self):
        plan, _ = self.draft()
        elements = elements_of(self.spec(plan))
        for key in ("pnlPivot", "glTable", "assumpTable"):
            self.assertEqual(elements[key]["tableStyle"]["preset"], "presentation", key)
            self.assertEqual(elements[key]["style"], ps.card_style(), key)

    def test_the_inputs_page_shows_the_drivers_and_the_editable_inputs(self):
        plan, _ = self.draft(intent="read-only")
        spec = self.spec(plan)
        self.assertEqual([e["id"] for e in next(p for p in spec["pages"] if p["id"] == "inputs")["elements"]],
                         ["assumpTable"])
        plan, _ = self.draft(intent="what-if")
        spec = self.spec(plan)
        self.assertEqual([e["id"] for e in next(p for p in spec["pages"] if p["id"] == "inputs")["elements"]],
                         ["assumpTable", "rateInput", "manualInput"])

    def test_an_inputs_page_with_nothing_to_show_is_omitted(self):
        plan, _ = self.draft(intent="read-only")
        plan["assumptions"] = []
        spec = self.spec(plan)
        self.assertEqual([p["id"] for p in spec["pages"]], ["pnl", "detail", "model"])
        self.assertNotIn("assumpTable", elements_of(spec))
        self.assertNotIn("inputs", layout_pages(spec["layout"]))
        self.assert_covers_every_element_once(spec)

    def test_the_hidden_model_and_detail_pages_keep_their_roles(self):
        plan, _ = self.draft()
        spec = self.spec(plan)
        pages = {p["id"]: p for p in spec["pages"]}
        self.assertEqual(pages["model"]["visibility"], "hidden")
        self.assertEqual([e["id"] for e in pages["detail"]["elements"]], ["glTable"])
        self.assertEqual([e["id"] for e in pages["model"]["elements"]],
                         ["actuals", "spineWrap", "forecast", "pnlCombined"])

    def test_text_spans_only_use_allowed_properties(self):
        plan, _ = self.draft()
        for element in elements_of(self.spec(plan)).values():
            if element["kind"] == "text":
                self.assertEqual(ps.disallowed_span_styles(element["body"]), [], element["id"])

    def test_the_spec_wires_for_posting(self):
        plan, _ = self.draft(mutate=add_metric_chart(["Total Revenue"], title="Revenue"))
        wired = workbook_wire.wire_workbook(self.spec(plan))
        self.assertIn("settings", wired["contents"])
        ids = [e["id"] for e in wired["contents"]["elements"]]
        self.assertEqual(len(ids), len(set(ids)))


# ================================================================== 6. --look
def walk(value):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from walk(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk(child)


class TestLook(Base):
    def test_a_look_is_recorded_in_the_plan(self):
        plan, _ = self.draft(look=LOOK)
        self.assertEqual(plan["presentation"]["look"], LOOK)

    def test_a_look_that_is_not_version_one_is_refused(self):
        fixture.build_pnl(self.path)
        with self.assertRaisesRegex(PlanDraftError, "version 1"):
            draft_plan(self.path, intent="what-if", connection="c", folder="f", look={"version": 2})
        plan, _ = self.draft()
        plan["presentation"]["look"] = {"version": 2}
        with self.assertRaisesRegex(BUILD.PlanError, "version 1"):
            self.spec(plan)

    def test_the_look_styles_the_page_and_ours_does_not_leak(self):
        plan, _ = self.draft(look=LOOK, mutate=style_headers("DDE7F5"))
        spec = self.spec(plan)
        elements = elements_of(spec)
        self.assertEqual(spec["settings"], {"theme": LOOK["theme"]})
        self.assertEqual(elements["pnlHeaderBand"]["style"], {"borderRadius": "pill"})
        self.assertIn("font-size: 30px", elements["pnlHeader"]["body"])
        self.assertIn("font-size: 13px", elements["pnlHeader"]["body"])
        for key in ("kpiRev", "kpiEbitda", "kpiNI", "kpiGM"):
            card = elements[key]
            self.assertEqual(card["style"], {"borderRadius": "pill"})
            self.assertEqual(card["value"]["fontSize"], 30)
            self.assertEqual(card["name"], {"text": card["name"]["text"], "color": "#121b39", "fontSize": 13,
                                            "fontWeight": "bold"})
        for key in ("pnlPivot", "glTable", "assumpTable"):
            self.assertEqual(elements[key]["style"], {"borderRadius": "pill"}, key)
            self.assertEqual(elements[key]["tableStyle"]["preset"], "presentation", key)
            # the theme styles the headers: only the source's bold survives
            self.assertEqual(elements[key]["tableStyle"]["textStyles"]["header"], {"fontWeight": "bold"}, key)
        self.assertEqual(elements["monthlyTrend"]["style"], {"padding": "none"})
        ours = {v.lower() for v in list(ps.ROLES.values()) + [h for hue in ps.HUES.values() for h in hue]
                + list(ps.CATEGORICAL_MID)}
        text = json.dumps(spec).lower()
        for value in ours:
            self.assertNotIn(f'"{value}"', text.replace("'", '"'), f"{value} leaked under a look")
        self.assertNotIn("color: #0f172a", text)
        self.assertEqual(ps.disallowed_span_styles(elements["pnlHeader"]["body"]), [])

    def test_excel_formats_titles_and_order_still_win_under_a_look(self):
        plan, _ = self.draft(look=LOOK, mutate=set_formats({}, PAREN_RED))
        elements = elements_of(self.spec(plan))
        self.assertEqual(elements["pnlPivot"]["columns"][-1]["format"]["formatString"], "(,.0f")
        self.assertIn("Synthetic Co - rolling forecast", elements["pnlHeader"]["body"])
        self.assertEqual(elements["pnlPivot"]["rowsBy"][0]["sort"]["by"], "p_so")

    def test_a_look_never_changes_a_value(self):
        plain, _ = self.draft()
        styled, _ = self.draft(look=LOOK)
        self.assertEqual(value_snapshot(self, plain), value_snapshot(self, styled))

    def test_the_cli_takes_look_json(self):
        fixture.build_pnl(self.path)
        look_path = self.dir / "look.json"
        look_path.write_text(json.dumps(LOOK))
        out = self.dir / "plan.json"
        stdout = io.StringIO()
        with contextlib.redirect_stdout(stdout):
            code = DRAFT.main([str(self.path), "--connection", "c", "--folder", "f", "--look", str(look_path),
                               "--out", str(out)])
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(out.read_text())["presentation"]["look"], LOOK)
        self.assertIn("look", stdout.getvalue())
        built = self.dir / "built"
        with contextlib.redirect_stdout(io.StringIO()):
            code = BUILD.main([str(out), "--source", str(self.path), "--out-dir", str(built)])
        self.assertEqual(code, 0)
        wb = json.loads((built / "wb_spec.json").read_text())
        self.assertEqual(wb["contents"]["settings"], {"theme": LOOK["theme"]})


# ================================================================== 7. values are untouched
def value_snapshot(test, plan):
    """The DM spec, parity and every emitted column formula of a plan's build, ignoring all presentation."""
    spec = BUILD.build_wb_spec(plan, "<dm>", ELS)
    elements = {}
    for page in spec["pages"]:
        for element in page["elements"]:
            elements[element["id"]] = (element["kind"], [c.get("formula") for c in element.get("columns", [])])
    with contextlib.redirect_stdout(io.StringIO()):
        parity = BUILD.parity(plan)
    return (hashlib.sha256(json.dumps(BUILD.build_dm_spec(plan), sort_keys=True).encode()).hexdigest(),
            plan.get("parityTargets"), parity, elements)


class TestValuesAreUntouched(Base):
    def assert_same_values(self, plan, dropped=()):
        legacy = {k: v for k, v in plan.items() if k != "presentation"}
        new, old = value_snapshot(self, plan), value_snapshot(self, legacy)
        self.assertEqual(new[:3], old[:3])
        self.assertEqual(set(old[3]) - set(new[3]), set(dropped))
        for element_id, built in new[3].items():
            if element_id in old[3]:
                self.assertEqual(built, old[3][element_id], element_id)

    def test_synthetic_plans(self):
        for name, kwargs in (("default", {}),
                             ("formatted", {"mutate": chain(set_formats({"EBITDA": PAREN_DASH}, PAREN_RED),
                                                            style_headers("DDE7F5"))}),
                             ("chart", {"mutate": add_metric_chart(EVERY_SERIES, title="All three")})):
            with self.subTest(name):
                plan, _ = self.draft(**kwargs)
                self.assert_same_values(plan, dropped=("monthlyTrend",) if name == "chart" else ())

    def test_a_quarterly_plan(self):
        fixture.build_quarterly(self.path)
        plan, _ = draft_plan(self.path, intent="what-if", connection="c", folder="f")
        self.assertIn("presentation", plan)
        self.assert_same_values(plan)

    def test_02_and_03_equal_what_main_emitted(self):
        for key, source, dropped in (("02-saas", SAAS, {"monthlyTrend"}), ("03-retail", RETAIL, set())):
            with self.subTest(key):
                plan, _ = draft_plan(source, intent="what-if", connection="c", folder="f")
                digest, targets, parity, elements = value_snapshot(self, plan)
                golden = GOLDEN[key]
                self.assertEqual(digest, golden["dmSha256"])
                self.assertEqual(plan["parityTargets"], golden["parityTargets"])
                self.assertEqual(parity, golden["parity"])
                self.assertEqual(set(golden["elements"]) - set(elements), dropped)
                for element_id, (kind, formulas) in elements.items():
                    if element_id in golden["elements"]:
                        self.assertEqual({"kind": kind, "formulas": formulas}, golden["elements"][element_id],
                                         f"{key} {element_id}")

    def test_plans_drafted_before_the_port_emit_main_s_workbook_unchanged(self):
        from test_builders import forecast_plan
        sample = json.loads((SCRIPTS / "sample-model-plan.json").read_text())
        for key, plan in (("sample-model-plan", sample), ("builders-plan", forecast_plan())):
            with self.subTest(key):
                self.assertNotIn("presentation", plan)
                spec = BUILD.build_wb_spec(plan, "<dm>", ELS)
                digest = hashlib.sha256(json.dumps(spec, sort_keys=True).encode()).hexdigest()
                self.assertEqual(digest, GOLDEN["legacy"][key]["wbSpecSha256"])
                self.assertEqual(value_snapshot(self, plan)[3],
                                 {i: (e["kind"], e["formulas"]) for i, e in GOLDEN["legacy"][key]["snapshot"]["elements"].items()})

    def test_the_drafter_leaves_every_other_plan_field_alone(self):
        for source in (SAAS, RETAIL):
            plan, _ = draft_plan(source, intent="what-if", connection="c", folder="f")
            self.assertEqual(plan["parityTargets"], GOLDEN["02-saas" if source == SAAS else "03-retail"]["parityTargets"])


# ================================================================== 8. row-height pixel maths
class TestRowHeightMaths(Base):
    """Live measurement of the presentation preset at pixelWidth 1400 (screenshots come back 1600 px wide, so
    divide by 1.143): body rows 48 px in the image = 42 CSS px, pivot header rows ~25 CSS px, grid row 23 px."""
    BODY_PX, HEADER_PX = 42, 25

    def test_the_constants_pin_the_measurement(self):
        self.assertEqual(ps.ROOMY_FACTOR, 2.15)               # 42 px / ROW_PX 20 = 2.1, rounded up
        self.assertEqual(ps.HEADER_ROOMY_FACTOR, 1.3)         # 25 px / 20 = 1.25, rounded up
        self.assertGreaterEqual(ps.ROW_PX * ps.ROOMY_FACTOR, self.BODY_PX)
        self.assertGreaterEqual(ps.ROW_PX * ps.HEADER_ROOMY_FACTOR, self.HEADER_PX)
        self.assertEqual(ps.SPACIOUS_FACTOR, 2.5)             # medium / large cell spacing, unmeasured

    def test_the_presentation_preset_is_tighter_than_spacious_themes(self):
        presentation = {"tablePreset": "presentation"}
        self.assertEqual(ps.roomy_tables(None), ps.ROOMY_FACTOR)
        self.assertEqual(ps.roomy_tables(presentation), ps.ROOMY_FACTOR)
        for spacing in ("medium", "large"):
            theme = {"overrides": {"tableStyles": {"cellSpacing": spacing}}}
            self.assertEqual(ps.roomy_tables({**presentation, "theme": theme}), ps.SPACIOUS_FACTOR, spacing)
            self.assertEqual(ps.roomy_tables({"theme": theme}), ps.SPACIOUS_FACTOR, spacing)
        small = {"overrides": {"tableStyles": {"cellSpacing": "small"}}}
        self.assertEqual(ps.roomy_tables({**presentation, "theme": small}), ps.ROOMY_FACTOR)
        self.assertEqual(ps.roomy_tables({"theme": {"name": "org-theme-id"}}), ps.SPACIOUS_FACTOR)   # unknown
        self.assertFalse(ps.roomy_tables({"tablePreset": None, "theme": small}))

    def test_every_visible_row_fits_by_the_measured_pixels(self):
        for factor, body, header in ((ps.ROOMY_FACTOR, self.BODY_PX, self.HEADER_PX),
                                     (ps.SPACIOUS_FACTOR, ps.ROW_PX * ps.SPACIOUS_FACTOR,
                                      ps.ROW_PX * ps.HEADER_ROOMY_FACTOR)):
            for rows in range(1, 16):
                for extra in (0, 1, 3):
                    height = ps.table_height(rows, extra_header_rows=extra, roomy=factor)
                    self.assertGreaterEqual((height - 1) * ps.GRID_ROW_PX,
                                            ps.CHROME_PX + rows * body + extra * header, (factor, rows, extra))
        self.assertEqual(ps.table_height(5, roomy=True), ps.table_height(5, roomy=ps.ROOMY_FACTOR))

    def test_no_table_or_pivot_of_any_fixture_is_cut_off(self):
        from test_table_report_presentation import SOURCES, draft, wb_spec
        calls = []
        real = ps.table_height

        def record(rows, extra_header_rows=0, roomy=False):
            height = real(rows, extra_header_rows, roomy)
            calls.append((rows, extra_header_rows, roomy, height))
            return height
        with mock.patch.object(ps, "table_height", record):
            for source in SOURCES.values():
                wb_spec(draft(source)[0])
            for source in (SAAS, RETAIL):
                plan, _ = draft_plan(source, intent="what-if", connection="c", folder="f")
                self.spec(plan)
        self.assertGreater(len(calls), 10)
        for rows, extra, roomy, height in calls:
            body = self.BODY_PX if roomy == ps.ROOMY_FACTOR else ps.ROW_PX * (roomy or 1)
            head = self.HEADER_PX if roomy == ps.ROOMY_FACTOR else ps.ROW_PX * (ps.HEADER_ROOMY_FACTOR if roomy else 1)
            visible = min(max(rows, 1), ps.MAX_VISIBLE_ROWS)
            self.assertGreaterEqual((height - 1) * ps.GRID_ROW_PX, ps.CHROME_PX + visible * body + extra * head,
                                    (rows, extra, roomy))


if __name__ == "__main__":
    unittest.main()
