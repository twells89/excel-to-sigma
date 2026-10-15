#!/usr/bin/env python3
from __future__ import annotations

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts" / "lib"))

from xlsx_audit import audit_workbook  # noqa: E402


class TestXlsxAudit(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def save(self, workbook, name="book.xlsx"):
        path = Path(self.tmp.name) / name
        workbook.save(path)
        return path

    def test_plain_workbook_is_automatic(self):
        wb = Workbook()
        wb.active["A1"] = "value"
        report = audit_workbook(self.save(wb))
        self.assertEqual(report["status"], "AUTO")
        self.assertEqual(report["issues"], [])

    def test_unknown_and_volatile_formulas_require_review(self):
        wb = Workbook()
        wb.active["A1"] = "=OFFSET(B1,0,0)"
        wb.active["A2"] = "=MYSTERYFUNC(1)"
        report = audit_workbook(self.save(wb))
        codes = {item["code"] for item in report["issues"]}
        self.assertEqual(report["status"], "REVIEW_REQUIRED")
        self.assertIn("VOLATILE_FORMULAS", codes)
        self.assertIn("UNKNOWN_FORMULAS", codes)

    def test_chart_validation_and_protection_are_routed(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Category", "Value"])
        ws.append(["A", 1])
        ws.append(["B", 2])
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        ws.add_chart(chart, "D2")
        validation = DataValidation(type="list", formula1='"A,B"')
        ws.add_data_validation(validation)
        validation.add(ws["A2"])
        ws.protection.sheet = True

        report = audit_workbook(self.save(wb))
        codes = {item["code"] for item in report["issues"]}
        self.assertTrue({
            "CHARTS_PRESENT", "DATA_VALIDATION_PRESENT", "WORKBOOK_PROTECTION"
        }.issubset(codes))

    def test_connection_and_custom_xml_are_not_silent(self):
        wb = Workbook()
        path = self.save(wb)
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr("xl/connections.xml", "<connections/>")
            archive.writestr("customXml/item1.xml", "<business-metadata/>")
        report = audit_workbook(path)
        codes = {item["code"] for item in report["issues"]}
        self.assertIn("POWER_QUERY_OR_CONNECTIONS", codes)
        self.assertIn("CUSTOM_XML", codes)

    def test_legacy_binary_is_blocked_before_parsing(self):
        path = Path(self.tmp.name) / "book.xls"
        path.write_bytes(b"not an OOXML package")
        report = audit_workbook(path)
        self.assertEqual(report["status"], "BLOCKED")
        self.assertEqual(report["issues"][0]["code"], "LEGACY_OR_UNSUPPORTED_FORMAT")


if __name__ == "__main__":
    unittest.main()
