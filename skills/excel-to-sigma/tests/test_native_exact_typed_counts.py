"""Typed range counts executed through the complete emitted native graph."""
import copy
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts/lib'))
from native_exact_plan import build_exact_native_calculation_plan, lower_exact_native_calculations
from source_inputs import build_source_input_plan
from native_graph_reference import NativeGraphReference
import test_native_calculations


class TypedCalculatedColumnCountTests(unittest.TestCase):
    source = test_native_calculations.NativeCalculationTests.source

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def lower_counts(self, formula, *, blank_amount=False, result_column='Result', raw_column_name=None):
        path, old = self.source()
        book = load_workbook(path)
        for row in (2, 3):
            book['Ledger Renamed'][f'D{row}'] = formula.format(row=row)
        if blank_amount:
            book['Ledger Renamed']['B2'] = None
        if raw_column_name:
            book['Ledger Renamed']['B1'] = raw_column_name
            book['Ledger Renamed'].tables['tblLedger'].tableColumns[1].name = raw_column_name
        for cell, function in [('A1', 'COUNT'), ('A2', 'COUNTA')]:
            book['Summary Renamed'][cell] = f'={function}(tblLedger[{result_column}])'
        book.save(path)
        decisions = copy.deepcopy(old['decisions'])
        if raw_column_name:
            feature = decisions['features'][0]
            feature['literalColumns'] = [raw_column_name if c == 'Amount' else c
                                         for c in feature['literalColumns']]
            feature['columnTypes'][raw_column_name] = feature['columnTypes'].pop('Amount')
        decisions['sourceSha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
        if blank_amount:
            decisions['features'][0]['reviewedBlankColumns'] = [raw_column_name or 'Amount']
        inputs = build_source_input_plan(path, decisions, connection_id='connection', folder_id='folder')
        plan = build_exact_native_calculation_plan(path, inputs)
        rows = [{**row, 'Row ID': f'physical-{i}'} for i, row in
                enumerate(inputs['sharedInputArchitecture']['inputs']['ledger']['rows'])]
        assigned = {'ledger': {'elementId': 'observed-input', 'name': 'tblLedger',
                              'rowIdColumn': 'Row ID',
                              'rows': {str(i): row['Row ID'] for i, row in enumerate(rows)}}}
        lowered = lower_exact_native_calculations(plan, assigned, max_columns=8)
        self.assertEqual(lowered['rangeAggregateCount'], 2)
        self.assertEqual(lowered['formalTemplateObligations'], [])
        for shard in lowered['shards']:
            if 'rangeDomain' in shard:
                self.assertEqual(shard['rangeDomain']['appendPolicy'], 'live-formal-table')
        return plan, lowered, rows

    def assert_counts(self, lowered, rows, expected):
        graph = NativeGraphReference({'observed-input': ('tblLedger', rows)}).materialize(lowered['shards'])
        for cell, number in zip(('A1', 'A2'), expected):
            result = graph.output(lowered, 'Summary Renamed!' + cell)
            self.assertEqual(result['error'], '')
            self.assertEqual(result['kind'], 'n')
            self.assertEqual(result['number'], number)

    def test_calculated_counts_ignore_text_empty_text_error_and_logical_for_count(self):
        # Each mixed column has one numeric result. COUNTA also counts its
        # nonnumeric result, including empty text and Excel errors.
        for branch in ('"text"', '""', '1/0', 'TRUE'):
            with self.subTest(branch=branch):
                plan, lowered, rows = self.lower_counts(f'=IF(B{{row}}>3,{branch},B{{row}})')
                self.assertEqual(plan['formulas']['Summary Renamed!A1']['localExpected'], 1)
                self.assertEqual(plan['formulas']['Summary Renamed!A2']['localExpected'], 2)
                self.assert_counts(lowered, rows, (1, 2))
                # Reuse the same lowering after an existing row changes type.
                changed = copy.deepcopy(rows)
                changed[0]['Amount'] = 9
                self.assert_counts(lowered, changed, (0, 2))
                # New physical rows participate without rebuilding the plan.
                appended = rows + [{'Account': 'c', 'Amount': 8, 'Rate': 1, 'Row ID': 'new-row'}]
                self.assert_counts(lowered, appended, (1, 3))
                appended.append({'Account': 'd', 'Amount': 2, 'Rate': 1, 'Row ID': 'another-row'})
                self.assert_counts(lowered, appended, (2, 4))

    def test_formula_blank_result_is_numeric_zero_and_counts_after_append(self):
        plan, lowered, rows = self.lower_counts('=B{row}', blank_amount=True)
        self.assertEqual(plan['formulas']['Ledger Renamed!D2']['localExpected'], 0)
        self.assert_counts(lowered, rows, (2, 2))
        rows.append({'Account': 'c', 'Amount': None, 'Rate': 1, 'Row ID': 'new-row'})
        self.assert_counts(lowered, rows, (3, 3))

    def test_raw_input_blanks_remain_excluded_after_append(self):
        _, lowered, rows = self.lower_counts('=B{row}', blank_amount=True, result_column='Amount')
        self.assert_counts(lowered, rows, (1, 1))
        rows.append({'Account': 'c', 'Amount': None, 'Rate': 1, 'Row ID': 'new-row'})
        self.assert_counts(lowered, rows, (1, 1))

    def test_raw_counts_with_helper_named_source_columns_do_not_clobber_helpers(self):
        for column in ('native_key', 'native_match', 'native_error', 'native_order',
                       'native_value', 'native_count_value', 'native_count_value_1'):
            with self.subTest(column=column):
                plan, lowered, rows = self.lower_counts('=B{row}', result_column=column,
                                                       raw_column_name=column)
                self.assertEqual(plan['formulas']['Summary Renamed!A1']['localExpected'], 2)
                self.assertEqual(plan['formulas']['Summary Renamed!A2']['localExpected'], 2)
                self.assert_counts(lowered, rows, (2, 2))
                for shard in lowered['shards']:
                    if 'rangeDomain' not in shard:
                        continue
                    ids = [c['id'] for c in shard['columns']]
                    names = [c['name'] for c in shard['columns']]
                    self.assertEqual(len(ids), len(set(ids)))
                    self.assertEqual(len(names), len(set(names)))
                    copied = [c for c in shard['columns']
                              if c['formula'] == f'[tblLedger/{column}]']
                    self.assertEqual(len(copied), 1)
                    self.assertNotIn(copied[0]['id'], rows[0])
                    self.assertNotIn(copied[0]['name'], rows[0])
                changed = copy.deepcopy(rows)
                changed[0][column] = None
                self.assert_counts(lowered, changed, (1, 1))
                changed.append({'Account': 'c', column: 7, 'Rate': 1, 'Row ID': 'new-row'})
                self.assert_counts(lowered, changed, (2, 2))


if __name__ == '__main__':
    unittest.main()
