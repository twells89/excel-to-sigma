"""Generic source presentation behavior over the eight independent sources."""
from __future__ import annotations

import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, ScatterChart, Reference, Series

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts" / "lib"))

import presentation  # noqa: E402
from input_orchestration import gate_input_staging  # noqa: E402
from migration_gate import gate_plan_staging_source  # noqa: E402
from migration_manifest import build_manifest  # noqa: E402
from source_presentation import build_source_plan, emit_approved, emit_preview, audit_donut_openapi, _digest  # noqa: E402

SOURCES = Path(__file__).resolve().parents[4] / "broader-eight-source-replay-20261020" / "source-originals"


def synthetic_bindings(plan):
    """Explicit, source-derived *logical* facts; never a live runtime receipt."""
    charts = {}
    for chart in plan["charts"]:
        if chart["chartType"] not in {"bar", "line", "doughnut"}:
            continue
        charts[chart["id"]] = {
            "sourceElementId": f"live-{chart['id']}",
            "categorySourceCells": list(chart["categoryCells"]),
            "categoryFormula": "[Live/Category]",
            "orderFormula": "Min([Live/Source Order])",
            "series": [{"sourceCells": list(series["sourceCells"]),
                        "formula": f"Sum([Live/Metric {index}])"}
                       for index, series in enumerate(chart["series"], 1)],
            "points": [
                {"sourceCategoryCell": chart["categoryCells"][point],
                 "sourceValueCells": [series["sourceCells"][point] for series in chart["series"]],
                 "targetCategoryKey": f"key-{point + 1}",
                 "targetOrder": point + 1,
                 "targetCategoryExpression": "[Live/Category]",
                 "targetValueExpressions": [f"Sum([Live/Metric {index}])"
                                            for index in range(1, len(chart["series"]) + 1)]}
                for point in range(chart["pointCount"])
            ],
        }
    return {"sourceHash": plan["sourceHash"], "presentationPlanHash": plan["planHash"],
            "charts": charts, "documentId": "logical", "documentVersion": 1,
            "viewSourceElementId": "view-source", "mainElementId": "main",
            "detailElementId": "detail", "editElementId": "edit",
            "mainViewColumns": [{"id": "category", "name": "Category", "formula": "[Live/Category]"}],
            "detailViewColumns": [{"id": "category", "name": "Category", "formula": "[Live/Category]"}]}


class SourcePresentationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.paths = sorted(path for path in SOURCES.iterdir() if path.suffix in {".xlsx", ".xlsm"})
        cls.plans = [build_source_plan(path) for path in cls.paths]

    def test_eight_sources_inventory_every_chart_and_point_including_doughnut(self):
        self.assertEqual(len(self.plans), 8)
        self.assertEqual([len(plan["charts"]) for plan in self.plans], [2, 1, 1, 2, 1, 1, 1, 1])
        self.assertEqual(sum(chart["pointCount"] * len(chart["series"])
                             for plan in self.plans for chart in plan["charts"]), 253)
        doughnut = self.plans[3]["charts"][1]
        self.assertEqual(doughnut["chartType"], "doughnut")
        self.assertEqual(doughnut["pointCount"], 6)
        self.assertEqual(doughnut["sourceGeometry"], {"holePercent": 10, "firstSliceAngle": 0})
        self.assertEqual(doughnut["series"][0]["sourceCells"][0], "Forecast!C12")
        self.assertEqual(doughnut["categoryCells"][0], "Forecast!A12")

    def test_combination_chart_keeps_both_plots_and_never_emits_partial_chart(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        for row in (("Category", "Bars", "Line"), ("A", 10, 7), ("B", 20, 9)):
            sheet.append(row)
        categories = Reference(sheet, min_col=1, min_row=2, max_row=3)
        bars = BarChart()
        bars.title = "Combination"
        bars.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        bars.set_categories(categories)
        line = LineChart()
        line.add_data(Reference(sheet, min_col=3, min_row=1, max_row=3), titles_from_data=True)
        line.set_categories(categories)
        bars += line
        sheet.add_chart(bars, "E2")
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "combo.xlsx"
            workbook.save(source)
            plan = build_source_plan(source)
            chart = plan["charts"][0]
            self.assertEqual(chart["chartType"], "combo")
            self.assertEqual(chart["status"], "unresolved")
            self.assertEqual([plot["chartType"] for plot in chart["plots"]], ["bar", "line"])
            self.assertEqual([series["name"] for series in chart["series"]], ["Bars", "Line"])
            self.assertEqual([series["sourceCells"] for series in chart["series"]],
                             [["Data!B2", "Data!B3"], ["Data!C2", "Data!C3"]])
            artifact = emit_preview(plan, synthetic_bindings(plan))
            self.assertFalse(any(item["kind"].endswith("-chart") for item in artifact["elements"]))
            self.assertTrue(any(item["kind"] == "chart" and item["sourceId"] == chart["id"]
                                for item in artifact["unresolved"]))

    def test_unsupported_scatter_keeps_x_and_y_ranges_cells_and_points(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        for row in (("X", "Y"), (1, 11), (2, 22)):
            sheet.append(row)
        chart = ScatterChart()
        chart.title = "XY"
        chart.series.append(Series(Reference(sheet, min_col=2, min_row=2, max_row=3),
                                   Reference(sheet, min_col=1, min_row=2, max_row=3),
                                   title="Points"))
        sheet.add_chart(chart, "D2")
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "scatter.xlsx"
            workbook.save(source)
            extracted = build_source_plan(source)["charts"][0]
            self.assertEqual((extracted["chartType"], extracted["status"]),
                             ("scatter", "unresolved"))
            self.assertEqual(extracted["pointCount"], 2)
            self.assertEqual(extracted["categoryCells"], ["Data!A2", "Data!A3"])
            self.assertEqual(extracted["series"][0]["sourceCells"], ["Data!B2", "Data!B3"])
            self.assertEqual(extracted["series"][0]["sourceReferences"],
                             {"xVal": "'Data'!$A$2:$A$3", "yVal": "'Data'!$B$2:$B$3"})

    def test_present_malformed_rule_target_is_rejected_but_absent_is_unresolved(self):
        plan = next(plan for plan in self.plans if plan["validations"])
        rule = plan["validations"][0]
        bindings = synthetic_bindings(plan)
        artifact = emit_preview(plan, bindings)
        self.assertTrue(any(item["kind"] == "validation" and item["sourceId"] == rule["id"]
                            for item in artifact["unresolved"]))
        bindings["validationTargets"] = {rule["id"]: 42}
        with self.assertRaisesRegex(ValueError, rule["id"]):
            emit_preview(plan, bindings)
        bindings["validationTargets"] = 42
        with self.assertRaisesRegex(ValueError, "validationTargets"):
            emit_preview(plan, bindings)

    def test_source_range_bindings_emit_supported_charts_without_filename_routes(self):
        artifacts = [emit_preview(plan, synthetic_bindings(plan)) for plan in self.plans]
        emitted = [element for artifact in artifacts for element in artifact["elements"]
                   if element["kind"] in {"bar-chart", "line-chart", "donut-chart"}]
        self.assertEqual(len(emitted), 10)
        self.assertEqual(sum(len(item["aggregateChecks"]) for item in artifacts), 253)
        self.assertTrue(all(check["sourceCachedExpectationAuthority"] == "DIAGNOSTIC_SOURCE_CACHE"
                            and check["referenceQualificationStatus"] == "PENDING_INDEPENDENT_ORACLE"
                            and check["status"] == "PENDING_RUNTIME_QUERY"
                            for artifact in artifacts for check in artifact["aggregateChecks"]))
        self.assertTrue(all(item["stagingGate"]["status"] == "LOGICAL_FIXTURE_ONLY"
                            for item in artifacts))
        donut = next(item for item in artifacts[3]["elements"] if item["kind"] == "donut-chart")
        self.assertEqual(donut["color"]["columnId"], donut["columns"][0]["id"])
        self.assertEqual(donut["value"]["columnId"], donut["columns"][1]["id"])
        self.assertEqual(donut["innerRadius"], 0.4)
        self.assertEqual(donut["color"]["sort"], {
            "direction": "ascending", "by": donut["columns"][-1]["id"], "aggregation": "min"})
        self.assertNotIn("xAxis", donut)
        self.assertNotIn("yAxis", donut)
        self.assertTrue(any(item["kind"] == "chart-geometry" and item["sourceId"] == "chart-2"
                            for item in artifacts[3]["unresolved"]))
        retail = artifacts[2]
        self.assertIn("Quarterly", retail["pages"][0]["name"])
        self.assertEqual(retail["elements"][3]["name"], "Net Sales by Quarter")

    def test_donut_schema_audit_requires_native_fields_and_column_sort(self):
        donut = {"allOf": [
            {"required": ["id", "kind", "source", "columns", "color", "value"],
             "properties": {"kind": {"enum": ["donut-chart"]},
                            "color": {"allOf": [
                                {"required": ["columnId"], "properties": {"columnId": {}}},
                                {"properties": {"sort": {"allOf": [
                                    {"properties": {"direction": {"enum": ["ascending", "descending"]}}},
                                    {"properties": {"by": {}, "aggregation": {"enum": ["min"]}}}]}}}]},
                            "value": {"required": ["columnId"], "properties": {"columnId": {}}}}},
            {"properties": {"innerRadius": {"type": "number"}}}]}
        openapi = {"components": {"schemas": {"CommonElement": {"oneOf": [donut]}}}}
        self.assertEqual(audit_donut_openapi(openapi)["status"], "PASS")
        broken = copy.deepcopy(openapi)
        del broken["components"]["schemas"]["CommonElement"]["oneOf"][0]["allOf"][1]["properties"]["innerRadius"]
        self.assertEqual(audit_donut_openapi(broken)["status"], "FAIL")
        wrong_sort = copy.deepcopy(openapi)
        wrong_sort["components"]["schemas"]["CommonElement"]["oneOf"][0]["allOf"][0]["properties"]["color"]["allOf"][1]["properties"]["sort"]["allOf"][1]["properties"]["aggregation"]["enum"] = ["sum"]
        self.assertEqual(audit_donut_openapi(wrong_sort)["status"], "FAIL")

    def test_chart_format_comes_from_source_cells_when_chart_cache_has_none(self):
        retail = self.plans[2]
        series = retail["charts"][0]["series"][0]
        self.assertEqual(series["numberFormat"], "#,##0;[Red](#,##0)")
        artifact = emit_preview(retail, synthetic_bindings(retail))
        chart = next(item for item in artifact["elements"] if item["kind"] == "bar-chart")
        self.assertEqual(chart["columns"][1]["format"]["formatString"], ",.0f")
        self.assertTrue(any(item["kind"] == "chart-format" and item["sourceId"] == "chart-1"
                            for item in artifact["unresolved"]))
        demand = self.plans[4]
        demand_chart = next(item for item in emit_preview(demand, synthetic_bindings(demand))["elements"]
                            if item["kind"] == "line-chart")
        self.assertNotIn("format", demand_chart["columns"][1])
        regional = self.plans[0]
        dated = next(item for item in emit_preview(regional, synthetic_bindings(regional))["elements"]
                     if item["kind"] == "line-chart")
        self.assertEqual(dated["columns"][0]["format"],
                         {"kind": "datetime", "formatString": "%b-%y"})

    def test_missing_point_or_literal_target_cannot_be_presented_as_live_coverage(self):
        plan = self.plans[0]
        bindings = synthetic_bindings(plan)
        first = bindings["charts"][plan["charts"][0]["id"]]
        first["points"].pop()
        with self.assertRaisesRegex(ValueError, "point coverage"):
            emit_preview(plan, bindings)
        first["points"] = synthetic_bindings(plan)["charts"][plan["charts"][0]["id"]]["points"]
        first["series"][0]["formula"] = "12345"
        with self.assertRaisesRegex(ValueError, "live expression"):
            emit_preview(plan, bindings)

    def test_source_rules_retain_whole_number_and_priority_without_invented_mapping(self):
        inventory = self.plans[7]
        self.assertTrue(all(rule["sourceXml"]["sourceHash"] == inventory["sourceHash"]
                            for rule in inventory["validations"]))
        whole = [rule for plan in self.plans for rule in plan["validations"]
                 if rule["sourceType"] == "whole"]
        self.assertEqual(len(whole), 1)
        self.assertEqual((whole[0]["operator"], whole[0]["formula1"]), ("greaterThanOrEqual", "0"))
        self.assertEqual(whole[0]["status"], "unresolved")
        self.assertIsNone(whole[0]["emission"])
        finite = build_source_plan(ROOT / "tests" / "fixtures" / "fixture-inventory-reorder.xlsx")
        finite_whole = next(rule for rule in finite["validations"] if rule["sourceType"] == "whole")
        self.assertEqual(finite_whole["emission"], {"values": list(range(1, 91))})
        cf = [rule for plan in self.plans for rule in plan["conditionalFormats"]]
        self.assertTrue(all("relativeReference" in rule and "priority" in rule for rule in cf))
        self.assertTrue(any(rule["status"] == "unresolved" for rule in finite["conditionalFormats"]))

    def test_condition_rules_keep_excel_priority_and_relative_reference_provenance(self):
        plan = build_source_plan(ROOT / "tests" / "fixtures" / "fixture-expense-approval.xlsx")
        bindings = synthetic_bindings(plan)
        rules = [rule for rule in plan["conditionalFormats"] if rule["range"] == "O2:O49"]
        self.assertEqual([rule["priority"] for rule in rules], [1, 2, 3])
        bindings["conditionalTargets"] = {
            rule["id"]: [{"elementId": "detail", "columnId": "status",
                          "sourceRange": rule["range"], "sourceSheet": rule["sheet"]}]
            for rule in rules}
        artifact = emit_preview(plan, bindings)
        native = artifact["elementPatches"]["detail"]["conditionalFormats"]
        self.assertEqual([item["style"]["backgroundColor"] for item in native],
                         ["#E2F0D9", "#FFF2CC", "#FCE4D6"])
        retained = [rule for rule in artifact["ruleProvenance"] if rule["range"] == "O2:O49"]
        self.assertTrue(all("relativeReference" in rule for rule in retained))

    def test_cli_exports_source_plan_and_refuses_preview_as_put(self):
        source = self.paths[2]
        plan = self.plans[2]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output, binding_path = root / "output.json", root / "bindings.json"
            script = str(ROOT / "scripts" / "build-source-presentation.py")
            plan_call = subprocess.run([sys.executable, script, str(source), "--out", str(output)],
                                       capture_output=True, text=True)
            self.assertEqual(plan_call.returncode, 0, plan_call.stderr)
            exported = json.loads(output.read_text())
            self.assertEqual(exported["planHash"], plan["planHash"])
            self.assertEqual(exported["charts"][0]["title"], "Net Sales by Quarter")
            binding_path.write_text(json.dumps(synthetic_bindings(plan)))
            preview = subprocess.run([sys.executable, script, str(source), "--preview",
                                      "--bindings", str(binding_path), "--contents", str(output),
                                      "--out", str(root / "put.json")],
                                     capture_output=True, text=True)
            self.assertNotEqual(preview.returncode, 0)
            self.assertFalse((root / "put.json").exists())

    def test_approved_emission_reuses_staging_gate_and_requires_reviewed_chart_facts(self):
        source = ROOT / "tests" / "fixtures" / "fixture-sales-pipeline.xlsx"
        extracted = presentation.extract_presentation(source)
        architecture = presentation.build_business_input_architecture(
            extracted, connection_id="conn", folder_id="folder")
        logical = presentation.build_business_input_plan(extracted, architecture)
        manifest = build_manifest(source, logical)
        manifest["target"] = {"documentId": "logical", "version": 1}
        _source, evidence = gate_plan_staging_source(logical, source_arg=source)
        target_bindings = {feature: {"documentId": "logical", "version": 1, "pageId": "edit"}
                           for feature in architecture["inputs"]}
        approval = {"sourceHash": manifest["sourceHash"], "planHash": manifest["planHash"],
                    "owner": "Owner", "reviewer": "Reviewer",
                    "reviewItems": [item["id"] for item in manifest["formulaCells"]],
                    "mappedFeatures": architecture["mappedFeatures"],
                    "targetBindings": target_bindings,
                    "sourceReviewApproval": {"status": "approved_for_staging",
                                             "issues": evidence.get("issues") or []}}
        receipt = gate_input_staging(manifest, approval, logical, evidence)
        plan = build_source_plan(source)
        bindings = synthetic_bindings(plan)
        bindings.update(manifestPlanHash=manifest["planHash"], stagingReceiptHash=_digest(receipt))
        openapi = {"info": {"version": "test"}, "paths": {"/v2/workbooks/{workbookId}/contents": {}},
                   "components": {"schemas": {
                       "Kinds": {"enum": ["bar-chart", "line-chart", "single", "backgroundScale",
                                          "fontScale", "number", "datetime", "center", "hidden", "none"]},
                       "Chart": {"properties": {"xAxis": {}, "yAxis": {}, "columns": {}, "columnIds": {}}},
                       "Validation": {"properties": {"values": {}, "valuesFrom": {}, "range": {}}},
                       "Conditional": {"properties": {"columnIds": {}, "scheme": {}, "domain": {},
                                                      "order": {}, "nullBehavior": {}}},
                       "Number": {"properties": {"formatString": {}, "currencySymbol": {}}},
                       "Text": {"properties": {"verticalAlign": {}, "text": {}, "visibility": {}}}}}}
        with self.assertRaisesRegex(ValueError, "reviewed binding facts"):
            emit_approved(plan, bindings, manifest, approval, receipt, openapi)
        for chart in bindings["charts"].values():
            chart["review"] = {"status": "approved", "reviewer": "Reviewer",
                               "evidenceId": "reviewed-source-range", "sourceHash": plan["sourceHash"],
                               "presentationPlanHash": plan["planHash"],
                               "manifestPlanHash": manifest["planHash"],
                               "stagingReceiptHash": _digest(receipt),
                               "documentId": "logical", "documentVersion": 1,
                               "bindingHash": _digest(chart)}
        artifact = emit_approved(plan, bindings, manifest, approval, receipt, openapi)
        self.assertEqual(artifact["stagingGate"]["status"], "APPROVED_STAGING")
        self.assertEqual(len(artifact["aggregateChecks"]), 16)
        self.assertTrue(all(check["status"] == "PENDING_RUNTIME_QUERY"
                            for check in artifact["aggregateChecks"]))
        rule = plan["validations"][0]
        target = {"elementId": "edit", "columnId": "probability", "sourceRange": rule["range"],
                  "sourceSheet": rule["sheet"]}
        bindings["validationTargets"] = {rule["id"]: [target]}
        with self.assertRaisesRegex(ValueError, "reviewed target"):
            emit_approved(plan, bindings, manifest, approval, receipt, openapi)
        target["review"] = {"status": "approved", "reviewer": "Reviewer",
                            "evidenceId": "range-column-review", "sourceHash": plan["sourceHash"],
                            "presentationPlanHash": plan["planHash"],
                            "manifestPlanHash": manifest["planHash"],
                            "stagingReceiptHash": _digest(receipt),
                            "documentId": "logical", "documentVersion": 1,
                            "targetBindingHash": _digest(target)}
        artifact = emit_approved(plan, bindings, manifest, approval, receipt, openapi)
        self.assertTrue(artifact["elementPatches"]["edit"]["columnUpdates"]["probability"])
        stale_chart = copy.deepcopy(bindings)
        next(iter(stale_chart["charts"].values()))["review"]["documentVersion"] = 2
        with self.assertRaisesRegex(ValueError, "reviewed binding facts"):
            emit_approved(plan, stale_chart, manifest, approval, receipt, openapi)
        changed = copy.deepcopy(receipt)
        changed["reviewer"] = "tampered"
        with self.assertRaisesRegex(ValueError, "stale"):
            emit_approved(plan, bindings, manifest, approval, changed, openapi)


if __name__ == "__main__":
    unittest.main()
