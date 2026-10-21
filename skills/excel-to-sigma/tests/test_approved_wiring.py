"""Actual offline public paths; runtime Sigma execution remains an acceptance gate."""
import copy
import importlib.util
import json
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / 'scripts'
sys.path.insert(0, str(SCRIPTS / 'lib'))
from input_sync import (InputSyncError, build_artifact_records, emit_workbook_wiring,
                        plan_seed, plan_input_extension, checkpoint_input_extension,
                        plan_extension_seed, emit_write_programs)
import presentation

FIXTURES = ROOT / 'tests' / 'fixtures'
CLI = SCRIPTS / 'presentation-inputs.py'


def receipt(manifest, assigned=True):
    architecture = manifest['planProjection']['sharedInputArchitecture']
    bindings = {}
    for feature, spec in architecture['inputs'].items():
        bindings[feature] = {'documentId': 'doc', 'version': 8, 'pageId': 'host'}
        if assigned:
            bindings[feature].update(elementId='assigned-' + feature, columnIds={
                c['name']: 'assigned-' + str(i) for i, c in enumerate(spec['columns'])})
    result = {'status': 'APPROVED_STAGING', 'sourceHash': manifest['sourceHash'],
              'planHash': manifest['planHash'], 'targetBindings': bindings}
    result['artifacts'] = build_artifact_records(manifest, architecture, bindings)
    return result


class TestApprovedWiring(unittest.TestCase):
    def prepare(self, workflow, tmp):
        names = {'sales': 'sales-pipeline', 'inventory': 'inventory-reorder',
                 'expense': 'expense-approval', 'forecast': 'forecast-model'}
        paths = [tmp / 'plan.json', tmp / 'manifest.json']
        subprocess.run([sys.executable, str(CLI), 'prepare',
                        str(FIXTURES / f'fixture-{names[workflow]}.xlsx'),
                        '--connection', 'conn', '--folder', 'folder',
                        '--out-plan', str(paths[0]), '--out-manifest', str(paths[1])],
                       check=True, capture_output=True, text=True)
        return json.loads(paths[0].read_text()), json.loads(paths[1].read_text())

    def test_approved_emit_owns_operational_derivations_and_real_check_targets(self):
        for workflow, expected in [('sales', {'Weighted ACV', 'Cycle Days'}),
                                   ('inventory', {'Available', 'Reorder Point', 'Days Cover',
                                                  'Reorder Qty', 'Status'}),
                                   ('expense', set())]:
            with self.subTest(workflow=workflow), tempfile.TemporaryDirectory() as directory:
                tmp = Path(directory)
                plan, manifest = self.prepare(workflow, tmp)
                staging = receipt(manifest)
                result = emit_workbook_wiring(manifest, staging)
                elements = {e['id']: e for e in result['wiring']['elements']}
                target = {'sales': 'salesInputBase', 'inventory': 'inventory',
                          'expense': 'expenseNumericAmountCount'}[workflow]
                self.assertIn(target, elements)
                self.assertEqual(manifest['planProjection']['presentationConditionalFormats'],
                                 plan['presentationConditionalFormats'])
                for rule in plan['presentationConditionalFormats']:
                    if rule['sourceType']=='colorScale':
                        self.assertIn('<colorScale>', rule['sourceRuleXml'])
                columns = {c['name']: c for c in elements[target]['columns']}
                self.assertTrue(expected <= columns.keys())
                self.assertNotIn('input-table', {e['kind'] for e in elements.values()},
                                 'assigned inputs must never be duplicated by the adapter')
                if workflow == 'inventory':
                    self.assertEqual(columns['Reorder Qty']['id'], 'reorder_qty')
                    self.assertIn('RoundUp(', columns['Reorder Qty']['formula'])
                    self.assertEqual(elements[target]['source']['elementId'],
                                     'assigned-inventoryBusinessInput')
                if workflow == 'sales':
                    self.assertIn('[Probability] * [ACV]', columns['Weighted ACV']['formula'])
                    self.assertNotIn('/Weighted ACV]', columns['Weighted ACV']['formula'])
                if workflow == 'expense':
                    self.assertEqual(columns['Numeric Amount Count']['formula'],
                                     'Count([Expense Business Inputs/Amount])')
                    check = result['wiring']['checks'][0]
                    self.assertEqual(check['logicalCheckId'], 'expense-numeric-amount-count')
                    self.assertEqual(check['elementId'], target)
                    self.assertEqual(check['columnId'], columns['Numeric Amount Count']['id'])
                    self.assertEqual(check['status'], 'pending-runtime-query')
                # The public emission returns the hash-bound logical graph even before IDs exist.
                (tmp / 'receipt.json').write_text(json.dumps(receipt(manifest, assigned=False)))
                subprocess.run([sys.executable, str(CLI), 'emit', '--manifest', str(tmp/'manifest.json'),
                                '--receipt', str(tmp/'receipt.json'), '--feature', workflow+'BusinessInput',
                                '--out', str(tmp/'emitted.json')], check=True, capture_output=True)
                emitted = json.loads((tmp/'emitted.json').read_text())
                self.assertIn('wiringTemplate', emitted)
                self.assertTrue(emitted['wiringTemplate']['elements'])
                # Exact target/version and architecture swaps invalidate the approval.
                foreign = copy.deepcopy(manifest)
                foreign['target'] = {'documentId': 'foreign', 'version': 8}
                with self.assertRaisesRegex(InputSyncError, 'target|stale'):
                    emit_workbook_wiring(foreign, staging)
                altered = copy.deepcopy(manifest)
                altered['planProjection']['sharedInputArchitecture']['inputs'][workflow+'BusinessInput']['rows'][0]['foreign'] = 1
                with self.assertRaisesRegex(InputSyncError, 'artifact|stale'):
                    emit_workbook_wiring(altered, staging)

    def test_public_expense_policy_excess_uses_scalar_clamp_and_retains_missing_policy_guard(self):
        from native_graph_oracle import scalar
        with tempfile.TemporaryDirectory() as directory:
            _,manifest=self.prepare('expense',Path(directory))
            result=emit_workbook_wiring(manifest,receipt(manifest))
            calc=next(e for e in result['wiring']['elements'] if e['id']=='expensesWithPolicy')
            formula=next(c['formula'] for c in calc['columns'] if c['name']=='PolicyExcess')
            self.assertIn('Greatest(',formula)
            self.assertNotIn('Max(',formula)
            for category,amount,limit,expected in [('Meals',100,90,10),('Meals',80,90,0),
                                                  ('Meals',90,90,0),(None,100,None,None)]:
                self.assertEqual(scalar(formula,{'Policy/Category':category,'ExpenseJoin/Amount':amount,
                                                'Policy/Per-item Limit':limit}),expected)
            architecture=manifest['planProjection']['sharedInputArchitecture']
            self.assertIn('Greatest(',architecture['calculations']['PolicyExcess'])

    def test_public_actual_bindings_integrate_all_four_without_logical_companion_targets(self):
        names={'forecast':'forecast-model','sales':'sales-pipeline','inventory':'inventory-reorder','expense':'expense-approval'}
        for workflow,stem in names.items():
            with self.subTest(workflow=workflow), tempfile.TemporaryDirectory() as directory:
                _,manifest=self.prepare(workflow,Path(directory))
                manifest['target']={'documentId':'doc','version':8}
                staging=receipt(manifest)
                graph=emit_workbook_wiring(manifest,staging)['wiring']['elements']
                inputs=[{'id':target['elementId'],'kind':'input-table','name':
                    manifest['planProjection']['sharedInputArchitecture']['inputs'][feature]['name'],
                    'columns':[{'id':column_id,'name':name}for name,column_id in target['columnIds'].items()]}
                    for feature,target in staging['targetBindings'].items()]
                current={'schemaVersion':1,'pages':[{'id':'host','name':'Host'}],
                         'elements':inputs+graph,'layout':'<Page id="host"></Page>'}
                extracted=presentation.extract_presentation(FIXTURES/('fixture-'+stem+'.xlsx'))
                from migration_gate import gate_plan_staging_source,source_issue_hash
                _,evidence=gate_plan_staging_source({**copy.deepcopy(manifest['planProjection']),
                    'sourceSha256':manifest['sourceHash'],'sourcePath':extracted['source']['path']},
                    source_arg=extracted['source']['path'])
                issues=evidence.get('issues') or []
                staging.update(pendingSourceIssues=issues,sourceReviewStatus='PENDING_STAGING' if issues else 'NOT_REQUIRED',
                    sourceReviewRequirements=[{'issue':issue,'issueHash':source_issue_hash(issue)}for issue in issues])
                bindings=presentation.approved_business_bindings(extracted,manifest,staging)
                artifact=presentation.emit_presentation(extracted,bindings,manifest,staging)
                integrated=presentation.integrate_presentation(current,artifact)
                by_id={e['id']:e for e in integrated['elements']}
                for element_id,patch in artifact['elementPatches'].items():
                    self.assertIn(element_id,by_id)
                    ids={c['id']for c in by_id[element_id]['columns']}
                    self.assertTrue(set(patch['columnUpdates'])<=ids)
                    self.assertTrue(all(set(rule['columnIds'])<=ids for rule in patch['conditionalFormats']))
                for chart in artifact['elements']:
                    if chart['kind'].endswith('-chart'):self.assertIn(chart['source']['elementId'],by_id)
                if workflow=='forecast':
                    self.assertEqual(bindings['columns']['Actuals Detail!Amount']['elementId'],'actuals')
                    self.assertNotIn('Actuals Detail!Date',bindings['columns'])
                if workflow=='expense':
                    self.assertEqual(bindings['columns']['Policy!Per-item Limit']['elementId'],'assigned-expensePolicy')
                for sheet in {'sales':['Assumptions','Rep Targets'],'inventory':['Suppliers']}.get(workflow,[]):
                    self.assertFalse(any(key.startswith(sheet+'!')for key in bindings['columns']))
                    self.assertTrue(any(u['sourceId'].startswith(sheet+'!')for u in artifact['unresolved']))

    def test_forecast_public_adapter_consumes_assigned_wide_input_and_same_pnl_chain(self):
        with tempfile.TemporaryDirectory() as directory:
            plan, manifest = self.prepare('forecast', Path(directory))
            emitted = emit_workbook_wiring(manifest, receipt(manifest))
            elements = {e['id']: e for e in emitted['wiring']['elements']}
            self.assertTrue({'actuals', 'spineWrap', 'forecast', 'pnlCombined'} <= elements.keys())
            forecast = elements['forecast']
            join = forecast['source']['joins'][0]
            self.assertEqual(join['right']['elementId'], 'assigned-forecastBusinessInput')
            self.assertEqual(join['columns'], [{'left': '[Line Item]', 'right': '[Line Item]'}])
            formulas = {c['name']: c['formula'] for c in forecast['columns']}
            self.assertEqual(formulas['Kind'], '[Inputs/Method]')
            self.assertEqual(formulas['Eff Rate'], '[Inputs/Monthly Growth]')
            for month in ['Jul-26', 'Aug-26', 'Sep-26', 'Oct-26', 'Nov-26', 'Dec-26']:
                self.assertIn('[Inputs/'+month+']', formulas['Eff Manual'])
            self.assertNotIn('Coalesce', json.dumps(forecast))
            self.assertEqual(elements['pnlCombined']['source']['sources'][1]['elementId'], 'forecast')
            self.assertNotIn('rateInput', json.dumps(emitted))
            self.assertNotIn('manualInput', json.dumps(emitted))

    def test_public_wiring_merges_only_approved_nodes_and_rejects_foreign_collisions(self):
        with tempfile.TemporaryDirectory() as directory:
            tmp = Path(directory)
            plan, manifest = self.prepare('inventory', tmp)
            staging = receipt(manifest)
            manifest['target'] = {'documentId': 'doc', 'version': 8}
            (tmp/'manifest.json').write_text(json.dumps(manifest))
            (tmp/'receipt.json').write_text(json.dumps(staging))
            current = {'documentVersion': 8, 'workbookId': 'doc', 'contents': {
                'schemaVersion': 1, 'pages': [{'id': 'host', 'name': 'Host'}],
                'elements': [{'id': 'assigned-inventoryBusinessInput', 'kind': 'input-table',
                              'name': 'Inventory Business Inputs', 'columns': [
                                  {'id': column_id, 'name': name, 'type': next(c['type'] for c in
                                   plan['sharedInputArchitecture']['inputs']['inventoryBusinessInput']['columns'] if c['name']==name)}
                                  for name, column_id in staging['targetBindings']['inventoryBusinessInput']['columnIds'].items()]}],
                'layout': '<Page id="host"></Page>\n \n'}}
            (tmp/'current.json').write_text(json.dumps(current))
            command = [sys.executable, str(CLI), 'wiring', '--manifest', str(tmp/'manifest.json'),
                       '--receipt', str(tmp/'receipt.json'), '--contents', str(tmp/'current.json'),
                       '--out', str(tmp/'put.json')]
            first = subprocess.run(command, capture_output=True, text=True)
            self.assertEqual(first.returncode, 0, first.stderr)
            put = json.loads((tmp/'put.json').read_text())
            self.assertEqual(put['documentVersion'], 8)
            self.assertIn('inventory', {e['id'] for e in put['contents']['elements']})
            self.assertIn('elementId="inventory"', put['contents']['layout'])
            self.assertTrue(put['contents']['layout'].startswith(current['contents']['layout'] + '\n'))
            # Feed the GET-shaped result back into the same public path: idempotent.
            (tmp/'current.json').write_text(json.dumps({'workbookId': 'doc', **put}))
            second = subprocess.run(command, capture_output=True, text=True)
            self.assertEqual(second.returncode, 0, second.stderr)
            self.assertEqual(json.loads((tmp/'put.json').read_text())['contents'], put['contents'])
            # Owned native elements elsewhere cannot confer ownership on a
            # user page with the deterministic native-model page ID.
            foreign_page = copy.deepcopy(current)
            owned_node = next(e for e in put['contents']['elements'] if e['id']=='inventory')
            foreign_page['contents']['elements'] += [owned_node,
                {'id':'foreign-widget', 'kind':'text', 'body':'User analysis'}]
            foreign_page['contents']['pages'].append(
                {'id':'inventory-native-model', 'name':'User Analysis'})
            foreign_page['contents']['layout'] += (
                '<Page id="inventory-native-model"><Element elementId="foreign-widget" '
                'gridColumn="1 / 25" gridRow="1 / 5"/></Page>')
            # Also reject a user placement added to an otherwise owned page.
            polluted = {'workbookId':'doc', **copy.deepcopy(put)}
            polluted['contents']['elements'].append(
                {'id':'foreign-widget', 'kind':'text', 'body':'User analysis'})
            polluted['contents']['layout'] = polluted['contents']['layout'].replace(
                '</Page>', '<Element elementId="foreign-widget" gridColumn="1 / 25" '
                'gridRow="13 / 17"/></Page>')
            for unsafe in (foreign_page, polluted):
                original = json.dumps(unsafe)
                (tmp/'current.json').write_text(original)
                (tmp/'put.json').write_text('do not replace existing output')
                rejected = subprocess.run(command, capture_output=True, text=True)
                self.assertNotEqual(rejected.returncode, 0, 'foreign page content was accepted')
                self.assertRegex(rejected.stderr, 'foreign|ownership|collid')
                self.assertEqual((tmp/'current.json').read_text(), original)
                self.assertEqual((tmp/'put.json').read_text(), 'do not replace existing output')
            swapped = copy.deepcopy(current)
            swapped['contents']['elements'][0]['columns'][0]['id'] = 'foreign-column'
            (tmp/'current.json').write_text(json.dumps(swapped))
            rejected = subprocess.run(command, capture_output=True, text=True)
            self.assertNotEqual(rejected.returncode, 0)
            self.assertIn('column', rejected.stderr)
            current['contents']['elements'].append({'id': 'inventory', 'name': 'Foreign', 'kind': 'table'})
            (tmp/'current.json').write_text(json.dumps(current))
            rejected = subprocess.run(command, capture_output=True, text=True)
            self.assertNotEqual(rejected.returncode, 0)
            self.assertIn('collid', rejected.stderr)

    def test_forecast_builder_can_consume_only_exact_approved_wide_adapter(self):
        from test_builders import FORECAST, forecast_plan
        with tempfile.TemporaryDirectory() as directory:
            plan, manifest = self.prepare('forecast', Path(directory))
            legacy = forecast_plan()
            legacy.update(intent='editable', sourceSha256=manifest['sourceHash'], connectionId='conn')
            wb = FORECAST.build_wb_spec(legacy, 'dm',
                {'glDetail': 'gl', 'spine': 'sp', 'assumptions': 'as'},
                approved_inputs={'manifest': manifest, 'receipt': receipt(manifest)})
            nodes = [e for p in wb['pages'] for e in p['elements']]
            self.assertNotIn('rateInput', {e['id'] for e in nodes})
            self.assertNotIn('manualInput', {e['id'] for e in nodes})
            self.assertIn('assigned-forecastBusinessInput', json.dumps(nodes))
            legacy['sourceSha256'] = 'foreign'
            with self.assertRaisesRegex(ValueError, 'source|stale'):
                FORECAST.build_wb_spec(legacy, 'dm',
                    {'glDetail': 'gl', 'spine': 'sp', 'assumptions': 'as'},
                    approved_inputs={'manifest': manifest, 'receipt': receipt(manifest)})

    def test_emitted_forecast_matches_144_source_values_and_36_trends_after_real_edits(self):
        from native_graph_oracle import run_forecast
        from formula_graph import FormulaGraph, CellIdentity
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            tmp = Path(directory)
            plan, manifest = self.prepare('forecast', tmp)
            architecture = plan['sharedInputArchitecture']
            elements = emit_workbook_wiring(manifest, receipt(manifest))['wiring']['elements']
            workbook = load_workbook(FIXTURES/'fixture-forecast-model.xlsx', data_only=False)
            for edited in (False, True):
                wide = copy.deepcopy(architecture['inputs']['forecastBusinessInput']['rows'])
                if edited:
                    wide[0]['Monthly Growth'] = 0.05
                    wide[2]['Jul-26'] = 0
                    workbook['Forecast Inputs']['C6'] = 0.05
                    workbook['Forecast Inputs']['D8'] = 0
                native = run_forecast(elements, wide)
                expected = FormulaGraph.from_workbook(workbook).evaluate()
                monthly = {}
                for record in native:
                    key = (record['Line Item'], record['Month'])
                    monthly[key] = monthly.get(key, 0) + record['Amount']
                checked = trends = 0
                for r in range(8, 20):
                    line = workbook['Monthly Model'].cell(r, 2).value
                    for c in range(3, 15):
                        month = workbook['Monthly Model'].cell(5, c).value
                        self.assertAlmostEqual(monthly[line, month],
                            expected[CellIdentity('Monthly Model', workbook['Monthly Model'].cell(r,c).coordinate)], delta=1e-6)
                        checked += 1
                for c in range(3, 15):
                    month = workbook['Monthly Model'].cell(5, c).value
                    records = [row for row in native if row['Month'] == month]
                    revenue = sum(row['Revenue Amt'] for row in records)
                    ebitda = revenue - sum(row['COGS Amt']+row['OpEx Amt'] for row in records)
                    income = ebitda - sum(row['Below Amt'] for row in records)
                    for rownum, value in [(22, revenue), (27, ebitda), (28, income)]:
                        self.assertAlmostEqual(value, expected[CellIdentity('Monthly Model',
                            workbook['Monthly Model'].cell(rownum,c).coordinate)], delta=1e-6)
                        trends += 1
                self.assertEqual((checked, trends), (144, 36))

    def test_method_edits_and_different_year_use_authored_wide_cells(self):
        import datetime as dt
        from openpyxl import load_workbook
        from native_graph_oracle import run_forecast
        from migration_manifest import build_manifest
        with tempfile.TemporaryDirectory() as directory:
            tmp = Path(directory)
            workbook = load_workbook(FIXTURES/'fixture-forecast-model.xlsx', data_only=False)
            for sheet in workbook:
                for row in sheet:
                    for cell in row:
                        if isinstance(cell.value, dt.datetime): cell.value = cell.value.replace(year=2027)
                        elif isinstance(cell.value, str) and re.fullmatch(r'[A-Z][a-z]{2}-26', cell.value):
                            cell.value = cell.value.replace('-26', '-27')
            source = tmp/'different-forecast.xlsx'
            workbook.save(source)
            extracted = presentation.extract_presentation(source)
            architecture = presentation.build_business_input_architecture(extracted, connection_id='conn', folder_id='folder')
            logical = presentation.build_business_input_plan(extracted, architecture)
            manifest = build_manifest(source, logical)
            elements = emit_workbook_wiring(manifest, receipt(manifest))['wiring']['elements']
            wide = copy.deepcopy(architecture['inputs']['forecastBusinessInput']['rows'])
            # Growth→manual, manual→growth, and growth→flat must read the edited Method.
            wide[0]['Method'] = 'manual'
            wide[0]['Jul-27'] = 0
            wide[0]['Aug-27'] = 4321
            wide[2]['Method'] = 'growth'
            wide[2]['Monthly Growth'] = 0.1
            wide[1]['Method'] = 'flat'
            native = run_forecast(elements, wide)
            forecast = {(r['Line Item'], r['Month']): r['Amount'] for r in native if r['Period']=='Forecast'}
            self.assertEqual(forecast[wide[0]['Line Item'], 'Jul-27'], 0)
            self.assertEqual(forecast[wide[0]['Line Item'], 'Aug-27'], 4321)
            for index in (1,2):
                base = sum(r['Amount'] for r in architecture['forecastModel']['actualRows']
                           if r['Line Item']==wide[index]['Line Item'] and r['Month']=='Jun-27')
                for exponent, month in enumerate(architecture['forecastModel']['months'],1):
                    wanted = base if index == 1 else base * 1.1**exponent
                    self.assertAlmostEqual(forecast[wide[index]['Line Item'], month['month']], wanted)

    def test_public_bindings_resolve_views_charts_and_rule_fields_to_emitted_targets(self):
        for workflow in ('forecast', 'sales', 'inventory', 'expense'):
            with self.subTest(workflow=workflow), tempfile.TemporaryDirectory() as directory:
                tmp = Path(directory)
                plan, manifest = self.prepare(workflow, tmp)
                extracted = presentation.extract_presentation(plan['sourcePath'])
                manifest['target'] = {'documentId': 'doc', 'version': 8}
                staging = receipt(manifest)
                bindings = presentation.approved_business_bindings(extracted, manifest, staging)
                graph = emit_workbook_wiring(manifest, staging)['wiring']
                targets = {e['id']: e for e in graph['elements']}
                view = targets[bindings['viewSourceElementId']]
                names = {c['name'] for c in view['columns']}
                for field in bindings['mainViewColumns'] + bindings['detailViewColumns']:
                    references = re.findall(r'\[([^\]]+)\]', field['formula'])
                    self.assertTrue(references)
                    for ref in references:
                        namespace, name = ref.split('/',1)
                        self.assertEqual(namespace, view['name'])
                        self.assertIn(name, names)
                for chart in bindings['charts'].values():
                    target = targets[chart['sourceElementId']]
                    chart_names = {c['name'] for c in target['columns']}
                    for formula in [chart['categoryFormula'], chart['orderFormula'], *chart['seriesFormulas']]:
                        for ref in re.findall(r'\[([^\]]+)\]', formula):
                            self.assertIn(ref.split('/',1)[1], chart_names)
                primary = workflow+'BusinessInput'
                for rule in extracted['validations']:
                    for target in bindings['validationTargets'][rule['id']]:
                        if target['elementId'] == staging['targetBindings'][primary]['elementId']:
                            self.assertIn(target['columnId'], staging['targetBindings'][primary]['columnIds'].values())
                source_rules = [rule for rule in extracted['conditionalFormats']
                    if rule['status'] == 'supported'
                    and rule['sheet'] == presentation.BUSINESS_TABLES[workflow][0]]
                for rule in source_rules:
                    element_ids = {target['elementId']
                        for target in bindings['conditionalTargets'][rule['id']]}
                    expected_views = set() if workflow == 'forecast' else {
                        workflow+'-main-view', workflow+'-detail-view'}
                    self.assertEqual(element_ids & {
                        workflow+'-main-view', workflow+'-detail-view'}, expected_views)

    def test_pending_replacement_cannot_claim_an_unemitted_target(self):
        with tempfile.TemporaryDirectory() as directory:
            plan, manifest = self.prepare('inventory', Path(directory))
            extracted = presentation.extract_presentation(FIXTURES/'fixture-inventory-reorder.xlsx')
            item = next(f for f in extracted['formulaProvenance'] if f['cell']=='Inventory!N2')
            replacement = {'id': item['cell'], 'mappingStatus': 'live',
                'targetMapping': {'logicalElementId': 'nonexistent', 'logicalColumnId': 'reorder_qty'},
                'reviewedReplacement': {'originalFormula': item['formula'],
                    'replacementFormula': '0', 'reviewer': 'Test Reviewer', 'reviewStatus': 'approved'}}
            with self.assertRaisesRegex(ValueError, 'target|emitted'):
                presentation.build_business_input_plan(extracted, plan['sharedInputArchitecture'], [replacement])

    def test_writer_rejects_stale_source_or_tampered_reconciliation_actions(self):
        with tempfile.TemporaryDirectory() as directory:
            logical, manifest = self.prepare('inventory', Path(directory))
            spec = logical['sharedInputArchitecture']['inputs']['inventoryBusinessInput']
            staging = receipt(manifest)
            seed = plan_seed(spec['rows'], [], spec['keyColumns'], 'Row ID')
            self.assertTrue(emit_write_programs(seed, manifest, staging, 'inventoryBusinessInput'))
            stale = copy.deepcopy(spec['rows'])
            stale[0]['On Hand'] += 1
            stale_seed = plan_seed(stale, [], spec['keyColumns'], 'Row ID')
            with self.assertRaisesRegex(InputSyncError, 'source|seed'):
                emit_write_programs(stale_seed, manifest, staging, 'inventoryBusinessInput')
            tampered = copy.deepcopy(seed)
            tampered['inserts'][0]['source']['On Hand'] += 1000
            with self.assertRaisesRegex(InputSyncError, 'digest|reconciliation|source'):
                emit_write_programs(tampered, manifest, staging, 'inventoryBusinessInput')

    def test_native_inventory_and_sales_derivations_respond_to_literal_edits(self):
        import datetime as dt
        import math
        from native_graph_oracle import project
        for workflow in ('inventory', 'sales'):
            with tempfile.TemporaryDirectory() as directory:
                logical, manifest = self.prepare(workflow, Path(directory))
                architecture = logical['sharedInputArchitecture']
                spec = architecture['inputs'][workflow+'BusinessInput']
                elements = emit_workbook_wiring(manifest, receipt(manifest))['wiring']['elements']
                target = next(e for e in elements if e['id']==('inventory' if workflow=='inventory' else 'salesInputBase'))
                for edited in (False, True):
                    rows = copy.deepcopy(spec['rows'])
                    if edited and workflow=='inventory':
                        rows[0]['On Hand'] = 0
                        rows[1]['Avg Daily Use'] = 3.25
                    elif edited:
                        rows[0]['Probability'] = 0.75
                        rows[0]['ACV'] = 200000
                    for source in rows:
                        if workflow=='sales':
                            for name in ('Created Date', 'Expected Close'):
                                source[name] = dt.datetime.fromisoformat(source[name]).date()
                        computed = project(target['columns'], {spec['name']+'/'+key: value for key,value in source.items()})
                        if workflow=='inventory':
                            available = source['On Hand']-source['Allocated']
                            reorder_point = source['Avg Daily Use']*source['Lead Days']+source['Safety Stock']
                            self.assertEqual(computed['Available'], available)
                            self.assertEqual(computed['Reorder Point'], reorder_point)
                            self.assertAlmostEqual(computed['Days Cover'], available/source['Avg Daily Use'])
                            self.assertEqual(computed['Reorder Qty'], math.ceil(max(0,reorder_point+source['Avg Daily Use']*30-available)))
                            self.assertEqual(computed['Status'], 'REORDER' if available<=reorder_point else 'OK')
                        else:
                            self.assertEqual(computed['Weighted ACV'], source['Probability']*source['ACV'])
                            self.assertEqual(computed['Cycle Days'], (source['Expected Close']-source['Created Date']).days)

    def test_public_prepare_rejects_a_foreign_source_review_payload(self):
        with tempfile.TemporaryDirectory() as directory:
            tmp = Path(directory)
            (tmp/'review.json').write_text(json.dumps({'sourceHash':'foreign','formulaCells':[]}))
            result = subprocess.run([sys.executable, str(CLI), 'prepare',
                       str(FIXTURES/'fixture-inventory-reorder.xlsx'), '--connection','conn','--folder','folder',
                       '--reviewed-formulas',str(tmp/'review.json'), '--out-plan',str(tmp/'plan.json'),
                       '--out-manifest',str(tmp/'manifest.json')], capture_output=True,text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('source', result.stderr)

    def test_library_rejects_foreign_and_ambiguous_prior_shapes(self):
        rows = [{'Key': 'A', 'Value': 1}]
        for foreign in ({}, {'programs': []}, {'plan': {}, 'seedRecords': []},
                        {'seedRecords': 'not-list'}, {'unexpected': True},
                        {'seedRecords': [], 'keyColumns': ['Key'], 'rowIdColumn': 'Row ID', 'foreign': True}):
            with self.subTest(prior=foreign), self.assertRaises(InputSyncError):
                plan_seed(rows, [], ['Key'], 'Row ID', prior_seed=foreign)
        baseline = plan_seed(rows, [], ['Key'], 'Row ID')
        with self.assertRaisesRegex(InputSyncError, 'key|scope'):
            plan_seed(rows, [], ['Value'], 'Row ID', prior_seed=baseline)

    def test_cli_persisted_wrapper_cannot_replay_cleared_null_extension(self):
        feature = 'salesBusinessInput'
        architecture = {'targetScope': {'connectionId': 'conn'},
                        'presentationInputFeature': feature,
                        'inputs': {feature: {'name': 'Sales', 'connectionId': 'conn',
                            'columns': [{'name': 'Key', 'type': 'text'}, {'name': 'New', 'type': 'number'}],
                            'keyColumns': ['Key'], 'rows': [{'Key': 'A', 'New': 10}]}}}
        manifest = {'sourceHash': 'a'*64, 'planHash': 'b'*64,
                    'planProjection': {'sharedInputArchitecture': architecture}}
        pre = receipt(manifest)
        pre['targetBindings'][feature]['columnIds'].pop('New')
        pre['artifacts'] = build_artifact_records(manifest, architecture, pre['targetBindings'])
        extension = plan_input_extension(manifest, pre, feature, [
            {'name': 'Key', 'type': 'text', 'id': 'assigned-0'}])
        proof = checkpoint_input_extension(extension, {'New': 'assigned-1'}, result_version=9,
                 readback={'readbackId': 'query-1', 'elementId': 'assigned-'+feature, 'version': 9,
                           'columns': {'New': {'columnId': 'assigned-1', 'value': None}}})
        post = receipt(manifest)
        post['targetBindings'][feature]['version'] = 9
        post['artifacts'] = build_artifact_records(manifest, architecture, post['targetBindings'])
        prior = plan_seed([{'Key': 'A'}], [{'Key': 'A', 'Row ID': 'row'}], ['Key'], 'Row ID')
        current = [{'Key': 'A', 'New': None, 'Row ID': 'row'}]
        foreign_rows = [{'Key': 'A', 'New': 999}]
        with self.assertRaisesRegex(InputSyncError, 'source|seed'):
            plan_extension_seed(foreign_rows, current, ['Key'], 'Row ID', prior_seed=prior,
                                manifest=manifest, staging_receipt=post, feature=feature,
                                extension_proof=proof)
        with self.assertRaisesRegex(InputSyncError, 'readback|omits'):
            plan_extension_seed(architecture['inputs'][feature]['rows'], [{'Key':'A','Row ID':'row'}],
                                ['Key'], 'Row ID', prior_seed=prior, manifest=manifest,
                                staging_receipt=post, feature=feature, extension_proof=proof)
        with tempfile.TemporaryDirectory() as directory:
            tmp = Path(directory)
            for name, value in [('manifest', manifest), ('receipt', post), ('proof', proof),
                                ('prior', prior), ('current', current)]:
                (tmp/(name+'.json')).write_text(json.dumps(value))
            command = [sys.executable, str(CLI), 'seed-extension', '--manifest', str(tmp/'manifest.json'),
                       '--receipt', str(tmp/'receipt.json'), '--feature', feature,
                       '--extension-proof', str(tmp/'proof.json'), '--current-rows', str(tmp/'current.json'),
                       '--row-id', 'Row ID', '--snapshot-id', 'query-2', '--snapshot-age-seconds', '0']
            first = subprocess.run(command + ['--prior-seed', str(tmp/'prior.json'),
                                    '--out', str(tmp/'first.json')], capture_output=True, text=True)
            self.assertEqual(first.returncode, 0, first.stderr)
            written = json.loads((tmp/'first.json').read_text())
            self.assertEqual(written['plan']['updates'][0]['changes'], {'New': 10})
            self.assertNotIn('New', written['plan']['seedRecords'][0]['confirmedMigrationValue'])
            self.assertEqual(written['plan']['seedRecords'][0]['pendingMigrationValue']['New'], 10)
            again = subprocess.run(command + ['--prior-seed', str(tmp/'first.json'),
                                    '--out', str(tmp/'second.json')], capture_output=True, text=True)
            self.assertNotEqual(again.returncode, 0, 'public persisted wrapper replayed creation proof')
            self.assertIn('replay', again.stderr)
            # Normal reconciliation must carry the consumption ledger forward, too.
            reconciled = plan_seed(architecture['inputs'][feature]['rows'], current, ['Key'], 'Row ID',
                                   prior_seed=written)
            self.assertEqual(reconciled['extensionSeedCheckpoints'], written['plan']['extensionSeedCheckpoints'])
            # d19 checkpoints have a consumption ledger but no extension scope.
            # A fresh ordinary reconciliation must remain usable after upgrade.
            legacy = copy.deepcopy(written['plan'])
            legacy.pop('extensionSeedScope')
            applied = [{'Key': 'A', 'New': 10, 'Row ID': 'row'}]
            ordinary = plan_seed(architecture['inputs'][feature]['rows'], applied,
                                 ['Key'], 'Row ID', prior_seed=legacy)
            self.assertEqual(emit_write_programs(ordinary, manifest, post, feature), [])
            conflicted = plan_seed(architecture['inputs'][feature]['rows'], current,
                                   ['Key'], 'Row ID', prior_seed=legacy)
            self.assertTrue(conflicted['conflicts'])
            self.assertEqual(emit_write_programs(conflicted, manifest, post, feature), [])
            with self.assertRaisesRegex(InputSyncError, 'scope'):
                emit_write_programs(legacy, manifest, post, feature)
