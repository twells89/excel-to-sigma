"""Hand-derived multirow criteria results checked through all emitted tables."""
import hashlib
import tempfile
import sys
import unittest
from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts/lib'))
from source_inputs import build_source_input_plan
from native_exact_plan import build_exact_native_calculation_plan, lower_exact_native_calculations
from native_graph_reference import NativeGraphReference


class CompleteNativeGraphTests(unittest.TestCase):
    def test_sumifs_multirow_multicriteria_zero_blank_and_selected_error_lanes(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / 'criteria.xlsx'; book = Workbook(); data = book.active; data.title = 'Renamed'
            data.append(['Group','Amount','Rate','Period','Result'])
            for index, row in enumerate([['a',3,2,1], ['a',14,2,2], ['a',None,2,2], ['b',5,0,2]], 2):
                data.append(row + [f'=B{index}/C{index}'])
            data.add_table(Table(displayName='tblCases', ref='A1:E5'))
            summary = book.create_sheet('Outputs')
            formulas = {
                'A1': '=SUMIFS(tblCases[Result],tblCases[Group],"a",tblCases[Period],">=2")',
                'A2': '=SUMIFS(tblCases[Result],tblCases[Group],"missing")',
                'A3': '=SUMIFS(tblCases[Result],tblCases[Group],"b")',
                'A4': '=IFERROR(A3,99)',
                'A5': '=COUNTIFS(tblCases[Group],"a",tblCases[Period],">=2")',
            }
            for cell, formula in formulas.items(): summary[cell] = formula
            book.save(path)
            decisions = {'sourceSha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'features': [{
                'featureId':'cases','sheet':'Renamed','table':'tblCases','role':'input','reviewed':True,
                'keyColumns':['Group','Amount','Rate','Period'],
                'literalColumns':['Group','Amount','Rate','Period'],
                'columnTypes':{'Group':'text','Amount':'number','Rate':'number','Period':'number'},
                'reviewedBlankColumns':['Amount']}]}
            # A blank business field must not be used as part of a natural key.
            decisions['features'][0]['keyColumns'] = ['Group','Period','Rate']
            # Distinguish the two a/period2 rows without changing their amount.
            data['C4']=3; book.save(path); decisions['sourceSha256']=hashlib.sha256(path.read_bytes()).hexdigest()
            inputs = build_source_input_plan(path, decisions, connection_id='connection', folder_id='folder')
            native = build_exact_native_calculation_plan(path, inputs)
            rows = [{**row, 'Row ID':f'physical-{i}'} for i,row in enumerate(inputs['sharedInputArchitecture']['inputs']['cases']['rows'])]
            assigned = {'cases':{'elementId':'observed-cases','name':'tblCases','rowIdColumn':'Row ID',
                                 'rows':{str(i):row['Row ID'] for i,row in enumerate(rows)}}}
            lower = lower_exact_native_calculations(native, assigned)
            graph = NativeGraphReference({'observed-cases':('tblCases',rows)}).materialize(lower['shards'])
            for cell, number, error in [('A1',7,''),('A2',0,''),('A3',0,'#DIV/0!'),('A4',99,''),('A5',2,'')]:
                with self.subTest(cell=cell):
                    result=graph.output(lower,'Outputs!'+cell)
                    self.assertEqual(result['error'],error)
                    if not error:
                        self.assertEqual(result['kind'],'n');self.assertEqual(result['number'],number)
