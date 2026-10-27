#!/usr/bin/env python3
"""Fixes from the independent review of the table-report presentation: lookup KPIs, the emitted-formula gate,
accounting / scaling formats, ledger chart periods, chart-row layout and label edge cases."""
from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl.chart import BarChart, LineChart, Reference

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "scripts" / "lib"))
sys.path.insert(0, str(ROOT / "tests"))

import excel_formats as ef  # noqa: E402
import presentation_style as ps  # noqa: E402
import table_report_fixture as fx  # noqa: E402
import table_report_model as trm  # noqa: E402
import table_report_presentation as pres  # noqa: E402
from test_table_report_presentation import (Emitted, clustered_bars, doughnut, draft, element,  # noqa: E402
                                            layout_tree, placements, stacked_columns, wb_spec)

PIPELINE = ROOT / "tests" / "fixtures" / "broader-eight" / "04 Sales Pipeline Forecast.xlsx"


def num(string, **extra):
    return {"kind": "number", "formatString": string, **extra}


# --------------------------------------------------------------------------- 1. lookup KPIs
def build_budget(path, regions=("East", "West", "North"), extra="South", duplicate=False):
    """Deals summarised by region next to a budget LOOKUP; the budget table may list an extra region."""
    deals = [(r, 100.0 * (i + 1)) for i, r in enumerate(regions)] + [(regions[0], 50.0)]
    budget = {r: 1000.0 * (i + 1) for i, r in enumerate(regions)}
    if extra:
        budget[extra] = 99999.0
    book = fx._Book()
    book.table(book.sheet("Deals"), "tblDeals", ["Region", "Amount"], [list(x) for x in deals])
    book.table(book.sheet("Budget"), "tblBudget", ["Region", "Budget"], [[k, v] for k, v in budget.items()])
    sheet = book.sheet("Summary")
    sheet["A1"] = "Pipeline vs budget"
    for col, head in enumerate(["Region", "Total Amount", "Budget", "Variance"], 1):
        sheet.cell(3, col, head)
    for row, region in enumerate(regions, 4):
        sheet.cell(row, 1, region)
        amount = sum(a for k, a in deals if k == region)
        book.put(sheet, f"B{row}", f"=SUMIFS(tblDeals[Amount],tblDeals[Region],$A{row})", amount)
        book.put(sheet, f"C{row}", f"=INDEX(tblBudget[Budget],MATCH($A{row},tblBudget[Region],0))", budget[region])
        book.put(sheet, f"D{row}", f"=B{row}-C{row}", amount - budget[region])
    return book.save(path)


class TestLookupKpis(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = build_budget(Path(self.tmp.name) / "budget.xlsx")
        self.plan, self.info = draft(self.path)

    def test_a_lookup_kpi_sums_only_the_frames_labels(self):
        budget = next(k for k in self.plan["presentation"]["kpis"] if k["label"] == "Budget")
        self.assertEqual(budget["expected"], 6000.0)                       # not 6000 + the extra South row
        emitted = Emitted(self.plan)
        self.assertAlmostEqual(emitted.kpi(budget), 6000.0)
        formula = element(wb_spec(self.plan), budget["id"])["columns"][0]["formula"]
        self.assertIn("SumIf([tblBudget/Budget]", formula)
        self.assertIn('[tblBudget/Region] = "East"', formula)

    def test_the_lookup_restriction_is_always_on_even_without_extra_keys(self):
        path = build_budget(Path(self.tmp.name) / "nobody.xlsx", extra=None)
        plan, _ = draft(path)
        budget = next(k for k in plan["presentation"]["kpis"] if k["label"] == "Budget")
        formula = element(wb_spec(plan), budget["id"])["columns"][0]["formula"]
        self.assertIn('[tblBudget/Region] = "East"', formula)

    def test_a_lookup_table_with_duplicate_keys_drops_the_kpi_and_says_why(self):
        plan = copy.deepcopy(self.plan)
        table = next(t for t in plan["tables"] if t["name"] == "tblBudget")
        table["rows"].append(["East", 777.0])                              # a second East: the sum would double
        model = trm.compute(plan)
        notes = pres.drop_unexpressible(plan, model)
        self.assertNotIn("Budget", [k["label"] for k in plan["presentation"]["kpis"]])
        self.assertTrue(any("tblBudget[Region]" in n["message"] and "duplicate" in n["message"] for n in notes))

    def test_the_gate_evaluates_the_emitted_formulas_and_drops_what_does_not_match(self):
        model = trm.compute(self.plan)
        lines, problems, unverified = pres.check_emitted(self.plan, model)
        self.assertEqual((problems, unverified), ([], []))
        self.assertTrue(any(line.startswith("EMITTED PASS KPI Budget") for line in lines))
        with mock.patch.object(trm, "restrict_aggs", lambda node, member: node), \
                mock.patch.object(trm, "member_ir", lambda by, labels, key_type="text": ["b", True]):
            broken = copy.deepcopy(self.plan)
            table = next(t for t in broken["tables"] if t["name"] == "tblDeals")
            table["rows"].append(["South", 5.0])
            broken["presentation"]["restrict"] = {f["id"]: f["labels"] for f in broken["frames"]}
            model = trm.compute(broken)
            _, problems, _ = pres.check_emitted(broken, model)
            self.assertTrue(any("Total Amount" in p for p in problems), problems)
            notes = pres.drop_unexpressible(broken, model)
            self.assertNotIn("Total Amount", [k["label"] for k in broken["presentation"]["kpis"]])
            self.assertTrue(any(n["code"] == "PRESENTATION_ELEMENT_DROPPED" for n in notes))

    def test_a_failed_kpi_with_a_colon_does_not_drop_its_shorter_peer(self):
        plan = copy.deepcopy(self.plan)
        short = next(k for k in plan["presentation"]["kpis"] if k["label"] == "Budget")
        colon = copy.deepcopy(short)
        colon.update(id=f"{short['id']}_actual", label="Budget: Actual")
        plan["presentation"]["kpis"].append(colon)
        model = trm.compute(plan)
        problem = "KPI Budget: Actual: emitted 1 vs oracle 2"
        with mock.patch.object(pres, "check_emitted", return_value=([], [problem], [])):
            pres.drop_unexpressible(plan, model)
        labels = [k["label"] for k in plan["presentation"]["kpis"]]
        self.assertIn("Budget", labels)
        self.assertNotIn("Budget: Actual", labels)

    def test_a_draft_carries_a_drop_as_a_review_issue(self):
        path = build_budget(Path(self.tmp.name) / "braces.xlsx", regions=("East", "We}}st", "North"))
        plan, info = draft(path)
        self.assertNotIn("Variance", [k["label"] for k in plan["presentation"]["kpis"]])   # two tables: a text card
        issue = next(i for i in plan["draftReviewIssues"] if i["code"] == "PRESENTATION_ELEMENT_DROPPED")
        self.assertEqual(issue["severity"], "review")
        self.assertTrue(any("Variance" in e and "}}" in e for e in issue["evidence"]))
        self.assertIn("PRESENTATION_ELEMENT_DROPPED", info["reviewIssues"])
        self.assertIn("Total Amount", [k["label"] for k in plan["presentation"]["kpis"]])  # single-source: kept

    def test_the_build_gate_fails_a_plan_whose_emitted_formulas_disagree_with_the_oracle(self):
        import contextlib
        import io
        import json
        from test_table_report_presentation import BUILD_CLI, DRAFT_CLI, run_cli
        plan_path = Path(self.tmp.name) / "plan.json"
        code, out, err = run_cli(DRAFT_CLI, [str(PIPELINE), "--connection", "c", "--folder", "f",
                                             "--out", str(plan_path)])
        self.assertEqual(code, 0, err)
        plan = json.loads(plan_path.read_text())
        frame = next(f for f in plan["frames"] if f["title"] == "H2 FY26 Weighted Forecast")
        plan["presentation"]["restrict"] = {frame["id"]: ["Jordan Lee"]}   # the emitted KPIs would cover one owner
        plan_path.write_text(json.dumps(plan))
        code, out, err = run_cli(BUILD_CLI, [str(plan_path), "--source", str(PIPELINE),
                                             "--out-dir", str(Path(self.tmp.name) / "build")])
        self.assertEqual(code, 1)
        self.assertIn("EMITTED PRESENTATION FAILED", err)
        self.assertFalse((Path(self.tmp.name) / "build" / "wb_spec.json").exists())


# --------------------------------------------------------------------------- 2/3/6. formats
class TestAccountingAndScaling(unittest.TestCase):
    def test_openpyxl_accounting_builtins_translate_cleanly(self):
        from openpyxl.styles.numbers import BUILTIN_FORMATS
        want = {41: num("(,.0f"), 42: num("($,.0f"), 43: num("(,.2f"), 44: num("($,.2f")}
        for number, fmt in want.items():
            result = ef.translate(BUILTIN_FORMATS[number])
            self.assertEqual(result.format, fmt, number)
            self.assertNotIn("suffix", result.format)
            self.assertNotIn("prefix", result.format)

    def test_escaped_parentheses_are_negative_parentheses(self):
        self.assertEqual(ef.translate(r"#,##0.00_);\(#,##0.00\)").format, num("(,.2f"))
        self.assertEqual(ef.translate(r"_(* #,##0_);_(* \(#,##0\);_(* ""-""_);_(@_)").format, num("(,.0f"))

    def test_a_prefix_or_suffix_never_carries_format_tokens(self):
        for code in ('0.00" "@', '@" sq ft"', '#,##0 @', '#,##0 0'):
            result = ef.translate(code)
            for part in (result.format or {}).get("prefix", ""), (result.format or {}).get("suffix", ""):
                self.assertNotRegex(part, r"[@0-9#?]", code)

    def test_scaling_commas_show_the_full_number_without_the_unit(self):
        for code in ('#,##0,"K"', '#,##0.0,,"M"', '[>=1000000]0.0,,"M";[>=1000]0.0,"K";0'):
            result = ef.translate(code)
            self.assertNotIn("suffix", result.format, code)
            self.assertEqual(result.format["formatString"].count("."), 1)
            self.assertTrue(result.format["formatString"].endswith(".0f"), code)    # the full number, no decimals
            self.assertTrue(any("scal" in note for note in result.approximated), code)
            shown = f"{1234567:,.0f}" + result.format.get("suffix", "")
            self.assertEqual(shown, "1,234,567")                                    # never "1,234,567K"

    def test_a_red_single_section_and_other_font_colors_are_noted(self):
        red = ef.translate("[Red]#,##0")
        self.assertFalse(red.negative_red)
        self.assertTrue(any("[Red]" in n for n in red.approximated))
        blue = ef.translate("[Blue]#,##0;[Red]-#,##0")
        self.assertTrue(blue.negative_red)
        self.assertTrue(any("[Blue]" in n for n in blue.approximated))
        color = ef.translate("[Color10]#,##0;[Color3]-#,##0")
        self.assertTrue(any("Color" in n for n in color.approximated))
        self.assertEqual(ef.translate("#,##0;[Red]-#,##0").approximated, [])

    def test_a_currency_after_the_number_is_a_suffix(self):
        self.assertEqual(ef.translate("#,##0.00 [$€-407]").format, num(",.2f", suffix=" €"))
        self.assertEqual(ef.translate('#,##0.00\\ "€"').format, num(",.2f", suffix=" €"))
        self.assertEqual(ef.translate("[$€-407] #,##0.00").format, num(",.2f", prefix="€ "))
        self.assertEqual(ef.translate("€#,##0").format, num(",.0f", prefix="€"))

    def test_yyy_is_a_four_digit_year_like_excel(self):
        self.assertEqual(ef.translate("yyy").format["formatString"], "%Y")
        self.assertEqual(ef.translate("yyy-mm-dd").format["formatString"], "%Y-%m-%d")
        self.assertEqual(ef.translate("yy").format["formatString"], "%y")

    def test_fractional_seconds_are_a_time_format_not_a_number(self):
        result = ef.translate("h:mm:ss.000")
        self.assertEqual(result.category, "date")
        self.assertEqual(result.format["formatString"], "%-H:%M:%S")
        self.assertTrue(any("fractional" in n for n in result.approximated))


class TestKpiTextCardLiterals(unittest.TestCase):
    def test_a_text_card_keeps_the_prefix_and_suffix(self):
        _, text = ps.kpi_text_card("k", "REVENUE", "Sum([T/x])", num(",.0f", prefix="€", suffix=" k"))
        value = text["body"].split("\n\n")[1]
        self.assertEqual(value, '<span style="font-size: 28px">**€{{Sum([T/x]) | ,.0f}} k**</span>')
        _, plain = ps.kpi_text_card("k", "REVENUE", "Sum([T/x])", num(".1%"))
        self.assertIn("**{{Sum([T/x]) | .1%}}**", plain["body"])

    def test_a_label_with_template_delimiters_skips_that_text_card_kpi(self):
        plan, _ = draft(PIPELINE)
        plan = copy.deepcopy(plan)
        pipeline = next(t for t in plan["tables"] if t["name"] == "tblPipeline")
        frame = next(f for f in plan["frames"] if f["title"] == "H2 FY26 Weighted Forecast")
        frame["labels"][0] = "Jordan }} Lee"
        pipeline["rows"] = [[*r[:2], "Jordan }} Lee" if r[2] == "Jordan Lee" else r[2], *r[3:]]
                            for r in pipeline["rows"]]
        for t in plan["tables"]:
            if t["name"].endswith("Inputs"):
                t["rows"] = [["Jordan }} Lee" if r[0] == "Jordan Lee" else r[0], r[1]] for r in t["rows"]]
        plan["presentation"]["restrict"] = {frame["id"]: frame["labels"]}
        model = trm.compute(plan)
        notes = pres.drop_unexpressible(plan, model)
        self.assertNotIn("Coverage", [k["label"] for k in plan["presentation"]["kpis"]])
        self.assertTrue(any("Coverage" in n["message"] and "}}" in n["message"] for n in notes))


# --------------------------------------------------------------------------- label edge cases
class TestLabelFormulas(unittest.TestCase):
    def test_non_ascii_labels_are_literal_utf8_in_this_prs_formulas(self):
        formula = trm.order_formula("[k]", ["Zürich", "São Paulo"], "text")
        self.assertIn('"Zürich"', formula)
        self.assertNotIn("\\u", formula)
        emitter = trm.Emitter(col=lambda n: f"[T/{n}]", utf8_strings=True)
        self.assertEqual(emitter.emit(["s", "Zürich"]), '"Zürich"')
        self.assertEqual(trm.Emitter(col=lambda n: n).emit(["s", "Zürich"]), '"Z\\u00fcrich"')   # unchanged

    def test_numeric_labels_are_quoted_when_the_key_column_is_text(self):
        self.assertEqual(trm.order_formula("[k]", [10, 2], "text"),
                         'If([k] = "10", 1, If([k] = "2", 2, 3))')
        self.assertEqual(trm.order_formula("[k]", [10, 2], "number"), "If([k] = 10, 1, If([k] = 2, 2, 3))")
        self.assertEqual(trm.member_ir("k", [10], "text"), ["op", "=", ["c", "k"], ["s", "10"]])
        self.assertEqual(trm.member_ir("k", [10], "number"), ["op", "=", ["c", "k"], ["n", 10]])

    def test_source_order_is_capped_and_a_note_says_so(self):
        self.assertEqual(trm.MAX_ORDER_LABELS, 40)
        frame = {"kind": "summary", "labels": [f"r{i}" for i in range(41)]}
        self.assertIsNone(trm.order_labels(frame))
        self.assertIsNotNone(trm.order_labels({"kind": "summary", "labels": ["a"] * 40}))


class TestOrderCapNote(unittest.TestCase):
    def test_a_grid_over_the_cap_raises_the_review_issue(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "many.xlsx"
            regions = [f"R{i:02d}" for i in range(41)]
            build_budget(path, regions=regions, extra=None)
            plan, info = draft(path)
            self.assertIn("SOURCE_ORDER_OMITTED", info["reviewIssues"])
            issue = next(i for i in plan["draftReviewIssues"] if i["code"] == "SOURCE_ORDER_OMITTED")
            self.assertIn("41 rows", issue["evidence"][0])
            summary = element(wb_spec(plan), "f_s_pipeline_vs_budget")
            self.assertNotIn("Source Order", [c["name"] for c in summary["columns"]])


# --------------------------------------------------------------------------- 4. ledger chart periods
class TestLedgerChartPeriods(unittest.TestCase):
    def draft_with_chart(self, last_col):
        original = fx._Book.save

        def save(book, path):
            ws = book.wb["Cash"]
            chart = LineChart()
            chart.title = "Collections"
            chart.add_data(Reference(ws, min_col=2, max_col=last_col, min_row=7, max_row=7), from_rows=True)
            chart.set_categories(Reference(ws, min_col=2, max_col=last_col, min_row=4, max_row=4))
            ws.add_chart(chart, "Z2")
            return original(book, path)
        self.addCleanup(setattr, fx._Book, "save", original)
        fx._Book.save = save
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "cash.xlsx"
        fx.build_cash(path)
        return draft(path)[0]

    def test_a_chart_over_the_first_three_weeks_filters_the_ledger_to_those_weeks(self):
        plan = self.draft_with_chart(4)
        chart = plan["presentation"]["charts"][0]
        self.assertEqual(len(chart["categories"]), 3)
        el = element(wb_spec(plan), chart["id"])
        self.assertEqual(len(el["filters"]), 1)
        flt = el["filters"][0]
        self.assertEqual((flt["kind"], flt["mode"], flt["values"]), ("list", "include", [1, 2, 3]))
        column = next(c for c in el["columns"] if c["id"] == flt["columnId"])
        axis = next(f for f in plan["frames"] if f["kind"] == "ledger")["axis"]
        self.assertTrue(column["formula"].endswith(f"/{axis}]"))
        self.assertTrue(column.get("hidden"))

    def test_the_period_column_honors_the_hidden_column_surface(self):
        plan = self.draft_with_chart(4)
        self.addCleanup(ps.SURFACES.update, dict(ps.SURFACES))
        ps.SURFACES["chart_hidden_column"] = False
        el = element(wb_spec(plan), plan["presentation"]["charts"][0]["id"])
        flt = el["filters"][0]
        column = next(c for c in el["columns"] if c["id"] == flt["columnId"])
        self.assertNotIn("hidden", column)

    def test_a_failed_chart_with_a_colon_does_not_drop_its_shorter_peer(self):
        plan = self.draft_with_chart(4)
        short = plan["presentation"]["charts"][0]
        colon = copy.deepcopy(short)
        colon.update(id=f"{short['id']}_first", title="Collections: First")
        plan["presentation"]["charts"].append(colon)
        model = trm.compute(plan)
        problem = "chart Collections: First: emitted series do not match"
        with mock.patch.object(pres, "check_emitted", return_value=([], [problem], [])):
            pres.drop_unexpressible(plan, model)
        titles = [chart["title"] for chart in plan["presentation"]["charts"]]
        self.assertIn("Collections", titles)
        self.assertNotIn("Collections: First", titles)

    def test_a_chart_over_every_period_needs_no_filter(self):
        plan = self.draft_with_chart(6)
        el = element(wb_spec(plan), plan["presentation"]["charts"][0]["id"])
        self.assertNotIn("filters", el)

    def test_without_the_list_filter_surface_the_subset_chart_is_unmapped(self):
        self.addCleanup(ps.SURFACES.update, dict(ps.SURFACES))
        ps.SURFACES["chart_list_filter"] = False
        plan = self.draft_with_chart(4)
        self.assertEqual(plan["presentation"]["charts"], [])
        self.assertTrue(any("Collections" in m for m in plan["unmappedCharts"]))


# --------------------------------------------------------------------------- 5. chart rows
def many_charts(ws):
    stacked_columns(ws)
    clustered_bars(ws)
    doughnut(ws)
    for number, col in ((1, 5), (2, 7)):
        chart = BarChart()
        chart.title = f"Extra {number}"
        chart.add_data(Reference(ws, min_col=col, min_row=3, max_row=6), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=6))
        ws.add_chart(chart, f"Z{number * 20}")


class TestChartRows(unittest.TestCase):
    def test_five_charts_form_aligned_rows_without_overlap(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "orders.xlsx"
            fx.build_orders(path, charts=many_charts)
            plan, _ = draft(path)
            self.assertEqual(len(plan["presentation"]["charts"]), 5)
            spec = wb_spec(plan)
            nodes = placements(spec)
            ids = {n[0]: n for n in nodes}
            charts = [ids[c["id"]] for c in plan["presentation"]["charts"] if c["id"] in ids]
            self.assertEqual(len(charts), 5)
            for a in nodes:
                for b in nodes:
                    if a is not b:
                        self.assertFalse(a[1] < b[2] and b[1] < a[2] and a[3] < b[4] and b[3] < a[4], (a, b))
            loose = [n for n in charts if n[1] in (1, 13) and n[2] - n[1] == 12]
            self.assertEqual(len(loose), 2)                           # one pair of half-width charts
            self.assertEqual((loose[0][3], loose[0][4]), (loose[1][3], loose[1][4]))     # same top, same bottom
            self.assertEqual(loose[0][4] - loose[0][3], ps.CHART_ROWS)
            tops = sorted({n[3] for n in charts})
            self.assertEqual(len(tops), len(set(tops)))
            hero = min(charts, key=lambda n: n[3])
            self.assertEqual((hero[1], hero[2]), (1, 25))
            self.assertEqual(loose[0][3], hero[4])                   # the pair starts where the hero ends


# --------------------------------------------------------------------------- merged summaries (bindings)
class TestMergedSummaryBindings(unittest.TestCase):
    """A summary frame merged into a keyed host table (08) is shown by the HOST's element: the chart over the
    summary's labels still pairs with it and the summary's red negatives still reach it."""

    def setUp(self):
        from test_table_report_presentation import SOURCES
        plan, _ = draft(SOURCES["08"])
        self.plan = copy.deepcopy(plan)
        self.summary = next(f for f in self.plan["frames"] if f["id"] == "s_projection_summary")
        self.second = copy.deepcopy(self.plan["presentation"]["charts"][0])
        self.second.update(id="ch_second", title="Minimum projected (again)")
        self.plan["presentation"]["charts"].append(self.second)           # not the hero: it pairs with a table
        next(c for c in self.summary["columns"] if c["name"] == "Min Proj")["negativeRed"] = True

    def test_the_host_table_is_bound_to_the_merged_summary(self):
        parts = trm.build_parts(self.plan, "<dm>", trm.local_dm_elements(self.plan))
        host = next(m for m in parts["visibleMeta"] if m["id"] == "f_k_projection")
        self.assertEqual(host["frame"], "k_projection")
        self.assertEqual(host["bindings"], ["s_projection_summary"])

    def test_a_chart_over_the_merged_summary_pairs_with_its_host_table(self):
        spec = wb_spec(self.plan)
        nodes = {n[0]: n for n in placements(spec)}
        host, chart = nodes["f_k_projection"], nodes["ch_second"]
        self.assertEqual((host[1], host[2], chart[1], chart[2]), (1, 15, 15, 25))     # table | its chart
        self.assertEqual(host[3:], chart[3:])                                         # same grid rows

    def test_the_merged_summarys_red_negatives_reach_the_host_table(self):
        host = element(wb_spec(self.plan), "f_k_projection")
        ids = {c["name"]: c["id"] for c in host["columns"]}
        rules = host["conditionalFormats"]
        self.assertEqual(len(rules), 1)
        self.assertEqual(rules[0]["columnIds"], [ids["Min Proj"]])

    def test_the_host_still_leads_when_its_summary_is_the_main_frame(self):
        self.assertEqual(self.plan["presentation"]["main"], "s_projection_summary")
        order = [n[0] for n in placements(wb_spec(self.plan))]
        self.assertLess(order.index("f_k_projection"), order.index("pv_projection"))


if __name__ == "__main__":
    unittest.main()
