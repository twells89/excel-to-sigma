"""Synthetic cell-drawn P&L workbook (SUMIFS actuals + driver growth + flat + manual).

openpyxl writes formulas without cached values, but the drafter reads Excel's
cached results. ``build_pnl`` therefore computes every formula cell's value in
Python while writing it, saves the workbook, then injects those values as ``<v>``
elements into the saved sheet XML. The result is a normal workbook whose cached
values are exactly what Excel would have stored, so the tests exercise the real
``load_workbook(data_only=True)`` path with no mocking.
"""
from __future__ import annotations

import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

MONTHS = ["Jan-26", "Feb-26", "Mar-26", "Apr-26", "May-26", "Jun-26"]
N_ACT = 3          # Jan-Mar actual, Apr-Jun forecast
HEADER_ROW = 3

# (section, label, kind, parameter, actual amounts per actual month)
DEFAULT_LINES = [
    ("Revenue", "Subscription", "growth", 0.05, [1000.0, 1100.0, 1200.0]),
    ("Revenue", "Services", "manual", [300.0, 320.0, 340.0], [250.0, 260.0, 270.0]),
    ("Cost of Revenue", "Hosting", "flat", None, [200.0, 210.0, 220.0]),
    ("Cost of Revenue", "Support", "growth", 0.02, [100.0, 105.0, 110.0]),
    ("Operating Expenses", "R&D", "manual", [400.0, 410.0, 420.0], [380.0, 385.0, 390.0]),
    ("Below the Line", "Interest", "flat", None, [20.0, 20.0, 21.0]),
]
SECTION_ORDER = ["Revenue", "Cost of Revenue", "Operating Expenses", "Below the Line"]


def _inject_cached_values(path: Path, cache: dict[tuple[str, str], float], sheet_names: list[str]):
    temp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
            if match:
                sheet = sheet_names[int(match.group(1)) - 1]

                def fill(found):
                    value = cache.get((sheet, found.group(2)))
                    if value is None:
                        return found.group(0)
                    return f"{found.group(1)}{found.group(3)}<v>{value!r}</v>"

                data = re.sub(r'(<c r="([A-Z]+\d+)"[^>]*>)(<f>.*?</f>)<v />',
                              fill, data.decode("utf-8")).encode("utf-8")
            target.writestr(item, data)
    shutil.move(temp, path)


def build_pnl(path, *, lines=None, mutate=None, cached_overrides=None, gl_sheet="GL"):
    """Write the fixture to ``path``; returns the expected values for assertions.

    ``mutate(pnl, gl, drivers)`` edits cells after the clean layout is written (cached
    values stay as computed for the clean layout). ``cached_overrides`` replaces cached
    values by ``(sheet, coord)`` to simulate a workbook whose caches disagree.
    """
    path = Path(path)
    lines = lines or DEFAULT_LINES
    wb = Workbook()
    pnl = wb.active
    pnl.title = "P&L"
    gl = wb.create_sheet(gl_sheet)
    drivers = wb.create_sheet("Drivers")
    cache: dict[tuple[str, str], float] = {}

    def put(ws, coord, formula, value):
        ws[coord] = formula
        cache[(ws.title, coord)] = float(value)

    # ---- drivers
    drivers["A1"] = "Forecast drivers"
    drivers["A3"] = "Driver"
    drivers["B3"] = "Monthly Rate"
    drivers["C3"] = "Basis"
    driver_row = {}
    for label, kind, param in ((l[1], l[2], l[3]) for l in lines):
        if kind == "growth":
            row = 4 + len(driver_row)
            driver_row[label] = row
            drivers.cell(row, 1, f"{label} growth")
            drivers.cell(row, 2, param)
            drivers.cell(row, 3, f"{label} basis")

    # ---- GL: two postings per actual line-month
    gl["A1"] = "GL detail"
    for col, head in enumerate(["Txn Date", "Month", "Line Item", "Department", "Memo / Vendor", "Amount"], 1):
        gl.cell(4, col, head)
    row = 5
    actual = {}
    for section, label, kind, param, amounts in lines:
        for month_index, amount in enumerate(amounts):
            half = round(amount / 2, 2)
            for part, value in enumerate((half, amount - half)):
                gl.cell(row, 1, datetime(2026, month_index + 1, 5 + part))
                gl.cell(row, 2, MONTHS[month_index])
                gl.cell(row, 3, label)
                gl.cell(row, 4, section)
                gl.cell(row, 5, f"vendor {part}")
                gl.cell(row, 6, value)
                row += 1
            actual[(label, month_index)] = amount
    gl_last = row - 1

    # ---- P&L
    pnl["A1"] = "Synthetic Co - rolling forecast"
    for index, month in enumerate(MONTHS):
        pnl.cell(HEADER_ROW, 2 + index, month)
    fy_col = 2 + len(MONTHS)
    pnl.cell(HEADER_ROW, fy_col, "FY2026")
    values = {}
    totals = {}
    row = HEADER_ROW + 1
    section_rows = {}
    for section in SECTION_ORDER:
        members = [l for l in lines if l[0] == section]
        if not members:
            continue
        pnl.cell(row, 1, section)
        row += 1
        first = row
        for _, label, kind, param, amounts in members:
            pnl.cell(row, 1, label)
            series = []
            for month_index in range(len(MONTHS)):
                col = get_column_letter(2 + month_index)
                coord = f"{col}{row}"
                if month_index < N_ACT:
                    formula = (f"=SUMIFS({gl_sheet}!$F$5:$F${gl_last},{gl_sheet}!$C$5:$C${gl_last},"
                               f"$A{row},{gl_sheet}!$B$5:$B${gl_last},{col}${HEADER_ROW})")
                    value = actual[(label, month_index)]
                    put(pnl, coord, formula, value)
                else:
                    prev = get_column_letter(1 + month_index)
                    if kind == "growth":
                        value = series[-1] * (1 + param)
                        put(pnl, coord, f"={prev}{row}*(1+Drivers!$B${driver_row[label]})", value)
                    elif kind == "flat":
                        value = series[-1]
                        put(pnl, coord, f"={prev}{row}", value)
                    else:
                        value = param[month_index - N_ACT]
                        pnl[coord] = value
                series.append(value)
            put(pnl, f"{get_column_letter(fy_col)}{row}",
                f"=SUM(B{row}:{get_column_letter(fy_col - 1)}{row})", sum(series))
            values[label] = series
            row += 1
        last = row - 1
        total = [sum(values[l[1]][m] for l in members) for m in range(len(MONTHS))]
        pnl.cell(row, 1, f"Total {section}")
        for month_index in range(len(MONTHS)):
            col = get_column_letter(2 + month_index)
            put(pnl, f"{col}{row}", f"=SUM({col}{first}:{col}{last})", total[month_index])
        put(pnl, f"{get_column_letter(fy_col)}{row}", f"=SUM(B{row}:{get_column_letter(fy_col - 1)}{row})", sum(total))
        totals[section] = (row, total)
        row += 2

    def derived(label, expr_rows, signs, series):
        nonlocal row
        pnl.cell(row, 1, label)
        for month_index in range(len(MONTHS)):
            col = get_column_letter(2 + month_index)
            formula = "=" + "".join(f"{s}{col}{r}" for s, r in zip(signs, expr_rows)
                                    if r is not None).lstrip("+")
            put(pnl, f"{col}{row}", formula, series[month_index])
        put(pnl, f"{get_column_letter(fy_col)}{row}", f"=SUM(B{row}:{get_column_letter(fy_col - 1)}{row})", sum(series))
        row_here = row
        row += 1
        return row_here

    zero = [0.0] * len(MONTHS)
    rev = totals["Revenue"]
    cogs = totals.get("Cost of Revenue", (None, zero))
    opex = totals.get("Operating Expenses", (None, zero))
    below = totals.get("Below the Line", (None, zero))
    gp = [r - c for r, c in zip(rev[1], cogs[1])]
    gp_row = derived("Gross Profit", [rev[0], cogs[0]], ["+", "-"], gp)
    ebitda = [g - o for g, o in zip(gp, opex[1])]
    ebitda_row = derived("EBITDA", [gp_row, opex[0]], ["+", "-"], ebitda)
    net = [e - b for e, b in zip(ebitda, below[1])]
    derived("Net Income", [ebitda_row, below[0]], ["+", "-"], net)
    expected = {"Revenue": sum(rev[1]), "Gross Profit": sum(gp), "EBITDA": sum(ebitda),
                "Net Income": sum(net)}

    if mutate:
        mutate(pnl, gl, drivers)
    for key, value in (cached_overrides or {}).items():
        cache[key] = float(value)
    wb.save(path)
    _inject_cached_values(path, cache, wb.sheetnames)
    return {"targets": expected, "values": values, "gl_last": gl_last, "driver_row": driver_row}


# ---------------------------------------------------------------------------
# Quarterly budget workbook: discovered sections, quarter labels, year-over-year,
# percent-of-total, alternative growth syntax, an external-link cell, a zero actual
# line, ratio rows and an FY column that covers only some periods.
# ---------------------------------------------------------------------------
QUARTERS = ["Q1 FY25", "Q2 FY25", "Q3 FY25", "Q4 FY25", "Q1 FY26 Bud", "Q2 FY26 Bud"]
Q_ACTUAL = 4
Q_FY_FIRST = 2   # the FY column sums the last four periods (index 2..5)

# (section, [(label, kind, parameter, four actual amounts)])
#   kinds: yoy(rate) | growth_alt(rate) | pct(rate, of section) | flat | manual(list of 2)
#          | external(list of 2) | zero(list of 2 forecast values via literal)
QUARTERLY_SECTIONS = [
    ("Net Sales", [
        ("In-Store", "yoy", 0.04, [1000.0, 1100.0, 1200.0, 1300.0]),
        ("Online", "growth_alt", 0.10, [400.0, 420.0, 450.0, 480.0]),
    ]),
    ("COGS", [
        ("Merchandise", "pct", (0.6, "Net Sales"), [900.0, 950.0, 1000.0, 1050.0]),
        ("Shrink", "pct", (0.02, "Net Sales"), [30.0, 31.0, 33.0, 35.0]),
    ]),
    ("SG&A", [
        ("Payroll", "growth_alt", 0.02, [200.0, 204.0, 208.0, 212.0]),
        ("Marketing", "manual", [90.0, 95.0], [80.0, 85.0, 88.0, 92.0]),
        ("Allocation", "external", [61.0, 64.0], [0.0, 0.0, 0.0, 0.0]),
    ]),
    ("Other Income/Expense", [
        ("Interest", "manual", [-5.0, -5.0], [-8.0, -8.0, -7.0, -6.0]),
        ("Depreciation", "flat", None, [45.0, 45.0, 45.0, 45.0]),
    ]),
]
QUARTERLY_DERIVED = [
    ("Gross Margin", "{Net Sales}-{COGS}"),
    ("Gross Margin %", "IF({Net Sales}=0,0,{Gross Margin}/{Net Sales})"),
    ("Operating Profit", "{Gross Margin}-{SG&A}"),
    ("Pre-tax Income", "{Operating Profit}-{Other Income/Expense}"),
]


def build_quarterly(path, *, sections=None, derived=None, mutate=None, cached_overrides=None,
                    labels=None):
    """Write the quarterly fixture; returns ``{"targets": {...}, "values": {...}}``.

    ``targets`` are the FY-column values by row label (what the drafter should emit).
    """
    path = Path(path)
    sections = sections or QUARTERLY_SECTIONS
    derived = derived if derived is not None else QUARTERLY_DERIVED
    labels = labels or QUARTERS
    n = len(labels)
    wb = Workbook()
    rpt = wb.active
    rpt.title = "Budget"
    gl = wb.create_sheet("GL Detail")
    drivers = wb.create_sheet("Drivers")
    cache = {}

    def put(ws, coord, formula, value):
        ws[coord] = formula
        cache[(ws.title, coord)] = float(value)

    cols = [get_column_letter(2 + i) for i in range(n)]
    fy = get_column_letter(2 + n)

    # drivers: one rate per growth / yoy / pct line
    drivers["A1"] = "Budget drivers"
    drivers["A3"], drivers["B3"], drivers["C3"] = "Driver", "Value", "Notes"
    driver_row = {}
    for _, members in sections:
        for label, kind, param, _ in members:
            if kind in ("yoy", "growth_alt", "pct"):
                row = 4 + len(driver_row)
                driver_row[label] = row
                drivers.cell(row, 1, f"{label} driver")
                drivers.cell(row, 2, param[0] if kind == "pct" else param)

    # GL
    gl["A1"], gl["B1"], gl["C1"], gl["D1"] = "Period", "Store", "Account", "Amount"
    gl_row = 2
    actual = {}
    for _, members in sections:
        for label, kind, param, amounts in members:
            for index, amount in enumerate(amounts):
                if amount == 0:
                    continue
                half = round(amount / 2, 2)
                for part, value in enumerate((half, amount - half)):
                    gl.cell(gl_row, 1, labels[index])
                    gl.cell(gl_row, 2, f"Store {part}")
                    gl.cell(gl_row, 3, label)
                    gl.cell(gl_row, 4, value)
                    gl_row += 1
            actual[label] = list(amounts)

    rpt["A1"] = "Quarterly budget"
    for i, text in enumerate(labels):
        rpt.cell(4, 2 + i, text)
    rpt.cell(4, 2 + n, "FY26 Total")
    rpt.cell(4, 3 + n, "YoY %")           # an unmodelled trailing column
    values, totals, rows = {}, {}, {}
    row = 6
    for name, members in sections:
        rpt.cell(row, 1, name)
        row += 1
        first = row
        for label, kind, param, amounts in members:
            rpt.cell(row, 1, label)
            rows[label] = row
            series = list(amounts)
            for i in range(Q_ACTUAL):
                coord = f"{cols[i]}{row}"
                if amounts[i] == 0 and kind == "external":
                    put(rpt, coord, "=0", 0)
                else:
                    put(rpt, coord,
                        f"=SUMIFS('GL Detail'!$D:$D,'GL Detail'!$C:$C,$A{row},'GL Detail'!$A:$A,{cols[i]}$4)",
                        amounts[i])
            values[label] = series
            row += 1
        last = row - 1
        rpt.cell(row, 1, f"Total {name}")
        totals[name] = row
        rows[f"Total {name}"] = row
        row += 2
    # forecast cells need the section totals of the same column (pct), so go column by column
    drv = lambda label: f"Drivers!$B${driver_row[label]}"
    for i in range(Q_ACTUAL, n):
        col, prev = cols[i], cols[i - 1]
        for _, members in sections:
            for label, kind, param, amounts in members:
                r = rows[label]
                j = i - Q_ACTUAL
                if kind == "yoy":
                    base = values[label][i - 4]
                    value = base * (1 + param)
                    put(rpt, f"{col}{r}", f"={cols[i - 4]}{r}*(1+{drv(label)})", value)
                elif kind == "growth_alt":
                    value = values[label][-1] * (1 + param)
                    put(rpt, f"{col}{r}", f"={prev}{r}+{prev}{r}*{drv(label)}", value)
                elif kind == "flat":
                    value = values[label][-1]
                    put(rpt, f"{col}{r}", f"={prev}{r}", value)
                elif kind == "manual":
                    value = param[j]
                    rpt[f"{col}{r}"] = value
                elif kind == "external":
                    value = param[j]
                    put(rpt, f"{col}{r}", f"='[Alloc.xlsx]Alloc'!${'CD'[j]}$4", value)
                else:
                    continue  # pct lines are filled in below, once the totals exist
                values[label].append(value)
        # percent-of-total lines read the total of the base section in the same column
        for _, members in sections:
            for label, kind, param, amounts in members:
                if kind == "pct":
                    base_total = sum(values[m[0]][i] for s, ms in sections if s == param[1] for m in ms
                                     if m[1] != "pct")
                    value = base_total * param[0]
                    put(rpt, f"{col}{rows[label]}", f"={col}{totals[param[1]]}*{drv(label)}", value)
                    values[label].append(value)
    # totals + derived rows, per period
    sec_values = {}
    for name, members in sections:
        r = totals[name]
        first, last = rows[members[0][0]], rows[members[-1][0]]
        sec_values[name] = [sum(values[m[0]][i] for m in members) for i in range(n)]
        for i in range(n):
            put(rpt, f"{cols[i]}{r}", f"=SUM({cols[i]}{first}:{cols[i]}{last})", sec_values[name][i])
    row = max(totals.values()) + 2
    derived_values = {}
    import re as _re
    for label, template in derived:
        rpt.cell(row, 1, label)
        rows[label] = row
        series = []
        for i in range(n):
            def ref(found, i=i):
                key = found.group(1)
                return f"{cols[i]}{totals[key] if key in totals else rows[key]}"

            def val(found, i=i):
                key = found.group(1)
                return repr(sec_values[key][i] if key in sec_values else derived_values[key][i])
            formula = "=" + _re.sub(r"\{([^}]+)\}", ref, template)
            python = _re.sub(r"(?<![<>])=", "==", _re.sub(r"\{([^}]+)\}", val, template)).replace("IF(", "_if(")
            result = eval(python, {"_if": lambda c, a, b: a if c else b})
            series.append(float(result))
            put(rpt, f"{cols[i]}{row}", formula, series[-1])
        derived_values[label] = series
        row += 1
    # FY column over the last four periods, per line / total / derived row
    fy_cols = f"{cols[Q_FY_FIRST]}{{r}}:{cols[-1]}{{r}}"
    targets = {}
    for label, r in rows.items():
        series = (values.get(label) or derived_values.get(label)
                  or (sec_values[label[6:]] if label.startswith("Total ") else None))
        if label.endswith("%"):
            continue
        put(rpt, f"{fy}{r}", f"=SUM({fy_cols.format(r=r)})", sum(series[Q_FY_FIRST:]))
        if label in derived_values or label.startswith("Total "):
            targets[label] = sum(series[Q_FY_FIRST:])
    if mutate:
        mutate(rpt, gl, drivers)
    for key, value in (cached_overrides or {}).items():
        cache[key] = float(value)
    wb.save(path)
    _inject_cached_values(path, cache, wb.sheetnames)
    return {"targets": targets, "values": values, "sections": sec_values, "derived": derived_values,
            "rows": rows}
