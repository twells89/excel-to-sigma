#!/usr/bin/env python3
from __future__ import annotations

import copy
import datetime as dt
import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / "lib"))

from input_sync import (  # noqa: E402
    InputSyncError,
    build_artifact_records,
    build_operational_input_elements,
    build_shared_input_architecture,
    emit_construction_program,
    emit_workbook_wiring,
    emit_write_programs,
    evaluate_expense_exception,
    evaluate_timing_status,
    plan_seed,
    sigma_input_type,
)
from input_orchestration import gate_input_completion, gate_input_staging  # noqa: E402


def load_script(filename, module_name):
    spec = importlib.util.spec_from_file_location(module_name, SCRIPTS / filename)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def shared_plan(source_inputs=None):
    return {
        "sourceInputs": list(source_inputs or []),
        "sharedInputSources": {
            "salesSettings": {"sourceBinding": "Assumptions!F2",
                              "rows": [{"Setting": "AsOfDate", "AsOfDate": "2026-06-30"}]},
            "expensePolicy": {"sourceRange": "Policy!A1:D7", "rows": [
                {"Category": "Airfare", "Per-item Limit": 1200, "Receipt Threshold": 25,
                 "Policy Owner": "Travel"},
                {"Category": "Lodging", "Per-item Limit": 350, "Receipt Threshold": 25,
                 "Policy Owner": "Travel"},
                {"Category": "Meals", "Per-item Limit": 90, "Receipt Threshold": 25,
                 "Policy Owner": "Finance"},
                {"Category": "Ground Transport", "Per-item Limit": 180,
                 "Receipt Threshold": 25, "Policy Owner": "Finance"},
                {"Category": "Software", "Per-item Limit": 500, "Receipt Threshold": 0,
                 "Policy Owner": "Finance"},
                {"Category": "Supplies", "Per-item Limit": 250, "Receipt Threshold": 25,
                 "Policy Owner": "Finance"},
            ]},
        },
    }


class TestSeedPlanner(unittest.TestCase):
    def test_rerun_inserts_nothing_and_partial_retry_reconciles_server_rows(self):
        source = [{"Business ID": "A", "Amount": 10}, {"Business ID": "B", "Amount": 0}]
        first = plan_seed(source, [{"Business ID": "A", "Amount": 10, "Row ID": "r-a"}],
                          ["Business ID"], "Row ID", snapshot={"id": "q1", "ageSeconds": 2})
        self.assertEqual([row["source"]["Business ID"] for row in first["inserts"]], ["B"])
        self.assertEqual(first["updates"], [])

        current = [{"Business ID": "A", "Amount": 10, "Row ID": "r-a"},
                   {"Business ID": "B", "Amount": 0, "Row ID": "r-b"}]
        rerun = plan_seed(source, current, ["Business ID"], "Row ID", prior_seed=first,
                          snapshot={"id": "q2", "ageSeconds": 0})
        self.assertEqual(rerun["inserts"], [])
        self.assertEqual(rerun["updates"], [])
        self.assertEqual(len(rerun["unchanged"]), 2)
        self.assertEqual(rerun["snapshot"]["id"], "q2")

    def test_source_change_updates_by_real_row_id_but_preserves_user_edit_as_conflict(self):
        baseline = plan_seed([{"Key": "A", "Value": 1}], [], ["Key"], "Row ID")
        safe = plan_seed([{"Key": "A", "Value": 2}],
                         [{"Key": "A", "Value": 1, "Row ID": "server-42"}],
                         ["Key"], "Row ID", prior_seed=baseline)
        self.assertEqual(safe["updates"][0]["rowId"], "server-42")
        self.assertEqual(safe["updates"][0]["changes"], {"Value": 2})

        edited = plan_seed([{"Key": "A", "Value": 2}],
                           [{"Key": "A", "Value": 9, "Row ID": "server-42"}],
                           ["Key"], "Row ID", prior_seed=baseline)
        self.assertEqual(edited["updates"], [])
        self.assertEqual(edited["conflicts"][0]["columns"], ["Value"])
        self.assertEqual(edited["conflicts"][0]["current"]["Value"], 9)

    def test_interrupted_update_keeps_confirmed_baseline_until_readback(self):
        initial = plan_seed([{"Key": "A", "Value": 1}, {"Key": "B", "Value": 1}], [],
                            ["Key"], "Row ID")
        confirmed = plan_seed(
            [{"Key": "A", "Value": 1}, {"Key": "B", "Value": 1}],
            [{"Key": "A", "Value": 1, "Row ID": "a"},
             {"Key": "B", "Value": 1, "Row ID": "b"}],
            ["Key"], "Row ID", prior_seed=initial)
        pending = plan_seed(
            [{"Key": "A", "Value": 2}, {"Key": "B", "Value": 2}],
            [{"Key": "A", "Value": 1, "Row ID": "a"},
             {"Key": "B", "Value": 1, "Row ID": "b"}],
            ["Key"], "Row ID", prior_seed=confirmed)
        records = {r["key"][0]["value"]: r for r in pending["seedRecords"]}
        self.assertEqual(records["A"]["confirmedMigrationValue"]["Value"], 1)
        self.assertEqual(records["A"]["pendingMigrationValue"]["Value"], 2)

        retry = plan_seed(
            [{"Key": "A", "Value": 2}, {"Key": "B", "Value": 2}],
            [{"Key": "A", "Value": 2, "Row ID": "a"},
             {"Key": "B", "Value": 1, "Row ID": "b"}],
            ["Key"], "Row ID", prior_seed=pending)
        self.assertEqual([u["rowId"] for u in retry["updates"]], ["b"])
        self.assertEqual(retry["conflicts"], [])
        retried = {r["key"][0]["value"]: r for r in retry["seedRecords"]}
        self.assertEqual(retried["A"]["confirmedMigrationValue"]["Value"], 2)
        self.assertEqual(retried["B"]["confirmedMigrationValue"]["Value"], 1)

    def test_typed_keys_reject_duplicates_and_missing_but_keep_zero_distinct(self):
        with self.assertRaisesRegex(InputSyncError, "missing business key"):
            plan_seed([{"Key": None}], [], ["Key"], "Row ID")
        with self.assertRaisesRegex(InputSyncError, "duplicate source business key"):
            plan_seed([{"Key": 1}, {"Key": 1}], [], ["Key"], "Row ID")
        with self.assertRaisesRegex(InputSyncError, "duplicate current business key"):
            plan_seed([{"Key": "A"}], [{"Key": "A", "Row ID": "1"},
                                          {"Key": "A", "Row ID": "2"}], ["Key"], "Row ID")
        plan = plan_seed([{"Key": 0}, {"Key": "0"}], [], ["Key"], "Row ID")
        self.assertEqual([entry["key"][0]["type"] for entry in plan["inserts"]],
                         ["int", "str"])

    def test_orphans_are_reported_and_never_deleted(self):
        plan = plan_seed([], [{"Key": "old", "Row ID": "r-old"}], ["Key"], "Row ID")
        self.assertEqual(plan["orphans"][0]["rowId"], "r-old")
        self.assertNotIn("deletes", plan)

    def test_date_only_seed_matches_datetime_readback(self):
        source = [{"Setting": "AsOfDate", "AsOfDate": "2026-06-30"}]
        initial = plan_seed(source, [], ["Setting"], "Row ID")
        rerun = plan_seed(
            source,
            [{"Setting": "AsOfDate", "AsOfDate": dt.datetime(2026, 6, 30, 12),
              "Row ID": "server-1"}],
            ["Setting"], "Row ID", prior_seed=initial,
        )

        self.assertEqual(rerun["conflicts"], [])
        self.assertEqual(rerun["updates"], [])
        self.assertEqual(len(rerun["unchanged"]), 1)


class TestMcpPrograms(unittest.TestCase):
    def setUp(self):
        self.architecture = build_shared_input_architecture(
            shared_plan(), connection_id="conn-1", folder_id="folder-1",
            workflows=["sales"])
        self.manifest = {"sourceHash": "a" * 64, "planHash": "b" * 64,
                         "planProjection": {"sharedInputArchitecture": self.architecture}}
        self.bindings = {
            name: {"documentId": "doc-1", "version": 7, "pageId": "page-1"}
            for name in self.architecture["inputs"]
        }
        self.receipt = {
            "status": "APPROVED_STAGING", "sourceHash": "a" * 64,
            "planHash": "b" * 64, "targetBindings": self.bindings,
            "artifacts": build_artifact_records(self.manifest, self.architecture, self.bindings),
        }

    def test_construction_creates_columns_then_refetches_and_seeds_once(self):
        program = emit_construction_program(
            self.manifest, self.receipt, "salesSettings",
        )
        self.assertLess(program.index("add_editable_column('AsOfDate'"),
                        program.index("upsert_rows(["))
        self.assertLess(program.index("add_system_columns(['row_id'])"),
                        program.index("upsert_rows(["))
        self.assertIn("sigma.workbook.get_page('page-1')", program)
        self.assertIn("get_element(created_id)", program)
        self.assertIn("get_column(created_column_ids[1])", program)
        self.assertIn("isinstance(table, InputTable)", program)
        self.assertIn("isinstance(c1, LiteralEditableColumn)", program)
        self.assertIn("datetime(2026, 6, 30, 12)", program)
        self.assertEqual(program.count("upsert_rows(["), 1)
        self.assertNotIn(".query(", program)

    def test_emission_rejects_swapped_manifest_artifact_or_target(self):
        swapped = copy.deepcopy(self.receipt)
        swapped["targetBindings"]["salesSettings"]["pageId"] = "other-page"
        with self.assertRaisesRegex(InputSyncError, "artifact binding"):
            emit_construction_program(self.manifest, swapped, "salesSettings")
        stale = copy.deepcopy(self.receipt)
        stale["planHash"] = "c" * 64
        with self.assertRaisesRegex(InputSyncError, "stale"):
            emit_construction_program(self.manifest, stale, "salesSettings")
        wiring = emit_workbook_wiring(self.manifest, self.receipt)
        self.assertEqual(wiring["artifacts"], self.receipt["artifacts"])
        self.assertIn("salesWithSettings",
                      {item["id"] for item in wiring["wiring"]["elements"]})

    def test_write_programs_use_actual_ids_chunk_at_2000_and_preserve_timestamp(self):
        rows = [{"Key": str(index), "At": dt.datetime(2026, 6, 30, 8, 45, 12)}
                for index in range(2001)]
        seed = plan_seed(rows, [], ["Key"], "Row ID")
        architecture = copy.deepcopy(self.architecture)
        architecture["inputs"]["salesSettings"]["rows"] = rows
        architecture["inputs"]["salesSettings"]["keyColumns"] = ["Key"]
        architecture["inputs"]["salesSettings"]["columns"] = [
            {"name": "Key", "type": "text", "provenance": "source"},
            {"name": "At", "type": "datetime", "provenance": "source"},
        ]
        manifest = {**self.manifest,
                    "planProjection": {"sharedInputArchitecture": architecture}}
        binding = {"documentId": "doc-1", "version": 7, "pageId": "p",
                   "elementId": "e", "columnIds": {"Key": "c-key", "At": "c-at"}}
        bindings = {**self.bindings, "salesSettings": binding}
        receipt = {**self.receipt, "targetBindings": bindings,
                   "artifacts": build_artifact_records(manifest, architecture, bindings)}
        programs = emit_write_programs(
            seed, manifest, receipt, "salesSettings",
        )
        self.assertEqual(len(programs), 2)
        self.assertTrue(all("get_element('e')" in item["program"] for item in programs))
        self.assertIn("get_column('c-at')", programs[0]["program"])
        self.assertIn("datetime(2026, 6, 30, 8, 45, 12)", programs[0]["program"])
        self.assertEqual([item["chunkIndex"] for item in programs], [0, 1])
        self.assertTrue(all(item["checkpointId"] for item in programs))

    def test_write_program_uses_approved_schema_types(self):
        spec = self.architecture["inputs"]["salesSettings"]
        seed = plan_seed(spec["rows"], [], spec["keyColumns"], "Row ID")
        binding = {
            "documentId": "doc-1", "version": 7, "pageId": "page-1",
            "elementId": "sales-settings",
            "columnIds": {"Setting": "setting", "AsOfDate": "as-of-date"},
        }
        bindings = {**self.bindings, "salesSettings": binding}
        receipt = {
            **self.receipt,
            "targetBindings": bindings,
            "artifacts": build_artifact_records(
                self.manifest, self.architecture, bindings),
        }

        program = emit_write_programs(
            seed, self.manifest, receipt, "salesSettings")[0]["program"]

        self.assertIn("datetime(2026, 6, 30, 12)", program)
        self.assertNotIn("'2026-06-30'", program)


class TestSharedArchitecture(unittest.TestCase):
    def test_exact_sales_expense_and_research_bindings_are_plan_hash_material(self):
        source_inputs = [
            {"id": "i_model_inputs_c2", "sourceCellId": "Model Inputs!C2",
             "sheet": "Model Inputs", "address": "C2", "defaultValue": 0.05,
             "values": {"2025": 0.05}, "role": "cross_sheet_driver"},
            {"id": "i_fy_results_b8", "sourceCellId": "FY results!B8",
             "sheet": "FY results", "address": "B8", "defaultValue": 100,
             "values": {"2024": 100}, "role": "source_literal"},
        ]
        architecture = build_shared_input_architecture(
            shared_plan(source_inputs), connection_id="conn", folder_id="folder",
            workflows=["sales", "expense", "research"])
        self.assertEqual(architecture["targetScope"],
                         {"connectionId": "conn", "folderId": "folder"})
        sales = architecture["inputs"]["salesSettings"]
        self.assertEqual(sales["rows"][0]["AsOfDate"], "2026-06-30")
        self.assertEqual(sales["sourceBinding"], "Assumptions!F2")
        self.assertEqual(architecture["calculations"]["TimingStatus"],
                         'If([Stage] = "Closed Won" or [Stage] = "Closed Lost", "Closed", If(DateTrunc("day", [Close Date]) < DateTrunc("day", [AsOfDate]), "Slipped", "Open"))')
        self.assertIn("dayGrainTimingDecision", architecture["provenance"])
        self.assertNotIn("referenceQuotingRepair", architecture["provenance"])

        policy = architecture["inputs"]["expensePolicy"]
        self.assertEqual([row["Category"] for row in policy["rows"]],
                         ["Airfare", "Lodging", "Meals", "Ground Transport", "Software", "Supplies"])
        self.assertEqual(next(row for row in policy["rows"] if row["Category"] == "Meals")["Per-item Limit"], 90)
        self.assertEqual(policy["sourceRange"], "Policy!A1:D7")
        self.assertEqual(policy["unknownCategoryBehavior"], "explicit-null-error")

        research = architecture["inputs"]["researchModelInputs"]
        self.assertEqual(research["keyColumns"], ["Scope"])
        self.assertEqual(len(research["rows"]), 1)
        self.assertEqual(research["rows"][0]["i_model_inputs_c2"], 0.05)
        self.assertEqual(set(research["editableColumnBindings"]),
                         {"i_model_inputs_c2", "i_fy_results_b8"})
        self.assertEqual(research["sourceCellBindings"]["i_fy_results_b8"], "FY results!B8")
        self.assertEqual([column["type"] for column in research["columns"]],
                         ["text", "number", "number"])
        self.assertEqual(research["joinGrain"], ["Scope"])
        self.assertNotIn("Row ID", research["joinGrain"])

    def test_operational_fragments_emit_ordinary_inputs_native_joins_and_calculations(self):
        architecture = build_shared_input_architecture(
            shared_plan(), connection_id="conn", folder_id="folder",
            workflows=["sales", "expense"])
        fragments = build_operational_input_elements(architecture)
        by_id = {item["id"]: item for item in fragments["elements"]}
        self.assertEqual(by_id["salesSettings"]["source"],
                         {"kind": "empty", "connectionId": "conn"})
        self.assertEqual(by_id["expensePolicy"]["source"]["kind"], "empty")
        sales = by_id["salesWithSettings"]
        self.assertEqual(sales["source"]["kind"], "join")
        self.assertEqual(sales["source"]["joins"][0]["right"]["elementId"],
                         "salesSettings")
        self.assertIn('[Settings/AsOfDate]',
                      next(c["formula"] for c in sales["columns"] if c["name"] == "TimingStatus"))
        self.assertIn('DateTrunc("day", [Settings/AsOfDate])',
                      next(c["formula"] for c in sales["columns"] if c["name"] == "TimingStatus"))
        self.assertIn("Weighted ACV", {c["name"] for c in sales["columns"]})
        expense = by_id["expensesWithPolicy"]
        self.assertEqual(expense["source"]["joins"][0]["right"]["elementId"],
                         "expensePolicy")
        excess = next(c["formula"] for c in expense["columns"] if c["name"] == "PolicyExcess")
        self.assertIn("IsNull([Policy/Category])", excess)
        self.assertNotIn("Receipt Threshold", excess)
        exception = next(c["formula"] for c in expense["columns"] if c["name"] == "Exception")
        self.assertIn("IsNull([Policy/Category])", exception)
        self.assertIn("[PolicyExcess] > 0", exception)
        self.assertIn('"MISSING RECEIPT"', exception)
        self.assertIn("Department", {c["name"] for c in expense["columns"]})
        self.assertEqual(fragments["consumerSources"],
                         {"sales": "salesWithSettings", "expense": "expensesWithPolicy"})

    def test_shared_type_mapping_covers_null_bool_datetime_number_and_text(self):
        self.assertEqual([sigma_input_type(value) for value in
                          [None, True, dt.date(2026, 1, 1), 1.5, "x"]],
                         ["text", "boolean", "datetime", "number", "text"])

    def test_workflow_selection_does_not_require_foreign_source_features(self):
        research = build_shared_input_architecture(
            {"sourceInputs": []}, connection_id="conn", folder_id="folder",
            workflows=["research"])
        self.assertEqual(set(research["inputs"]), {"researchModelInputs"})
        self.assertEqual(set(research["mappedFeatures"]),
                         {"research-model-inputs", "research-literal-overrides"})
        with self.assertRaisesRegex(InputSyncError, "Assumptions"):
            build_shared_input_architecture(
                {"sourceInputs": []}, connection_id="conn", folder_id="folder",
                workflows=["sales"])
        duplicate_policy = shared_plan()
        duplicate_policy["sharedInputSources"]["expensePolicy"]["rows"].append(
            copy.deepcopy(duplicate_policy["sharedInputSources"]["expensePolicy"]["rows"][0]))
        with self.assertRaisesRegex(InputSyncError, "duplicate source business key"):
            build_shared_input_architecture(
                duplicate_policy, connection_id="conn", folder_id="folder",
                workflows=["expense"])
        missing_policy = shared_plan()
        missing_policy["sharedInputSources"]["expensePolicy"]["rows"][0].pop("Policy Owner")
        with self.assertRaisesRegex(InputSyncError, "preserve all"):
            build_shared_input_architecture(
                missing_policy, connection_id="conn", folder_id="folder",
                workflows=["expense"])

    def test_sales_date_boundaries_and_closed_precedence_match_emitted_formula(self):
        june_30_noon = dt.datetime(2026, 6, 30, 12)
        self.assertEqual(evaluate_timing_status("Open", dt.datetime(2026, 6, 29),
                                                june_30_noon), "Slipped")
        self.assertEqual(evaluate_timing_status("Open", dt.datetime(2026, 6, 30),
                                                june_30_noon), "Open")
        self.assertEqual(evaluate_timing_status("Open", june_30_noon,
                                                june_30_noon), "Open")
        self.assertEqual(evaluate_timing_status("Open", dt.datetime(2026, 7, 1),
                                                june_30_noon), "Open")
        july_1_noon = dt.datetime(2026, 7, 1, 12)
        self.assertEqual(evaluate_timing_status("Open", dt.datetime(2026, 6, 30),
                                                july_1_noon), "Slipped")
        self.assertEqual(evaluate_timing_status("Closed Won", dt.datetime(2026, 6, 29),
                                                july_1_noon), "Closed")
        self.assertEqual(evaluate_timing_status("Closed Lost", dt.datetime(2026, 7, 2),
                                                july_1_noon), "Closed")

    def test_expense_policy_edit_recomputes_all_matching_exceptions_without_changing_spend(self):
        meals = [
            {"Amount": 125, "Card Amount": 125, "Ledger Amount": 125,
             "Receipt Status": "Attached", "Approval Status": "Approved"},
            {"Amount": 300, "Card Amount": 300, "Ledger Amount": 300,
             "Receipt Status": "Attached", "Approval Status": "Approved"},
        ]
        spend = sum(row["Amount"] for row in meals)
        before = [evaluate_expense_exception(row, 90)["exception"] for row in meals]
        after = [evaluate_expense_exception(row, 400)["exception"] for row in meals]
        self.assertEqual(before, ["OVER LIMIT", "OVER LIMIT"])
        self.assertEqual(after, ["CLEAR", "CLEAR"])
        self.assertEqual(sum(row["Amount"] for row in meals), spend)
        self.assertEqual(evaluate_expense_exception(meals[0], None),
                         {"policyExcess": None, "exception": None})


class TestInputOrchestration(unittest.TestCase):
    @mock.patch("input_orchestration.gate_staging_manifest", return_value=True)
    def test_staging_calls_manifest_gate_and_requires_exact_feature_scope(self, gate):
        architecture = build_shared_input_architecture(
            {"sourceInputs": []}, connection_id="conn", folder_id="folder",
            workflows=["research"])
        manifest = {"sourceHash": "a" * 64, "planHash": "b" * 64,
                    "planProjection": {"sharedInputArchitecture": architecture}}
        approval = {"sourceHash": "a" * 64, "planHash": "b" * 64,
                    "reviewer": "Named Reviewer",
                    "mappedFeatures": list(reversed(architecture["mappedFeatures"])),
                    "targetBindings": {"researchModelInputs": {
                        "documentId": "doc", "version": 1, "pageId": "page"}}}
        source_gate = {"sourceSha256": "a" * 64, "status": "REVIEW_REQUIRED",
                       "reviewStatus": "DISPOSITIONED"}
        result = gate_input_staging(manifest, approval,
                                    {"sharedInputArchitecture": architecture}, source_gate)
        gate.assert_called_once_with(manifest, approval)
        self.assertEqual(result["status"], "APPROVED_STAGING")
        self.assertEqual(result["sourceStatus"], "REVIEW_REQUIRED")
        for field in ("sourceHash", "planHash", "reviewer", "mappedFeatures"):
            bad = copy.deepcopy(approval); bad.pop(field)
            with self.assertRaisesRegex(ValueError, field):
                gate_input_staging(manifest, bad, architecture, source_gate)
        with self.assertRaisesRegex(ValueError, "BLOCKED"):
            gate_input_staging(manifest, approval, architecture,
                               {**source_gate, "status": "BLOCKED", "reviewStatus": "BLOCKED"})

    @mock.patch("input_orchestration.gate_completion_manifest", return_value=True)
    def test_completion_calls_manifest_gate_and_consumes_recorded_runtime_evidence(self, gate):
        architecture = build_shared_input_architecture(
            shared_plan(), connection_id="conn", folder_id="folder",
            workflows=["sales", "expense", "research"])
        bindings = {name: {"documentId": "doc", "version": 3,
                           "pageId": "p-" + name, "elementId": "e-" + name,
                           "columnIds": {c["name"]: "c-" + str(i)
                                         for i, c in enumerate(spec["columns"])}}
                    for name, spec in architecture["inputs"].items()}
        manifest = {"sourceHash": "a" * 64, "planHash": "b" * 64,
                    "target": {"documentId": "doc", "version": 3},
                    "planProjection": {"sharedInputArchitecture": architecture}}
        receipt = {"status": "APPROVED_STAGING", "sourceHash": "a" * 64,
                   "planHash": "b" * 64, "targetBindings": bindings,
                   "mappedFeatures": architecture["mappedFeatures"],
                   "artifacts": build_artifact_records(manifest, architecture, bindings)}
        evidence = {"sourceHash": "a" * 64, "planHash": "b" * 64,
                    "status": "passed", "target": {"documentId": "doc", "version": 3},
                    "mappedFeatures": architecture["mappedFeatures"],
                    "targetBindings": bindings,
                    "tables": {name: {"querySnapshotId": "q-" + name,
                                       "readbackId": "r-" + name,
                                       "persistenceReadbackId": "p-" + name,
                                       "rows": len(spec["rows"]),
                                       "distinct": len(spec["rows"]),
                                       "keyColumns": spec["keyColumns"],
                                       "naturalKeys": [
                                           {key: row[key] for key in spec["keyColumns"]}
                                           for row in spec["rows"]]}
                               for name, spec in architecture["inputs"].items()},
                    "checkpointBindings": [
                        {"checkpointId": "chunk-" + name, "feature": name,
                         "artifactHash": receipt["artifacts"][name]["artifactHash"],
                         "documentId": "doc", "version": 3,
                         "elementId": bindings[name]["elementId"]}
                        for name in architecture["inputs"]],
                    "acceptanceChecks": [
                        {"id": check_id, "status": "passed", "restored": True,
                         "evidenceId": "ev-" + check_id, "documentId": "doc", "version": 3,
                         "sourceHash": "a" * 64, "planHash": "b" * 64}
                        for check_id in architecture["acceptanceChecks"]]}
        result = gate_input_completion(manifest, receipt, evidence)
        gate.assert_called_once()
        self.assertEqual(result["status"], "VERIFIED_COMPLETE")

        partial = copy.deepcopy(evidence)
        partial["tables"].pop("expensePolicy")
        with self.assertRaisesRegex(ValueError, "exactly cover"):
            gate_input_completion(manifest, receipt, partial)
        swapped = copy.deepcopy(evidence)
        swapped["targetBindings"]["salesSettings"]["elementId"] = "other"
        with self.assertRaisesRegex(ValueError, "target bindings"):
            gate_input_completion(manifest, receipt, swapped)

        incomplete = copy.deepcopy(evidence)
        incomplete["tables"]["expensePolicy"].update({
            "rows": 1, "distinct": 1,
            "naturalKeys": [{"Category": "Airfare"}],
        })
        with self.assertRaisesRegex(ValueError, "approved row count"):
            gate_input_completion(manifest, receipt, incomplete)

        wrong_keys = copy.deepcopy(evidence)
        wrong_keys["tables"]["expensePolicy"]["naturalKeys"] = [
            {"Category": "Wrong " + str(index)} for index in range(6)]
        with self.assertRaisesRegex(ValueError, "natural keys"):
            gate_input_completion(manifest, receipt, wrong_keys)

        expense = evidence["tables"]["expensePolicy"]
        self.assertEqual(expense["rows"], 6)
        self.assertEqual(len(expense["naturalKeys"]), 6)


class TestTypedWorkbookExtraction(unittest.TestCase):
    def test_business_id_is_retained_and_system_id_requires_explicit_provenance(self):
        module = load_script("xlsx-to-input-csv.py", "typed_input_csv")
        headers = ["ID", "Amount", "Updated At"]
        self.assertEqual(module.select_columns(headers, system_columns={"Updated At"}),
                         [(0, "ID"), (1, "AMOUNT")])

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "fixture.xlsx"
            wb = Workbook(); ws = wb.active
            ws.append(["ID", "Amount"]); ws.append([0, None]); ws.append([1, 0])
            tab = Table(displayName="Entries", ref="A1:B3")
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2")
            ws.add_table(tab); wb.save(path)
            loaded = load_workbook(path, data_only=True)
            _, found = module.find_table(loaded, "Entries")
            rows = module.extract_typed_rows(loaded.active, found)
            self.assertEqual(rows, [{"ID": 0, "AMOUNT": None}, {"ID": 1, "AMOUNT": 0}])


class TestInputSyncCli(unittest.TestCase):
    def test_plan_command_writes_reviewable_plan_and_program_bundle(self):
        module = load_script("input-sync.py", "input_sync_cli")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            architecture = build_shared_input_architecture(
                shared_plan(), connection_id="conn", folder_id="folder",
                workflows=["sales"])
            architecture["inputs"]["salesSettings"]["rows"] = [{"Key": "A", "Value": 1}]
            architecture["inputs"]["salesSettings"]["keyColumns"] = ["Key"]
            architecture["inputs"]["salesSettings"]["columns"] = [
                {"name": "Key", "type": "text", "provenance": "source"},
                {"name": "Value", "type": "number", "provenance": "source"},
            ]
            manifest = {"sourceHash": "a" * 64, "planHash": "b" * 64,
                        "planProjection": {"sharedInputArchitecture": architecture}}
            bindings = {name: {"documentId": "doc", "version": 1, "pageId": "p",
                               "elementId": "e-" + name,
                               "columnIds": {c["name"]: "c-" + str(i)
                                             for i, c in enumerate(spec["columns"])}}
                        for name, spec in architecture["inputs"].items()}
            receipt = {"status": "APPROVED_STAGING", "sourceHash": "a" * 64,
                       "planHash": "b" * 64, "targetBindings": bindings,
                       "artifacts": build_artifact_records(manifest, architecture, bindings)}
            (root / "manifest.json").write_text(json.dumps(manifest))
            (root / "receipt.json").write_text(json.dumps(receipt))
            (root / "current.json").write_text("[]")
            code = module.main([
                "plan", "--manifest", str(root / "manifest.json"),
                "--receipt", str(root / "receipt.json"),
                "--feature", "salesSettings", "--current-rows", str(root / "current.json"),
                "--row-id", "Row ID",
                "--snapshot-id", "query-1", "--snapshot-age-seconds", "4",
                "--out", str(root / "bundle.json"),
            ])
            self.assertEqual(code, 0)
            bundle = json.loads((root / "bundle.json").read_text())
            self.assertEqual(bundle["plan"]["snapshot"], {"id": "query-1", "ageSeconds": 4})
            self.assertEqual(bundle["programs"][0]["actionCount"], 1)
            self.assertIn("checkpointId=", bundle["programs"][0]["program"])

    def test_stage_command_passes_architecture_output_plan_shape(self):
        module = load_script("input-sync.py", "input_sync_stage_cli")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.xlsx"
            wb = Workbook(); wb.active["A1"] = "={1,2,3}"; wb.save(source)
            plan_path = root / "plan.json"
            plan_path.write_text(json.dumps({"sourcePath": str(source),
                                             "sourceSha256": "placeholder"}))
            replacements = root / "replacements.json"
            replacements.write_text(json.dumps({"formulaCells": [{
                "id": "Sheet!A1", "mappingStatus": "live",
                "targetMapping": {"logicalElementId": "calc", "logicalColumnId": "a1"},
                "reviewedReplacement": {"originalFormula": "={1,2,3}",
                                        "replacementFormula": "[Input/Value]",
                                        "reviewer": "Owner", "reviewStatus": "approved"},
            }]}))
            out_plan = root / "out-plan.json"; manifest = root / "manifest.json"
            module.main(["architecture", "--plan", str(plan_path), "--source", str(source),
                         "--connection", "conn", "--folder", "folder",
                         "--workflow", "research",
                         "--reviewed-replacements", str(replacements),
                         "--out-plan", str(out_plan), "--out-manifest", str(manifest)])
            approval = root / "approval.json"; approval.write_text("{}")
            with (mock.patch.object(module, "gate_plan_staging_source",
                                    return_value=(source, {"sourceSha256": "x"})),
                  mock.patch.object(module, "gate_input_staging",
                                    return_value={"status": "APPROVED_STAGING"}) as stage):
                module.main(["stage", "--manifest", str(manifest), "--plan", str(out_plan),
                             "--approval", str(approval), "--out", str(root / "receipt.json")])
            passed_plan = stage.call_args.args[2]
            self.assertIn("sharedInputArchitecture", passed_plan)
            self.assertIn("mappedFeatures", passed_plan["sharedInputArchitecture"])
            manifest_payload = json.loads(manifest.read_text())
            self.assertEqual(manifest_payload["formulaCells"][0]["mappingStatus"], "live")


if __name__ == "__main__":
    unittest.main()
