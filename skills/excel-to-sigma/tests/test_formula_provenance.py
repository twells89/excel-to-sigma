import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName


ROOT = Path(__file__).resolve().parents[1]
LIB = ROOT / "scripts" / "lib"
if str(LIB) not in sys.path:
    sys.path.insert(0, str(LIB))

from excel_expression import (  # noqa: E402
    UnsupportedExpression,
    emit_sigma,
    evaluate,
    extract_references,
    parse_expression,
)
from formula_graph import CellIdentity, FormulaGraph  # noqa: E402
from migration_manifest import (  # noqa: E402
    ManifestError,
    build_manifest,
    validate_completion,
    validate_manifest,
    validate_staging_approval,
)


class TestExcelExpression(unittest.TestCase):
    def test_one_ast_drives_references_evaluation_and_sigma_emission(self):
        expression = parse_expression("=SUM((B2+B3))*104.5%+ROUND(B4/2,1)", sheet="FY results")
        refs = {(ref.sheet, ref.address) for ref in extract_references(expression)}
        values = {("FY results", "B2"): 10, ("FY results", "B3"): 6,
                  ("FY results", "B4"): 5}

        self.assertEqual(refs, {("FY results", "B2"), ("FY results", "B3"),
                                ("FY results", "B4")})
        self.assertAlmostEqual(evaluate(expression, lambda ref: values[(ref.sheet, ref.address)]),
                               19.22)
        self.assertEqual(
            emit_sigma(expression, lambda ref: f"[{ref.address}]"),
            "Sum(([B2] + [B3])) * (104.5 * 0.01) + Round([B4] / 2, 1)",
        )

    def test_if_absolute_cross_sheet_and_named_references_are_preserved(self):
        expression = parse_expression(
            "=IF($B2>0,$B2*(1+'Model Inputs'!$C$2),TaxRate)", sheet="FY results"
        )
        refs = list(extract_references(expression))
        self.assertEqual(
            [(ref.sheet, ref.address, ref.column_absolute, ref.row_absolute, ref.name)
             for ref in refs],
            [("FY results", "B2", True, False, None),
             ("FY results", "B2", True, False, None),
             ("Model Inputs", "C2", True, True, None),
             (None, None, False, False, "TaxRate")],
        )
        values = {("FY results", "B2"): 100, ("Model Inputs", "C2"): 0.05,
                  (None, "TaxRate"): 20}
        self.assertEqual(
            evaluate(expression, lambda ref: values[(ref.sheet, ref.address or ref.name)]),
            105,
        )

    def test_unsupported_arrays_and_external_references_fail_loudly(self):
        for formula in ("={1,2,3}", "='[other.xlsx]Sheet1'!A1"):
            with self.subTest(formula=formula):
                with self.assertRaises(UnsupportedExpression):
                    parse_expression(formula, sheet="Sheet1")

    def test_sumifs_ranges_and_text_if_support_audit_checks(self):
        expression = parse_expression(
            '=IF(SUMIFS(\'Segment Detail\'!$D$2:$D$4,\'Segment Detail\'!$A$2:$A$4,2024)'
            '-\'FY results\'!B8=0,"PASS","REVIEW")',
            sheet="Audit",
        )
        values = {
            ("Segment Detail", "A2"): 2023, ("Segment Detail", "A3"): 2024,
            ("Segment Detail", "A4"): 2024, ("Segment Detail", "D2"): 9,
            ("Segment Detail", "D3"): 40, ("Segment Detail", "D4"): 60,
            ("FY results", "B8"): 100,
        }

        self.assertEqual(
            evaluate(expression, lambda ref: values[(ref.sheet, ref.address)]), "PASS"
        )


class TestFormulaGraph(unittest.TestCase):
    def test_cell_period_graph_does_not_invent_a_row_cycle(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "FY results"
        ws["B8"] = 760
        ws["B11"] = -452
        ws["B13"] = "=B8+B11"
        ws["H8"] = 1242
        ws["H13"] = 602
        ws["H11"] = "=H13-H8"

        graph = FormulaGraph.from_workbook(wb)

        self.assertEqual(graph.cycles(), [])
        self.assertEqual(
            set(graph.topological_order()),
            {CellIdentity("FY results", "B13"), CellIdentity("FY results", "H11")},
        )

    def test_true_cell_cycle_and_missing_reference_are_explicit(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "=B1"
        ws["B1"] = "=A1"
        ws["C1"] = "=Z99"

        graph = FormulaGraph.from_workbook(wb)

        self.assertEqual(
            graph.cycles(),
            [(CellIdentity("Sheet1", "A1"), CellIdentity("Sheet1", "B1"))],
        )
        self.assertEqual(graph.unresolved_references,
                         {CellIdentity("Sheet1", "C1"): (CellIdentity("Sheet1", "Z99"),)})

    def test_defined_cell_name_is_a_real_dependency_and_evaluates(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Model"
        inputs = wb.create_sheet("Inputs")
        inputs["A1"] = 0.25
        wb.defined_names.add(DefinedName("TaxRate", attr_text="'Inputs'!$A$1"))
        ws["A1"] = "=100*TaxRate"

        graph = FormulaGraph.from_workbook(wb)

        self.assertEqual(graph.dependencies[CellIdentity("Model", "A1")],
                         (CellIdentity("Inputs", "A1"),))
        self.assertEqual(graph.evaluate()[CellIdentity("Model", "A1")], 25)

    def test_unknown_names_and_transitive_consumers_are_unresolved(self):
        wb = Workbook(); ws = wb.active; ws.title = "Model"
        ws["A1"] = "=MissingName+1"; ws["B1"] = "=A1*2"
        graph = FormulaGraph.from_workbook(wb)
        self.assertEqual(graph.unresolved_names[CellIdentity("Model", "A1")],
                         ("MissingName",))
        self.assertEqual(graph.unresolved_formula_nodes(),
                         {CellIdentity("Model", "A1"), CellIdentity("Model", "B1")})

    def test_cross_sheet_computed_reference_remains_a_formula_dependency(self):
        wb = Workbook(); model = wb.active; model.title = "Model"
        inputs = wb.create_sheet("Inputs")
        inputs["A1"] = 5; inputs["A2"] = "=A1*2"; model["A1"] = "=Inputs!A2+1"
        graph = FormulaGraph.from_workbook(wb)
        self.assertEqual(graph.dependencies[CellIdentity("Model", "A1")],
                         (CellIdentity("Inputs", "A2"),))
        self.assertEqual(graph.evaluate()[CellIdentity("Model", "A1")], 11)

    def test_fixture_graph_evaluates_all_display_and_audit_formulas(self):
        fixture = ROOT / "tests" / "fixtures" / "fixture-research-model.xlsx"
        expected_path = ROOT / "tests" / "fixtures" / "fixture-research-expected.json"
        if not fixture.exists() or not expected_path.exists():
            self.skipTest("research fixture not installed")
        expected = json.loads(expected_path.read_text(encoding="utf-8"))
        from openpyxl import load_workbook
        wb = load_workbook(fixture, data_only=False)
        cached = load_workbook(fixture, data_only=True)
        graph = FormulaGraph.from_workbook(wb)
        values = graph.evaluate()

        self.assertEqual(sum(len(row) for row in expected["rows"].values()), 420)
        for row, wanted in expected["rows"].items():
            for column, value in enumerate(wanted, 2):
                got = values.get(CellIdentity("FY results", f"{chr(64 + column)}{row}"))
                if value is None:
                    self.assertIsNone(got)
                else:
                    source_value = cached["FY results"].cell(int(row), column).value
                    self.assertAlmostEqual(got, source_value, delta=1e-9)
                    self.assertAlmostEqual(got, value, delta=1e-3)
        self.assertEqual(sum(1 for cell in graph.formulas if cell.sheet == "FY results"), 327)
        self.assertEqual(sum(1 for cell in graph.formulas if cell.sheet == "Audit"), 16)
        self.assertEqual([values[CellIdentity("Audit", f"F{row}")] for row in range(5, 9)],
                         ["PASS", "PASS", "PASS", "PASS"])


class TestMigrationManifest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "model.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "FY results"
        ws["A1"] = 2024
        ws["B1"] = 10
        ws["C1"] = "=B1*2"
        audit = wb.create_sheet("Audit")
        audit["A1"] = "='FY results'!C1-20"
        wb.save(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_manifest_accounts_for_each_formula_once_and_keeps_literals_distinct(self):
        manifest = build_manifest(self.path)
        formulas = manifest["formulaCells"]

        self.assertEqual([item["id"] for item in formulas],
                         ["Audit!A1", "FY results!C1"])
        self.assertEqual(len({item["id"] for item in formulas}), 2)
        self.assertEqual([item["id"] for item in manifest["literalInputs"]],
                         ["FY results!A1", "FY results!B1"])
        self.assertEqual({item["id"]: item["mappingStatus"] for item in formulas},
                         {"Audit!A1": "verification", "FY results!C1": "unresolved"})
        self.assertTrue(validate_manifest(manifest, self.path))

    def test_source_staleness_and_plan_hash_are_independent_of_runtime_evidence(self):
        plan = {"lines": [{"row": 1, "canonical": "[c2]"}],
                "evidence": [{"runId": "old"}], "targetVersions": [1]}
        manifest = build_manifest(self.path, plan)
        changed = copy.deepcopy(plan)
        changed["evidence"] = [{"runId": "new"}]
        changed["targetVersions"] = [2, 3]
        rebuilt = build_manifest(self.path, changed)
        self.assertEqual(manifest["planHash"], rebuilt["planHash"])

        assigned = copy.deepcopy(plan)
        assigned["formulaCells"] = [{
            "id": "FY results!C1",
            "targetMapping": {"logicalFeatureId": "revenue", "workbookId": "wb-old",
                              "elementId": "element-old"},
        }]
        reassigned = copy.deepcopy(assigned)
        reassigned["formulaCells"][0]["targetMapping"].update(
            {"workbookId": "wb-new", "elementId": "element-new"})
        self.assertNotEqual(build_manifest(self.path, assigned)["planHash"],
                            build_manifest(self.path, reassigned)["planHash"])

        for key, value in (
            ("lines", [{"row": 1, "canonical": "[c3]"}]),
            ("sourceInputs", [{"id": "driver", "timestamp": "2026-10-15T12:00:00Z"}]),
            ("features", [{"sourceFeatureId": "FY results!C1", "mapping": "Revenue"}]),
            ("targetScope", {"connectionId": "logical-connection", "folderId": "logical-folder"}),
        ):
            with self.subTest(key=key):
                logical_change = copy.deepcopy(plan)
                logical_change[key] = value
                self.assertNotEqual(manifest["planHash"], build_manifest(self.path, logical_change)["planHash"])

        self.path.write_bytes(self.path.read_bytes() + b"changed")
        with self.assertRaisesRegex(ManifestError, "source hash"):
            validate_manifest(manifest, self.path)

    def test_unsupported_source_formula_is_explicitly_unresolved(self):
        from openpyxl import load_workbook
        workbook = load_workbook(self.path)
        workbook["FY results"]["D1"] = "={1,2,3}"
        workbook.save(self.path)

        manifest = build_manifest(self.path)
        item = next(item for item in manifest["formulaCells"] if item["id"] == "FY results!D1")

        self.assertEqual(item["mappingStatus"], "unresolved")
        self.assertIn("unsupported", item["unresolvedReason"])

    def test_staging_review_and_completion_have_separate_evidence_contracts(self):
        plan = {"formulaCells": [
            {"id": "Audit!A1", "mappingStatus": "live",
             "targetMapping": {"elementId": "logical-audit", "checkId": "a1"}},
            {"id": "FY results!C1", "mappingStatus": "live",
             "targetMapping": {"elementId": "logical-calc", "columnId": "c1"}},
        ]}
        manifest = build_manifest(self.path, plan)
        approval = {
            "sourceHash": manifest["sourceHash"],
            "planHash": manifest["planHash"],
            "owner": "Migration Owner",
            "reviewItems": [item["id"] for item in manifest["formulaCells"]],
        }

        self.assertTrue(validate_staging_approval(manifest, approval))
        manifest["target"] = {"documentId": "wb-1", "version": 7}
        with self.assertRaisesRegex(ManifestError, "verification evidence"):
            validate_completion(manifest)

        for item in manifest["formulaCells"]:
            item["targetBinding"] = {
                "documentId": "wb-1", "elementId": "server-element", "version": 7,
                "columnId": "server_" + item["id"].replace("!", "_"),
            }
            item["verification"] = {
                "status": "passed", "sourceCellId": item["id"],
                "targetId": item["targetBinding"]["columnId"], "evidenceId": "run-1",
                "sourceHash": manifest["sourceHash"], "planHash": manifest["planHash"],
                "targetDocumentId": "wb-1", "targetVersion": 7,
            }
        self.assertTrue(validate_completion(manifest))

    def test_manifest_rejects_mapping_tampering_omissions_and_cached_staging(self):
        plan = {"formulaCells": [
            {"id": "Audit!A1", "mappingStatus": "verification", "targetMapping": {"checkId": "a1"}},
            {"id": "FY results!C1", "mappingStatus": "live", "targetMapping": {"columnId": "c1"}},
        ]}
        manifest = build_manifest(self.path, plan)
        for mutate in (
            lambda value: value["formulaCells"][1].update(mappingStatus="cached"),
            lambda value: value["formulaCells"][1]["targetMapping"].update(columnId="changed"),
            lambda value: value["formulaCells"].pop(),
        ):
            changed = copy.deepcopy(manifest); mutate(changed)
            with self.assertRaises(ManifestError):
                validate_manifest(changed)

        cached_plan = copy.deepcopy(plan)
        cached_plan["formulaCells"][1]["mappingStatus"] = "cached"
        cached = build_manifest(self.path, cached_plan)
        approval = {"sourceHash": cached["sourceHash"], "planHash": cached["planHash"],
                    "owner": "Owner", "reviewItems": [i["id"] for i in cached["formulaCells"]]}
        with self.assertRaisesRegex(ManifestError, "unresolved"):
            validate_staging_approval(cached, approval)


if __name__ == "__main__":
    unittest.main()
