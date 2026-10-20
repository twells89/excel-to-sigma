"""Offline construction/math proofs only; never runtime acceptance or VBA execution."""
import copy
import importlib.util
import json
import math
import sys
import unittest
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'/'lib'))


class CustomPortTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec = importlib.util.find_spec('custom_ports')
        if spec is None:
            return
        import custom_ports
        cls.port = custom_ports

    def setUp(self):
        self.assertTrue(hasattr(self, 'port'), 'reviewed custom-port implementation is missing')
        self.model = json.loads((ROOT/'tests/fixtures/task5-controlled-model.json').read_text())
        self.request = dict(OperationId='solve-A', State='PENDING', Target=2000000,
                            LowerBound=0, UpperBound=.25, ToleranceCents=1,
                            MaxIterations=64, ModelBinding='reviewed-model-1')

    def solve(self, **kwargs):
        return self.port.solve_forecast(self.model['actuals'], self.model['inputs'],
            self.model['definitions'], {**self.request, **kwargs}, expected_binding='reviewed-model-1')

    def test_warehouse_sigma_accessors_and_output_names_are_compile_time_literals(self):
        import ast
        for port,artifact in self.port.emit_python_programs('reviewed-model-1').items():
            tree=ast.parse(artifact['code'])
            for node in ast.walk(tree):
                if not isinstance(node,ast.Call) or not isinstance(node.func,ast.Attribute) or not isinstance(node.func.value,ast.Name) or node.func.value.id!='sigma':continue
                if node.func.attr in {'get_element','get_control_value','output'}:
                    self.assertTrue(node.args and isinstance(node.args[0],ast.Constant) and isinstance(node.args[0].value,str),(port,ast.unparse(node)))
            for function in [n for n in ast.walk(tree) if isinstance(n,(ast.FunctionDef,ast.AsyncFunctionDef))]:
                self.assertFalse(any(isinstance(n,ast.Call) and isinstance(n.func,ast.Attribute) and isinstance(n.func.value,ast.Name) and n.func.value.id=='sigma' for n in ast.walk(function)),function.name)

    def test_port_buttons_emit_explicit_observed_filled_appearance(self):
        fragment=self.port.build_port_fragments('binding','connection')
        buttons=[e for e in fragment['elements'] if e['kind']=='button']
        self.assertTrue(buttons)
        self.assertTrue(all(e.get('appearance')=='filled' for e in buttons))

    def test_allocation_one_based_radians_unrounded_and_not_normalized(self):
        rows=self.port.allocate_overhead([40,30,20,10],100000,'allocation-A')
        golden=[56829.41969615794,43639.46140238523,21411.20008059867,6215.987523460359]
        self.assertEqual([r['Department'] for r in rows], ['Eng','Sales','G&A','Support'])
        self.assertEqual([r['SourceIndex'] for r in rows], [1,2,3,4])
        for row,want in zip(rows,golden):self.assertAlmostEqual(row['Allocation'],want,places=9)
        self.assertAlmostEqual(sum(r['Allocation'] for r in rows),128096.06870260219,places=8)

    def test_allocation_failure_is_one_explicit_null_business_row(self):
        for hc,overhead in [([0]*4,1),([None,1,2,3],1),(['4',1,2,3],1),
                ([float('inf'),1,2,3],1),([-1,1,2,3],1),([1]*4,-1),([1]*4,float('nan'))]:
            before=copy.deepcopy(hc)
            result=self.port.allocate_overhead(hc,overhead,'allocation-B')
            self.assertEqual(len(result),1);self.assertEqual(result[0]['Status'],'FAILURE')
            self.assertIsNone(result[0]['Allocation']);self.assertEqual(hc,before)

    def test_live_equations_and_deterministic_bisection_match_independent_goldens(self):
        result=self.solve()
        self.assertEqual(result['Status'],'SUCCESS');self.assertTrue(result['Converged'])
        self.assertAlmostEqual(result['LowerNetIncome'],482220.9186962107,places=6)
        self.assertAlmostEqual(result['UpperNetIncome'],4689314.644282148,places=6)
        self.assertEqual(result['Iterations'],28);self.assertEqual(result['SolvedGrowth'],.114270488732)
        self.assertAlmostEqual(result['SolvedNetIncome'],1999999.9961722782,places=6)
        # Independent high precision collapsed fixture equation, no port helper calls.
        g=Decimal(str(result['SolvedGrowth']))
        want=Decimal('-2644379.0813037893')+Decimal('521100')*sum((1+g)**n for n in range(1,7))
        self.assertLess(abs(Decimal(str(result['SolvedNetIncome']))-want),Decimal('.000001'))
        self.model['inputs'][1]['Monthly Growth']=.08
        self.assertNotEqual(self.solve()['SolvedGrowth'],result['SolvedGrowth'])

    def test_endpoint_invalid_bracket_iteration_and_model_failures_never_solve(self):
        for target,growth in [(482220.9186962107,0),(4689314.644282148,.25)]:
            result=self.solve(Target=target)
            self.assertEqual((result['SolvedGrowth'],result['Iterations'],result['Endpoint']),(growth,0,True))
        for opts,code in [({'Target':1e9},'INVALID_BRACKET'),({'MaxIterations':1},'ITERATION_LIMIT'),
                ({'LowerBound':-.01},'INVALID_REQUEST'),({'ToleranceCents':2},'INVALID_REQUEST')]:
            result=self.solve(**opts);self.assertEqual(result['ErrorCode'],code)
            self.assertFalse(result['Converged']);self.assertIsNone(result['SolvedGrowth'])
        for mutate in [lambda m:m['inputs'].append(m['inputs'][0]),
                lambda m:m['inputs'][0].update({'Monthly Growth':None}),
                lambda m:m['actuals'][0].update(Amount=float('inf')),
                lambda m:m['definitions'].pop(),
                lambda m:[r.update(Amount=0) for r in m['actuals'] if r['Line Item']=='Subscription Revenue']]:
            original=copy.deepcopy(self.model);mutate(self.model)
            result=self.solve();self.assertEqual(result['Status'],'FAILURE');self.assertIsNone(result['SolvedGrowth'])
            self.model=original

    def test_apply_guard_blocks_stale_pending_failed_inputs_model_row_and_tampered_solution(self):
        result=self.solve()
        current={**self.request,'State':'APPLY_PENDING'}
        def guard(r=result, request=current):
            return self.port.guard_goal_seek(self.model['actuals'],self.model['inputs'],
                self.model['definitions'],request,r,expected_binding='reviewed-model-1')
        self.assertTrue(guard()['ReadyToApply'])
        for request in [{**current,'OperationId':'solve-B'}, {**current,'State':'PENDING'},
                {**current,'ModelBinding':'changed'}, {**current,'Target':2100000}]:
            self.assertFalse(guard(request=request)['ReadyToApply'])
        for change in [{'Converged':False},{'SolvedGrowth':.2},{'Status':'FAILURE'},
                {'SubscriptionRowId':'other'},{'InputFingerprint':'old'}]:
            self.assertFalse(guard(r={**result,**change})['ReadyToApply'])
        self.model['inputs'][2]['Jul-26']+=1
        self.assertFalse(guard()['ReadyToApply'])

    def test_verified_write_plan_changes_only_subscription_and_retry_keeps_operation_identity(self):
        result=self.solve();req={**self.request,'State':'APPLY_PENDING'}
        before=copy.deepcopy(self.model['inputs'])
        change=self.port.plan_goal_seek_write(self.model['actuals'],before,self.model['definitions'],req,result,
                                              expected_binding='reviewed-model-1')
        self.assertEqual(change,{'operationId':'solve-A','rowId':'controlled-row-1',
                                'key':{'Line Item':'Subscription Revenue'},'values':{'Monthly Growth':.114270488732}})
        self.assertEqual(before,self.model['inputs'])
        after=copy.deepcopy(before);after[0]['Monthly Growth']=result['SolvedGrowth']
        receipt=self.port.verify_goal_seek_write(before,after,change)
        self.assertEqual(receipt['status'],'VERIFIED_COMPONENT_WRITE')
        retry={**req,'State':'APPLIED','AppliedOperationId':'solve-A','AppliedGrowth':result['SolvedGrowth']}
        self.assertIsNone(self.port.plan_goal_seek_write(self.model['actuals'],after,self.model['definitions'],retry,result,
                                                       expected_binding='reviewed-model-1'))
        after[1]['Monthly Growth']=0
        with self.assertRaises(ValueError):self.port.verify_goal_seek_write(before,after,change)

    def test_decimal_warehouse_write_and_retry_match_exact_persisted_growth(self):
        result=self.solve();request={**self.request,'State':'APPLY_PENDING'}
        before=copy.deepcopy(self.model['inputs'])
        before[0]['Monthly Growth']=Decimal('0.032')
        change=self.port.plan_goal_seek_write(self.model['actuals'],before,self.model['definitions'],request,result,
            expected_binding='reviewed-model-1')
        after=copy.deepcopy(before);after[0]['Monthly Growth']=Decimal('0.114270488732')
        self.assertEqual(self.port.verify_goal_seek_write(before,after,change)['status'],'VERIFIED_COMPONENT_WRITE')
        retry={**request,'State':'APPLIED','AppliedOperationId':'solve-A','AppliedGrowth':Decimal('0.114270488732')}
        self.assertIsNone(self.port.plan_goal_seek_write(self.model['actuals'],after,self.model['definitions'],retry,result,
            expected_binding='reviewed-model-1'))
        with self.assertRaises(ValueError):self.port.verify_goal_seek_write(before,after,
            {**change,'values':{**change['values'],'Method':'manual'}})

    def test_numeric_fingerprint_is_order_stable_and_value_sensitive(self):
        r=self.solve();self.model['actuals'].reverse();self.model['inputs'].reverse()
        self.assertEqual(r['InputFingerprint'],self.solve()['InputFingerprint'])
        self.model['actuals'][0]['Amount']=Decimal(str(self.model['actuals'][0]['Amount']))
        self.assertEqual(r['InputFingerprint'],self.solve()['InputFingerprint'])
        self.model['actuals'][0]['Amount']+=Decimal('.01')
        self.assertNotEqual(r['InputFingerprint'],self.solve()['InputFingerprint'])



class ConstructionTests(unittest.TestCase):
    def module(self):
        import custom_ports
        return custom_ports

    def test_controlled_source_has_complete_call_graph_and_raw_xlsx_companion(self):
        from macro_behavior import analyze_macro_source
        from openpyxl import load_workbook
        source=ROOT/'tests/fixtures/task5-controlled-business-macros.bas'
        analyzed=analyze_macro_source(source)
        by_name={p['name']:p for p in analyzed['procedures']}
        required={'AllocateOverheadByHeadcount','SolveForTarget','TryFiniteDouble','IsFiniteDouble',
                  'TryLoadActualRows','TryLoadForecastInputs','TryJuneBase','TryForecastAmount',
                  'TrySectionSign','TryAnnualNetIncome','TrySolveBounded','Task5TypedEcho'}
        self.assertEqual(set(by_name),required)
        inventory=self.module().source_helper_inventory(analyzed)
        self.assertFalse(inventory['missingHelpers'])
        self.assertIn('TryAnnualNetIncome',inventory['definedHelpers'])
        # Exact source procedure identities remain unresolved for automatic native emission.
        self.assertEqual(by_name['AllocateOverheadByHeadcount']['disposition']['status'],'UNRESOLVED')
        wb=load_workbook(ROOT/'tests/fixtures/task5-controlled-business-fixture.xlsx')
        self.assertIsNone(wb.vba_archive);self.assertEqual(wb['Alloc']['B2'].value,-999)
        self.assertEqual(wb['Forecast Inputs']['C6'].value,'=Assumptions!B5')
        self.assertEqual(len(wb['Actuals']['A2:D163']),162)
        self.assertTrue(wb['P&L']['N33'].value.startswith('='));wb.close()

    def test_new_operation_tables_enable_published_explore_without_inheriting_companion_mode(self):
        from native_actions import _spec
        for template in [{'connectionId':'c'}, {'connectionId':'c','editMode':'draft'}]:
            before=copy.deepcopy(template)
            for title in ['Workflow Operations','Saved Scenarios','Forecast Snapshots']:
                value=_spec(template,title,[('Operation ID','text')])
                self.assertEqual(value.get('editMode'),'published-explore')
            self.assertEqual(template,before)

    def test_standalone_runtime_collects_live_inputs_and_always_emits_typed_failure(self):
        p=self.module()
        self.assertTrue(hasattr(p,'emit_python_programs'),'warehouse runtime emitter is missing')
        programs=p.emit_python_programs('reviewed-model-1')
        # Host doubles model the documented .columns/.collect()/Row.as_dict contract;
        # execute actual emitted source and assert business output, not mock behavior.
        model=json.loads((ROOT/'tests/fixtures/task5-controlled-model.json').read_text())
        request=dict(OperationId='solve-A',State='PENDING',Target=2000000,LowerBound=0,UpperBound=.25,
                     ToleranceCents=1,MaxIterations=64,ModelBinding='reviewed-model-1',RequestKey='goal-seek')
        from types import SimpleNamespace
        class Frame:
            def __init__(self,rows):self.rows=rows;self.columns=list(rows[0]) if rows else []
            def collect(self):return [SimpleNamespace(as_dict=lambda row=row:row) for row in self.rows]
        data={'Actuals':model['actuals'],'Forecast Inputs View':model['inputs'],
              'Forecast Line Definitions':model['definitions'],'Goal Seek Invocation':[request],
              'Task 5 Headcount Inputs':[{'Department':d,'Headcount':h} for d,h in zip(['Eng','Sales','G&A','Support'],[40,30,20,10])],
              'Task 5 Allocation Assumptions':[{'Key':'Overhead','Value':100000}]}
        output=[]
        sigma=SimpleNamespace(get_element=lambda name:Frame(data[name]),output=lambda name,df:output.append((name,df)))
        session=SimpleNamespace(create_dataframe=lambda rows,schema:rows)
        exec(programs['goalSeek']['code'],{'sigma':sigma,'session':session})
        row=dict(zip(programs['goalSeek']['columns'],output[-1][1][0]))
        self.assertEqual(row['SolvedGrowth'],.114270488732)
        data[programs['goalSeek']['output']]=[row];request['State']='APPLY_PENDING'
        exec(programs['applyGuard']['code'],{'sigma':sigma,'session':session})
        ready=dict(zip(programs['applyGuard']['columns'],output[-1][1][0]))
        self.assertTrue(ready['ReadyToApply'])
        request.update(OperationId='solve-B',State='PENDING')
        exec(programs['applyGuard']['code'],{'sigma':sigma,'session':session})
        self.assertFalse(dict(zip(programs['applyGuard']['columns'],output[-1][1][0]))['ReadyToApply'])
        data['Actuals']=[]
        exec(programs['goalSeek']['code'],{'sigma':sigma,'session':session})
        failure=dict(zip(programs['goalSeek']['columns'],output[-1][1][0]))
        self.assertEqual(failure['OperationId'],'solve-B');self.assertEqual(failure['Status'],'FAILURE')
        exec(programs['allocation']['code'],{'sigma':sigma,'session':session})
        self.assertEqual(len(output[-1][1]),4)
        data['Task 5 Headcount Inputs'][0]['Headcount']=None
        exec(programs['allocation']['code'],{'sigma':sigma,'session':session})
        self.assertEqual(len(output[-1][1]),1);self.assertEqual(output[-1][1][0][1],'FAILURE')

    def test_guard_reads_solver_python_output_reference_with_observed_column_pairs(self):
        from types import SimpleNamespace
        p=self.module()
        model=json.loads((ROOT/'tests/fixtures/task5-controlled-model.json').read_text())
        request=dict(RequestKey='goal-seek',OperationId='solve-A',State='PENDING',Target=2000000,
                     LowerBound=0,UpperBound=.25,ToleranceCents=1,MaxIterations=64,ModelBinding='reviewed-model-1')
        result_columns=[name for name,_ in p.OUTPUT_SCHEMAS['goalSeek']]
        observed={name:{'identifier':'"'+name+'"','rowKey':name} for name in result_columns}
        programs=p.emit_python_programs('reviewed-model-1',column_maps={'GoalSeekResult':observed})
        data={'Actuals':model['actuals'],'Forecast Inputs View':model['inputs'],
              'Forecast Line Definitions':model['definitions'],'Goal Seek Invocation':[request]}
        class Frame:
            def __init__(self,name):
                self.rows=data[name]
                self.columns=['"'+key+'"' for key in self.rows[0]] if name=='GoalSeekResult' else list(self.rows[0])
            def collect(self):return [SimpleNamespace(as_dict=lambda row=row:row) for row in self.rows]
        output=[]
        sigma=SimpleNamespace(get_element=Frame,output=lambda name,rows:output.append((name,rows)))
        session=SimpleNamespace(create_dataframe=lambda rows,schema:rows)
        exec(programs['goalSeek']['code'],{'sigma':sigma,'session':session})
        self.assertEqual(output[-1][0],'GoalSeekResult')
        solved=dict(zip(programs['goalSeek']['columns'],output[-1][1][0]))
        self.assertEqual(solved['Status'],'SUCCESS')
        data[output[-1][0]]=[solved]
        request['State']='APPLY_PENDING'
        exec(programs['applyGuard']['code'],{'sigma':sigma,'session':session})
        guard=dict(zip(programs['applyGuard']['columns'],output[-1][1][0]))
        self.assertTrue(guard['ReadyToApply'],guard)

    def test_runtime_column_identifier_and_collected_row_key_are_separately_explicit(self):
        from types import SimpleNamespace
        maps={'Goal Seek Invocation':{'RequestKey':{'identifier':'"RequestKey"','rowKey':'RequestKey'}}}
        programs=self.module().emit_python_programs('binding',column_maps=maps)
        # Real metadata can quote frame identifiers while Row.as_dict keys differ.
        request={'RequestKey':'goal-seek','OperationId':'allocation','State':'PENDING','Target':1,'LowerBound':0,'UpperBound':.25,'ToleranceCents':1,'MaxIterations':64,'ModelBinding':'binding'}
        data={'Goal Seek Invocation':[request],'Task 5 Headcount Inputs':[{'Department':d,'Headcount':h} for d,h in zip(['Eng','Sales','G&A','Support'],[40,30,20,10])],'Task 5 Allocation Assumptions':[{'Key':'Overhead','Value':100000}]}
        class Frame:
            def __init__(self,title):
                self.rows=data[title];self.columns=['"RequestKey"' if title=='Goal Seek Invocation' and k=='RequestKey' else k for k in self.rows[0]]
            def collect(self):return [SimpleNamespace(as_dict=lambda r=r:r) for r in self.rows]
        output=[];sigma=SimpleNamespace(get_element=Frame,output=lambda name,rows:output.append(rows));session=SimpleNamespace(create_dataframe=lambda rows,schema:rows)
        exec(programs['allocation']['code'],{'sigma':sigma,'session':session})
        self.assertEqual(len(output[-1]),4)
        wrong={'Goal Seek Invocation':{'RequestKey':{'identifier':'"RequestKey"','rowKey':'wrong'}}}
        bad=self.module().emit_python_programs('binding',column_maps=wrong)
        exec(bad['allocation']['code'],{'sigma':sigma,'session':session})
        self.assertEqual(output[-1][0][1:3],['FAILURE','RUNTIME_ROW_KEY_MISMATCH'])

    def test_procedure_datetime_argument_preserves_api_date_constant_and_milliseconds(self):
        p=self.module();contract=p.controlled_procedure_contract('conn')
        contract.update(signature='TASK5_TYPED_ECHO(VARCHAR, NUMBER, NUMBER, BOOLEAN, TIMESTAMP_NTZ)',
            catalogObjectId='catalog-object',lookupEvidenceId='actual-lookup',callerEvidenceId='caller-grants',
            scalarMappingEvidenceId='actual-editor-output',datetimeEvidenceId='actual-timestamp')
        contract['lookup']={'connectionId':'conn','path':list(p.PROCEDURE_PATH),'signature':contract['signature'],
            'returnType':'VARCHAR','parameters':list(p.PROCEDURE_TYPES),'overloads':1,'defaults':False}
        timestamp='2026-10-19T14:30:12.123'
        action=p.build_procedure_action(contract,['timestamp-operation',1234.56,7,True,timestamp])
        # OpenAPI Actions / call-stored-procedure / arguments: the parameter
        # type is datetime, while its fixed timestamp constant uses date.
        self.assertEqual(action['effects'][0]['arguments'][4],{
            'parameterIndex':4,'type':'datetime',
            'value':{'type':'constant','value':{'type':'date','value':timestamp+'Z'}}})

    def test_typed_procedure_rejects_pending_signature_bad_order_types_and_changed_retry(self):
        p=self.module();self.assertTrue(hasattr(p,'build_procedure_action'),'typed procedure mapper is missing')
        contract=p.controlled_procedure_contract('conn')
        args=['task5-sp-0001',1234.56,7,True,'2026-10-15T14:30:00.000']
        with self.assertRaises(ValueError):p.build_procedure_action(contract,args)
        contract.update(signature='TASK5_TYPED_ECHO(VARCHAR, NUMBER, NUMBER, BOOLEAN, TIMESTAMP_NTZ)',
            catalogObjectId='catalog-object',lookupEvidenceId='actual-lookup',callerEvidenceId='caller-grants',
            scalarMappingEvidenceId='actual-editor-output',datetimeEvidenceId='actual-timestamp')
        contract['lookup']={'connectionId':'conn','path':list(p.PROCEDURE_PATH),'signature':contract['signature'],
            'returnType':'VARCHAR','parameters':list(p.PROCEDURE_TYPES),'overloads':1,'defaults':False}
        effect=p.build_procedure_action(contract,args)
        self.assertEqual([a['parameterIndex'] for a in effect['effects'][0]['arguments']],list(range(5)))
        self.assertEqual([a['type'] for a in effect['effects'][0]['arguments']],['text','number','integer','boolean','datetime'])
        self.assertEqual(effect['successToast'],{'showMessage':'hidden'})
        for idx,value in [(0,'\t'),(0,' id '),(1,1234.567),(1,1e15),(2,1.5),(2,-1),(3,1),(4,'2026-02-30T14:30:00.000')]:
            bad=list(args);bad[idx]=value
            with self.assertRaises(ValueError):p.build_procedure_action(contract,bad)
        with self.assertRaises(ValueError):p.build_procedure_action({**contract,'signature':'another signature'},args)
        changed=list(args);changed[2]=8
        with self.assertRaises(ValueError):p.build_procedure_action(contract,changed,prior_invocation=effect['invocation'])
        receipt=p.audit_procedure_result(effect['invocation'],{'status':'FAILURE','operationId':args[0],
            'scalar':None,'dependentEffectsExecuted':False,'successToastShown':False,'evidenceId':'raised'})
        self.assertEqual(receipt['status'],'VERIFIED_COMPONENT_FAILURE');self.assertEqual(receipt['productionStatus'],'BLOCKED')
        with self.assertRaises(ValueError):p.audit_procedure_result(effect['invocation'],{'status':'FAILURE',
            'operationId':args[0],'scalar':None,'dependentEffectsExecuted':True,'successToastShown':False,'evidenceId':'bad'})

class PortGateTests(unittest.TestCase):
    def case(self):
        import custom_ports as p
        self.assertTrue(hasattr(p,'build_custom_port_plan'),'source-bound custom-port gate is missing')
        from migration_manifest import build_manifest
        from presentation import extract_presentation,build_business_input_architecture,build_business_input_plan
        from test_native_actions import staging
        fixture=ROOT/'tests/fixtures/fixture-forecast-model.xlsx'
        extracted=extract_presentation(fixture)
        arch=build_business_input_architecture(extracted,connection_id='conn',folder_id='folder')
        base=build_manifest(fixture,build_business_input_plan(extracted,arch));base['target']={'documentId':'doc','version':8}
        from macro_behavior import analyze_macro_source
        module=ROOT/'tests/fixtures/task5-controlled-business-macros.bas'
        analysis=analyze_macro_source(module)
        reviews={name:{'mappingKind':'reviewed-custom-port','reviewStatus':'approved','reviewer':'Named reviewer',
            'decisionId':'controlled-'+name,'procedureId':proc['procedureId'], 'procedureSha256':proc['procedureSha256'],
            'acknowledgeUnsupportedImperative':True,'sourceOrigin':'new-controlled-fixture'}
            for name,pname in [('allocation','AllocateOverheadByHeadcount'),('goalSeek','SolveForTarget'),('procedure','Task5TypedEcho')]
            for proc in analysis['procedures'] if proc['name']==pname}
        scope={'status':'APPROVED_STAGING','documentId':'doc','version':8,'folderId':'folder','connectionId':'conn','owner':'reviewer'}
        manifest=p.build_custom_port_plan(base,staging(base),module_path=module,
            fixture_path=ROOT/'tests/fixtures/task5-controlled-business-fixture.xlsx',
            expected_path=ROOT/'tests/fixtures/task5-controlled-business-expected.json',reviews=reviews,scope=scope)
        receipt=staging(manifest)
        current={'workbookId':'doc','documentVersion':8,'folderId':'folder','name':'Controlled stage',
            'contents':{'kind':'workbook','schemaVersion':1,'pages':[{'id':'existing','name':'Existing'}],
            'elements':[{'id':eid,'kind':'table','name':title,'source':{'kind':'sql','connectionId':'conn','statement':'reviewed'},'columns':[]}
                for eid,title in [('actuals','Actuals'),('view','Forecast Inputs View'),('defs','Forecast Line Definitions'),
                    ('invocation','Goal Seek Invocation'),('hc','Task 5 Headcount Inputs'),('overhead','Task 5 Allocation Assumptions'),('forecast','Forecast'),('pnl','PnL Combined')]],
            'layout':'<Page id="existing"><Element elementId="actuals" /></Page>','opaquePreserved':{'keep':True}}}
        sources={e['name']:{'elementId':e['id'],'elementHash':p._digest(e)} for e in current['contents']['elements']}
        return p,manifest,receipt,current,sources

    def test_explicit_custom_review_required_and_manifest_cannot_complete_from_unit_evidence(self):
        p,m,receipt,current,sources=self.case()
        self.assertTrue(p.validate_custom_port_plan(m))
        with self.assertRaisesRegex(ValueError,'runtime|acceptance'):
            p.validate_custom_port_completion(m)
        m['planProjection']['customPortArchitecture']['reviews']['allocation']['mappingKind']='native'
        with self.assertRaises(ValueError):p.validate_custom_port_plan(m)

    def test_rest_additions_preserve_unrelated_and_require_hash_bound_dryrun_fresh_version_no_replay(self):
        p,m,receipt,current,sources=self.case()
        self.assertTrue(hasattr(p,'prepare_port_update'),'guarded REST authoring is missing')
        original=copy.deepcopy(current)
        from port_phase_fixtures import context
        captured=context(m,receipt,current)
        candidate=p.prepare_port_update(m,receipt,current,sources,phase_context=captured)
        self.assertEqual(current,original)
        proposed=candidate['dryRunRequest']['contents']
        self.assertEqual(proposed['elements'][:len(current['contents']['elements'])],current['contents']['elements'])
        self.assertEqual(proposed['opaquePreserved'],{'keep':True})
        dry={'httpStatus':200,'dryRun':True,'warnings':[],'proposalHash':candidate['proposalHash'],'evidenceId':'actual-dry-run'}
        req=p.finalize_port_update(candidate,current,dry,fresh_phase_context=captured)
        self.assertEqual(req['body']['documentVersion'],8)
        for fresh in [{**current,'documentVersion':9},{**current,'workbookId':'other'},
                {**current,'contents':{**current['contents'],'opaquePreserved':{'keep':False}}}]:
            with self.assertRaises(ValueError):p.finalize_port_update(candidate,fresh,dry,fresh_phase_context=captured)
        with self.assertRaises(ValueError):p.finalize_port_update(candidate,current,{**dry,'proposalHash':'other'},fresh_phase_context=captured)
        with self.assertRaises(ValueError):p.finalize_port_update(candidate,current,dry,consumed_proposals=[candidate['proposalHash']],fresh_phase_context=captured)
        with self.assertRaises(ValueError):p.finalize_port_update(candidate,current,{**dry,'warnings':['review me']},fresh_phase_context=captured)
        unrelated=copy.deepcopy(current);unrelated['contents']['elements'].append({'id':'task5AllocationPython','kind':'text','text':'unknown'})
        with self.assertRaises(ValueError):p.prepare_port_update(m,receipt,unrelated,sources)
        tampered=copy.deepcopy(candidate);tampered['dryRunRequest']['contents']['elements'][0]['source']['statement']='changed'
        with self.assertRaises(ValueError):p.finalize_port_update(tampered,current,dry,fresh_phase_context=captured)

class InvocationActionTests(unittest.TestCase):
    def test_invocation_actions_only_write_approved_single_assumption_after_guard(self):
        import custom_ports as p
        self.assertTrue(hasattr(p,'build_goal_seek_actions'),'goal seek native action compiler is missing')
        model=json.loads((ROOT/'tests/fixtures/task5-controlled-model.json').read_text())
        request=dict(RequestKey='goal-seek',OperationId='op-A',State='PENDING',Target=2000000,LowerBound=0,
            UpperBound=.25,ToleranceCents=1,MaxIterations=64,ModelBinding='model-A')
        result=p.solve_forecast(model['actuals'],model['inputs'],model['definitions'],request,expected_binding='model-A')
        binding={'invocation':{'elementId':'inv','name':'Goal Seek Invocation',
            'columnIds':{k:'inv-'+k for k in request}},'assumption':{'elementId':'business','name':'Forecast Business Inputs',
            'columnIds':{'Line Item':'line','Monthly Growth':'growth','Input Row ID':'rowid'}},
            'guard':{'elementId':'guard','name':'Goal Seek Apply Guard','columnIds':{k:'guard-'+k for k,t in p.OUTPUT_SCHEMAS['applyGuard']}},
            'inputView':{'elementId':'view','name':'Forecast Inputs View'}}
        actions=p.build_goal_seek_actions(binding,request,model['actuals'],model['inputs'],model['definitions'],result,
                                         expected_binding='model-A')
        self.assertEqual([e['effect'] for e in actions['compute']['effects']],['update-rows','run-python-element'])
        self.assertEqual(actions['compute']['effects'][0]['values']['inv-State']['value']['value'],'PENDING')
        apply=actions['apply']['effects'][0]
        self.assertEqual(apply['effect'],'if-else');self.assertNotIn('else',apply)
        writes=apply['if']['effects'];self.assertEqual(len(writes),1)
        self.assertEqual(writes[0]['tableElementId'],'business')
        self.assertEqual(set(writes[0]['values']),{'growth'})
        self.assertEqual(set(writes[0]['whichRows']['primaryKeys']),{'line','rowid'})
        self.assertEqual(actions['apply']['successToast'],{'showMessage':'hidden'})
        self.assertNotIn('APPLIED',json.dumps(actions['apply']))
        with self.assertRaises(ValueError):p.build_goal_seek_actions(binding,{**request,'OperationId':'op-B'},
            model['actuals'],model['inputs'],model['definitions'],result,expected_binding='model-A')

    def test_controlled_workbook_equations_match_solver_at_endpoints_and_solution(self):
        from openpyxl import load_workbook
        from formula_graph import FormulaGraph,CellIdentity
        fixture=ROOT/'tests/fixtures/task5-controlled-business-fixture.xlsx'
        for growth,want in [(0,482220.9186962107),(.25,4689314.644282148),(.114270488732,1999999.9961722782)]:
            wb=load_workbook(fixture);wb['Assumptions']['B5']=growth
            values=FormulaGraph.from_workbook(wb).evaluate();wb.close()
            self.assertAlmostEqual(values[CellIdentity('P&L','N33')],want,places=6)

class InputAndModelBindingTests(unittest.TestCase):
    def test_port_inputs_use_shared_constructor_published_explore_and_preserve_forecast_spec(self):
        import custom_ports as p
        self.assertTrue(hasattr(p,'build_port_input_specs'),'port input specs are missing')
        from input_sync import emit_construction_program,build_artifact_records
        from test_native_actions import staging
        _,manifest,_,_,_=PortGateTests().case()
        original=copy.deepcopy(manifest['planProjection']['sharedInputArchitecture']['inputs']['forecastBusinessInput'])
        specs=p.build_port_input_specs(original)
        self.assertEqual(original,manifest['planProjection']['sharedInputArchitecture']['inputs']['forecastBusinessInput'])
        for spec in specs.values():self.assertEqual(spec['editMode'],'published-explore')
        self.assertEqual([r['Headcount'] for r in specs['task5Headcount']['rows']],[40,30,20,10])
        self.assertEqual(specs['task5Invocation']['rows'][0]['State'],'UNPREPARED')
        # A real emitted constructor is type-safe at the SDK boundary.
        from migration_manifest import build_manifest
        projection=copy.deepcopy(manifest['planProjection'])
        projection['sharedInputArchitecture']['inputs'].update(specs)
        projection['formulaCells']=copy.deepcopy(manifest['formulaCells'])
        rebuilt=build_manifest(manifest['source']['path'],projection);rebuilt['target']=manifest['target']
        receipt=staging(rebuilt)
        program=emit_construction_program(rebuilt,receipt,'task5Invocation')
        self.assertIn("created.set_edit_mode('published-explore')",program)
        self.assertNotIn("created.set_edit_mode({",program)

class FollowupAuthoringTests(unittest.TestCase):
    def setup_followup(self):
        import custom_ports as p
        self.assertTrue(hasattr(p,'prepare_goal_action_update'),'follow-up action guard is missing')
        p,m,receipt,current,sources=PortGateTests().case()
        inv=next(e for e in current['contents']['elements'] if e['id']=='invocation')
        inv_fields=['RequestKey','OperationId','State','Target','LowerBound','UpperBound','ToleranceCents','MaxIterations','ModelBinding']
        inv['kind']='input-table';inv['source']={'kind':'empty','connectionId':'conn'}
        inv['columns']=[{'id':'inv-'+n,'name':n,'type':'text' if n in ('RequestKey','OperationId','State','ModelBinding') else 'number'} for n in inv_fields]
        sources['Goal Seek Invocation']['elementHash']=p._digest(inv)
        business={'id':'forecastBusinessInput','kind':'input-table','name':'Forecast Business Inputs',
            'source':{'kind':'empty','connectionId':'conn'},'columns':[{'id':'ID' if n=='Input Row ID' else receipt['targetBindings']['forecastBusinessInput']['columnIds'][n],
                'name':n,'type':'text' if n!='Monthly Growth' else 'number'} for n in ('Line Item','Input Row ID','Monthly Growth')]}
        current['contents']['elements'].append(business)
        from input_sync import build_artifact_records
        receipt['targetBindings']['task5Invocation'].update(elementId=inv['id'],
            columnIds={c['name']:c['id'] for c in inv['columns']})
        # Include the two reconciliation fields present in the approved input schema.
        for n in ('AppliedOperationId','AppliedGrowth'):
            inv['columns'].append({'id':'inv-'+n,'name':n,'type':'text' if n=='AppliedOperationId' else 'number'})
            receipt['targetBindings']['task5Invocation']['columnIds'][n]='inv-'+n
        receipt['artifacts']=build_artifact_records(m,m['planProjection']['sharedInputArchitecture'],receipt['targetBindings'])
        sources['Goal Seek Invocation']['elementHash']=p._digest(inv)
        from port_phase_fixtures import complete
        m,receipt,created,initial=complete(m,receipt,current,sources)
        model=json.loads((ROOT/'tests/fixtures/task5-controlled-model.json').read_text())
        request=dict(RequestKey='goal-seek',OperationId='op-A',State='APPLY_PENDING',Target=2000000,
            LowerBound=0,UpperBound=.25,ToleranceCents=1,MaxIterations=64,ModelBinding=initial['modelBinding'])
        result=p.solve_forecast(model['actuals'],model['inputs'],model['definitions'],request,expected_binding=initial['modelBinding'])
        guard=next(e for e in created['contents']['elements'] if e['id']=='task5GoalSeekApplyGuard')
        binding={'invocation':{'elementId':'invocation','name':inv['name'],'columnIds':{c['name']:c['id'] for c in inv['columns']}},
            'assumption':{'elementId':business['id'],'name':business['name'],'columnIds':{c['name']:c['id'] for c in business['columns']}},
            'guard':{'elementId':guard['id'],'name':guard['name'],'columnIds':{c['name']:c['id'] for c in guard['columns']}},
            'inputView':{'elementId':'view','name':'Forecast Inputs View'}}
        return p,m,receipt,created,initial,binding,request,model,result

    def test_followup_action_update_is_guarded_append_and_rejects_changed_model_or_foreign_action(self):
        p,m,receipt,created,initial,binding,request,model,result=self.setup_followup()
        follow=p.prepare_goal_action_update(m,receipt,created,initial,binding,request,model['actuals'],model['inputs'],model['definitions'],result)
        self.assertEqual(follow['documentVersion'],11)
        self.assertEqual(follow['dryRunRequest']['contents']['elements'][:-3],created['contents']['elements'])
        self.assertEqual(len(follow['addedElementIds']),3)
        # Failure must stop before accepting any unverified action target binding.
        with self.assertRaises(ValueError):p.prepare_goal_action_update(m,receipt,created,initial,{},request,
            model['actuals'],model['inputs'],model['definitions'],result)
        changed=copy.deepcopy(created);changed['contents']['elements'][0]['source']['statement']='changed'
        with self.assertRaisesRegex(ValueError,'MODEL|source|SOURCE'):
            p.prepare_goal_action_update(m,receipt,changed,initial,{},request,model['actuals'],model['inputs'],model['definitions'],result)

    def test_followup_rejects_foreign_guard_invocation_and_input_view_even_same_connection(self):
        p,m,receipt,created,initial,binding,request,model,result=self.setup_followup()
        for role in ('guard','invocation','inputView'):
            for connection in ('conn','other-connection'):
                with self.subTest(role=role,connection=connection):
                    current=copy.deepcopy(created);bad=copy.deepcopy(binding)
                    source=next(e for e in current['contents']['elements'] if e['id']==binding[role]['elementId'])
                    foreign=copy.deepcopy(source)
                    foreign.update(id='foreign'+role,name='Unreviewed '+role,
                        source={'kind':'empty' if role=='invocation' else 'sql','connectionId':connection,
                                **({} if role=='invocation' else {'statement':'SELECT stale guards'})})
                    current['contents']['elements'].append(foreign)
                    bad[role].update(elementId=foreign['id'],name=foreign['name'])
                    with self.assertRaises(ValueError):
                        p.prepare_goal_action_update(m,receipt,current,initial,bad,request,
                            model['actuals'],model['inputs'],model['definitions'],result)

    def test_followup_rejects_duplicate_runtime_titles_and_wrong_assigned_invocation(self):
        p,m,receipt,created,initial,binding,request,model,result=self.setup_followup()
        for title in ('Goal Seek Apply Guard','Goal Seek Invocation','Forecast Inputs View','Goal Seek Result','Forecast Business Inputs'):
            with self.subTest(title=title):
                current=copy.deepcopy(created)
                duplicate=copy.deepcopy(next(e for e in current['contents']['elements'] if e.get('name')==title))
                duplicate['id']='duplicate-'+duplicate['id'];current['contents']['elements'].append(duplicate)
                with self.assertRaises(ValueError):
                    p.prepare_goal_action_update(m,receipt,current,initial,binding,request,
                        model['actuals'],model['inputs'],model['definitions'],result)
        from input_sync import build_artifact_records
        other=copy.deepcopy(receipt)
        other['targetBindings']['task5Invocation']['elementId']='foreignInvocation'
        other['artifacts']=build_artifact_records(m,m['planProjection']['sharedInputArchitecture'],other['targetBindings'])
        with self.assertRaises(ValueError):
            p.prepare_goal_action_update(m,other,created,initial,binding,request,
                model['actuals'],model['inputs'],model['definitions'],result)

class ProcedureDecimalTransportTests(unittest.TestCase):
    def contract(self):
        import custom_ports as p
        contract=p.controlled_procedure_contract('conn')
        contract.update(signature='TASK5_TYPED_ECHO(VARCHAR, NUMBER, NUMBER, BOOLEAN, TIMESTAMP_NTZ)',
            catalogObjectId='catalog',lookupEvidenceId='lookup',callerEvidenceId='caller',
            scalarMappingEvidenceId='scalar',datetimeEvidenceId='datetime')
        contract['lookup']={k:contract[k] for k in ('connectionId','path','signature','returnType','overloads','defaults')}
        contract['lookup']['parameters']=list(p.PROCEDURE_TYPES)
        return p,contract

    def test_large_accepted_amounts_expect_exact_json_decimal_cents(self):
        p,contract=self.contract()
        for amount in (100000000000000.1,999999999999999.1,100000000000000.12,
                -100000000000000.1,999999999999999.9,0.01,-0.0,0,1234.56):
            with self.subTest(amount=amount):
                action=p.build_procedure_action(contract,['op',amount,7,True,'2026-10-15T14:30:00.000'])
                literal=action['effects'][0]['arguments'][1]['value']['value']['value']
                wire=json.dumps(literal,allow_nan=False)
                decimal=Decimal(wire)
                if decimal==0:decimal=Decimal(0)
                expected='operation_id=op|amount='+format(decimal,'.2f')+'|count=7|enabled=true|occurred_at=2026-10-15T14:30:00.000'
                self.assertEqual(action['invocation']['expectedScalar'],expected)
                accepted=p.audit_procedure_result(action['invocation'],{'status':'SUCCESS','operationId':'op',
                    'scalar':expected,'evidenceId':'actual-runtime-required'})
                self.assertEqual(accepted['status'],'VERIFIED_COMPONENT_SUCCESS')
                self.assertEqual(json.dumps(json.loads(wire)),wire)

    def test_decimal_inputs_that_lose_cents_in_json_double_conversion_are_rejected(self):
        p,contract=self.contract()
        for amount in (Decimal('999999999999999.11'),Decimal('100000000000000.13'),
                Decimal('999999999999999.99'),Decimal('1000000000000000'),Decimal('1.001')):
            with self.subTest(amount=str(amount)),self.assertRaises(ValueError):
                p.build_procedure_action(contract,['op',amount,7,True,'2026-10-15T14:30:00.000'])
        # JSON-native numbers which round-trip exactly remain accepted.
        for amount in (Decimal('1234.56'),Decimal('999999999999999.10')):
            action=p.build_procedure_action(contract,['op',amount,7,True,'2026-10-15T14:30:00.000'])
            literal=action['effects'][0]['arguments'][1]['value']['value']['value']
            self.assertEqual(Decimal(json.dumps(literal)),amount)
            reloaded=json.loads(json.dumps(action,allow_nan=False))
            self.assertEqual(reloaded['invocation']['expectedScalar'],action['invocation']['expectedScalar'])

if __name__=='__main__':unittest.main()
