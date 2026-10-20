import copy
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts" / "lib"))
from source_workbook import build_source_model, validate_source_model  # noqa: E402


class SourceWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.source = Path(self.tmp.name) / "source.xlsx"

    def save(self, workbook):
        workbook.save(self.source)
        return hashlib.sha256(self.source.read_bytes()).hexdigest()

    @staticmethod
    def cells(model):
        return {(c["sheet"], c["address"]): c for c in model["cells"]}

    def test_source_facts_and_table_boundaries_are_preserved(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "O'Brien & Plan"
        ws.append(["Text", "Number", "Calc"])
        ws.append(["a", 0, "=B2+1"])
        ws.append(["b", 2, 9])
        ws.append(["Total", None, "=SUM(C2:C3)"])
        ws["A2"].value = "=not a formula"
        ws["A2"].data_type = "s"
        ws["E1"] = False
        ws["E2"] = date(2026, 1, 2)
        ws["E3"] = datetime(2026, 1, 2, 3, 4, 5)
        ws["E4"] = ""
        ws.merge_cells("G1:H2")
        ws.protection.sheet = True
        table = Table(displayName="tblPlan", ref="A1:C4", totalsRowCount=1)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9")
        ws.add_table(table)
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3))
        ws.add_chart(chart, "J1")
        hidden = wb.create_sheet("Scratch")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "=1+2"
        before = self.save(wb)

        model = build_source_model(self.source)
        self.assertEqual(model, build_source_model(self.source))
        self.assertEqual(json.loads(json.dumps(model)), model)
        self.assertEqual(model["sourceSha256"], before)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), before)
        self.assertEqual(model["counts"], {"formulas": 3, "charts": 1, "tables": 1})
        cells = self.cells(model)
        self.assertEqual(cells[(ws.title, "A2")]["kind"], "literal")
        self.assertEqual(cells[(ws.title, "A2")]["value"], "=not a formula")
        self.assertEqual(cells[(ws.title, "B2")]["value"], 0)
        self.assertIs(cells[(ws.title, "E1")]["value"], False)
        self.assertEqual(cells[(ws.title, "E2")]["value"], {"type": "date", "iso": "2026-01-02"})
        self.assertEqual(cells[(ws.title, "E3")]["value"], {"type": "datetime", "iso": "2026-01-02T03:04:05"})
        self.assertEqual(cells[(ws.title, "C2")]["expression"], "=B2+1")
        self.assertEqual(cells[("Scratch", "A1")]["expression"], "=1+2")
        self.assertEqual(model["sheets"][0]["mergedRanges"], ["G1:H2"])
        self.assertTrue(model["sheets"][0]["protected"])
        self.assertEqual(model["sheets"][1]["state"], "hidden")
        table_model = model["tables"][0]
        self.assertEqual(table_model["headers"], ["Text", "Number", "Calc"])
        self.assertEqual(table_model["bodyRange"], "A2:C3")
        self.assertEqual(table_model["totalsRange"], "A4:C4")
        self.assertEqual(table_model["columns"][2]["formulaCells"], ["C2"])
        self.assertEqual(table_model["columns"][2]["literalCells"], ["C3"])
        self.assertEqual(table_model["columns"][2]["intent"], "ambiguous")
        validate_source_model(model, self.source)

    def test_distant_styled_empty_cell_does_not_expand_inventory_rectangle(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "=1+2"
        ws["Z1000000"].font = Font(bold=True)
        before = self.save(wb)
        model = build_source_model(self.source)
        self.assertEqual([(cell["address"], cell["kind"]) for cell in model["cells"]], [("A1", "formula")])
        self.assertEqual(model["sourceSha256"], before)

    def test_inline_empty_string_is_distinct_from_absent_and_styled_empty_cells(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = ""
        ws["B1"].font = Font(bold=True)
        self.save(wb)
        cells = self.cells(build_source_model(self.source))
        self.assertEqual(cells[("Sheet", "A1")]["dataType"], "inlineStr")
        self.assertEqual(cells[("Sheet", "A1")]["value"], "")
        self.assertEqual(cells[("Sheet", "A1")]["kind"], "literal")
        self.assertNotIn(("Sheet", "B1"), cells)
        self.assertNotIn(("Sheet", "C1"), cells)

    def test_array_formula_text_survives_and_other_formula_objects_are_unresolved(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = ArrayFormula("A1:A2", text="=SUM(B1:B2)")
        ws["D1"] = DataTableFormula("D1:D2", r1="B1")
        self.save(wb)
        model = build_source_model(self.source)
        cells = self.cells(model)
        array = cells[("Sheet", "A1")]
        self.assertEqual(array["kind"], "formula")
        self.assertEqual(array["expression"], "=SUM(B1:B2)")
        self.assertEqual(array["formulaObjectType"], "ArrayFormula")
        self.assertEqual([d["address"] for d in array["dependencies"]], ["B1", "B2"])
        data_table = cells[("Sheet", "D1")]
        self.assertEqual(data_table["kind"], "formula")
        self.assertEqual(data_table["formulaObjectType"], "DataTableFormula")
        self.assertEqual(data_table["dependencyStatus"], "unresolved")
        self.assertTrue(any(u["cellId"] == data_table["id"] for u in model["unresolved"]))

    def test_dependencies_include_range_interiors_blanks_cycles_and_names(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Calc"
        ws["A1"] = "=SUM(B1:B3)"
        ws["A2"] = "='O''Brien'!C7"
        ws["A3"] = "=A4"
        ws["A4"] = "=A3"
        ws["A5"] = "=Rate"
        ws["A6"] = "=Global"
        ws["A7"] = "=Missing"
        ws["A8"] = "=SUM(Table1[Amount])"
        ws["A9"] = "='[other.xlsx]Sheet1'!A1"
        ws["A10"] = "=Ghost!A1"
        ws["A11"] = "=Constant"
        other = wb.create_sheet("O'Brien")
        other["C7"] = 4
        wb.defined_names.add(DefinedName("Rate", attr_text="'O''Brien'!$C$7"))
        wb.defined_names.add(DefinedName("Global", attr_text="Calc!$B$1"))
        wb.defined_names.add(DefinedName("Constant", attr_text="42"))
        ws.defined_names.add(DefinedName("Rate", attr_text="Calc!$B$3", localSheetId=0))
        self.save(wb)
        model = build_source_model(self.source)
        cells = self.cells(model)

        def targets(address):
            return {(d["sheet"], d["address"], d["targetKind"]) for d in cells[("Calc", address)]["dependencies"]}

        self.assertEqual(targets("A1"), {("Calc", "B1", "blank"), ("Calc", "B2", "blank"), ("Calc", "B3", "blank")})
        self.assertEqual(targets("A2"), {("O'Brien", "C7", "populated")})
        self.assertEqual(targets("A3"), {("Calc", "A4", "populated")})
        self.assertEqual(targets("A4"), {("Calc", "A3", "populated")})
        self.assertEqual(targets("A5"), {("Calc", "B3", "blank")})
        self.assertEqual(targets("A6"), {("Calc", "B1", "blank")})
        self.assertEqual(cells[("Calc", "A7")]["dependencyStatus"], "unresolved")
        self.assertEqual(cells[("Calc", "A8")]["dependencyStatus"], "unresolved")
        self.assertEqual(cells[("Calc", "A9")]["dependencyStatus"], "unresolved")
        self.assertEqual(cells[("Calc", "A10")]["dependencyStatus"], "unresolved")
        self.assertEqual(cells[("Calc", "A11")]["dependencyStatus"], "unresolved")
        self.assertGreaterEqual(len(model["unresolved"]), 5)
        validate_source_model(model, self.source)

    def test_reviewed_blocks_keep_cells_and_reject_invalid_decisions(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        ws["A1"] = "=Scratch!A1"
        scratch = wb.create_sheet("Scratch")
        scratch.sheet_state = "hidden"
        scratch["A1"] = "=Main!A1"
        sha = self.save(wb)
        decision = {"sourceSha256": sha, "blocks": [{"sheet": "Scratch", "range": "A1:B2", "role": "scratch"}]}
        model = build_source_model(self.source, block_decisions=decision)
        self.assertEqual(model["counts"]["formulas"], 2)
        self.assertEqual(self.cells(model)[("Main", "A1")]["dependencies"][0]["sheet"], "Scratch")
        self.assertEqual(next(b for b in model["blocks"] if b["reviewed"])["role"], "scratch")
        for changed in [
            {**decision, "sourceSha256": "0" * 64},
            {**decision, "blocks": [{"sheet": "Unknown", "range": "A1", "role": "scratch"}]},
            {**decision, "blocks": [{"sheet": "Scratch", "range": "A1:B2", "role": "edit"}]},
            {**decision, "blocks": [{"sheet": "Scratch", "range": "A:A", "role": "scratch"}]},
            {**decision, "blocks": [{"sheet": "Scratch", "range": "B2:A1", "role": "scratch"}]},
            {**decision, "blocks": [decision["blocks"][0], {"sheet": "Scratch", "range": "B2:C3", "role": "input"}]},
        ]:
            with self.subTest(changed=changed), self.assertRaises(ValueError):
                build_source_model(self.source, block_decisions=changed)

    def test_validator_and_cli_reject_tampering_without_overwriting_output(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = 1
        ws["B1"] = "=A1"
        self.save(wb)
        model = build_source_model(self.source)
        for mutate in [
            lambda m: m["cells"].append(copy.deepcopy(m["cells"][0])),
            lambda m: m.update(modelVersion=2),
            lambda m: m["counts"].update(formulas=0),
            lambda m: m["cells"][1].update(kind="literal"),
            lambda m: m["cells"][1]["dependencies"][0].update(targetId="bad"),
            lambda m: m["cells"][1]["dependencies"][0].update(targetKind="blank"),
            lambda m: m["cells"][1].update(dependencies=[]),
            lambda m: m["sheets"][0].update(id="bad"),
            lambda m: m["cells"][0].update(id="bad"),
            lambda m: m["counts"].update(charts=True),
        ]:
            bad = copy.deepcopy(model)
            mutate(bad)
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                validate_source_model(bad, self.source)
        out = Path(self.tmp.name) / "model.json"
        cli = ROOT / "scripts" / "source-workbook.py"
        okay = subprocess.run([sys.executable, str(cli), str(self.source), "--out", str(out)], capture_output=True, text=True)
        self.assertEqual(okay.returncode, 0, okay.stderr)
        self.assertEqual(json.loads(out.read_text()), model)
        out.write_text("sentinel")
        bad_decisions = Path(self.tmp.name) / "bad.json"
        bad_decisions.write_text(json.dumps({"sourceSha256": "0" * 64, "blocks": []}))
        failed = subprocess.run([sys.executable, str(cli), str(self.source), "--out", str(out), "--block-decisions", str(bad_decisions)], capture_output=True, text=True)
        self.assertNotEqual(failed.returncode, 0)
        self.assertEqual(out.read_text(), "sentinel")
        self.source.write_bytes(self.source.read_bytes() + b" ")
        with self.assertRaises(ValueError):
            validate_source_model(model, self.source)

    def test_validator_cannot_relabel_unsupported_formula_as_resolved(self):
        wb = Workbook()
        wb.active["A1"] = "=SUM(Table1[Amount])"
        self.save(wb)
        model = build_source_model(self.source)
        model["cells"][0]["dependencyStatus"] = "resolved"
        model["unresolved"] = []
        with self.assertRaises(ValueError):
            validate_source_model(model)

    def test_unresolved_identity_must_match_cell_and_reason_ordinal(self):
        wb = Workbook()
        wb.active["A1"] = "=Missing+Other"
        self.save(wb)
        model = build_source_model(self.source)
        self.assertEqual(len(model["unresolved"]), 2)
        for replacement in ("arbitrary", 'unresolved:["cell:[\\"Sheet\\",\\"A1\\"]",2]'):
            bad = copy.deepcopy(model)
            bad["unresolved"][0]["id"] = replacement
            with self.subTest(replacement=replacement), self.assertRaises(ValueError):
                validate_source_model(bad, self.source)

    def test_cli_refuses_source_as_output_through_direct_and_alias_paths(self):
        wb = Workbook()
        wb.active["A1"] = 7
        before = self.save(wb)
        alias_symlink = Path(self.tmp.name) / "symlink.xlsx"
        alias_symlink.symlink_to(self.source)
        alias_hardlink = Path(self.tmp.name) / "hardlink.xlsx"
        alias_hardlink.hardlink_to(self.source)
        cli = ROOT / "scripts" / "source-workbook.py"
        for output in (self.source, alias_symlink, alias_hardlink):
            with self.subTest(output=output):
                result = subprocess.run([sys.executable, str(cli), str(self.source), "--out", str(output)],
                                        capture_output=True, text=True)
                self.assertNotEqual(result.returncode, 0)
                self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), before)

    def test_validation_rejects_stale_or_overlapping_review_bindings(self):
        wb = Workbook()
        wb.active["A1"] = 1
        sha = self.save(wb)
        decisions = {"sourceSha256": sha, "blocks": [{"sheet": "Sheet", "range": "A1:B2", "role": "input"}]}
        model = build_source_model(self.source, block_decisions=decisions)
        stale = copy.deepcopy(model)
        stale["blocks"][0]["sourceSha256"] = "0" * 64
        with self.assertRaises(ValueError):
            validate_source_model(stale)
        overlap = copy.deepcopy(model)
        overlap["blocks"].append({"id": 'block:["reviewed","Sheet","B2:C3"]',
                                  "sheet": "Sheet", "range": "B2:C3", "role": "scratch",
                                  "reviewed": True, "sourceSha256": sha})
        with self.assertRaises(ValueError):
            validate_source_model(overlap)


if __name__ == "__main__":
    unittest.main()
