#!/usr/bin/env python3
from __future__ import annotations

import json
import copy
import importlib.util
import re
import sys
import subprocess
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / "lib"))

import presentation  # noqa: E402
from input_sync import (build_artifact_records, build_shared_input_architecture,
                        include_research_line_input_literals)  # noqa: E402
from input_orchestration import gate_input_staging  # noqa: E402
from migration_gate import gate_plan_staging_source, source_issue_hash  # noqa: E402
from migration_manifest import build_manifest  # noqa: E402


FIXTURES = ROOT / "tests" / "fixtures"


def research_runtime_plan():
    spec = importlib.util.spec_from_file_location(
        "presentation_research_inference", SCRIPTS / "infer-canonical-formulas.py")
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    plan = module.infer(FIXTURES / "fixture-research-model.xlsx")
    include_research_line_input_literals(plan)
    return plan


def presentation_openapi_fixture():
    return {"info": {"version": "portable-current-contract"},
            "paths": {"/v2/workbooks/{workbookId}/contents": {}},
            "components": {"schemas": {
                "Kinds": {"enum": ["bar-chart", "line-chart", "single",
                                     "backgroundScale", "fontScale", "number",
                                     "datetime", "center", "hidden", "none"]},
                "Chart": {"properties": {"xAxis": {}, "yAxis": {}, "columns": {},
                                                   "columnIds": {}}},
                "Validation": {"properties": {"values": {}, "valuesFrom": {}, "range": {}}},
                "Conditional": {"properties": {"columnIds": {}, "scheme": {}, "domain": {},
                                                         "order": {}, "nullBehavior": {}}},
                "Number": {"properties": {"formatString": {}, "currencySymbol": {}}},
                "Text": {"properties": {"verticalAlign": {}, "text": {}, "visibility": {}}},
            }}}


class TestPresentationExtraction(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.plans = {
            path.name: presentation.extract_presentation(path)
            for path in sorted(FIXTURES.glob("fixture-*.xlsx"))
        }

    def test_extracts_all_native_charts_with_source_relationships_and_series_order(self):
        charts = [chart for plan in self.plans.values() for chart in plan["charts"]]
        self.assertEqual(len(charts), 10)
        self.assertTrue(all(chart["sourceLocation"]["worksheet"] for chart in charts))
        self.assertTrue(all(chart["sourceLocation"]["drawing"] for chart in charts))
        self.assertEqual(
            [series["name"] for series in self.plans["fixture-forecast-model.xlsx"]["charts"][0]["series"]],
            ["Revenue", "EBITDA", "Net Income"],
        )
        self.assertEqual(
            [series["name"] for series in self.plans["fixture-research-model.xlsx"]["charts"][1]["series"]],
            ["EPS", "Net debt / EBITDA"],
        )
        self.assertEqual(
            self.plans["fixture-sales-pipeline.xlsx"]["charts"][1]["categoryRef"],
            "'Overview'!$D$12:$D$21",
        )
        self.assertEqual(
            self.plans["fixture-inventory-reorder.xlsx"]["charts"][1]["unit"], "currency")

    def test_current_openapi_supports_emitted_alignment_description_and_style_contract(self):
        audit = presentation.audit_presentation_openapi(
            presentation_openapi_fixture())
        self.assertEqual(audit["status"], "PASS")
        self.assertTrue(audit["capabilities"]["presentationTextAndDescription"])

    def test_extracts_validation_contracts_and_exact_whole_number_list(self):
        inventory = self.plans["fixture-inventory-reorder.xlsx"]
        lead_days = next(rule for rule in inventory["validations"] if rule["range"] == "E2:E37")
        self.assertEqual(lead_days["sourceType"], "whole")
        self.assertEqual(lead_days["emission"], {"values": list(range(1, 91))})
        for invalid in (1.5, 0, 91):
            self.assertFalse(presentation.validation_accepts(lead_days, invalid))
        self.assertTrue(presentation.validation_accepts(lead_days, 1))
        self.assertTrue(presentation.validation_accepts(lead_days, 90))

        sales = self.plans["fixture-sales-pipeline.xlsx"]
        probability = next(rule for rule in sales["validations"] if rule["range"] == "I2:I49")
        self.assertEqual(probability["emission"], {"range": {"min": 0.0, "max": 1.0}})
        self.assertIsNone(probability["allowBlank"])
        self.assertEqual(probability["sourceXml"]["rawAttributes"]["allowBlank"],
                         {"present": False, "value": None})

    def test_preserves_conditional_rule_priority_relative_references_and_scales(self):
        rules = [rule for plan in self.plans.values() for rule in plan["conditionalFormats"]]
        priorities = {(rule["sheet"], rule["range"], rule["priority"]) for rule in rules}
        self.assertIn(("Expense Register", "O2:O49", 1), priorities)
        self.assertIn(("Expense Register", "O2:O49", 3), priorities)
        self.assertTrue(all("relativeReference" in rule for rule in rules))
        scales = [rule for rule in rules if rule["sourceType"] == "colorScale"]
        self.assertGreaterEqual(len(scales), 3)
        self.assertTrue(all(rule["status"] == "unresolved" for rule in scales))
        self.assertTrue(all(rule["emission"] is None for rule in scales))
        self.assertTrue(all('type="percentile"' in rule["sourceRuleXml"] for rule in scales))

        expense = self.plans["fixture-expense-approval.xlsx"]
        status_rules = [rule for rule in expense["conditionalFormats"]
                        if rule["range"] == "O2:O49"]
        self.assertEqual(
            [rule["emission"]["style"]["backgroundColor"] for rule in status_rules],
            ["#FCE4D6", "#FFF2CC", "#E2F0D9"],
        )

    def test_chart_series_unit_follows_direct_source_reference_format(self):
        research = self.plans["fixture-research-model.xlsx"]
        chart = next(chart for chart in research["charts"] if chart["title"] == "EPS and leverage")
        by_name = {series["name"]: series for series in chart["series"]}
        self.assertEqual(by_name["EPS"]["unit"], "currency")
        self.assertEqual(by_name["EPS"]["numberFormat"], "$0.00;($0.00);-")
        self.assertEqual(by_name["Net debt / EBITDA"]["unit"], "ratio")
        self.assertEqual(by_name["Net debt / EBITDA"]["numberFormat"], "0.0x")
        artifact = presentation.emit_presentation(
            research, presentation.fixture_bindings(research), logical_fixture=True)
        emitted = next(element for element in artifact["elements"]
                       if element.get("name") == "EPS and leverage")
        columns = {column["name"]: column for column in emitted["columns"]}
        self.assertEqual(columns["EPS"]["format"]["formatString"], "($.2f")
        self.assertEqual(columns["Net debt / EBITDA"]["format"]["formatString"], ".1f")     # `0.0x` has no grouping
        self.assertEqual(columns["Net debt / EBITDA"]["format"]["suffix"], "x")

    def test_extracts_source_precision_dates_and_all_formula_provenance(self):
        expected = {
            "fixture-sales-pipeline.xlsx": 170,
            "fixture-expense-approval.xlsx": 632,
            "fixture-inventory-reorder.xlsx": 442,
            "fixture-forecast-model.xlsx": 278,
            "fixture-research-model.xlsx": 343,
        }
        self.assertEqual(
            {name: len(plan["formulaProvenance"]) for name, plan in self.plans.items()}, expected)
        self.assertEqual(
            self.plans["fixture-sales-pipeline.xlsx"]["semanticFormats"]["Pipeline!B2"]["dateGrain"],
            "day",
        )
        self.assertEqual(
            self.plans["fixture-sales-pipeline.xlsx"]["semanticFormats"]["Pipeline!I2"]["precision"],
            0,
        )
        sales_timing = next(
            item for item in self.plans["fixture-sales-pipeline.xlsx"]["formulaProvenance"]
            if item["cell"] == "Pipeline!N2"
        )
        self.assertNotIn("\\\\", sales_timing["formula"])
        chart_cells = [cell for plan in self.plans.values() for chart in plan["charts"]
                       for series in chart["series"] for cell in series["sourceFormulaCells"]]
        self.assertTrue(chart_cells)


class TestPresentationFormatTranslation(unittest.TestCase):
    """Number formats reach Sigma only through ``excel_formats`` (the same translator the table-report builder uses)."""

    @staticmethod
    def semantic(cells, cached=None):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = "S"
        for index, (value, code) in enumerate(cells, 1):
            sheet.cell(index, 1, value).number_format = code
        return presentation._semantic_formats(book, cached)

    def test_cell_formats_follow_the_translator(self):
        got = self.semantic([(1234, "0"), (1234.5, "#,##0.0"), (0.5, "0%"), (-3, "#,##0;[Red](#,##0)"),
                             (3, "€#,##0.00"), (2.5, "0.0x"), (12, "$#,##0_);($#,##0)")])
        emissions = {cell: item["emission"] for cell, item in got.items()}
        self.assertEqual(emissions["S!A1"], {"kind": "number", "formatString": ".0f"})        # no grouping the source lacks
        self.assertEqual(emissions["S!A2"], {"kind": "number", "formatString": ",.1f"})
        self.assertEqual(emissions["S!A3"], {"kind": "number", "formatString": ".0%"})
        self.assertEqual(emissions["S!A4"], {"kind": "number", "formatString": "(,.0f"})
        self.assertEqual(emissions["S!A5"], {"kind": "number", "formatString": ",.2f", "prefix": "€"})
        self.assertEqual(emissions["S!A6"], {"kind": "number", "formatString": ".1f", "suffix": "x"})
        self.assertEqual(emissions["S!A7"], {"kind": "number", "formatString": "($,.0f", "currencySymbol": "$"})
        self.assertEqual({cell: item["semantic"] for cell, item in got.items()},
                         {"S!A1": "number", "S!A2": "number", "S!A3": "percentage", "S!A4": "number",
                          "S!A5": "currency", "S!A6": "number", "S!A7": "currency"})
        self.assertEqual(got["S!A2"]["precision"], 1)
        self.assertTrue(got["S!A4"]["negativeRed"])
        self.assertNotIn("negativeRed", got["S!A1"])

    def test_general_numbers_keep_the_decimals_they_hold(self):
        got = self.semantic([(0.15, "General"), (7, "General"), ("=A1*2", "General")])
        self.assertEqual(got["S!A1"]["emission"], {"kind": "number", "formatString": ",.2f"})
        self.assertEqual(got["S!A2"]["emission"], {"kind": "number", "formatString": ",.0f"})
        self.assertEqual(got["S!A3"]["emission"], {"kind": "number", "formatString": ",.0f"})
        cached = {"S": {"A3": 0.3}}
        class Cached:
            sheetnames = ["S"]
            def __getitem__(self, name):
                return type("Sheet", (), {"__getitem__": lambda _, address: type("C", (), {"value": cached[name][address]})()})()
        self.assertEqual(self.semantic([(0.15, "General"), (7, "General"), ("=A1*2", "General")], Cached())
                         ["S!A3"]["emission"], {"kind": "number", "formatString": ",.1f"})

    def test_dates_use_their_own_mask_including_on_formula_cells(self):
        import datetime as dt
        got = self.semantic([(dt.datetime(2026, 7, 31), "mmm-yy"), (dt.datetime(2026, 7, 31), "yyyy-mm-dd"),
                             ("=A1+30", "mm/dd")])
        self.assertEqual(got["S!A1"]["emission"], {"kind": "datetime", "formatString": "%b-%y"})
        self.assertEqual(got["S!A2"]["emission"], {"kind": "datetime", "formatString": "%Y-%m-%d"})
        self.assertEqual((got["S!A3"]["semantic"], got["S!A3"]["emission"]),
                         ("date", {"kind": "datetime", "formatString": "%m/%d"}))
        self.assertEqual(got["S!A2"]["dateGrain"], "day")

    def test_text_cells_get_no_format(self):
        self.assertEqual(self.semantic([("Revenue", "General"), ("n/a", "@")]), {})

    def test_series_unit_comes_from_the_format_then_the_title(self):
        unit = presentation._unit
        self.assertEqual(unit("$#,##0"), "currency")
        self.assertEqual(unit("€#,##0.00"), "currency")
        self.assertEqual(unit("[$$-409]#,##0"), "currency")
        self.assertEqual(unit("0.0%"), "percentage")
        self.assertEqual(unit("General", "Revenue (USD m)"), "currency")
        self.assertEqual(unit("General", "Margin %"), "percentage")
        self.assertEqual(unit("General", "Order count"), "count")
        self.assertEqual(unit("0.0x", "EPS and leverage"), "ratio")
        self.assertEqual(unit("#,##0"), "number")
        self.assertEqual(unit("mmm-yy"), "number")
        self.assertEqual(unit(None), "number")

    def test_series_formats(self):
        def series(code, unit, values=()):
            return {"numberFormat": code, "unit": unit, "expectedValues": list(values)}
        fmt = presentation._series_format
        self.assertEqual(fmt(series("$#,##0.0;($#,##0.0);-", "currency")),
                         {"kind": "number", "formatString": "($,.1f", "currencySymbol": "$"})
        self.assertEqual(fmt(series("€#,##0", "currency")), {"kind": "number", "formatString": ",.0f", "prefix": "€"})
        self.assertEqual(fmt(series("0.0x", "ratio")), {"kind": "number", "formatString": ".1f", "suffix": "x"})
        self.assertEqual(fmt(series("0.0%", "percentage")), {"kind": "number", "formatString": ".1%"})
        self.assertEqual(fmt(series("General", "number", [1, 2])), {"kind": "number", "formatString": ",.0f"})
        self.assertEqual(fmt(series(None, "number", [1.25, 2])), {"kind": "number", "formatString": ",.2f"})
        # only the title says the unit: the format is restyled, not re-parsed
        self.assertEqual(fmt(series("General", "currency", [10.5])),
                         {"kind": "number", "formatString": "$,.1f", "currencySymbol": "$"})
        self.assertEqual(fmt(series("General", "percentage", [0.25])), {"kind": "number", "formatString": ",.2%"})
        self.assertEqual(fmt(series("#,##0", "count")), {"kind": "number", "formatString": ",.0f"})

    def test_red_negatives_reach_the_table_as_a_conditional_rule_beneath_source_rules(self):
        import shutil
        from openpyxl import load_workbook as load
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "fixture-sales-pipeline.xlsx"
            shutil.copy(FIXTURES / "fixture-sales-pipeline.xlsx", source)
            book = load(source)
            book["Pipeline"]["J2"].number_format = "#,##0;[Red](#,##0)"          # ACV
            book.save(source)
            plan = presentation.extract_presentation(source)
            bindings = presentation.fixture_bindings(plan)
            artifact = presentation.emit_presentation(plan, bindings, logical_fixture=True)
            target = bindings["columns"]["Pipeline!ACV"]
            self.assertEqual(plan["semanticFormats"]["Pipeline!J2"]["negativeRed"], True)
            patch = artifact["elementPatches"][target["elementId"]]
            self.assertEqual(patch["columnUpdates"][target["columnId"]]["format"],
                             {"kind": "number", "formatString": "(,.0f"})
            red = patch["conditionalFormats"][0]                      # later rules win in Sigma: source rules follow it
            self.assertEqual((red["type"], red["condition"], red["value"], red["columnIds"]),
                             ("single", "<", 0, [target["columnId"]]))
            self.assertEqual(sum(1 for rule in patch["conditionalFormats"] if rule == red), 1)
            # a column without [Red] gets no such rule
            plain = bindings["columns"]["Pipeline!Probability"]
            self.assertFalse(any(rule.get("condition") == "<" and plain["columnId"] in rule.get("columnIds", [])
                                 for rule in artifact["elementPatches"][plain["elementId"]]["conditionalFormats"]))


class TestPresentationEmission(unittest.TestCase):
    def test_source_table_rules_reach_both_views_without_cross_sheet_propagation(self):
        expected = {
            "fixture-sales-pipeline.xlsx": {
                "Pipeline": ["sales-input", "sales-main-view", "sales-detail-view"],
            },
            "fixture-inventory-reorder.xlsx": {
                "Inventory": ["inventory-input", "inventory-main-view", "inventory-detail-view"],
                "Purchase Plan": ["inventory-main-view"],
            },
            "fixture-expense-approval.xlsx": {
                "Expense Register": ["expense-input", "expense-main-view", "expense-detail-view"],
                "Reconciliation": ["expense-main-view"],
            },
        }
        for filename, by_sheet in expected.items():
            plan = presentation.extract_presentation(FIXTURES / filename)
            bindings = presentation.fixture_bindings(plan)
            for rule in plan["conditionalFormats"]:
                if rule["status"] != "supported":
                    continue
                with self.subTest(filename=filename, source=rule["sheet"], rule=rule["id"]):
                    self.assertEqual(
                        [target["elementId"] for target in bindings["conditionalTargets"][rule["id"]]],
                        by_sheet[rule["sheet"]],
                    )
            artifact = presentation.emit_presentation(plan, bindings, logical_fixture=True)
            source_sheet = next(iter(by_sheet))
            source_rules = [rule for rule in plan["conditionalFormats"]
                            if rule["sheet"] == source_sheet and rule["status"] == "supported"]
            for view in (f"{plan['workflow']}-main-view", f"{plan['workflow']}-detail-view"):
                emitted = artifact["elementPatches"][view]["conditionalFormats"]
                for rule in source_rules:
                    self.assertIn(rule["emission"]["style"], [item["style"] for item in emitted])
            if plan["workflow"] == "expense":
                self.assertEqual(
                    [rule["value"] for rule in artifact["elementPatches"]["expense-detail-view"]["conditionalFormats"]],
                    ["CLEAR", "OVER", "RECON"],
                )

    def test_partial_or_unrelated_business_rule_does_not_spread_to_views(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-inventory-reorder.xlsx")
        source_rule = next(rule for rule in plan["conditionalFormats"]
                           if rule["sheet"] == "Inventory" and rule["status"] == "supported")
        partial = copy.deepcopy(source_rule)
        partial["id"] = "partial-status"
        partial["range"] = "O2:O18"
        unrelated = copy.deepcopy(source_rule)
        unrelated["id"] = "unrelated-supplier"
        unrelated["sheet"] = "Suppliers"
        unrelated["range"] = "A2:A5"
        unrelated["field"] = "Supplier"
        unrelated["fields"] = ["Supplier"]
        plan["conditionalFormats"].extend([partial, unrelated])
        targets = presentation.fixture_bindings(plan)["conditionalTargets"]
        self.assertEqual([target["elementId"] for target in targets["partial-status"]],
                         ["inventory-input"])
        self.assertEqual([target["elementId"] for target in targets["unrelated-supplier"]],
                         ["suppliers"])

    def test_multi_area_source_rule_keeps_base_target_without_widening(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "fixture-inventory-multi-area.xlsx"
            workbook = load_workbook(FIXTURES / "fixture-inventory-reorder.xlsx")
            workbook["Inventory"].conditional_formatting.add(
                "O2:O18 O20:O37",
                CellIsRule(operator="equal", formula=['"REORDER"']),
            )
            workbook.save(source)
            workbook.close()
            plan = presentation.extract_presentation(source)
            rule = next(rule for rule in plan["conditionalFormats"]
                        if rule["range"] == "O2:O18 O20:O37")
            targets = presentation.fixture_bindings(plan)["conditionalTargets"][rule["id"]]
            self.assertEqual([(target["elementId"], target["columnId"]) for target in targets],
                             [("inventory-input", "status")])
            # A full first area does not authorize projection when another
            # area targets a different source field.
            rule["range"] = "O2:O37 A2:A37"
            targets = presentation.fixture_bindings(plan)["conditionalTargets"][rule["id"]]
            self.assertEqual([(target["elementId"], target["columnId"]) for target in targets],
                             [("inventory-input", "status")])

    def test_cli_contents_uses_integer_document_version_after_staging_gate(self):
        source = FIXTURES / "fixture-research-model.xlsx"
        plan = presentation.extract_presentation(source)
        runtime_plan = research_runtime_plan()
        architecture = presentation.build_business_input_architecture(
            plan, connection_id="conn", folder_id="folder", research_plan=runtime_plan)
        feature = "researchModelInputs"
        columns = architecture["inputs"][feature]["columns"]
        column_ids = {column["name"]: f"id-{index}"
                      for index, column in enumerate(columns)}
        target_binding = {"documentId": "doc", "version": 7, "pageId": "inputs",
                          "elementId": "input-live", "columnIds": column_ids}
        manifest = build_manifest(plan["source"]["path"],
                                  {**copy.deepcopy(runtime_plan), "sharedInputArchitecture": architecture})
        manifest["target"] = {"documentId": "doc", "version": 7}
        _source, evidence = gate_plan_staging_source(
            {**copy.deepcopy(runtime_plan), "sourcePath": plan["source"]["path"],
             "sourceSha256": plan["source"]["sha256"]},
            source_arg=plan["source"]["path"])
        issues = evidence.get("issues") or []
        receipt = {"status": "APPROVED_STAGING", "sourceHash": manifest["sourceHash"],
                   "planHash": manifest["planHash"],
                   "sourceReviewStatus": "PENDING_STAGING" if issues else "NOT_REQUIRED",
                   "pendingSourceIssues": issues,
                   "sourceReviewRequirements": [
                       {"issue": issue, "issueHash": source_issue_hash(issue)} for issue in issues],
                   "targetBindings": {feature: target_binding}}
        receipt["artifacts"] = build_artifact_records(
            manifest, architecture, receipt["targetBindings"])
        bindings = presentation.fixture_bindings(plan)
        bindings.update({"planHash": manifest["planHash"], "documentId": "doc",
                         "documentVersion": 7, "editElementId": "input-live"})
        bindings["businessInput"].update({
            "elementId": "input-live",
            "currentColumns": [{"name": column["name"], "type": column["type"],
                                "id": column_ids[column["name"]]}
                               for column in columns],
        })
        bindings["stagingReceiptHash"] = presentation._digest(receipt)
        for rule in plan["validations"]:
            if rule["sheet"] == "Model Inputs":
                bindings["validationTargets"][rule["id"]] = [
                    {"elementId": "input-live", "columnId": column_ids[f"i_model_inputs_c{row}"]}
                    for row in range(2, 8)]

        artifact = presentation.emit_presentation(plan, bindings, manifest, receipt)
        current = {"contents": {
            "pages": [], "layout": "", "elements": [
                {"id": element_id,
                 "description": {"text": artifact["ownershipMarker"]}
                                if element_id in {item["id"] for item in artifact["elements"]} else {},
                 "name": artifact.get("calculationSource", {}).get("name")
                         if element_id == artifact.get("calculationSource", {}).get("elementId")
                         else element_id,
                 "columns": [{"id": column_id} for column_id in
                             artifact["elementPatches"].get(element_id, {}).get("columnUpdates", {})]}
                for element_id in set(artifact["referencedElementIds"]) | set(artifact["elementPatches"])
            ],
        }}
        with tempfile.TemporaryDirectory() as directory:
            tmp = Path(directory)
            paths = {name: tmp / f"{name}.json"
                     for name in ("bindings", "manifest", "receipt", "contents")}
            for name, value in (("bindings", bindings), ("manifest", manifest),
                                ("receipt", receipt), ("contents", current)):
                paths[name].write_text(json.dumps(value), encoding="utf-8")
            out = tmp / "put.json"
            command = [sys.executable, str(SCRIPTS / "build-presentation.py"), str(source),
                       "--bindings", str(paths["bindings"]), "--manifest", str(paths["manifest"]),
                       "--staging-receipt", str(paths["receipt"]), "--contents", str(paths["contents"]),
                       "--document-version", "7", "--out", str(out)]
            result = subprocess.run(command, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            put = json.loads(out.read_text())
            self.assertEqual(put["documentVersion"], 7)
            self.assertIsInstance(put["documentVersion"], int)
            self.assertEqual(put["sourceHash"], manifest["sourceHash"])
            self.assertGreater(len(put["contents"]["elements"]), len(current["contents"]["elements"]))

            out.unlink()
            paths["manifest"].write_text(json.dumps({**manifest, "sourceHash": "0" * 64}),
                                         encoding="utf-8")
            stale = subprocess.run(command, capture_output=True, text=True)
            self.assertNotEqual(stale.returncode, 0)
            self.assertIn("source", stale.stderr.lower())
            self.assertFalse(out.exists())
            paths["manifest"].write_text(json.dumps(manifest), encoding="utf-8")

            invalid = subprocess.run(command[:command.index("7")] + ["not-an-integer"]
                                     + command[command.index("7") + 1:],
                                     capture_output=True, text=True)
            self.assertEqual(invalid.returncode, 2)
            self.assertIn("invalid int value", invalid.stderr)
            self.assertFalse(out.exists())

    def test_business_input_programs_seed_only_literal_columns(self):
        expectations = {
            "fixture-forecast-model.xlsx": ("Line Item", None),
            "fixture-sales-pipeline.xlsx": ("Probability", "Weighted ACV"),
            "fixture-inventory-reorder.xlsx": ("Lead Days", "Reorder Qty"),
            "fixture-expense-approval.xlsx": ("Approval Status", "Exception"),
        }
        for filename, (editable, derived) in expectations.items():
            plan = presentation.extract_presentation(FIXTURES / filename)
            spec = presentation.extract_business_input(plan)
            program = presentation.emit_business_input_program(spec, presentation.fixture_bindings(plan)["businessInput"])
            column_names = [column["name"] for column in spec["columns"]]
            self.assertIn(editable, column_names)
            if derived is not None:
                self.assertNotIn(derived, column_names)
            self.assertIn(f"add_editable_column({editable!r}", program["program"])
            self.assertIn("set_edit_mode('published-explore')", program["program"])
            if derived is not None:
                self.assertNotIn(f"add_editable_column({derived!r}", program["program"])
                self.assertTrue(any(column["name"] == derived for column in spec["derivedColumns"]))
                if filename == "fixture-sales-pipeline.xlsx":
                    self.assertIn("add_formula_column('Weighted ACV'", program["program"])
                    self.assertIn("add_formula_column('Cycle Days'", program["program"])
                if filename == "fixture-inventory-reorder.xlsx":
                    self.assertIn("sigma.formula.round_up", program["program"])
                    self.assertEqual(program["liveFormulaColumns"],
                                     ["Available", "Reorder Point", "Days Cover", "Reorder Qty", "Status"])
            retry_binding = {**presentation.fixture_bindings(plan)["businessInput"], "elementId": "existing-input"}
            retry = presentation.emit_business_input_program(spec, retry_binding)
            self.assertEqual(retry["mode"], "reconcile-existing")
            self.assertNotIn("program", retry)

    def test_research_reuses_existing_shared_input_identity(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-research-model.xlsx")
        spec = presentation.extract_business_input(plan)
        current = [{"name": column["name"], "type": column["type"]}
                   for column in spec["additionalColumns"][:3]]
        program = presentation.emit_business_input_program(spec, {
            "pageId": "input-page", "elementId": "0CI0ooHYzw", "rowId": "row-1",
            "scopeColumnId": "scope-id", "planHash": "plan", "stagingReceiptHash": "receipt",
            "currentColumns": current,
        })
        self.assertEqual(program["mode"], "extend-existing")
        self.assertEqual(program["elementId"], "0CI0ooHYzw")
        self.assertEqual(program["expectedExistingEditableColumns"], 134)
        self.assertEqual(program["expectedEditableColumns"], 146)
        self.assertEqual(len(program["additionalColumns"]), 9)
        self.assertEqual(program["additionalColumns"][0]["sourceCellId"], "FY results!E56")
        self.assertEqual(program["additionalColumns"][-1]["sourceCellId"], "FY results!M56")
        self.assertEqual(program["additionalColumns"][0]["name"], "i_fy_results_e56")
        self.assertNotIn("update_rows", program["program"])

        complete = presentation.emit_business_input_program(spec, {
            "pageId": "input-page", "elementId": "0CI0ooHYzw", "rowId": "row-1",
            "scopeColumnId": "scope-id", "planHash": "plan", "stagingReceiptHash": "receipt",
            "currentColumns": [{"name": column["name"], "type": column["type"]}
                               for column in spec["additionalColumns"]],
        })
        self.assertEqual(complete["mode"], "reconcile-existing")
        self.assertNotIn("program", complete)

    def test_sales_and_expense_consume_reviewed_operational_wiring_builder(self):
        sales_plan = presentation.extract_presentation(FIXTURES / "fixture-sales-pipeline.xlsx")
        sales_binding = presentation.fixture_bindings(sales_plan)["businessInput"]
        sales_binding["elementId"] = "sales-input"
        sales = presentation.build_business_operational_wiring(sales_plan, sales_binding)
        sales_by_id = {item["id"]: item for item in sales["fragments"]["elements"]}
        timing = next(column for column in sales_by_id["salesWithSettings"]["columns"]
                      if column["name"] == "TimingStatus")
        self.assertIn('DateTrunc("day", [Settings/AsOfDate])', timing["formula"])

        expense_plan = presentation.extract_presentation(FIXTURES / "fixture-expense-approval.xlsx")
        expense_binding = presentation.fixture_bindings(expense_plan)["businessInput"]
        expense_binding["elementId"] = "expense-input"
        expense = presentation.build_business_operational_wiring(expense_plan, expense_binding)
        expense_by_id = {item["id"]: item for item in expense["fragments"]["elements"]}
        expense_calc = expense_by_id["expensesWithPolicy"]
        self.assertEqual(expense_calc["source"]["kind"], "join")
        exception = next(column for column in expense_calc["columns"] if column["name"] == "Exception")
        self.assertIn("[Policy/Per-item Limit]", json.dumps(expense_calc))
        self.assertIn("MISSING RECEIPT", exception["formula"])

    def test_emits_live_charts_formats_rules_and_three_usable_pages(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-sales-pipeline.xlsx")
        bindings = presentation.fixture_bindings(plan)
        artifact = presentation.emit_presentation(plan, bindings, logical_fixture=True)

        charts = [element for element in artifact["elements"] if element["kind"].endswith("-chart")]
        self.assertEqual([chart["kind"] for chart in charts], ["bar-chart", "line-chart"])
        self.assertEqual(charts[0]["yAxis"]["columnIds"], ["chart_1_series_1"])
        self.assertEqual(charts[0]["columns"][1]["formula"], "Sum([Pipeline/Weighted ACV])")
        self.assertEqual(charts[0]["xAxis"]["sort"]["by"], "chart_1_order")
        self.assertNotIn("cached", json.dumps(charts).lower())

        self.assertEqual([page["name"] for page in artifact["pages"]], ["Sales Overview", "Sales Edit", "Sales Detail"])
        self.assertEqual(artifact["pageRoles"], {"overview": "sales-overview", "edit": "sales-edit", "detail": "sales-detail"})
        self.assertIn("Probability", artifact["editableColumns"])
        self.assertIn("ACV", artifact["editableColumns"])
        self.assertIn("Weighted ACV", artifact["principalCalculatedColumns"])
        self.assertIn("Timing Status", artifact["principalCalculatedColumns"])
        self.assertLessEqual(len(artifact["mainGridColumns"]), 12)
        self.assertGreater(len(artifact["detailColumns"]), len(artifact["mainGridColumns"]))
        self.assertEqual(len(artifact["formulaAccounting"]), len(plan["formulaProvenance"]))
        self.assertIn("mapped-live-chart-check",
                      {item["targetStatus"] for item in artifact["formulaAccounting"]})

        probability_patch = artifact["elementPatches"]["sales-input"]["columnUpdates"]["probability"]
        self.assertEqual(probability_patch["range"], {"min": 0.0, "max": 1.0})

        page_spans = {}
        for page_id, body in re.findall(r'<Page\b[^>]*\bid="([^"]+)"[^>]*>(.*?)</Page>', artifact["layout"], re.S):
            spans = []
            for element_id, c1, c2, r1, r2 in re.findall(
                r'<Element elementId="([^"]+)" gridColumn="(\d+) / (\d+)" gridRow="(\d+) / (\d+)"/>', body
            ):
                span = tuple(map(int, (c1, c2, r1, r2)))
                for other_id, other in spans:
                    overlap = span[0] < other[1] and other[0] < span[1] and span[2] < other[3] and other[2] < span[3]
                    self.assertFalse(overlap, f"{page_id}: {element_id} overlaps {other_id}")
                spans.append((element_id, span))
            page_spans[page_id] = spans
        self.assertEqual(set(page_spans), {"sales-overview", "sales-edit", "sales-detail"})

    def test_nonfixture_emission_rejects_unverified_hash_labels(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-sales-pipeline.xlsx")
        bindings = presentation.fixture_bindings(plan)
        bindings["logicalFixture"] = True  # A caller label cannot enable the fixture bypass.
        with self.assertRaisesRegex(ValueError, "manifest and staging receipt"):
            presentation.emit_presentation(plan, bindings)
        with self.assertRaisesRegex(ValueError, "stagingReceiptHash"):
            presentation.emit_presentation(
                plan, bindings,
                {"sourceHash": plan["source"]["sha256"], "planHash": "a" * 64},
                {"status": "APPROVED_STAGING"},
            )

    def test_actual_emission_recomputes_approved_146_column_gate(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-research-model.xlsx")
        runtime_plan = research_runtime_plan()
        architecture = presentation.build_business_input_architecture(
            plan, connection_id="conn", folder_id="folder", research_plan=runtime_plan)
        feature = "researchModelInputs"
        column_ids = {column["name"]: f"id-{index}"
                      for index, column in enumerate(architecture["inputs"][feature]["columns"])}
        target_binding = {"documentId": "doc", "version": 7, "pageId": "inputs",
                          "elementId": "input-live", "columnIds": column_ids}
        manifest = build_manifest(plan["source"]["path"],
                                  {**copy.deepcopy(runtime_plan), "sharedInputArchitecture": architecture})
        manifest["target"] = {"documentId": "doc", "version": 7}
        _source, evidence = gate_plan_staging_source(
            {**copy.deepcopy(runtime_plan), "sourcePath": plan["source"]["path"],
             "sourceSha256": plan["source"]["sha256"]},
            source_arg=plan["source"]["path"])
        issues = evidence.get("issues") or []
        receipt = {"status": "APPROVED_STAGING", "sourceHash": manifest["sourceHash"],
                   "planHash": manifest["planHash"],
                   "sourceReviewStatus": "PENDING_STAGING" if issues else "NOT_REQUIRED",
                   "pendingSourceIssues": issues,
                   "sourceReviewRequirements": [
                       {"issue": issue, "issueHash": source_issue_hash(issue)} for issue in issues],
                   "targetBindings": {feature: target_binding}}
        receipt["artifacts"] = build_artifact_records(
            manifest, architecture, receipt["targetBindings"])
        bindings = presentation.fixture_bindings(plan)
        bindings.update({"planHash": manifest["planHash"], "documentId": "doc",
                         "documentVersion": 7, "editElementId": "input-live"})
        bindings["businessInput"].update({
            "elementId": "input-live",
            "currentColumns": [{"name": column["name"], "type": column["type"],
                                "id": column_ids[column["name"]]}
                               for column in architecture["inputs"][feature]["columns"]],
        })
        bindings["stagingReceiptHash"] = presentation._digest(receipt)
        for rule in plan["validations"]:
            if rule["sheet"] == "Model Inputs":
                bindings["validationTargets"][rule["id"]] = [
                    {"elementId": "input-live", "columnId": column_ids[f"i_model_inputs_c{row}"]}
                    for row in range(2, 8)]

        artifact = presentation.emit_presentation(plan, bindings, manifest, receipt)

        self.assertEqual(artifact["stagingGate"]["status"], "APPROVED_STAGING")
        self.assertEqual(len(architecture["inputs"][feature]["columns"]), 146)
        self.assertEqual(artifact["businessInputProgram"]["mode"], "reconcile-existing")
        updates = artifact["elementPatches"]["input-live"]["columnUpdates"]
        self.assertTrue(all("hidden" not in update and "format" not in update for update in updates.values()))
        self.assertEqual(updates[column_ids["i_model_inputs_c2"]]["range"], {"min":0.0,"max":0.25})
        display = artifact["inputDisplay"]
        visible = {column["name"] for column in display["columns"] if column["visible"]}
        self.assertEqual(visible, {"Scope", *(f"i_model_inputs_c{row}" for row in range(2, 8))})
        self.assertEqual(display["status"], "PENDING_SDK_READBACK")
        self.assertTrue(all(column["format"] == {"pattern":".1%"}
                            for column in display["columns"] if column["name"] in visible-{"Scope"}))

    def test_business_input_architecture_covers_primary_inputs_for_all_five_workflows(self):
        research_plan = research_runtime_plan()
        expected = {
            "fixture-forecast-model.xlsx": {"forecastBusinessInput"},
            "fixture-research-model.xlsx": {"researchModelInputs"},
            "fixture-sales-pipeline.xlsx": {"salesBusinessInput", "salesSettings"},
            "fixture-inventory-reorder.xlsx": {"inventoryBusinessInput"},
            "fixture-expense-approval.xlsx": {"expenseBusinessInput", "expensePolicy"},
        }
        for filename, features in expected.items():
            plan = presentation.extract_presentation(FIXTURES / filename)
            architecture = presentation.build_business_input_architecture(
                plan, connection_id="conn", folder_id="folder",
                research_plan=research_plan if plan["workflow"] == "research" else None)
            self.assertEqual(set(architecture["inputs"]), features)
            primary = architecture["presentationInputFeature"]
            self.assertIn(primary, features)
            self.assertEqual(architecture["inputs"][primary]["editMode"], "published-explore")
        self.assertEqual(len(presentation.build_business_input_architecture(
            presentation.extract_presentation(FIXTURES / "fixture-research-model.xlsx"),
            connection_id="conn", folder_id="folder", research_plan=research_plan,
        )["inputs"]["researchModelInputs"]["columns"]), 146)

    def test_public_approved_entrypoint_constructs_or_extends_all_five_workflows(self):
        research_plan = research_runtime_plan()
        modes = {}
        for filename in (
                "fixture-forecast-model.xlsx", "fixture-research-model.xlsx",
                "fixture-sales-pipeline.xlsx", "fixture-inventory-reorder.xlsx",
                "fixture-expense-approval.xlsx"):
            extracted = presentation.extract_presentation(FIXTURES / filename)
            architecture = presentation.build_business_input_architecture(
                extracted, connection_id="conn", folder_id="folder",
                research_plan=research_plan if extracted["workflow"] == "research" else None)
            logical = presentation.build_business_input_plan(extracted, architecture)
            manifest = build_manifest(extracted["source"]["path"], logical)
            unresolved = [item for item in manifest["formulaCells"]
                          if item["id"] in manifest["unsupported"]]
            if unresolved:
                reviewed = []
                for item in unresolved:
                    mapping = architecture["formulaTargets"][item["id"]]
                    graph = architecture["nativeWiring"]
                    if "logicalElementId" in mapping:
                        element_id, column_id = mapping["logicalElementId"], mapping["logicalColumnId"]
                    else:
                        check = next(c for c in graph["checks"] if c["logicalCheckId"] == mapping["logicalCheckId"])
                        element_id, column_id = check["elementId"], check["columnId"]
                    formula = next(c["formula"] for e in graph["elements"] if e["id"] == element_id
                                   for c in e["columns"] if c["id"] == column_id)
                    reviewed.append({"id": item["id"], "mappingStatus": "live", "targetMapping": mapping,
                                     "reviewedReplacement": {"originalFormula": item["originalFormula"],
                                      "replacementFormula": formula, "reviewer": "Fixture Reviewer",
                                      "reviewStatus": "approved"}})
                logical = presentation.build_business_input_plan(
                    extracted, architecture, reviewed)
                manifest = build_manifest(extracted["source"]["path"], logical)
            bindings = {}
            for feature, spec in architecture["inputs"].items():
                binding = {"documentId": "doc", "version": 1,
                           "pageId": f"{extracted['workflow']}-edit"}
                if extracted["workflow"] == "research":
                    columns = [column for column in spec["columns"]
                               if not re.fullmatch(r"i_fy_results_[b-m]56", column["name"])]
                    binding.update({"elementId": "research-input", "columnIds": {
                        column["name"]: f"id-{index}" for index, column in enumerate(columns)}})
                bindings[feature] = binding
            _source, evidence = gate_plan_staging_source(
                logical, source_arg=extracted["source"]["path"])
            issues = evidence.get("issues") or []
            approval = {
                "sourceHash": manifest["sourceHash"], "planHash": manifest["planHash"],
                "owner": "Migration Owner", "reviewer": "Migration Reviewer",
                "reviewItems": [item["id"] for item in manifest["formulaCells"]],
                "mappedFeatures": architecture["mappedFeatures"],
                "targetBindings": bindings,
            }
            if issues:
                approval["sourceReviewApproval"] = {
                    "status": "approved_for_staging", "issues": issues}
            receipt = gate_input_staging(manifest, approval, logical, evidence)
            feature = architecture["presentationInputFeature"]
            current = None
            if extracted["workflow"] == "research":
                by_name = {column["name"]: column for column in architecture["inputs"][feature]["columns"]}
                current = [{"name": name, "type": by_name[name]["type"], "id": column_id}
                           for name, column_id in bindings[feature]["columnIds"].items()]
            result = presentation.emit_approved_business_input_program(
                manifest, receipt, feature, current_columns=current)
            modes[extracted["workflow"]] = result["mode"]
        self.assertEqual(modes, {"forecast": "construct-approved", "research": "extend-existing",
                                 "sales": "construct-approved", "inventory": "construct-approved",
                                 "expense": "construct-approved"})

    def test_research_edit_view_keeps_shared_drivers_prominent_and_133_fields_in_detail(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-research-model.xlsx")
        bindings = presentation.fixture_bindings(plan)
        bindings["sharedInputColumns"] = [f"Source Cell {i}" for i in range(1, 134)]
        bindings["prominentEditableColumns"] = [f"Source Cell {i}" for i in range(1, 7)]
        artifact = presentation.emit_presentation(plan, bindings, logical_fixture=True)

        self.assertEqual(len(artifact["editColumns"]), 6)
        self.assertEqual(len(artifact["detailColumns"]), 133)
        self.assertEqual(artifact["sharedInputScope"], "shared")
        main = {column["name"]: column for element in artifact["elements"]
                if element["id"] == "research-main-view" for column in element["columns"]}
        detail_view = next(element for element in artifact["elements"]
                           if element["id"] == "research-detail-view")
        self.assertEqual(len(detail_view["columns"]), 36)
        self.assertNotIn("Line Item", {column["name"] for column in detail_view["columns"]})
        self.assertEqual(main["Revenue"]["format"]["currencySymbol"], "$")
        self.assertEqual(main["EPS"]["format"]["formatString"], "($.2f")
        self.assertEqual(main["Net debt / EBITDA (x)"]["format"]["suffix"], "x")
        self.assertFalse(any(item["kind"] == "mixed-unit-row-format"
                             for item in artifact["unresolved"]))
        updates = artifact["elementPatches"]["research-input-existing"]["columnUpdates"]
        self.assertEqual(
            {name for name, update in updates.items() if update.get("range") == {"min": 0.0, "max": 0.25}},
            {f"i_model_inputs_c{row}" for row in range(2, 8)},
        )

    def test_research_uses_runtime_calculation_source_name_for_views_and_charts(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-research-model.xlsx")
        bindings = presentation.fixture_bindings(plan)
        bindings["calculationSource"] = {"elementId": "calc", "name": "Calc"}
        bindings["viewSourceElementId"] = "calc"
        for chart in bindings["charts"].values():
            chart["sourceElementId"] = "calc"
        bindings["mainViewColumns"], bindings["detailViewColumns"] = (
            presentation.research_view_columns(plan, bindings["calculationSource"]))

        artifact = presentation.emit_presentation(plan, bindings, logical_fixture=True)

        self.assertTrue(artifact["researchMetricFormatsExact"])
        formulas = [column["formula"] for element in artifact["elements"]
                    if element.get("source", {}).get("elementId") == "calc"
                    for column in element.get("columns", []) if column.get("formula")]
        self.assertTrue(formulas)
        self.assertTrue(all("[Calculation/" not in formula for formula in formulas))
        self.assertTrue(any("[Calc/c8]" in formula for formula in formulas))

    def test_research_horizontal_rules_bind_metric_rows_and_audit_stays_unresolved(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-research-model.xlsx")
        rules = {r["range"]: r for r in plan["conditionalFormats"] if r["sheet"] == "FY results"}
        self.assertEqual(rules["B31:M31"]["fields"], ["Net income"])
        self.assertEqual(rules["B57:M57"]["fields"], ["Net debt / EBITDA (x)"])
        self.assertEqual(rules["B31:M31"]["periodScope"]["values"], list(range(2019, 2031)))
        bindings = presentation.fixture_bindings(plan)
        self.assertEqual(bindings["conditionalTargets"][rules["B31:M31"]["id"]], [
            {"elementId": "research-detail-view", "columnId": "research_c31"}])
        self.assertEqual({t["elementId"] for t in bindings["conditionalTargets"][rules["B57:M57"]["id"]]},
                         {"research-main-view", "research-detail-view"})
        artifact = presentation.emit_presentation(plan, bindings, logical_fixture=True)
        audit_ids = {r["id"] for r in plan["conditionalFormats"] if r["sheet"] == "Audit"}
        self.assertTrue(audit_ids)
        self.assertTrue(audit_ids <= {u["sourceId"] for u in artifact["unresolved"] if u["kind"] == "conditionalFormat"})
        self.assertNotIn("researchCalc", artifact["elementPatches"])

    def test_public_emission_omits_dynamic_midpoints_and_keeps_finite_literal_midpoint(self):
        from openpyxl import load_workbook
        import xml.etree.ElementTree as ET
        with tempfile.TemporaryDirectory() as directory:
            tmp = Path(directory)
            for mode, value in (("percentile", 50), ("formula", "A1"), ("num", 2.5)):
                workbook = load_workbook(FIXTURES / "fixture-research-model.xlsx")
                rule = next(rule for key, rules in workbook["FY results"].conditional_formatting._cf_rules.items()
                            for rule in rules if rule.type == "colorScale")
                rule.colorScale.cfvo[1].type = mode
                rule.colorScale.cfvo[1].val = value
                source = tmp / ("fixture-research-" + mode + ".xlsx")
                workbook.save(source)
                command = [sys.executable, str(SCRIPTS / "build-presentation.py"), str(source),
                           "--fixture-bindings", "--out", str(tmp / "artifact.json")]
                result = subprocess.run(command, capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)
                artifact = json.loads((tmp / "artifact.json").read_text())
                provenance = next(r for r in artifact["ruleProvenance"] if r["sourceType"]=="colorScale")
                patches = [r for p in artifact["elementPatches"].values()
                           for r in p["conditionalFormats"] if r["type"]=="backgroundScale"]
                if mode == "num":
                    self.assertEqual(len(patches), 2)
                    self.assertTrue(all(r["domain"]["mid"]==2.5 for r in patches))
                else:
                    self.assertEqual(patches, [], 'dynamic midpoint must not be approximated or emitted')
                    unresolved = next(r for r in artifact["unresolved"] if r["sourceId"]==provenance["id"])
                    self.assertEqual(unresolved["kind"], "conditionalFormat")
                    self.assertIn("midpoint", unresolved["reason"])
                    self.assertIsNone(provenance["emission"])
                for patch in patches:
                    self.assertIsInstance(patch["domain"].get("mid"), (int,float))
                if mode == "num":
                    # Reject a legacy/API-invalid domain even if a caller marks
                    # the extracted rule supported. This reproduces the actual
                    # store diagnostic at the emitter boundary.
                    legacy = presentation.extract_presentation(source)
                    legacy_rule = next(r for r in legacy["conditionalFormats"] if r["sourceType"]=="colorScale")
                    legacy_rule["emission"]["domain"]["mid"] = "Percentile([research_c57], 0.5)"
                    guarded = presentation.emit_presentation(legacy, presentation.fixture_bindings(legacy), logical_fixture=True)
                    self.assertTrue(any(u["sourceId"]==legacy_rule["id"] and "midpoint" in u["reason"]
                                        for u in guarded["unresolved"]))
                    self.assertFalse(any(r["type"]=="backgroundScale" for p in guarded["elementPatches"].values()
                                         for r in p["conditionalFormats"]))
                xml = ET.fromstring(provenance["sourceRuleXml"])
                threshold = xml.find("colorScale/cfvo[@type='" + mode + "']")
                self.assertIsNotNone(threshold)
                if mode == "formula":
                    self.assertEqual(threshold.get("val"), "A1")
                else:
                    self.assertEqual(float(threshold.get("val")), 50 if mode=="percentile" else 2.5)

    def test_forecast_manual_validation_covers_each_month_column(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-forecast-model.xlsx")
        artifact = presentation.emit_presentation(plan, presentation.fixture_bindings(plan), logical_fixture=True)
        updates = artifact["elementPatches"]["forecast-input"]["columnUpdates"]
        self.assertEqual(
            {name for name, update in updates.items()
             if update.get("range") == {"min": 0.0, "max": 1000000.0}},
            {"jul_26", "aug_26", "sep_26", "oct_26", "nov_26", "dec_26"},
        )
        zero_rules = [rule for rule in artifact["elementPatches"]["forecast-input"]["conditionalFormats"]
                      if rule["style"].get("backgroundColor") == "#FFE699"]
        self.assertEqual({rule["columnIds"][0] for rule in zero_rules},
                         {"jul_26", "aug_26", "sep_26", "oct_26", "nov_26", "dec_26"})

    def test_schema_safe_text_descriptions_and_styles(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-research-model.xlsx")
        artifact = presentation.emit_presentation(plan, presentation.fixture_bindings(plan), logical_fixture=True)
        text = [element for element in artifact["elements"] if element["kind"] == "text"]
        self.assertTrue(all(element["verticalAlign"] == "center" for element in text))
        self.assertNotIn("<!--", json.dumps(text))
        described = [element for element in artifact["elements"]
                     if element["kind"] == "table" or element["kind"].endswith("-chart")]
        self.assertTrue(all(element["description"]["visibility"] == "hidden" for element in described))
        self.assertTrue(all("padding" not in element.get("style", {}) for element in described))

    def test_research_metric_exactness_is_derived_from_columns(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-research-model.xlsx")
        bindings = presentation.fixture_bindings(plan)
        bindings["researchMetricFormatsExact"] = True
        bindings["mainViewColumns"][1]["format"] = {"kind": "number", "formatString": ".1%"}
        artifact = presentation.emit_presentation(plan, bindings, logical_fixture=True)
        self.assertFalse(artifact["researchMetricFormatsExact"])
        self.assertTrue(any(item["kind"] == "mixed-unit-row-format"
                            for item in artifact["unresolved"]))

    def test_formula_accounting_has_owned_checks_or_explicit_unresolved_scope(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-sales-pipeline.xlsx")
        artifact = presentation.emit_presentation(plan, presentation.fixture_bindings(plan), logical_fixture=True)
        scopes = {item["sourceId"]: item for item in artifact["unresolved"]
                  if item["kind"] == "formula-scope"}
        self.assertIn("formula-scope-sales", scopes)
        for item in artifact["formulaAccounting"]:
            if item["targetStatus"] == "mapped-live-chart-check":
                self.assertTrue(item["targetOwnership"]["checkIds"])
            else:
                self.assertEqual(item["unresolvedScopeId"], "formula-scope-sales")

    def test_unbound_source_chart_remains_explicitly_unresolved(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-expense-approval.xlsx")
        bindings = presentation.fixture_bindings(plan)
        bindings["charts"].pop(plan["charts"][0]["id"])
        artifact = presentation.emit_presentation(plan, bindings, logical_fixture=True)

        self.assertEqual(len([e for e in artifact["elements"] if e["kind"].endswith("-chart")]), 1)
        unresolved = next(item for item in artifact["unresolved"] if item["sourceId"] == plan["charts"][0]["id"])
        self.assertEqual(unresolved["reason"], "missing live target binding")

    def test_full_contents_integration_applies_validation_and_rule_patches(self):
        plan = presentation.extract_presentation(FIXTURES / "fixture-inventory-reorder.xlsx")
        bindings = presentation.fixture_bindings(plan)
        artifact = presentation.emit_presentation(plan, bindings, logical_fixture=True)
        generated_ids = {element["id"] for element in artifact["elements"]}
        element_ids = set(artifact["referencedElementIds"]) | (set(artifact["elementPatches"]) - generated_ids)
        patch_columns = {
            element_id: set(patch["columnUpdates"])
            | {column_id for rule in patch["conditionalFormats"] for column_id in rule["columnIds"]}
            for element_id, patch in artifact["elementPatches"].items()
        }
        contents = {
            "schemaVersion": 1, "kind": "workbook",
            "pages": [{"id": "existing", "name": "Existing"}],
            "elements": [
                {"id": element_id, "kind": "input-table" if element_id in artifact["elementPatches"] else "table",
                 "columns": [{"id": column_id, "type": "number"} for column_id in sorted(patch_columns.get(element_id, set()))],
                 "source": {"kind": "table", "elementId": "source"}, "inputMode": "edit"}
                for element_id in sorted(element_ids)
            ],
            "layout": ('<?xml version="1.0" encoding="utf-8"?>\n<Page id="existing">'
                       f'<Element elementId="{artifact["editElementId"]}" gridColumn="1 / 2" gridRow="1 / 2"/>'
                       '</Page>'),
        }
        merged = presentation.integrate_presentation(contents, artifact)
        lead_days = next(
            column for element in merged["elements"] if element["id"] == "inventory-input"
            for column in element["columns"] if column["id"] == "lead_days"
        )
        self.assertEqual(lead_days["values"], list(range(1, 91)))
        self.assertEqual(len(merged["pages"]), 4)
        self.assertEqual(len(re.findall(r"<Page\b", merged["layout"])), 4)
        self.assertEqual(len(re.findall(
            rf'<Element\b[^>]*elementId="{re.escape(artifact["editElementId"])}"', merged["layout"])), 1)

        repeated = presentation.integrate_presentation(merged, artifact)
        self.assertEqual(len(repeated["pages"]), 4)
        self.assertEqual(len(re.findall(r"<Page\b", repeated["layout"])), 4)
        self.assertEqual(len({element["id"] for element in repeated["elements"]}), len(repeated["elements"]))

        foreign = json.loads(json.dumps(contents))
        foreign["elements"].append({"id": artifact["elements"][0]["id"], "kind": "text", "body": "foreign"})
        with self.assertRaisesRegex(ValueError, "foreign contents"):
            presentation.integrate_presentation(foreign, artifact)


if __name__ == "__main__":
    unittest.main()
