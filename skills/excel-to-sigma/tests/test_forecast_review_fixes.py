#!/usr/bin/env python3
"""Review fixes for the forecast drafter and builder.

The central check evaluates the formulas the builder EMITS (KPI cards, the trend chart and
every amount derived row) over the plan's own rows with ``emitted_formula_eval`` and compares
them with the oracle, so a plan that verifies against the workbook but would show different
numbers in Sigma fails here. Also covers the router never crashing, arity-checked row
formulas, the tight comparison tolerance, formula-valued drivers and KPI role selection.
"""
from __future__ import annotations

import contextlib
import copy
import importlib.util
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / "lib"))
sys.path.insert(0, str(ROOT / "tests"))

import emitted_formula_eval as emitted  # noqa: E402
import forecast_calc as calc  # noqa: E402
import forecast_plan  # noqa: E402
import forecast_pnl_fixture as fixture  # noqa: E402
from forecast_plan import PlanDraftError, draft_plan  # noqa: E402

FIXTURES = ROOT / "tests" / "fixtures"
RETAIL = FIXTURES / "broader-eight" / "03 Retail Quarterly Budget.xlsx"
SAAS = FIXTURES / "broader-eight" / "02 SaaS Driver Forecast FY26.xlsx"


def load_script(filename, name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / filename)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


BUILD = load_script("build-forecast-model.py", "review_fix_builder")
DRAFT = load_script("draft-forecast-plan.py", "review_fix_draft_cli")
MIGRATE = load_script("migrate.py", "review_fix_migrate")
ELS = {"glDetail": "glDetail", "spine": "spine", "assumptions": "assumptions"}


def quiet_parity(plan, computed, tolerance=None):
    """``BUILD.check_parity`` without its per-metric report lines."""
    with contextlib.redirect_stdout(io.StringIO()):
        return BUILD.check_parity(plan, computed, tolerance)


def card_name(element):
    """A KPI card's title text, case-folded: the card carries a styled ``name`` object with an uppercase label."""
    name = element["name"]
    return (name["text"] if isinstance(name, dict) else name).casefold()


def label_rows(ws):
    return {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}


def check_emitted(test, plan):
    """Emitted KPI, trend and per-row formulas == the oracle, for every derived amount row."""
    ev = calc.Evaluation(plan)
    spec = BUILD.build_wb_spec(plan, "dm", ELS)
    rows = emitted.combined_rows(plan, ev, spec)
    fy = plan.get("fyPeriods")
    idx = [i for i, label in enumerate(ev.labels) if fy is None or label in fy]
    fy_nums = sorted(ev.periods[i]["num"] for i in idx) if fy is not None else []
    periods, n_act, sections = BUILD.plan_model(plan)
    sec_cols = BUILD.section_columns(sections)
    amount = {d["label"]: d for d in plan["derivedRows"] if d["kind"] == "amount"}

    # 1. every amount derived row, whole horizon / FY window / each period
    for label, d in amount.items():
        row = {"coefficients": d["coefficients"], "lineCoefficients": d.get("lineCoefficients", {})}
        whole = emitted.evaluate(BUILD.amount_formula(row, "PnL Combined", sec_cols, [], False), rows)
        test.assertAlmostEqual(whole or 0.0, ev.aggregate(label, range(len(ev.labels))), places=6,
                               msg=f"{label}: emitted formula over every period")
        windowed = emitted.evaluate(BUILD.amount_formula(row, "PnL Combined", sec_cols, fy_nums, True), rows)
        test.assertAlmostEqual(windowed or 0.0, ev.aggregate(label, idx), places=6,
                               msg=f"{label}: emitted formula over the FY periods")
        for i, period in enumerate(ev.periods):
            in_period = [r for r in rows if r["Month Num"] == period["num"]]
            test.assertAlmostEqual(
                emitted.evaluate(BUILD.amount_formula(row, "PnL Combined", sec_cols, [], False), in_period) or 0.0,
                ev.derived[label][i], places=6, msg=f"{label}: emitted formula in {period['label']}")

    # 2. the KPI cards and trend series the spec actually carries
    fy_values = {}
    for sec in plan["sections"]:
        if sec.get("totalLabel"):
            fy_values[sec["totalLabel"]] = ev.section_total(sec["name"], idx)
    for label in amount:
        fy_values[label] = ev.aggregate(label, idx)
    first_positive = next(s["name"] for s in sections if s["sign"] > 0)
    kpis = [el for page in spec["pages"] for el in page["elements"] if el["kind"] == "kpi-chart"]
    test.assertGreaterEqual(len(kpis), 3)
    for kpi in kpis:
        name, value = card_name(kpi), emitted.evaluate(kpi["columns"][0]["formula"], rows)
        if name.endswith("%"):
            first = next(iter(amount))
            expected = fy_values[first] / ev.section_total(first_positive, idx)
        else:
            key = max((k for k in fy_values if name.endswith(k.casefold())), key=len)
            expected = fy_values[key]
        test.assertAlmostEqual(value, expected, places=6, msg=f"KPI {name!r}")
    trend = next((el for page in spec["pages"] for el in page["elements"] if el["id"] == "monthlyTrend"), None)
    for column in (trend["columns"] if trend else []):
        if column["name"] in ("Month", "Period", "Period Order") or column["id"] in ("t_month", "t_ord"):
            continue
        for i, period in enumerate(ev.periods):
            in_period = [r for r in rows if r["Month Num"] == period["num"]]
            value = emitted.evaluate(column["formula"], in_period)
            series = ev.derived[column["name"]] if column["name"] in amount else ev.totals[column["name"]]
            test.assertAlmostEqual(value or 0.0, series[i], places=6,
                                   msg=f"trend {column['name']!r} in {period['label']}")

    # 3. the mapped source charts: each series over every period equals the oracle's row total
    def metric_series(name):
        for label, values in ev.derived.items():
            if calc.TARGET_LABELS.get(label.lower(), label) == name:
                return values
        for sec, label in ev.total_labels.items():
            if label and calc.TARGET_LABELS.get(label.lower(), label) == name:
                return ev.totals[sec]
        return ev.totals.get(name)
    for chart in (el for page in spec["pages"] for el in page["elements"] if el["id"].startswith("sourceChart")):
        if chart.get("filters") and any(f["startDate"] > ev.periods[0]["end"][:8] + "01" or f["endDate"] < ev.periods[-1]["end"]
                                        for f in chart["filters"] if f["kind"] == "date-range"):
            continue            # a chart over part of the horizon
        for column in chart["columns"]:
            if "_y" not in column["id"]:
                continue
            series = metric_series(column["name"])
            test.assertIsNotNone(series, f"source chart series {column['name']!r}")
            test.assertAlmostEqual(emitted.evaluate(column["formula"], rows) or 0.0, sum(series), places=6,
                                   msg=f"source chart {chart['id']} series {column['name']!r}")
    return ev, spec


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        self.path = self.dir / "pnl.xlsx"

    def draft(self, path=None, **kwargs):
        return draft_plan(path or self.path, intent="what-if", connection="c", folder="f", **kwargs)

    def plan(self, path):
        stdout = io.StringIO()
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(io.StringIO()):
            code = MIGRATE.main(["plan", str(path), "--out-dir", str(self.dir / "work")])
        return code, json.loads(stdout.getvalue())


# ------------------------------------------------------------------ 1. line -> section collapse
class TestLineLevelTerms(Base):
    def contribution_workbook(self):
        """Contribution = Subscription (one revenue line) - Total Cost of Revenue."""
        clean = self.dir / "clean.xlsx"
        base = fixture.build_pnl(clean)
        cogs = [h + s for h, s in zip(base["values"]["Hosting"], base["values"]["Support"])]
        contrib = [s - c for s, c in zip(base["values"]["Subscription"], cogs)]

        def mutate(pnl, gl, drivers):
            rows = label_rows(pnl)
            new = max(rows.values()) + 2
            pnl.cell(new, 1, "Contribution")
            for i in range(6):
                col = chr(ord("B") + i)
                pnl[f"{col}{new}"] = f"={col}{rows['Subscription']}-{col}{rows['Total Cost of Revenue']}"
            pnl[f"H{new}"] = f"=SUM(B{new}:G{new})"
            self.new_row = new
        # the row number is only known inside mutate; the override keys follow the layout
        new = 26
        overrides = {("P&L", f"{chr(ord('B') + i)}{new}"): contrib[i] for i in range(6)}
        overrides[("P&L", f"H{new}")] = sum(contrib)
        fixture.build_pnl(self.path, mutate=mutate, cached_overrides=overrides)
        self.assertEqual(self.new_row, new)
        return contrib

    def test_line_reference_stays_a_line_term_in_the_plan(self):
        contrib = self.contribution_workbook()
        plan, summary = self.draft()
        row = next(d for d in plan["derivedRows"] if d["label"] == "Contribution")
        self.assertEqual(row["kind"], "amount")
        self.assertEqual(row["coefficients"], {"Cost of Revenue": -1})
        self.assertEqual(row["lineCoefficients"], {"Subscription": 1})
        # the oracle and the workbook agree on the real value, not Revenue - COGS
        ev = calc.Evaluation(plan)
        self.assertAlmostEqual(ev.aggregate("Contribution", range(6)), sum(contrib), places=6)
        self.assertNotAlmostEqual(sum(contrib), ev.section_total("Revenue") - ev.section_total("Cost of Revenue"))
        # classic plans keep their four canonical targets and add the extra amount row
        self.assertEqual(list(plan["parityTargets"]),
                         ["Revenue", "Gross Profit", "EBITDA", "Net Income", "Contribution"])
        self.assertAlmostEqual(plan["parityTargets"]["Contribution"], sum(contrib), places=6)

    def test_emitted_formulas_match_the_oracle_for_every_derived_row(self):
        self.contribution_workbook()
        plan, _ = self.draft()
        check_emitted(self, plan)
        # the line-level formula really is line-level
        sec_cols = BUILD.section_columns(BUILD.plan_model(plan)[2])
        row = {"coefficients": {"Cost of Revenue": -1}, "lineCoefficients": {"Subscription": 1}}
        self.assertEqual(BUILD.amount_formula(row, "PnL Combined", sec_cols, []),
                         '-Sum([PnL Combined/COGS Amt]) + SumIf([PnL Combined/Amount], '
                         '[PnL Combined/Line Item] = "Subscription")')

    def test_the_old_section_collapse_is_caught_by_the_emitted_evaluator(self):
        """The check has teeth: a plan that widens the line to its section disagrees."""
        self.contribution_workbook()
        plan, _ = self.draft()
        broken = copy.deepcopy(plan)
        row = next(d for d in broken["derivedRows"] if d["label"] == "Contribution")
        row["coefficients"] = {"Revenue": 1, "Cost of Revenue": -1}
        row.pop("lineCoefficients")
        with self.assertRaises(AssertionError):
            check_emitted(self, broken)

    def test_kpi_cards_are_chosen_by_role_not_sheet_position(self):
        self.contribution_workbook()
        plan, _ = self.draft()
        spec = BUILD.build_wb_spec(plan, "dm", ELS)
        cards = {el["id"]: card_name(el) for page in spec["pages"] for el in page["elements"]
                 if el["kind"] == "kpi-chart"}
        self.assertEqual(cards, {"kpiRev": "total revenue", "kpiEbitda": "ebitda",
                                 "kpiNI": "net income", "kpiGM": "gross margin %"})

    def test_line_names_the_formulas_cannot_carry_are_rejected_naming_the_cell(self):
        lines = [list(line) for line in fixture.DEFAULT_LINES]
        lines[0][1] = 'Sub "A"'
        clean = self.dir / "clean.xlsx"
        base = fixture.build_pnl(clean, lines=lines)

        def mutate(pnl, gl, drivers):
            rows = label_rows(pnl)
            pnl.cell(26, 1, "Odd")
            for i in range(6):
                col = chr(ord("B") + i)
                pnl[f"{col}26"] = f"={col}{rows['Sub \"A\"']}-{col}{rows['Total Cost of Revenue']}"
        with self.assertRaisesRegex(PlanDraftError, r"P&L!B26 \('Odd'\).*Sub \"A\".*quote or backslash"):
            fixture.build_pnl(self.path, lines=lines, mutate=mutate)
            self.draft()


    def test_builder_rejects_a_hand_edited_plan_with_an_inexpressible_line_name(self):
        fixture.build_pnl(self.path)
        plan, _ = self.draft()
        next(d for d in plan["derivedRows"] if d["label"] == "EBITDA")["lineCoefficients"] = {'Sub "A"': 1}
        plan_path = self.dir / "plan.json"
        plan_path.write_text(json.dumps(plan), encoding="utf-8")
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = BUILD.main([str(plan_path), "--out-dir", str(self.dir / "build")])
        self.assertEqual(code, 1)
        self.assertIn("PLAN ERROR: line 'Sub \"A\"' contains a quote or backslash", err.getvalue())
        self.assertFalse((self.dir / "build" / "wb_spec.json").exists())


# ------------------------------------------------------------------ 4. false rejections
class TestFalseRejections(Base):
    def test_net_income_over_both_lines_of_a_section(self):
        lines = [list(line) for line in fixture.DEFAULT_LINES]
        lines.append(["Below the Line", "Tax", "flat", None, [5.0, 5.0, 6.0]])

        def mutate(pnl, gl, drivers):
            rows = label_rows(pnl)
            for i in range(6):
                col = chr(ord("B") + i)
                pnl[f"{col}{rows['Net Income']}"] = (
                    f"={col}{rows['EBITDA']}-{col}{rows['Interest']}-{col}{rows['Tax']}")
        built = fixture.build_pnl(self.path, lines=lines, mutate=mutate)
        plan, _ = self.draft()
        net = next(d for d in plan["derivedRows"] if d["label"] == "Net Income")
        self.assertEqual(net["coefficients"], {"Revenue": 1, "Cost of Revenue": -1,
                                               "Operating Expenses": -1, "Below the Line": -1})
        self.assertNotIn("lineCoefficients", net)
        self.assertEqual([s["sign"] for s in plan["sections"]], [1, -1, -1, -1])
        self.assertAlmostEqual(plan["parityTargets"]["Net Income"], built["targets"]["Net Income"], places=6)
        check_emitted(self, plan)

    def helper_total_workbook(self):
        """Total Costs = COGS + Opex; Profit = Revenue - Total Costs (no Net Income row)."""
        lines = fixture.DEFAULT_LINES[:5]
        clean = self.dir / "clean.xlsx"
        base = fixture.build_pnl(clean, lines=lines)
        values = base["values"]
        rev = [a + b for a, b in zip(values["Subscription"], values["Services"])]
        cogs = [a + b for a, b in zip(values["Hosting"], values["Support"])]
        opex = values["R&D"]
        costs = [c + o for c, o in zip(cogs, opex)]
        profit = [r - c for r, c in zip(rev, costs)]
        shape = {}

        def mutate(pnl, gl, drivers):
            rows = label_rows(pnl)
            gp, ebitda, net = rows["Gross Profit"], rows["EBITDA"], rows["Net Income"]
            shape.update(costs=gp, profit=ebitda)
            pnl.cell(gp, 1, "Total Costs")
            pnl.cell(ebitda, 1, "Profit")
            for i in range(6):
                col = chr(ord("B") + i)
                pnl[f"{col}{gp}"] = f"={col}{rows['Total Cost of Revenue']}+{col}{rows['Total Operating Expenses']}"
                pnl[f"{col}{ebitda}"] = f"={col}{rows['Total Revenue']}-{col}{gp}"
            for col in range(1, 9):
                pnl.cell(net, col).value = None
        from openpyxl import load_workbook
        clean_rows = label_rows(load_workbook(clean)["P&L"])
        gp_row, ebitda_row = clean_rows["Gross Profit"], clean_rows["EBITDA"]
        overrides = {}
        for i in range(6):
            col = chr(ord("B") + i)
            overrides[("P&L", f"{col}{gp_row}")] = costs[i]
            overrides[("P&L", f"{col}{ebitda_row}")] = profit[i]
        overrides[("P&L", f"H{gp_row}")], overrides[("P&L", f"H{ebitda_row}")] = sum(costs), sum(profit)
        fixture.build_pnl(self.path, lines=lines, mutate=mutate, cached_overrides=overrides)
        return costs, profit

    def test_helper_total_row_does_not_make_signs_ambiguous(self):
        costs, profit = self.helper_total_workbook()
        plan, _ = self.draft()
        self.assertEqual([(s["name"], s["sign"], s["signSource"]) for s in plan["sections"]],
                         [("Revenue", 1, "derived-rows"), ("Cost of Revenue", -1, "derived-rows"),
                          ("Operating Expenses", -1, "derived-rows")])
        by_label = {d["label"]: d for d in plan["derivedRows"]}
        self.assertEqual(by_label["Total Costs"]["coefficients"], {"Cost of Revenue": 1, "Operating Expenses": 1})
        self.assertEqual(by_label["Profit"]["coefficients"],
                         {"Revenue": 1, "Cost of Revenue": -1, "Operating Expenses": -1})
        self.assertAlmostEqual(plan["parityTargets"]["Total Costs"], sum(costs), places=6)
        self.assertAlmostEqual(plan["parityTargets"]["Profit"], sum(profit), places=6)
        self.assertEqual(BUILD.parity(plan)["Profit"], plan["parityTargets"]["Profit"])
        check_emitted(self, plan)

    def test_helper_total_layout_cards(self):
        self.helper_total_workbook()
        plan, _ = self.draft()
        spec = BUILD.build_wb_spec(plan, "dm", ELS)
        cards = {el["id"]: card_name(el) for page in spec["pages"] for el in page["elements"]
                 if el["kind"] == "kpi-chart"}
        self.assertEqual(cards["kpiNI"], "profit")        # reads every section
        self.assertEqual(cards["kpiEbitda"], "total costs")

    def test_opposite_signs_across_bottom_line_rows_are_still_ambiguous(self):
        derived = [{"label": "A", "kind": "amount", "coefficients": {"S": 1}, "expr": ["sec", "S"]},
                   {"label": "B", "kind": "amount", "coefficients": {"S": -1},
                    "expr": ["neg", ["sec", "S"]]}]
        with self.assertRaisesRegex(PlanDraftError, "section 'S' has ambiguous sign"):
            forecast_plan._infer_signs(["S"], derived)

    def test_partial_line_use_is_not_folded_into_the_section(self):
        terms = {("line", "A"): 1.0, ("line", "B"): 1.0}
        self.assertEqual(calc.fold_terms(terms, {"S": ["A", "B"]}), ({"S": 1.0}, {}))
        self.assertEqual(calc.fold_terms(terms, {"S": ["A", "B", "C"]}), ({}, {"A": 1.0, "B": 1.0}))
        # a section total plus one extra line keeps both, not 2x the line's share
        mixed = {("sec", "S"): 1.0, ("line", "A"): 1.0}
        self.assertEqual(calc.fold_terms(mixed, {"S": ["A", "B"]}), ({"S": 1.0}, {"A": 1.0}))
        self.assertIsNone(calc.linear_over_sections(["line", "A"], {"A": "S", "B": "S"}, {}))
        self.assertEqual(calc.linear_over_sections(["line", "A"], {"A": "S"}, {}), {"S": 1.0})


# ------------------------------------------------------------------ 2. KPI roles
class TestKpiRoles(unittest.TestCase):
    def rows(self, spec):
        out = []
        for label, coefficients, lines in spec:
            row = {"label": label, "coefficients": coefficients}
            if lines:
                row["lineCoefficients"] = lines
            out.append(row)
        return out

    def plan(self, rows):
        # each row reads the one before it, like GP -> Pretax in a real P&L
        return {"lines": [{"line": "L1", "section": "A"}, {"line": "L2", "section": "A"},
                          {"line": "M1", "section": "B"}, {"line": "N1", "section": "C"}],
                "derivedRows": [{"label": r["label"],
                                 "expr": ["drv", self.reads[r["label"]]] if r["label"] in self.reads
                                 else ["num", 0]} for r in rows]}

    reads = {"Pretax": "GP"}

    def test_the_row_that_reads_every_section_takes_the_bottom_line_card(self):
        rows = self.rows([("GP", {"A": 1, "B": -1}, {}), ("Pretax", {"A": 1, "B": -1, "C": -1}, {}),
                          ("Odd", {"B": -1}, {"L1": 1})])
        self.assertEqual(BUILD.kpi_rows(self.plan(rows), rows), (0, 1))

    def test_a_net_income_label_wins_a_tie(self):
        rows = self.rows([("Net Income", {"A": 1, "B": -1, "C": -1}, {}),
                          ("After Tax Adj", {"A": 1, "B": -1, "C": -1}, {})])
        self.assertEqual(BUILD.kpi_rows(self.plan(rows), rows)[1], 0)

    def test_classic_rows_keep_the_historical_cards(self):
        rows = self.rows([("Gross Profit", {"A": 1, "B": -1}, {}),
                          ("EBITDA", {"A": 1, "B": -1, "C": -1}, {}),
                          ("Net Income", {"A": 1, "B": -1, "C": -1, "D": -1}, {})])
        self.assertEqual(BUILD.kpi_rows({"lines": [], "derivedRows": []}, rows), (1, 2))

    def test_a_single_row_has_only_a_bottom_card(self):
        rows = self.rows([("Net Result", {"A": 1, "B": -1}, {})])
        self.assertEqual(BUILD.kpi_rows(self.plan(rows), rows), (None, 0))

    def test_regression_fixtures_keep_their_cards(self):
        for source, expected in ((SAAS, ["Total Revenue", "EBITDA", "Net Income", "Gross Margin %"]),
                                 (RETAIL, ["FY26 Total Net Sales", "FY26 Operating Profit",
                                           "FY26 Pre-tax Income", "FY26 Gross Margin %"])):
            with self.subTest(source=source.name):
                plan, _ = draft_plan(source, intent="what-if", connection="c", folder="f")
                spec = BUILD.build_wb_spec(plan, "dm", ELS)
                self.assertEqual([card_name(el) for page in spec["pages"] for el in page["elements"]
                                  if el["kind"] == "kpi-chart"], [name.casefold() for name in expected])
                check_emitted(self, plan)


# ------------------------------------------------------------------ 3. crashes
class TestRowFormulaArity(Base):
    def test_arity_errors_are_expression_errors(self):
        for text in ("=IF(B7)", "=IF(B7,1,2,3)", "=ABS()", "=ABS(1,2)", "=MAX()", "=MIN()", "=IFERROR(B7)",
                     "=SUM()"):
            with self.subTest(text=text):
                with self.assertRaisesRegex(calc.ExprError, "takes"):
                    calc.parse_formula(text)

    def test_hand_edited_plan_expressions_are_arity_checked_too(self):
        for expr in (["call", "IF", [["num", 1]]], ["call", "ABS", []], ["call", "MAX", []],
                     ["call", "IFERROR", [["num", 1]]]):
            with self.subTest(expr=expr):
                with self.assertRaises(calc.ExprError):
                    calc.eval_expr(expr, None)

    def test_two_argument_if_is_false_without_an_else(self):
        env = mock.Mock()
        expr = calc.to_entities(calc.parse_formula("=IF(1>2,5)"), mock.Mock())
        self.assertIs(calc.eval_expr(expr, env), False)
        expr = calc.to_entities(calc.parse_formula("=IF(3>2,5)"), mock.Mock())
        self.assertEqual(calc.eval_expr(expr, env), 5.0)
        # FALSE reads as 0 in arithmetic, as in Excel
        expr = calc.to_entities(calc.parse_formula("=IF(1>2,5)+7"), mock.Mock())
        self.assertEqual(calc.eval_expr(expr, env), 7.0)

    def if2_workbook(self, cached):
        def mutate(pnl, gl, drivers):
            pnl.cell(26, 1, "Big month")
            for i in range(6):
                col = chr(ord("B") + i)
                pnl[f"{col}26"] = f"=IF({col}7>100000,1)"
        overrides = {("P&L", f"{chr(ord('B') + i)}26"): cached for i in range(6)} if cached is not None else {}
        fixture.build_pnl(self.path, mutate=mutate, cached_overrides=overrides)

    def test_two_argument_if_with_cached_false_values_drafts(self):
        self.if2_workbook(0.0)
        plan, _ = self.draft()
        row = next(d for d in plan["derivedRows"] if d["label"] == "Big month")
        self.assertEqual(row["kind"], "ratio")
        self.assertEqual(row["formula"], "IF([Revenue] > 100000, 1)")

    def test_a_derived_row_without_cached_values_is_rejected_not_skipped(self):
        self.if2_workbook(None)
        with self.assertRaisesRegex(PlanDraftError, r"P&L!B26 \('Big month', Jan-26\).*no numeric cached value"):
            self.draft()

    def test_router_routes_the_probe_to_cell_model_instead_of_crashing(self):
        self.if2_workbook(None)
        code, result = self.plan(self.path)
        self.assertEqual(code, 0)
        self.assertEqual(result["archetype"], "cell-model")
        self.assertIn("forecast drafter rejects the shape", result["routeReason"])

    def test_draft_cli_reports_a_failure_instead_of_a_traceback(self):
        self.if2_workbook(None)
        err = io.StringIO()
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(err):
            code = DRAFT.main([str(self.path), "--connection", "c", "--folder", "f",
                               "--out", str(self.dir / "plan.json")])
        self.assertEqual(code, 1)
        self.assertIn("DRAFT FAILED: ", err.getvalue())

    def test_any_unexpected_failure_is_a_draft_error_naming_the_last_cell_read(self):
        fixture.build_pnl(self.path)
        with mock.patch.object(forecast_plan, "_derive_rows", side_effect=AttributeError("boom")):
            with self.assertRaisesRegex(PlanDraftError, r"P&L![A-Z]+\d+ .*unexpected AttributeError: boom"):
                self.draft()
        with mock.patch.object(forecast_plan, "find_period_header", side_effect=KeyError("x")):
            with self.assertRaisesRegex(PlanDraftError, "unexpected KeyError"):
                self.draft()

    def test_unreadable_workbook_is_a_draft_error(self):
        self.path.write_text("not a workbook", encoding="utf-8")
        with self.assertRaisesRegex(PlanDraftError, "unexpected"):
            self.draft()

    def test_router_survives_a_drafter_defect(self):
        fixture.build_pnl(self.path)
        with mock.patch.object(MIGRATE, "draft_plan", side_effect=ZeroDivisionError("division by zero")):
            code, result = self.plan(self.path)
        self.assertEqual(code, 0)
        self.assertEqual(result["archetype"], "cell-model")
        self.assertIn("could not run (ZeroDivisionError: division by zero)", result["routeReason"])

    def test_evaluation_errors_name_the_cell(self):
        def mutate(pnl, gl, drivers):
            rows = label_rows(pnl)
            pnl.cell(26, 1, "Ratio")
            for i in range(6):
                col = chr(ord("B") + i)
                pnl[f"{col}26"] = f"={col}{rows['Total Revenue']}/({col}{rows['Interest']}-{col}{rows['Interest']})"
        fixture.build_pnl(self.path, mutate=mutate)
        with self.assertRaisesRegex(PlanDraftError, r"P&L!B26 \('Ratio'\).*derived row 'Ratio' in Jan-26"):
            self.draft()

    def test_an_fy_division_by_zero_does_not_crash_the_builder_parity(self):
        fixture.build_pnl(self.path)
        plan, _ = self.draft()
        builder_calc = BUILD.calc          # the builder imports the lib copy of forecast_calc
        real = builder_calc.Evaluation.aggregate

        def aggregate(self, label, idxs):
            if label == "EBITDA":
                raise builder_calc.RowEvalError(label, 0, "the summed periods",
                                                ZeroDivisionError("division by zero"))
            return real(self, label, idxs)
        with mock.patch.object(builder_calc.Evaluation, "aggregate", aggregate):
            computed = BUILD.parity(plan)
        self.assertNotIn("EBITDA", computed)
        self.assertFalse(quiet_parity(plan, computed))   # a target without a row fails loudly


# ------------------------------------------------------------------ 5. tolerance
class TestTolerance(Base):
    def test_close_is_float_noise_only(self):
        self.assertTrue(calc.close(14402043.63 + 1e-9, 14402043.63))
        self.assertTrue(calc.close(0.3, 0.3 + 1e-12))
        self.assertFalse(calc.close(0.3, 0.304))          # four tenths of a point on a ratio
        self.assertFalse(calc.close(1000.0, 1000.005))     # half a cent on a small value
        self.assertTrue(calc.close(1000.0, 1000.005, tolerance=0.01))   # an explicit bound is honoured

    def test_a_half_cent_cache_drift_is_now_a_mismatch(self):
        fixture.build_pnl(self.path, cached_overrides={("P&L", "B5"): 1000.005})
        with self.assertRaisesRegex(PlanDraftError, r"P&L!B5 \('Subscription', Jan-26\)"):
            self.draft()
        plan, _ = self.draft(tolerance=0.01)               # the old bound stays available
        self.assertEqual(plan["lines"][0]["line"], "Subscription")

    def test_a_ratio_cache_drift_is_a_mismatch(self):
        built = fixture.build_quarterly(self.path)
        row = built["rows"]["Gross Margin %"]
        ratio = built["derived"]["Gross Margin %"][0]
        fixture.build_quarterly(self.path, cached_overrides={("Budget", f"B{row}"): ratio + 0.004})
        with self.assertRaisesRegex(PlanDraftError, rf"Budget!B{row} \('Gross Margin %', Q1 FY25\)"):
            self.draft()

    def test_fixtures_verify_at_the_tight_default(self):
        for source in (SAAS, RETAIL):
            with self.subTest(source=source.name):
                plan, _ = self.draft(source)
                self.assertTrue(quiet_parity(plan, BUILD.parity(plan)))

    def test_builder_default_parity_tolerance_is_tight(self):
        self.assertIsNone(BUILD.parse_args(["plan.json"]).parity_tolerance)
        self.assertIsNone(DRAFT.parse_args(["x.xlsx", "--connection", "c", "--folder", "f",
                                            "--out", "o"]).tolerance)


# ------------------------------------------------------------------ 6. formula-valued drivers
class TestFormulaDrivers(Base):
    def test_a_driver_formula_is_flagged_with_the_cell_and_formula(self):
        def mutate(pnl, gl, drivers):
            drivers["B4"] = "=0.6/12"
        fixture.build_pnl(self.path, mutate=mutate, cached_overrides={("Drivers", "B4"): 0.05})
        plan, summary = self.draft()
        issue = next(i for i in plan["draftReviewIssues"] if i["code"] == "DRIVER_FORMULA_FROZEN")
        self.assertEqual(issue["severity"], "review")
        self.assertEqual(len(issue["evidence"]), 1)
        self.assertIn("Drivers!B4: =0.6/12 (cached 0.05; read by Subscription)", issue["evidence"][0])
        self.assertIn("DRIVER_FORMULA_FROZEN", summary["reviewIssues"])
        self.assertEqual(plan["lines"][0]["rate"], 0.05)   # the cached rate is what the plan carries

    def test_typed_drivers_raise_no_issue(self):
        fixture.build_pnl(self.path)
        plan, _ = self.draft()
        self.assertNotIn("DRIVER_FORMULA_FROZEN", [i["code"] for i in plan["draftReviewIssues"]])

    def test_the_issue_reaches_the_build_gate(self):
        from migration_gate import plan_review_issues

        def mutate(pnl, gl, drivers):
            drivers["B4"] = "=0.6/12"
        fixture.build_pnl(self.path, mutate=mutate, cached_overrides={("Drivers", "B4"): 0.05})
        plan, _ = self.draft()
        self.assertIn("DRIVER_FORMULA_FROZEN", [i["code"] for i in plan_review_issues(plan)])


if __name__ == "__main__":
    unittest.main()
