"""Conditional errors and nullable logic through the exact emitted graph."""
import copy
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts/lib'))
from excel_expression import ExcelCalcError
from native_calculations import _ast
from native_exact import Literal, compile_expression, evaluate_bound
from native_exact_plan import build_exact_native_calculation_plan, lower_exact_native_calculations
from native_formula_reference import evaluate_native
from native_graph_reference import NativeGraphReference
from source_inputs import build_source_input_plan
from workbook_expression import WorkbookContext
import test_native_calculations


class ExactConditionalLogicTests(unittest.TestCase):
    source = test_native_calculations.NativeCalculationTests.source

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def lower_formulas(self, formulas):
        path, old = self.source()
        book = load_workbook(path)
        for cell, formula in formulas.items():
            book['Summary Renamed'][cell] = formula
        book.save(path)
        decisions = copy.deepcopy(old['decisions'])
        decisions['sourceSha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
        decisions['features'][0]['reviewedBlankColumns'] = ['Amount', 'Rate']
        inputs = build_source_input_plan(path, decisions, connection_id='connection', folder_id='folder')
        plan = build_exact_native_calculation_plan(path, inputs)
        rows = [{**row, 'Row ID': f'physical-{i}'} for i, row in
                enumerate(inputs['sharedInputArchitecture']['inputs']['ledger']['rows'])]
        assigned = {'ledger': {'elementId': 'observed-input', 'name': 'tblLedger',
                              'rowIdColumn': 'Row ID',
                              'rows': {str(i): row['Row ID'] for i, row in enumerate(rows)}}}
        lowered = lower_exact_native_calculations(plan, assigned, max_columns=8)
        return plan, lowered, rows

    def test_vector_if_propagates_condition_error_into_each_selected_cell(self):
        values = {'A1': 1, 'A2': 2, 'B1': 4, 'B2': 8}
        book = Workbook()
        context = WorkbookContext('memory.xlsx', book, 'test')
        for condition, error in [('1/0', '#DIV/0!'), ('"invalid"', '#VALUE!')]:
            formula = f'=SUM(IF({condition},A1:A2,B1:B2))'
            node = _ast(context.bind_formula(formula, sheet='Sheet', address='Z1'))
            with self.subTest(condition=condition):
                with self.assertRaisesRegex(ExcelCalcError, error):
                    evaluate_bound(node, lambda cell: values[cell.split('!')[1]])
                result = compile_expression(node, lambda cell: Literal.value(values[cell.split('!')[1]]))
                self.assertEqual(evaluate_native(result.error), error)

    def test_vector_if_error_survives_plan_lowering_and_outer_iferror_catches_it(self):
        plan, lowered, rows = self.lower_formulas({
            'A1': "=SUM(IF(1/0,'Ledger Renamed'!B2:B3,'Ledger Renamed'!C2:C3))",
            'A2': '=IFERROR(A1,17)',
        })
        self.assertEqual(plan['formulas']['Summary Renamed!A1']['localExpectedError'], '#DIV/0!')
        self.assertEqual(plan['formulas']['Summary Renamed!A2']['localExpected'], 17)
        graph = NativeGraphReference({'observed-input': ('tblLedger', rows)}).materialize(lowered['shards'])
        self.assertEqual(graph.output(lowered, 'Summary Renamed!A1')['error'], '#DIV/0!')
        caught = graph.output(lowered, 'Summary Renamed!A2')
        self.assertEqual((caught['kind'], caught['number'], caught['error']), ('n', 17, ''))

    def test_and_or_nullable_range_rechecks_eligible_values_after_runtime_edits(self):
        formulas = {
            'A1': "=AND('Ledger Renamed'!B2:B3)",
            'A2': "=OR('Ledger Renamed'!B2:B3)",
            'A3': "=AND('Ledger Renamed'!B2:B3,TRUE)",
            'A4': "=OR('Ledger Renamed'!B2:B3,FALSE)",
            'A5': "=AND('Ledger Renamed'!B2:B3,1/0)",
        }
        plan, lowered, rows = self.lower_formulas(formulas)
        for amounts, expected in [((3, 4), (True, True)), ((None, 0), (False, False)),
                                  ((None, 4), (True, True)), ((None, None), ('#VALUE!', '#VALUE!'))]:
            changed = copy.deepcopy(rows)
            for row, amount in zip(changed, amounts):
                row['Amount'] = amount
            graph = NativeGraphReference({'observed-input': ('tblLedger', changed)}).materialize(lowered['shards'])
            for cell, want in zip(('A1', 'A2'), expected):
                record = plan['formulas']['Summary Renamed!' + cell]
                source_values = {'Ledger Renamed!B2': amounts[0], 'Ledger Renamed!B3': amounts[1]}
                with self.subTest(amounts=amounts, cell=cell):
                    result = graph.output(lowered, 'Summary Renamed!' + cell)
                    if isinstance(want, str):
                        with self.assertRaisesRegex(ExcelCalcError, want):
                            evaluate_bound(record['ast'], source_values.__getitem__)
                        self.assertEqual(result['error'], want)
                    else:
                        self.assertEqual(evaluate_bound(record['ast'], source_values.__getitem__), want)
                        self.assertEqual((result['logical'], result['error']), (want, ''))
            # Direct logical arguments remain eligible even when every referenced
            # cell is blank; real calculation errors retain precedence.
            if amounts == (None, None):
                for cell, want in [('A3', True), ('A4', False)]:
                    result = graph.output(lowered, 'Summary Renamed!' + cell)
                    self.assertEqual((result['logical'], result['error']), (want, ''))
                self.assertEqual(graph.output(lowered, 'Summary Renamed!A5')['error'], '#DIV/0!')


if __name__ == '__main__':
    unittest.main()
