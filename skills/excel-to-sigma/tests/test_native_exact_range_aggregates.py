"""Exact emitted aggregate graphs retain source domains and date serials."""
import copy
import datetime as dt
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts/lib'))
from native_exact_plan import build_exact_native_calculation_plan, lower_exact_native_calculations
from native_graph_reference import NativeGraphReference
from source_inputs import build_source_input_plan
import test_native_calculations


class ExactRangeAggregateTests(unittest.TestCase):
    source = test_native_calculations.NativeCalculationTests.source

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def lower(self, path, decisions):
        decisions = copy.deepcopy(decisions)
        decisions['sourceSha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
        inputs = build_source_input_plan(path, decisions, connection_id='connection', folder_id='folder')
        plan = build_exact_native_calculation_plan(path, inputs)
        rows = [{**row, 'Row ID': f'physical-{i}'} for i, row in
                enumerate(inputs['sharedInputArchitecture']['inputs']['ledger']['rows'])]
        assigned = {'ledger': {'elementId': 'observed-input', 'name': 'tblLedger',
                              'rowIdColumn': 'Row ID',
                              'rows': {str(i): row['Row ID'] for i, row in enumerate(rows)}}}
        return plan, lower_exact_native_calculations(plan, assigned, max_columns=8), rows

    def assert_outputs(self, lowered, rows, expected):
        graph = NativeGraphReference({'observed-input': ('tblLedger', rows)}).materialize(lowered['shards'])
        for address, number in expected.items():
            with self.subTest(address=address):
                result = graph.output(lowered, 'Summary Renamed!' + address)
                self.assertEqual(result['error'], '')
                self.assertEqual(result['kind'], 'n')
                self.assertEqual(result['number'], number)

    def test_calculated_subranges_use_the_requested_physical_rows(self):
        path, old = self.source()
        book = load_workbook(path)
        ledger = book['Ledger Renamed']
        for row in (2, 3):
            ledger[f'D{row}'] = f'=B{row}'
        ledger.append(['c', 9, 1, '=B4'])
        ledger.tables['tblLedger'].ref = 'A1:D4'
        formulas = {
            'A1': "=SUM('Ledger Renamed'!D2:D2)",
            'A2': "=COUNT('Ledger Renamed'!D2:D2)",
            'A3': "=SUM('Ledger Renamed'!D3:D3)",
            'A4': "=COUNT('Ledger Renamed'!D3:D3)",
            'A5': "=SUM('Ledger Renamed'!D2:D4)",
            'A6': "=COUNT('Ledger Renamed'!D2:D4)",
            'A7': "=SUM('Ledger Renamed'!D4:D2)",
            'A8': "=SUM('Ledger Renamed'!D2:D2,'Ledger Renamed'!D2:D2)",
            'A9': "=SUM('Ledger Renamed'!D4:D4,'Ledger Renamed'!D2:D2)",
            'A10': '=SUM(tblLedger[Result])',
            'A11': '=COUNT(tblLedger[Result])',
            'A12': '=SUMIFS(\'Ledger Renamed\'!D3:D3,\'Ledger Renamed\'!A3:A3,"b")',
            'A13': "=COUNT('Ledger Renamed'!D2:D2,'Ledger Renamed'!D2:D2)",
            'A14': "=COUNT('Ledger Renamed'!D4:D4,'Ledger Renamed'!D2:D2)",
        }
        for address, formula in formulas.items():
            book['Summary Renamed'][address] = formula
        book.save(path)
        plan, lowered, rows = self.lower(path, old['decisions'])
        expected = {'A1': 3, 'A2': 1, 'A3': 4, 'A4': 1, 'A5': 16, 'A6': 3,
                    'A7': 16, 'A8': 6, 'A9': 12, 'A10': 16, 'A11': 3, 'A12': 4,
                    'A13': 2, 'A14': 2}
        for address, number in expected.items():
            self.assertEqual(plan['formulas']['Summary Renamed!' + address]['localExpected'], number)
        self.assert_outputs(lowered, rows, expected)
        # Physical IDs stay authoritative when queried rows arrive in another
        # order; only formal data-body domains include appended rows.
        rows[1]['Amount'] = 10
        rows = [rows[2], rows[1], rows[0],
                {'Account': 'd', 'Amount': 20, 'Rate': 1, 'Row ID': 'appended'}]
        self.assert_outputs(lowered, rows, {**expected, 'A3': 10, 'A5': 22, 'A7': 22,
                                           'A10': 42, 'A11': 4, 'A12': 10})

    def test_ambiguous_physical_row_ids_are_rejected(self):
        path, inputs = self.source()
        plan = build_exact_native_calculation_plan(path, inputs)
        assigned = {'ledger': {'elementId': 'observed-input', 'name': 'tblLedger',
                              'rowIdColumn': 'Row ID',
                              'rows': {'0': 'same-physical-id', '1': 'same-physical-id'}}}
        with self.assertRaisesRegex(ValueError, 'unique.*row ID|row ID.*unique'):
            lower_exact_native_calculations(plan, assigned)

    def test_raw_datetime_sum_uses_source_epoch_serials_and_ignores_blanks(self):
        for epoch, fraction, blank, total in [('1900', False, False, 90585),
                                               ('1904', False, False, 87661),
                                               ('1900', True, False, 90585.75),
                                               ('1904', True, False, 87661.75),
                                               ('1900', False, True, 45293)]:
            with self.subTest(epoch=epoch, fraction=fraction, blank=blank):
                path, old = self.source()
                book = load_workbook(path)
                if epoch == '1904':
                    book.epoch = CALENDAR_MAC_1904
                book['Ledger Renamed']['B2'] = None if blank else dt.datetime(2024, 1, 1, 12 if fraction else 0)
                book['Ledger Renamed']['B3'] = dt.datetime(2024, 1, 2, 6 if fraction else 0)
                for row in (2, 3):
                    book['Ledger Renamed'][f'D{row}'] = '=1'
                book['Summary Renamed']['A1'] = "=SUM('Ledger Renamed'!B2:B3)"
                book['Summary Renamed']['A2'] = '=SUMIFS(tblLedger[Amount],tblLedger[Account],"<>")'
                book.save(path)
                decisions = copy.deepcopy(old['decisions'])
                decisions['features'][0]['columnTypes']['Amount'] = 'datetime'
                if blank:
                    decisions['features'][0]['reviewedBlankColumns'] = ['Amount']
                plan, lowered, rows = self.lower(path, decisions)
                for address in ('A1', 'A2'):
                    self.assertEqual(plan['formulas']['Summary Renamed!' + address]['localExpected'], total)
                for row in rows:
                    if row['Amount'] is not None:
                        row['Amount'] = dt.datetime.fromisoformat(row['Amount'])
                self.assert_outputs(lowered, rows, {'A1': total, 'A2': total})


if __name__ == '__main__':
    unittest.main()
