"""A cleared numeric source literal is raw NULL, but Excel arithmetic reads it as zero."""
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / "lib"))
from excel_expression import UnsupportedExpression, emit_sigma, evaluate, parse_expression  # noqa: E402
from formula_graph import CellIdentity, FormulaGraph  # noqa: E402
from native_graph_oracle import scalar  # noqa: E402


def load_script(name):
    spec = importlib.util.spec_from_file_location(name.replace("-", "_"), SCRIPTS / name)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


INFER = load_script("infer-canonical-formulas.py")
RESEARCH = load_script("build-research-model.py")
FIXTURE = ROOT / "tests" / "fixtures" / "fixture-research-model.xlsx"


class NumericBlankReferenceTests(unittest.TestCase):
    def test_source_c3_blank_reference_is_zero_in_arithmetic_but_raw_input_stays_null(self):
        source = load_workbook(FIXTURE, data_only=False)
        self.assertEqual(source["FY results"]["I8"].value,
                         "=H8*(1+'Model Inputs'!$C$3)")
        expression = parse_expression(source["FY results"]["I8"].value, sheet="FY results")
        values = {("FY results", "H8"): 1158 * (1 + 0.0725),
                  ("Model Inputs", "C3"): None}
        numeric_source_literal = lambda ref: (ref.sheet, ref.address) == ("Model Inputs", "C3")
        def resolve(ref):
            return values[(ref.sheet, ref.address)]
        def arithmetic_value(ref, value):
            return 0 if numeric_source_literal(ref) and value is None else value
        def field(ref):
            return "[" + ref.sheet + "!" + ref.address + "]"
        def arithmetic_field(ref):
            return "Coalesce(" + field(ref) + ", 0)" if numeric_source_literal(ref) else field(ref)

        direct = parse_expression("='Model Inputs'!C3", sheet="FY results")
        self.assertIsNone(evaluate(direct, resolve,
                                   arithmetic_reference_resolver=arithmetic_value))
        self.assertAlmostEqual(evaluate(expression, resolve,
                                        arithmetic_reference_resolver=arithmetic_value), 1241.955)
        emitted = emit_sigma(expression, field,
                             arithmetic_reference_renderer=arithmetic_field)
        self.assertIn("Coalesce([Model Inputs!C3], 0)", emitted)
        self.assertAlmostEqual(scalar(emitted, {"FY results!H8": values[("FY results", "H8")],
                                                "Model Inputs!C3": None}), 1241.955)
        self.assertEqual(emit_sigma(direct, field,
                                    arithmetic_reference_renderer=arithmetic_field),
                         "[Model Inputs!C3]")

        graph_values = FormulaGraph.from_workbook(source).evaluate({
            CellIdentity("Model Inputs", "C3"): None,
        })
        self.assertIsNone(graph_values[CellIdentity("Model Inputs", "C3")])
        self.assertAlmostEqual(graph_values[CellIdentity("FY results", "I8")], 1241.955)
        self.assertAlmostEqual(graph_values[CellIdentity("FY results", "J8")], 1313.98839)

    def test_derived_error_is_not_coalesced_and_aggregate_blank_is_ignored(self):
        raw = lambda ref: "[" + ref.address + "]"
        arithmetic = lambda ref: "Coalesce(" + raw(ref) + ", 0)" if ref.address == "C3" else raw(ref)
        formula = emit_sigma(parse_expression("=A1+C3", sheet="Model"), raw,
                             arithmetic_reference_renderer=arithmetic)
        self.assertEqual(formula, "[A1] + Coalesce([C3], 0)")
        self.assertIsNone(scalar(formula, {"A1": None, "C3": None}))
        with self.assertRaises(ZeroDivisionError):
            evaluate(parse_expression("=1/(A1-1)+C3", sheet="Model"),
                     lambda ref: {"A1": 1, "C3": None}[ref.address],
                     arithmetic_reference_resolver=lambda ref, value:
                     0 if ref.address == "C3" and value is None else value)
        for name in ("MIN", "MAX"):
            expression = parse_expression(f"={name}(A1,B1)", sheet="Model")
            self.assertEqual(emit_sigma(expression, raw,
                                        arithmetic_reference_renderer=arithmetic),
                             f"{name.title()}([A1], [B1])")
            self.assertEqual(evaluate(expression, lambda ref: {"A1": None, "B1": 5}[ref.address],
                                      arithmetic_reference_resolver=lambda ref, value: value), 5)
            self.assertEqual(evaluate(expression, lambda _ref: None), 0)
        summed = parse_expression("=SUM(A1,C3)", sheet="Model")
        self.assertEqual(emit_sigma(summed, raw,
                                    arithmetic_reference_renderer=arithmetic),
                         "Sum([A1], [C3])")
        self.assertEqual(evaluate(summed, lambda ref: {"A1": 5, "C3": None}[ref.address]), 5)
        rounded = parse_expression("=ROUND(C3,2)", sheet="Model")
        self.assertEqual(emit_sigma(rounded, raw,
                                    arithmetic_reference_renderer=arithmetic),
                         "Round(Coalesce([C3], 0), 2)")
        self.assertEqual(evaluate(rounded, lambda _ref: None,
                                  arithmetic_reference_resolver=lambda _ref, value:
                                  0 if value is None else value), 0)
        for name in ("COUNT", "COUNTA", "AVERAGE"):
            with self.assertRaises(UnsupportedExpression):
                parse_expression(f"={name}(A1,B1)", sheet="Model")

    def test_formula_links_compute_zero_for_blank_source_without_coercing_raw_inputs_or_errors(self):
        source = Workbook()
        ws = source.active
        ws.title = "Model"
        other = source.create_sheet("Other")
        ws["A1"] = 7
        ws["B1"] = "=A1"
        ws["C1"] = "=B1+1"
        other["A1"] = "='Model'!B1"
        other["B1"] = "=A1+1"
        other["C1"] = "=IF(1,A1,1/0)+1"
        ws["D1"] = "=1/0"
        ws["E1"] = "=D1"
        other["D1"] = "='Model'!E1+A1"
        other["E1"] = "=MIN(A1,5)"
        ws["F1"] = "=MIN(A1,5)"
        values = FormulaGraph.from_workbook(source).evaluate({CellIdentity("Model", "A1"): None})
        self.assertIsNone(values[CellIdentity("Model", "A1")])
        for sheet, address in (("Model", "B1"), ("Other", "A1"), ("Other", "E1")):
            with self.subTest(sheet=sheet, address=address):
                self.assertEqual(values[CellIdentity(sheet, address)], 0)
        self.assertEqual(values[CellIdentity("Model", "F1")], 5)
        for sheet, address in (("Model", "C1"), ("Other", "B1"), ("Other", "C1")):
            with self.subTest(sheet=sheet, address=address):
                self.assertEqual(values.get(CellIdentity(sheet, address)), 1)
        self.assertNotIn(CellIdentity("Other", "D1"), values)

    def test_research_carried_c3_blank_is_zero_in_source_and_emitted_arithmetic(self):
        source = load_workbook(FIXTURE, data_only=False)
        source["Model Inputs"]["D3"] = "=C3"
        source["FY results"]["I8"] = "=H8*(1+'Model Inputs'!$D$3)"
        values = FormulaGraph.from_workbook(source).evaluate({CellIdentity("Model Inputs", "C3"): None})
        self.assertEqual(values[CellIdentity("Model Inputs", "D3")], 0)
        with self.subTest(path="source evaluation"):
            self.assertEqual(values.get(CellIdentity("FY results", "I8")), 1241.955)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "carried-blank.xlsx"
            source.save(path)
            plan = INFER.infer(path, sheet="FY results")
        revenue = next(line for line in plan["lines"] if line["row"] == 8)
        formula = revenue["periodFormulas"]["2026"]
        native_value = scalar(formula, {"i_fy_results_g8": 1158,
                                        "i_model_inputs_c2": .0725,
                                        "i_model_inputs_c3": None})
        self.assertIsNotNone(native_value)
        self.assertAlmostEqual(native_value, 1241.955)

    def test_research_projected_alias_is_zero_and_aggregates_distinguish_literal_blank(self):
        for formula, expected in (("='Model Inputs'!D3", 0),
                                  ("=MIN('Model Inputs'!D3,5)", 0),
                                  ("=MIN('Model Inputs'!C3,5)", 5)):
            with self.subTest(formula=formula):
                source = load_workbook(FIXTURE, data_only=False)
                source["Model Inputs"]["D3"] = "=C3"
                source["FY results"]["I8"] = formula
                values = FormulaGraph.from_workbook(source).evaluate({CellIdentity("Model Inputs", "C3"): None})
                self.assertIsNone(values[CellIdentity("Model Inputs", "C3")])
                with self.subTest(path="source formula result"):
                    self.assertEqual(values.get(CellIdentity("FY results", "I8")), expected)
                with tempfile.TemporaryDirectory() as directory:
                    path = Path(directory) / "projected-alias.xlsx"
                    source.save(path)
                    plan = INFER.infer(path, sheet="FY results")
                revenue = next(line for line in plan["lines"] if line["row"] == 8)
                emitted = revenue["periodFormulas"]["2026"]
                self.assertEqual(scalar(emitted, {"i_model_inputs_c3": None}), expected)

    def test_aggregates_propagate_derived_errors_but_ignore_true_blank_literals(self):
        source = Workbook()
        ws = source.active
        ws.title = "Model"
        ws["A1"] = 7
        ws["D1"] = "=1/0"
        ws["D2"] = "=D1"
        ws["D3"] = "=AVERAGE(A1,5)"
        for row, dependency in enumerate(("A1", "D1", "D2", "D3", "Z9"), 1):
            for column, name in (("E", "MIN"), ("F", "MAX"), ("G", "SUM")):
                ws[f"{column}{row}"] = f"={name}({dependency},5)"
        values = FormulaGraph.from_workbook(source).evaluate({CellIdentity("Model", "A1"): None})
        for column in "EFG":
            self.assertEqual(values[CellIdentity("Model", f"{column}1")], 5)
            for row in range(2, 6):
                with self.subTest(column=column, row=row):
                    self.assertFalse(CellIdentity("Model", f"{column}{row}") in values)

    def test_range_aggregates_ignore_empty_cells_but_propagate_formula_errors(self):
        source = Workbook()
        ws = source.active
        ws.title = "Model"
        ws["A1"] = 3
        ws["A3"] = 7
        ws["B1"] = 1
        ws["B3"] = 1
        ws["C1"] = "=SUM(A1:A3)"
        ws["C2"] = "=MIN(A1:A3)"
        ws["C3"] = "=MAX(A1:A3)"
        ws["C4"] = "=SUMIFS(A1:A3,B1:B3,1)"
        ws["D1"] = "=1/0"
        ws["D2"] = "=SUM(D1:D1)"

        values = FormulaGraph.from_workbook(source).evaluate()

        self.assertEqual([values[CellIdentity("Model", f"C{row}")] for row in range(1, 5)],
                         [10, 3, 7, 10])
        self.assertNotIn(CellIdentity("Model", "D2"), values)

    def test_research_range_sum_coalesces_numeric_source_leaves(self):
        source = load_workbook(FIXTURE, data_only=False)
        source["FY results"]["I8"] = "=SUM('Model Inputs'!C2:C3)"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "range-sum.xlsx"
            source.save(path)
            plan = INFER.infer(path, sheet="FY results")

        revenue = next(line for line in plan["lines"] if line["row"] == 8)
        formula = revenue["periodFormulas"]["2026"]
        self.assertEqual(
            formula,
            "Sum((Coalesce([i_model_inputs_c2], 0) + Coalesce([i_model_inputs_c3], 0)))",
        )
        self.assertEqual(scalar(formula, {"i_model_inputs_c2": 0.0725,
                                          "i_model_inputs_c3": None}), 0.0725)

    def test_research_aggregate_emission_preserves_division_error_through_formula_link(self):
        for name in ("MIN", "MAX", "SUM"):
            with self.subTest(name=name):
                source = load_workbook(FIXTURE, data_only=False)
                source["Model Inputs"]["D3"] = "=1/0"
                source["Model Inputs"]["E3"] = "=D3"
                source["FY results"]["I8"] = f"={name}('Model Inputs'!$E$3,5)"
                values = FormulaGraph.from_workbook(source).evaluate()
                with self.subTest(path="source evaluation"):
                    self.assertFalse(CellIdentity("FY results", "I8") in values)
                with tempfile.TemporaryDirectory() as directory:
                    path = Path(directory) / "derived-error.xlsx"
                    source.save(path)
                    plan = INFER.infer(path, sheet="FY results")
                revenue = next(line for line in plan["lines"] if line["row"] == 8)
                formula = revenue["periodFormulas"]["2026"]
                with self.assertRaises(ZeroDivisionError):
                    scalar(formula, {})

    def test_fixture_emission_keeps_raw_input_and_coerces_only_arithmetic_reads(self):
        plan = INFER.infer(FIXTURE)
        plan["targetScope"] = {"connectionId": "connection", "folderId": "folder"}
        _dm, data_lines = RESEARCH.build_dm(plan, "connection", "folder", "Research")
        workbook = RESEARCH.build_wb(plan, "dm", {"data": "data", "linedim": "linedim"},
                                     data_lines, "folder", "Research")
        calc = next(e for page in workbook["pages"] for e in page["elements"] if e["id"] == "calc")
        formulas = {column["id"]: column["formula"] for column in calc["columns"]}
        self.assertEqual(formulas["i_model_inputs_c3"],
                         "If(IsNull([Inputs/Scope]), [CJ/i_model_inputs_c3], [Inputs/i_model_inputs_c3])")
        self.assertIn("Coalesce([i_model_inputs_c3], 0)", formulas["c8"])
        self.assertNotIn("Coalesce([i_model_inputs_c3], 0)", formulas["i_model_inputs_c3"])


if __name__ == "__main__":
    unittest.main()
