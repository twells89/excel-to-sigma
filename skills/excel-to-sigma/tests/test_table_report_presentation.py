#!/usr/bin/env python3
"""table-report presentation: source number formats, row order, header styles, mapped charts, KPI cards and
the Report page layout - on the 04 Sales Pipeline fixture and on synthetic workbooks."""
from __future__ import annotations

import contextlib
import copy
import importlib.util
import io
import json
import re
import sys
import xml.etree.ElementTree as ET
import tempfile
import unittest
from pathlib import Path

from openpyxl.chart import BarChart, DoughnutChart, Reference, ScatterChart, Series

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / "lib"))
sys.path.insert(0, str(ROOT / "tests"))

import excel_formats as ef  # noqa: E402
import presentation_style as ps  # noqa: E402
import table_report as tr  # noqa: E402
import table_report_fixture as fx  # noqa: E402
import table_report_model as trm  # noqa: E402
import table_report_presentation as pres  # noqa: E402
import workbook_wire  # noqa: E402

FIXTURES = ROOT / "tests" / "fixtures" / "broader-eight"
SOURCES = {"04": FIXTURES / "04 Sales Pipeline Forecast.xlsx",
           "06": FIXTURES / "06 13-Week Cash Flow.xlsx",
           "07": FIXTURES / "07 Headcount and Opex Plan.xlsx",
           "08": FIXTURES / "08 Inventory Replenishment Plan.xlsm"}
PIPELINE = SOURCES["04"]


def load_script(name, filename):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


DRAFT_CLI = load_script("draft_table_report_cli_pres", "draft-table-report.py")
BUILD_CLI = load_script("build_table_report_cli_pres", "build-table-report.py")


def run_cli(module, argv):
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        code = module.main(argv)
    return code, out.getvalue(), err.getvalue()


def draft(path, **options):
    return tr.draft_table_report(path, intent="what-if", connection="conn", folder="folder", **options)


def wb_spec(plan):
    return workbook_wire.wire_workbook(trm.build_wb_spec(plan, "<dm>", trm.local_dm_elements(plan)))


def element(spec, element_id):
    return next(e for e in spec["contents"]["elements"] if e["id"] == element_id)


def layout_tree(spec, page="report"):
    """Top-level layout nodes of one page: [(id, c0, c1, r0, r1, [children])] (children rows are relative)."""
    layout = spec["contents"]["layout"]
    root = ET.fromstring("<Root>" + layout.split("?>", 1)[1] + "</Root>")
    node = next(p for p in root if p.get("id") == page)

    def rect(el):
        c0, c1 = map(int, el.get("gridColumn").split(" / "))
        r0, r1 = map(int, el.get("gridRow").split(" / "))
        return el.get("elementId"), c0, c1, r0, r1
    return [(*rect(el), [rect(child) for child in el]) for el in node]


def placements(spec, page="report"):
    """Top-level (id, c0, c1, r0, r1) of one page (a container counts as one block)."""
    return [node[:5] for node in layout_tree(spec, page)]


def layout_ids(spec, page="report"):
    return [i for node in layout_tree(spec, page) for i in [node[0]] + [c[0] for c in node[5]]]


def order_positions(formula):
    """[(label, position)] a nested ``If([Key] = "x", n, ...)`` source-order formula assigns."""
    return [(m.group(1), int(m.group(2))) for m in re.finditer(r'= "([^"]+)", (\d+),', formula)]


def number_columns(plan):
    for table in plan["tables"]:
        for column in table["columns"]:
            if column["type"] == "number":
                yield table["name"], column
    for frame in plan["frames"]:
        for column in frame["columns"]:
            if column["type"] == "number":
                yield frame["title"], column


# --------------------------------------------------------------------------- 04 Sales Pipeline, end to end
class TestPipelinePresentation(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.plan, cls.info = draft(PIPELINE)
        cls.spec = wb_spec(cls.plan)
        cls.pres = cls.plan["presentation"]
        cls.frame = next(f for f in cls.plan["frames"] if f["title"] == "H2 FY26 Weighted Forecast")
        cls.stages = next(f for f in cls.plan["frames"] if f["title"] == "Open deals by stage")

    # ---- 1/2: formats on every numeric column, no inherited percent
    def test_source_formats_are_translated_into_the_plan(self):
        formats = {col["name"]: col["format"]["formatString"] for _, col in number_columns(self.plan)
                   if col["name"] in ("Weighted Amount", "Probability", "Coverage", "# Deals", "Pipeline $",
                                      "Commit", "Quota", "Amount", "Jul-26", "Total Weighted")}
        self.assertEqual(formats["Weighted Amount"], "(,.0f")          # #,##0;[Red](#,##0)
        self.assertEqual(formats["Jul-26"], "(,.0f")
        self.assertEqual(formats["Probability"], ".0%")                # 0%
        self.assertEqual(formats["Coverage"], ".1%")                   # 0.0%
        self.assertEqual(formats["Quota"], ",.0f")                     # General -> plain number
        self.assertEqual(formats["# Deals"], ",.0f")
        close = next(c for t in self.plan["tables"] if t["name"] == "tblPipeline" for c in t["columns"]
                     if c["name"] == "Close Date")
        self.assertEqual(close["format"], {"kind": "datetime", "formatString": "%Y-%m-%d"})

    def test_every_numeric_column_has_an_explicit_format(self):
        for owner, column in number_columns(self.plan):
            self.assertIsInstance(column.get("format"), dict, f"{owner}.{column['name']}")
        elements = {e["id"]: e for e in self.spec["contents"]["elements"] if "columns" in e}
        types = {(t["name"], c["name"]): c["type"] for t in self.plan["tables"] for c in t["columns"]}
        types.update({(f["title"], c["name"]): c["type"] for f in self.plan["frames"] for c in f["columns"]})
        checked = 0
        for el in elements.values():
            for column in el["columns"]:
                if isinstance(el.get("name"), str) and types.get((el["name"], column.get("name"))) == "number":
                    checked += 1
                    self.assertIn("format", column, f"{el['id']}.{column['name']}")
        self.assertGreater(checked, 15)

    def test_a_product_of_a_percent_and_a_money_column_does_not_inherit_the_percent(self):
        wrapper = element(self.spec, "w_tblpipeline")
        columns = {c["name"]: c for c in wrapper["columns"]}
        self.assertEqual(columns["Probability"]["format"]["formatString"], ".0%")
        self.assertEqual(columns["Weighted Amount"]["formula"], "([Amount] * [Probability])")
        self.assertEqual(columns["Weighted Amount"]["format"], {"kind": "number", "formatString": "(,.0f"})
        summary = element(self.spec, "f_s_h2_fy26_weighted_forecast")
        for column in summary["columns"]:
            if column["name"] in ("Jul-26", "Total Weighted", "Commit", "Quota"):
                self.assertFalse(column["format"]["formatString"].endswith("%"), column["name"])

    def test_red_negatives_become_a_conditional_format_rule(self):
        summary = element(self.spec, "f_s_h2_fy26_weighted_forecast")
        ids = {c["name"]: c["id"] for c in summary["columns"]}
        rule = summary["conditionalFormats"][0]
        self.assertEqual((rule["type"], rule["condition"], rule["value"]), ("single", "<", 0))
        self.assertIn(ids["Jul-26"], rule["columnIds"])
        self.assertNotIn(ids["Quota"], rule["columnIds"])             # General: no red section
        self.assertNotIn(ids["Coverage"], rule["columnIds"])
        self.assertEqual(rule["style"], {"color": ps.NEGATIVE})

    # ---- 3: row order
    def test_summary_rows_keep_the_sheet_order(self):
        summary = element(self.spec, "f_s_h2_fy26_weighted_forecast")
        order = next(c for c in summary["columns"] if c["name"] == "Source Order")
        self.assertTrue(order["hidden"])
        self.assertEqual(order_positions(order["formula"]),
                         [("Jordan Lee", 1), ("Sam Ortiz", 2), ("Priya Nair", 3), ("Alex Kim", 4)])
        self.assertEqual(summary["sort"], [{"columnId": order["id"], "direction": "ascending"}])
        self.assertIn(order["id"], summary["groupings"][0]["calculations"])
        stages = element(self.spec, "f_s_open_deals_by_stage")
        stage_order = next(c for c in stages["columns"] if c["name"] == "Source Order")
        self.assertEqual([label for label, _ in order_positions(stage_order["formula"])],
                         ["Prospect", "Discovery", "Proposal", "Negotiation", "Closed Won", "Closed Lost"])

    def test_the_order_formula_evaluates_to_each_row_s_position(self):
        from sigma_formula_eval import Ctx, Evaluator, parse
        summary = element(self.spec, "f_s_h2_fy26_weighted_forecast")
        formula = next(c["formula"] for c in summary["columns"] if c["name"] == "Source Order")
        evaluator = Evaluator(lambda *args: None)
        for position, owner in enumerate(["Jordan Lee", "Sam Ortiz", "Priya Nair", "Alex Kim"], 1):
            self.assertEqual(evaluator.ev(parse(formula), Ctx(lambda name, owner=owner: owner)), position)
        self.assertEqual(evaluator.ev(parse(formula), Ctx(lambda name: "A New Rep")), 5)   # unlisted rows last

    def test_the_order_column_is_not_a_plan_column_and_changes_no_value(self):
        self.assertNotIn("Source Order", [c["name"] for f in self.plan["frames"] for c in f["columns"]])
        model = trm.compute(self.plan)
        self.assertEqual([r["Owner"] for r in model.frame_rows(self.frame["id"])],
                         ["Jordan Lee", "Sam Ortiz", "Priya Nair", "Alex Kim"])

    # ---- 4: header style + title
    def test_source_header_style_reaches_the_table_header(self):
        self.assertEqual(self.frame["headerStyle"], {"fill": "#DDE7F5", "bold": True})
        for element_id in ("f_s_h2_fy26_weighted_forecast", "f_s_open_deals_by_stage"):
            self.assertEqual(element(self.spec, element_id)["tableStyle"],
                             {"textStyles": {"header": {"backgroundColor": "#DDE7F5", "fontWeight": "bold"}}})

    def test_the_header_band_carries_the_source_title_not_the_plan_name(self):
        self.assertEqual(self.pres["title"], "H2 FY26 Weighted Forecast")
        header, band = element(self.spec, "trHeader"), element(self.spec, "trHeaderBand")
        self.assertIn("H2 FY26 Weighted Forecast", header["body"])
        self.assertIn("04 Sales Pipeline Forecast.xlsx", header["body"])
        self.assertNotIn(self.plan["name"] + "\n", header["body"])
        self.assertEqual(band["kind"], "container")        # a text element has no style: the band is a container
        self.assertNotIn("style", header)
        self.assertEqual(band["style"], {"backgroundColor": self.pres["accent"]["tint"], "borderRadius": "round"})
        tree = {n[0]: n for n in layout_tree(self.spec)}
        self.assertEqual([c[0] for c in tree["trHeaderBand"][5]], ["trHeader"])

    def test_the_accent_is_the_palette_hue_nearest_the_source_header_fill(self):
        self.assertEqual(self.pres["accent"]["source"], "#DDE7F5")
        self.assertEqual(self.pres["accent"]["hex"], ps.ROLES["primary"])
        self.assertEqual(ps.accent_from_fill("#FDE9D9")["hue"], "amber")
        self.assertEqual(ps.accent_from_fill("#E2EFDA")["hue"], "green")
        self.assertEqual(ps.accent_from_fill("#F2F2F2")["hue"], "blue")      # grey: no hue to match
        self.assertEqual(ps.accent_from_fill(None)["hue"], "blue")

    # ---- 5: charts
    def test_both_source_charts_are_mapped(self):
        self.assertEqual(self.plan["unmappedCharts"], [])
        self.assertEqual([(c["kind"], c["stacking"], c["frame"]) for c in self.pres["charts"]],
                         [("bar", "stacked", self.frame["id"]), ("donut", "none", self.stages["id"])])
        bar, donut = self.pres["charts"]
        self.assertEqual(bar["title"], "Weighted Pipeline by Owner & Month")
        self.assertEqual([s["title"] for s in bar["series"]],
                         ["Jul-26", "Aug-26", "Sep-26", "Oct-26", "Nov-26", "Dec-26"])
        self.assertEqual(bar["categories"], ["Jordan Lee", "Sam Ortiz", "Priya Nair", "Alex Kim"])
        self.assertEqual(donut["categories"], ["Prospect", "Discovery", "Proposal", "Negotiation",
                                               "Closed Won", "Closed Lost"])

    def test_chart_elements_read_the_frame_element_at_its_grouping(self):
        bar = element(self.spec, "ch_weighted_pipeline_by_owner_month")
        self.assertEqual(bar["kind"], "bar-chart")
        self.assertEqual(bar["source"], {"kind": "table", "elementId": "f_s_h2_fy26_weighted_forecast",
                                         "groupingId": "f_s_h2_fy26_weighted_forecast_by"})
        self.assertEqual(bar["stacking"], "stacked")
        names = {c["id"]: c for c in bar["columns"]}
        self.assertEqual([names[i]["name"] for i in bar["yAxis"]["columnIds"]],
                         ["Jul-26", "Aug-26", "Sep-26", "Oct-26", "Nov-26", "Dec-26"])
        self.assertEqual(names[bar["yAxis"]["columnIds"][0]]["formula"],
                         "Sum([H2 FY26 Weighted Forecast/Jul-26])")
        order = names[bar["xAxis"]["sort"]["by"]]
        self.assertEqual((order["name"], order["hidden"], bar["xAxis"]["sort"]["aggregation"]),
                         ("Source Order", True, "min"))
        donut = element(self.spec, "ch_pipeline_by_stage")
        self.assertEqual(donut["kind"], "donut-chart")
        self.assertEqual(donut["source"]["groupingId"], "f_s_open_deals_by_stage_by")
        self.assertEqual(donut["value"]["columnId"], "ch_pipeline_by_stage_y1")
        self.assertEqual(donut["color"]["sort"]["direction"], "ascending")

    def test_chart_colours_come_from_the_workbook_theme_palette(self):
        scheme = self.spec["contents"]["settings"]["theme"]["overrides"]["categoricalScheme"]
        self.assertEqual(scheme[0], self.pres["accent"]["hex"])
        self.assertEqual(scheme[-1], ps.NEGATIVE)                         # red only once the hues run out
        self.assertEqual(len(set(scheme)), len(scheme))

    def test_chart_points_are_verified_against_the_source_cells(self):
        model = trm.compute(self.plan)
        lines, problems = pres.check_charts(self.plan, model)
        self.assertEqual(problems, [])
        self.assertEqual(lines, ["CHART PASS Weighted Pipeline by Owner & Month: 24/24 points equal verified "
                                 "s_h2_fy26_weighted_forecast cells",
                                 "CHART PASS Pipeline $ by Stage: 6/6 points equal verified "
                                 "s_open_deals_by_stage cells"])
        broken = copy.deepcopy(self.plan)
        broken["presentation"]["charts"][1]["series"][0]["values"][2] += 1
        _, problems = pres.check_charts(broken, model)
        self.assertEqual(len(problems), 1)
        self.assertIn("Forecast!C14", problems[0])

    # ---- 6: KPIs, layout
    def test_kpi_row_is_proposed_from_the_main_summary(self):
        kpis = self.pres["kpis"]
        self.assertEqual([(k["label"], k["kind"]) for k in kpis],
                         [("Total Weighted", "total"), ("Commit", "total"), ("Quota", "total"),
                          ("Coverage", "ratio")])
        total = kpis[0]
        self.assertEqual(total["formula"], "Sum([H2 FY26 Weighted Forecast/Total Weighted])")
        self.assertEqual(total["format"], {"kind": "number", "formatString": "(,.0f"})
        self.assertEqual(total["evidence"]["cells"], ["Forecast!H4", "Forecast!H5", "Forecast!H6", "Forecast!H7"])
        ratio = kpis[3]
        self.assertEqual((ratio["num"], ratio["den"]), ("Total Weighted", "Quota"))
        self.assertIn("Sum([H2 FY26 Weighted Forecast/Total Weighted]) / Sum([H2 FY26 Weighted Forecast/Quota])",
                      ratio["formula"])
        self.assertEqual(ratio["format"]["formatString"], ".1%")

    def test_kpi_values_equal_the_frame_totals(self):
        model = trm.compute(self.plan)
        lines, problems = pres.check_kpis(self.plan, model)
        self.assertEqual(problems, [])
        self.assertEqual(len(lines), 4)
        rows = model.frame_rows(self.frame["id"])
        by_label = {k["label"]: k["expected"] for k in self.pres["kpis"]}
        self.assertAlmostEqual(by_label["Total Weighted"], sum(r["Total Weighted"] for r in rows))
        self.assertAlmostEqual(by_label["Commit"], sum(r["Commit"] for r in rows))
        self.assertAlmostEqual(by_label["Quota"], 3300000)
        # the ratio is over the totals, not the sum (or mean) of the row ratios
        self.assertAlmostEqual(by_label["Coverage"], 3265400 / 3300000)
        self.assertNotAlmostEqual(by_label["Coverage"], sum(r["Coverage"] for r in rows))
        broken = copy.deepcopy(self.plan)
        broken["presentation"]["kpis"][1]["expected"] += 1
        self.assertEqual(len(pres.check_kpis(broken, model)[1]), 1)

    def test_kpi_elements_are_peer_cards_over_the_grouped_frame(self):
        cards = [element(self.spec, k["id"]) for k in self.pres["kpis"]]
        self.assertEqual([c["kind"] for c in cards], ["kpi-chart"] * 4)
        self.assertEqual(cards[0]["name"], {"text": "TOTAL WEIGHTED", "color": self.pres["accent"]["hex"]})
        self.assertEqual(cards[1]["name"]["color"], ps.ROLES["muted"])       # one accent, not sprayed
        self.assertEqual(cards[0]["source"]["groupingId"], "f_s_h2_fy26_weighted_forecast_by")
        self.assertEqual(cards[0]["style"], ps.card_style())
        spans = {p[0]: (p[2] - p[1], p[4] - p[3]) for p in placements(self.spec)}
        self.assertEqual({spans[c["id"]] for c in cards}, {(6, ps.KPI_ROWS)})   # equal peers

    def test_report_layout_references_every_element_once_without_overlap(self):
        report = placements(self.spec)
        ids = [p[0] for p in report]
        self.assertEqual(len(ids), len(set(ids)))
        known = {e["id"] for e in self.spec["contents"]["elements"]}
        self.assertTrue(set(ids) <= known)
        for a in report:
            self.assertTrue(1 <= a[1] < a[2] <= 25 and 1 <= a[3] < a[4], a)
            for b in report:
                if a is not b:
                    self.assertFalse(a[1] < b[2] and b[1] < a[2] and a[3] < b[4] and b[3] < a[4], (a, b))
        every = [i for page in ("report", "inputs", "model") for i in layout_ids(self.spec, page)]
        self.assertEqual(sorted(every), sorted(known))                 # every element exactly once

    def test_hero_proportion_and_tables_sized_to_their_rows(self):
        spans = {p[0]: p[1:] for p in placements(self.spec)}
        bar, donut = spans["ch_weighted_pipeline_by_owner_month"], spans["ch_pipeline_by_stage"]
        self.assertEqual((bar[0], bar[1], donut[0], donut[1]), (1, 17, 17, 25))        # 2/3 : 1/3, not 50/50
        self.assertEqual(bar[2:], donut[2:])
        summary = spans["f_s_h2_fy26_weighted_forecast"]
        stages = spans["f_s_open_deals_by_stage"]
        self.assertEqual((summary[0], summary[1]), (1, 25))
        self.assertEqual((stages[0], stages[1]), (1, 17))                              # 3 columns: not stretched
        self.assertLessEqual(summary[3] - summary[2], 12)        # 4 rows, not the old 20-row card
        self.assertLessEqual(stages[3] - stages[2], 14)          # 6 rows
        self.assertGreater(stages[3] - stages[2], summary[3] - summary[2])
        order = [p[0] for p in placements(self.spec)]
        self.assertEqual(order[:5], ["trHeaderBand", "kpi_total_weighted", "kpi_commit", "kpi_quota",
                                     "kpi_coverage"])
        self.assertLess(order.index("trDivider"), order.index("trDetail"))
        self.assertLess(order.index("trDetail"), order.index("f_s_h2_fy26_weighted_forecast"))

    def test_hidden_model_and_input_pages_are_unchanged(self):
        pages = {p["id"]: p for p in self.spec["contents"]["pages"]}
        self.assertEqual(pages["model"]["visibility"], "hidden")
        self.assertEqual([p[0] for p in placements(self.spec, "inputs")],
                         ["in_tblpipeline", "in_tblstages", "in_h2_fy26_weighted_forecast_inputs"])

    def test_build_cli_prints_the_presentation_parity_and_fails_when_a_point_differs(self):
        with tempfile.TemporaryDirectory() as tmp:
            plan_path = Path(tmp) / "plan.json"
            code, out, err = run_cli(DRAFT_CLI, [str(PIPELINE), "--connection", "c", "--folder", "f",
                                                 "--out", str(plan_path)])
            self.assertEqual(code, 0, err)
            self.assertIn("source chart mapped to a verified frame: Pipeline $ by Stage (donut)", out)
            self.assertIn("KPI cards proposed: Total Weighted, Commit, Quota, Coverage", out)
            build = Path(tmp) / "build"
            code, out, err = run_cli(BUILD_CLI, [str(plan_path), "--source", str(PIPELINE),
                                                 "--out-dir", str(build)])
            self.assertEqual(code, 0, err)
            self.assertIn("KPI PASS Total Weighted", out)
            self.assertIn("CHART PASS Pipeline $ by Stage: 6/6", out)
            plan = json.loads(plan_path.read_text())
            plan["presentation"]["kpis"][0]["expected"] += 5
            plan_path.write_text(json.dumps(plan))
            (build / "wb_spec.json").unlink()
            code, out, err = run_cli(BUILD_CLI, [str(plan_path), "--source", str(PIPELINE),
                                                 "--out-dir", str(build)])
            self.assertEqual(code, 1)
            self.assertIn("PRESENTATION PARITY FAILED", err)
            self.assertFalse((build / "wb_spec.json").exists())


# --------------------------------------------------------------------------- the four fixtures
class TestBroaderEightLayouts(unittest.TestCase):
    def test_every_fixture_lays_out_cleanly_and_keeps_its_values(self):
        for number, path in SOURCES.items():
            with self.subTest(fixture=number):
                plan, info = draft(path)
                spec = wb_spec(plan)
                report = placements(spec)
                ids = [p[0] for p in report]
                self.assertEqual(len(ids), len(set(ids)))
                self.assertEqual(ids[0], "trHeaderBand")
                for a in report:
                    for b in report:
                        if a is not b:
                            self.assertFalse(a[1] < b[2] and b[1] < a[2] and a[3] < b[4] and b[3] < a[4], (a, b))
                known = {e["id"] for e in spec["contents"]["elements"]}
                every = [i for page in ("report", "inputs", "model") for i in layout_ids(spec, page)]
                self.assertEqual(sorted(every), sorted(known))
                model = trm.compute(plan)
                self.assertEqual(trm.check_cells(plan, model)[1], [])
                self.assertEqual(pres.check_presentation(plan, model)[1], [])
                for _, column in number_columns(plan):
                    self.assertIsInstance(column.get("format"), dict)

    def test_the_other_fixtures_map_their_charts_and_propose_kpis_where_the_source_has_totals(self):
        plan07, _ = draft(SOURCES["07"])
        self.assertEqual([k["label"] for k in plan07["presentation"]["kpis"]],
                         ["FY26 Plan", "FY26 Budget", "Variance", "Var %"])
        self.assertEqual([(c["kind"], c["stacking"]) for c in plan07["presentation"]["charts"]],
                         [("bar", "none")])                                       # clustered
        plan06, _ = draft(SOURCES["06"])
        self.assertEqual([c["kind"] for c in plan06["presentation"]["charts"]], ["line"])
        self.assertEqual(plan06["presentation"]["kpis"], [])
        plan08, _ = draft(SOURCES["08"])
        self.assertEqual(plan08["presentation"]["kpis"], [])       # min / text / conditional columns: none additive
        self.assertEqual(plan08["unmappedCharts"], [])

    def test_the_summary_leads_the_page_over_the_expansion_pivot(self):
        plan, _ = draft(SOURCES["07"])
        report = [p[0] for p in placements(wb_spec(plan))]
        self.assertLess(report.index("f_s_plan_vs_budget_by_department"), report.index("pv_monthly_cost"))

    def test_keyed_pivot_rows_keep_the_grid_order(self):
        plan, _ = draft(SOURCES["08"])
        spec = wb_spec(plan)
        pivot = element(spec, "pv_projection")
        order = next(c for c in pivot["columns"] if c["name"] == "Source Order")
        self.assertEqual(pivot["rowsBy"][0]["sort"], {"direction": "ascending", "by": order["id"],
                                                      "aggregation": "min"})
        keyed = element(spec, "f_k_projection")
        base = next(c for c in keyed["columns"] if c["name"] == "Source Order")
        self.assertEqual([label for label, _ in order_positions(base["formula"])],
                         ["SKU-200", "SKU-300", "SKU-400", "SKU-500", "SKU-600", "SKU-700"])


# --------------------------------------------------------------------------- synthetic workbooks
def stacked_columns(ws):
    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
    chart.title = "Revenue by region and month"
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=3, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=6))
    ws.add_chart(chart, "K2")


def clustered_bars(ws):
    chart = BarChart()
    chart.type, chart.grouping = "bar", "clustered"
    chart.title = "Attainment"
    chart.add_data(Reference(ws, min_col=8, min_row=3, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=6))
    ws.add_chart(chart, "K20")


def doughnut(ws):
    chart = DoughnutChart()
    chart.title = "Revenue share"
    chart.add_data(Reference(ws, min_col=4, min_row=3, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=6))
    ws.add_chart(chart, "K38")


def unmappable(ws):
    subset = BarChart()                                       # two of the three regions only
    subset.title = "Two regions"
    subset.add_data(Reference(ws, min_col=4, min_row=3, max_row=5), titles_from_data=True)
    subset.set_categories(Reference(ws, min_col=1, min_row=4, max_row=5))
    ws.add_chart(subset, "T2")
    scatter = ScatterChart()                                  # a chart kind with no mapping
    scatter.title = "Scatter"
    scatter.series.append(Series(Reference(ws, min_col=4, min_row=4, max_row=6),
                                 Reference(ws, min_col=2, min_row=4, max_row=6)))
    ws.add_chart(scatter, "T20")
    data = ws.parent["Orders"]                                # values that are not a verified frame cell
    raw = BarChart()
    raw.title = "Raw quantities"
    raw.add_data(Reference(data, min_col=5, min_row=1, max_row=9), titles_from_data=True)
    raw.set_categories(Reference(data, min_col=1, min_row=2, max_row=9))
    ws.add_chart(raw, "T38")


class SyntheticCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name).resolve()

    def drafted(self, name="orders", **options):
        path = self.dir / f"{name}.xlsx"
        fx.build_orders(path, **options)
        plan, info = draft(path)
        return path, plan, info


class TestFormatsFromTheSource(SyntheticCase):
    FORMATS = {"Summary": {"B4:C6": "$#,##0", "D4:D6": "#,##0;[Red](#,##0)", "H4:H6": "0.0%",
                           "F4:F6": "# ?/?"},
               "Orders": {"G2:G9": "0%", "H2:H9": "$#,##0.00"}}

    def test_cell_formats_reach_frame_and_table_columns(self):
        _, plan, info = self.drafted(formats=self.FORMATS)
        frame = next(f for f in plan["frames"] if f["kind"] == "summary")
        formats = {c["name"]: c for c in frame["columns"]}
        self.assertEqual(formats["Jan"]["format"], {"kind": "number", "formatString": "$,.0f"})
        self.assertEqual(formats["Total"]["format"]["formatString"], "(,.0f")
        self.assertTrue(formats["Total"]["negativeRed"])
        self.assertEqual(formats["Attain"]["format"]["formatString"], ".1%")
        self.assertEqual(formats["Big"]["format"]["formatString"], ",.0f")          # General count

    def test_a_percent_operand_does_not_leak_into_the_money_column(self):
        _, plan, _ = self.drafted(formats=self.FORMATS)
        spec = wb_spec(plan)
        wrapper = element(spec, "w_tblorders")
        columns = {c["name"]: c for c in wrapper["columns"]}
        self.assertEqual(columns["Rate"]["format"]["formatString"], ".0%")
        self.assertIn("[Rate]", columns["Revenue"]["formula"])
        self.assertEqual(columns["Revenue"]["format"], {"kind": "number", "formatString": "$,.2f"})
        self.assertEqual(columns["Qty"]["format"]["formatString"], ",.0f")

    def test_an_untranslatable_code_is_listed_and_the_column_gets_a_plain_number(self):
        _, plan, info = self.drafted(formats=self.FORMATS)
        issue = next(i for i in plan["draftReviewIssues"] if i["code"] == "NUMBER_FORMAT_APPROXIMATED")
        self.assertEqual(issue["severity"], "review")
        self.assertEqual(len(issue["evidence"]), 1)
        self.assertIn("# ?/?", issue["evidence"][0])
        self.assertIn("fraction formats are not supported", issue["evidence"][0])
        self.assertIn("Summary!F4", issue["evidence"][0])
        frame = next(f for f in plan["frames"] if f["kind"] == "summary")
        avg = next(c for c in frame["columns"] if c["name"] == "Avg Rev")
        self.assertEqual(avg["format"]["kind"], "number")
        self.assertIn("NUMBER_FORMAT_APPROXIMATED", info["reviewIssues"])

    def test_exactly_translated_formats_raise_no_issue(self):
        _, plan, info = self.drafted(formats={"Summary": {"B4:C6": "$#,##0", "H4:H6": "0.0%"}})
        self.assertNotIn("NUMBER_FORMAT_APPROXIMATED", info["reviewIssues"])


class TestRowOrderAndHeaders(SyntheticCase):
    def test_a_grid_that_is_not_alphabetical_keeps_its_row_order(self):
        _, plan, _ = self.drafted()
        spec = wb_spec(plan)
        summary = element(spec, "f_s_revenue_by_region")
        order = next(c for c in summary["columns"] if c["name"] == "Source Order")
        self.assertEqual(order_positions(order["formula"]), [("East", 1), ("West", 2), ("North", 3)])
        self.assertEqual(summary["sort"][0]["columnId"], order["id"])

    def test_header_fill_bold_and_the_title_cell_come_from_the_source(self):
        _, plan, _ = self.drafted(header_style=True)
        frame = next(f for f in plan["frames"] if f["kind"] == "summary")
        self.assertEqual(frame["headerStyle"], {"fill": "#DDE7F5", "bold": True})
        self.assertEqual(plan["presentation"]["title"], "Revenue by region")
        spec = wb_spec(plan)
        self.assertEqual(element(spec, "f_s_revenue_by_region")["tableStyle"]["textStyles"]["header"],
                         {"backgroundColor": "#DDE7F5", "fontWeight": "bold"})
        self.assertIn("Revenue by region", element(spec, "trHeader")["body"])

    def test_no_source_style_means_no_table_style(self):
        _, plan, _ = self.drafted()
        spec = wb_spec(plan)
        self.assertNotIn("tableStyle", element(spec, "f_s_revenue_by_region"))

    def test_pivot_headers_use_the_column_header_style(self):
        style = ps.table_style({"fill": "#DDE7F5", "bold": True, "color": "#1F3864"}, "pivot-table")
        self.assertEqual(style["textStyles"]["columnHeader"],
                         {"backgroundColor": "#DDE7F5", "color": "#1F3864", "fontWeight": "bold"})
        self.assertEqual(style["textStyles"]["header"], style["textStyles"]["columnHeader"])
        self.assertIsNone(ps.table_style(None))
        self.assertEqual(set(ps.table_style({"fill": "#DDE7F5"})["textStyles"]), {"header"})


class TestChartMapping(SyntheticCase):
    def charts(self, *builders):
        def add(ws):
            for builder in builders:
                builder(ws)
        return self.drafted(charts=add)

    def test_stacked_column_clustered_bar_and_doughnut_map_with_their_modes(self):
        _, plan, info = self.charts(stacked_columns, clustered_bars, doughnut)
        by_title = {c["title"]: c for c in plan["presentation"]["charts"]}
        self.assertEqual(set(by_title), {"Revenue by region and month", "Attainment", "Revenue share"})
        stacked = by_title["Revenue by region and month"]
        self.assertEqual((stacked["kind"], stacked["stacking"], stacked["orientation"]),
                         ("bar", "stacked", "vertical"))
        self.assertEqual([s["title"] for s in stacked["series"]], ["Jan", "Feb"])
        self.assertEqual(stacked["categories"], ["East", "West", "North"])         # the sheet's order
        bars = by_title["Attainment"]
        self.assertEqual((bars["kind"], bars["stacking"], bars["orientation"]), ("bar", "none", "horizontal"))
        self.assertEqual(by_title["Revenue share"]["kind"], "donut")
        self.assertEqual(info["unmappedCharts"], [])
        spec = wb_spec(plan)
        horizontal = element(spec, "ch_attainment")
        self.assertEqual((horizontal["orientation"], horizontal["stacking"]), ("horizontal", "none"))
        self.assertEqual(element(spec, "ch_revenue_by_region_and_month")["stacking"], "stacked")
        self.assertEqual(element(spec, "ch_revenue_share")["kind"], "donut-chart")

    def test_a_single_series_chart_takes_the_accent_and_hides_its_legend(self):
        _, plan, _ = self.charts(clustered_bars)
        chart = element(wb_spec(plan), "ch_attainment")
        self.assertEqual(chart["color"], {"by": "single", "value": plan["presentation"]["accent"]["hex"]})
        self.assertEqual(chart["legend"], {"visibility": "hidden"})

    def test_every_mapped_point_equals_a_verified_frame_cell(self):
        _, plan, _ = self.charts(stacked_columns, doughnut)
        model = trm.compute(plan)
        lines, problems = pres.check_charts(plan, model)
        self.assertEqual(problems, [])
        self.assertEqual(len(lines), 2)
        checked = {cell[0] for group in plan["checks"] for cell in group["cells"]}
        for chart in plan["presentation"]["charts"]:
            for series in chart["series"]:
                self.assertTrue(set(series["cells"]) <= checked)
        tampered = copy.deepcopy(plan)
        tampered["presentation"]["charts"][0]["series"][1]["values"][0] += 0.5
        self.assertEqual(len(pres.check_charts(tampered, model)[1]), 1)
        shifted = copy.deepcopy(plan)
        shifted["tables"][0]["rows"][0][4] += 1                 # data moves: the oracle's point no longer matches
        self.assertTrue(pres.check_charts(shifted, trm.compute(shifted))[1])

    def test_unmappable_charts_stay_listed_with_the_reason(self):
        _, plan, info = self.charts(stacked_columns, unmappable)
        self.assertEqual([c["title"] for c in plan["presentation"]["charts"]], ["Revenue by region and month"])
        listed = {t: msg for t in ("Two regions", "Scatter", "Raw quantities")
                  for msg in plan["unmappedCharts"] if msg.startswith(t)}
        self.assertEqual(len(listed), 3)
        self.assertIn("not all of the grid's rows", listed["Two regions"])
        self.assertIn("ScatterChart is not a mapped chart kind", listed["Scatter"])
        self.assertIn("not row labels or verified period cells", listed["Raw quantities"])
        self.assertIn("rebuild from the presentation plan", listed["Scatter"])
        self.assertEqual(len(info["unmappedCharts"]), 3)

    def test_chart_ids_are_unique_and_every_chart_reads_an_existing_column(self):
        _, plan, _ = self.charts(stacked_columns, clustered_bars, doughnut)
        spec = wb_spec(plan)
        elements = {e["id"]: e for e in spec["contents"]["elements"]}
        for chart in plan["presentation"]["charts"]:
            el = elements[chart["id"]]
            source = elements[el["source"]["elementId"]]
            self.assertIn(el["source"]["groupingId"], {g["id"] for g in source["groupings"]})
            for column in el["columns"]:
                for ref, col in re.findall(r"\[([^\[\]/]+)/([^\[\]]+)\]", column["formula"]):
                    self.assertEqual(ref, source["name"])
                    self.assertIn(col, {c["name"] for c in source["columns"]})


class TestKpiProposals(SyntheticCase):
    def test_totals_and_a_ratio_of_totals_are_proposed(self):
        _, plan, _ = self.drafted()
        kpis = plan["presentation"]["kpis"]
        self.assertEqual([(k["label"], k["kind"]) for k in kpis],
                         [("Total", "total"), ("Target", "total"), ("Attain", "ratio")])
        model = trm.compute(plan)
        rows = model.frame_rows(next(f for f in plan["frames"] if f["kind"] == "summary")["id"])
        self.assertAlmostEqual(kpis[0]["expected"], sum(r["Total"] for r in rows))
        self.assertAlmostEqual(kpis[2]["expected"], sum(r["Total"] for r in rows) / 700.0)
        self.assertEqual(pres.check_kpis(plan, model)[1], [])

    def test_a_ratio_of_totals_differs_from_the_mean_of_ratios_and_a_zero_denominator_is_zero(self):
        _, plan, _ = self.drafted()
        ratio = plan["presentation"]["kpis"][2]
        self.assertIn("If(IsNull(Sum(", ratio["formula"])
        self.assertIn("= 0, 0,", ratio["formula"])
        self.assertIn("not the sum of the row ratios", ratio["reason"])

    def test_period_columns_averages_and_counts_are_not_kpis(self):
        _, plan, _ = self.drafted()
        labels = [k["label"] for k in plan["presentation"]["kpis"]]
        for skipped in ("Jan", "Feb", "Big", "Avg Rev"):
            self.assertNotIn(skipped, labels)

    def test_a_total_row_in_the_source_is_evidence_and_is_checked(self):
        _, plan, _ = self.drafted(totals="sum")
        by_label = {k["label"]: k for k in plan["presentation"]["kpis"]}
        self.assertEqual(by_label["Total"]["evidence"]["totalCell"], "Summary!D7")
        model = trm.compute(plan)
        self.assertEqual(pres.check_kpis(plan, model)[1], [])
        broken = copy.deepcopy(plan)
        for group in broken["checks"]:
            if group["kind"] == "summary-total" and group["column"] == "Total":
                group["cells"][0][2] += 10
        self.assertEqual(len(pres.check_kpis(broken, model)[1]), 1)

    def test_fewer_than_two_kpis_means_no_kpi_row(self):
        plan = {"frames": [], "tables": [], "checks": []}
        self.assertEqual(pres.propose_kpis(plan), [])
        _, plan, _ = self.drafted()
        trimmed = copy.deepcopy(plan)
        trimmed["frames"] = [f for f in trimmed["frames"]]
        frame = next(f for f in trimmed["frames"] if f["kind"] == "summary")
        frame["columns"] = [c for c in frame["columns"] if c["name"] in ("Jan", "Feb", "Total")]
        self.assertEqual(pres.propose_kpis(trimmed), [])

    def test_the_kpi_row_never_exceeds_four_cards(self):
        _, plan, _ = self.drafted()
        self.assertLessEqual(len(plan["presentation"]["kpis"]), pres.MAX_KPIS)


class TestPageComposition(SyntheticCase):
    def test_page_without_charts_or_kpis_is_just_the_header_and_the_tables(self):
        _, plan, _ = self.drafted()
        plan = copy.deepcopy(plan)
        plan["presentation"]["kpis"], plan["presentation"]["charts"] = [], []
        report = [p[0] for p in placements(wb_spec(plan))]
        self.assertEqual(report[0], "trHeaderBand")
        self.assertNotIn("trDivider", report)
        self.assertNotIn("settings", wb_spec(plan)["contents"])

    def test_an_old_plan_without_a_presentation_block_still_builds(self):
        _, plan, _ = self.drafted()
        plan = copy.deepcopy(plan)
        del plan["presentation"]
        spec = wb_spec(plan)
        self.assertIn(plan["name"], element(spec, "trHeader")["body"])
        legacy = copy.deepcopy(plan)
        legacy["frames"][0]["columns"][0]["format"] = "percent"        # the pre-presentation plan shape
        self.assertEqual(trm.number_format(legacy, "percent"), {"kind": "number", "formatString": ".1%"})
        wb_spec(legacy)

    def test_table_height_follows_the_row_count_within_bounds(self):
        heights = [ps.table_height(n) for n in (1, 4, 6, 13, 30, 500)]
        self.assertEqual(heights, sorted(heights))
        self.assertEqual(heights[0], ps.MIN_TABLE_ROWS)
        self.assertEqual(heights[-1], ps.MAX_TABLE_ROWS)
        self.assertLess(ps.table_height(4), 14)

    def test_surfaces_can_be_switched_off_and_the_builder_simply_omits_the_field(self):
        self.addCleanup(ps.SURFACES.update, dict(ps.SURFACES))
        ps.SURFACES.update(table_text_styles=False, conditional_single=False, grouping_source=False)
        _, plan, _ = self.drafted(header_style=True, formats={"Summary": {"D4:D6": "#,##0;[Red](#,##0)"}},
                                  charts=stacked_columns)
        spec = wb_spec(plan)
        summary = element(spec, "f_s_revenue_by_region")
        self.assertNotIn("tableStyle", summary)
        self.assertNotIn("conditionalFormats", summary)
        self.assertNotIn("groupingId", element(spec, "ch_revenue_by_region_and_month")["source"])

    def test_presentation_validation_rejects_a_chart_with_a_short_series(self):
        _, plan, _ = self.drafted(charts=stacked_columns)
        broken = copy.deepcopy(plan)
        broken["presentation"]["charts"][0]["series"][0]["values"].pop()
        with self.assertRaises(trm.PlanError):
            trm.validate_plan(broken)
        broken = copy.deepcopy(plan)
        broken["presentation"]["kpis"][0]["frame"] = "nope"
        with self.assertRaises(trm.PlanError):
            trm.validate_plan(broken)


if __name__ == "__main__":
    unittest.main()
