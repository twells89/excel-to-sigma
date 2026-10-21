"""Real source/input planning plus synthetic protocol receipts (no target calls)."""
import ast
import copy
import datetime
import json
import tempfile
from openpyxl import Workbook, load_workbook
import hashlib
import unittest
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts/lib'))
from source_inputs import build_source_input_plan
import test_native_calculations
from test_regional_canary_inputs import SDK_FORMULA, check_formula_api_contract
from native_exact_plan import (build_exact_native_calculation_plan,lower_exact_native_calculations,
    prepare_exact_native_scaffold,confirm_exact_native_scaffold,prepare_exact_native_update,confirm_exact_native_readback,
    _formula_result,input_scalar)
from native_exact import compile_expression
from native_formula_reference import evaluate_native
from native_calculations import build_native_calculation_backup
from migration_manifest import build_manifest
from input_sync import build_artifact_records

class ExactPlanningTests(unittest.TestCase):
    source=test_native_calculations.NativeCalculationTests.source
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
    def test_complete_plan_error_class_and_bounded_dependency_graph(self):
        path,inputs=self.source();plan=build_exact_native_calculation_plan(path,inputs)
        self.assertEqual(sum(r['status']=='ready-to-bind'for r in plan['formulas'].values()),4)
        self.assertEqual(plan['formulas']['Summary Renamed!A2']['localExpected'],99)
        self.assertEqual(plan['formulas']['Ledger Renamed!D2']['localExpected'],9)
        self.assertEqual(plan['formulas']['Ledger Renamed!D3']['localExpected'],5)
        self.assertEqual(plan['formulas']['Summary Renamed!A1']['localExpected'],14)
        assigned={'ledger':{'elementId':'observed-input','name':'tblLedger','rowIdColumn':'Row ID','rows':{'0':'observed-a','1':'observed-b'}}}
        lowered=lower_exact_native_calculations(plan,assigned,max_columns=8)
        self.assertEqual(len(lowered['formulas']),4)
        for shard in lowered['shards']:
            self.assertLessEqual(len(shard['columns']),8)
            for col in shard['columns']:
                self.assertLessEqual(len(col['formula']),6000)
                self.assertNotIn('__native_',col['formula'])
        self.assertEqual(lowered['runtimeStatus'],'UNVERIFIED_CANDIDATE')
        self.assertEqual(set(lowered['formalTemplateObligations']),{'Ledger Renamed!D2','Ledger Renamed!D3'})
        with self.assertRaisesRegex(ValueError,'row ID'):lower_exact_native_calculations(plan,{})
        plan['formulas']['Summary Renamed!A2']['localExpected']=100
        with self.assertRaisesRegex(ValueError,'hash'):lower_exact_native_calculations(plan,assigned)
    def test_formal_this_row_is_scalar_and_appended_rows_get_symbolic_formulas(self):
        path,old=self.source();book=load_workbook(path)
        for row in (2,3):book['Ledger Renamed'][f'D{row}']='=tblLedger[[#This Row],[Amount]]*tblLedger[[#This Row],[Rate]]'
        book.save(path)
        decisions=copy.deepcopy(old['decisions']);decisions['sourceSha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        inputs=build_source_input_plan(path,decisions,connection_id='connection',folder_id='folder')
        p=build_exact_native_calculation_plan(path,inputs)
        self.assertEqual(p['formulas']['Ledger Renamed!D2']['localExpected'],6)
        self.assertEqual(p['formulas']['Ledger Renamed!D3']['localExpected'],4)
        assigned={'ledger':{'elementId':'actual-input','name':'tblLedger','rowIdColumn':'Row ID','rows':{'0':'actual-a','1':'actual-b'}}}
        lower=lower_exact_native_calculations(p,assigned)
        domains=[e['calculatedColumnDomain']for e in lower['shards']if 'calculatedColumnDomain'in e]
        self.assertTrue(domains)
        self.assertEqual(domains[0]['appendPolicy'],'live-formal-table')
        self.assertEqual(lower['formalTemplateObligations'],[])

    def test_formal_row_formula_result_converts_blank_without_changing_input_kind(self):
        path,old=self.source();book=load_workbook(path)
        book['Ledger Renamed']['B2']=None
        for row in (2,3):book['Ledger Renamed'][f'D{row}']=f'=B{row}'
        book.save(path)
        decisions=copy.deepcopy(old['decisions']);decisions['sourceSha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        decisions['features'][0]['reviewedBlankColumns']=['Amount']
        inputs=build_source_input_plan(path,decisions,connection_id='connection',folder_id='folder')
        plan=build_exact_native_calculation_plan(path,inputs)
        self.assertEqual(plan['formulas']['Ledger Renamed!D2']['localExpected'],0)
        self.assertEqual(plan['formulas']['Ledger Renamed!D3']['localExpected'],4)
        assigned={'ledger':{'elementId':'actual-input','name':'tblLedger','rowIdColumn':'Row ID','rows':{'0':'actual-a','1':'actual-b'}}}
        lower=lower_exact_native_calculations(plan,assigned)
        row_tables=[e for e in lower['shards'] if 'calculatedColumnDomain' in e]
        self.assertEqual(len(row_tables),1)
        self.assertEqual(row_tables[0]['calculatedColumnDomain']['appendPolicy'],'live-formal-table')
        kinds=[c['formula'] for c in row_tables[0]['columns'] if c['name']=='kind']
        self.assertEqual(len(kinds),1)
        self.assertIn('"n"',kinds[0])

    def test_formula_text_that_looks_like_internal_helper_is_literal(self):
        path,inputs=self.source();book=load_workbook(path);book['Summary Renamed']['A2']='="__native_h0__"';book.save(path)
        decisions=copy.deepcopy(inputs['decisions']);decisions['sourceSha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        inputs=build_source_input_plan(path,decisions,connection_id='connection',folder_id='folder')
        p=build_exact_native_calculation_plan(path,inputs)
        assigned={'ledger':{'elementId':'actual-input','name':'tblLedger','rowIdColumn':'Row ID','rows':{'0':'actual-a','1':'actual-b'}}}
        lower=lower_exact_native_calculations(p,assigned)
        self.assertEqual(lower['formulas']['Summary Renamed!A2']['lanes']['text'],'"__native_h0__"')

    def test_blank_boolean_inputs_remain_unproven_through_formula_references(self):
        path,old=self.source();book=load_workbook(path)
        for row in (2,3):
            book['Ledger Renamed'][f'C{row}']=None
            book['Ledger Renamed'][f'D{row}']=f'=C{row}'
        book['Summary Renamed']['A1']="=COUNTIFS('Ledger Renamed'!D2:D3,\">0\")"
        book.save(path)
        decisions=copy.deepcopy(old['decisions']);decisions['sourceSha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        decisions['features'][0]['columnTypes']['Rate']='boolean'
        decisions['features'][0]['reviewedBlankColumns']=['Rate']
        inputs=build_source_input_plan(path,decisions,connection_id='connection',folder_id='folder')
        native=build_exact_native_calculation_plan(path,inputs)
        record=native['formulas']['Summary Renamed!A1']
        self.assertEqual(record['status'],'unresolved')
        self.assertIn('Boolean criteria',record['reason'])

    def test_sumifs_with_materialized_criteria_helpers_still_sums_numeric_values(self):
        path,old=self.source();book=load_workbook(path)
        for row in (2,3):book['Ledger Renamed'][f'D{row}']=f'=B{row}*C{row}'
        book['Summary Renamed']['A1']='=SUMIFS(tblLedger[Result],tblLedger[Account],"a")'
        book.save(path)
        decisions=copy.deepcopy(old['decisions']);decisions['sourceSha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        inputs=build_source_input_plan(path,decisions,connection_id='connection',folder_id='folder')
        native=build_exact_native_calculation_plan(path,inputs)
        self.assertEqual(native['formulas']['Summary Renamed!A1']['localExpected'],6)
        assigned={'ledger':{'elementId':'actual-input','name':'tblLedger','rowIdColumn':'Row ID','rows':{'0':'actual-a','1':'actual-b'}}}
        lower=lower_exact_native_calculations(native,assigned)
        aggregate=[e for e in lower['shards'] if 'rangeDomain' in e]
        self.assertEqual(len(aggregate),1)
        self.assertTrue(any(' Criteria ' in e['name'] for e in lower['shards']))
        self.assertIn('native_value',{c['name'] for c in aggregate[0]['columns']})
        expressions=[c['formula'] for e in lower['shards'] for c in e['columns']]
        self.assertTrue(any('SumIf(' in f for f in expressions))
        self.assertFalse(any('CountIf(' in f for f in expressions))

    def test_source_excel_error_remains_disclosed_candidate_error_class(self):
        path=Path(self.temp.name)/'errors.xlsx';book=Workbook();book.active['A1']='=1/0';book.active['A2']='=IFERROR(A1,5)';book.save(path)
        p=build_exact_native_calculation_plan(path)
        self.assertEqual(p['formulas']['Sheet!A1']['sourceResultClass'],'excel-error')
        self.assertEqual(p['formulas']['Sheet!A1']['localExpectedError'],'#DIV/0!')
        self.assertEqual(p['formulas']['Sheet!A2']['localExpected'],5)

    def test_reviewed_numeric_blank_formula_result_is_zero_for_downstream_aggregates(self):
        path,old=self.source();book=load_workbook(path)
        book['Ledger Renamed']['B2']=None
        book['Ledger Renamed']['D2']='=B2'
        book['Ledger Renamed']['E1']='Note'
        book['Ledger Renamed']['E2']=None
        book['Ledger Renamed']['E3']='kept'
        book['Ledger Renamed'].tables['tblLedger'].ref='A1:E3'
        book['Summary Renamed']['A1']="=COUNT('Ledger Renamed'!D2)"
        book['Summary Renamed']['A2']="=MIN('Ledger Renamed'!D2,5)"
        book['Summary Renamed']['A3']='=IF(TRUE,"",0)'
        book['Summary Renamed']['A4']='=1/0'
        book['Summary Renamed']['A5']="=COUNT('Ledger Renamed'!B2)"
        book['Summary Renamed']['A6']="=MIN('Ledger Renamed'!B2,5)"
        book['Summary Renamed']['A7']="=IF(TRUE,'Ledger Renamed'!B2,\"\")"
        book['Summary Renamed']['A8']="=IF(FALSE,0,'Ledger Renamed'!B2)"
        book['Summary Renamed']['A9']="=IF(FALSE,'Ledger Renamed'!B2,\"\")"
        book['Summary Renamed']['A10']='=COUNT(A7)'
        book['Summary Renamed']['A11']='=COUNT(A9)'
        book['Summary Renamed']['A12']="='Ledger Renamed'!E2"
        book['Summary Renamed']['A13']='=COUNT(A12)'
        book['Summary Renamed']['A14']="=IF(TRUE,1/0,'Ledger Renamed'!B2)"
        book['Summary Renamed']['A15']="=IF(TRUE,TRUE,'Ledger Renamed'!B2)"
        book.save(path)
        decisions=copy.deepcopy(old['decisions']);decisions['sourceSha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        decisions['features'][0]['reviewedBlankColumns']=['Amount','Note']
        decisions['features'][0]['literalColumns'].append('Note')
        decisions['features'][0]['columnTypes']['Note']='text'
        inputs=build_source_input_plan(path,decisions,connection_id='connection',folder_id='folder')
        plan=build_exact_native_calculation_plan(path,inputs)
        for cell,expected,kind in [('Ledger Renamed!D2',0,'number'),
                                    ('Summary Renamed!A1',1,'number'),
                                    ('Summary Renamed!A2',0,'number'),
                                    ('Summary Renamed!A3','','text'),
                                    ('Summary Renamed!A5',0,'number'),
                                    ('Summary Renamed!A6',5,'number'),
                                    ('Summary Renamed!A7',0,'number'),
                                    ('Summary Renamed!A8',0,'number'),
                                    ('Summary Renamed!A9','','text'),
                                    ('Summary Renamed!A10',1,'number'),
                                    ('Summary Renamed!A11',0,'number'),
                                    ('Summary Renamed!A12',0,'number'),
                                    ('Summary Renamed!A13',1,'number'),
                                    ('Summary Renamed!A15',True,'logical')]:
            with self.subTest(cell=cell):
                record=plan['formulas'][cell]
                self.assertEqual(record['status'],'ready-to-bind')
                self.assertEqual(record['localExpected'],expected)
                self.assertEqual(record['sourceResultClass'],kind)
        self.assertEqual(plan['formulas']['Summary Renamed!A4']['sourceResultClass'],'excel-error')
        self.assertEqual(plan['formulas']['Summary Renamed!A14']['sourceResultClass'],'excel-error')
        self.assertEqual(plan['formulas']['Summary Renamed!A14']['localExpectedError'],'#DIV/0!')
        assigned={'ledger':{'elementId':'actual-input','name':'tblLedger','rowIdColumn':'Row ID','rows':{'0':'actual-a','1':'actual-b'}}}
        lowered=lower_exact_native_calculations(plan,assigned)
        direct=lowered['formulas']['Ledger Renamed!D2']['lanes']
        self.assertIn('"n"',direct['kind'])
        self.assertNotEqual(direct['kind'],'"z"')

    def test_blank_datetime_formula_emits_zero_number_lane_and_preserves_raw_blank(self):
        path,old=self.source();book=load_workbook(path)
        book['Ledger Renamed']['E1']='At'
        book['Ledger Renamed']['E2']=None
        book['Ledger Renamed']['E3']=datetime.datetime(2024,1,1)
        book['Ledger Renamed'].tables['tblLedger'].ref='A1:E3'
        book['Ledger Renamed']['D2']='=E2'
        book['Summary Renamed']['A1']="=COUNT('Ledger Renamed'!D2)"
        book['Summary Renamed']['A2']="=MIN('Ledger Renamed'!D2,5)"
        book['Summary Renamed']['A3']="=IF(TRUE,'Ledger Renamed'!E2,\"\")"
        book['Summary Renamed']['A4']='=COUNT(A3)'
        book['Summary Renamed']['A5']="=IF(FALSE,0,'Ledger Renamed'!E2)"
        book['Summary Renamed']['A6']="=COUNT('Ledger Renamed'!E2)"
        book['Summary Renamed']['A7']="=MIN('Ledger Renamed'!E2,5)"
        book['Summary Renamed']['A8']="=IF(TRUE,1/0,'Ledger Renamed'!E2)"
        book.save(path)
        decisions=copy.deepcopy(old['decisions']);decisions['sourceSha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        decisions['features'][0]['literalColumns'].append('At')
        decisions['features'][0]['columnTypes']['At']='datetime'
        decisions['features'][0]['reviewedBlankColumns']=['At']
        inputs=build_source_input_plan(path,decisions,connection_id='connection',folder_id='folder')
        plan=build_exact_native_calculation_plan(path,inputs)
        for cell,expected,kind in [('Ledger Renamed!D2',0,'number'),('Summary Renamed!A1',1,'number'),
                                   ('Summary Renamed!A2',0,'number'),('Summary Renamed!A3',0,'number'),
                                   ('Summary Renamed!A4',1,'number'),('Summary Renamed!A5',0,'number'),
                                   ('Summary Renamed!A6',0,'number'),('Summary Renamed!A7',5,'number')]:
            with self.subTest(cell=cell):
                record=plan['formulas'][cell]
                self.assertEqual(record['status'],'ready-to-bind')
                self.assertEqual(record['localExpected'],expected)
                self.assertEqual(record['sourceResultClass'],kind)
        self.assertEqual(plan['formulas']['Summary Renamed!A8']['localExpectedError'],'#DIV/0!')
        raw=input_scalar('Ledger Renamed!E2',plan['bindings'],lambda _binding:'Null')
        self.assertEqual(evaluate_native(raw.kind),'z')
        self.assertIn('DateDiff(',raw.number)
        for cell in ('Ledger Renamed!D2','Summary Renamed!A3','Summary Renamed!A5'):
            with self.subTest(emitted_cell=cell):
                compiled=_formula_result(compile_expression(plan['formulas'][cell]['ast'],
                    lambda _reference:raw))
                self.assertEqual(evaluate_native(compiled.kind),'n')
                self.assertEqual(evaluate_native(compiled.number),0)
                self.assertEqual(evaluate_native(compiled.error),'')
        fault=_formula_result(compile_expression(plan['formulas']['Summary Renamed!A8']['ast'],
            lambda _reference:raw))
        self.assertEqual(evaluate_native(fault.error),'#DIV/0!')

    def test_assigned_scaffold_ids_guarded_formula_update_and_runtime_separation(self):
        path,inputs=self.source();native=build_exact_native_calculation_plan(path,inputs)
        assigned={'ledger':{'elementId':'actual-input','name':'tblLedger','rowIdColumn':'Row ID','rows':{'0':'actual-a','1':'actual-b'}}}
        lowered=lower_exact_native_calculations(native,assigned)
        plan={'inputPlanHash':inputs['inputPlanHash'],'sharedInputArchitecture':inputs['sharedInputArchitecture'],
              'nativeCalculationPlan':native,'nativeReviewApproval':{'status':'reviewed','owner':'owner','reviewer':'reviewer','sourceSha256':native['sourceSha256'],'nativePlanHash':native['nativePlanHash']},
              'formulaCells':[{'id':c,'mappingStatus':'live','targetMapping':f['targetMapping']}for c,f in lowered['formulas'].items()]}
        manifest=build_manifest(path,plan);manifest['target']={'documentId':'actual-doc','version':7}
        target={'documentId':'actual-doc','version':7,'pageId':'page','elementId':'actual-input','columnIds':{'Account':'ca','Amount':'cb','Rate':'cc'},'rowIdColumnId':'rid'}
        receipt={'status':'APPROVED_STAGING','sourceHash':manifest['sourceHash'],'planHash':manifest['planHash'],'targetBindings':{'ledger':target}}
        receipt['artifacts']=build_artifact_records(manifest,inputs['sharedInputArchitecture'],{'ledger':target})
        current={'workbookId':'actual-doc','documentVersion':7,'contents':{'kind':'workbook','schemaVersion':1,'pages':[{'id':'page','name':'Page'}],
          'elements':[{'id':'actual-input','kind':'input-table','name':'tblLedger','columns':[{'id':'ca','name':'Account','type':'text'},{'id':'cb','name':'Amount','type':'number'},{'id':'cc','name':'Rate','type':'number'},{'id':'rid','name':'Row ID','type':None}]},
                      {'id':'actual-seed','kind':'table','name':'Observed Singleton','source':{'kind':'sql','statement':'SELECT 1 AS C0','connectionId':native['targetScope']['connectionId']},'columns':[]}],
          'layout':'<Page id="page"></Page>'}}
        rows={'ledger':[{'Account':'a','Amount':3,'Rate':2,'Row ID':'actual-a'},{'Account':'b','Amount':4,'Rate':1,'Row ID':'actual-b'}]}
        bindings={}
        def refresh():
            manifest['target']['version']=current['documentVersion'];target['version']=current['documentVersion']
            receipt['artifacts']=build_artifact_records(manifest,inputs['sharedInputArchitecture'],{'ledger':target})
        for index,shard in enumerate(lowered['shards']):
            refresh();backup=build_native_calculation_backup(current,rows)
            proposal=prepare_exact_native_scaffold(native,inputs,manifest,receipt,current,rows,backup=backup,
              seed_binding={'elementId':'actual-seed','documentId':'actual-doc','documentVersion':current['documentVersion'],'queriedRows':[{'C0':1}]},page_id='page',shard_index=index)
            ast.parse(proposal['program']);self.assertNotIn('parse_formula',proposal['program'])
            formula_calls=[call for call in ast.walk(ast.parse(proposal['program']))
                           if isinstance(call,ast.Call) and isinstance(call.func,ast.Attribute)
                           and ast.unparse(call.func.value)=='sigma.formula']
            self.assertTrue(formula_calls)
            for call in formula_calls:
                self.assertEqual(call.func.attr,'const_num')
                self.assertEqual(len(call.args),1)
                self.assertEqual(call.keywords,[])
                self.assertIsInstance(call.args[0],ast.Constant)
                self.assertEqual(call.args[0].value,0)
            # Local MCP SDK documentation offers a second, independent check;
            # a clean source checkout still runs the structural contract above.
            if SDK_FORMULA.is_file():check_formula_api_contract(proposal['program'])
            self.assertIn('sigma.formula.const_num(0)',proposal['program'])
            saved=copy.deepcopy(current);saved['documentVersion']+=1
            saved['contents']['elements'].append({'id':'assigned-'+str(index),'kind':'table','name':shard['name'],
               'source':{'kind':'table','elementId':'actual-seed'} if shard['source']['kind']=='sql' else shard['source'],
               'columns':[{'id':f'assigned-{index}-{j}','name':c['name'],'formula':'0'}for j,c in enumerate(shard['columns'])]})
            placement=f'<Element elementId="assigned-{index}" gridColumn="1 / 25" gridRow="{1+12*index} / {13+12*index}" />'
            saved['contents']['layout']=saved['contents']['layout'].replace('</Page>',placement+'</Page>',1)
            binding=confirm_exact_native_scaffold(proposal,saved)
            self.assertNotEqual(binding['elementId'],shard['id']);bindings[shard['id']]=binding;current=saved
            if index == 0:
                refresh()
                canary=prepare_exact_native_update(native,inputs,manifest,receipt,current,rows,
                    backup=build_native_calculation_backup(current,rows),native_bindings=bindings)
                self.assertEqual(canary['shardIndex'],0)
                self.assertEqual(canary['actualElementId'],'assigned-0')
                bad_prefix={lowered['shards'][-1]['id']:binding}
                with self.assertRaisesRegex(ValueError,'prefix'):
                    prepare_exact_native_update(native,inputs,manifest,receipt,current,rows,
                        backup=build_native_calculation_backup(current,rows),native_bindings=bad_prefix)
        refresh()
        bad=copy.deepcopy(current)
        bad['contents']['elements'][-1]['columns'].append({'id':'unowned-extra','name':'unowned-extra','formula':'5'})
        with self.assertRaisesRegex(ValueError,'scaffold|column'):
            prepare_exact_native_update(native,inputs,manifest,receipt,bad,rows,
                backup=build_native_calculation_backup(bad,rows),native_bindings=bindings)
        proposal=prepare_exact_native_update(native,inputs,manifest,receipt,current,rows,
            backup=build_native_calculation_backup(current,rows),native_bindings=bindings)
        self.assertEqual(proposal['actualElementId'],'assigned-0')
        saved={'workbookId':'actual-doc','documentVersion':current['documentVersion']+1,'contents':copy.deepcopy(proposal['dryRunRequest']['contents'])}
        result=confirm_exact_native_readback(proposal,saved)
        self.assertEqual(result['runtimeStatus'],'UNVERIFIED_CANDIDATE')
        saved['contents']['elements'][-1]['name']='tampered'
        with self.assertRaisesRegex(ValueError,'readback'):confirm_exact_native_readback(proposal,saved)

if __name__=='__main__':unittest.main()
