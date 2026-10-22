#!/usr/bin/env python3
"""table-report: recognisers, the Python oracle, Sigma emission, builder gates, and the
broader-eight fixtures (08 Inventory, 04 Sales Pipeline, 07 Headcount) end to end."""
from __future__ import annotations

import contextlib
import copy
import importlib.util
import io
import json
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / "lib"))
sys.path.insert(0, str(ROOT / "tests"))

import table_report as tr  # noqa: E402
import table_report_fixture as fx  # noqa: E402
import table_report_model as trm  # noqa: E402

FIXTURES = ROOT / "tests" / "fixtures" / "broader-eight"
INVENTORY = FIXTURES / "08 Inventory Replenishment Plan.xlsm"
PIPELINE = FIXTURES / "04 Sales Pipeline Forecast.xlsx"
HEADCOUNT = FIXTURES / "07 Headcount and Opex Plan.xlsx"


def load_script(name, filename):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


DRAFT_CLI = load_script("draft_table_report_cli", "draft-table-report.py")
BUILD_CLI = load_script("build_table_report_cli", "build-table-report.py")


def run_cli(module, argv):
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        code = module.main(argv)
    return code, out.getvalue(), err.getvalue()


def element_by_id(spec, element_id):
    return next(e for e in spec["contents"]["elements"] if e["id"] == element_id)


def element_by_name(spec, name):
    return next(e for e in spec["contents"]["elements"] if e.get("name") == name)


def check_spec_refs(test, spec):
    """Every ``[Element/Column]`` reference must name a real element (or join alias) and column."""
    elements = spec["contents"]["elements"]
    ids = [e["id"] for e in elements]
    test.assertEqual(len(ids), len(set(ids)), "duplicate element ids")
    names = {e["name"]: e for e in elements if e.get("name")}
    layout = spec["contents"]["layout"]
    for element_id in re.findall(r'elementId="([^"]+)"', layout):
        test.assertIn(element_id, ids, f"layout references unknown element {element_id}")
    test.assertTrue(set(ids) - {"trHeader"} <= set(re.findall(r'elementId="([^"]+)"', layout)),
                    "an element is missing from the layout")
    for element in elements:
        source = element.get("source") or {}
        aliases = set()
        if source.get("kind") == "join":
            aliases = {source["name"]} | {j["name"] for j in source["joins"]}
        if source.get("kind") == "table":
            test.assertIn(source["elementId"], ids)
        if source.get("kind") == "table":
            aliases.add(next(e for e in elements if e["id"] == source["elementId"])["name"])
        for column in element.get("columns", []):
            for ref, col in re.findall(r"\[([^\[\]/]+)/([^\[\]]+)\]", column.get("formula") or ""):
                if ref in aliases or ref == "Custom SQL":
                    continue
                test.assertIn(ref, names, f"{element['id']}.{column.get('name')}: unknown element {ref!r}")
                test.assertIn(col, {c.get("name") for c in names[ref].get("columns", [])},
                              f"{element['id']}.{column.get('name')}: {ref!r} has no column {col!r}")
        ids_in = {c["id"] for c in element.get("columns", [])}
        for group in element.get("groupings", []):
            test.assertTrue(set(group["groupBy"]) <= ids_in)
            test.assertTrue(set(group["calculations"]) <= ids_in)


class Workdir(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name).resolve()

    def draft(self, path, intent="what-if"):
        return tr.draft_table_report(path, intent=intent, connection="conn", folder="folder")

    def drafted(self, builder, name, **options):
        path = self.dir / f"{name}.xlsx"
        builder(path, **options)
        plan, info = self.draft(path)
        return path, plan, info

    def rejected(self, path):
        with self.assertRaises(tr.PlanDraftError) as caught:
            self.draft(path)
        return str(caught.exception)


# --------------------------------------------------------------------------- calculated columns
class TestCalculatedColumns(Workdir):
    def setUp(self):
        super().setUp()
        self.path, self.plan, self.info = self.drafted(fx.build_orders, "orders")
        self.table = next(t for t in self.plan["tables"] if t["name"] == "tblOrders")
        self.columns = {c["name"]: c for c in self.table["columns"]}

    def test_literal_and_calculated_columns_are_separated(self):
        self.assertEqual([c["name"] for c in self.table["columns"] if "expr" not in c],
                         ["Order", "Region", "Product", "Ship Date", "Qty", "Price"])
        self.assertEqual([c["name"] for c in self.table["columns"] if "expr" in c],
                         ["Rate", "Revenue", "Bucket", "Month End", "Flag", "Cap", "Boxes"])
        self.assertEqual(len(self.table["rows"]), 8)
        self.assertEqual(self.table["rows"][0][3], "2026-01-05")  # dates are ISO strings in the plan

    def test_xlookup_becomes_a_keyed_lookup_with_its_default(self):
        self.assertEqual(self.columns["Rate"]["expr"],
                         ["lk", "tblRates", ["c", "Product"], "Product", "Rate", ["n", 0]])

    def test_index_match_self_lookup_in_the_lookup_table_itself(self):
        rates = next(t for t in self.plan["tables"] if t["name"] == "tblRates")
        pct = next(c for c in rates["columns"] if c["name"] == "Pct")
        self.assertEqual(pct["expr"][0:2], ["op", "*"])
        model = trm.compute(self.plan)
        self.assertEqual([r["Pct"] for r in model.table_rows("tblRates")], [10.0, 20.0])

    def test_row_formulas_are_evaluated_like_excel(self):
        rows = trm.compute(self.plan).table_rows("tblOrders")
        first, fifth = rows[0], rows[4]
        self.assertAlmostEqual(first["Revenue"], 2 * 60 * 0.9)
        self.assertEqual(first["Bucket"], "Big")
        self.assertEqual(first["Month End"], date(2026, 1, 31))
        self.assertEqual(first["Flag"], "N")      # January: MONTH >= 2 fails
        self.assertEqual(fifth["Flag"], "Y")
        self.assertEqual(first["Cap"], 2 + 60)    # MIN(2,3) + MAX(60,50)
        self.assertEqual(first["Boxes"], 5)       # CEILING(2, 5)
        self.assertEqual(rows[3]["Rate"], 0)      # Gizmo is not in tblRates -> XLOOKUP default

    def test_calculated_columns_are_emitted_as_sigma_formulas(self):
        workbook = trm.build_wb_spec(self.plan, "<dm>", trm.local_dm_elements(self.plan))
        wrapper = next(e for e in trm.build_parts(self.plan, "<dm>", trm.local_dm_elements(self.plan))["hidden"]
                       if e["name"] == "tblOrders")
        formulas = {c["name"]: c["formula"] for c in wrapper["columns"]}
        self.assertEqual(formulas["Rate"], "Coalesce(Lookup([tblRates/Rate], [Product], [tblRates/Product]), 0)")
        self.assertEqual(formulas["Bucket"], 'If(([Revenue] >= 100), "Big", "Small")')
        self.assertEqual(formulas["Month End"],
                         'DateAdd("day", -1, DateAdd("month", 1, DateTrunc("month", [Ship Date])))')
        self.assertEqual(formulas["Flag"],
                         'If((([Qty] > 1) and (DatePart("month", [Ship Date]) >= 2)), "Y", "N")')
        self.assertEqual(formulas["Cap"], "(Least([Qty], 3) + Greatest([Price], 50))")
        self.assertEqual(formulas["Boxes"], "If(5 = 0, 0, (Ceiling([Qty] / 5) * 5))")
        self.assertIn("pages", workbook)

    def test_same_row_a1_references_in_a_calculated_column_read_the_tables_own_columns(self):
        _, plan, _ = self.drafted(fx.build_a1_table, "a1")
        table = plan["tables"][0]
        amount = next(c for c in table["columns"] if c["name"] == "Amount")
        self.assertEqual(amount["expr"], ["op", "*", ["c", "Qty"], ["c", "Price"]])
        frame = plan["frames"][0]
        rows = {r["Kind"]: r for r in trm.compute(plan).frame_rows(frame["id"])}
        self.assertEqual((rows["food"]["Amount"], rows["gear"]["Amount"]), (22.0, 20.0))

    def test_a_calculated_column_reading_another_row_is_refused(self):
        path = self.dir / "a1.xlsx"
        fx.build_a1_table(path)
        bad = self.dir / "a1-bad.xlsx"
        for row, ref in ((2, "B3*C2"), (3, "B4*C3"), (4, "B5*C4")):
            step = self.dir / f"a1-step-{row}.xlsx"
            fx.patch_cell(path if row == 2 else prev, step, "Lines", f"D{row}", formula=ref)
            prev = step
        message = self.rejected(prev)
        self.assertIn("another row of tblLines", message)

    def test_a_typed_value_among_calculated_cells_is_refused_with_its_cell(self):
        path = self.dir / "mixed.xlsx"
        fx.build_orders(path, mixed=True)
        message = self.rejected(path)
        self.assertIn("Orders!M3", message)
        self.assertIn("mixes formulas and literal values", message)


# --------------------------------------------------------------------------- summary grids
class TestSummaryGrid(Workdir):
    def setUp(self):
        super().setUp()
        self.path, self.plan, self.info = self.drafted(fx.build_orders, "orders")
        self.frame = next(f for f in self.plan["frames"] if f["kind"] == "summary")
        self.columns = {c["name"]: c for c in self.frame["columns"]}

    def test_rows_group_by_the_label_column_of_the_source_table(self):
        self.assertEqual((self.frame["by"], self.frame["source"]), ("Region", {"kind": "table", "name": "tblOrders"}))
        self.assertEqual(self.frame["labels"], ["East", "West", "North"])
        self.assertEqual([c["name"] for c in self.frame["columns"]],
                         ["Jan", "Feb", "Total", "Big", "Avg Rev", "Target", "Attain"])

    def test_literal_date_windows_fold_to_date_literals(self):
        jan = self.columns["Jan"]["expr"]
        self.assertEqual(jan, ["agg", "sum", "Revenue", ["f", "AND",
                               ["op", ">=", ["c", "Ship Date"], ["d", "2026-01-01"]],
                               ["op", "<=", ["c", "Ship Date"], ["d", "2026-01-31"]]]])

    def test_countifs_averageifs_and_row_arithmetic(self):
        self.assertEqual(self.columns["Big"]["expr"],
                         ["agg", "count", None, ["op", "=", ["c", "Bucket"], ["s", "Big"]]])
        self.assertEqual(self.columns["Avg Rev"]["expr"], ["agg", "avg", "Revenue", None])
        self.assertEqual(self.columns["Total"]["expr"], ["f", "SUM", ["c", "Jan"], ["c", "Feb"]])
        self.assertEqual(self.columns["Attain"]["expr"][0:2], ["f", "IFERROR"])

    def test_typed_quota_cells_become_an_editable_keyed_input_table(self):
        inputs = next(t for t in self.plan["tables"] if t.get("synthetic"))
        self.assertTrue(inputs["editable"])
        self.assertEqual(inputs["rows"], [["East", 400.0], ["West", 0.0], ["North", 300.0]])
        self.assertEqual(self.columns["Target"]["expr"],
                         ["lk", inputs["name"], ["key"], "Region", "Target", None])

    def test_every_cell_is_recomputed_and_ratio_guards_division_by_zero(self):
        model = trm.compute(self.plan)
        checked, problems = trm.check_cells(self.plan, model)
        self.assertEqual((checked, problems), (self.info["verifiedCells"], []))
        rows = {r["Region"]: r for r in model.frame_rows(self.frame["id"])}
        self.assertEqual(rows["West"]["Attain"], 0)  # Target 0 -> IFERROR(.., 0)
        self.assertAlmostEqual(rows["East"]["Total"], 108 + 120 + 189)
        self.assertEqual(rows["East"]["Big"], 3)

    def test_summary_is_emitted_as_a_grouped_table_with_conditional_aggregates(self):
        parts = trm.build_parts(self.plan, "<dm>", trm.local_dm_elements(self.plan))
        element = next(e for e in parts["visible"] if e["name"] == self.frame["title"])
        formulas = {c["name"]: c["formula"] for c in element["columns"]}
        self.assertEqual(formulas["Region"], "[tblOrders/Region]")
        self.assertEqual(formulas["Jan"], 'Coalesce(SumIf([tblOrders/Revenue], (([tblOrders/Ship Date] >= '
                                          'Date("2026-01-01")) and ([tblOrders/Ship Date] <= Date("2026-01-31")))), 0)')
        self.assertEqual(formulas["Big"], 'Coalesce(CountIf(([tblOrders/Bucket] = "Big")), 0)')
        self.assertEqual(formulas["Avg Rev"], "Avg([tblOrders/Revenue])")
        self.assertEqual(formulas["Attain"], "If(IsNull([Target]) or [Target] = 0, 0, ([Total] / [Target]))")
        group = element["groupings"][0]
        self.assertEqual(group["groupBy"], [element["columns"][0]["id"]])
        self.assertEqual(len(group["calculations"]), 7)

    def test_a_total_row_is_verified_against_the_frame_but_not_emitted(self):
        path, plan, info = self.drafted(fx.build_orders, "totals", totals="sum")
        frame = next(f for f in plan["frames"] if f["kind"] == "summary")
        self.assertEqual(frame["labels"], ["East", "West", "North"])  # Total is not a data row
        kinds = [g["kind"] for g in plan["checks"]]
        self.assertEqual(kinds.count("summary-total"), 5)
        self.assertIn("TOTAL_ROW_VERIFIED_ONLY", info["reviewIssues"])
        model = trm.compute(plan)
        self.assertEqual(trm.check_cells(plan, model)[1], [])
        rows = model.frame_rows(frame["id"])
        self.assertAlmostEqual(model.cell(frame["id"], "Total", "summary-total", []),
                               sum(r["Total"] for r in rows))
        total = model.cell(frame["id"], "Target", "summary-total", [])
        self.assertEqual(total, 700.0)

    def test_a_total_row_that_is_not_a_sum_of_the_rows_above_is_refused(self):
        path = self.dir / "badtotal.xlsx"
        fx.build_orders(path, totals="bad")
        message = self.rejected(path)
        self.assertIn("Summary!B7", message)
        self.assertIn("total row cell must be SUM of the grid rows above it", message)

    def test_empty_matches_are_zero_in_excel_so_every_conditional_aggregate_is_coalesced(self):
        # criteria that match no rows must feed totals as 0, not NULL (Sigma's SumIf is NULL there)
        model = trm.compute(self.plan)
        empty = ["agg", "sum", "Revenue", ["op", "=", ["c", "Bucket"], ["s", "Nope"]]]
        self.assertEqual(model.ev(["op", "+", empty, ["n", 5]], trm.Env(group=model.table_rows("tblOrders"))), 5)
        parts = trm.build_parts(self.plan, "<dm>", trm.local_dm_elements(self.plan))
        element = next(e for e in parts["visible"] if e["name"] == self.frame["title"])
        for column in element["columns"][1:]:
            formula = column["formula"]
            if re.search(r"\b(Sum|Count|Min|Max)(If)?\(", formula) and "Avg" not in formula:
                bare = re.sub(r"Coalesce\((?:Sum|Count|Min|Max)(?:If)?\(.*?\), 0\)", "", formula)
                self.assertNotRegex(bare, r"\b(Sum|Count|Min|Max)(If)?\(", column["name"])
        self.assertEqual(element["columns"][5]["formula"], "Avg([tblOrders/Revenue])")  # #DIV/0! stays an error

    def test_a_label_without_rows_is_refused_because_a_grouped_table_would_drop_it(self):
        path = self.dir / "ghost.xlsx"
        fx.build_orders(path, ghost=True)
        message = self.rejected(path)
        self.assertIn("Summary!A7", message)
        self.assertIn("'South' has no rows in tblOrders[Region]", message)

    def test_a_tampered_cached_value_names_the_cell(self):
        bad = self.dir / "tampered.xlsx"
        fx.patch_cell(self.path, bad, "Summary", "B4", value=999.0)
        message = self.rejected(bad)
        self.assertIn("Summary!B4", message)
        self.assertIn("does not reproduce the workbook's cached values", message)

    def test_an_unsupported_function_names_the_cell(self):
        bad = self.dir / "vlookup.xlsx"
        step = self.path
        for row in (4, 5, 6):  # keep the column uniform so the function itself is what is refused
            target = self.dir / f"vlookup-{row}.xlsx"
            fx.patch_cell(step, target, "Summary", f"E{row}", formula=f"VLOOKUP($A{row},tblRates,2,FALSE)")
            step = target
        bad = step
        message = self.rejected(bad)
        self.assertIn("Summary!E4", message)
        self.assertIn("unsupported function VLOOKUP", message)

    def test_a_formula_that_is_not_copied_down_is_refused_with_its_cell(self):
        bad = self.dir / "uneven.xlsx"
        fx.patch_cell(self.path, bad, "Summary", "E5",
                      formula='COUNTIFS(tblOrders[Region],$A5,tblOrders[Bucket],"Small")')
        message = self.rejected(bad)
        self.assertIn("Summary!E5", message)
        self.assertIn("one formula copied down", message)

    def test_a_formula_cell_outside_any_grid_is_never_skipped(self):
        extra = self.dir / "stray.xlsx"
        fx.build_orders(extra)
        import zipfile
        with zipfile.ZipFile(extra) as z:
            files = fx._sheet_files(z)
        with zipfile.ZipFile(extra) as source, zipfile.ZipFile(self.dir / "stray2.xlsx", "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == files["Summary"]:
                    data = data.replace(b"</sheetData>", b'<row r="12"><c r="B12"><f>SUM(B4:B6)</f><v>1</v></c>'
                                        b"</row></sheetData>")
                target.writestr(item, data)
        message = self.rejected(self.dir / "stray2.xlsx")
        self.assertIn("Summary!B12", message)
        self.assertIn("outside any recognised grid", message)


class TestDraftBoundaries(Workdir):
    def test_a_plain_data_entry_workbook_is_not_a_table_report(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.table import Table
        wb = Workbook()
        ws = wb.active
        ws.append(["Account", "Amount"])
        for i in range(1, 6):
            ws.append([f"A{i}", i * 10])
        ws.add_table(Table(displayName="tblEntries", ref="A1:B6"))
        path = self.dir / "entries.xlsx"
        wb.save(path)
        with self.assertRaises(tr.NoTableReportShape):
            self.draft(path)

    def test_a_workbook_without_formal_tables_is_not_a_table_report(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active["A1"] = "x"
        path = self.dir / "none.xlsx"
        wb.save(path)
        with self.assertRaises(tr.NoTableReportShape):
            self.draft(path)

    def test_a_missing_cached_value_is_reported_with_the_recalculate_hint(self):
        from openpyxl import load_workbook
        path = self.dir / "uncached.xlsx"
        fx.build_orders(path)
        # openpyxl rewrites formulas without cached results: exactly an un-recalculated workbook
        wb = load_workbook(path)
        wb.save(path)
        message = self.rejected(path)
        self.assertIn("no cached value", message)

    def test_duplicate_lookup_keys_are_refused_because_first_match_semantics_would_differ(self):
        path = self.dir / "dups.xlsx"
        fx.build_orders(path)
        from openpyxl import load_workbook
        wb = load_workbook(path)
        wb["Rates"]["A3"] = "Widget"   # tblRates[Product] now holds Widget twice
        wb.save(self.dir / "dups3.xlsx")
        with self.assertRaises(tr.PlanDraftError) as caught:
            self.draft(self.dir / "dups3.xlsx")
        self.assertIn("duplicate values", str(caught.exception))

    def test_a_parameter_cell_that_became_a_formula_is_refused(self):
        path = self.dir / "staff.xlsx"
        fx.build_staff(path)
        bad = self.dir / "staff-bad.xlsx"
        import zipfile
        with zipfile.ZipFile(path) as z:
            files = fx._sheet_files(z)
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(bad, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == files["Rates"]:
                    data = re.sub(rb'<c r="B1"[^>]*>.*?</c>', b'<c r="B1"><f>0.1+0</f><v>0.1</v></c>', data)
                target.writestr(item, data)
        message = self.rejected(bad)
        self.assertIn("Rates!B1", message)
        self.assertIn("formula cell outside any recognised grid", message)


# --------------------------------------------------------------------------- keyed grids
class TestKeyedGrid(Workdir):
    def test_running_balance_becomes_opening_plus_cumulative_sum(self):
        _, plan, info = self.drafted(fx.build_stock, "stock")
        keyed = next(f for f in plan["frames"] if f["kind"] == "keyed")
        self.assertEqual((keyed["table"], keyed["rowKey"], keyed["periodKey"]), ("tblFlow", "Item", "Week"))
        self.assertEqual([c["name"] for c in keyed["columns"]], ["Net Change", "Projected balance"])
        self.assertEqual(keyed["columns"][0]["expr"],
                         ["op", "+", ["neg", ["c", "Out"]], ["c", "In"]])
        self.assertEqual(keyed["columns"][1]["expr"],
                         ["op", "+", ["lk", "tblItem", ["key"], "Item", "Opening", None],
                          ["cumsum", ["c", "Net Change"], "Item", "Week"]])
        self.assertIn("RUNNING_BALANCE_PARTITION_REVIEW", info["reviewIssues"])
        model = trm.compute(plan)
        rows = model.frame_rows(keyed["id"])
        a = [r["Projected balance"] for r in rows if r["Item"] == "A"]
        self.assertEqual(a, [70, 20, 10])
        summary = {r["Item"]: r for r in model.frame_rows("s_balance_summary")}
        self.assertEqual((summary["A"]["Low"], summary["A"]["Short?"]), (10, "LOW"))
        self.assertEqual((summary["B"]["Low"], summary["B"]["Short?"]), (25, ""))

    def test_emission_partitions_the_running_sum_by_the_grouping_and_merges_the_summary(self):
        _, plan, _ = self.drafted(fx.build_stock, "stock")
        parts = trm.build_parts(plan, None, {})
        element = next(e for e in parts["visible"] if e["name"] == "Balance")
        formulas = {c["name"]: c["formula"] for c in element["columns"]}
        self.assertEqual(formulas["Net Change"], "((-[Out]) + [In])")
        self.assertEqual(formulas["Projected balance"],
                         "(Lookup([tblItem/Opening], [Item], [tblItem/Item]) + CumulativeSum([Net Change]))")
        self.assertEqual(formulas["Low"], "Coalesce(Min([Projected balance]), 0)")
        self.assertEqual(formulas["Short?"], 'If(([Low] < 20), "LOW", "")')
        group = element["groupings"][0]
        self.assertEqual(len(group["calculations"]), 2)  # the summary rides on the keyed grouping
        self.assertEqual(element["sort"][0]["direction"], "ascending")
        self.assertNotIn("Balance Summary", [e.get("name") for e in parts["visible"]])
        self.assertTrue(any(e["kind"] == "pivot-table" for e in parts["visible"]))

    def test_a_plain_keyed_grid_has_no_running_sum(self):
        _, plan, info = self.drafted(fx.build_stock, "plain", running=False)
        keyed = next(f for f in plan["frames"] if f["kind"] == "keyed")
        self.assertEqual([c["name"] for c in keyed["columns"]], ["Projected balance"])
        self.assertEqual(keyed["columns"][0]["expr"], ["c", "In"])
        self.assertNotIn("RUNNING_BALANCE_PARTITION_REVIEW", info["reviewIssues"])

    def test_a_sparse_grid_without_a_spine_is_refused(self):
        path = self.dir / "sparse.xlsx"
        fx.build_stock(path, sparse=True)
        message = self.rejected(path)
        self.assertIn("do not pair one-to-one on (Item, Week)", message)
        self.assertIn("no spine for sparse grids", message)


# --------------------------------------------------------------------------- expansion grids
class TestExpansionGrid(Workdir):
    def setUp(self):
        super().setUp()
        self.path, self.plan, self.info = self.drafted(fx.build_staff, "staff")
        self.frame = next(f for f in self.plan["frames"] if f["kind"] == "expansion")

    def test_ordinal_rows_expand_the_table_across_the_period_axis(self):
        self.assertEqual((self.frame["base"], self.frame["attrs"]), ("tblPeople", ["Name", "Dept"]))
        self.assertEqual(self.frame["axis"]["values"], ["2026-01-01", "2026-02-01", "2026-03-01"])
        self.assertEqual(self.frame["measure"], "Cost by person")
        self.assertEqual(self.plan["params"], [{"name": "Uplift", "value": 0.1, "cell": "Rates!B1"}])
        expr = self.frame["columns"][0]["expr"]
        self.assertIn(["p", "Uplift"], list(trm.walk(expr)))
        self.assertIn(["c", "Period"], list(trm.walk(expr)))

    def test_every_cell_including_attributes_and_row_totals_is_verified(self):
        model = trm.compute(self.plan)
        checked, problems = trm.check_cells(self.plan, model)
        self.assertEqual(problems, [])
        kinds = {g["kind"] for g in self.plan["checks"]}
        self.assertTrue({"expansion", "rowtotal", "table", "summary"} <= kinds)
        rows = {r["Dept"]: r for r in model.frame_rows("s_cost_by_department")}
        self.assertAlmostEqual(rows["Eng"]["Cost"], 120000 / 12 * 1.1 * 3 + 96000 / 12 * 1.1 * 2)
        self.assertEqual(rows["Eng"]["Active"], 2)

    def test_summary_over_the_expansion_filters_period_columns_by_their_header_date(self):
        summary = next(f for f in self.plan["frames"] if f["kind"] == "summary")
        columns = {c["name"]: c["expr"] for c in summary["columns"]}
        self.assertEqual(columns["Cost"], ["agg", "sum", "Cost by person", None])
        self.assertEqual(columns["Active"], ["agg", "count", None, ["f", "AND",
                         ["op", "=", ["c", "Period"], ["d", "2026-03-01"]],
                         ["op", ">", ["c", "Cost by person"], ["n", 0]]]])
        self.assertEqual(summary["source"], {"kind": "frame", "name": self.frame["id"]})

    def test_emission_cross_joins_on_a_constant_key_and_reads_parameters_by_name(self):
        elements = trm.local_dm_elements(self.plan)
        parts = trm.build_parts(self.plan, "<dm>", elements)
        expansion = next(e for e in parts["hidden"] if e["name"] == "Cost")
        self.assertEqual(expansion["source"]["kind"], "join")
        self.assertEqual(expansion["source"]["joins"][0]["joinType"], "inner")
        self.assertEqual(expansion["source"]["joins"][0]["columns"], [{"left": "[Join Key]", "right": "[Join Key]"}])
        base = next(e for e in parts["hidden"] if e["name"] == "tblPeople")
        self.assertIn({"id": "w_tblpeople_jk", "name": "Join Key", "formula": "1"}, base["columns"])
        measure = next(c for c in expansion["columns"] if c["name"] == "Cost by person")["formula"]
        self.assertIn('Lookup([Parameters/Value], "Uplift", [Parameters/Parameter])', measure)
        self.assertIn("[Period]", measure)
        pivot = next(e for e in parts["visible"] if e["kind"] == "pivot-table")
        self.assertEqual(pivot["axisTotals"], [{"axisId": "grand-row-total", "visibility": "hidden"}])

    def test_read_only_tables_live_in_the_data_model_and_editable_ones_do_not(self):
        read_only = tr.draft_table_report(self.path, intent="read-only", connection="c", folder="f")[0]
        self.assertTrue(all(not t["editable"] for t in read_only["tables"]))
        dm = trm.build_dm_spec(read_only)
        names = [c["name"] for e in dm["pages"][0]["elements"] for c in e["columns"]]
        self.assertIn("Salary", names)
        self.assertIn("Value", names)  # parameters are VALUES too
        sql = dm["pages"][0]["elements"][0]["source"]["statement"]
        self.assertIn("'2026-01-01'::date", sql)
        editable_dm = trm.build_dm_spec(self.plan)
        self.assertEqual([c["name"] for c in editable_dm["pages"][0]["elements"][0]["columns"]], ["Period"])

    def test_an_ordinal_grid_that_misses_a_table_row_is_refused(self):
        path = self.dir / "short.xlsx"
        fx.build_staff(path, short=True)
        message = self.rejected(path)
        self.assertIn("do not cover all 4 rows of tblPeople", message)


# --------------------------------------------------------------------------- model internals
class TestModelPieces(unittest.TestCase):
    def test_blank_and_text_comparison_follow_excel(self):
        self.assertTrue(trm._compare("=", None, ""))
        self.assertTrue(trm._compare("=", None, 0))
        self.assertTrue(trm._compare(">=", date(2026, 1, 1), None))
        self.assertFalse(trm._compare("=", "a", 1))
        self.assertTrue(trm._compare(">", "a", 1))  # text sorts above numbers

    def test_values_match_uses_tolerance_for_numbers_and_exact_text(self):
        self.assertTrue(trm.values_match(1.004, 1.0, 0.01))
        self.assertFalse(trm.values_match(1.02, 1.0, 0.01))
        self.assertTrue(trm.values_match("", None, 0.01))
        self.assertFalse(trm.values_match("ORDER", "order", 0.01))
        self.assertTrue(trm.values_match(date(2026, 1, 31), "2026-01-31", 0.01))

    def test_division_by_zero_is_an_excel_error_that_iferror_catches(self):
        plan = {"tables": [], "frames": [], "params": []}
        model = trm.Model(plan)
        div = ["op", "/", ["n", 1], ["n", 0]]
        with self.assertRaises(trm.EvalError):
            model.ev(div, trm.Env(row={}))
        self.assertEqual(model.ev(["f", "IFERROR", div, ["n", 7]], trm.Env(row={})), 7)

    def test_invalid_ir_is_rejected(self):
        for bad in (["op", "%", ["n", 1], ["n", 2]], ["f", "VLOOKUP", ["n", 1]], ["zzz"], ["lk", "t"], "x"):
            with self.assertRaises(trm.PlanError):
                trm.validate_ir(bad)

    def test_emitter_handles_blank_text_comparisons_by_column_type(self):
        em = trm.Emitter(col=lambda n: f"[{n}]", coltype={"End": "datetime", "Note": "text"}.get)
        self.assertEqual(em.emit(["op", "=", ["c", "End"], ["s", ""]]), "IsNull([End])")
        self.assertEqual(em.emit(["op", "=", ["c", "Note"], ["s", ""]]), '(IsNull([Note]) or [Note] = "")')
        self.assertEqual(em.emit(["op", "<>", ["c", "End"], ["s", ""]]), "Not(IsNull([End]))")

    def test_values_sql_types_nulls_dates_and_text(self):
        sql = trm.values_sql([["it's", 1.5, date(2026, 1, 2), None], [None, 2, None, 3]],
                             [("A", "text"), ("B", "number"), ("C", "datetime"), ("D", "number")])
        self.assertIn("'it''s',1.5,'2026-01-02'::date,NULL::number", sql)
        self.assertIn("NULL::varchar,2,NULL::date,3", sql)
        with self.assertRaises(trm.PlanError):
            trm.values_sql([], [("A", "text")])

    def test_identical_column_sets_cannot_be_told_apart_in_the_data_model_readback(self):
        plan = {"name": "x", "folderId": "f", "connectionId": "c", "frames": [], "params": [],
                "tables": [{"name": n, "editable": False, "columns": [{"name": "A", "type": "text"}],
                            "rows": [["v"]]} for n in ("t1", "t2")]}
        with self.assertRaisesRegex(trm.PlanError, "identical column sets"):
            trm.build_dm_spec(plan)


# --------------------------------------------------------------------------- builder gates
class TestBuilderCli(Workdir):
    def setUp(self):
        super().setUp()
        self.path, self.plan, _ = self.drafted(fx.build_staff, "staff")
        self.plan_path = self.dir / "plan.json"
        self.plan_path.write_text(json.dumps(self.plan))
        self.out = self.dir / "build"

    def build(self, *extra, plan_path=None):
        return run_cli(BUILD_CLI, [str(plan_path or self.plan_path), "--out-dir", str(self.out), *extra])

    def test_dry_run_writes_specs_seeds_and_reports_parity(self):
        code, out, err = self.build()
        self.assertEqual(code, 0, err)
        self.assertIn("MIGRATION GATE: ", out)
        self.assertIn("LOCAL PARITY: 28/28 source cells reproduced", out)
        self.assertIn("PARITY PASS Cost by department: 4/4 summary cells", out)
        self.assertIn("DRY RUN", out)
        for name in ("dm_spec.json", "wb_spec.json", "seed_tblpeople.csv", "seed_parameters.csv"):
            self.assertTrue((self.out / name).is_file(), name)
        seed = (self.out / "seed_tblpeople.csv").read_text().splitlines()
        self.assertEqual(seed[0], "Name,Dept,Salary,Start")
        self.assertEqual(seed[1], "Ann,Eng,120000,2026-01-01")
        spec = json.loads((self.out / "wb_spec.json").read_text())
        check_spec_refs(self, spec)
        self.assertEqual([p["id"] for p in spec["contents"]["pages"]], ["report", "inputs", "model"])
        hidden = next(p for p in spec["contents"]["pages"] if p["id"] == "model")
        self.assertEqual(hidden["visibility"], "hidden")

    def test_a_cell_mismatch_exits_1_before_any_spec_is_written(self):
        bad = copy.deepcopy(self.plan)
        bad["checks"][0]["cells"][0][2] = "tampered"
        bad_path = self.dir / "bad-plan.json"
        bad_path.write_text(json.dumps(bad))
        code, out, err = self.build(plan_path=bad_path)
        self.assertEqual(code, 1)
        self.assertIn("LOCAL PARITY FAILED", err)
        self.assertFalse((self.out / "wb_spec.json").exists())

    def test_a_parity_target_mismatch_exits_1_before_any_spec_is_written(self):
        bad = copy.deepcopy(self.plan)
        bad["parityTargets"][0]["value"] += 5
        bad_path = self.dir / "bad-target.json"
        bad_path.write_text(json.dumps(bad))
        code, out, err = self.build(plan_path=bad_path)
        self.assertEqual(code, 1)
        self.assertIn("PARITY FAIL", out)
        self.assertFalse((self.out / "wb_spec.json").exists())

    def test_edited_table_data_that_breaks_parity_is_caught(self):
        bad = copy.deepcopy(self.plan)
        next(t for t in bad["tables"] if t["name"] == "tblPeople")["rows"][0][2] = 1.0
        bad_path = self.dir / "bad-data.json"
        bad_path.write_text(json.dumps(bad))
        code, _, err = self.build(plan_path=bad_path)
        self.assertEqual(code, 1)
        self.assertIn("source cells differ", err)

    def test_a_changed_source_workbook_fails_the_gate(self):
        other = self.dir / "other.xlsx"
        fx.build_staff(other, short=True)
        code, _, err = self.build("--source", str(other))
        self.assertEqual(code, 1)
        self.assertIn("SHA256 does not match", err)

    def test_post_keeps_the_strict_disposition_contract_and_never_reaches_the_network(self):
        code, _, err = self.build("--post")
        self.assertEqual(code, 1)
        self.assertIn("MIGRATION GATE", err)
        self.assertIn("review is unresolved", err)

    def test_post_creates_the_data_model_reads_its_ids_back_and_dry_runs_the_workbook_first(self):
        from migration_gate import gate_plan_source
        _, report = gate_plan_source(self.plan, plan_path=self.plan_path, draft=True)
        dispositions = self.dir / "review.json"
        dispositions.write_text(json.dumps({
            "sourceSha256": self.plan["sourceSha256"],
            "issues": [{"code": code, "owner": "tester", "disposition": "accepted", "verified": True}
                       for code in report["pendingReview"]]}))
        calls = []

        class FakeHelpers:
            @staticmethod
            def env():
                return {"SIGMA_BASE_URL": "https://sigma.example"}

            @staticmethod
            def token(env):
                return "token"

            @staticmethod
            def api(base, tok, method, path, body=None):
                calls.append((method, path, body))
                if (method, path) == ("POST", "/v2/dataModels/spec"):
                    return {"dataModelId": "dm-123"}
                if (method, path) == ("GET", "/v2/dataModels/dm-123/spec"):
                    live = copy.deepcopy(calls[0][2])
                    for number, element in enumerate(live["pages"][0]["elements"], 1):
                        element["id"] = f"live-{number}"   # ids are reassigned on POST
                    return live
                if method == "POST" and path == "/v2/workbooks" and body.get("dryRun"):
                    return {"dryRun": True, "warnings": []}
                if method == "POST" and path == "/v2/workbooks":
                    return {"workbookId": "wb-1", "url": "https://sigma.example/wb-1"}
                raise AssertionError((method, path))

        original = BUILD_CLI._forecast_helpers
        BUILD_CLI._forecast_helpers = lambda: FakeHelpers
        self.addCleanup(setattr, BUILD_CLI, "_forecast_helpers", original)
        code, out, err = self.build("--post", "--review-dispositions", str(dispositions))
        self.assertEqual(code, 0, err)
        self.assertEqual([(m, p) for m, p, _ in calls],
                         [("POST", "/v2/dataModels/spec"), ("GET", "/v2/dataModels/dm-123/spec"),
                          ("POST", "/v2/workbooks"), ("POST", "/v2/workbooks")])
        self.assertTrue(calls[2][2]["dryRun"])
        self.assertNotIn("dryRun", calls[3][2])
        wire = json.dumps(calls[3][2])
        self.assertIn('"dataModelId": "dm-123"', wire)
        self.assertIn('"elementId": "live-1"', wire)
        self.assertNotIn('"elementId": "dm1"', wire)
        self.assertIn("workbookId: wb-1", out)

    def test_post_without_any_read_only_table_skips_the_data_model(self):
        stock = self.dir / "stock.xlsx"
        fx.build_stock(stock)
        plan = tr.draft_table_report(stock, intent="what-if", connection="c", folder="f")[0]
        self.assertIsNone(trm.build_dm_spec(plan))
        plan_path = self.dir / "stock-plan.json"
        plan_path.write_text(json.dumps(plan))
        code, out, err = run_cli(BUILD_CLI, [str(plan_path), "--out-dir", str(self.dir / "stock-build")])
        self.assertEqual(code, 0, err)
        self.assertIn("no read-only tables: no data model is needed", out)
        self.assertFalse((self.dir / "stock-build" / "dm_spec.json").exists())

    def test_draft_cli_reports_verified_cells_and_review_issues(self):
        out_path = self.dir / "cli-plan.json"
        code, out, err = run_cli(DRAFT_CLI, [str(self.path), "--connection", "c", "--folder", "f",
                                             "--out", str(out_path)])
        self.assertEqual(code, 0, err)
        self.assertIn("verified: 28 source cells reproduced", out)
        self.assertIn("frame expansion: Cost", out)
        self.assertIn("review issues: GROUPED_ROWS_FROM_DATA", out)
        self.assertEqual(json.loads(out_path.read_text())["archetype"], "table-report")

    def test_draft_cli_failure_exits_1_with_the_cell(self):
        bad = self.dir / "bad.xlsx"
        fx.patch_cell(self.path, bad, "Dept", "B4", value=1.0)
        code, _, err = run_cli(DRAFT_CLI, [str(bad), "--connection", "c", "--folder", "f",
                                           "--out", str(self.dir / "never.json")])
        self.assertEqual(code, 1)
        self.assertIn("DRAFT FAILED", err)
        self.assertIn("Dept!B4", err)
        self.assertFalse((self.dir / "never.json").exists())


# --------------------------------------------------------------------------- broader-eight fixtures
class TestBroaderEightFixtures(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.plans = {}
        for key, path in (("08", INVENTORY), ("04", PIPELINE), ("07", HEADCOUNT)):
            cls.plans[key] = tr.draft_table_report(path, intent="what-if", connection="conn", folder="folder")

    def model(self, key):
        return trm.compute(self.plans[key][0])

    def summary_rows(self, key, frame_id):
        return {r[self.model(key).frames[frame_id]["by"]]: r for r in self.model(key).frame_rows(frame_id)}

    def test_08_inventory_projection_reproduces_every_cell(self):
        plan, info = self.plans["08"]
        self.assertEqual((info["verifiedCells"], info["formulaCells"]), (90, 90))
        self.assertEqual([k for k, _ in info["frames"]], ["keyed", "summary"])
        self.assertEqual([t["name"] for t in plan["tables"]], ["tblSKU", "tblDemand", "tblSnapshots"])
        rows = self.summary_rows("08", "s_projection_summary")
        self.assertEqual((rows["SKU-200"]["Min Proj"], rows["SKU-200"]["Reorder?"],
                          rows["SKU-200"]["Suggested Qty"]), (-399, "ORDER", 575))
        self.assertEqual((rows["SKU-400"]["Min Proj"], rows["SKU-400"]["Reorder?"],
                          rows["SKU-400"]["Suggested Qty"]), (475, "", 0))
        self.assertEqual(rows["SKU-300"]["Suggested Qty"], 1150)
        self.assertEqual(sum(r["Suggested Qty"] for r in rows.values()), 2550)
        self.assertEqual(sum(r["Min Proj"] for r in rows.values()), -1137)
        projection = [r for r in self.model("08").frame_rows("k_projection") if r["SKU"] == "SKU-200"]
        self.assertEqual([r["Projected ending inventory by week"] for r in projection][:3], [555, 449, 385])
        self.assertTrue(plan["tables"][2]["editable"] and plan["tables"][2]["rows"] == [])

    def test_08_empty_snapshot_table_keeps_its_structure_with_guessed_types(self):
        plan, info = self.plans["08"]
        snap = plan["tables"][2]
        self.assertEqual([c["name"] for c in snap["columns"]],
                         ["Committed At", "SKU", "Week Start", "Forecast Units", "Committed By"])
        self.assertEqual([c["type"] for c in snap["columns"]],
                         ["datetime", "text", "datetime", "number", "text"])
        self.assertIn("TABLE_COLUMN_TYPES_GUESSED", info["reviewIssues"])
        self.assertIn("RUNNING_BALANCE_PARTITION_REVIEW", info["reviewIssues"])

    def test_04_pipeline_summaries_match_the_recorded_parity_values(self):
        plan, info = self.plans["04"]
        self.assertEqual((info["verifiedCells"], info["formulaCells"]), (232, 228))  # +4 typed quota cells
        owners = self.summary_rows("04", "s_h2_fy26_weighted_forecast")
        self.assertEqual({k: round(v["Total Weighted"], 2) for k, v in owners.items()},
                         {"Jordan Lee": 687400.0, "Sam Ortiz": 280200.0, "Priya Nair": 1342500.0,
                          "Alex Kim": 955300.0})
        self.assertEqual(owners["Jordan Lee"]["Commit"], 605000)
        self.assertAlmostEqual(owners["Priya Nair"]["Coverage"], 1.3425)
        self.assertEqual(owners["Jordan Lee"]["Quota"], 900000)
        stages = self.summary_rows("04", "s_open_deals_by_stage")
        self.assertEqual(sum(r["Pipeline $"] for r in stages.values()), 7673000)
        self.assertEqual(sum(r["# Deals"] for r in stages.values()), 60)
        self.assertEqual(stages["Closed Lost"]["# Deals"], 7)
        pipeline = next(t for t in plan["tables"] if t["name"] == "tblPipeline")
        self.assertEqual([c["name"] for c in pipeline["columns"] if "expr" in c],
                         ["Probability", "Weighted Amount", "Forecast Category"])
        rows = self.model("04").table_rows("tblPipeline")
        lost = next(r for r in rows if r["Stage"] == "Closed Lost")
        self.assertEqual((lost["Probability"], lost["Forecast Category"]), (0, "Omitted"))  # tblStages row

    def test_07_headcount_plan_matches_budget_variance_and_december_headcount(self):
        plan, info = self.plans["07"]
        self.assertEqual((info["verifiedCells"], info["formulaCells"]), (625, 625))
        self.assertEqual([p["name"] for p in plan["params"]],
                         ["Benefits load", "Merit effective month", "Merit increase"])
        rows = self.summary_rows("07", "s_plan_vs_budget_by_department")
        self.assertEqual({k: round(v["FY26 Plan"], 2) for k, v in rows.items()},
                         {"Engineering": 1237257.92, "Sales": 1036293.42, "Marketing": 1394753.82,
                          "G&A": 933163.77, "Support": 1155056.35})
        self.assertEqual({k: v["FY26 Budget"] for k, v in rows.items()},
                         {"Engineering": 1251000, "Sales": 1052000, "Marketing": 1371000,
                          "G&A": 932000, "Support": 1170000})
        self.assertEqual({k: v["Dec Headcount"] for k, v in rows.items()},
                         {"Engineering": 8, "Sales": 7, "Marketing": 7, "G&A": 8, "Support": 7})
        self.assertAlmostEqual(rows["Marketing"]["Variance"], -23753.816667, places=4)
        wide = next(t for t in plan["tables"] if t["name"] == "tblDeptBudget")
        self.assertEqual(wide["columns"][-1]["name"], "Total (Jan to Dec)")
        self.assertTrue(wide["columns"][-1]["synthetic"])

    def test_04_empty_month_windows_and_missing_categories_are_coalesced_to_zero(self):
        plan = self.plans["04"][0]
        parts = trm.build_parts(plan, "<dm>", trm.local_dm_elements(plan))
        element = next(e for e in parts["visible"] if e["name"] == "H2 FY26 Weighted Forecast")
        formulas = {c["name"]: c["formula"] for c in element["columns"]}
        self.assertTrue(formulas["Jul-26"].startswith("Coalesce(SumIf("))   # Jordan Lee has no July deals
        self.assertEqual(formulas["Commit"].count("Coalesce(SumIf("), 2)   # both terms of the addition
        rows = self.summary_rows("04", "s_h2_fy26_weighted_forecast")
        self.assertEqual(rows["Jordan Lee"]["Jul-26"], 0)
        self.assertEqual(rows["Sam Ortiz"]["Dec-26"], 0)
        self.assertEqual(round(rows["Sam Ortiz"]["Total Weighted"], 2), 280200.0)

    def test_no_emitted_element_asks_for_visible_axis_totals(self):
        import workbook_wire
        for key, plan_info in self.plans.items():
            plan = plan_info[0]
            spec = workbook_wire.wire_workbook(
                trm.build_wb_spec(plan, "<dm>", trm.local_dm_elements(plan)))
            for element in spec["contents"]["elements"]:
                self.assertNotIn('"visibility": "visible"', json.dumps(element), (key, element["id"]))
                for total in element.get("axisTotals", []):
                    self.assertEqual(total["visibility"], "hidden")

    def test_every_fixture_plan_builds_specs_whose_references_resolve(self):
        for key, plan_info in self.plans.items():
            plan = plan_info[0]
            elements = trm.local_dm_elements(plan)
            spec = trm.build_wb_spec(plan, "<dm>", elements)
            import workbook_wire
            check_spec_refs(self, workbook_wire.wire_workbook(spec))

    def test_read_only_intent_hosts_every_non_empty_table_in_the_data_model(self):
        import workbook_wire
        for key, path, expected_inputs in (("04", PIPELINE, 0), ("07", HEADCOUNT, 0), ("08", INVENTORY, 1)):
            plan = tr.draft_table_report(path, intent="read-only", connection="c", folder="f")[0]
            dm = trm.build_dm_spec(plan)
            elements = trm.map_dm_elements(plan, dm) if dm else {}
            parts = trm.build_parts(plan, "<dm>", elements)
            self.assertEqual(len(parts["inputs"]), expected_inputs, key)  # only the empty snapshot log
            spec = workbook_wire.wire_workbook(trm.build_wb_spec(plan, "<dm>", elements))
            check_spec_refs(self, spec)
            if dm:
                sources = [e["source"] for e in spec["contents"]["elements"] if e.get("source", {}).get("kind") == "data-model"]
                self.assertGreaterEqual(len(sources), 1)
                self.assertTrue(all(src["elementId"].startswith("dm") for src in sources))
            sql = [e["source"]["statement"] for e in (dm or {"pages": [{"elements": []}]})["pages"][0]["elements"]]
            self.assertTrue(all(statement.startswith("select * from (values") for statement in sql))

    def test_flipping_a_table_editable_flag_in_the_plan_moves_it_between_input_table_and_data_model(self):
        import workbook_wire
        plan = copy.deepcopy(self.plans["08"][0])
        next(t for t in plan["tables"] if t["name"] == "tblSKU")["editable"] = False   # reference data
        dm = trm.build_dm_spec(plan)
        self.assertEqual([c["name"] for c in dm["pages"][0]["elements"][0]["columns"]][:3],
                         ["SKU", "Description", "On Hand"])
        elements = trm.map_dm_elements(plan, dm)
        parts = trm.build_parts(plan, "<dm>", elements)
        self.assertEqual(sorted(e["name"] for e in parts["inputs"]), ["tblDemand Input", "tblSnapshots Input"])
        spec = workbook_wire.wire_workbook(trm.build_wb_spec(plan, "<dm>", elements))
        check_spec_refs(self, spec)
        wrapper = element_by_name(spec, "tblSKU")
        self.assertEqual(wrapper["source"]["kind"], "data-model")
        self.assertEqual(wrapper["columns"][0]["formula"], "[Custom SQL/SKU]")
        self.assertEqual(trm.compute(plan).frame_rows("s_projection_summary")[0]["Min Proj"], -399)

    def test_fixture_builds_run_through_both_clis(self):
        with tempfile.TemporaryDirectory() as tmp:
            for key, path in (("08", INVENTORY), ("04", PIPELINE), ("07", HEADCOUNT)):
                plan = Path(tmp) / f"{key}.json"
                code, out, err = run_cli(DRAFT_CLI, [str(path), "--connection", "c", "--folder", "f",
                                                     "--out", str(plan)])
                self.assertEqual(code, 0, err)
                code, out, err = run_cli(BUILD_CLI, [str(plan), "--out-dir", str(Path(tmp) / key)])
                self.assertEqual(code, 0, err)
                self.assertNotIn("FAIL", out)
                self.assertIn("PARITY PASS", out)
                self.assertIn("source cells reproduced", out)

    def test_every_summary_cell_of_every_fixture_is_a_parity_target(self):
        counts = {k: len(p[0]["parityTargets"]) for k, p in self.plans.items()}
        self.assertEqual(counts, {"08": 18, "04": 52, "07": 25})

    def test_the_oracle_notices_edited_table_data(self):
        plan = copy.deepcopy(self.plans["04"][0])
        pipeline = next(t for t in plan["tables"] if t["name"] == "tblPipeline")
        amount = [c["name"] for c in pipeline["columns"] if "expr" not in c].index("Amount")
        pipeline["rows"][0][amount] += 1000
        model = trm.compute(plan)
        checked, problems = trm.check_cells(plan, model)
        self.assertGreater(len(problems), 0)
        self.assertIn("Forecast!", problems[0] + problems[-1] + "".join(problems))


if __name__ == "__main__":
    unittest.main()
