"""Source-bound workbook calculations, with hand-derived Excel expectations."""
import hashlib
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts/lib"))
from excel_expression import Function, UnsupportedExpression, emit_sigma, extract_references, parse_expression  # noqa: E402
from workbook_expression import WorkbookContext  # noqa: E402
from formula_graph import _expression_cells  # noqa: E402
from source_workbook import _walk_references  # noqa: E402


class WorkbookExpressionTests(unittest.TestCase):
    def save(self, workbook):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        path = Path(temporary.name) / "source.xlsx"
        workbook.save(path)
        return path

    def table_book(self):
        book = Workbook()
        data = book.active
        data.title = "Sales O'Brien"
        data.append(["Product Line", "Month", "Units", "Unit Price", "Revenue"])
        data.append(["North", datetime(2026, 1, 1), 2, 10, "=tblOrders[[#This Row],[Units]]*tblOrders[[#This Row],[Unit Price]]"])
        data.append(["South", datetime(2026, 2, 1), 3, 7, "=tblOrders[[#This Row],[Units]]*tblOrders[[#This Row],[Unit Price]]"])
        data.append(["North", datetime(2026, 3, 1), 4, 5, "=tblOrders[[#This Row],[Units]]*tblOrders[[#This Row],[Unit Price]]"])
        table = Table(displayName="tblOrders", ref="A1:E4")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        data.add_table(table)
        report = book.create_sheet("FY Board")
        report["A1"] = "North"
        report["B1"] = '=SUMIFS(tblOrders[Revenue],tblOrders[Product Line],A1,tblOrders[Month],">="&DATE(2026,1,1),tblOrders[Month],"<"&EDATE(DATE(2026,1,1),3))'
        report["B2"] = '=VLOOKUP(A1,tblOrders,3,FALSE)'
        report["B3"] = '=_xlfn.XLOOKUP(A1,tblOrders[Product Line],tblOrders[Revenue],0)'
        report["B4"] = '=INDEX(tblOrders[Revenue],MATCH("South",tblOrders[Product Line],0))'
        report["B5"] = '=COUNTIFS(tblOrders[Product Line],"N*")'
        report["B6"] = '=IFERROR(1/0,99)'
        report["B7"] = '=IF(FALSE,1/0,12)'
        return book

    def test_table_binding_and_lookups_use_source_cells_and_report_all_dependencies(self):
        path = self.save(self.table_book())
        context = WorkbookContext.from_path(path)
        self.assertEqual(context.evaluate_cell("FY Board", "B1"), 40)
        self.assertEqual(context.evaluate_cell("FY Board", "B2"), 2)
        self.assertEqual(context.evaluate_cell("FY Board", "B3"), 20)
        self.assertEqual(context.evaluate_cell("FY Board", "B4"), 21)
        self.assertEqual(context.evaluate_cell("FY Board", "B5"), 2)
        self.assertEqual(context.evaluate_cell("FY Board", "B6"), 99)
        self.assertEqual(context.evaluate_cell("FY Board", "B7"), 12)
        self.assertEqual(context.evaluate_cell("FY Board", "B1", overrides={("Sales O'Brien", "C2"): 5}), 70)
        report = context.formula_report()
        self.assertEqual(report["sourceSha256"], hashlib.sha256(path.read_bytes()).hexdigest())
        self.assertEqual(report["formulaCount"], 10)
        self.assertEqual(report["counts"], {"verified": 10, "unresolved": 0, "error": 0, "cycle": 0})
        self.assertIn("Sales O'Brien!E2", report["formulas"]["FY Board!B1"]["dependencies"])
        self.assertIn("Sales O'Brien!B4", report["formulas"]["FY Board!B1"]["dependencies"])

    def test_column_domains_are_bounded_only_where_blank_rows_cannot_change_result(self):
        book = Workbook()
        data = book.active
        data.title = "Facts"
        data.append(["Code", "Amount"])
        data.append(["A", 7])
        data.append(["B", 9])
        data["Z1048576"].number_format = "0.00"
        report = book.create_sheet("Result")
        report["A1"] = '=SUMIFS(Facts!$B:$B,Facts!$A:$A,"A")'
        report["A2"] = '=COUNTA(Facts!$A:$A)'
        report["A3"] = '=COUNTIFS(Facts!$A:$A,"")'
        report["B3"] = None
        report["A4"] = '=COUNTIFS(Facts!$A:$A,B3)'
        context = WorkbookContext.from_path(self.save(book))
        self.assertEqual(context.active_last_row, 4)
        self.assertEqual(context.evaluate_cell("Result", "A1"), 7)
        self.assertEqual(context.evaluate_cell("Result", "A2"), 3)
        report = context.formula_report()
        self.assertEqual(report["formulas"]["Result!A3"]["status"], "unresolved")
        self.assertEqual(report["formulas"]["Result!A4"]["status"], "unresolved")
        self.assertEqual(report["formulas"]["Result!A1"]["bindingEvidence"][0]["declaredDomain"], "1:1048576")
        with self.assertRaises(UnsupportedExpression):
            context.bind_formula("=SUM(Facts!A1:A1048576)", sheet="Result", address="C1")

    def test_names_scenarios_external_bindings_and_cycles_remain_source_bound(self):
        book = Workbook()
        values = book.active
        values.title = "Inputs"
        values["B3"] = "Base"
        values["B4"] = 0.5
        values["B5"] = 0.25
        book.defined_names.add(DefinedName("Scenario", attr_text="Inputs!$B$3"))
        book.defined_names.add(DefinedName("Coll_Base", attr_text="Inputs!$B$4"))
        book.defined_names.add(DefinedName("Coll_Downside", attr_text="Inputs!$B$5"))
        result = book.create_sheet("Cash")
        result["A1"] = '=100*INDIRECT("Coll_"&Scenario)'
        result["A2"] = "='[FY26 Corporate Allocations.xlsx]Alloc'!$C$4"
        result["A3"] = '=A4+1'
        result["A4"] = '=A3+1'
        path = self.save(book)
        context = WorkbookContext.from_path(path)
        self.assertEqual(context.evaluate_cell("Cash", "A1"), 50)
        self.assertEqual(context.evaluate_cell("Cash", "A1", overrides={("Inputs", "B3"): "Downside"}), 25)
        report = context.formula_report()
        self.assertEqual(report["formulas"]["Cash!A2"]["status"], "unresolved")
        self.assertEqual(report["formulas"]["Cash!A3"]["status"], "cycle")
        self.assertEqual(report["formulas"]["Cash!A4"]["status"], "cycle")
        self.assertEqual(set(report["formulas"]["Cash!A1"]["dependencies"]),
                         {"Inputs!B3", "Inputs!B4", "Inputs!B5"})
        bound = WorkbookContext.from_path(path, external_inputs={("FY26 Corporate Allocations.xlsx", "Alloc", "C4"): 17})
        self.assertEqual(bound.evaluate_cell("Cash", "A2"), 17)
        reviewed_selector = WorkbookContext.from_path(path, indirect_bindings={
            ("Cash", "A1"): {"Coll_Base": "Inputs!$B$5", "Coll_Downside": "Inputs!$B$4"}})
        self.assertEqual(reviewed_selector.evaluate_cell("Cash", "A1"), 25)
        invalid_selector = WorkbookContext.from_path(path, indirect_bindings={
            ("Cash", "A1"): {"Coll_Base": "=1+1"}})
        with self.assertRaises(UnsupportedExpression):
            invalid_selector.bind_formula(result["A1"].value, sheet="Cash", address="A1")

    def test_new_bound_nodes_fail_explicitly_in_old_sigma_emitter(self):
        context = WorkbookContext.from_path(self.save(self.table_book()))
        bound = context.bind_formula("=tblOrders[Revenue]", sheet="FY Board", address="C1")
        with self.assertRaises(UnsupportedExpression):
            emit_sigma(bound, lambda ref: ref.address)
        with self.assertRaises(UnsupportedExpression):
            parse_expression("=tblOrders[Revenue]", sheet="FY Board")
        with self.assertRaises(UnsupportedExpression):
            tuple(_expression_cells(bound))
        with self.assertRaises(UnsupportedExpression):
            tuple(_walk_references(bound))
        with self.assertRaises(UnsupportedExpression):
            extract_references(parse_expression("=tblOrders[Revenue]", sheet="FY Board", workbook_mode=True))
        with self.assertRaises(UnsupportedExpression):
            extract_references(bound)
        with self.assertRaises(UnsupportedExpression):
            emit_sigma(parse_expression('="A"&"B"', workbook_mode=True), lambda ref: ref.address)
        with self.assertRaises(UnsupportedExpression):
            emit_sigma(parse_expression("='[FY26 Corporate Allocations.xlsx]Alloc'!C4", workbook_mode=True),
                       lambda ref: ref.address)

    def test_only_known_xlfn_prefixes_are_normalized(self):
        self.assertEqual(parse_expression("=_xlfn.XLOOKUP(1,A1:A2,B1:B2)", workbook_mode=True).name,
                         "XLOOKUP")
        unknown = parse_expression("=_xlfn.MYSTERY(1)", workbook_mode=True)
        self.assertIsInstance(unknown, Function)
        self.assertEqual(unknown.name, "_XLFN.MYSTERY")

    def test_blank_end_date_and_index_zero_column_preserve_row_calculation(self):
        book = Workbook()
        roster = book.active
        roster.title = "Team"
        roster.append(["Employee", "End Date", "Pay"])
        roster.append(["A", None, 120])
        roster.append(["B", datetime(2026, 1, 1), 240])
        roster.add_table(Table(displayName="tblTeam", ref="A1:C3"))
        result = book.create_sheet("Plan")
        result["A1"] = '=IF(OR(INDEX(tblTeam[End Date],1)="",INDEX(tblTeam[End Date],1)>=DATE(2026,2,1)),INDEX(tblTeam[Pay],1)/12,0)'
        result["A2"] = '=SUM(INDEX(tblTeam[[Employee]:[Pay]],1,0))'
        context = WorkbookContext.from_path(self.save(book))
        self.assertEqual(context.evaluate_cell("Plan", "A1"), 10)
        self.assertEqual(context.evaluate_cell("Plan", "A2"), 120)

    def test_scoped_names_and_1904_date_epoch_are_bound_at_use_site(self):
        book = Workbook()
        book.epoch = CALENDAR_MAC_1904
        one = book.active
        one.title = "East"
        two = book.create_sheet("West")
        one["A1"] = datetime(2026, 1, 1)
        one["B1"] = '=DATE(2026,2,1)-A1'
        one["B2"] = '=MONTH(EDATE(A1,1))'
        one["B3"] = '=EOMONTH(A1,1)-DATE(2026,2,1)'
        one["C1"] = '=Rate*100'
        two["C1"] = '=Rate*100'
        book.defined_names.add(DefinedName("Rate", attr_text="0.2"))
        one.defined_names.add(DefinedName("Rate", attr_text="0.3"))
        context = WorkbookContext.from_path(self.save(book))
        self.assertEqual(context.evaluate_cell("East", "B1"), 31)
        self.assertEqual(context.evaluate_cell("East", "B2"), 2)
        self.assertEqual(context.evaluate_cell("East", "B3"), 27)
        self.assertEqual(context.evaluate_cell("East", "C1"), 30)
        self.assertEqual(context.evaluate_cell("West", "C1"), 20)

    def test_iferror_catches_calculation_errors_but_not_missing_bindings(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Review"
        sheet["A1"] = "#DIV/0!"
        sheet["A1"].data_type = "e"
        sheet["B1"] = '=IFERROR(A1,7)'
        sheet["B2"] = '=IFERROR(UnknownName,7)'
        sheet["B3"] = '=IF(FALSE,1/0,9)'
        sheet["B4"] = '=OFFSET(A1,1,0)'
        context = WorkbookContext.from_path(self.save(book))
        self.assertEqual(context.evaluate_cell("Review", "B1"), 7)
        self.assertEqual(context.evaluate_cell("Review", "B3"), 9)
        report = context.formula_report()
        self.assertEqual(report["formulas"]["Review!B2"]["status"], "unresolved")
        self.assertEqual(report["formulas"]["Review!B2"]["dependencyStatus"], "incomplete")
        self.assertEqual(report["formulas"]["Review!B1"]["dependencyStatus"], "complete")
        self.assertEqual(report["formulas"]["Review!B4"]["dependencyStatus"], "incomplete")

    def test_table_data_and_totals_selectors_keep_separate_rows(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Actuals"
        sheet.append(["Category", "Value"])
        sheet.append(["A", 4])
        sheet.append(["B", 6])
        sheet.append(["Total", 10])
        sheet.add_table(Table(displayName="tblActuals", ref="A1:B4", totalsRowCount=1))
        report = book.create_sheet("Analysis")
        report["A1"] = '=SUM(tblActuals[[#Data],[Value]])'
        report["A2"] = '=SUM(tblActuals[[#Totals],[Value]])'
        report["A3"] = '=SUM(tblActuals[[#All],[Value]])'
        context = WorkbookContext.from_path(self.save(book))
        self.assertEqual([context.evaluate_cell("Analysis", f"A{row}") for row in range(1,4)],
                         [10, 10, 20])

    def test_existing_acceptance_fixture_count_and_roundup_functions(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Inputs"
        sheet["A1"] = 1.21
        sheet["A2"] = "text"
        sheet["B1"] = '=ROUNDUP(-A1,1)'
        sheet["B2"] = '=COUNT(A1:A3)'
        context = WorkbookContext.from_path(self.save(book))
        self.assertEqual(context.evaluate_cell("Inputs", "B1"), -1.3)
        self.assertEqual(context.evaluate_cell("Inputs", "B2"), 1)

    def test_counta_counts_empty_text_but_not_source_blank(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Counts"
        sheet["A1"] = ""
        sheet["A2"] = None
        sheet["B1"] = '=COUNTA(A1:A2)'
        sheet["B2"] = '=A1+1'
        sheet["B3"] = '=A2+1'
        context = WorkbookContext.from_path(self.save(book))
        self.assertEqual(context.evaluate_cell("Counts", "B1"), 1)
        self.assertEqual(context.formula_report()["formulas"]["Counts!B2"]["status"], "error")
        self.assertEqual(context.evaluate_cell("Counts", "B3"), 1)


if __name__ == "__main__":
    unittest.main()
