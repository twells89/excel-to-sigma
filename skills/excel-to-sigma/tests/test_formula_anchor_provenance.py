"""Formula anchors must follow OOXML cell type, including formula objects."""

import copy
from contextlib import redirect_stdout
import importlib.util
import io
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.table import Table


ROOT = Path(__file__).resolve().parents[1]
LIB = ROOT / "scripts" / "lib"
sys.path.insert(0, str(LIB))
sys.path.insert(0, str(ROOT / "scripts"))

from formula_graph import CellIdentity, FormulaGraph  # noqa: E402
from migration_manifest import (ManifestError, build_manifest, validate_completion,
                                validate_manifest)  # noqa: E402
from source_workbook import build_source_model  # noqa: E402

spec = importlib.util.spec_from_file_location("xlsx_discover", ROOT / "scripts" / "xlsx-discover.py")
discover = importlib.util.module_from_spec(spec)
spec.loader.exec_module(discover)


class TestFormulaAnchorProvenance(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.path = Path(temporary.name) / "formula-anchors.xlsx"
        book = Workbook()
        table = book.active
        table.title = "TableSource"
        table.append(["Text", "Array", "Scalar", "Data Table"])
        table["A2"] = "=literal source text"
        table["A2"].data_type = "s"
        table["B2"] = ArrayFormula(ref="B2:B3", text="=C2:C3*2")
        table["C2"] = "=1+1"
        table["D2"] = DataTableFormula(ref="D2:E3", dt2D=True, r1="A2", r2="C2")
        table.add_table(Table(displayName="Inputs", ref="A1:D3"))
        report = book.create_sheet("Report")
        report["A1"] = "=literal report text"
        report["A1"].data_type = "s"
        report["B1"] = ArrayFormula(ref="B1:B2", text="=TableSource!C2:C3")
        report["C1"] = "=TableSource!C2"
        book.save(self.path)
        book.close()

    def test_graph_and_manifest_account_for_every_formula_anchor(self):
        book = load_workbook(self.path, data_only=False)
        graph = FormulaGraph.from_workbook(book)
        model = build_source_model(self.path)
        expected = {"TableSource!B2", "TableSource!C2", "TableSource!D2",
                    "Report!B1", "Report!C1"}
        self.assertEqual({cell.id for cell in graph.formulas}, expected)
        self.assertEqual({CellIdentity("TableSource", "A2"), CellIdentity("Report", "A1")}
                         & set(graph.literals),
                         {CellIdentity("TableSource", "A2"), CellIdentity("Report", "A1")})
        self.assertEqual(model["counts"]["formulas"], len(expected))
        self.assertEqual(graph.literals[CellIdentity("TableSource", "A2")],
                         "=literal source text")
        for cell in ("TableSource!B2", "TableSource!D2", "Report!B1"):
            identity = next(key for key in graph.formulas if key.id == cell)
            self.assertIn("requires review", graph.unsupported_formulas[identity])
            self.assertNotIn(identity, graph.evaluate())
        book.close()

        manifest = build_manifest(self.path)
        self.assertEqual({cell["id"] for cell in manifest["formulaCells"]}, expected)
        self.assertEqual({cell["id"] for cell in manifest["literalInputs"]}
                         & {"TableSource!A2", "Report!A1"},
                         {"TableSource!A2", "Report!A1"})
        for cell in manifest["formulaCells"]:
            if cell["id"] in {"TableSource!B2", "TableSource!D2", "Report!B1"}:
                self.assertEqual(cell["mappingStatus"], "unresolved")
                self.assertIn("requires review", cell["unresolvedReason"])
                self.assertIsInstance(cell["originalFormula"], dict)
        array = next(cell for cell in manifest["formulaCells"] if cell["id"] == "TableSource!B2")
        self.assertEqual(array["originalFormula"],
                         {"objectType": "ArrayFormula", "attributes":
                          {"ref": "B2:B3", "text": "=C2:C3*2"}})
        data_table = next(cell for cell in manifest["formulaCells"] if cell["id"] == "TableSource!D2")
        self.assertEqual(data_table["originalFormula"]["objectType"], "DataTableFormula")
        self.assertEqual(data_table["originalFormula"]["attributes"]["r1"], "A2")
        self.assertEqual(next(cell for cell in manifest["formulaCells"]
                              if cell["id"] == "TableSource!C2")["originalFormula"], "=1+1")
        self.assertTrue(validate_manifest(manifest, self.path))
        with self.assertRaises(ManifestError):
            validate_completion(manifest)

    def test_manifest_rejects_missing_or_reclassified_object_anchor(self):
        manifest = build_manifest(self.path)
        missing = copy.deepcopy(manifest)
        missing["formulaCells"] = [cell for cell in missing["formulaCells"]
                                   if cell["id"] != "TableSource!B2"]
        from migration_manifest import _plan_hash
        missing["planHash"] = _plan_hash(missing["sourceHash"], missing["planProjection"],
                                          missing["formulaCells"])
        with self.assertRaises(ManifestError):
            validate_manifest(missing, self.path)
        promoted = copy.deepcopy(manifest)
        item = next(cell for cell in promoted["formulaCells"] if cell["id"] == "TableSource!B2")
        item["mappingStatus"] = "live"
        # A self-consistent candidate hash cannot make an unsupported object executable.
        promoted["planHash"] = _plan_hash(promoted["sourceHash"], promoted["planProjection"],
                                           promoted["formulaCells"])
        with self.assertRaisesRegex(ManifestError, "formula object"):
            validate_manifest(promoted, self.path)

    def test_discovery_uses_cell_type_for_table_and_report(self):
        book = load_workbook(self.path, data_only=False)
        table = book["TableSource"]
        analysis = discover.analyze_table([cell.value for cell in table[1]],
                                          list(table.iter_rows(min_row=2, max_row=3,
                                                               min_col=1, max_col=4)),
                                          2, False)
        columns = analysis["columns"]
        self.assertEqual(columns[0]["literalCellRefs"], ["A2"])
        self.assertEqual(columns[1]["formulaCellRefs"], ["B2"])
        self.assertEqual(columns[2]["formulaCellRefs"], ["C2"])
        self.assertEqual(columns[3]["formulaCellRefs"], ["D2"])
        self.assertEqual(analysis["route"], "MIXED")
        report, _, _ = discover.classify_report_cells(book)
        self.assertEqual(report["Report"]["OTHER_FORMULA"], 1)
        self.assertEqual(report["Report"]["DEPENDENCY_LINK"], 1)
        book.close()
        output = io.StringIO()
        with redirect_stdout(output):
            discover.main(str(self.path))
        self.assertIn("5 formula cells", output.getvalue())


if __name__ == "__main__":
    unittest.main()
