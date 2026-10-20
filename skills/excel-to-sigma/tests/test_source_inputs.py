"""Generated workbook contracts for ordinary source inputs and explicit reshapes."""
from __future__ import annotations

import copy
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "scripts"), str(ROOT / "scripts/lib")]

from input_sync import emit_construction_program, emit_write_programs
from migration_manifest import build_manifest
from source_inputs import build_source_input_plan, plan_source_seed, validate_source_input_plan
from source_reshape import unpivot_source_table
from input_orchestration import gate_input_staging
from migration_gate import gate_plan_staging_source


class SourceFixture(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales Input"
        ws.append(["Business ID", "When", "Enabled", "Entered", "Computed"])
        ws.append(["A", "2026-01-01", True, 4, "=D2*2"])
        ws.append(["B", "2026-02-01", False, 7, None])
        ws.add_table(Table(displayName="tblSales", ref="A1:E3"))
        assumptions = wb.create_sheet("Assumptions")
        assumptions["B2"] = "Rate"
        assumptions["C2"] = .25
        assumptions["D2"] = "=C2*2"
        self.source = self.folder / "fixture.xlsx"
        wb.save(self.source)
        wb.close()
        self.sha = hashlib.sha256(self.source.read_bytes()).hexdigest()

    def decisions(self):
        return {"sourceSha256": self.sha, "features": [{
            "featureId": "sales-entry", "sheet": "Sales Input", "table": "tblSales",
            "role": "input", "reviewed": True, "keyColumns": ["Business ID"],
            "literalColumns": ["Business ID", "When", "Enabled", "Entered"],
            "columnTypes": {"Business ID": "text", "When": "datetime",
                            "Enabled": "boolean", "Entered": "number"},
        }]}

    def build(self, decisions=None):
        return build_source_input_plan(self.source, decisions or self.decisions(),
                                       connection_id="connection", folder_id="folder")


class TestSourceInputs(SourceFixture):
    def test_mixed_table_preserves_literals_and_formula_obligations(self):
        plan = self.build()
        validate_source_input_plan(plan, self.source)
        spec = plan["sharedInputArchitecture"]["inputs"]["sales-entry"]
        self.assertEqual(spec["name"], "tblSales")
        self.assertEqual(spec["rows"], [
            {"Business ID": "A", "When": "2026-01-01", "Enabled": True, "Entered": 4},
            {"Business ID": "B", "When": "2026-02-01", "Enabled": False, "Entered": 7},
        ])
        self.assertNotIn("Computed", spec["rows"][0])
        self.assertEqual(spec["sourceCells"]["0"]["Entered"], "Sales Input!D2")
        obligations = {item["sourceCellId"]: item for item in plan["formulaObligations"]}
        self.assertEqual(obligations["Sales Input!E2"]["sourceColumn"], "Computed")
        self.assertIn("Assumptions!D2", obligations)
        self.assertEqual(plan["status"], "REVIEWED_PLAN_PENDING_STAGING")
        self.assertEqual(plan, self.build())

    def test_source_hash_and_typed_keys_fail_closed(self):
        bad = self.decisions(); bad["sourceSha256"] = "0" * 64
        with self.assertRaisesRegex(ValueError, "hash"):
            self.build(bad)
        bad = self.decisions(); bad["features"][0]["keyColumns"] = ["Missing"]
        with self.assertRaises(ValueError):
            self.build(bad)
        from openpyxl import load_workbook
        wb = load_workbook(self.source); wb["Sales Input"]["A3"] = "A"; wb.save(self.source); wb.close()
        self.sha = hashlib.sha256(self.source.read_bytes()).hexdigest()
        bad = self.decisions()
        with self.assertRaisesRegex(ValueError, "duplicate"):
            self.build(bad)
        bad = self.decisions(); bad["features"][0]["columnTypes"]["Business ID"] = "number"
        with self.assertRaisesRegex(ValueError, "type"):
            self.build(bad)

    def test_formula_cache_request_and_unreviewed_blank_type_fail(self):
        bad = self.decisions(); bad["features"][0]["literalColumns"].append("Computed")
        bad["features"][0]["columnTypes"]["Computed"] = "number"
        with self.assertRaisesRegex(ValueError, "formula"):
            self.build(bad)
        bad = self.decisions(); bad["features"][0]["literalColumns"] = ["Business ID", "Computed"]
        bad["features"][0]["columnTypes"] = {"Business ID": "text"}
        with self.assertRaises(ValueError):
            self.build(bad)

    def test_reviewed_literal_override_in_computed_column_is_editable_only_at_named_cell(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.source); wb["Sales Input"]["E3"] = 9; wb.save(self.source); wb.close()
        self.sha = hashlib.sha256(self.source.read_bytes()).hexdigest()
        decisions = self.decisions()
        feature = decisions["features"][0]
        feature["literalColumns"].append("Computed")
        feature["columnTypes"]["Computed"] = "number"
        with self.assertRaisesRegex(ValueError, "formula|override"):
            self.build(decisions)
        feature["overrideCells"] = {"Computed": ["E3"]}
        plan = self.build(decisions)
        rows = plan["sharedInputArchitecture"]["inputs"]["sales-entry"]["rows"]
        self.assertNotIn("Computed", rows[0])
        self.assertNotIn("Computed", rows[1])
        main = plan["sharedInputArchitecture"]["inputs"]["sales-entry"]
        self.assertIn("Computed", main["derivedColumns"])
        self.assertNotIn("Computed", [column["name"] for column in main["columns"]])
        override_feature = main["overrideInputFeatures"]["Computed"]
        override = plan["sharedInputArchitecture"]["inputs"][override_feature]
        self.assertEqual(override["rows"], [{"Business ID": "B", "Computed": 9}])
        self.assertEqual(override["sourceCells"]["0"]["Computed"], "Sales Input!E3")
        self.assertEqual(len(plan_source_seed(plan, override_feature, [], row_id_column="Row ID")["inserts"]), 1)
        self.assertIn("Sales Input!E2", {item["sourceCellId"] for item in plan["formulaObligations"]})
        from input_sync import build_artifact_records
        manifest = build_manifest(self.source, plan)
        architecture = plan["sharedInputArchitecture"]
        bindings = {name: {"documentId": "isolated", "version": 1, "pageId": "inputs"}
                    for name in architecture["inputs"]}
        receipt = {"status": "APPROVED_STAGING", "sourceHash": manifest["sourceHash"],
                   "planHash": manifest["planHash"], "targetBindings": bindings,
                   "artifacts": build_artifact_records(manifest, architecture, bindings)}
        main_program = emit_construction_program(manifest, receipt, "sales-entry")
        override_program = emit_construction_program(manifest, receipt, override_feature)
        self.assertNotIn("add_editable_column('Computed'", main_program)
        self.assertIn("add_editable_column('Computed'", override_program)

    def test_excel_error_cells_are_never_entered_literals(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.source)
        cell = wb["Sales Input"]["D2"]
        cell.value = "#DIV/0!"; cell.data_type = "e"
        wb["Sales Input"]["D3"] = "seven"
        wb.save(self.source); wb.close()
        self.sha = hashlib.sha256(self.source.read_bytes()).hexdigest()
        decisions = self.decisions()
        decisions["features"][0]["columnTypes"]["Entered"] = "text"
        with self.assertRaisesRegex(ValueError, "error|unsupported"):
            self.build(decisions)

    def test_unpivot_rejects_error_key_and_measure(self):
        from openpyxl import load_workbook
        wb = Workbook(); ws = wb.active; ws.title = "Data"
        ws.append(["ID", "A", "B"]); ws.append(["x", 1, 2])
        ws.add_table(Table(displayName="tblWide", ref="A1:C2"))
        source = self.folder / "errors.xlsx"; wb.save(source); wb.close()
        options = dict(sheet="Data", table="tblWide", key_columns=["ID"],
                       value_columns=["A", "B"], axis_name="Kind", value_name="Value")
        for address in ("A2", "B2"):
            wb = load_workbook(source); cell = wb["Data"][address]
            cell.value = "#N/A"; cell.data_type = "e"; wb.save(source); wb.close()
            with self.subTest(address=address), self.assertRaisesRegex(ValueError, "error|unsupported"):
                unpivot_source_table(source, **options)
            wb = load_workbook(source); wb["Data"][address] = "x" if address == "A2" else 1
            wb.save(source); wb.close()

    def test_source_change_updates_physical_id_and_user_edit_conflicts(self):
        baseline = self.build()
        first = plan_source_seed(baseline, "sales-entry", [], row_id_column="Row ID")
        current = [{**row, "Row ID": f"physical-{index}"} for index, row in
                   enumerate(baseline["sharedInputArchitecture"]["inputs"]["sales-entry"]["rows"])]
        confirmed = plan_source_seed(baseline, "sales-entry", current,
                                     row_id_column="Row ID", prior_seed=first)
        from openpyxl import load_workbook
        wb = load_workbook(self.source); wb["Sales Input"]["D2"] = 5; wb.save(self.source); wb.close()
        self.sha = hashlib.sha256(self.source.read_bytes()).hexdigest()
        updated = self.build()
        safe = plan_source_seed(updated, "sales-entry", current,
                                row_id_column="Row ID", prior_seed=confirmed)
        self.assertEqual(safe["updates"][0]["rowId"], "physical-0")
        self.assertEqual(safe["updates"][0]["changes"], {"Entered": 5})
        current[0]["Entered"] = 8
        conflict = plan_source_seed(updated, "sales-entry", current,
                                    row_id_column="Row ID", prior_seed=confirmed)
        self.assertEqual(conflict["updates"], [])
        self.assertEqual(conflict["conflicts"][0]["columns"], ["Entered"])

    def test_blank_table_schema_requires_reviewed_types_and_has_no_seed_rows(self):
        wb = Workbook(); ws = wb.active; ws.title = "Snapshots"
        ws.append(["SKU", "Committed At", "Forecast Units"])
        ws.append([None, None, None])
        ws.add_table(Table(displayName="tblSnapshots", ref="A1:C2"))
        source = self.folder / "blank.xlsx"; wb.save(source); wb.close()
        sha = hashlib.sha256(source.read_bytes()).hexdigest()
        decision = {"sourceSha256": sha, "features": [{
            "featureId": "snapshots", "sheet": "Snapshots", "table": "tblSnapshots",
            "role": "input", "reviewed": True, "keyColumns": ["SKU"],
            "literalColumns": ["SKU", "Committed At", "Forecast Units"],
            "columnTypes": {"SKU": "text", "Committed At": "datetime", "Forecast Units": "number"},
            "emptySchema": True,
        }]}
        with self.assertRaisesRegex(ValueError, "review"):
            build_source_input_plan(source, decision, connection_id="c", folder_id="f")
        decision["features"][0]["reviewedBlankColumns"] = ["SKU", "Committed At", "Forecast Units"]
        decision["features"][0]["typeReviewEvidence"] = "Reviewed VBA snapshot append assignments"
        plan = build_source_input_plan(source, decision, connection_id="c", folder_id="f")
        spec = plan["sharedInputArchitecture"]["inputs"]["snapshots"]
        self.assertEqual(spec["rows"], [])
        self.assertEqual([c["type"] for c in spec["columns"]], ["text", "datetime", "number"])

    def test_reviewed_supplied_replacement_is_source_formula_bound(self):
        decision = {"sourceSha256": self.sha, "features": [{
            "featureId": "corporate-allocations", "role": "replacement", "reviewed": True,
            "displayName": "Corporate Allocations", "replacementFor": "external rate",
            "keyColumns": ["Period"], "literalColumns": ["Period", "Allocation"],
            "columnTypes": {"Period": "text", "Allocation": "number"},
            "suppliedRows": [{"Period": "FY26-Q3", "Allocation": .5}],
            "replacementBindings": [{"key": {"Period": "FY26-Q3"},
                                     "sourceCells": {"Period": "Assumptions!B2", "Allocation": "Assumptions!D2"},
                                     "sourceLabel": "Rate", "originalFormula": "=C2*2"}],
            "replacementReview": {"approved": True, "basis": "User-approved editable value"},
        }]}
        plan = self.build(decision)
        spec = plan["sharedInputArchitecture"]["inputs"]["corporate-allocations"]
        self.assertEqual(spec["rows"], [{"Period": "FY26-Q3", "Allocation": .5}])
        self.assertEqual(spec["sourceCells"]["0"]["Allocation"], "Assumptions!D2")
        self.assertEqual(spec["replacementBindings"][0]["originalFormula"], "=C2*2")
        self.assertEqual(plan_source_seed(plan, "corporate-allocations", [],
                                          row_id_column="Row ID")["inserts"][0]["source"]["Allocation"], .5)
        broken = copy.deepcopy(decision)
        broken["features"][0]["replacementBindings"][0]["originalFormula"] = "=wrong"
        with self.assertRaisesRegex(ValueError, "formula"):
            self.build(broken)
        broken = copy.deepcopy(decision)
        broken["features"][0]["replacementReview"]["approved"] = False
        with self.assertRaisesRegex(ValueError, "review"):
            self.build(broken)

    def test_target_identifiers_are_safe_and_distinct_for_similar_feature_names(self):
        decisions = self.decisions()
        decisions["features"][0]["featureId"] = "sales-input"
        second = copy.deepcopy(decisions["features"][0]); second["featureId"] = "sales_input"
        decisions["features"].append(second)
        plan = self.build(decisions)
        ids = [item["id"] for item in plan["sharedInputArchitecture"]["nativeWiring"]["elements"]]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertTrue(all(value.replace("_", "").isalnum() for value in ids))

    def test_explicit_empty_string_and_boolean_are_preserved(self):
        wb = Workbook(); ws = wb.active; ws.title = "Entries"
        ws.append(["Key", "Note", "Enabled"]); ws.append(["A", "", False])
        ws.add_table(Table(displayName="tblEntries", ref="A1:C2"))
        source = self.folder / "typed.xlsx"; wb.save(source); wb.close()
        decisions = {"sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                     "features": [{"featureId": "entries", "sheet": "Entries", "table": "tblEntries",
                                   "role": "input", "reviewed": True, "keyColumns": ["Key"],
                                   "literalColumns": ["Key", "Note", "Enabled"],
                                   "columnTypes": {"Key": "text", "Note": "text", "Enabled": "boolean"}}]}
        plan = build_source_input_plan(source, decisions, connection_id="c", folder_id="f")
        self.assertEqual(plan["sharedInputArchitecture"]["inputs"]["entries"]["rows"][0],
                         {"Key": "A", "Note": "", "Enabled": False})

    def test_datetime_type_rejects_invalid_calendar_text(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.source); wb["Sales Input"]["B2"] = "2026-13-01"
        wb.save(self.source); wb.close()
        self.sha = hashlib.sha256(self.source.read_bytes()).hexdigest()
        with self.assertRaisesRegex(ValueError, "type"):
            self.build()

    def test_queried_datetime_key_matches_source_date_only_grain(self):
        from datetime import datetime
        decisions = self.decisions()
        decisions["features"][0]["keyColumns"] = ["Business ID", "When"]
        plan = self.build(decisions)
        current = [{"Business ID": "A", "When": datetime(2026, 1, 1, 12),
                    "Enabled": True, "Entered": 4, "Row ID": "physical-A"},
                   {"Business ID": "B", "When": datetime(2026, 2, 1, 12),
                    "Enabled": False, "Entered": 7, "Row ID": "physical-B"}]
        seed = plan_source_seed(plan, "sales-entry", current, row_id_column="Row ID")
        self.assertEqual(seed["inserts"], [])
        self.assertEqual(seed["orphans"], [])
        self.assertEqual(len(seed["unchanged"]), 2)

    def test_cell_block_assumption_and_seed_resume_user_edits(self):
        decisions = self.decisions()
        decisions["features"] = [{"featureId": "rate-input", "sheet": "Assumptions",
                                  "range": "B2:C2", "role": "input", "reviewed": True,
                                  "keyColumns": ["Setting"], "literalColumns": ["Setting", "Value"],
                                  "headers": ["Setting", "Value"],
                                  "columnTypes": {"Setting": "text", "Value": "number"}}]
        plan = self.build(decisions)
        spec = plan["sharedInputArchitecture"]["inputs"]["rate-input"]
        self.assertEqual(spec["rows"], [{"Setting": "Rate", "Value": .25}])
        self.assertEqual(spec["sourceCells"]["0"]["Value"], "Assumptions!C2")
        first = plan_source_seed(plan, "rate-input", [], row_id_column="Row ID")
        self.assertEqual(len(first["inserts"]), 1)
        current = [{"Setting": "Rate", "Value": .25, "Row ID": "physical-1"}]
        resumed = plan_source_seed(plan, "rate-input", current,
                                   row_id_column="Row ID", prior_seed=first)
        self.assertEqual(resumed["inserts"], [])
        self.assertEqual(resumed["updates"], [])
        altered = copy.deepcopy(plan)
        altered["sharedInputArchitecture"]["inputs"]["rate-input"]["rows"][0]["Value"] = .3
        # A modified plan must be rebuilt and rebound to source before planning writes.
        with self.assertRaises(ValueError):
            plan_source_seed(altered, "rate-input", current,
                             row_id_column="Row ID", prior_seed=resumed)

    def test_existing_emitter_requires_staging_receipt(self):
        plan = self.build()
        manifest = build_manifest(self.source, plan)
        feature = "sales-entry"
        with self.assertRaisesRegex(ValueError, "APPROVED_STAGING"):
            emit_construction_program(manifest, {}, feature)
        with self.assertRaisesRegex(ValueError, "APPROVED_STAGING"):
            emit_write_programs(plan_source_seed(plan, feature, [], row_id_column="Row ID"),
                                manifest, {}, feature)

    def test_existing_constructor_and_writer_accept_approved_artifact(self):
        from input_sync import build_artifact_records
        plan = self.build(); manifest = build_manifest(self.source, plan)
        architecture = plan["sharedInputArchitecture"]
        binding = {"documentId": "isolated", "version": 1, "pageId": "input-page",
                   "elementId": "assigned-table", "columnIds": {
                       column["name"]: f"assigned-{index}" for index, column in
                       enumerate(architecture["inputs"]["sales-entry"]["columns"])}}
        receipt = {"status": "APPROVED_STAGING", "sourceHash": manifest["sourceHash"],
                   "planHash": manifest["planHash"], "targetBindings": {"sales-entry": binding},
                   "artifacts": build_artifact_records(manifest, architecture, {"sales-entry": binding})}
        program = emit_construction_program(manifest, receipt, "sales-entry")
        self.assertIn("add_input_table", program)
        self.assertIn("upsert_rows", program)
        self.assertNotIn("Computed", program)
        wiring = architecture["nativeWiring"]["elements"][0]
        self.assertEqual(set(wiring["columns"][0]), {"id", "name", "type"})
        seed = plan_source_seed(plan, "sales-entry", [], row_id_column="Row ID")
        writes = emit_write_programs(seed, manifest, receipt, "sales-entry")
        self.assertEqual(len(writes), 1)
        self.assertIn("upsert_rows", writes[0]["program"])

    def test_cli_refuses_source_output_alias(self):
        decisions_path = self.folder / "decisions.json"
        decisions_path.write_text(json.dumps(self.decisions()))
        result = subprocess.run([sys.executable, str(ROOT / "scripts/source-inputs.py"), "plan",
                                 "--source", str(self.source), "--decisions", str(decisions_path),
                                 "--connection", "connection", "--folder", "folder",
                                 "--out", str(self.source)], capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.sha)


class TestReshape(SourceFixture):
    def test_generated_headcount_and_history_counts(self):
        from datetime import date
        from openpyxl import load_workbook
        wb = Workbook(); budget = wb.active; budget.title = "Budget"
        months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
        budget.append(["Department", *months])
        for index in range(5):
            budget.append([f"D{index}", *[100 + index * 12 + month for month in range(12)]])
        budget.add_table(Table(displayName="tblDeptBudget", ref="A1:M6"))
        history = wb.create_sheet("History")
        history.append(["Month", "A", "B", "C"])
        for month in range(36):
            history.append([date(2023 + month // 12, month % 12 + 1, 1), month, month + 1, month + 2])
        history.add_table(Table(displayName="tblHistory", ref="A1:D37"))
        source = self.folder / "full-shapes.xlsx"; wb.save(source); wb.close()
        budget_rows = unpivot_source_table(source, sheet="Budget", table="tblDeptBudget",
                                           key_columns=["Department"], value_columns=months,
                                           axis_name="Month", value_name="Amount",
                                           axis_values={month: f"2026-{index + 1:02d}-01" for index, month in enumerate(months)})
        self.assertEqual(len(budget_rows["rows"]), 60)
        self.assertEqual(budget_rows["sourceCells"]["59"]["Amount"], "Budget!M6")
        history_rows = unpivot_source_table(source, sheet="History", table="tblHistory",
                                            key_columns=["Month"], value_columns=["A", "B", "C"],
                                            axis_name="SKU", value_name="Units")
        self.assertEqual(len(history_rows["rows"]), 108)
        self.assertEqual(history_rows["sourceCells"]["107"]["Units"], "History!D37")

        decisions = {"sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest(), "features": [
            {"featureId": "budget-long", "role": "input", "reviewed": True,
             "keyColumns": ["Department", "Month"],
             "literalColumns": ["Department", "Month", "Amount"],
             "columnTypes": {"Department": "text", "Month": "datetime", "Amount": "number"},
             "reshape": {"kind": "unpivot", "sheet": "Budget", "table": "tblDeptBudget",
                         "keyColumns": ["Department"], "valueColumns": months,
                         "axisName": "Month", "valueName": "Amount",
                         "axisValues": {month: f"2026-{index + 1:02d}-01" for index, month in enumerate(months)}}},
            {"featureId": "history-long", "role": "input", "reviewed": True,
             "keyColumns": ["Month", "SKU"], "literalColumns": ["Month", "SKU", "Units"],
             "columnTypes": {"Month": "datetime", "SKU": "text", "Units": "number"},
             "reshape": {"kind": "unpivot", "sheet": "History", "table": "tblHistory",
                         "keyColumns": ["Month"], "valueColumns": ["A", "B", "C"],
                         "axisName": "SKU", "valueName": "Units"}},
        ]}
        plan = build_source_input_plan(source, decisions, connection_id="c", folder_id="f")
        validate_source_input_plan(plan, source)
        architecture = plan["sharedInputArchitecture"]
        self.assertEqual(len(architecture["inputs"]["budget-long"]["rows"]), 60)
        self.assertEqual(len(architecture["inputs"]["history-long"]["rows"]), 108)
        self.assertEqual(architecture["inputs"]["budget-long"]["sourceCells"]["59"]["Amount"], "Budget!M6")
        self.assertEqual(architecture["inputs"]["budget-long"]["sourceCells"]["59"]["Month"], "Budget!M1")
        self.assertEqual(len(plan_source_seed(plan, "budget-long", [], row_id_column="Row ID")["inserts"]), 60)
        manifest = build_manifest(source, plan)
        _source, source_evidence = gate_plan_staging_source(plan, source_arg=source)
        bindings = {name: {"documentId": "isolated", "version": 1, "pageId": "inputs",
                           "elementId": f"assigned-{name}", "columnIds": {
                               column["name"]: f"{name}-{index}" for index, column in enumerate(spec["columns"])}}
                    for name, spec in architecture["inputs"].items()}
        approval = {"sourceHash": manifest["sourceHash"], "planHash": manifest["planHash"],
                    "owner": "Fixture Owner", "reviewer": "Fixture Reviewer", "reviewItems": [],
                    "mappedFeatures": architecture["mappedFeatures"], "targetBindings": bindings}
        if source_evidence.get("issues"):
            approval["sourceReviewApproval"] = {"status": "approved_for_staging", "issues": source_evidence["issues"]}
        receipt = gate_input_staging(manifest, approval, plan, source_evidence)
        self.assertEqual(receipt["status"], "APPROVED_STAGING")
        program = emit_construction_program(manifest, receipt, "budget-long")
        self.assertIn("upsert_rows", program)
        self.assertIn("2026", program)
        seed = plan_source_seed(plan, "history-long", [], row_id_column="Row ID")
        self.assertEqual(len(emit_write_programs(seed, manifest, receipt, "history-long")), 1)

    def test_explicit_month_mapping_and_reverse_provenance(self):
        wb = Workbook(); ws = wb.active; ws.title = "Budget"
        ws.append(["Department", "Jan", "Feb"])
        ws.append(["Ops", 10, 20]); ws.append(["Sales", 30, 40])
        ws.add_table(Table(displayName="tblBudget", ref="A1:C3"))
        source = self.folder / "budget.xlsx"; wb.save(source); wb.close()
        result = unpivot_source_table(source, sheet="Budget", table="tblBudget",
                                      key_columns=["Department"], value_columns=["Jan", "Feb"],
                                      axis_name="Month", value_name="Amount",
                                      axis_values={"Jan": "2026-01-01", "Feb": "2026-02-01"})
        self.assertEqual(result["rows"], [
            {"Department": "Ops", "Month": "2026-01-01", "Amount": 10},
            {"Department": "Ops", "Month": "2026-02-01", "Amount": 20},
            {"Department": "Sales", "Month": "2026-01-01", "Amount": 30},
            {"Department": "Sales", "Month": "2026-02-01", "Amount": 40},
        ])
        self.assertEqual(result["sourceCells"]["0"]["Amount"], "Budget!B2")
        self.assertEqual(result["keyColumns"], ["Department", "Month"])
        self.assertEqual(result["columnTypes"], {"Department": "text", "Month": "datetime", "Amount": "number"})
        with self.assertRaisesRegex(ValueError, "axis"):
            unpivot_source_table(source, sheet="Budget", table="tblBudget",
                                 key_columns=["Department"], value_columns=["Jan", "Feb"],
                                 axis_name="Month", value_name="Amount")

    def test_conflicting_axis_and_overlapping_columns_fail(self):
        wb = Workbook(); ws = wb.active; ws.title = "History"
        ws.append(["Month", "SKU A", "SKU B"]); ws.append(["2026-01-01", 1, 2])
        ws.add_table(Table(displayName="tblHistory", ref="A1:C2"))
        source = self.folder / "history.xlsx"; wb.save(source); wb.close()
        options = dict(sheet="History", table="tblHistory", key_columns=["Month"],
                       value_columns=["SKU A", "SKU B"], axis_name="SKU", value_name="Units")
        result = unpivot_source_table(source, **options)
        self.assertEqual(result["rows"][1], {"Month": "2026-01-01", "SKU": "SKU B", "Units": 2})
        self.assertEqual(result["columnTypes"]["SKU"], "text")
        with self.assertRaisesRegex(ValueError, "overlap"):
            unpivot_source_table(source, **{**options, "value_columns": ["Month"]})
        with self.assertRaisesRegex(ValueError, "duplicate"):
            unpivot_source_table(source, **{**options, "axis_values": {"SKU A": "X", "SKU B": "X"}})
