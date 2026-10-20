"""Regression tests for source roles inferred before any Sigma build."""
from __future__ import annotations

import importlib.util
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "xlsx-discover.py"
spec = importlib.util.spec_from_file_location("xlsx_discover", SCRIPT)
discover = importlib.util.module_from_spec(spec)
spec.loader.exec_module(discover)


def inspect(wb):
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "scratch.xlsx"
        wb.save(path)
        output = io.StringIO()
        with redirect_stdout(output):
            discover.main(str(path), json_output=True)
        return json.loads(output.getvalue())


def pipeline_fixture(sheet="Pipeline", table="tblPipeline"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["Opp ID", "Account", "Owner", "Stage", "Close Date", "Amount",
               "Probability", "Weighted Amount", "Forecast Category"])
    for number in range(1, 61):
        ws.append([f"OPP-{number}", "Acme", "Sam", "Discovery", "2026-07-01", 100,
                   f"=VLOOKUP(D{number + 1},Stages!A:C,2,FALSE)",
                   f"=F{number + 1}*G{number + 1}",
                   f"=VLOOKUP(D{number + 1},Stages!A:C,3,FALSE)"])
    ws.add_table(Table(displayName=table, ref="A1:I61"))
    return wb


class TestDiscoveryRouting(unittest.TestCase):
    def test_mixed_table_keeps_literal_columns_separate_from_formula_columns(self):
        # Catches a workbook-wide formula-ratio route replacing six editable sources with caches.
        table = inspect(pipeline_fixture())["tables"][0]
        self.assertEqual(table["route"], "MIXED")
        self.assertEqual([column["role"] for column in table["columns"]],
                         ["SOURCE"] * 6 + ["DERIVED"] * 3)
        self.assertEqual(table["columns"][5]["literalCells"], 60)
        self.assertEqual(table["columns"][6]["formulaCells"], 60)
        self.assertEqual(table["columns"][5]["literalCellRefs"][0], "F2")
        self.assertEqual(table["columns"][6]["formulaCellRefs"][0], "G2")
        self.assertIn("review", table["action"].lower())
        self.assertNotIn("cache", table["action"].lower())

    def test_partly_formula_column_requires_review(self):
        # Catches a majority-cell heuristic that silently assigns one column to a build route.
        wb = pipeline_fixture()
        wb.active["G2"] = 0.5
        table = inspect(wb)["tables"][0]
        self.assertEqual(table["columns"][6]["role"], "MIXED_REVIEW")
        self.assertEqual(table["columns"][6]["literalCells"], 1)
        self.assertEqual(table["columns"][6]["formulaCells"], 59)
        self.assertEqual(table["columns"][6]["literalCellRefs"], ["G2"])
        self.assertEqual(table["columns"][6]["formulaCellRefs"][0], "G3")

    def test_renamed_table_uses_same_column_roles(self):
        table = inspect(pipeline_fixture("Opportunities", "tblOpportunities"))["tables"][0]
        self.assertEqual(table["route"], "MIXED")
        self.assertEqual([c["role"] for c in table["columns"]], ["SOURCE"] * 6 + ["DERIVED"] * 3)

    def test_derived_only_formal_table_is_not_routed_as_editable(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Calc A", "Calc B"])
        for row in range(2, 5):
            ws.append([f"=Inputs!A{row}", f"=A{row}*2"])
        ws.add_table(Table(displayName="tblResults", ref="A1:B4"))
        table = inspect(wb)["tables"][0]
        self.assertEqual(table["route"], "DERIVED")
        self.assertEqual([c["role"] for c in table["columns"]], ["DERIVED", "DERIVED"])

    def test_blank_template_requires_review_before_routing(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["SKU", "Committed At"])
        ws.append([None, None])
        ws.add_table(Table(displayName="tblSnapshots", ref="A1:B2"))
        table = inspect(wb)["tables"][0]
        self.assertEqual(table["route"], "REVIEW")
        self.assertEqual([c["role"] for c in table["columns"]], ["EMPTY_REVIEW", "EMPTY_REVIEW"])

    def test_roll_forward_references_keep_cell_dependencies(self):
        # Catches classifying =prior-ending-cash as a frozen FLAT input.
        wb = Workbook()
        ws = wb.active
        ws.title = "13WCF"
        for index in range(3, 15):
            column = ws.cell(7, index)
            column.value = f"={ws.cell(23, index - 1).coordinate}"
        report = inspect(wb)["reportSheets"][0]
        self.assertEqual(report["classes"]["DEPENDENCY_LINK"], 12)
        self.assertEqual(report["dependencyLinks"][0], {"cell": "C7", "sourceSheet": "13WCF", "sourceCell": "B23"})
        self.assertEqual(report["dependencyLinks"][-1], {"cell": "N7", "sourceSheet": "13WCF", "sourceCell": "M23"})
        self.assertNotIn("FLAT", report["classes"])

    def test_renamed_sheet_and_cross_sheet_reference_keep_dependencies(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Liquidity"
        ws["C7"] = "=B23"
        ws["D7"] = "='Cash Inputs'!$B$6"
        report = inspect(wb)["reportSheets"][0]
        self.assertEqual(report["classes"]["DEPENDENCY_LINK"], 2)
        self.assertEqual(report["dependencyLinks"][1],
                         {"cell": "D7", "sourceSheet": "Cash Inputs", "sourceCell": "B6"})

    def test_sku_measure_axis_is_reviewed_without_calendar_headers(self):
        # Catches overlooking a wide measure axis when headers are SKU labels.
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Month", "SKU-A100", "SKU-B200", "SKU-C300"])
        for month in range(1, 13):
            ws.append([f"2026-{month:02d}-01", 100 + month, 200 + month, 300 + month])
        ws.add_table(Table(displayName="tblSalesHistory", ref="A1:D13"))
        table = inspect(wb)["tables"][0]
        self.assertEqual(table["wideAxis"], "POSSIBLE_MEASURE_AXIS")
        self.assertIn("review", table["action"].lower())

    def test_report_without_drivers_does_not_claim_a_driver_forecast_architecture(self):
        # Catches attaching a P&L-specific architecture to an unrelated cash model.
        wb = Workbook()
        wb.active["C7"] = "=B23"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "cash.xlsx"
            wb.save(path)
            output = io.StringIO()
            with redirect_stdout(output):
                discover.main(str(path))
        self.assertIn("live dependency=1", output.getvalue())
        self.assertNotIn("Architecture: union", output.getvalue())
        self.assertNotIn("FLAT (carry)", output.getvalue())
        self.assertNotIn("SUMIFS/rollups become", output.getvalue())
        self.assertNotIn("driver-override", output.getvalue())

if __name__ == "__main__":
    unittest.main()
