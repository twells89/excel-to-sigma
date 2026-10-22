#!/usr/bin/env python3
"""scripts/migrate.py plan: archetype routing, BLOCKED exit code, review template."""
from __future__ import annotations

import contextlib
import importlib.util
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / "lib"))
sys.path.insert(0, str(ROOT / "tests"))

import forecast_pnl_fixture as fixture  # noqa: E402
from migration_gate import source_sha256  # noqa: E402

spec = importlib.util.spec_from_file_location("migrate_router", SCRIPTS / "migrate.py")
MIGRATE = importlib.util.module_from_spec(spec)
spec.loader.exec_module(MIGRATE)


class TestRouter(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name).resolve()
        self.out = self.dir / "work"

    def plan(self, path, *extra):
        stdout = io.StringIO()
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(io.StringIO()):
            code = MIGRATE.main(["plan", str(path), "--out-dir", str(self.out), *extra])
        return code, (json.loads(stdout.getvalue()) if stdout.getvalue() else None)

    def table_workbook(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Entries"
        ws.append(["Account", "Amount"])
        for index in range(1, 21):
            ws.append([f"A{index}", index * 10])
        ws.add_table(Table(displayName="tblEntries", ref="A1:B21"))
        path = self.dir / "table.xlsx"
        wb.save(path)
        return path

    def test_forecast_archetype_routes_to_drafter_then_builder(self):
        path = self.dir / "pnl.xlsx"
        fixture.build_pnl(path)
        code, result = self.plan(path)
        self.assertEqual(code, 0)
        self.assertEqual(result["archetype"], "forecast")
        self.assertEqual(result["auditStatus"], "AUTO")
        self.assertIn("refs/forecast-recipes.md", result["readRefs"])
        self.assertEqual(result["beforePost"], ["refs/coverage-contract.md"])
        self.assertNotIn("refs/coverage-contract.md", result["readRefs"])
        commands = result["nextCommands"]
        self.assertIn("draft-forecast-plan.py", commands[0])
        self.assertIn(str(self.out / "plan.json"), commands[0])
        self.assertIn("build-forecast-model.py", commands[1])
        self.assertNotIn("--post", commands[1])
        self.assertIn("--post --review-dispositions", commands[2])
        self.assertEqual(result["modelMap"]["P&L"]["ACTUAL"], 18)
        self.assertTrue((self.out / "migration-audit.json").is_file())
        self.assertTrue((self.out / "discovery.json").is_file())
        self.assertIsNone(result["reviewTemplate"])  # AUTO -> nothing to disposition
        self.assertFalse((self.out / "review.template.json").exists())

    def test_input_table_archetype(self):
        code, result = self.plan(self.table_workbook())
        self.assertEqual(code, 0)
        self.assertEqual(result["archetype"], "input-table")
        self.assertEqual(result["tables"], [{"name": "tblEntries", "sheet": "Entries", "rows": 20}])
        self.assertIn("refs/input-tables.md", result["readRefs"])
        self.assertIn("xlsx-to-input-csv.py", result["nextCommands"][0])

    def test_unknown_archetype_asks_the_user(self):
        wb = Workbook()
        wb.active["A1"] = "notes"
        path = self.dir / "unknown.xlsx"
        wb.save(path)
        code, result = self.plan(path)
        self.assertEqual(code, 0)
        self.assertEqual(result["archetype"], "unknown")
        self.assertTrue(result["askUser"])

    def test_research_archetype_uses_draft_inference(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "FY results"
        ws.append(["PROFIT AND LOSS", 2022, 2023, 2024])
        ws.append(["Revenue", 10, 11, 12])
        path = self.dir / "research.xlsx"
        wb.save(path)
        code, result = self.plan(path)
        self.assertEqual(result["archetype"], "research")
        self.assertIn("--draft", result["nextCommands"][0])
        self.assertIn("refs/research-recipes.md", result["readRefs"])

    def test_blocked_workbook_exits_two_and_still_prints_the_record(self):
        wb = Workbook()
        wb.active["A1"] = "x"
        path = self.dir / "blocked.xlsx"
        wb.save(path)
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr("xl/activeX/activeX1.xml", "<activeXControl/>")
        code, result = self.plan(path)
        self.assertEqual(code, 2)
        self.assertEqual(result["auditStatus"], "BLOCKED")
        self.assertEqual(result["archetype"], "blocked")
        self.assertEqual(result["nextCommands"], [])
        self.assertIn("BLOCKED stops everything.", result["notes"])

    def test_review_template_carries_source_sha_and_one_entry_per_review_issue(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Category", "Value"])
        ws.append(["A", 1])
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        ws.add_chart(chart, "D2")
        path = self.dir / "chart.xlsx"
        wb.save(path)
        code, result = self.plan(path)
        self.assertEqual(code, 0)
        self.assertEqual(result["auditStatus"], "REVIEW_REQUIRED")
        self.assertIn("refs/presentation-migration.md", result["readRefs"])
        template = json.loads((self.out / "review.template.json").read_text())
        self.assertEqual(result["reviewTemplate"], str(self.out / "review.template.json"))
        self.assertEqual(template["sourceSha256"], source_sha256(path))
        self.assertEqual(result["sourceSha256"], source_sha256(path))
        self.assertEqual(template["issues"], [{"code": "CHARTS_PRESENT", "owner": "", "disposition": "",
                                               "verified": False}])
        # The template is not itself a valid disposition (verified must be true).
        from migration_gate import MigrationGateError, gate_source
        with self.assertRaises(MigrationGateError):
            gate_source(path, str(self.out / "review.template.json"))

    def test_missing_source_is_an_error_not_a_traceback(self):
        code, result = self.plan(self.dir / "nope.xlsx")
        self.assertEqual(code, 1)
        self.assertIsNone(result)


if __name__ == "__main__":
    unittest.main()
