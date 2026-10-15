#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import importlib.util
import json
import sys
import tempfile
import types
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / "lib"))

from migration_gate import (  # noqa: E402
    MigrationGateError,
    gate_plan_source,
    gate_source,
    source_sha256,
)


def load_script(filename, name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class WorkbookFixtures(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.tmp_path = Path(self.tmp.name)

    def save_plain(self, name="book.xlsx"):
        wb = Workbook()
        wb.active["A1"] = "value"
        path = self.tmp_path / name
        wb.save(path)
        return path

    def save_research_grid(self, name="research.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.title = "FY results"
        ws.append(["PROFIT AND LOSS", 2022, 2023, 2024])
        ws.append(["Revenue", 10, 11, 12])
        ws.append(["Net income", 2, 3, 4])
        path = self.tmp_path / name
        wb.save(path)
        return path

    def save_review(self, name="review.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.append(["Category", "Value"])
        ws.append(["A", 1])
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        ws.add_chart(chart, "D2")
        path = self.tmp_path / name
        wb.save(path)
        return path

    def save_blocked(self, name="blocked.xlsx"):
        path = self.save_plain(name)
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr("xl/activeX/activeX1.xml", "<activeXControl/>")
        return path

    def dispositions(self, source, codes, *, digest=None):
        path = self.tmp_path / f"{source.stem}.review.json"
        path.write_text(json.dumps({
            "sourceSha256": digest or hashlib.sha256(source.read_bytes()).hexdigest(),
            "issues": [
                {"code": code, "owner": "Migration owner", "disposition": "Implemented and checked",
                 "verified": True}
                for code in codes
            ],
        }), encoding="utf-8")
        return path


class TestMigrationGate(WorkbookFixtures):
    def test_auto_source_records_sha256(self):
        source = self.save_plain()

        result = gate_source(source)

        self.assertEqual(result["status"], "AUTO")
        self.assertEqual(result["sourceSha256"], hashlib.sha256(source.read_bytes()).hexdigest())

    def test_blocker_cannot_be_overridden_by_disposition(self):
        source = self.save_blocked()
        review = self.dispositions(source, ["ADDINS_OR_ACTIVEX"])

        with self.assertRaises(MigrationGateError) as caught:
            gate_source(source, review)

        self.assertEqual(caught.exception.exit_code, 2)
        self.assertEqual(caught.exception.report["status"], "BLOCKED")

    def test_review_requires_complete_named_verified_dispositions(self):
        source = self.save_review()

        with self.assertRaises(MigrationGateError) as caught:
            gate_source(source)

        self.assertEqual(caught.exception.exit_code, 1)
        self.assertIn("CHARTS_PRESENT", caught.exception.unresolved_codes)

    def test_complete_source_bound_review_can_proceed(self):
        source = self.save_review()
        review = self.dispositions(source, ["CHARTS_PRESENT"])

        result = gate_source(source, review)

        self.assertEqual(result["status"], "REVIEW_REQUIRED")
        self.assertEqual(result["reviewStatus"], "DISPOSITIONED")

    def test_blank_or_unverified_dispositions_remain_unresolved(self):
        source = self.save_review()
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        for field, value in (("owner", ""), ("disposition", "  "), ("verified", False)):
            with self.subTest(field=field):
                item = {
                    "code": "CHARTS_PRESENT",
                    "owner": "Migration owner",
                    "disposition": "Implemented and checked",
                    "verified": True,
                }
                item[field] = value
                path = self.tmp_path / f"invalid-{field}.json"
                path.write_text(json.dumps({"sourceSha256": digest, "issues": [item]}),
                                encoding="utf-8")
                with self.assertRaises(MigrationGateError) as caught:
                    gate_source(source, path)
                self.assertIn("CHARTS_PRESENT", caught.exception.unresolved_codes)

    def test_stale_dispositions_fail_after_source_changes(self):
        source = self.save_review()
        review = self.dispositions(source, ["CHARTS_PRESENT"])
        with zipfile.ZipFile(source, "a") as archive:
            archive.writestr("customXml/item1.xml", "<changed/>")

        with self.assertRaises(MigrationGateError) as caught:
            gate_source(source, review)

        self.assertIn("sourceSha256", str(caught.exception))

    def test_plan_requires_source_path_and_sha_identity(self):
        source = self.save_plain("source.xlsx")
        digest = source_sha256(source)
        plan_path = self.tmp_path / "plan.json"
        plan_path.write_text("{}", encoding="utf-8")

        with self.assertRaises(MigrationGateError):
            gate_plan_source({"sourceSha256": digest}, plan_path=plan_path)
        with self.assertRaises(MigrationGateError):
            gate_plan_source({"sourcePath": "source.xlsx"}, plan_path=plan_path)
        with self.assertRaises(MigrationGateError):
            gate_plan_source(
                {"sourcePath": "source.xlsx", "sourceSha256": "0" * 64},
                plan_path=plan_path,
            )

        resolved, result = gate_plan_source(
            {"sourcePath": "source.xlsx", "sourceSha256": digest},
            plan_path=plan_path,
        )
        self.assertEqual(resolved, source.resolve())
        self.assertEqual(result["status"], "AUTO")

    def test_frozen_formula_fallback_requires_its_own_disposition(self):
        source = self.save_plain("source.xlsx")
        digest = source_sha256(source)
        plan = {"sourcePath": str(source), "sourceSha256": digest, "parity": {"live": 10, "frozen": 2}}

        with self.assertRaises(MigrationGateError) as caught:
            gate_plan_source(plan)
        self.assertIn("CACHED_FORMULA_FALLBACK", caught.exception.unresolved_codes)

        review = self.dispositions(source, ["CACHED_FORMULA_FALLBACK"])
        _, result = gate_plan_source(plan, dispositions_path=review)
        self.assertEqual(result["reviewStatus"], "DISPOSITIONED")


class TestGateCliIntegration(WorkbookFixtures):
    def test_standalone_inference_blocks_activex_before_infer(self):
        source = self.save_blocked()
        output = self.tmp_path / "blocked-plan.json"
        inference = load_script("infer-canonical-formulas.py", "blocked_inference_gate_test")

        with mock.patch.object(inference, "infer", side_effect=AssertionError("inference reached")):
            exit_code = inference.main([str(source), "--out", str(output), "--quiet"])

        self.assertEqual(exit_code, 2)
        self.assertFalse(output.exists())

    def test_standalone_inference_requires_review_before_infer(self):
        source = self.save_review()
        output = self.tmp_path / "review-plan.json"
        inference = load_script("infer-canonical-formulas.py", "review_inference_gate_test")

        with mock.patch.object(inference, "infer", side_effect=AssertionError("inference reached")):
            exit_code = inference.main([str(source), "--out", str(output), "--quiet"])

        self.assertEqual(exit_code, 1)
        self.assertFalse(output.exists())

    def test_standalone_inference_accepts_source_bound_review(self):
        source = self.save_research_grid()
        loaded = load_workbook(source)
        sheet = loaded["FY results"]
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        sheet.add_chart(chart, "F2")
        loaded.save(source)
        review = self.dispositions(source, ["CHARTS_PRESENT"])
        output = self.tmp_path / "reviewed-plan.json"
        inference = load_script("infer-canonical-formulas.py", "reviewed_inference_gate_test")

        exit_code = inference.main([
            str(source), "--review-dispositions", str(review),
            "--out", str(output), "--quiet",
        ])

        self.assertEqual(exit_code, 0)
        self.assertEqual(json.loads(output.read_text(encoding="utf-8"))["source"], source.name)

    def test_inference_plan_records_source_identity_with_spaced_options(self):
        source = self.save_research_grid()
        output = self.tmp_path / "plan.json"
        inference = load_script("infer-canonical-formulas.py", "inference_source_identity_test")

        exit_code = inference.main([
            str(source), "--sheet", "FY results", "--out", str(output), "--quiet",
        ])

        plan = json.loads(output.read_text(encoding="utf-8"))
        self.assertEqual(exit_code, 0)
        self.assertEqual(plan["source"], source.name)
        self.assertEqual(plan["sourcePath"], str(source.resolve()))
        self.assertEqual(plan["sourceSha256"], hashlib.sha256(source.read_bytes()).hexdigest())

    def test_batch_marks_activex_blocked_without_inference_or_auth(self):
        source = self.save_blocked()
        manifest = self.tmp_path / "manifest.json"
        batch = load_script("batch-convert.py", "batch_gate_test")

        with mock.patch.object(batch.INF, "infer", side_effect=AssertionError("inference reached")):
            exit_code = batch.main([
                str(source), "--out", str(manifest), "--post", "--post-review",
                "--conn", "conn", "--folder", "folder",
            ])

        payload = json.loads(manifest.read_text(encoding="utf-8"))
        self.assertEqual(exit_code, 2)
        self.assertEqual(payload["files"][0]["bucket"], "BLOCKED")
        self.assertEqual(payload["files"][0]["auditStatus"], "BLOCKED")
        self.assertEqual(payload["files"][0]["auditIssues"][0]["code"], "ADDINS_OR_ACTIVEX")

    def test_batch_post_review_does_not_bypass_unresolved_review(self):
        source = self.save_review()
        batch = load_script("batch-convert.py", "batch_review_gate_test")

        with mock.patch.object(batch.INF, "infer", side_effect=AssertionError("inference reached")), \
             mock.patch.object(batch, "load_mod", side_effect=AssertionError("auth/API path reached")):
            exit_code = batch.main([str(source), "--post", "--post-review"])

        self.assertEqual(exit_code, 1)

    def test_batch_stale_review_keeps_full_audit_in_manifest(self):
        source = self.save_review()
        review = self.dispositions(source, ["CHARTS_PRESENT"], digest="0" * 64)
        manifest = self.tmp_path / "stale-manifest.json"
        batch = load_script("batch-convert.py", "batch_stale_review_test")

        with mock.patch.object(batch.INF, "infer", side_effect=AssertionError("inference reached")):
            exit_code = batch.main([
                str(source), "--review-dispositions", str(review), "--out", str(manifest),
            ])

        entry = json.loads(manifest.read_text(encoding="utf-8"))["files"][0]
        self.assertEqual(exit_code, 1)
        self.assertEqual(entry["auditStatus"], "REVIEW_REQUIRED")
        self.assertEqual(entry["auditIssues"][0]["code"], "CHARTS_PRESENT")
        self.assertEqual(entry["reviewStatus"], "UNRESOLVED")

    def test_batch_reports_frozen_formula_count_when_review_blocks(self):
        source = self.save_plain()
        manifest = self.tmp_path / "frozen-manifest.json"
        batch = load_script("batch-convert.py", "batch_frozen_review_test")
        plan = {
            "source": source.name,
            "sheet": "FY results",
            "periods": [{"year": 2024, "col": "B"}],
            "lines": [],
            "anchors": {},
            "summary": {"input": 0, "derived": 0, "ratio": 0, "needs_review": 0},
            "parity": {"live": 1, "frozen": 1},
        }

        with mock.patch.object(batch, "strip_report", return_value={"sheets": ["FY results"]}), \
             mock.patch.object(batch.INF, "infer", return_value=plan):
            exit_code = batch.main([str(source), "--out", str(manifest)])

        entry = json.loads(manifest.read_text(encoding="utf-8"))["files"][0]
        self.assertEqual(exit_code, 1)
        self.assertEqual(entry["bucket"], "NEEDS_REVIEW")
        self.assertEqual(entry["frozen"], 1)
        self.assertEqual(entry["auditIssues"][0]["code"], "CACHED_FORMULA_FALLBACK")

    def test_batch_post_keeps_distinct_plans_for_duplicate_basenames(self):
        first_dir = self.tmp_path / "first"
        second_dir = self.tmp_path / "second"
        first_dir.mkdir()
        second_dir.mkdir()
        first = self.save_plain()
        first.rename(first_dir / "model.xlsx")
        first = first_dir / "model.xlsx"
        second = self.save_plain("second.xlsx")
        second.rename(second_dir / "model.xlsx")
        second = second_dir / "model.xlsx"
        batch = load_script("batch-convert.py", "batch_duplicate_basename_test")

        plans = {
            str(first): {"marker": "first", "lines": [{"typed": {2024: 1}}]},
            str(second): {"marker": "second", "lines": [{"typed": {2024: 2}}]},
        }
        consumed = []
        dm_counter = iter(("dm-first", "dm-second"))

        def fake_triage(path, _sheet, _review):
            return ({"file": Path(path).name, "bucket": "AUTO_PARITY"}, plans[path])

        class FakeWire:
            @staticmethod
            def wire_workbook(body):
                return body

            @staticmethod
            def dry_run_request(body):
                return body

            @staticmethod
            def assert_dry_run_success(_response):
                return None

        def build_dm(plan, _conn, _folder, _name):
            consumed.append(plan["marker"])
            return ({"marker": plan["marker"]}, [])

        def api(_base, _token, method, route, _body=None):
            if method == "POST" and route == "/v2/dataModels/spec":
                return {"dataModelId": next(dm_counter)}
            if method == "GET":
                return {"pages": []}
            return {"workbookId": "workbook"}

        fake_builder = types.SimpleNamespace(
            env=lambda: {"SIGMA_BASE_URL": "https://example.invalid"},
            token=lambda _env: "token",
            api=api,
            build_dm=build_dm,
            map_dm=lambda _full: {},
            build_wb=lambda *_args: {},
            workbook_wire=FakeWire(),
        )

        with mock.patch.object(batch, "triage", side_effect=fake_triage), \
             mock.patch.object(batch, "load_mod", return_value=fake_builder):
            exit_code = batch.main([
                str(first), str(second), "--post", "--conn", "conn", "--folder", "folder",
            ])

        self.assertEqual(exit_code, 0)
        self.assertEqual(consumed, ["first", "second"])

    def test_research_plan_missing_source_fails_before_auth(self):
        plan_path = self.tmp_path / "research-plan.json"
        plan_path.write_text(json.dumps({"sourceSha256": "0" * 64}), encoding="utf-8")
        builder = load_script("build-research-model.py", "research_plan_gate_test")

        with mock.patch.object(builder, "env", side_effect=AssertionError("auth reached")):
            exit_code = builder.main([str(plan_path), "--post"])

        self.assertEqual(exit_code, 1)

    def test_forecast_plan_stale_source_fails_before_auth(self):
        source = self.save_plain("forecast-source.xlsx")
        plan_path = self.tmp_path / "forecast-plan.json"
        plan_path.write_text(json.dumps({
            "sourcePath": str(source),
            "sourceSha256": "0" * 64,
        }), encoding="utf-8")
        builder = load_script("build-forecast-model.py", "forecast_plan_gate_test")

        with mock.patch.object(builder, "env", side_effect=AssertionError("auth reached")):
            exit_code = builder.main([str(plan_path), "--post"])

        self.assertEqual(exit_code, 1)

    def test_direct_builder_blocks_before_inference_even_for_dry_run(self):
        source = self.save_blocked()
        builder = load_script("build-research-workbook.py", "direct_gate_test")

        with mock.patch.object(builder.INF, "load", side_effect=AssertionError("inference reached")):
            exit_code = builder.main([str(source), "--post", "--conn", "conn", "--folder", "folder"])

        self.assertEqual(exit_code, 2)

    def test_direct_builder_unresolved_review_blocks_before_inference(self):
        source = self.save_review()
        builder = load_script("build-research-workbook.py", "direct_review_gate_test")

        with mock.patch.object(builder.INF, "load", side_effect=AssertionError("inference reached")):
            exit_code = builder.main([str(source)])

        self.assertEqual(exit_code, 1)


if __name__ == "__main__":
    unittest.main()
