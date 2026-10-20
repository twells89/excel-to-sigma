"""Source-cell native lowering, checked against hand-derived workbook cases."""
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts/lib"))
from source_inputs import build_source_input_plan  # noqa: E402
from native_calculations import (build_native_calculation_plan, lower_native_calculations,
                                 prepare_native_calculation_update, confirm_native_calculation_readback,
                                 build_native_calculation_backup, _render)  # noqa: E402
from excel_expression import UnsupportedExpression  # noqa: E402
from input_sync import build_artifact_records  # noqa: E402
from migration_manifest import build_manifest  # noqa: E402


class NativeCalculationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def source(self, sheet="Ledger Renamed"):
        book = Workbook()
        ws = book.active
        ws.title = sheet
        ws.append(["Account", "Amount", "Rate", "Result"])
        ws.append(["a", 3, 2, "=B2*(1+C2)"])
        ws.append(["b", 4, 1, "=B3+C3"])
        ws.add_table(Table(displayName="tblLedger", ref="A1:D3"))
        result = book.create_sheet("Summary Renamed")
        result["A1"] = f"='{sheet}'!D2+'{sheet}'!D3"
        result["A2"] = "=IFERROR(1/0,99)"
        path = Path(self.temp.name) / "renamed.xlsx"
        book.save(path)
        decisions = {"sourceSha256": hashlib.sha256(path.read_bytes()).hexdigest(), "features": [{
            "featureId": "ledger", "sheet": sheet, "table": "tblLedger", "role": "input",
            "reviewed": True, "keyColumns": ["Account"],
            "literalColumns": ["Account", "Amount", "Rate"],
            "columnTypes": {"Account": "text", "Amount": "number", "Rate": "number"}}]}
        plan = build_source_input_plan(path, decisions, connection_id="connection", folder_id="folder")
        return path, plan

    def test_exact_source_bindings_and_cross_sheet_dependency_order(self):
        path, inputs = self.source()
        plan = build_native_calculation_plan(path, inputs)
        self.assertEqual(plan["sourceSha256"], inputs["sourceSha256"])
        self.assertEqual(plan["formulaCount"], 4)
        self.assertEqual(plan["bindings"]["Ledger Renamed!B2"], {
            "kind": "input", "featureId": "ledger", "rowIndex": 0,
            "column": "Amount", "valueType": "number"})
        self.assertEqual(plan["formulas"]["Ledger Renamed!D2"]["status"], "ready-to-bind")
        self.assertEqual(plan["formulas"]["Ledger Renamed!D2"]["localExpected"], 9)
        self.assertEqual(plan["formulas"]["Summary Renamed!A1"]["localExpected"], 14)
        self.assertEqual(plan["formulas"]["Summary Renamed!A1"]["dependencies"],
                         ["Ledger Renamed!D2", "Ledger Renamed!D3"])
        self.assertLess(plan["formulas"]["Ledger Renamed!D2"]["layer"],
                        plan["formulas"]["Summary Renamed!A1"]["layer"])
        self.assertEqual(plan["formulas"]["Summary Renamed!A2"]["status"], "unresolved")
        self.assertIn("IFERROR", plan["formulas"]["Summary Renamed!A2"]["reason"])

    def test_parenthesized_live_lookup_uses_observed_physical_rows(self):
        path, inputs = self.source()
        plan = build_native_calculation_plan(path, inputs)
        assigned = {"ledger": {"elementId": "observed-element", "name": "Ledger Inputs",
                               "rowIdColumn": "ID", "rows": {"0": "observed-row-a", "1": "observed-row-b"}}}
        result = lower_native_calculations(plan, assigned, max_columns=2)
        first = result["formulas"]["Ledger Renamed!D2"]
        self.assertIn('Lookup([Ledger Inputs/Amount], "observed-row-a", [Ledger Inputs/ID])', first)
        self.assertIn('Lookup([Ledger Inputs/Rate], "observed-row-a", [Ledger Inputs/ID])', first)
        self.assertIn(' * ((1 + ', first)
        self.assertEqual(result["formulas"]["Summary Renamed!A1"],
                         "([Native Calculations 1/f_ledger_renamed_d2] + [Native Calculations 1/f_ledger_renamed_d3])")
        self.assertEqual(len(result["shards"]), 2)
        self.assertEqual(result["shards"][1]["source"]["elementId"], result["shards"][0]["id"])

    def test_missing_row_or_changed_source_fails_closed(self):
        path, inputs = self.source()
        plan = build_native_calculation_plan(path, inputs)
        with self.assertRaisesRegex(ValueError, "row ID"):
            lower_native_calculations(plan, {"ledger": {"name": "Ledger Inputs", "rowIdColumn": "ID",
                                                         "rows": {"0": "a"}}})
        book = Workbook()
        book.active["A1"] = 1
        book.save(path)
        with self.assertRaisesRegex(ValueError, "source hash"):
            build_native_calculation_plan(path, inputs)

    def test_ready_constant_without_reviewed_target_scope_cannot_construct(self):
        book = Workbook()
        book.active["A1"] = "=1+2"
        path = Path(self.temp.name) / "constant.xlsx"
        book.save(path)
        plan = build_native_calculation_plan(path)
        self.assertEqual(plan["formulas"]["Sheet!A1"]["status"], "ready-to-bind")
        with self.assertRaisesRegex(ValueError, "target scope"):
            lower_native_calculations(plan, {})

    def test_unmapped_business_literal_and_error_are_explicit_obligations(self):
        path, inputs = self.source()
        book = Workbook()
        ws = book.active
        ws.title = "Cases"
        ws["A1"] = 7
        ws["B1"] = "=A1+1"
        ws["C1"] = "#DIV/0!"
        ws["C1"].data_type = "e"
        ws["D1"] = "=C1+1"
        other = Path(self.temp.name) / "cases.xlsx"
        book.save(other)
        empty = build_native_calculation_plan(other)
        self.assertEqual(empty["formulas"]["Cases!B1"]["status"], "unresolved")
        self.assertIn("Cases!A1", empty["formulas"]["Cases!B1"]["reason"])
        self.assertEqual(empty["formulas"]["Cases!D1"]["status"], "unresolved")
        self.assertIn("Excel error", empty["formulas"]["Cases!D1"]["reason"])

    def test_source_calculation_error_never_becomes_native_candidate(self):
        book = Workbook()
        ws = book.active
        ws.title = "Errors"
        ws["A1"] = "=1/0"
        ws["A2"] = "=A1+1"
        path = Path(self.temp.name) / "errors.xlsx"
        book.save(path)
        plan = build_native_calculation_plan(path)
        self.assertEqual(plan["formulas"]["Errors!A1"]["status"], "unresolved")
        self.assertIn("#DIV/0!", plan["formulas"]["Errors!A1"]["reason"])
        self.assertEqual(plan["formulas"]["Errors!A2"]["status"], "unresolved")

    def test_excel_sum_ignoring_inline_text_is_not_rendered_as_text_addition(self):
        book = Workbook()
        book.active["A1"] = '=SUM("ignored",2)'
        path = Path(self.temp.name) / "sum-text.xlsx"
        book.save(path)
        plan = build_native_calculation_plan(path)
        self.assertEqual(plan["formulas"]["Sheet!A1"]["status"], "unresolved")
        self.assertIn("text", plan["formulas"]["Sheet!A1"]["reason"])

    def test_finite_indirect_needs_exact_invalid_selector_error_semantics(self):
        bound = {"kind": "indirect", "selector": {"kind": "text", "value": "Base"},
                 "alternatives": [["Base", {"kind": "number", "source": "1"}]]}
        with self.assertRaisesRegex(UnsupportedExpression, "invalid-selector/error"):
            _render(bound, lambda _cell, _arithmetic: "unused")

    def test_reviewed_external_replacement_feeds_downstream_live_formula(self):
        book = Workbook()
        ws = book.active
        ws.title = "Budget"
        ws["F4"] = "Q3 2026"
        ws["F21"] = "='[Corporate.xlsx]Alloc'!$C$4"
        ws["F22"] = "=F21+1"
        path = Path(self.temp.name) / "replacement.xlsx"
        book.save(path)
        decisions = {"sourceSha256": hashlib.sha256(path.read_bytes()).hexdigest(), "features": [{
            "featureId": "allocation", "role": "replacement", "reviewed": True,
            "displayName": "Corporate Allocations", "replacementFor": "Corporate.xlsx",
            "replacementReview": {"approved": True, "basis": "explicit supplied value"},
            "keyColumns": ["Period"], "literalColumns": ["Period", "Allocation"],
            "columnTypes": {"Period": "text", "Allocation": "number"},
            "suppliedRows": [{"Period": "Q3", "Allocation": 61}],
            "replacementBindings": [{"key": {"Period": "Q3"},
                                     "sourceCells": {"Period": "Budget!F4", "Allocation": "Budget!F21"},
                                     "sourceLabel": "Q3 2026", "originalFormula": ws["F21"].value}]}]}
        inputs = build_source_input_plan(path, decisions, connection_id="conn", folder_id="folder")
        native = build_native_calculation_plan(path, inputs)
        self.assertEqual(native["formulas"]["Budget!F21"]["status"], "reviewed-replacement-input")
        self.assertEqual(native["formulas"]["Budget!F22"]["status"], "ready-to-bind")
        compiled = lower_native_calculations(native, {"allocation": {
            "elementId": "real-input", "name": "Corporate Allocations", "rowIdColumn": "Row ID",
            "rows": {"0": "real-row"}}})
        self.assertIn('Lookup([Corporate Allocations/Allocation], "real-row",',
                      compiled["formulas"]["Budget!F22"])

    def test_one_cell_structured_row_reference_lowers_to_exact_input_cell(self):
        book = Workbook()
        ws = book.active
        ws.title = "Orders"
        ws.append(["ID", "Units", "Price", "Revenue"])
        ws.append(["x", 2, 7, "=tblOrders[[#This Row],[Units]]*tblOrders[[#This Row],[Price]]"])
        ws.add_table(Table(displayName="tblOrders", ref="A1:D2"))
        path = Path(self.temp.name) / "orders.xlsx"
        book.save(path)
        input_plan = build_source_input_plan(path, {"sourceSha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "features": [{"featureId": "orders", "sheet": "Orders", "table": "tblOrders",
                          "role": "input", "reviewed": True, "keyColumns": ["ID"],
                          "literalColumns": ["ID", "Units", "Price"],
                          "columnTypes": {"ID": "text", "Units": "number", "Price": "number"}}]},
            connection_id="conn", folder_id="folder")
        native = build_native_calculation_plan(path, input_plan)
        self.assertEqual(native["formulas"]["Orders!D2"]["status"], "ready-to-bind")
        formula = lower_native_calculations(native, {"orders": {"elementId": "e", "name": "Orders Input",
            "rowIdColumn": "ID", "rows": {"0": "physical-x"}}})["formulas"]["Orders!D2"]
        self.assertIn("[Orders Input/Units]", formula)
        self.assertIn("[Orders Input/Price]", formula)

    def test_guarded_rest_shard_preserves_foreign_contents_and_requires_readback(self):
        path, inputs = self.source()
        native = build_native_calculation_plan(path, inputs)
        approvals = [{"id": cell, "mappingStatus": "live",
                      "targetMapping": {"logicalElementId": "native_calc_" + native["sourceSha256"][:8] +
                                        ("_1" if record["layer"] == 0 else "_2"),
                                        "logicalColumnId": record["columnId"]}}
                     for cell, record in native["formulas"].items() if record["status"] == "ready-to-bind"]
        manifest = build_manifest(path, {"sharedInputArchitecture": inputs["sharedInputArchitecture"],
                                         "inputPlanHash": inputs["inputPlanHash"], "formulaCells": approvals})
        manifest["target"] = {"documentId": "actual-doc", "version": 7}
        target = {"documentId": "actual-doc", "version": 7, "pageId": "page",
                  "elementId": "actual-input", "columnIds": {"Account": "ca", "Amount": "cb", "Rate": "cc"},
                  "rowIdColumnId": "rid"}
        architecture = inputs["sharedInputArchitecture"]
        receipt = {"status": "APPROVED_STAGING", "sourceHash": manifest["sourceHash"],
                   "planHash": manifest["planHash"], "targetBindings": {"ledger": target},
                   "artifacts": build_artifact_records(manifest, architecture, {"ledger": target})}
        current = {"workbookId": "actual-doc", "documentVersion": 7, "contents": {
            "schemaVersion": 1, "kind": "workbook", "pages": [{"id": "page", "name": "Page"}],
            "elements": [{"id": "actual-input", "kind": "input-table", "name": "tblLedger",
                          "columns": [{"id": "ca", "name": "Account", "type": "text"},
                                      {"id": "cb", "name": "Amount", "type": "number"},
                                      {"id": "cc", "name": "Rate", "type": "number"},
                                      {"id": "rid", "name": "Row ID", "type": None}]}],
            "layout": '<Page id="page"></Page>', "foreignEvents": [{"id": "keep"}]}}
        observed = {"ledger": [{"Account": "a", "Amount": 3, "Rate": 2, "Row ID": "actual-row-a"},
                               {"Account": "b", "Amount": 4, "Rate": 1, "Row ID": "actual-row-b"}]}
        backup = build_native_calculation_backup(current, observed)
        self.assertEqual(backup["inputRows"], observed)
        with self.assertRaisesRegex(ValueError, "backup"):
            prepare_native_calculation_update(native, inputs, manifest, receipt,
                                              current, observed, shard_index=0, max_columns=2)
        proposal = prepare_native_calculation_update(native, inputs, manifest, receipt,
                                                     current, observed, backup=backup,
                                                     shard_index=0, max_columns=2)
        self.assertEqual(proposal["method"], "PUT")
        self.assertEqual(proposal["dryRunRequest"]["documentVersion"], 7)
        self.assertEqual(proposal["dryRunRequest"]["contents"]["foreignEvents"], [{"id": "keep"}])
        self.assertEqual(proposal["dryRunRequest"]["contents"]["elements"][0], current["contents"]["elements"][0])
        self.assertIn("actual-row-a", proposal["expectedFormulaColumns"]["Ledger Renamed!D2"])
        with self.assertRaisesRegex(ValueError, "readback"):
            confirm_native_calculation_readback(proposal, current)
        saved = {"workbookId": "actual-doc", "documentVersion": 8,
                 "contents": proposal["dryRunRequest"]["contents"]}
        confirmation = confirm_native_calculation_readback(proposal, saved)
        manifest["target"]["version"] = 8
        target["version"] = 8
        receipt["targetBindings"]["ledger"]["version"] = 8
        receipt["artifacts"] = build_artifact_records(manifest, architecture, {"ledger": target})
        receipt["planHash"] = manifest["planHash"]
        with self.assertRaisesRegex(ValueError, "prior shard readback"):
            prepare_native_calculation_update(native, inputs, manifest, receipt, saved, observed,
                backup=build_native_calculation_backup(saved, observed), shard_index=1,
                max_columns=2, prior_readbacks=[dict(confirmation, documentId="other-document")])
        second = prepare_native_calculation_update(native, inputs, manifest, receipt,
                                                   saved, observed, backup=build_native_calculation_backup(saved, observed),
                                                   shard_index=1, max_columns=2,
                                                   prior_readbacks=[confirmation])
        self.assertEqual(second["expectedFormulaColumns"]["Summary Renamed!A1"],
                         "([Native Calculations 1/f_ledger_renamed_d2] + [Native Calculations 1/f_ledger_renamed_d3])")
        self.assertEqual(len(second["dryRunRequest"]["contents"]["elements"]), 3)
        stale = dict(current, documentVersion=9)
        with self.assertRaisesRegex(ValueError, "target/version"):
            prepare_native_calculation_update(native, inputs, manifest, receipt, stale, observed,
                                              backup=build_native_calculation_backup(stale, observed),
                                              shard_index=0, max_columns=2)

    def test_cli_writes_source_bound_plan_without_target_calls(self):
        path, inputs = self.source()
        plan_path = Path(self.temp.name) / "inputs.json"
        output = Path(self.temp.name) / "native.json"
        plan_path.write_text(json.dumps(inputs))
        program = Path(__file__).resolve().parents[1] / "scripts/native-calculations.py"
        result = subprocess.run([sys.executable, str(program), "plan", "--source", str(path),
                                 "--input-plan", str(plan_path), "--out", str(output)],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(output.read_text())["formulaCount"], 4)


if __name__ == "__main__":
    unittest.main()
