#!/usr/bin/env python3
"""Generate deterministic Excel workbooks for chart-migration examples.

Each workbook contains an editable formal Table on its first sheet and three
charts on a dashboard sheet. The fixtures use synthetic data only.
"""

from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import (
    BarChart,
    DoughnutChart,
    LineChart,
    Reference,
    ScatterChart,
    Series,
)
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT_DIR = Path(__file__).resolve().parent / "chart-examples"
MONTHS = [date(2026, month, 1) for month in range(1, 7)]


def add_table(ws, name):
    table = Table(displayName=name, ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = table.ref


def format_data_sheet(ws, date_columns=(), currency_columns=()):
    for column in range(1, ws.max_column + 1):
        values = [str(ws.cell(row, column).value or "") for row in range(1, ws.max_row + 1)]
        ws.column_dimensions[get_column_letter(column)].width = min(
            24, max(12, max(map(len, values)) + 2)
        )
    for row in range(2, ws.max_row + 1):
        for column in date_columns:
            ws.cell(row, column).number_format = "mmm-yy"
        for column in currency_columns:
            ws.cell(row, column).number_format = '$#,##0.00'


def prepare_dashboard(ws, title, subtitle, accent):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(size=20, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=accent)
    ws.merge_cells("A1:H1")
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="475569")
    ws.merge_cells("A2:H2")


def add_line_chart(ws, data_range, category_range, title, anchor, colors=()):
    chart = LineChart()
    chart.title = title
    chart.height = 8
    chart.width = 15
    chart.add_data(data_range, titles_from_data=True)
    chart.set_categories(category_range)
    chart.legend.position = "b"
    chart.y_axis.title = "Value"
    for series, color in zip(chart.series, colors):
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 24000
    ws.add_chart(chart, anchor)


def add_bar_chart(ws, data_range, category_range, title, anchor, color):
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = title
    chart.height = 8
    chart.width = 11
    chart.add_data(data_range, titles_from_data=True)
    chart.set_categories(category_range)
    chart.legend = None
    chart.series[0].graphicalProperties.solidFill = color
    ws.add_chart(chart, anchor)


def add_doughnut_chart(ws, data_range, category_range, title, anchor):
    chart = DoughnutChart()
    chart.title = title
    chart.holeSize = 55
    chart.height = 8
    chart.width = 11
    chart.add_data(data_range, titles_from_data=True)
    chart.set_categories(category_range)
    chart.legend.position = "b"
    ws.add_chart(chart, anchor)


def make_sales():
    wb = Workbook()
    data = wb.active
    data.title = "Sales Data"
    data.append(["Month", "Region", "Product", "Units", "Unit Price", "Revenue"])

    regions = ["West", "Central", "East"]
    products = [("Hardware", 425), ("Services", 210), ("Software", 160)]
    rows = []
    for month_index, month in enumerate(MONTHS):
        for region_index, region in enumerate(regions):
            for product_index, (product, price) in enumerate(products):
                units = 24 + month_index * 3 + region_index * 4 + product_index * 5
                revenue = units * price
                row = [month, region, product, units, price, revenue]
                rows.append(row)
                data.append(row)
    add_table(data, "tblSales")
    format_data_sheet(data, date_columns=(1,), currency_columns=(5, 6))

    dashboard = wb.create_sheet("Sales Dashboard")
    prepare_dashboard(
        dashboard,
        "Sales Performance",
        "54 editable source rows · January–June 2026",
        "2563EB",
    )

    monthly = defaultdict(float)
    by_region = defaultdict(float)
    by_product = defaultdict(float)
    for month, region, product, _units, _price, revenue in rows:
        monthly[month] += revenue
        by_region[region] += revenue
        by_product[product] += revenue

    dashboard.append([])
    dashboard.append(["Month", "Revenue", "", "Region", "Revenue", "", "Product", "Revenue"])
    for index in range(max(len(monthly), len(by_region), len(by_product))):
        month_items = list(monthly.items())
        region_items = list(by_region.items())
        product_items = list(by_product.items())
        dashboard.append([
            month_items[index][0] if index < len(month_items) else None,
            month_items[index][1] if index < len(month_items) else None,
            None,
            region_items[index][0] if index < len(region_items) else None,
            region_items[index][1] if index < len(region_items) else None,
            None,
            product_items[index][0] if index < len(product_items) else None,
            product_items[index][1] if index < len(product_items) else None,
        ])
    for row in range(5, 11):
        dashboard.cell(row, 1).number_format = "mmm-yy"
    for column in (2, 5, 8):
        for row in range(5, 11):
            dashboard.cell(row, column).number_format = '$#,##0'

    add_line_chart(
        dashboard,
        Reference(dashboard, min_col=2, min_row=4, max_row=10),
        Reference(dashboard, min_col=1, min_row=5, max_row=10),
        "Monthly Revenue",
        "A12",
        ("2563EB",),
    )
    add_bar_chart(
        dashboard,
        Reference(dashboard, min_col=5, min_row=4, max_row=7),
        Reference(dashboard, min_col=4, min_row=5, max_row=7),
        "Revenue by Region",
        "P12",
        "14B8A6",
    )
    add_doughnut_chart(
        dashboard,
        Reference(dashboard, min_col=8, min_row=4, max_row=7),
        Reference(dashboard, min_col=7, min_row=5, max_row=7),
        "Product Mix",
        "P28",
    )
    return wb, len(rows)


def make_expenses():
    wb = Workbook()
    data = wb.active
    data.title = "Expense Data"
    data.append(["Month", "Department", "Category", "Budget", "Actual"])

    departments = ["Engineering", "Sales", "Operations", "Finance"]
    categories = ["Payroll", "Software", "Travel"]
    rows = []
    for month_index, month in enumerate(MONTHS):
        for department_index, department in enumerate(departments):
            for category_index, category in enumerate(categories):
                budget = 18000 + department_index * 7200 + category_index * 4300 + month_index * 950
                variance_points = month_index + department_index - category_index - 2
                actual = round(budget * (100 + variance_points) / 100, 2)
                row = [month, department, category, budget, actual]
                rows.append(row)
                data.append(row)
    add_table(data, "tblExpenses")
    format_data_sheet(data, date_columns=(1,), currency_columns=(4, 5))

    dashboard = wb.create_sheet("Expense Dashboard")
    prepare_dashboard(
        dashboard,
        "Expense Planning",
        "72 editable source rows · budget-to-actual tracking",
        "EA580C",
    )

    monthly = defaultdict(lambda: [0.0, 0.0])
    by_department = defaultdict(lambda: [0.0, 0.0])
    by_category = defaultdict(float)
    for month, department, category, budget, actual in rows:
        monthly[month][0] += budget
        monthly[month][1] += actual
        by_department[department][0] += budget
        by_department[department][1] += actual
        by_category[category] += actual

    dashboard.append([])
    dashboard.append([
        "Month", "Budget", "Actual", "", "Department", "Variance", "", "Category", "Actual",
    ])
    longest = max(len(monthly), len(by_department), len(by_category))
    for index in range(longest):
        month_items = list(monthly.items())
        department_items = list(by_department.items())
        category_items = list(by_category.items())
        department = department_items[index] if index < len(department_items) else None
        dashboard.append([
            month_items[index][0] if index < len(month_items) else None,
            month_items[index][1][0] if index < len(month_items) else None,
            month_items[index][1][1] if index < len(month_items) else None,
            None,
            department[0] if department else None,
            department[1][1] - department[1][0] if department else None,
            None,
            category_items[index][0] if index < len(category_items) else None,
            category_items[index][1] if index < len(category_items) else None,
        ])
    for row in range(5, 11):
        dashboard.cell(row, 1).number_format = "mmm-yy"
    for column in (2, 3, 6, 9):
        for row in range(5, 11):
            dashboard.cell(row, column).number_format = '$#,##0'

    monthly_chart = BarChart()
    monthly_chart.type = "col"
    monthly_chart.style = 10
    monthly_chart.title = "Monthly Budget vs Actual"
    monthly_chart.height = 8
    monthly_chart.width = 15
    monthly_chart.add_data(
        Reference(dashboard, min_col=2, max_col=3, min_row=4, max_row=10),
        titles_from_data=True,
    )
    monthly_chart.set_categories(Reference(dashboard, min_col=1, min_row=5, max_row=10))
    monthly_chart.legend.position = "b"
    dashboard.add_chart(monthly_chart, "A12")
    add_bar_chart(
        dashboard,
        Reference(dashboard, min_col=6, min_row=4, max_row=8),
        Reference(dashboard, min_col=5, min_row=5, max_row=8),
        "Variance by Department",
        "P12",
        "EA580C",
    )
    add_doughnut_chart(
        dashboard,
        Reference(dashboard, min_col=9, min_row=4, max_row=7),
        Reference(dashboard, min_col=8, min_row=5, max_row=7),
        "Actual Expense Mix",
        "P28",
    )
    return wb, len(rows)


def make_headcount():
    wb = Workbook()
    data = wb.active
    data.title = "Headcount Data"
    data.append(["Month", "Department", "Planned Headcount", "Actual Headcount", "Avg Salary"])

    departments = [
        ("Engineering", 32, 142000, 2),
        ("Sales", 24, 118000, 1),
        ("Operations", 18, 96000, 1),
        ("Finance", 11, 126000, 1),
    ]
    rows = []
    for month_index, month in enumerate(MONTHS):
        for department_index, (department, base, salary, growth) in enumerate(departments):
            planned = base + month_index * growth
            actual = planned + department_index - 1
            avg_salary = salary + month_index * 500
            row = [month, department, planned, actual, avg_salary]
            rows.append(row)
            data.append(row)
    add_table(data, "tblHeadcount")
    format_data_sheet(data, date_columns=(1,), currency_columns=(5,))

    dashboard = wb.create_sheet("Headcount Dashboard")
    prepare_dashboard(
        dashboard,
        "Headcount Planning",
        "24 editable source rows · plan-to-actual workforce view",
        "7C3AED",
    )

    monthly = defaultdict(lambda: [0.0, 0.0])
    current = {}
    for month, department, planned, actual, salary in rows:
        monthly[month][0] += planned
        monthly[month][1] += actual
        if month == MONTHS[-1]:
            current[department] = (actual, salary)

    dashboard.append([])
    dashboard.append([
        "Month", "Planned", "Actual", "", "Department", "Actual Headcount", "Average Salary",
    ])
    longest = max(len(monthly), len(current))
    for index in range(longest):
        month_items = list(monthly.items())
        current_items = list(current.items())
        current_row = current_items[index] if index < len(current_items) else None
        dashboard.append([
            month_items[index][0] if index < len(month_items) else None,
            month_items[index][1][0] if index < len(month_items) else None,
            month_items[index][1][1] if index < len(month_items) else None,
            None,
            current_row[0] if current_row else None,
            current_row[1][0] if current_row else None,
            current_row[1][1] if current_row else None,
        ])
    for row in range(5, 11):
        dashboard.cell(row, 1).number_format = "mmm-yy"
        dashboard.cell(row, 7).number_format = '$#,##0'

    add_line_chart(
        dashboard,
        Reference(dashboard, min_col=2, max_col=3, min_row=4, max_row=10),
        Reference(dashboard, min_col=1, min_row=5, max_row=10),
        "Planned vs Actual Headcount",
        "A12",
        ("94A3B8", "7C3AED"),
    )
    add_bar_chart(
        dashboard,
        Reference(dashboard, min_col=6, min_row=4, max_row=8),
        Reference(dashboard, min_col=5, min_row=5, max_row=8),
        "Current Headcount",
        "P12",
        "7C3AED",
    )

    scatter = ScatterChart()
    scatter.title = "Headcount vs Average Salary"
    scatter.x_axis.title = "Actual Headcount"
    scatter.y_axis.title = "Average Salary"
    scatter.height = 8
    scatter.width = 11
    for row in range(5, 9):
        series = Series(
            Reference(dashboard, min_col=7, min_row=row, max_row=row),
            Reference(dashboard, min_col=6, min_row=row, max_row=row),
            title=dashboard.cell(row, 5).value,
        )
        scatter.series.append(series)
    scatter.legend.position = "b"
    dashboard.add_chart(scatter, "P28")
    return wb, len(rows)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    examples = [
        ("Sales Performance.xlsx", make_sales),
        ("Expense Planning.xlsx", make_expenses),
        ("Headcount Planning.xlsx", make_headcount),
    ]
    total_rows = 0
    for filename, build in examples:
        workbook, row_count = build()
        path = OUTPUT_DIR / filename
        workbook.save(path)
        total_rows += row_count
        print(f"wrote {path} ({row_count} rows, 3 charts)")
    print(f"generated {len(examples)} workbooks, {total_rows} rows, 9 charts")


if __name__ == "__main__":
    main()
