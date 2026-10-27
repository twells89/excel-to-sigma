"""Presentation of a table-report plan: KPI proposals, source-chart mapping, the Report page composition.

Everything here is presentation only. Values are never recomputed differently: KPIs read the verified frame
columns, charts read verified frame cells, and ``check_presentation`` re-derives each of them from the Python
oracle and from the source workbook's cached cells before a spec is written.

Plan field ``presentation``::

    {"version": 1, "title", "subtitle", "accent": {...}, "main": <frame id>,
     "kpis":   [{id, label, kind: total|ratio, frame, column | num/den, formula, format, expected, evidence}],
     "charts": [{id, title, kind: bar|line|donut|pie, orientation, stacking, frame, categoryColumn, categories,
                 series: [{title, column, kind, cells, sels, values}]}]}
"""
from __future__ import annotations

import re
from datetime import date, datetime

try:
    from . import excel_formats as ef
    from . import presentation_style as ps
    from . import sigma_formula_eval as se
    from . import table_report_model as trm
except ImportError:  # direct import when scripts/lib itself is on sys.path
    import excel_formats as ef
    import presentation_style as ps
    import sigma_formula_eval as se
    import table_report_model as trm

KPI_NAME = re.compile(r"total|commit|quota|target|goal|plan|budget|forecast|revenue|bookings|pipeline|variance"
                      r"|actual", re.I)
PERIOD_NAME = re.compile(r"^(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[- ']?\d{0,4}$"
                         r"|^\d{4}-\d{2}(-\d{2})?$|^q[1-4]\b|^w(eek)?\s*\d+$|^fy\d+$|^\d+$", re.I)
MAX_KPIS = 4
MIN_KPIS = 2


# ---------------------------------------------------------------- helpers
def main_frame(plan):
    """The frame the page leads with: the first visible summary grid, else the first visible frame."""
    frames = [f for f in plan["frames"] if not f.get("hidden") and not f.get("helper")]
    summaries = [f for f in frames if f["kind"] == "summary"]
    return (summaries or frames or [None])[0]


def _frame(plan, frame_id):
    return next(f for f in plan["frames"] if f["id"] == frame_id)


def _column(frame, name):
    return next((c for c in frame["columns"] if c["name"] == name), None)


def _num(value):
    return 0 if value is None else value


def _is_percent(column):
    return ef.is_percent_format(column.get("format"))


def _slug(text):
    return trm.slug(text)


# ---------------------------------------------------------------- KPIs
def _additive(expr, columns, tables, named=False):
    """True when the column of a summary frame can be summed over its rows into a meaningful total.

    ``named``: the column's own name says it is a headline measure (Budget, Quota, ...), which is what lets a
    per-row lookup of a number from a data table (a department budget) be summed like an amount."""
    tag = expr[0]
    if tag == "agg":
        return expr[1] in ("sum", "count")
    if tag == "c":
        other = columns.get(expr[1])
        return other is not None and _additive(other["expr"], columns, tables, bool(KPI_NAME.search(expr[1])))
    if tag == "f" and expr[1] == "SUM":
        return all(_additive(arg, columns, tables, named) for arg in expr[2:])
    if tag == "op" and expr[1] in ("+", "-"):
        return _additive(expr[2], columns, tables, named) and _additive(expr[3], columns, tables, named)
    if tag == "neg":
        return _additive(expr[1], columns, tables, named)
    if tag == "lk":     # a typed per-row input of the grid itself (quota), or a per-row measure by name (budget)
        table = tables.get(expr[1])
        column = next((c for c in (table or {}).get("columns", []) if c["name"] == expr[4]), None)
        if not (table and column and column.get("type") == "number") or ef.is_percent_format(column.get("format")):
            return False
        return bool(table.get("synthetic")) or named
    return False


def _ratio_parts(expr):
    """``(numerator column, denominator column)`` of ``IFERROR(a / b, 0)`` / ``a / b`` over two columns."""
    if expr[0] == "f" and expr[1] == "IFERROR" and len(expr) == 4 and expr[3] == ["n", 0]:
        expr = expr[2]
    if expr[0] == "op" and expr[1] == "/" and expr[2][0] == "c" and expr[3][0] == "c":
        return expr[2][1], expr[3][1]
    return None


def _cells_of(plan, frame_id, column, kind="summary"):
    for group in plan["checks"]:
        if group["frame"] == frame_id and group["column"] == column and group["kind"] == kind:
            return group["cells"]
    return None


def _sum_cached(cells):
    return sum(_num(cell[2]) for cell in cells)


def propose_kpis(plan):
    """KPI cards for the main summary grid (Recipe 2: equal peers): additive total columns the source totals
    or that are named Total / Commit / Quota-like, plus a ratio of two of them (ratio of totals)."""
    frame = main_frame(plan)
    if frame is None or frame["kind"] != "summary":
        return []
    tables = {t["name"]: t for t in plan["tables"]}
    columns = {c["name"]: c for c in frame["columns"]}
    candidates = {}
    for column in frame["columns"]:
        name = column["name"]
        if column.get("type") != "number" or _is_percent(column) or PERIOD_NAME.match(name):
            continue
        named = bool(KPI_NAME.search(name))
        if not _additive(column["expr"], columns, tables, named):
            continue
        total_cells = _cells_of(plan, frame["id"], name, "summary-total")
        if named or total_cells:
            candidates[name] = {"column": column, "named": named, "totalCell": total_cells[0][0]
                                if total_cells else None}
    ratio = None
    for column in frame["columns"]:
        parts = _ratio_parts(column["expr"])
        if parts and parts[0] in candidates and parts[1] in candidates and parts[0] != parts[1]:
            ratio = (column, *parts)
            break
    limit = MAX_KPIS - (1 if ratio else 0)
    priority = sorted(candidates, key=lambda n: (0 if ratio and n in ratio[1:] else 1,
                                                 0 if re.search("total", n, re.I) else 1,
                                                 0 if candidates[n]["named"] else 1,
                                                 [c["name"] for c in frame["columns"]].index(n)))
    chosen = set(priority[:limit])
    kpis = []
    for column in frame["columns"]:
        name = column["name"]
        if name in chosen:
            kpis.append(_total_kpi(plan, frame, column, candidates[name]))
        elif ratio and column is ratio[0]:
            kpis.append(_ratio_kpi(plan, frame, ratio))
    seen = set()
    for kpi in kpis:                       # "Var" and "Var %" slug alike: keep every id unique
        base, number = kpi["id"], 2
        while kpi["id"] in seen:
            kpi["id"], number = f"{base}_{number}", number + 1
        seen.add(kpi["id"])
    return kpis if len(kpis) >= MIN_KPIS else []


def _kpi_id(label):
    return f"kpi_{_slug(label)}"


def _total_kpi(plan, frame, column, info):
    name, title = column["name"], frame["title"]
    cells = _cells_of(plan, frame["id"], name) or []
    why = "the source totals this column" if info["totalCell"] else "additive column named like a headline total"
    return {"id": _kpi_id(name), "label": name, "kind": "total", "frame": frame["id"], "column": name,
            "format": _kpi_format(column),
            "expected": _sum_cached(cells), "reason": why,
            "evidence": {"cells": [c[0] for c in cells], "totalCell": info["totalCell"]}}


def _ratio_kpi(plan, frame, ratio):
    column, num, den = ratio
    num_cells = _cells_of(plan, frame["id"], num) or []
    den_cells = _cells_of(plan, frame["id"], den) or []
    denominator = _sum_cached(den_cells)
    return {"id": _kpi_id(column["name"]), "label": column["name"], "kind": "ratio", "frame": frame["id"],
            "column": column["name"], "num": num, "den": den,
            "format": _kpi_format(column),
            "expected": 0 if denominator == 0 else _sum_cached(num_cells) / denominator,
            "reason": f"{column['name']} = {num} / {den}, taken over the totals (not the sum of the row ratios)",
            "evidence": {"cells": [c[0] for c in num_cells] + [c[0] for c in den_cells], "totalCell": None}}


def _kpi_format(column):
    fmt = column.get("format")
    return dict(fmt) if isinstance(fmt, dict) else trm.number_format({}, fmt) or ps.format_for("integer")


def check_kpis(plan, model, tolerance=None):
    """``(lines, problems)``: every KPI re-derived from the oracle's frame rows against the source cells."""
    lines, problems = [], []
    for kpi in (plan.get("presentation") or {}).get("kpis", []):
        frame = _frame(plan, kpi["frame"])
        rows = [r for r in model.frame_rows(kpi["frame"]) if frame["kind"] != "summary"
                or r[frame["by"]] in frame["labels"]]          # the grid's rows, not extra keys in the data
        if kpi["kind"] == "total":
            got = sum(_num(row[kpi["column"]]) for row in rows)
        else:
            top = sum(_num(row[kpi["num"]]) for row in rows)
            bottom = sum(_num(row[kpi["den"]]) for row in rows)
            got = 0 if bottom == 0 else top / bottom
        ok = trm.values_match(got, kpi["expected"], tolerance if tolerance is not None
                              else trm.tolerance_for(kpi["expected"]) * max(1, len(rows)))
        total_cell = kpi["evidence"].get("totalCell")
        if ok and total_cell:
            want = next((cell[2] for group in plan["checks"] for cell in group["cells"] if cell[0] == total_cell),
                        None)
            ok = trm.values_match(got, want, tolerance if tolerance is not None
                                  else trm.tolerance_for(_num(want)) * max(1, len(rows)))
        state = "PASS" if ok else "FAIL"
        lines.append(f"KPI {state} {kpi['label']}: plan={got:,.4f} source={kpi['expected']:,.4f}")
        if not ok:
            problems.append(f"KPI {kpi['label']}: plan {got!r} vs source {kpi['expected']!r}")
    return lines, problems


# ---------------------------------------------------------------- source charts
def _range(ref):
    """Addresses of a chart range such as ``'Forecast'!$B$4:$B$7`` -> ("Forecast", ["B4", ...])."""
    from openpyxl.utils import get_column_letter, range_boundaries
    found = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^!]+))!(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)", ref or "")
    if not found:
        return None
    sheet = (found.group(1) or found.group(2)).replace("''", "'")
    c1, r1, c2, r2 = range_boundaries(found.group(3).replace("$", ""))
    return sheet, [f"{get_column_letter(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def _chart_title(chart):
    try:
        return "".join(run.t for p in chart.title.tx.rich.p for run in (p.r or [])).strip() or None
    except AttributeError:
        return None


def _series_text(wf, series):
    tx = getattr(series, "tx", None)
    if tx is None:
        return None
    if tx.strRef is not None and tx.strRef.f:
        found = _range(tx.strRef.f)
        if found and len(found[1]) == 1:
            value = wf[found[0]][found[1][0]].value
            return str(value) if value is not None else None
    return str(tx.v) if tx.v is not None else None


def _kind_of(chart):
    name = type(chart).__name__
    if name == "BarChart":
        grouping = getattr(chart, "grouping", None) or "clustered"
        return "bar", {"stacked": "stacked", "percentStacked": "normalized"}.get(grouping, "none"), \
            ("horizontal" if getattr(chart, "type", "col") == "bar" else "vertical")
    if name == "LineChart":
        if (getattr(chart, "grouping", None) or "standard") != "standard":
            return None, None, "a stacked line chart is not mapped"
        return "line", "none", "vertical"
    if name == "DoughnutChart":
        return "donut", "none", "vertical"
    if name == "PieChart":
        return "pie", "none", "vertical"
    return None, None, f"{name} is not a mapped chart kind"


def map_charts(plan, wf):
    """Source charts whose categories and values read verified cells of ONE frame -> ``(charts, unmapped)``.

    A chart maps when every category is a row label (or a verified axis cell) of the frame in the frame's own
    order and every series reads one verified frame column at those rows; the plan records the cells so
    ``check_charts`` can re-derive each point. Anything else stays listed as unmapped, with the reason.
    """
    owner = {}
    for group in plan["checks"]:
        for address, sel, value in group["cells"]:
            owner[address] = (group["frame"], group["column"], group["kind"], sel, value)
    label_cell = {}
    for frame in plan["frames"]:
        if frame["kind"] == "summary":
            for index, address in enumerate((frame.get("sourceCells") or {}).get("labels", [])):
                label_cell[address] = (frame["id"], index)
    charts, unmapped, taken = [], [], set()
    for ws in wf.worksheets:
        for source in getattr(ws, "_charts", []):
            title = _chart_title(source)
            kind, stacking, orientation = _kind_of(source)
            if kind is None:
                unmapped.append(_unmapped(title, source, ws, orientation))
                continue
            if len(getattr(source, "_charts", [source])) > 1:
                unmapped.append(_unmapped(title, source, ws, "a combined (multi-plot) chart is not mapped"))
                continue
            try:
                chart = _map_one(plan, wf, ws, source, title or type(source).__name__, kind, stacking,
                                 orientation, owner, label_cell)
            except _NotMapped as reason:
                unmapped.append(_unmapped(title, source, ws, str(reason), plan, owner))
                continue
            base = f"ch_{_slug(chart['title'])}"
            chart["id"], number = base, 2
            while chart["id"] in taken:
                chart["id"], number = f"{base}_{number}", number + 1
            taken.add(chart["id"])
            charts.append(chart)
    return charts, unmapped


class _NotMapped(Exception):
    pass


def _unmapped(title, source, ws, reason, plan=None, owner=None):
    kind = type(source).__name__ + (" stacked" if getattr(source, "grouping", None) == "stacked" else "")
    reads = []
    for series in getattr(source, "series", []):
        ref = series.val.numRef.f if series.val is not None and series.val.numRef is not None else None
        found = _range(ref) if ref else None
        if found and owner:
            frames = {owner.get(f"{found[0]}!{c}", (None, None))[:2] for c in found[1]}
            if len(frames) == 1 and (None, None) not in frames:
                frame, column = next(iter(frames))
                reads.append(f"{frame}.{column}")
    detail = f"; series read verified {', '.join(dict.fromkeys(reads))}" if reads else ""
    return (f"{title or type(source).__name__} ({kind}) on {ws.title!r}{detail}: not emitted ({reason}) - "
            "rebuild from the presentation plan (refs/presentation-migration.md)")


def _map_one(plan, wf, ws, source, title, kind, stacking, orientation, owner, label_cell):
    if not source.series:
        raise _NotMapped("the chart has no series")
    if kind in ("donut", "pie") and len(source.series) != 1:
        raise _NotMapped("a donut / pie chart maps one series")
    cat_refs = {(s.cat.numRef.f if s.cat is not None and s.cat.numRef is not None else
                 s.cat.strRef.f if s.cat is not None and s.cat.strRef is not None else None)
                for s in source.series}
    if len(cat_refs) != 1 or None in cat_refs:
        raise _NotMapped("the series do not share one category range")
    sheet, cat_cells = _range(next(iter(cat_refs))) or (None, None)
    if not cat_cells:
        raise _NotMapped("the category range is not a plain cell range")
    addresses = [f"{sheet}!{cell}" for cell in cat_cells]
    frame_id, category_column, sels, category_kind = None, None, [], None
    if all(a in label_cell for a in addresses):                      # row labels of a summary grid
        frames = {label_cell[a][0] for a in addresses}
        if len(frames) != 1:
            raise _NotMapped("the categories span several grids")
        frame_id = next(iter(frames))
        frame = _frame(plan, frame_id)
        if [label_cell[a][1] for a in addresses] != list(range(len(frame["labels"]))):
            raise _NotMapped("the categories are not all of the grid's rows in the grid's order")
        category_column, category_kind = frame["by"], "summary"
        categories = list(frame["labels"])
        sels = [[label] for label in categories]
    elif all(a in owner and owner[a][2] == "ledger" for a in addresses):    # a verified period axis
        found = {owner[a][:2] for a in addresses}
        if len(found) != 1:
            raise _NotMapped("the categories span several columns")
        frame_id, category_column = next(iter(found))
        category_kind = "ledger"
        sels = [owner[a][3] for a in addresses]
        categories = [owner[a][4] for a in addresses]
        if sels != sorted(sels):
            raise _NotMapped("the periods are not in ascending order")
    else:
        raise _NotMapped("the categories are not row labels or verified period cells of one grid")
    series = []
    for source_series in source.series:
        ref = source_series.val.numRef.f if source_series.val is not None and source_series.val.numRef else None
        found = _range(ref) if ref else None
        if not found or len(found[1]) != len(categories):
            raise _NotMapped("a series does not have one value per category")
        cells = [f"{found[0]}!{c}" for c in found[1]]
        if not all(c in owner for c in cells):
            missing = next(c for c in cells if c not in owner)
            raise _NotMapped(f"{missing} is not a verified frame cell")
        owners = {owner[c][:3] for c in cells}
        if len(owners) != 1 or next(iter(owners))[0] != frame_id:
            raise _NotMapped("a series does not read one verified column of the categories' grid")
        _, column, cell_kind = next(iter(owners))
        if [owner[c][3] for c in cells] != sels:
            raise _NotMapped("the series values are not in the categories' order")
        name = _series_text(wf, source_series) or column
        series.append({"title": name, "column": column, "kind": cell_kind, "cells": cells, "sels": sels,
                       "values": [_plain(owner[c][4]) for c in cells]})
    return {"title": title, "sheet": ws.title, "kind": kind, "orientation": orientation, "stacking": stacking,
            "frame": frame_id, "categoryColumn": category_column, "categoryKind": category_kind,
            "categories": [_plain(v) for v in categories], "categoryCells": addresses, "series": series}


def _plain(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def check_charts(plan, model, tolerance=None):
    """``(lines, problems)``: every chart point must equal the frame cell the oracle computes for it."""
    lines, problems = [], []
    for chart in (plan.get("presentation") or {}).get("charts", []):
        points = bad = 0
        for series in chart["series"]:
            for address, sel, want in zip(series["cells"], series["sels"], series["values"]):
                points += 1
                try:
                    got = model.cell(chart["frame"], series["column"], series["kind"], sel)
                except (KeyError, IndexError) as error:
                    got = None
                    problems.append(f"chart {chart['title']!r} {address}: {error}")
                    bad += 1
                    continue
                if not trm.values_match(got, want, tolerance):
                    bad += 1
                    problems.append(f"chart {chart['title']!r} {address}: plan {got!r} vs source {want!r}")
        state = "PASS" if not bad else "FAIL"
        lines.append(f"CHART {state} {chart['title']}: {points - bad}/{points} points equal verified "
                     f"{chart['frame']} cells")
    return lines, problems


class EmittedEvaluator:
    """Evaluate the EMITTED KPI / chart formulas (Sigma NULL logic) over the plan's base rows."""

    def __init__(self, plan):
        self.plan = plan
        self.model = se.EmittedModel(plan)
        self.eval = se.Evaluator(self.model._lookup)
        self.frames = self.model.parts["frames"]

    def rows(self, element_id):
        return self.model.source_rows({"source": {"kind": "table", "elementId": element_id}})

    def _ctx(self, rows):
        return se.Ctx(lambda ref: None, rows, lambda src: se.Ctx(self.model._ref(src, {})))

    def aggregate(self, element_id, formula):
        """The formula over every row of the element (a KPI)."""
        return self.eval.ev(se.parse(formula), self._ctx(self.rows(element_id)))

    def grouped(self, element_id, key_formula, formula):
        """``{key: value}``: the formula per group of the key column (a chart series)."""
        groups = {}
        for src in self.rows(element_id):
            key = self.eval.ev(se.parse(key_formula), se.Ctx(self.model._ref(src, {})))
            groups.setdefault(key, []).append(src)
        return {key: self.eval.ev(se.parse(formula), self._ctx(rows)) for key, rows in groups.items()}

    def kpi(self, kpi):
        """The KPI's value from its emitted terms (each term over its own source element)."""
        top, bottom = kpi_terms(kpi, self.frames)

        def side(terms):
            return sum(t["sign"] * (self.aggregate(t["element"], t["formula"]) or 0) for t in terms)
        value = side(top)
        if bottom is None:
            return value
        denominator = side(bottom)
        return 0 if denominator == 0 else value / denominator


def check_emitted(plan, model, tolerance=None):
    """Evaluate the EMITTED KPI and chart formulas over the plan's base rows (``sigma_formula_eval``) and
    compare them with the oracle's frame totals and cells. Returns ``(lines, problems, unverified)``;
    ``unverified`` names what the evaluator cannot model (running sums), which is reported, not guessed."""
    presentation = plan.get("presentation") or {}
    lines, problems, unverified = [], [], []
    if not presentation.get("kpis") and not presentation.get("charts"):
        return lines, problems, unverified
    parts = trm.build_parts(plan, "<dm>", trm.local_dm_elements(plan))
    page = compose_report(plan, parts)
    elements = {e["id"]: e for e in page["elements"]}
    try:
        evaluator = EmittedEvaluator(plan)
    except Exception as error:      # noqa: BLE001 - the evaluator could not even load the plan
        return lines, problems, [f"all ({type(error).__name__}: {error})"]
    for kpi in presentation.get("kpis", []):
        label = f"KPI {kpi['label']}"
        frame = _frame(plan, kpi["frame"])
        rows = [r for r in model.frame_rows(kpi["frame"])
                if frame["kind"] != "summary" or r[frame["by"]] in frame["labels"]]
        if kpi["kind"] == "total":
            want = sum(_num(r[kpi["column"]]) for r in rows)
        else:
            bottom = sum(_num(r[kpi["den"]]) for r in rows)
            want = 0 if bottom == 0 else sum(_num(r[kpi["num"]]) for r in rows) / bottom
        try:
            got = evaluator.kpi(kpi)
            card = elements.get(kpi["id"])
            if card and card["kind"] == "kpi-chart":     # the formula text actually emitted on the card
                emitted = evaluator.aggregate(card["source"]["elementId"], card["columns"][0]["formula"])
                if not trm.values_match(emitted, got, trm.tolerance_for(got) * 10):
                    problems.append(f"{label}: the card's formula gives {emitted!r}, its terms give {got!r}")
        except (se.Unsupported, RecursionError):
            unverified.append(label)
            continue
        except Exception as error:      # noqa: BLE001
            problems.append(f"{label}: the emitted formula does not evaluate ({error})")
            continue
        scale = tolerance if tolerance is not None else trm.tolerance_for(want) * max(1, len(rows))
        for reference, name in ((want, "oracle"), (kpi["expected"], "source")):
            if not trm.values_match(got, reference, scale):
                problems.append(f"{label}: emitted {got!r} vs {name} {reference!r}")
                break
        else:
            lines.append(f"EMITTED PASS {label}: emitted={got:,.4f} oracle={want:,.4f}")
    for chart in presentation.get("charts", []):
        label = f"chart {chart['title']}"
        element = elements.get(chart["id"])
        if element is None:
            continue
        names = {c["id"]: c for c in element["columns"]}
        key = names[element["xAxis"]["columnId"] if "xAxis" in element else element["color"]["columnId"]]
        y_ids = element["yAxis"]["columnIds"] if "yAxis" in element else [element["value"]["columnId"]]
        points = bad = 0
        try:
            for y_id, series in zip(y_ids, chart["series"]):
                got = evaluator.grouped(element["source"]["elementId"], key["formula"], names[y_id]["formula"])
                for sel, category, want in zip(series["sels"], chart["categories"], series["values"]):
                    wanted = trm.parse_date(category) if chart["categoryKind"] == "ledger" else sel[0]
                    points += 1
                    oracle = model.cell(chart["frame"], series["column"], series["kind"], sel)
                    value = got.get(wanted)
                    if value is None or not (trm.values_match(value, want, tolerance)
                                             and trm.values_match(value, oracle, tolerance)):
                        bad += 1
                        problems.append(f"{label}: {series['title']} at {category}: emitted {value!r} "
                                        f"vs verified {want!r}")
        except (se.Unsupported, RecursionError):
            unverified.append(label)
            continue
        except Exception as error:      # noqa: BLE001
            problems.append(f"{label}: the emitted series do not evaluate ({error})")
            continue
        periods = chart_periods(plan, chart)
        if periods is not None:
            flt = (element.get("filters") or [{}])[0]
            if sorted(flt.get("values") or []) != sorted(periods):
                bad += 1
                problems.append(f"{label}: its period filter {flt.get('values')!r} is not the chart's {periods!r}")
        if not bad:
            lines.append(f"EMITTED PASS {label}: {points} points equal the verified cells")
    return lines, problems, unverified


def check_presentation(plan, model, tolerance=None):
    """Local presentation parity: KPI totals and chart points. ``(lines, problems)``."""
    lines, problems = check_kpis(plan, model, tolerance)
    more_lines, more_problems = check_charts(plan, model, tolerance)
    return lines + more_lines, problems + more_problems


# ---------------------------------------------------------------- drafter side
def propose(plan, wf, source_path, model, sheet_title=None):
    """The plan's ``presentation`` block plus the list of source charts that stay unmapped."""
    frame = main_frame(plan)
    title = (frame or {}).get("sheetTitle") or sheet_title or plan["name"]
    fill = next(((f.get("headerStyle") or {}).get("fill") for f in [frame] if f), None)
    accent = ps.accent_from_fill(fill)
    modified = getattr(wf.properties, "modified", None)
    parts = [source_path.name] + ([f"{frame['sheet']} sheet"] if frame else [])
    if modified:
        parts.append(f"as of {modified.date().isoformat()}")
    charts, unmapped = map_charts(plan, wf)
    presentation = {"version": 1, "title": title, "subtitle": " · ".join(parts), "accent": accent,
                    "main": frame["id"] if frame else None, "kpis": propose_kpis(plan), "charts": charts,
                    # rows each grouped / ledger table shows (extra group keys included): sizes its card
                    "rows": {f["id"]: len(model.frame_rows(f["id"])) for f in plan["frames"]
                             if f["kind"] in ("summary", "ledger") and not f.get("hidden")},
                    # frames whose source has keys the grid does not list: KPIs / charts keep to the labels
                    "restrict": {f["id"]: list(f["labels"]) for f in plan["frames"]
                                 if f["kind"] == "summary" and not f.get("hidden")
                                 and len(model.frame_rows(f["id"])) > len(f["labels"])}}
    draft_plan = dict(plan, presentation=presentation)
    notes = drop_unexpressible(draft_plan, model)
    unmapped += [n["message"] for n in notes if n["kind"] == "chart"]
    presentation["notes"] = notes
    return presentation, unmapped


def drop_unexpressible(plan, model):
    """Drop every KPI / chart that cannot be emitted from the frame's source columns, or whose EMITTED formulas
    do not evaluate to the oracle's values; ``plan["presentation"]`` is edited in place. Returns the notes
    ``[{code, kind, message}]`` (each becomes a review issue; a dropped chart also stays in ``unmappedCharts``)."""
    presentation = plan["presentation"]
    parts = trm.build_parts(plan, "<dm>", trm.local_dm_elements(plan))
    frames = parts["frames"]
    notes = []

    def drop(kind, label, why):
        notes.append({"code": "PRESENTATION_ELEMENT_DROPPED", "kind": kind,
                      "message": f"{kind} {label!r}: {why}"})
    kept = []
    for kpi in presentation["kpis"]:
        pair = kpi_terms(kpi, frames)
        if pair is None:
            feed = (frames.get(kpi["frame"]) or {}).get("feed") or {}
            names = [kpi.get("column")] if kpi["kind"] == "total" else [kpi.get("num"), kpi.get("den")]
            why = next((feed.get("why", {}).get(n) for n in names if feed.get("why", {}).get(n)),
                       "its terms span several tables and text cards are off")
            drop("KPI", kpi["label"], f"not emitted ({why})")
            continue
        sources = {t["element"] for side in pair if side for t in side}
        formula = kpi_formula(pair)
        if len(sources) > 1 and ("}}" in formula or "|" in formula):
            drop("KPI", kpi["label"], "not emitted (a label holds '}}' or '|', which breaks the {{ }} text "
                                      "interpolation of a card spanning several tables)")
            continue
        kept.append(kpi)
    charts = []
    for chart in presentation["charts"]:
        why = chart_gap(chart, frames, plan)
        if why:
            notes.append({"code": "PRESENTATION_ELEMENT_DROPPED", "kind": "chart",
                          "message": f"{chart['title']} ({chart['kind']}) on {chart['sheet']!r}: not emitted "
                                     f"({why}) - rebuild from the presentation plan "
                                     "(refs/presentation-migration.md)"})
        else:
            charts.append(chart)
    presentation["kpis"], presentation["charts"] = kept, charts
    # the emitted formulas themselves, evaluated over the plan's base rows, against the oracle
    _, problems, _ = check_emitted(plan, model)
    failure_prefixes = {
        **{f"KPI {kpi['label']}:": [] for kpi in presentation["kpis"]},
        **{f"chart {chart['title']}:": [] for chart in presentation["charts"]},
    }
    for problem in problems:
        matches = [prefix for prefix in failure_prefixes if problem.startswith(prefix)]
        if matches:
            failure_prefixes[max(matches, key=len)].append(problem)
    for kpi in list(presentation["kpis"]):
        prefix = f"KPI {kpi['label']}:"
        if failure_prefixes[prefix]:
            presentation["kpis"].remove(kpi)
            drop("KPI", kpi["label"], "not emitted (its emitted formula does not match the oracle: "
                 + failure_prefixes[prefix][0][len(prefix):].lstrip() + ")")
    for chart in list(presentation["charts"]):
        if failure_prefixes[f"chart {chart['title']}:"]:
            presentation["charts"].remove(chart)
            notes.append({"code": "PRESENTATION_ELEMENT_DROPPED", "kind": "chart",
                          "message": f"{chart['title']} ({chart['kind']}) on {chart['sheet']!r}: not emitted "
                                     "(its emitted series do not match the verified cells) - rebuild from the "
                                     "presentation plan (refs/presentation-migration.md)"})
    if 0 < len(presentation["kpis"]) < MIN_KPIS:
        presentation["kpis"] = []
    return notes


def ledger_periods(plan, frame):
    """The period values of a ledger's spine, in order."""
    table = trm.plan_table(plan, frame["spine"])
    names = [c["name"] for c in trm.literal_columns(table)]
    index = names.index(frame["axis"]) if frame["axis"] in names else None
    return [row[index] for row in table.get("rows") or []] if index is not None else []


def chart_periods(plan, chart):
    """``None`` when a ledger chart covers every period, else the periods it covers (it then needs a filter)."""
    if chart["categoryKind"] != "ledger":
        return None
    frame = _frame(plan, chart["frame"])
    mine = [s[0] for s in chart["series"][0]["sels"]]
    return None if sorted(mine) == sorted(ledger_periods(plan, frame)) else mine


def chart_gap(chart, frames, plan=None):
    """Why a mapped chart cannot be emitted from its frame's source columns (``None`` when it can)."""
    feed = (frames.get(chart["frame"]) or {}).get("feed")
    if not feed:
        return "its frame has no source-level feed"
    if plan is not None and chart_periods(plan, chart) is not None and not ps.SURFACES["chart_list_filter"]:
        return "it covers only some of the ledger's periods and the list filter that would limit it is off"
    missing = [s["column"] for s in chart["series"] if not feed["series"].get(s["column"])]
    if missing:
        return f"{', '.join(missing)} cannot be computed from the frame's source columns"
    if chart["categoryKind"] == "summary" and not feed.get("key"):
        return "the frame has no key column"
    return None


# ---------------------------------------------------------------- KPI formulas from source columns
def _combine(terms):
    """Signed sum of leaf-term formulas as one Sigma expression."""
    if len(terms) == 1:
        return terms[0]["formula"] if terms[0]["sign"] > 0 else f"(-({terms[0]['formula']}))"
    text = ""
    for number, term in enumerate(terms):
        sign = "-" if term["sign"] < 0 else "+"
        text += (f"-({term['formula']})" if sign == "-" else f"({term['formula']})") if number == 0 \
            else f" {sign} ({term['formula']})"
    return f"({text})"


def kpi_terms(kpi, frames):
    """``(numerator terms, denominator terms | None)`` of a KPI, emitted from the frame's source columns, or
    ``None`` when it cannot be (or when it spans several sources and text cards are switched off)."""
    feed = (frames.get(kpi["frame"]) or {}).get("feed")
    if not feed:
        return None
    if kpi["kind"] == "total":
        pair = (feed["terms"].get(kpi["column"]), None)
        ok = bool(pair[0])
    else:
        pair = (feed["terms"].get(kpi["num"]), feed["terms"].get(kpi["den"]))
        ok = bool(pair[0]) and bool(pair[1])
    if not ok:
        return None
    sources = {t["element"] for side in pair if side for t in side}
    if len(sources) > 1 and not ps.SURFACES["kpi_text_card"]:
        return None
    return pair


def kpi_formula(pair):
    """The KPI's value as one Sigma formula over its terms (a ratio is over the totals, 0 for a blank / zero
    denominator)."""
    top, bottom = pair
    if bottom is None:
        return _combine(top)
    num, den = _combine(top), _combine(bottom)
    return f"If(IsNull({den}) or {den} = 0, 0, {num} / {den})"


# ---------------------------------------------------------------- Report page composition
def _frame_format(plan, frame_id, column):
    column_def = _column(_frame(plan, frame_id), column) or {}
    fmt = column_def.get("format")
    return trm.number_format(plan, fmt) if fmt else None


def _chart_elements(plan, chart, meta, accent):
    """One Sigma chart element for a mapped chart, computed from the frame's SOURCE columns.

    The chart groups by the key column of the frame's source and aggregates each series with the frame's own
    formula (``Coalesce(SumIf([src/Amount], <window>), 0)``); a grouped table's grouping-level calculations
    are not addressable from another element.
    """
    feed = meta["feed"]
    frame = _frame(plan, chart["frame"])
    cid, label = chart["id"], feed["label"]
    source = {"kind": "table", "elementId": feed["element"]}
    x_id = f"{cid}_x"
    if chart["categoryKind"] == "summary":
        x_column = {"id": x_id, "name": chart["categoryColumn"], "formula": feed["key"]}
    else:
        x_column = {"id": x_id, "name": chart["categoryColumn"],
                    "formula": f"[{label}/{chart['categoryColumn']}]"}
    cat_format = _frame_format(plan, chart["frame"], chart["categoryColumn"])
    if cat_format and cat_format.get("kind") == "datetime":
        x_column["format"] = cat_format
    columns = [x_column]
    sort = {"direction": "ascending"}
    labels = trm.order_labels(frame)
    if chart["categoryKind"] == "summary" and labels is not None:
        o_column = {"id": f"{cid}_o", "name": trm.ORDER_NAME,
                    "formula": trm.order_formula(feed["key"], labels, feed.get("keyType", "text"))}
        if ps.SURFACES["chart_hidden_column"]:
            o_column["hidden"] = True
        columns.append(o_column)
        sort = {"by": o_column["id"], "direction": "ascending", "aggregation": "min"}
    y_ids, used = [], {x_column["name"], trm.ORDER_NAME}
    for number, series in enumerate(chart["series"], 1):
        name, bump = series["title"], 2
        while name in used:
            name, bump = f"{series['title']} ({bump})", bump + 1
        used.add(name)
        y = {"id": f"{cid}_y{number}", "name": name, "formula": feed["series"][series["column"]]}
        fmt = _frame_format(plan, chart["frame"], series["column"])
        if fmt:
            y["format"] = fmt
        columns.append(y)
        y_ids.append(y["id"])
    element = {"id": cid, "name": chart["title"], "source": source, "columns": columns}
    look = (plan.get("presentation") or {}).get("look")
    style = ps.chart_card(look)
    if style:
        element["style"] = style
    restrict = (plan.get("presentation") or {}).get("restrict", {}).get(chart["frame"])
    if restrict and chart["categoryKind"] == "summary" and ps.SURFACES["chart_list_filter"]:
        element["filters"] = [{"id": f"{cid}_f", "columnId": x_id, "kind": "list", "mode": "include",
                               "values": list(restrict)}]
    periods = chart_periods(plan, chart)
    if periods is not None and ps.SURFACES["chart_list_filter"]:
        # a chart over some of a ledger's periods plots exactly those (the table has a row per period)
        axis = frame["axis"]
        if chart["categoryColumn"] == axis:
            p_id = x_id
        else:
            p_id = f"{cid}_p"
            p_column = {"id": p_id, "name": axis, "formula": f"[{label}/{axis}]"}
            if ps.SURFACES["chart_hidden_column"]:
                p_column["hidden"] = True
            columns.append(p_column)
        element["filters"] = [{"id": f"{cid}_f", "columnId": p_id, "kind": "list", "mode": "include",
                               "values": periods}]
    if chart["kind"] in ("donut", "pie"):
        element["kind"] = f"{chart['kind']}-chart"
        element["value"] = {"columnId": y_ids[0]}
        element["color"] = {"columnId": x_id, "sort": sort}
        element["legend"] = {"position": "right"}
        return element
    element["xAxis"] = {"columnId": x_id, "sort": sort}
    element["legend"] = {"visibility": "hidden"} if len(y_ids) == 1 else {"position": "bottom"}
    if len(y_ids) == 1 and look is None and ps.single_color(accent["hex"]):   # a look's theme colours the series
        element["color"] = ps.single_color(accent["hex"])
    if chart["kind"] == "line":
        element["kind"] = "line-chart"
        element["yAxis"] = {"columnIds": y_ids, "format": {"marks": "grid",
                                                           "scale": {"type": "linear", "zero": False}}}
        element["lineAreaStyle"] = {"interpolation": "monotone"}
    else:
        element["kind"] = "bar-chart"
        if len(chart["categories"]) <= 12 and ps.SURFACES["chart_label_angle"]:      # flat category labels
            axis_format = element["xAxis"].get("format") or {}
            element["xAxis"]["format"] = {**axis_format, "labels": {**(axis_format.get("labels") or {}),
                                                                    "labelAngle": 0}}
        element["yAxis"] = {"columnIds": y_ids, "format": {"marks": "grid"}}
        element["stacking"] = chart["stacking"]
        if chart["orientation"] == "horizontal":
            element["orientation"] = "horizontal"
    return element


def _visible_columns(element):
    return [c for c in element.get("columns", []) if not c.get("hidden")]


def compose_report(plan, parts):
    """The Report page: header band, KPI row, hero charts, section label, tables sized to their rows.

    Returns ``{"elements": [...], "placements": [(id, c0, c1, r0, r1)], "settings": {...}}``.
    """
    pres = plan.get("presentation") or {}
    accent = pres.get("accent") or ps.accent_from_fill(None)
    elements, placements = [], []
    cursor = 1

    def place(element, c0, c1, height, row=None):
        nonlocal cursor
        start = cursor if row is None else row
        elements.append(element)
        placements.append((element["id"], c0, c1, start, start + height))
        return start + height

    look = pres.get("look")           # a reference workbook's look: theme / card styles / KPI type, never our colours
    band, header = ps.header_band("trHeader", pres.get("title") or plan["name"], pres.get("subtitle"), accent,
                                  look)
    elements.extend([band, header])
    # the text sits inside its band (child rows are relative to the container's own grid)
    placements.append((band["id"], 1, ps.PAGE_COLS + 1, 1, 1 + ps.HEADER_ROWS,
                       [(header["id"], 1, ps.PAGE_COLS + 1, 1, 1 + ps.HEADER_ROWS)]))
    cursor = 1 + ps.HEADER_ROWS
    frames = parts.get("frames", {})
    cards = []
    for kpi in pres.get("kpis", []):
        pair = kpi_terms(kpi, frames)
        if pair is not None:
            cards.append((kpi, pair))
    cards = cards[:ps.PAGE_COLS // 6]
    kpis = [k for k, _ in cards]
    if cards:
        width = ps.PAGE_COLS // len(cards)
        top = cursor
        for number, (kpi, pair) in enumerate(cards):
            formula = kpi_formula(pair)
            sources = sorted({t["element"] for side in pair if side for t in side})
            c0, c1 = 1 + number * width, 1 + (number + 1) * width
            if len(sources) == 1:
                element = ps.kpi_card(kpi["id"], kpi["label"].upper(), {"kind": "table", "elementId": sources[0]},
                                      formula, kpi.get("format"), look)
                end = place(element, c0, c1, ps.KPI_ROWS, row=top)
            else:      # a kpi-chart has one source: a card whose text interpolates the cross-element formula
                card, text = ps.kpi_text_card(kpi["id"], kpi["label"].upper(), formula, kpi.get("format"), look)
                elements.extend([card, text])
                placements.append((card["id"], c0, c1, top, top + ps.KPI_ROWS,
                                   [(text["id"], 1, ps.PAGE_COLS + 1, 1, 1 + ps.KPI_ROWS)]))
                end = top + ps.KPI_ROWS
        cursor = end
    charts = [c for c in pres.get("charts", []) if c["frame"] in frames and not chart_gap(c, frames, plan)]
    detail = parts["visible"]
    by_id = {m["id"]: m for m in parts.get("visibleMeta", [])}

    def bound(meta, frame_id):
        return meta.get("frame") == frame_id or frame_id in meta.get("bindings", ())

    main = pres.get("main")
    detail = sorted(detail, key=lambda e: 0 if bound(by_id.get(e["id"], {}), main) else 1)
    hero = next((c for c in charts if c["kind"] not in ("donut", "pie")), charts[0] if charts else None)
    # a supporting chart whose categories are a grid's row labels sits beside that grid's table
    beside = {}
    for chart in charts:
        if chart is hero or chart["categoryKind"] != "summary":
            continue
        table = next((e for e in detail if e["kind"] == "table" and e["id"] not in beside
                      and bound(by_id.get(e["id"], {}), chart["frame"])
                      and by_id[e["id"]]["role"] == "grouped"), None)
        if table is not None:
            beside[table["id"]] = chart
    if hero is not None:
        loose = [c for c in charts if c is not hero and c not in beside.values()]
        top = cursor
        end = place(_chart_elements(plan, hero, frames[hero["frame"]], accent), 1, ps.PAGE_COLS + 1,
                    ps.CHART_ROWS, row=top)
        for index in range(0, len(loose), 2):
            row = loose[index:index + 2]
            splits = [(1, 25)] if len(row) == 1 else [(1, 13), (13, 25)]
            top = end                      # both charts of a pair start together and end together
            for chart, (c0, c1) in zip(row, splits):
                place(_chart_elements(plan, chart, frames[chart["frame"]], accent), c0, c1, ps.CHART_ROWS,
                      row=top)
            end = top + ps.CHART_ROWS
        cursor = end
    if detail and (kpis or charts):
        cursor = place(ps.divider("trDivider"), 1, ps.PAGE_COLS + 1, ps.DIVIDER_ROWS)
        cursor = place(ps.section_label("trDetail", "Detail", look), 1, ps.PAGE_COLS + 1, ps.LABEL_ROWS)
    negative = {f["id"]: {c["name"] for c in f["columns"] if c.get("negativeRed")} for f in plan["frames"]}
    for element in detail:
        meta = by_id.get(element["id"], {"frame": None, "role": "table", "rows": 8})
        frame = _frame(plan, meta["frame"]) if meta["frame"] else {}
        negative_names = set().union(*(negative.get(frame_id, set())
                                       for frame_id in [meta["frame"]] + meta.get("bindings", [])))
        _style_table(element, frame, negative_names, look)
        wide = len(_visible_columns(element)) > 4 or element["kind"] == "pivot-table"
        rows = meta["rows"]
        if meta["role"] in ("grouped", "ledger"):
            rows = (pres.get("rows") or {}).get(meta["frame"], rows)
        height = ps.table_height(rows, extra_header_rows=1 if element["kind"] == "pivot-table" else 0,
                                 roomy=ps.roomy_tables(look))
        if element["id"] in beside:                       # table 1 / 15, its chart 15 / 25, the taller height
            chart = beside[element["id"]]
            height = max(height, ps.CHART_ROWS)
            place(_chart_elements(plan, chart, frames[chart["frame"]], accent), 15, ps.PAGE_COLS + 1, height)
            cursor = place(element, 1, 15, height, row=cursor)
        else:
            cursor = place(element, 1, ps.PAGE_COLS + 1 if wide else 15, height)
    settings = ps.settings_theme(accent["hex"], look) if (charts or look is not None) else {}
    return {"elements": elements, "placements": placements, "settings": settings}


def _style_table(element, frame, negative_names, look=None):
    """Card chrome, source header style and Excel's red negatives on a Report-page table / pivot.

    With a reference ``look`` the card is the reference's table style (none when it has none) and the reference's
    table preset joins Excel's header fill / bold in ``tableStyle``."""
    if look is not None:
        element.pop("style", None)
        card = ps.look_style(look, "table")
        if card:
            element["style"] = card
    else:
        element["style"] = ps.card_style()
    style = ps.table_style(frame.get("headerStyle"), element["kind"], look)
    if style:
        element["tableStyle"] = style
    ids = [c["id"] for c in element.get("columns", []) if c.get("name") in negative_names]
    rule = ps.negative_red_rule(ids)
    if rule:
        element["conditionalFormats"] = [rule]
