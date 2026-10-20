"""Read-only, source-bound Excel workbook inventory for later migration review."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.formula import ArrayFormula

from excel_expression import (
    Binary, ExpressionError, Function, Group, Percent, Range, Reference,
    Unary, expand_range, parse_expression,
)


_ROLES = {"input", "calculation", "reference", "scratch"}
_ADDRESS = re.compile(r"^[A-Z]{1,3}[1-9][0-9]*$")
_SHA = re.compile(r"^[0-9a-f]{64}$")
_MAX_EXPANSION = 100_000


def _id(kind, *parts):
    return kind + ":" + json.dumps(parts, ensure_ascii=False, separators=(",", ":"))


def _cell_id(sheet, address):
    return _id("cell", sheet, address)


def _hash(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _literal(value, number_format=None):
    if isinstance(value, datetime):
        # OOXML dates reload as datetime at midnight; the format preserves date-only intent.
        if value.time() == time.min and number_format and not any(token in number_format.lower() for token in ("h", "s")):
            return {"type": "date", "iso": value.date().isoformat()}
        return {"type": "datetime", "iso": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "iso": value.isoformat()}
    if isinstance(value, time):
        return {"type": "time", "iso": value.isoformat()}
    if isinstance(value, timedelta):
        return {"type": "duration", "seconds": value.total_seconds()}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    raise ValueError(f"unsupported cell literal type {type(value).__name__}")


def _formula_source(value):
    if isinstance(value, str):
        return value, None, None
    kind = type(value).__name__
    if isinstance(value, ArrayFormula):
        return value.text, kind, value.ref
    return None, kind, getattr(value, "ref", None)


def _bounds(source_range):
    if not isinstance(source_range, str) or not source_range or "!" in source_range:
        raise ValueError(f"invalid bounded range: {source_range!r}")
    try:
        left, top, right, bottom = range_boundaries(source_range)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"invalid bounded range: {source_range!r}") from exc
    if any(part is None for part in (left, top, right, bottom)):
        raise ValueError(f"unbounded range: {source_range!r}")
    if not (1 <= left <= right <= 16384 and 1 <= top <= bottom <= 1048576):
        raise ValueError(f"invalid bounded range: {source_range!r}")
    return left, top, right, bottom


def _range_string(left, top, right, bottom):
    return f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"


def _walk_references(node):
    if isinstance(node, Reference):
        yield node.value
    elif isinstance(node, Range):
        bounds = _bounds(_range_string(*_reference_bounds(node)))
        if (bounds[2] - bounds[0] + 1) * (bounds[3] - bounds[1] + 1) > _MAX_EXPANSION:
            raise ValueError("range expansion exceeds inventory limit")
        yield from expand_range(node)
    elif isinstance(node, Group):
        yield from _walk_references(node.expression)
    elif isinstance(node, (Unary, Percent)):
        yield from _walk_references(node.operand)
    elif isinstance(node, Binary):
        yield from _walk_references(node.left)
        yield from _walk_references(node.right)
    elif isinstance(node, Function):
        for argument in node.arguments:
            yield from _walk_references(argument)


def _reference_bounds(node):
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    if node.start.sheet != node.end.sheet:
        raise ValueError("range crosses sheets")
    start_col, start_row = coordinate_from_string(node.start.address)
    end_col, end_row = coordinate_from_string(node.end.address)
    return (min(column_index_from_string(start_col), column_index_from_string(end_col)),
            min(start_row, end_row),
            max(column_index_from_string(start_col), column_index_from_string(end_col)),
            max(start_row, end_row))


def _defined_names(workbook):
    result = []
    lookup = {}
    scopes = [(None, workbook.defined_names)] + [(sheet.title, sheet.defined_names) for sheet in workbook.worksheets]
    for scope, container in scopes:
        for defined in sorted(container.values(), key=lambda item: item.name):
            expression = defined.attr_text or ""
            record = {"id": _id("name", scope, defined.name), "name": defined.name,
                      "scope": scope, "expression": expression, "resolution": "unresolved",
                      "targetId": None, "reason": None}
            try:
                parsed = parse_expression(expression if expression.startswith("=") else "=" + expression,
                                          sheet=scope)
                if isinstance(parsed, Reference) and parsed.value.name is None:
                    target_sheet = parsed.value.sheet or scope
                    if target_sheet not in workbook.sheetnames:
                        record["reason"] = f"missing sheet {target_sheet!r}"
                    else:
                        record["resolution"] = "cell"
                        record["targetId"] = _cell_id(target_sheet, parsed.value.address)
                else:
                    record["reason"] = "defined name is not a single-cell reference"
            except (ExpressionError, ValueError) as exc:
                record["reason"] = str(exc)
            result.append(record)
            lookup[(scope, defined.name.lower())] = record
    return result, lookup


def _dependencies(expression, sheet, names, sheet_names, populated, formula_object_type=None):
    dependencies = {}
    reasons = []
    if expression is None:
        return [], [f"{formula_object_type or 'formula object'} has no scalar source expression"]
    try:
        parsed = parse_expression(expression, sheet=sheet)
        references = tuple(_walk_references(parsed))
    except (ExpressionError, ValueError) as exc:
        return [], [str(exc)]
    for ref in references:
        if ref.name:
            name = names.get((sheet, ref.name.lower())) or names.get((None, ref.name.lower()))
            if name is None:
                reasons.append(f"unresolved defined name {ref.name!r}")
                continue
            if name["resolution"] != "cell":
                reasons.append(f"defined name {ref.name!r}: {name['reason']}")
                continue
            target_id = name["targetId"]
            target_sheet, address = json.loads(target_id[len("cell:"):])
        else:
            target_sheet = ref.sheet or sheet
            address = ref.address
            target_id = _cell_id(target_sheet, address)
        if target_sheet not in sheet_names:
            reasons.append(f"missing sheet {target_sheet!r}")
            continue
        dependencies[target_id] = {"targetId": target_id, "sheet": target_sheet,
                                   "address": address,
                                   "targetKind": "populated" if target_id in populated else "blank"}
    return list(dependencies.values()), sorted(set(reasons))


def _table_model(sheet, table):
    left, top, right, bottom = _bounds(table.ref)
    header_count = table.headerRowCount if table.headerRowCount is not None else 1
    totals_count = table.totalsRowCount or 0
    body_top = top + header_count
    body_bottom = bottom - totals_count
    headers = [column.name for column in table.tableColumns]
    if len(headers) != right - left + 1:
        headers = [str(sheet.cell(top, col).value or "") for col in range(left, right + 1)]
    columns = []
    for offset, header in enumerate(headers):
        coordinate_col = get_column_letter(left + offset)
        record = {"header": header, "index": offset + 1, "literalCells": [],
                  "formulaCells": [], "blankCells": [], "intent": "ambiguous"}
        for row in range(body_top, body_bottom + 1):
            cell = sheet.cell(row, left + offset)
            field = ("formulaCells" if cell.data_type == "f" else
                     "literalCells" if cell.value is not None or cell.data_type == "inlineStr" else "blankCells")
            record[field].append(f"{coordinate_col}{row}")
        columns.append(record)
    return {"id": _id("table", sheet.title, table.name), "sheet": sheet.title,
            "name": table.name, "range": _range_string(left, top, right, bottom),
            "headers": headers,
            "headerRange": _range_string(left, top, right, top + header_count - 1) if header_count else None,
            "bodyRange": _range_string(left, body_top, right, body_bottom) if body_top <= body_bottom else None,
            "totalsRange": _range_string(left, body_bottom + 1, right, bottom) if totals_count else None,
            "columns": columns}


def _reviewed_blocks(decisions, sha, sheet_names):
    if decisions is None:
        return []
    if isinstance(decisions, (str, Path)):
        decisions = json.loads(Path(decisions).read_text(encoding="utf-8"))
    if not isinstance(decisions, dict) or decisions.get("sourceSha256") != sha or not isinstance(decisions.get("blocks"), list):
        raise ValueError("block decisions must bind to the current source SHA256")
    result = []
    by_sheet = {}
    for decision in decisions["blocks"]:
        if not isinstance(decision, dict) or decision.get("sheet") not in sheet_names or decision.get("role") not in _ROLES:
            raise ValueError("block decision has unknown sheet or role")
        bounds = _bounds(decision.get("range"))
        for previous in by_sheet.get(decision["sheet"], []):
            if bounds[0] <= previous[2] and previous[0] <= bounds[2] and bounds[1] <= previous[3] and previous[1] <= bounds[3]:
                raise ValueError("reviewed block ranges overlap")
        by_sheet.setdefault(decision["sheet"], []).append(bounds)
        source_range = _range_string(*bounds)
        result.append({"id": _id("block", "reviewed", decision["sheet"], source_range),
                       "sheet": decision["sheet"], "range": source_range,
                       "role": decision["role"], "reviewed": True,
                       "sourceSha256": sha})
    return result


def build_source_model(source_path, *, block_decisions=None) -> dict:
    """Inventory source facts without saving, evaluating formulas, or running VBA."""
    source_path = Path(source_path)
    source_bytes = source_path.read_bytes()
    sha = hashlib.sha256(source_bytes).hexdigest()
    workbook = load_workbook(BytesIO(source_bytes), data_only=False, read_only=False, keep_links=True)
    try:
        sheets = []
        cells = []
        tables = []
        blocks = []
        sheet_names = set(workbook.sheetnames)
        for sheet in workbook.worksheets:
            sheets.append({"id": _id("sheet", sheet.title), "name": sheet.title,
                           "state": sheet.sheet_state, "protected": bool(sheet.protection.sheet),
                           "mergedRanges": sorted(str(r) for r in sheet.merged_cells.ranges),
                           "chartCount": len(sheet._charts),
                           "validationRuleCount": len(sheet.data_validations.dataValidation),
                           "conditionalFormattingCount": len(sheet.conditional_formatting)})
            for cell in sorted(sheet._cells.values(), key=lambda item: (item.row, item.column)):
                if cell.value is None and cell.data_type != "inlineStr":
                    continue
                is_formula = cell.data_type == "f"
                record = {"id": _cell_id(sheet.title, cell.coordinate),
                          "sheet": sheet.title, "address": cell.coordinate,
                          "dataType": cell.data_type, "numberFormat": cell.number_format,
                          "kind": "formula" if is_formula else "literal"}
                if is_formula:
                    expression, object_type, formula_range = _formula_source(cell.value)
                    record.update(expression=expression, dependencies=[], dependencyStatus="unresolved")
                    if object_type is not None:
                        record["formulaObjectType"] = object_type
                        record["formulaRange"] = formula_range
                else:
                    record["value"] = _literal("" if cell.data_type == "inlineStr" and cell.value is None else cell.value,
                                               cell.number_format)
                cells.append(record)
            for table in sorted(sheet.tables.values(), key=lambda item: item.name):
                item = _table_model(sheet, table)
                tables.append(item)
                blocks.append({"id": _id("block", "table", sheet.title, item["range"]),
                               "sheet": sheet.title, "range": item["range"],
                               "role": "unreviewed", "reviewed": False,
                               "sourceSha256": sha})
        populated = {cell["id"] for cell in cells}
        names, name_lookup = _defined_names(workbook)
        unresolved = []
        for cell in cells:
            if cell["kind"] != "formula":
                continue
            deps, reasons = _dependencies(cell["expression"], cell["sheet"], name_lookup, sheet_names,
                                          populated, cell.get("formulaObjectType"))
            cell["dependencies"] = deps
            cell["dependencyStatus"] = "unresolved" if reasons else "resolved"
            for index, reason in enumerate(reasons):
                unresolved.append({"id": _id("unresolved", cell["id"], index),
                                   "cellId": cell["id"], "reason": reason})
        blocks.extend(_reviewed_blocks(block_decisions, sha, sheet_names))
        model = {"modelVersion": 1, "sourceSha256": sha, "sourceName": source_path.name,
                 "sheets": sheets, "cells": cells, "tables": tables, "blocks": blocks,
                 "definedNames": names, "unresolved": unresolved,
                 "counts": {"formulas": sum(c["kind"] == "formula" for c in cells),
                            "charts": sum(s["chartCount"] for s in sheets),
                            "tables": len(tables)}}
        validate_source_model(model, source_path)
        return model
    finally:
        workbook.close()


def validate_source_model(model, source_path=None) -> None:
    """Check structure and source binding; this does not certify migration support."""
    if not isinstance(model, dict) or model.get("modelVersion") != 1:
        raise ValueError("invalid source model version")
    if (not isinstance(model.get("sourceName"), str) or
            not model["sourceName"] or Path(model["sourceName"]).name != model["sourceName"] or
            not isinstance(model.get("sourceSha256"), str) or not _SHA.fullmatch(model["sourceSha256"])):
        raise ValueError("invalid source identity")
    if source_path is not None and (_hash(source_path) != model["sourceSha256"] or Path(source_path).name != model["sourceName"]):
        raise ValueError("source bytes or name changed")
    collections = ("sheets", "cells", "tables", "blocks", "definedNames", "unresolved")
    if any(not isinstance(model.get(key), list) for key in collections):
        raise ValueError("invalid model collections")
    all_ids = []
    for key in collections:
        for item in model[key]:
            if not isinstance(item, dict) or not isinstance(item.get("id"), str):
                raise ValueError("malformed model identity")
            all_ids.append(item["id"])
    if len(all_ids) != len(set(all_ids)):
        raise ValueError("duplicate model identity")
    sheet_names = set()
    for sheet in model["sheets"]:
        name = sheet.get("name")
        if not isinstance(name, str) or sheet["id"] != _id("sheet", name) or name in sheet_names:
            raise ValueError("malformed sheet identity")
        if type(sheet.get("chartCount")) is not int or sheet["chartCount"] < 0:
            raise ValueError("invalid chart count")
        sheet_names.add(name)
    cell_ids = set()
    formula_ids = set()
    for cell in model["cells"]:
        sheet, address = cell.get("sheet"), cell.get("address")
        if sheet not in sheet_names or not isinstance(address, str) or not _ADDRESS.fullmatch(address) or cell["id"] != _cell_id(sheet, address):
            raise ValueError("malformed cell identity")
        cell_ids.add(cell["id"])
        if cell.get("dataType") == "f":
            expression = cell.get("expression")
            expression_valid = (isinstance(expression, str) and expression.startswith("=")) or (
                expression is None and isinstance(cell.get("formulaObjectType"), str) and
                bool(cell["formulaObjectType"]))
            if cell.get("kind") != "formula" or not expression_valid or "value" in cell:
                raise ValueError("misclassified formula")
            formula_ids.add(cell["id"])
            if cell.get("dependencyStatus") not in {"resolved", "unresolved"} or not isinstance(cell.get("dependencies"), list):
                raise ValueError("invalid formula dependency status")
            for dep in cell["dependencies"]:
                if (dep.get("sheet") not in sheet_names or
                        not isinstance(dep.get("address"), str) or not _ADDRESS.fullmatch(dep["address"]) or
                        dep.get("targetId") != _cell_id(dep["sheet"], dep["address"])):
                    raise ValueError("broken dependency identity")
        elif cell.get("kind") != "literal" or "value" not in cell or "expression" in cell:
            raise ValueError("misclassified literal")
    for cell in model["cells"]:
        if cell["kind"] == "formula":
            for dep in cell["dependencies"]:
                if dep.get("targetKind") != ("populated" if dep["targetId"] in cell_ids else "blank"):
                    raise ValueError("broken dependency target kind")
    unresolved_ids = set()
    unresolved_ordinals = {}
    for item in model["unresolved"]:
        if item.get("cellId") not in formula_ids or not isinstance(item.get("reason"), str) or not item["reason"]:
            raise ValueError("invalid unresolved reference")
        ordinal = unresolved_ordinals.get(item["cellId"], 0)
        if item["id"] != _id("unresolved", item["cellId"], ordinal):
            raise ValueError("malformed unresolved identity")
        unresolved_ordinals[item["cellId"]] = ordinal + 1
        unresolved_ids.add(item["cellId"])
    for cell in model["cells"]:
        if cell["kind"] == "formula" and (cell["dependencyStatus"] == "unresolved") != (cell["id"] in unresolved_ids):
            raise ValueError("unresolved formula lacks explicit reason")
    for table in model["tables"]:
        if table.get("sheet") not in sheet_names or table.get("id") != _id("table", table.get("sheet"), table.get("name")):
            raise ValueError("malformed table identity")
        _bounds(table.get("range"))
    reviewed_bounds = {}
    for block in model["blocks"]:
        if block.get("sheet") not in sheet_names or block.get("role") not in (_ROLES | {"unreviewed"}):
            raise ValueError("invalid block binding")
        bounds = _bounds(block.get("range"))
        source_range = _range_string(*bounds)
        origin = "reviewed" if block.get("reviewed") else "table"
        if block.get("id") != _id("block", origin, block["sheet"], source_range) or block.get("sourceSha256") != model["sourceSha256"]:
            raise ValueError("invalid block identity or source binding")
        if block.get("reviewed") != (block.get("role") != "unreviewed"):
            raise ValueError("invalid block review state")
        if block["reviewed"]:
            for previous in reviewed_bounds.get(block["sheet"], []):
                if bounds[0] <= previous[2] and previous[0] <= bounds[2] and bounds[1] <= previous[3] and previous[1] <= bounds[3]:
                    raise ValueError("reviewed block ranges overlap")
            reviewed_bounds.setdefault(block["sheet"], []).append(bounds)
    for name in model["definedNames"]:
        if name.get("id") != _id("name", name.get("scope"), name.get("name")) or name.get("scope") not in (sheet_names | {None}):
            raise ValueError("malformed defined name identity")
        if name.get("resolution") == "cell":
            if not isinstance(name.get("targetId"), str) or not name["targetId"].startswith("cell:"):
                raise ValueError("broken defined name target")
            try:
                target_sheet, target_address = json.loads(name["targetId"][len("cell:"):])
            except (TypeError, ValueError) as exc:
                raise ValueError("broken defined name target") from exc
            if (target_sheet not in sheet_names or not isinstance(target_address, str) or
                    not _ADDRESS.fullmatch(target_address) or
                    name["targetId"] != _cell_id(target_sheet, target_address)):
                raise ValueError("broken defined name target")
        elif name.get("resolution") != "unresolved" or not name.get("reason"):
            raise ValueError("invalid defined name resolution")
    name_lookup = {(name["scope"], name["name"].lower()): name for name in model["definedNames"]}
    reasons_by_cell = {}
    for item in model["unresolved"]:
        reasons_by_cell.setdefault(item["cellId"], []).append(item["reason"])
    for cell in model["cells"]:
        if cell["kind"] != "formula":
            continue
        expected_deps, expected_reasons = _dependencies(cell["expression"], cell["sheet"],
                                                        name_lookup, sheet_names, cell_ids,
                                                        cell.get("formulaObjectType"))
        if (cell["dependencies"] != expected_deps or
                reasons_by_cell.get(cell["id"], []) != expected_reasons or
                cell["dependencyStatus"] != ("unresolved" if expected_reasons else "resolved")):
            raise ValueError("formula dependency analysis does not match source expression")
    counts = model.get("counts")
    if (not isinstance(counts, dict) or any(type(counts.get(key)) is not int for key in ("formulas", "charts", "tables")) or
            counts != {"formulas": len(formula_ids),
                       "charts": sum(s["chartCount"] for s in model["sheets"]),
                       "tables": len(model["tables"])}):
        raise ValueError("invalid source counts")
