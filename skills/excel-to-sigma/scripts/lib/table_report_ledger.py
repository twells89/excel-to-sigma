"""Period ledger recogniser for ``table_report`` (the 13-week cash flow shape).

A *period ledger* is a report sheet with line items down column A and period ordinals across a
header row (``Week | 1 | 2 | ... | 13 | 13-Wk Total``). Each Excel row is one of

* a **manual** row - typed numbers (rent, debt service): an editable column of the period spine;
* a **table** row - ``SUMIFS(tbl[Value], tbl[Period], B$4)``, optionally times a scenario factor;
* a **recursion** row - ``seed`` then ``=prev*(1+rate)`` (lowered to the closed form
  ``seed*(1+rate)^(period-first)``), or a date row ``seed`` then ``=prev+7``;
* a **cadence** row - ``IF(MOD(B$4,2)=0, amount, 0)``;
* a **derived** row - arithmetic / ``SUM`` over other rows of the same period, ``=B23-MinCash``;
* a **running balance** - ``Beginning = previous Ending`` with ``Ending = Beginning + Net``
  (``Opening + CumulativeSum(Net)``; Beginning = Ending - Net) and running ``MIN($B$23:B23)``.

The whole ledger is lowered to ONE table with a row per period (the spine table: the period key
plus the manual columns) and a calculated column per Excel row, in the sheet's row order. Table
rows read a hidden grouped *helper* frame (the table summed by its period column) through
``Lookup`` with a ``0`` default, because Sigma has no row-level SUMIFS over another table. The
``INDIRECT("prefix"&<named cell>)`` scenario idiom becomes a ``Lookup`` into a small scenario
table built from the named-range family ``prefix*`` and the selected scenario (a text parameter).

Every row's cells are then recomputed by the Python oracle (``table_report_model``) from the plan
and compared with Excel's cached values; anything unrecognised raises ``PlanDraftError`` naming
``Sheet!Cell``.
"""
from __future__ import annotations

import copy
import re

from openpyxl.utils import get_column_letter

try:
    from . import table_report as tr
    from . import table_report_model as trm
    from .excel_expression import Binary, Function, Group, Number, Range, Reference, Text
except ImportError:  # direct import when scripts/lib itself is on sys.path
    import table_report as tr
    import table_report_model as trm
    from excel_expression import Binary, Function, Group, Number, Range, Reference, Text

PlanDraftError = tr.PlanDraftError
PARAM_KINDS = ("number", "text", "date")
NON_CONSTANT = {"c", "period", "key", "prev", "agg", "cumsum", "cummin"}


def _is_int(value):
    return (isinstance(value, int) and not isinstance(value, bool)) or \
        (isinstance(value, float) and value.is_integer())


def _constant(ir):
    return not any(node[0] in NON_CONSTANT for node in trm.walk(ir))


def _uses_row_labels(ws, sheet, lc, pcols, first_row, last_row):
    """A row-label criterion identifies a keyed/summary grid, not a period ledger."""
    for row in range(first_row, last_row + 1):
        for col in pcols:
            value = ws.cell(row, col).value
            if not tr._is_formula(value):
                continue
            ast = tr._parse(value, sheet, tr._addr(sheet, row, col))
            if tr._contains_ref_to(ast, lambda ref, row=row: bool(ref.address)
                                   and (ref.sheet or sheet) == sheet
                                   and tr._rc(ref) == (row, lc)):
                return True
    return False


class Region:
    """One detected ledger: header row, label column, period columns, optional total column."""

    def __init__(self, ws, hr, lc, run, tcol, total_label, first_row, last_row):
        self.ws, self.sheet = ws, ws.title
        self.hr, self.lc = hr, lc
        self.pcols = [c for c, _ in run]
        self.periods = [v for _, v in run]
        self.tcol, self.total_label = tcol, total_label
        self.first_row, self.last_row = first_row, last_row


def find_ledgers(reg):
    found = []
    for ws in reg.wf.worksheets:
        cov = reg.covered.get(ws.title, set())

        def cell(r, c, ws=ws, cov=cov):
            return None if (r, c) in cov else ws.cell(r, c).value

        used_until = 0
        for hr in range(1, ws.max_row + 1):
            if hr <= used_until:
                continue
            present = [c for c in range(1, ws.max_column + 1) if cell(hr, c) is not None]
            if len(present) < 4:
                continue
            lc = present[0]
            label = cell(hr, lc)
            if not (isinstance(label, str) and label.strip() and not tr._is_formula(label)):
                continue
            run, c = [], lc + 1
            while _is_int(cell(hr, c)) and (not run or cell(hr, c) == run[-1][1] + 1):
                run.append((c, cell(hr, c)))
                c += 1
            if len(run) < 3:
                continue
            tcol, total_label = None, None
            extra = cell(hr, c)
            if isinstance(extra, str) and extra.strip() and not tr._is_formula(extra):
                tcol, total_label = c, extra.strip()
            last, blank, r = hr, 0, hr + 1
            while r <= ws.max_row:
                if any(cell(r, k) is not None for k in range(lc, ws.max_column + 1)):
                    last, blank = r, 0
                else:
                    blank += 1
                    if blank >= 2:
                        break
                r += 1
            if _uses_row_labels(ws, ws.title, lc, [k for k, _ in run], hr + 1, last):
                continue
            found.append(Region(ws, hr, lc, run, tcol, total_label, hr + 1, last))
            used_until = last
    return found


# ---------------------------------------------------------------------------- helper frames
class Helper:
    """A hidden grouped table: ``table`` summed by its period column (``Lookup`` target)."""

    def __init__(self, drafter, table, key_col, sheet):
        keys = sorted({v for v in tr.col_values(drafter.tabs[table], key_col) if v is not None},
                      key=lambda v: (str(type(v)), str(v)))
        title = drafter.unique_title(tr._title(f"{table} by {key_col}"))
        self.frame = {"id": f"s_{trm.slug(title)}", "kind": "summary", "title": title, "sheet": sheet,
                      "helper": True, "hidden": True, "source": {"kind": "table", "name": table},
                      "by": key_col, "labels": keys, "columns": []}
        drafter.frames.append(self.frame)
        self.drafter, self.cols = drafter, {}

    @property
    def title(self):
        return self.frame["title"]

    def column(self, value_col, filters):
        key = (value_col, repr(filters))
        if key not in self.cols:
            taken = {c["name"] for c in self.frame["columns"]} | {self.frame["by"]}
            name, n = value_col, 2
            while name in taken:
                name, n = f"{value_col} ({n})", n + 1
            self.cols[key] = name
            self.frame["columns"].append({"name": name, "type": "number",
                                          "expr": ["agg", "sum", value_col, tr._and(filters)]})
        return self.cols[key]


# ---------------------------------------------------------------------------- lowering context
class LedgerCtx(tr.Ctx):
    def __init__(self, draft, row, col):
        self.ld, self.row, self.col = draft, row, col

    def named(self, low, name):
        sheet, address = self.ld.defined_cell(name, low.where)
        return self.ld.param_ir(low, sheet, address)

    def ref(self, low, ref):
        ld, led = self.ld, self.ld.led
        sheet = ref.sheet or low.sheet
        row, col = tr._rc(ref)
        if sheet == ld.sheet:
            if row == led.hr:
                if col == self.col:
                    return ["period"]
                raise low.fail(f"{ref.address}: a period header reference must be the formula's own column "
                               "(for example B$4 in column B)")
            if row in ld.names:
                if row == self.row:
                    raise low.fail(f"{ref.address}: a ledger row cannot read itself")
                if ref.column_absolute:
                    raise low.fail(f"{ref.address}: a fixed-column reference to a ledger row would read one "
                                   "period for every period (use a relative column)")
                if col == self.col:
                    return ["c", ld.names[row]]
                raise low.fail(f"{ref.address}: ledger formulas may read the same period only (previous "
                               "periods are supported as a running balance / recursion, not here)")
            if led.first_row <= row <= led.last_row and led.lc <= col <= (led.tcol or led.pcols[-1]):
                raise low.fail(f"{ref.address}: references a label or blank cell of the ledger")
        if not ref.column_absolute:
            raise low.fail(f"{ref.address}: a relative reference to a cell outside the ledger moves across the "
                           "periods when the formula is copied right; use an absolute column ($B$3) so it is "
                           "one parameter")
        return ld.param_ir(low, sheet, ref.address)

    def func(self, low, node):
        name = node.name
        if name == "SUMIFS":
            return self.ld.sumifs(low, node, self)
        if name == "INDIRECT":
            return self.ld.indirect(low, node)
        if name in ("SUM", "MIN", "MAX") and len(node.arguments) == 1 and isinstance(node.arguments[0], Range):
            return self.ld.range_function(low, node, self)
        return None


# ---------------------------------------------------------------------------- one ledger
class LedgerDraft:
    def __init__(self, drafter, led):
        self.d, self.led, self.reg = drafter, led, drafter.reg
        self.ws, self.sheet = led.ws, led.sheet
        self.names = {}            # excel row -> column name (line rows only)
        self.helpers = {}          # (table, key col) -> Helper
        self.scenario = None
        self.indirect_cells = []
        self.recursion_cells, self.balance_cells = [], []
        self.axis = None

    # ---- naming / parameters
    def defined_cell(self, name, where):
        found = next((v for k, v in self.reg.wf.defined_names.items() if k.casefold() == name.casefold()), None)
        if found is None:
            raise PlanDraftError(f"{where}: named range {name!r} is not defined in this workbook")
        targets = list(found.destinations)
        if len(targets) != 1 or ":" in targets[0][1]:
            raise PlanDraftError(f"{where}: named range {name!r} must be a single cell, got {found.attr_text!r}")
        return targets[0][0], targets[0][1].replace("$", "")

    def param_ir(self, low, sheet, address):
        name = self.reg.param(sheet, address, low.where, allow=PARAM_KINDS)
        entry = self.reg.params[(sheet, address.replace("$", ""))]
        if entry.get("kind") == "text" and "options" not in entry:
            options = self.dv_options(sheet, address)
            if options:
                entry["options"] = options
        return [{"number": "p", "text": "pt", "date": "pd"}[entry.get("kind", "number")], name]

    def dv_options(self, sheet, address):
        coord = address.replace("$", "")
        for dv in self.reg.wf[sheet].data_validations.dataValidation:
            if dv.type == "list" and coord in dv.sqref and dv.formula1 and dv.formula1.startswith('"'):
                return [v.strip() for v in dv.formula1.strip('"').split(",")]
        return None

    # ---- SUMIFS over a table, by the period header
    def sumifs(self, low, node, ctx):
        args = node.arguments
        if len(args) < 3 or len(args) % 2 == 0:
            raise low.fail("SUMIFS needs a value range then range/criterion pairs")
        table, value_col = low.table_ref(args[0], "SUMIFS value range")
        period_col, filters = None, []
        for rng, crit in zip(args[1::2], args[2::2]):
            t, col = low.table_ref(rng, "SUMIFS criteria range")
            if t is not table:
                raise low.fail("SUMIFS ranges belong to different tables")
            inner = crit.expression if isinstance(crit, Group) else crit
            if isinstance(inner, Reference) and inner.value.address and \
                    (inner.value.sheet or low.sheet) == self.sheet and \
                    tr._rc(inner.value) == (self.led.hr, ctx.col):
                if period_col is not None:
                    raise low.fail("SUMIFS has two criteria on the period header")
                period_col = col
                continue
            op, value = tr._crit(low, crit, lambda ref: False)
            if value[0] not in ("n", "s"):
                raise low.fail("extra SUMIFS criteria must be literals (the period header is the only cell "
                               "criterion a ledger row can have)")
            filters.append(["op", op, ["c", col], value])
        if period_col is None:
            raise low.fail("a ledger SUMIFS must be filtered by the period header (for example B$4)")
        if table.col(value_col).type not in (None, "number") and table.col(value_col).kind == "literal":
            raise low.fail(f"SUMIFS value column {table.name}[{value_col}] is not numeric")
        key = (table.name, period_col)
        if key not in self.helpers:
            self.helpers[key] = Helper(self.d, table.name, period_col, self.sheet)
        helper = self.helpers[key]
        return ["lk", helper.title, ["period"], period_col, helper.column(value_col, filters), ["n", 0]]

    # ---- INDIRECT("prefix"&Name): the scenario-lookup idiom
    def indirect(self, low, node):
        arg = node.arguments[0] if len(node.arguments) == 1 else None
        while isinstance(arg, Group):
            arg = arg.expression
        if not (isinstance(arg, Binary) and arg.operator == "&" and isinstance(arg.left, Text)
                and isinstance(arg.right, Reference) and arg.right.value.name):
            raise low.fail('INDIRECT is only supported as INDIRECT("prefix"&<named cell>) over a family of '
                           "named single cells (the scenario-lookup idiom)")
        self.indirect_cells.append(low.where)
        return self.scenario_factor(low, arg.left.value, arg.right.value.name)

    def scenario_factor(self, low, prefix, selector_name):
        if self.scenario is None or self.scenario["prefix"] != prefix:
            family = [(k, v) for k, v in self.reg.wf.defined_names.items()
                      if k.casefold().startswith(prefix.casefold()) and len(k) > len(prefix)]
            if not family:
                raise low.fail(f"INDIRECT prefix {prefix!r} matches no named range")
            rows, cells = [], []
            for name, _ in family:
                sheet, address = self.defined_cell(name, low.where)
                cell = self.reg.wf[sheet][address]
                if tr._is_formula(cell.value) or not tr._is_number(cell.value):
                    raise low.fail(f"named range {name!r} ({sheet}!{address}) must be a typed number to be a "
                                   "scenario value")
                rows.append([name[len(prefix):], cell.value])
                cells.append((sheet, address, cell))
            if len({r[0].casefold() for r in rows}) != len(rows):
                raise low.fail(f"named ranges {prefix}* repeat a scenario name")
            sheet, address, first = cells[0]
            ws = self.reg.wf[sheet]
            key_name = rate_name = None
            above, left = ws[address].offset(-1, 0), ws[address].offset(-1, -1)
            if isinstance(above.value, str) and isinstance(left.value, str) and above.value != left.value:
                rate_name, key_name = tr._clean_name(above.value, above.coordinate), \
                    tr._clean_name(left.value, left.coordinate)
            key_name, rate_name = key_name or "Scenario", rate_name or "Value"
            title = self.d.unique_title(tr._title(f"{rate_name} by {key_name}"))
            fmt = tr._fmt_kind(first)
            self.reg.synthetic[title] = {
                "name": title, "sheet": sheet,
                "ref": ", ".join(f"{a}" for _, a, _ in cells), "editable": True, "synthetic": True,
                "columns": [{"name": key_name, "type": "text"},
                            {"name": rate_name, "type": "number", **({"format": fmt} if fmt else {})}],
                "rows": rows,
                "note": f"Scenario values from the named ranges {prefix}* "
                        f"({', '.join(f'{n}={a}' for (n, _), (_, a, _) in zip(family, cells))})."}
            self.scenario = {"prefix": prefix, "table": title, "key": key_name, "rate": rate_name,
                             "options": [r[0] for r in rows], "selector": selector_name}
        selector_sheet, selector_address = self.defined_cell(selector_name, low.where)
        entry_name = self.reg.param(selector_sheet, selector_address, low.where, allow=("text",))
        entry = self.reg.params[(selector_sheet, selector_address)]
        if entry["value"] not in self.scenario["options"]:
            raise low.fail(f"selected scenario {entry['value']!r} ({selector_sheet}!{selector_address}) is not "
                           f"one of the named ranges {prefix}* ({', '.join(self.scenario['options'])})")
        self.scenario["parameter"] = entry_name
        return ["lk", self.scenario["table"], ["pt", entry_name], self.scenario["key"],
                self.scenario["rate"], None]

    # ---- SUM/MIN/MAX over rows of the same period, or the running window MIN($B$23:B23)
    def range_function(self, low, node, ctx):
        rng = node.arguments[0]
        sheet = rng.start.sheet or low.sheet
        r1, c1 = tr._rc(rng.start)
        r2, c2 = tr._rc(rng.end)
        led = self.led
        if sheet != self.sheet:
            raise low.fail(f"{node.name} over a range on another sheet is not supported in a ledger")
        running = r1 == r2 and rng.start.column_absolute and not rng.end.column_absolute
        if not running and c1 == c2 == ctx.col:
            lo, hi = min(r1, r2), max(r1, r2)
            items = [["c", self.names[r]] for r in range(lo, hi + 1) if r in self.names]
            if ctx.row in range(lo, hi + 1):
                raise low.fail(f"{node.name} range includes the formula's own row")
            if not items:
                raise low.fail(f"{node.name} range holds no ledger row")
            if len(items) == 1:
                return items[0]
            return ["f", node.name] + items
        if running and r1 in self.names and c1 == led.pcols[0] and c2 == ctx.col:
            if r1 == ctx.row:
                raise low.fail("a running aggregate cannot read its own row")
            if node.name == "SUM":
                tag = "cumsum"
            elif node.name == "MIN":
                tag = "cummin"
            else:
                raise low.fail("a running MAX over the period columns is not supported (SUM and MIN are)")
            return [tag, ["c", self.names[r1]], None, self.axis]
        raise low.fail(f"{node.name} over {rng.start.address}:{rng.end.address} is neither a column span of "
                       "ledger rows nor a running window from the first period")

    # ---- the ledger
    def cached(self, row, col):
        return self.reg.wv[self.sheet].cell(row, col).value

    def lower(self, ast, row, col, where):
        low = tr.Lowerer(self.reg, self.sheet, where, LedgerCtx(self, row, col))
        return low.lower(ast)

    def prev_self(self, node, row, col):
        node = node.expression if isinstance(node, Group) else node
        return isinstance(node, Reference) and bool(node.value.address) and \
            (node.value.sheet or self.sheet) == self.sheet and tr._rc(node.value) == (row, col - 1)

    def classify(self):
        led, ws = self.led, self.ws
        records, section = [], None
        for r in range(led.first_row, led.last_row + 1):
            label = ws.cell(r, led.lc).value
            cells = [ws.cell(r, c).value for c in led.pcols]
            total = ws.cell(r, led.tcol).value if led.tcol else None
            beyond = [c for c in range((led.tcol or led.pcols[-1]) + 1, ws.max_column + 1)
                      if ws.cell(r, c).value is not None]
            if beyond:
                raise PlanDraftError(f"{tr._addr(self.sheet, r, beyond[0])}: cell right of the ledger's total "
                                     "column is not supported")
            if all(v is None for v in cells) and total is None:
                if isinstance(label, str) and label.strip():
                    section = label.strip()
                continue
            if not (isinstance(label, str) and label.strip()) or tr._is_formula(label):
                raise PlanDraftError(f"{tr._addr(self.sheet, r, led.lc)}: ledger rows need a text label")
            blanks = [c for c, v in zip(led.pcols, cells) if v is None]
            if blanks:
                raise PlanDraftError(f"{tr._addr(self.sheet, r, blanks[0])}: empty cell in ledger row "
                                     f"{label.strip()!r}; every period needs a value or formula")
            name = tr._clean_name(label, tr._addr(self.sheet, r, led.lc))
            formulas = [tr._is_formula(v) for v in cells]
            record = {"row": r, "name": name, "section": section, "values": cells, "total": total}
            if not any(formulas):
                bad = next((c for c, v in zip(led.pcols, cells) if not tr._is_number(v)), None)
                if bad is not None:
                    raise PlanDraftError(f"{tr._addr(self.sheet, r, bad)}: manual ledger rows hold numbers only")
                record["kind"] = "manual"
            elif all(formulas):
                asts = [tr._parse(v, self.sheet, tr._addr(self.sheet, r, c))
                        for c, v in zip(led.pcols, cells)]
                norms = [tr._norm(a, r, c, self.sheet) for a, c in zip(asts, led.pcols)]
                record["asts"] = asts
                if len(set(norms)) == 1:
                    record["kind"] = "uniform"
                elif len(set(norms[1:])) == 1:
                    record["kind"] = "seeded"
                else:
                    odd = next(i for i in range(2, len(norms)) if norms[i] != norms[1])
                    raise PlanDraftError(
                        f"{tr._addr(self.sheet, r, led.pcols[odd])}: formula differs from "
                        f"{tr._addr(self.sheet, r, led.pcols[1])} ({cells[odd]!r}); a ledger row is one "
                        "formula copied across (only the first period may differ)")
            else:
                first_literal = next(c for c, f in zip(led.pcols, formulas) if not f)
                raise PlanDraftError(f"{tr._addr(self.sheet, r, first_literal)}: ledger row {name!r} mixes "
                                     "formulas and typed values")
            records.append(record)
        if len({rec["name"] for rec in records}) != len(records):
            dup = next(rec["name"] for rec in records if [x["name"] for x in records].count(rec["name"]) > 1)
            raise PlanDraftError(f"{self.sheet}: duplicate ledger row label {dup!r}")
        return records

    def run(self):
        led, d = self.led, self.d
        self.axis = tr._clean_name(self.ws.cell(led.hr, led.lc).value, tr._addr(self.sheet, led.hr, led.lc))
        records = self.classify()
        if not records:
            return False
        for rec in records:
            if rec["name"] == self.axis:
                raise PlanDraftError(f"{tr._addr(self.sheet, rec['row'], led.lc)}: row label {rec['name']!r} "
                                     "equals the period header label")
            self.names[rec["row"]] = rec["name"]
        by_row = {rec["row"]: rec for rec in records}
        columns, balance_of = {}, {}      # name -> {expr, type}; end row -> (begin rec, net name, opening)
        # running balances first: Ending = Opening + CumulativeSum(Net), Beginning = Ending - Net
        for rec in records:
            if rec["kind"] != "seeded":
                continue
            tail, col = rec["asts"][1], led.pcols[1]
            target = tail.expression if isinstance(tail, Group) else tail
            if isinstance(target, Reference) and target.value.address and \
                    (target.value.sheet or self.sheet) == self.sheet:
                trow, tcol = tr._rc(target.value)
                if tcol == col - 1 and trow in by_row and trow != rec["row"]:
                    rec["shape"] = ("carry", trow)
        for rec in records:
            if rec["kind"] == "seeded":
                self.seeded(rec, by_row, columns, balance_of)
        for rec in records:
            if rec["kind"] == "uniform" and rec["row"] not in balance_of:
                where = tr._addr(self.sheet, rec["row"], led.pcols[0])
                expr = self.lower(rec["asts"][0], rec["row"], led.pcols[0], where)
                columns[rec["name"]] = {"expr": expr, "row": rec["row"]}
        for end_row, (begin, net_name, opening) in balance_of.items():
            end_rec = by_row[end_row]
            columns[end_rec["name"]] = {"row": end_row, "expr": ["op", "+", opening,
                                        ["cumsum", ["c", net_name], None, self.axis]]}
            columns[begin["name"]] = {"row": begin["row"], "expr": ["op", "-", ["c", end_rec["name"]],
                                                                    ["c", net_name]]}
            self.balance_cells += [tr._addr(self.sheet, begin["row"], led.pcols[0]),
                                   tr._addr(self.sheet, end_row, led.pcols[0])]
        return self.assemble(records, columns)

    def constant(self, ir, where, what):
        if not _constant(ir):
            raise PlanDraftError(f"{where}: {what} must be a typed value, a parameter or a named cell (it "
                                 "cannot read other ledger rows or the period)")
        return ir

    def seeded(self, rec, by_row, columns, balance_of):
        led = self.led
        r, c0, c1 = rec["row"], led.pcols[0], led.pcols[1]
        first_where, tail_where = tr._addr(self.sheet, r, c0), tr._addr(self.sheet, r, c1)
        tail = rec["asts"][1]
        node = tail.expression if isinstance(tail, Group) else tail
        seed = self.constant(self.lower(rec["asts"][0], r, c0, first_where), first_where, "the first period")
        axis_first = ["n", led.periods[0]]
        offset = ["op", "-", ["period"], axis_first]
        if rec.get("shape", (None,))[0] == "carry":
            end_row = rec["shape"][1]
            end = by_row[end_row]
            if end["kind"] != "uniform":
                raise PlanDraftError(f"{tail_where}: carry-forward reads {end['name']!r}, which is not one "
                                     "formula across the periods")
            end_ast = end["asts"][0]
            end_ast = end_ast.expression if isinstance(end_ast, Group) else end_ast
            parts = []
            if isinstance(end_ast, Binary) and end_ast.operator == "+":
                for side in (end_ast.left, end_ast.right):
                    side = side.expression if isinstance(side, Group) else side
                    if isinstance(side, Reference) and side.value.address and \
                            (side.value.sheet or self.sheet) == self.sheet and tr._rc(side.value)[1] == c0:
                        parts.append(tr._rc(side.value)[0])
            if len(parts) != 2 or r not in parts:
                raise PlanDraftError(
                    f"{tr._addr(self.sheet, end_row, c0)}: a carry-forward row ({rec['name']!r}, "
                    f"{tail_where}) must read a row defined as {rec['name']} + <net row> in the same period")
            net_row = next(p for p in parts if p != r)
            if net_row not in by_row:
                raise PlanDraftError(f"{tr._addr(self.sheet, end_row, c0)}: the net row is not a ledger row")
            balance_of[end_row] = (rec, by_row[net_row]["name"], seed)
            return
        if isinstance(node, Binary) and node.operator == "*":
            for prev, other in ((node.left, node.right), (node.right, node.left)):
                other = other.expression if isinstance(other, Group) else other
                if self.prev_self(prev, r, c1) and isinstance(other, Binary) and other.operator == "+" \
                        and isinstance(other.left, Number) and other.left.value == 1:
                    rate = self.constant(self.lower(other.right, r, c1, tail_where), tail_where, "the growth rate")
                    growth = ["op", "^", ["op", "+", ["n", 1], rate], offset]
                    columns[rec["name"]] = {"row": r, "expr": ["op", "*", seed, growth]}
                    self.recursion_cells.append(tail_where)
                    return
        if isinstance(node, Binary) and node.operator == "+":
            for prev, step in ((node.left, node.right), (node.right, node.left)):
                if self.prev_self(prev, r, c1) and isinstance(step, Number) and float(step.value) == int(step.value):
                    columns[rec["name"]] = {"row": r, "expr": [
                        "f", "ADDDAYS", seed, ["op", "*", ["n", int(step.value)], offset]]}
                    self.recursion_cells.append(tail_where)
                    return
        raise PlanDraftError(
            f"{tail_where}: the first period of {rec['name']!r} differs and the next period's formula "
            f"{rec['values'][1]!r} is not a recognised recursion (=prev*(1+rate), =prev+<days>) or a "
            "carry-forward of another row's previous period")

    # ---- plan assembly
    def col_type(self, rec):
        led = self.led
        kinds = set()
        for c in led.pcols:
            v = self.cached(rec["row"], c)
            if v is None:
                raise PlanDraftError(f"{tr._addr(self.sheet, rec['row'], c)}: no cached value (open the workbook in "
                                     "Excel, recalculate and save)")
            kinds.add("datetime" if hasattr(v, "year") else "number" if tr._is_number(v) else "text")
        if len(kinds) != 1:
            raise PlanDraftError(f"{tr._addr(self.sheet, rec['row'], led.pcols[0])}: row {rec['name']!r} mixes "
                                 f"{' and '.join(sorted(kinds))} values")
        return kinds.pop()

    def assemble(self, records, columns):
        led, d = self.led, self.d
        title = d.unique_title(tr._title(self.sheet))
        frame_id = f"g_{trm.slug(title)}"
        manual = [rec for rec in records if rec["kind"] == "manual"]
        spine_title = d.unique_title(f"{title} Periods")
        spine_columns = [{"name": self.axis, "type": "number"}] + [
            {"name": rec["name"], "type": "number"} for rec in manual]
        spine_rows = [[p] + [rec["values"][i] for rec in manual] for i, p in enumerate(led.periods)]
        self.reg.synthetic[spine_title] = {
            "name": spine_title, "sheet": self.sheet,
            "ref": f"{get_column_letter(led.pcols[0])}{led.hr}:{get_column_letter(led.pcols[-1])}{led.last_row}",
            "editable": True, "synthetic": True, "columns": spine_columns, "rows": spine_rows,
            "note": f"The period spine of {self.sheet} (the header ordinals) with its typed line items "
                    f"({', '.join(r['name'] for r in manual) or 'none'}) as editable columns."}
        frame_columns, order, checks, targets, totals = [], [self.axis], [], [], []
        headline = []
        for rec in records:
            order.append(rec["name"])
            if rec["kind"] == "manual":
                continue
            column = columns[rec["name"]]
            ctype = self.col_type(rec)
            entry = {"name": rec["name"], "type": ctype, "expr": column["expr"], "row": rec["row"]}
            if ctype == "datetime":
                entry["format"] = "date"
            else:
                fmt = tr._fmt_kind(self.ws.cell(rec["row"], led.pcols[0]))
                if fmt:
                    entry["format"] = fmt
            frame_columns.append(entry)
            d.add_checks(frame_id, rec["name"], "ledger",
                         [[tr._addr(self.sheet, rec["row"], c), [p], self.cached(rec["row"], c)]
                          for c, p in zip(led.pcols, led.periods)])
        # total column: SUM across the period columns, verified and emitted as a one-row table
        for rec in records:
            if rec["total"] is None:
                continue
            where = tr._addr(self.sheet, rec["row"], led.tcol)
            ast = tr._parse(rec["total"], self.sheet, where) if tr._is_formula(rec["total"]) else None
            ok = isinstance(ast, Function) and ast.name == "SUM" and len(ast.arguments) == 1 \
                and isinstance(ast.arguments[0], Range)
            if ok:
                start, end = tr._rc(ast.arguments[0].start), tr._rc(ast.arguments[0].end)
                ok = start == (rec["row"], led.pcols[0]) and end == (rec["row"], led.pcols[-1])
            if not ok:
                raise PlanDraftError(f"{where}: the total column cell must be SUM across the period columns of "
                                     f"its row ({get_column_letter(led.pcols[0])}{rec['row']}:"
                                     f"{get_column_letter(led.pcols[-1])}{rec['row']}), got {rec['total']!r}")
            totals.append(rec["name"])
            d.add_checks(frame_id, rec["name"], "ledger-total", [[where, [], self.cached(rec["row"], led.tcol)]])
        # captions: ="Scenario: "&Scenario above the header
        captions = self.captions(frame_id)
        end_names = {rec["name"] for rec in records if rec["row"] in self._balance_end_rows(columns)}
        for name in totals:
            targets.append({"label": f"{title} / {name} / {led.total_label}", "frame": frame_id, "column": name,
                            "kind": "ledger-total", "sel": [],
                            "value": tr._plain(self.cached(next(r for r in records if r["name"] == name)["row"],
                                                           led.tcol))})
            headline.append(targets[-1])
        for column in frame_columns:
            expr = column["expr"]
            nodes = [n[0] for n in trm.walk(expr)]
            last_only = "cummin" in nodes
            full = "cumsum" in nodes and column["name"] in end_names
            if not (last_only or full):
                continue
            rec = next(r for r in records if r["name"] == column["name"])
            for c, p in zip(led.pcols, led.periods):
                if last_only and p != led.periods[-1]:
                    continue
                targets.append({"label": f"{title} / {column['name']} / {self.axis} {p}", "frame": frame_id,
                                "column": column["name"], "kind": "ledger", "sel": [p],
                                "value": tr._plain(self.cached(rec["row"], c))})
                if p == led.periods[-1]:
                    headline.append(targets[-1])
        d.targets.extend(targets)
        frame = {"id": frame_id, "kind": "ledger", "title": title, "sheet": self.sheet,
                 "spine": spine_title, "axis": self.axis, "columns": frame_columns, "order": order,
                 "totals": totals, "totalLabel": led.total_label or "Total", "captions": captions}
        if self.scenario:
            frame["scenario"] = {"parameter": self.scenario["parameter"], "table": self.scenario["table"],
                                 "options": self.scenario["options"]}
        d.frames.append(frame)
        for r0 in range(led.hr, led.last_row + 1):
            self.reg.covered.setdefault(self.sheet, set()).update(
                (r0, c) for c in range(led.lc, (led.tcol or led.pcols[-1]) + 1))
        self.issues(frame, manual, headline)
        return True

    @staticmethod
    def _balance_end_rows(columns):
        return {v["row"] for v in columns.values()
                if any(n[0] == "cumsum" for n in trm.walk(v["expr"]))}

    def captions(self, frame_id):
        led, out = self.led, []
        for r in range(1, led.hr):
            for c in range(1, self.ws.max_column + 1):
                value = self.ws.cell(r, c).value
                if not tr._is_formula(value):
                    continue
                where = tr._addr(self.sheet, r, c)
                ast = tr._parse(value, self.sheet, where)
                parts = []

                def flatten(node):
                    node = node.expression if isinstance(node, Group) else node
                    if isinstance(node, Binary) and node.operator == "&":
                        flatten(node.left)
                        flatten(node.right)
                    elif isinstance(node, Text):
                        parts.append(["s", node.value])
                    elif isinstance(node, Reference):
                        ctx = LedgerCtx(self, r, c)
                        low = tr.Lowerer(self.reg, self.sheet, where, ctx)
                        ir = low.lower(node)
                        if ir[0] != "pt":
                            raise PlanDraftError(f"{where}: a caption can join text and text parameters only")
                        parts.append(ir)
                    else:
                        raise PlanDraftError(f"{where}: formula above the ledger header is not a text caption "
                                             f"(\"text\"&<named cell>), got {value!r}")
                flatten(ast)
                out.append({"cell": where, "expr": ["f", "CONCAT"] + parts})
                self.d.add_checks(frame_id, "Caption", "ledger-caption",
                                  [[where, [len(out) - 1], self.reg.wv[self.sheet].cell(r, c).value]])
                self.reg.covered.setdefault(self.sheet, set()).add((r, c))
        return out

    def issues(self, frame, manual, headline):
        led, d = self.led, self.d
        sheet = self.sheet
        if self.indirect_cells:
            sc = self.scenario
            d.issues.append({
                "code": "INDIRECT_SCENARIO_LOOKUP", "severity": "review", "route": "scenario-lookup-review",
                "message": f"{sheet}: INDIRECT(\"{sc['prefix']}\"&{sc['selector']}) picks one named cell by the "
                           f"selected scenario. It is lowered to a Lookup into the editable "
                           f"'{sc['table']}' table ({', '.join(sc['options'])}) keyed by the {sc['parameter']!r} "
                           "text parameter (Parameters table); change the scenario by editing that parameter, "
                           "not a workbook control (controls are filters, not formula parameters).",
                "evidence": self.indirect_cells[:5]})
        d.issues.append({
            "code": "LEDGER_TRANSPOSED", "severity": "review", "route": "ledger-orientation-review",
            "message": f"{sheet}: Excel shows {self.axis}s across and line items down; the Sigma table has one row "
                       f"per {self.axis} and one column per Excel row, in the sheet's row order. The "
                       f"{led.total_label or 'total'} column becomes a one-row grouped table.",
            "evidence": [f"{sheet}!{get_column_letter(led.lc)}{led.hr}"]})
        if self.recursion_cells:
            d.issues.append({
                "code": "RECURSION_CLOSED_FORM", "severity": "review", "route": "recursion-closed-form-review",
                "message": f"{sheet}: =prev*(1+rate) / =prev+<days> rows were lowered to the closed form "
                           f"seed*(1+rate)^({self.axis}-first) / seed+days*({self.axis}-first); the oracle "
                           "verified every period against Excel's recursion.",
                "evidence": self.recursion_cells[:5]})
        if self.balance_cells or any("cummin" in [n[0] for n in trm.walk(c["expr"])] for c in frame["columns"]):
            d.issues.append({
                "code": "LEDGER_RUNNING_BALANCE_REVIEW", "severity": "review",
                "route": "cumulative-sum-live-validation",
                "message": f"{sheet}: the ending balance is Opening + CumulativeSum(net) ordered by {self.axis}, "
                           "the beginning balance is ending - net, and the running minimum is CumulativeMin; "
                           "confirm the live table matches the parity targets before relying on it.",
                "evidence": (self.balance_cells or [f"{sheet}!{get_column_letter(led.pcols[0])}{led.first_row}"])[:5]})
        if frame["captions"]:
            d.issues.append({
                "code": "CAPTION_FORMULA_VERIFIED_ONLY", "severity": "review", "route": "caption-review",
                "message": f"{sheet}: caption formulas were verified against their cached text but are not "
                           "emitted (a Sigma text element cannot read a parameter).",
                "evidence": [c["cell"] for c in frame["captions"]]})
        if manual:
            d.issues.append({
                "code": "LEDGER_MANUAL_LINES", "severity": "review", "route": "manual-lines-review",
                "message": f"{sheet}: typed per-{self.axis} amounts ({', '.join(m['name'] for m in manual)}) "
                           f"became editable columns of the '{frame['spine']}' input table; seed it from its CSV.",
                "evidence": [f"{sheet}!{get_column_letter(led.pcols[0])}{m['row']}" for m in manual]})
        keyless = []
        for helper in self.helpers.values():
            tab = self.d.tabs[helper.frame["source"]["name"]]
            unique = any(c.type == "text" and len({v for v in tr.col_values(tab, c.name)}) == tab.n
                         for c in tab.columns)
            if not unique and tab.name not in keyless:
                keyless.append(tab.name)
        if keyless:
            d.issues.append({
                "code": "KEYLESS_TABLE_ROWS", "severity": "review", "route": "keyless-table-review",
                "message": "These tables have no text column that is unique per row. Aggregating by period "
                           "needs no key, so no synthetic row key was added, but seeding the input table is "
                           "positional (upsert enforces no key): seed it once from its CSV and do not re-run "
                           "the seed.",
                "evidence": keyless})
        d.ledger_scenarios.append({
            "frame": frame["id"], "scenario": frame.get("scenario"), "headline": headline})


def draft_ledgers(drafter):
    drafter.ledger_scenarios = []
    for region in find_ledgers(drafter.reg):
        LedgerDraft(drafter, region).run()


def scenario_results(plan, drafter):
    """Recompute the headline parity targets under every scenario (informational: Excel cached one)."""
    out = {}
    for entry in drafter.ledger_scenarios:
        scenario = entry["scenario"]
        if not scenario:
            continue
        frame_id, results = entry["frame"], {}
        for option in scenario["options"]:
            candidate = copy.deepcopy(plan)
            for param in candidate["params"]:
                if param["name"] == scenario["parameter"]:
                    param["value"] = option
            try:
                model = trm.compute(candidate)
            except (trm.PlanError, trm.EvalError) as error:
                results[option] = {"error": str(error)}
                continue
            results[option] = {t["label"]: trm.iso(model.cell(t["frame"], t["column"], t["kind"], t["sel"]))
                               for t in entry["headline"]}
        out[frame_id] = {"parameter": scenario["parameter"], "selected": next(
            p["value"] for p in plan["params"] if p["name"] == scenario["parameter"]),
            "informational": "Excel cached only the selected scenario; the others are recomputed from the "
                             "plan and are not parity targets", "results": results}
    return out
