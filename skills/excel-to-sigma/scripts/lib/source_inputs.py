"""Reviewed, source-bound ordinary input tables; no target calls or writes."""
from __future__ import annotations

import copy
import datetime as dt
import hashlib
import json
import re
from pathlib import Path

from openpyxl.utils.cell import get_column_letter, range_boundaries

try:
    from .input_sync import plan_seed
    from .source_workbook import build_source_model
except ImportError:
    from input_sync import plan_seed
    from source_workbook import build_source_model


PLAN_VERSION = 1
_TYPES = {"text", "number", "datetime", "boolean"}
_SAFE = re.compile(r"^[A-Za-z][A-Za-z0-9_-]*$")


def _digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"),
                                     ensure_ascii=False).encode("utf-8")).hexdigest()


def _value(cell):
    if cell.get("dataType") == "e":
        raise ValueError("Excel error cell cannot be an entered value")
    if cell.get("dataType") not in {"n", "s", "b", "d", "inlineStr"}:
        raise ValueError(f"unsupported entered source cell dataType: {cell.get('dataType')}")
    value = cell["value"]
    if isinstance(value, dict):
        if value.get("type") in {"date", "datetime"}:
            return value["iso"]
        raise ValueError(f"unsupported source cell type: {value.get('type')}")
    if value is None or type(value) in {str, int, float, bool}:
        return value
    raise ValueError("unsupported source cell value")


def _matches(value, kind):
    if value is None:
        return True
    if kind == "text":
        return type(value) is str
    if kind == "number":
        return type(value) in {int, float}
    if kind == "boolean":
        return type(value) is bool
    if kind == "datetime":
        if not isinstance(value, str):
            return False
        try:
            if re.fullmatch(r"\d{4}-\d\d-\d\d", value):
                dt.date.fromisoformat(value)
            elif re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d(?:\.\d+)?(?:[+-]\d\d:\d\d|Z)?", value):
                dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
            else:
                return False
        except ValueError:
            return False
        return True
    return False


def _selection(model, decision):
    sheet = decision.get("sheet")
    if not isinstance(sheet, str) or sheet not in {item["name"] for item in model["sheets"] if item["kind"] == "worksheet"}:
        raise ValueError("unknown source sheet")
    table_name, block = decision.get("table"), decision.get("range")
    if bool(table_name) == bool(block):
        raise ValueError("select exactly one formal table or bounded cell block")
    if table_name:
        matches = [item for item in model["tables"] if item["sheet"] == sheet and item["name"] == table_name]
        if len(matches) != 1 or not matches[0]["bodyRange"]:
            raise ValueError("unknown or empty source table")
        item = matches[0]
        left, top, right, bottom = range_boundaries(item["bodyRange"])
        headers = item["headers"]
        source_range = item["range"]
    else:
        try:
            left, top, right, bottom = range_boundaries(block)
        except (TypeError, ValueError) as exc:
            raise ValueError("invalid bounded cell block") from exc
        if None in (left, top, right, bottom) or left > right or top > bottom:
            raise ValueError("invalid bounded cell block")
        headers = decision.get("headers")
        if not isinstance(headers, list) or len(headers) != right - left + 1:
            raise ValueError("cell block requires explicit headers")
        source_range = block
    if any(not isinstance(name, str) or not name for name in headers) or len(set(headers)) != len(headers):
        raise ValueError("source headers must be unique and nonempty")
    return sheet, headers, (left, top, right, bottom), source_range


def _build_replacement(model, decision, connection_id):
    feature = decision.get("featureId")
    if not isinstance(feature, str) or not _SAFE.fullmatch(feature):
        raise ValueError("replacement featureId must be stable and safe")
    review = decision.get("replacementReview") or {}
    if (decision.get("reviewed") is not True or review.get("approved") is not True
            or not str(review.get("basis") or "").strip()
            or not str(decision.get("replacementFor") or "").strip()
            or not str(decision.get("displayName") or "").strip()):
        raise ValueError("replacement requires explicit reviewed source and target label")
    keys, literal, types = (decision.get(name) for name in ("keyColumns", "literalColumns", "columnTypes"))
    if (not isinstance(keys, list) or len(keys) != 1 or not isinstance(literal, list)
            or len(literal) != 2 or keys[0] not in literal or not isinstance(types, dict)
            or set(types) != set(literal) or any(kind not in _TYPES for kind in types.values())):
        raise ValueError("replacement requires one typed natural key and one typed value")
    value_name = next(name for name in literal if name != keys[0])
    rows, bindings = decision.get("suppliedRows"), decision.get("replacementBindings")
    if (not isinstance(rows, list) or not rows or not isinstance(bindings, list)
            or len(bindings) != len(rows)):
        raise ValueError("replacement suppliedRows and source bindings must match")
    cells = {f"{cell['sheet']}!{cell['address']}": cell for cell in model["cells"]}
    provenance, obligations, seen_keys, seen_formulas = {}, [], set(), set()
    for index, (row, binding) in enumerate(zip(rows, bindings)):
        if (not isinstance(row, dict) or set(row) != set(literal)
                or not isinstance(binding, dict) or set(binding.get("key") or {}) != set(keys)
                or binding["key"] != {key: row[key] for key in keys}
                or set(binding.get("sourceCells") or {}) != set(literal)):
            raise ValueError("replacement row and source bindings are inconsistent")
        for name in literal:
            if row[name] is None or not _matches(row[name], types[name]):
                raise ValueError("replacement row disagrees with reviewed type")
        key_token = _digest([[type(row[key]).__name__, row[key]] for key in keys])
        if key_token in seen_keys:
            raise ValueError("duplicate replacement business key")
        seen_keys.add(key_token)
        key_cell_id = binding["sourceCells"][keys[0]]
        value_cell_id = binding["sourceCells"][value_name]
        key_cell, value_cell = cells.get(key_cell_id), cells.get(value_cell_id)
        if (not key_cell or key_cell["kind"] != "literal"
                or _value(key_cell) != binding.get("sourceLabel")):
            raise ValueError("replacement key source label is stale")
        if (not value_cell or value_cell["kind"] != "formula"
                or value_cell.get("expression") != binding.get("originalFormula")
                or value_cell_id in seen_formulas):
            raise ValueError("replacement source formula is stale or duplicated")
        seen_formulas.add(value_cell_id)
        provenance[str(index)] = dict(binding["sourceCells"])
        obligations.append({"featureId": feature, "sourceCellId": value_cell_id,
                            "sourceRow": int(re.search(r"\d+$", value_cell["address"]).group()),
                            "sourceColumn": value_name, "expression": binding["originalFormula"],
                            "dependencies": value_cell.get("dependencies") or [],
                            "mappingStatus": "reviewed-replacement-input", "replacementRowKey": binding["key"]})
    spec = {"kind": "input-table", "name": decision["displayName"],
            "keyColumns": keys, "connectionId": connection_id, "editMode": "published-explore",
            "sourceBinding": decision["replacementFor"], "rows": copy.deepcopy(rows),
            "columns": [{"name": name, "type": types[name], "provenance": "reviewed-replacement"}
                        for name in literal], "sourceCells": provenance,
            "sourceRowNumbers": [item["sourceRow"] for item in obligations],
            "sourceHeaders": literal, "derivedColumns": [], "overrideCells": {},
            "replacementBindings": copy.deepcopy(bindings), "replacementReview": copy.deepcopy(review)}
    return feature, spec, obligations, {}


def _build_reshape_feature(source_path, model, decision, connection_id):
    """Recompute a reviewed unpivot from exact source bytes before admitting it as input."""
    if (not isinstance(decision, dict) or decision.get("reviewed") is not True
            or decision.get("role") != "input"):
        raise ValueError("reshape input requires an explicit reviewed input role")
    feature = decision.get("featureId")
    if not isinstance(feature, str) or not _SAFE.fullmatch(feature):
        raise ValueError("reshape featureId must be stable and safe")
    reshape = decision.get("reshape")
    expected = {"kind", "sheet", "table", "keyColumns", "valueColumns", "axisName", "valueName"}
    if (not isinstance(reshape, dict) or reshape.get("kind") != "unpivot"
            or set(reshape) not in (expected, expected | {"axisValues"})
            or decision.get("overrideCells") or decision.get("emptySchema")):
        raise ValueError("reshape must explicitly select one reviewed unpivot")
    try:
        from .source_reshape import unpivot_source_table
    except ImportError:
        from source_reshape import unpivot_source_table
    result = unpivot_source_table(
        source_path, sheet=reshape["sheet"], table=reshape["table"],
        key_columns=reshape["keyColumns"], value_columns=reshape["valueColumns"],
        axis_name=reshape["axisName"], value_name=reshape["valueName"],
        axis_values=reshape.get("axisValues"))
    columns = result["keyColumns"] + [reshape["valueName"]]
    if (result["sourceSha256"] != model["sourceSha256"]
            or decision.get("keyColumns") != result["keyColumns"]
            or decision.get("literalColumns") != columns
            or decision.get("columnTypes") != result["columnTypes"]):
        raise ValueError("reviewed reshape keys, columns or types disagree with source")
    spec = {"kind": "input-table", "name": decision.get("displayName") or f"{reshape['table']} Long",
            "keyColumns": result["keyColumns"], "connectionId": connection_id,
            "editMode": "published-explore", "sourceBinding": result["sourceBinding"],
            "rows": result["rows"],
            "columns": [{"name": name, "type": result["columnTypes"][name], "provenance": "source-reshape"}
                        for name in columns],
            "sourceCells": result["sourceCells"], "sourceRowNumbers": result["sourceRowNumbers"],
            "sourceHeaders": reshape["keyColumns"] + reshape["valueColumns"],
            "axisMapping": result["axisMapping"], "reshape": copy.deepcopy(reshape),
            "derivedColumns": [], "overrideCells": {}}
    return feature, spec, [], {}


def _build_feature(model, decision, connection_id, source_path):
    if isinstance(decision, dict) and decision.get("role") == "replacement":
        return _build_replacement(model, decision, connection_id)
    if isinstance(decision, dict) and "reshape" in decision:
        return _build_reshape_feature(source_path, model, decision, connection_id)
    if not isinstance(decision, dict) or decision.get("reviewed") is not True or decision.get("role") not in {"input", "reference"}:
        raise ValueError("feature requires a reviewed input or maintained reference role")
    feature = decision.get("featureId")
    if not isinstance(feature, str) or not _SAFE.fullmatch(feature):
        raise ValueError("featureId must be a stable safe identifier")
    sheet, headers, bounds, source_range = _selection(model, decision)
    keys = decision.get("keyColumns")
    literal = decision.get("literalColumns")
    types = decision.get("columnTypes")
    if (not isinstance(keys, list) or not keys or len(keys) != len(set(keys))
            or not isinstance(literal, list) or not literal or len(literal) != len(set(literal))
            or not set(keys) <= set(literal) or not set(literal) <= set(headers)
            or not isinstance(types, dict) or set(types) != set(literal)
            or any(kind not in _TYPES for kind in types.values())):
        raise ValueError("input columns, typed keys, and explicit types must match source schema")
    overrides = decision.get("overrideCells") or {}
    if not isinstance(overrides, dict) or any(name not in literal or name in keys or not isinstance(cells, list)
                                                for name, cells in overrides.items()):
        raise ValueError("editable override mappings must name non-key literal columns")
    override_set = {name: set(cells) for name, cells in overrides.items()}
    if any(len(cells) != len(override_set[name]) for name, cells in overrides.items()):
        raise ValueError("duplicate editable override mapping")
    left, top, right, bottom = bounds
    cells = {(item["sheet"], item["address"]): item for item in model["cells"]}
    if decision.get("emptySchema") is True:
        if (set(decision.get("reviewedBlankColumns") or []) != set(literal)
                or not str(decision.get("typeReviewEvidence") or "").strip()):
            raise ValueError("empty schema requires reviewed types and typeReviewEvidence")
        if any((sheet, f"{get_column_letter(col)}{row}") in cells
               for row in range(top, bottom + 1) for col in range(left, right + 1)):
            raise ValueError("emptySchema cannot suppress populated source cells")
        return feature, {"kind": "input-table", "name": decision.get("displayName") or decision.get("table") or feature,
                         "keyColumns": keys, "connectionId": connection_id,
                         "editMode": "published-explore", "sourceBinding": f"{sheet}!{source_range}",
                         "rows": [], "columns": [{"name": name, "type": types[name], "provenance": "source"}
                                                 for name in literal],
                         "sourceCells": {}, "sourceRowNumbers": [], "sourceHeaders": headers,
                         "derivedColumns": [], "overrideCells": {}, "emptySchema": True,
                         "typeReviewEvidence": decision["typeReviewEvidence"]}, [], {}
    rows, source_cells, obligations, key_tokens = [], {}, [], set()
    override_records = {name: [] for name in overrides}
    for row_number in range(top, bottom + 1):
        row, provenance, row_overrides = {}, {}, {}
        for offset, name in enumerate(headers):
            address = f"{get_column_letter(left + offset)}{row_number}"
            cell = cells.get((sheet, address))
            cell_id = f"{sheet}!{address}"
            if cell and cell["kind"] == "formula":
                obligations.append({"featureId": feature, "sourceCellId": cell_id,
                                    "sourceRow": row_number, "sourceColumn": name,
                                    "expression": cell.get("expression"),
                                    "dependencies": cell.get("dependencies") or [],
                                    "mappingStatus": "live-calculation-required"})
                if name in literal and name not in override_set:
                    raise ValueError(f"formula cannot be seeded as literal entered data: {cell_id}")
                if name in override_set and address in override_set[name]:
                    raise ValueError(f"formula cell cannot be an editable override: {cell_id}")
                continue
            if name not in literal:
                if cell and cell["kind"] == "literal":
                    raise ValueError(f"unmapped literal source cell: {cell_id}")
                continue
            if name in override_set and cell is None:
                if address in override_set[name]:
                    raise ValueError(f"editable override has no literal source cell: {cell_id}")
                continue
            if cell is None:
                value = None
            else:
                value = _value(cell)
            if not _matches(value, types[name]):
                raise ValueError(f"source cell type disagrees with reviewed type: {cell_id}")
            if name in override_set and address not in override_set[name]:
                raise ValueError(f"literal override lacks explicit editable mapping: {cell_id}")
            if name in override_set:
                row_overrides[name] = (value, cell_id)
                continue
            row[name] = value
            provenance[name] = cell_id
        key = []
        for name in keys:
            if name not in row or row[name] is None or row[name] == "":
                raise ValueError(f"missing typed key {name!r} in {sheet} row {row_number}")
            key.append([type(row[name]).__name__, row[name]])
        token = _digest(key)
        if token in key_tokens:
            raise ValueError("duplicate typed business key")
        key_tokens.add(token)
        rows.append(row)
        source_cells[str(len(rows) - 1)] = provenance
        for name, (value, cell_id) in row_overrides.items():
            override_records[name].append({
                "row": {**{key_name: row[key_name] for key_name in keys}, name: value},
                "sourceCells": {**{key_name: provenance[key_name] for key_name in keys}, name: cell_id},
                "sourceRow": row_number})
    for name, addresses in override_set.items():
        seen = {item["sourceCells"][name].split("!")[1] for item in override_records[name]}
        if addresses != seen:
            raise ValueError("editable override mappings are incomplete or foreign")
        if not any(item["sourceColumn"] == name for item in obligations):
            raise ValueError(f"editable override column has no source formula: {name}")
    main_literal = [name for name in literal if name not in overrides]
    for name in main_literal:
        if all(row.get(name) is None for row in rows) and name not in (decision.get("reviewedBlankColumns") or []):
            raise ValueError(f"all-blank column {name!r} requires explicit reviewed type")
    if not rows:
        raise ValueError("source input selection has no rows")
    display = decision.get("displayName") or decision.get("table") or feature
    override_features = {name: f"{feature}__override_{_digest(name)[:8]}" for name in sorted(overrides)}
    derived = [name for name in headers if name not in main_literal]
    spec = {"kind": "input-table", "name": display,
            "keyColumns": keys, "connectionId": connection_id, "editMode": "published-explore",
            "sourceBinding": f"{sheet}!{source_range}", "rows": rows,
            "columns": [{"name": name, "type": types[name], "provenance": "source"} for name in main_literal],
            "sourceCells": source_cells, "sourceRowNumbers": list(range(top, bottom + 1)),
            "sourceHeaders": headers, "derivedColumns": derived,
            "derivedMappings": {name: {"sourceFormulaCells": [item["sourceCellId"] for item in obligations
                                                                  if item["sourceColumn"] == name],
                                       "overrideInputFeature": override_features.get(name)}
                                for name in derived},
            "overrideCells": overrides, "overrideInputFeatures": override_features}
    extra_inputs = {}
    for name, override_feature in override_features.items():
        records = override_records[name]
        if not records or all(item["row"][name] is None for item in records):
            raise ValueError(f"editable override column has no entered values: {name}")
        extra_inputs[override_feature] = {
            "kind": "input-table", "name": f"{display} — {name} Overrides",
            "keyColumns": keys, "connectionId": connection_id, "editMode": "published-explore",
            "sourceBinding": f"{sheet}!{source_range}",
            "rows": [item["row"] for item in records],
            "columns": [{"name": column, "type": types[column], "provenance": "source"}
                        for column in keys + [name]],
            "sourceCells": {str(index): item["sourceCells"] for index, item in enumerate(records)},
            "sourceRowNumbers": [item["sourceRow"] for item in records],
            "sourceHeaders": headers, "derivedColumns": [], "overrideCells": {name: overrides[name]},
            "overrideFor": {"featureId": feature, "sourceColumn": name,
                            "sourceFormulaCells": spec["derivedMappings"][name]["sourceFormulaCells"]}}
    return feature, spec, obligations, extra_inputs


def build_source_input_plan(source_path, decisions, *, connection_id, folder_id) -> dict:
    """Build reviewed ordinary input architecture and source-cell reverse mappings."""
    if isinstance(decisions, (str, Path)):
        decisions = json.loads(Path(decisions).read_text(encoding="utf-8"))
    if not isinstance(decisions, dict) or not isinstance(decisions.get("features"), list):
        raise ValueError("source input decisions require features")
    if not connection_id or not folder_id:
        raise ValueError("target connection and folder scope are required")
    model = build_source_model(source_path)
    if decisions.get("sourceSha256") != model["sourceSha256"]:
        raise ValueError("source hash does not match reviewed decisions")
    features = decisions["features"]
    if not features:
        raise ValueError("at least one reviewed source feature is required")
    inputs, obligations = {}, []
    for decision in features:
        feature, spec, formulas, extras = _build_feature(model, decision, connection_id, source_path)
        if set([feature, *extras]) & set(inputs):
            raise ValueError("duplicate featureId or generated override input")
        inputs[feature] = spec
        inputs.update(extras)
        obligations.extend(formulas)
    selected_formulas = {item["sourceCellId"] for item in obligations}
    for cell in model["cells"]:
        if cell["kind"] != "formula":
            continue
        source_cell = f"{cell['sheet']}!{cell['address']}"
        if source_cell not in selected_formulas:
            obligations.append({"featureId": None, "sourceCellId": source_cell,
                                "sourceRow": int(re.search(r"\d+$", cell["address"]).group()),
                                "sourceColumn": None, "expression": cell.get("expression"),
                                "dependencies": cell.get("dependencies") or [],
                                "mappingStatus": "live-calculation-required"})
    architecture = {"workflows": ["ordinary-source-inputs"],
                    "targetScope": {"connectionId": connection_id, "folderId": folder_id},
                    "mappedFeatures": sorted(inputs),
                    "featureToInput": {feature: feature for feature in sorted(inputs)},
                    "inputs": {feature: inputs[feature] for feature in sorted(inputs)},
                    "acceptanceChecks": [f"{feature}-edit-persist-restore" for feature in sorted(inputs)],
                    "nativeWiring": {"elements": [
                        {"id": f"input_{re.sub('[^A-Za-z0-9_]', '_', feature)}_{_digest(feature)[:8]}",
                         "kind": "input-table", "name": inputs[feature]["name"],
                         "inputMode": "explore", "source": {"kind": "empty", "connectionId": connection_id},
                         "columns": [{"id": f"c{index}", "name": column["name"], "type": column["type"]}
                                     for index, column in enumerate(inputs[feature]["columns"])]}
                        for feature in sorted(inputs)], "consumers": {}, "validations": {}}}
    logical = {"inputPlanVersion": PLAN_VERSION, "sourceSha256": model["sourceSha256"],
               "sourceName": model["sourceName"], "decisions": copy.deepcopy(decisions),
               "formulaObligations": sorted(obligations, key=lambda item: item["sourceCellId"]),
               "sharedInputArchitecture": architecture,
               "status": "REVIEWED_PLAN_PENDING_STAGING"}
    return {**logical, "inputPlanHash": _digest(logical)}


def validate_source_input_plan(plan, source_path=None) -> None:
    """Reject altered plans and, when supplied, changed source bytes or semantics."""
    if not isinstance(plan, dict) or plan.get("inputPlanVersion") != PLAN_VERSION:
        raise ValueError("unsupported source input plan")
    logical = {key: value for key, value in plan.items() if key != "inputPlanHash"}
    if _digest(logical) != plan.get("inputPlanHash"):
        raise ValueError("source input plan hash is stale")
    if source_path is not None:
        if hashlib.sha256(Path(source_path).read_bytes()).hexdigest() != plan["sourceSha256"]:
            raise ValueError("source hash is stale")
        scope = plan["sharedInputArchitecture"]["targetScope"]
        rebuilt = build_source_input_plan(source_path, plan["decisions"],
                                          connection_id=scope["connectionId"], folder_id=scope["folderId"])
        if rebuilt != plan:
            raise ValueError("source input plan differs from source bytes")


def plan_source_seed(plan, feature_id, current_rows, *, row_id_column, prior_seed=None) -> dict:
    """Delegate three-way, row-ID-aware reconciliation to the existing planner."""
    validate_source_input_plan(plan)
    inputs = plan["sharedInputArchitecture"]["inputs"]
    if feature_id not in inputs:
        raise ValueError("unknown source input feature")
    spec = inputs[feature_id]
    date_keys = {name for name in spec["keyColumns"]
                 if next(column["type"] for column in spec["columns"] if column["name"] == name) == "datetime"}
    date_only_keys = {name for name in date_keys if spec["rows"] and all(
        isinstance(row.get(name), str) and re.fullmatch(r"\d{4}-\d\d-\d\d", row[name])
        for row in spec["rows"])}
    normalized_current = []
    for current in current_rows:
        item = dict(current)
        for name in date_keys:
            value = item.get(name)
            if isinstance(value, dt.datetime):
                if value.tzinfo is not None:
                    value = value.astimezone(dt.timezone.utc).replace(tzinfo=None)
                if name in date_only_keys and value.time() in {dt.time.min, dt.time(12)}:
                    item[name] = value.date().isoformat()
                elif name in date_only_keys:
                    raise ValueError(f"queried date-only key has unexpected time: {name}")
                else:
                    item[name] = value.isoformat()
            elif isinstance(value, dt.date):
                item[name] = value.isoformat()
            elif isinstance(value, str) and name in date_only_keys and "T" in value:
                parsed = dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
                if parsed.time() not in {dt.time.min, dt.time(12)}:
                    raise ValueError(f"queried date-only key has unexpected time: {name}")
                item[name] = parsed.date().isoformat()
        normalized_current.append(item)
    return plan_seed(spec["rows"], normalized_current, spec["keyColumns"], row_id_column,
                     prior_seed=prior_seed,
                     column_types={column["name"]: column["type"] for column in spec["columns"]})
