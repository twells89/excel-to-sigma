"""Corrected Forecast snapshot definition and explicit retry reconciliation.

This module plans only a component. No function sends requests, accepts cached
Excel formula results as live values, or proves published workflow acceptance.
Native actions are not transactional: serialize writes and query after each run.
"""
import copy
import datetime as dt
import hashlib
import json
import math
from pathlib import Path
import re

PERIODS = ('Jul-26', 'Aug-26', 'Sep-26', 'Oct-26', 'Nov-26', 'Dec-26')
FINANCIAL_ROWS = (*range(8, 20), 22, 23, 24, 25, 27, 28)
SNAPSHOT_FIELDS = {'Operation ID':'text', 'Snapshot Key':'text', 'Source Sheet':'text',
    'Source Cell':'text', 'Source Row':'number', 'Line Item':'text', 'Section':'text',
    'Period':'text', 'Value':'number', 'Captured At':'datetime', 'Source Hash':'text',
    'Plan Hash':'text', 'Source Procedure Hash':'text', 'Replacement Decision ID':'text'}


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(',', ':'), allow_nan=False).encode()).hexdigest()


def _hash(value):
    if not re.fullmatch('[0-9a-f]{64}', str(value or '')):
        raise ValueError('source-bound SHA256 required')
    return value


def _number(value):
    if type(value) not in (int, float) or not math.isfinite(value):
        raise ValueError('finite actual numeric snapshot value required')
    return value


def _time(value):
    try:
        parsed = dt.datetime.fromisoformat(value.replace('Z', '+00:00'))
        if parsed.tzinfo is None or parsed.utcoffset() is None:raise ValueError()
    except (ValueError, TypeError, AttributeError):
        raise ValueError('one frozen timestamp with explicit timezone required')
    return parsed.astimezone(dt.timezone.utc).isoformat()


def build_snapshot_definition(workbook_source, *, plan_hash, procedure_hash, decision):
    """Bind the explicit replacement to source bytes, never Excel value caches."""
    from openpyxl import load_workbook
    scope = decision.get('scope', {})
    if (decision.get('status') != 'APPROVED_BUSINESS_SEMANTICS_ONLY' or not decision.get('decisionId')
            or scope.get('periods') != list(PERIODS) or scope.get('financialLineCount') != 18
            or scope.get('valueCount') != 108
            or scope.get('timestamp') != 'separate metadata; no financial value overwritten'):
        raise ValueError('explicit corrected July–December 18-line/108-value decision required')
    source = Path(workbook_source)
    wb = load_workbook(source, data_only=False)
    try:
        ws = wb['Monthly Model']
        # The reviewed source has June in H; I:N are the approved six months.
        if [ws.cell(5, c).value for c in range(9,15)] != list(PERIODS):
            raise ValueError('reviewed source July–December headers changed')
        cells = []
        for row in FINANCIAL_ROWS:
            label = ws.cell(row, 2).value
            if not isinstance(label,str) or not label.strip():raise ValueError('financial line label missing')
            section = ws.cell(row, 1).value if row < 20 else 'Derived'
            if not isinstance(section,str) or not section.strip():raise ValueError('financial section missing')
            for column,period in zip(range(9,15),PERIODS):
                cell = ws.cell(row,column)
                if cell.value is None:raise ValueError('source financial formula/value missing')
                key = str(row)+':'+period
                cells.append({'snapshotKey':key,'sourceSheet':'Monthly Model','sourceCell':cell.coordinate,
                    'sourceRow':row,'lineItem':label,'section':section,'period':period,
                    'operationFrozenColumn':'Frozen Snapshot / '+key})
        inputs = []
        inp = wb['Forecast Inputs']
        for row in range(6,18):
            for col,field in enumerate(('Method','Monthly Growth',*PERIODS),2):
                inputs.append({'sourceRow':row,'lineItem':inp.cell(row,1).value,'field':field,
                    'sourceCell':inp.cell(row,col).coordinate,
                    'operationColumn':f'Frozen Assumption / source row {row} / {field}',
                    'type':'text' if field=='Method' else 'number'})
    finally:wb.close()
    result = {'sourceHash':hashlib.sha256(source.read_bytes()).hexdigest(), 'planHash':_hash(plan_hash),
        'procedureHash':_hash(procedure_hash), 'decisionId':decision['decisionId'], 'cells':cells, 'inputFields':inputs,
        'literalMacroParity':False, 'timestampPolicy':'separate frozen metadata', 'productionStatus':'BLOCKED'}
    result['definitionHash'] = digest(result)
    return result


def _definition(definition):
    if definition.get('definitionHash') != digest({k:v for k,v in definition.items() if k!='definitionHash'}):
        raise ValueError('snapshot definition changed')
    cells = definition['cells']
    if (len(cells)!=108 or {c['sourceRow'] for c in cells}!=set(FINANCIAL_ROWS)
            or {(c['sourceRow'],c['period']) for c in cells}!={(r,p) for r in FINANCIAL_ROWS for p in PERIODS}
            or len({c['snapshotKey'] for c in cells})!=108):raise ValueError('complete 18×6 snapshot required')


def reconcile_snapshot_rows(definition, operation, rows):
    """Pure reconciliation, including exact existing rows before proposing inserts.

    An operation's actual native row ID is its identity. Retrying supplies that
    same row; a deliberate new commit supplies a new prepared operation row.
    This pure function makes no claim that its arguments are actual captures.
    """
    _definition(definition)
    op = operation.get('nativeRowId')
    if not isinstance(op,str) or not op.strip():raise ValueError('stored native operation row ID required')
    metadata = {'Source Hash':definition['sourceHash'], 'Plan Hash':definition['planHash'],
        'Source Procedure Hash':definition['procedureHash'], 'Replacement Decision ID':definition['decisionId']}
    if (any(operation.get(k)!=v for k,v in metadata.items()) or operation.get('Workflow')!='snapshot'
            or operation.get('Context')!='forecast-edit' or operation.get('State')!='PREPARED'
            or operation.get('Snapshot Definition Hash')!=definition['definitionHash']):
        raise ValueError('prepared source/plan/decision-bound snapshot operation required')
    timestamp = operation.get('Created At');_time(timestamp)
    expected = {}
    for cell in definition['cells']:
        row = {'Operation ID':op, 'Snapshot Key':cell['snapshotKey'], 'Source Sheet':cell['sourceSheet'],
            'Source Cell':cell['sourceCell'],'Source Row':cell['sourceRow'],'Line Item':cell['lineItem'],
            'Section':cell['section'],'Period':cell['period'],
            'Value':_number(operation.get(cell['operationFrozenColumn'])), 'Captured At':timestamp, **metadata}
        expected[cell['snapshotKey']] = row
    seen = set(); all_ids = set()
    for row in rows:
        row_id = row.get('nativeRowId')
        if not isinstance(row_id,str) or not row_id or row_id in all_ids:raise ValueError('unique actual stored snapshot row IDs required')
        all_ids.add(row_id)
        if set(row)!={*SNAPSHOT_FIELDS,'nativeRowId'}:raise ValueError('complete snapshot fields required')
        if row['Operation ID'] != op:continue
        key = row['Snapshot Key']
        if key in seen or key not in expected:raise ValueError('duplicate or foreign snapshot key; reconcile before effects')
        seen.add(key)
        actual = {k:v for k,v in row.items() if k!='nativeRowId'}
        _number(actual['Value'])
        # Timestamp spelling may differ after query, but the instant must match.
        actual['Captured At'] = _time(actual['Captured At'])
        wanted = {**expected[key], 'Captured At':_time(timestamp)}
        _number(actual['Source Row'])
        if any(actual[k]!=wanted[k] for k in wanted):raise ValueError('snapshot conflict; no overwrite or recomputed retry permitted')
    missing = [row for key,row in expected.items() if key not in seen]
    return {'status':'COMPONENT_RECONCILIATION_ONLY','operationId':op,'missingRows':missing,
        'existingCount':len(seen),'expectedCount':108,'complete':not missing,
        'operationHash':digest(operation),'snapshotRowsHash':digest(rows),'productionStatus':'BLOCKED'}


def plan_snapshot_reconciliation(definition, operation_id, operations, snapshots, target):
    """Require fresh complete actual destination-bound captures before planning."""
    required_ids=('documentId','operationsElementId','operationsRowIdColumnId',
                  'snapshotsElementId','snapshotsRowIdColumnId')
    if (not isinstance(target,dict)
            or any(not isinstance(target.get(k),str) or not target[k].strip() for k in required_ids)
            or type(target.get('version')) is not int or target['version']<=0):
        raise ValueError('nonempty target document/native IDs and strict positive integer version required')
    for role,capture in [('operations',operations),('snapshots',snapshots)]:
        if (capture.get('origin')!='sigma-query' or capture.get('complete') is not True
                or type(capture.get('version')) is not int or capture['version']<=0
                or type(capture.get('rowCount')) is not int or capture['rowCount']<0
                or not capture.get('evidenceId') or not capture.get('executionId')
                or capture.get('rowCount')!=len(capture.get('rows',[]))
                or not target.get(role+'ElementId') or not target.get(role+'RowIdColumnId')
                or capture.get('elementId')!=target[role+'ElementId']
                or capture.get('rowIdColumnId')!=target[role+'RowIdColumnId']
                or any(capture.get(k)!=target.get(k) for k in ('documentId','version'))
                or any(capture.get(k)!=definition[k] for k in ('sourceHash','planHash'))):
            raise ValueError('fresh complete actual source/plan/document/version/native-ID-bound '+role+' capture required')
    found = [r for r in operations['rows'] if r.get('nativeRowId')==operation_id]
    if len(found)!=1:raise ValueError('exactly one stored native operation row required')
    result = reconcile_snapshot_rows(definition,found[0],snapshots['rows'])
    result.update(operationReadbackHash=digest(operations),snapshotReadbackHash=digest(snapshots),target=copy.deepcopy(target))
    result['reconciliationHash'] = digest(result)
    return result


def _join(expressions, operator='and_'):
    values = list(expressions)
    if len(values)==1:return values[0]
    middle=len(values)//2
    return operator+'('+_join(values[:middle],operator)+', '+_join(values[middle:],operator)+')'


def _balance_lines(name, operator, count):
    """Pair adjacent boolean operands without changing order or duplicates."""
    following = 'next_'+name
    return [f'if len({name}) != {count}: raise RuntimeError("guard part count")',
        f'for balance_level in range({(count-1).bit_length()}):',
        f'    {following} = []',
        f'    for balance_index in range(0, len({name}), 2):',
        f'        if balance_index + 1 < len({name}):',
        f'            {following}.append({operator}({name}[balance_index], {name}[balance_index+1]))',
        '        else:', f'            {following}.append({name}[balance_index])',
        f'    {name} = {following}',
        f'if len({name}) != 1: raise RuntimeError("guard balance count")']


def render_snapshot_candidate(definition, current, bindings, *, phase, operation_control_id=None, prepared_contract=None):
    """Render one disabled, guarded native action construction phase.

    `bindings` records the full assigned SDK schemas and original input row IDs.
    The caller must tie those records to actual captured evidence. The returned
    controller precondition requires a fresh exact public source before submit;
    SDK exposes no version CAS. Both phases use the same definition and never
    enable/click actions. Confirm requires Prepare's actual assigned TextControl.
    """
    import xml.etree.ElementTree as ET
    _definition(definition)
    if phase not in ('prepare','confirm'):raise ValueError('explicit prepare or confirm phase required')
    if (not current.get('workbookId') or type(current.get('documentVersion')) is not int
            or current['documentVersion']<1 or set(bindings)!={'primary','finance','operations','snapshots'}):
        raise ValueError('current saved target and four complete assigned bindings required')
    contents = current['contents'];elements = {e['id']:e for e in contents['elements']}
    pages = {p['id']:p for p in contents['pages']}
    root = ET.fromstring('<Layout>'+re.sub(r'<\?xml[^?]*\?>','',contents['layout'])+'</Layout>')
    owners = {}
    for page in root:
        for node in page.iter('Element'):
            eid = node.get('elementId') or node.get('id')
            if eid in owners:raise ValueError('ambiguous assigned layout')
            owners[eid] = page.get('id')
    columns = {};lines = [
        '# DISABLED COMPONENT CANDIDATE. Exact controller precondition and independent review required.',
        '# No native transaction/CAS; stop on failure and capture partial assigned state before recovery.',
        'from sigma.page import Page', 'from sigma.framing import GridFramed',
        'from sigma.inputTable import InputTable, LiteralEditableColumn, SingleSelectEditableColumn, InputTableIDColumn',
        'from sigma.column import FormulaColumn', 'from sigma.levelTable import LevelTable',
        'from sigma.control import add_text_control', 'from sigma.controlHandle import TextControl',
        'from sigma.action.inputTableEffects import insert_row',
        'from sigma.action.setControlValueEffect import set_control_value',
        'from sigma.action.controlFlow import if_else',
        'from sigma.formula import const_text, const_num, eq, neq, and_, or_, not_, is_null, is_not_null, subtract, lookup, count_if, coalesce, now, as_num, as_text']
    required_ops = {'Workflow':'text','Context':'text','State':'text','Created At':'datetime',
        'Source Hash':'text','Plan Hash':'text','Source Procedure Hash':'text','Replacement Decision ID':'text',
        'Snapshot Definition Hash':'text','Snapshot Capture Contract Hash':'text',
        'Authoring Baseline Document Version':'number','Authoring Baseline Contents Hash':'text',
        'Source Graph Hash':'text','Primary Identity Contract Hash':'text','Snapshot Schema Contract Hash':'text'}
    required_ops.update({c['operationFrozenColumn']:'number' for c in definition['cells']})
    required_ops.update({c['operationColumn']:c['type'] for c in definition['inputFields']})
    required = {'primary':{'Line Item':'text','Method':'text','Monthly Growth':'number',**{p:'number' for p in PERIODS}},
        'finance':{n:None for n in ('Source Row','Section','Line Item',*PERIODS)},
        'operations':required_ops, 'snapshots':SNAPSHOT_FIELDS}
    for index,(role,b) in enumerate(bindings.items()):
        e = elements.get(b.get('elementId'),{});pid=b.get('pageId');schema=b.get('columns',[])
        ids = [c.get('id') for c in schema];by_name={c.get('name'):c for c in schema}
        public_columns = {c['id']:c for c in e.get('columns',[])}
        if role in ('operations','snapshots'):
            native=[c for c in schema if c.get('class')=='InputTableIDColumn' and c.get('kind')=='row_id']
            if len(native)==1 and public_columns.get('ID')=={'id':'ID'}:
                # Public code rep's ID placeholder differs from the actual SDK
                # InputTableIDColumn. Never use that placeholder as a selector.
                public_columns.pop('ID');public_columns[native[0]['id']]={'id':native[0]['id']}
        if (pid not in pages or owners.get(b.get('elementId'))!=pid or e.get('name')!=b.get('name')
                or e.get('kind')!=('table' if role=='finance' else 'input-table')
                or len(set(ids))!=len(ids) or not all(ids) or set(ids)!=set(public_columns)
                or len(by_name)!=len(schema) or not set(required[role])<=set(by_name)):
            raise ValueError('complete assigned element/page/native schema required: '+role)
        if role in ('operations','snapshots') and pages[pid].get('visibility')!='hidden':
            raise ValueError('technical operations/snapshot host must be hidden')
        if role=='primary' and pid!='forecast-edit':raise ValueError('primary action context differs')
        for c in schema:
            pub = public_columns[c['id']]
            if any(pub.get(k)!=c.get(k) for k in ('name','type') if k in pub):raise ValueError('assigned column differs from saved source')
        for name,kind in required[role].items():
            c=by_name[name]
            if role=='finance':
                if c.get('class')!='FormulaColumn' or public_columns[c['id']].get('formula')!=c.get('formula'):
                    raise ValueError('finance projection must be an exact saved live formula')
            elif c.get('type')!=kind or c.get('class') not in ('LiteralEditableColumn','SingleSelectEditableColumn'):
                raise ValueError('editable assigned datatype/class differs: '+name)
        if role in ('operations','snapshots'):
            system=[c for c in schema if c.get('class')=='InputTableIDColumn' and c.get('kind')=='row_id']
            if len(system)!=1:raise ValueError('actual native system row_id column required')
            if any(c.get('class')!='LiteralEditableColumn' for c in schema if c not in system):raise ValueError('ordinary audit input literal columns required')
        if role=='finance' and phase=='confirm':continue
        page_var='p'+str(index);table='t'+str(index)
        lines += [f'{page_var} = sigma.workbook.get_page({pid!r})',f'if not isinstance({page_var}, Page): raise RuntimeError("page")',
            f'{table} = {page_var}.get_element({b["elementId"]!r})',
            f'if not isinstance({table}, {"LevelTable" if role=="finance" else "InputTable"}): raise RuntimeError("table class")',
            f'if {table}.name() != {b["name"]!r}: raise RuntimeError("table name")']
        if role!='finance':lines.append(f'if {table}.edit_mode() != "published-explore": raise RuntimeError("mode")')
        tags={'LiteralEditableColumn':'L','SingleSelectEditableColumn':'S','InputTableIDColumn':'I','FormulaColumn':'F'}
        expected={x['id']:{'n':x['name'],'c':tags[x['class']],'t':x.get('type') or '',
            **({'a':x['allowedValues']} if 'allowedValues' in x else {}), **({'f':x['formula']} if 'formula' in x else {})} for x in schema}
        lines += [f'expected = {expected!r}', 'seen = {}',f'for column in {table}.list_columns():',
            '    cid = column.id()', '    if cid not in expected or cid in seen: raise RuntimeError("schema")',
            '    seen[cid] = True', f'    if column.name() != expected[cid]["n"] or column.element_id() != {b["elementId"]!r}: raise RuntimeError("column identity")',
            '    if isinstance(column, LiteralEditableColumn):',
            '        if expected[cid]["c"] != "L" or column.type() != expected[cid]["t"]: raise RuntimeError("literal type")',
            '    elif isinstance(column, SingleSelectEditableColumn):',
            '        if expected[cid]["c"] != "S" or column.type() != expected[cid]["t"] or column.allowed_values() != expected[cid]["a"]: raise RuntimeError("select type")',
            '    elif isinstance(column, InputTableIDColumn):',
            '        if expected[cid]["c"] != "I" or column.kind() != "row_id": raise RuntimeError("native ID")',
            '    elif isinstance(column, FormulaColumn):',
            '        if expected[cid]["c"] != "F": raise RuntimeError("formula class")']
        lines.append('        if column.formula() != expected[cid]["f"]: raise RuntimeError("source formula")')
        lines += ['    else: raise RuntimeError("column class")',f'if len(seen) != {len(ids)}: raise RuntimeError("schema count")']
        for i,column in enumerate(schema):
            cv='c'+str(index)+'_'+str(i);columns[(role,column['name'])]=cv
            lines += [f'{cv} = {table}.get_column({column["id"]!r})',
                f'if not isinstance({cv}, {column["class"]}): raise RuntimeError("column")']
    role_tables={role:'t'+str(i) for i,role in enumerate(bindings)}
    c=lambda role,name:columns[(role,name)]
    identities=bindings['primary'].get('identities',[])
    expected_identity={(x['sourceRow'],x['lineItem']) for x in definition['inputFields']}
    if ({(x.get('sourceRow'),x.get('lineItem')) for x in identities}!=expected_identity or len(identities)!=12
            or len({x.get('nativeRowId') for x in identities})!=12 or any(not x.get('nativeRowId') for x in identities)
            or 'ID' not in {x['name'] for x in bindings['primary']['columns']}):
        raise ValueError('exact twelve original physical/business input identities required')
    identity={x['sourceRow']:x for x in identities};primary_id=c('primary','ID')
    rid = next(columns[('operations',x['name'])] for x in bindings['operations']['columns'] if x['class']=='InputTableIDColumn')
    operation_table=role_tables['operations'];snapshot_table=role_tables['snapshots']
    op_page='p'+str(list(bindings).index('operations'));ui_page='p'+str(list(bindings).index('primary'))
    op_lookup=lambda name:f'lookup({c("operations",name)}, op, {rid})'
    finite=lambda value:f'and_(is_not_null({value}), eq(subtract({value}, {value}), const_num(0)))'
    nullable=lambda value:f'or_(is_null({value}), {finite(value)})'
    nullsame=lambda a,b:f'or_(and_(is_null({a}), is_null({b})), and_(and_(is_not_null({a}), is_not_null({b})), eq({a}, {b})))'
    primary_guards=[f'eq(count_if(is_not_null({primary_id})), const_num(12))']
    for item in identities:
        pred=_join([f'eq({primary_id}, const_text({item["nativeRowId"]!r}))',f'eq({c("primary","Line Item")}, const_text({item["lineItem"]!r}))'])
        primary_guards += [f'eq(count_if({pred}), const_num(1))',f'eq(count_if(eq({c("primary","Line Item")}, const_text({item["lineItem"]!r}))), const_num(1))']
    source_values={}
    for i,field in enumerate(definition['inputFields']):
        value=f'lookup({c("primary",field["field"])}, const_text({identity[field["sourceRow"]]["nativeRowId"]!r}), {primary_id})'
        v='a'+str(i);cast='as_text' if field['type']=='text' else 'as_num'
        lines.append(f'{v} = {cast}({value})');source_values[field['operationColumn']]=v
        primary_guards.append(_join([f'eq({v}, const_text({method!r}))' for method in ('growth','manual','flat')], 'or_') if field['type']=='text' else nullable(v))
    source_graph_hash=digest([e for e in contents['elements'] if e['id'] not in {b['elementId'] for b in bindings.values()} and e.get('kind') not in ('control','button','text')])
    contract = {'definitionHash':definition['definitionHash'],'primaryIdentities':identities,
        'bindingHash':digest(bindings),'sourceGraphHash':source_graph_hash}
    constants={'Workflow':'snapshot','Context':'forecast-edit','State':'PREPARED','Source Hash':definition['sourceHash'],
        'Plan Hash':definition['planHash'],'Source Procedure Hash':definition['procedureHash'],
        'Replacement Decision ID':definition['decisionId'],'Snapshot Definition Hash':definition['definitionHash'],
        'Snapshot Capture Contract Hash':digest(contract),'Authoring Baseline Document Version':current['documentVersion'],
        'Authoring Baseline Contents Hash':digest(contents),'Source Graph Hash':source_graph_hash,
        'Primary Identity Contract Hash':digest(identities),'Snapshot Schema Contract Hash':digest(bindings['snapshots'])}
    lines += ['ui_cursor = 0', 'technical_cursor = 0',
        f'for existing in {ui_page}.list_elements(flat=False):',
        '    if isinstance(existing, GridFramed): ui_cursor = max(ui_cursor, existing.y() + existing.height())',
        f'for existing in {op_page}.list_elements(flat=False):',
        '    if isinstance(existing, GridFramed): technical_cursor = max(technical_cursor, existing.y() + existing.height())']
    # Baseline metadata is captured at Prepare. Confirm binds the operation to
    # the same capture contract, without substituting its later source version.
    if phase=='prepare':
        if any(e.get('controlId')=='forecast-snapshot-operation' for e in elements.values()):raise ValueError('Prepare marker exists; capture partial state, never replay constructor')
        finance_guards=[]
        for row in FINANCIAL_ROWS:
            cell=next(x for x in definition['cells'] if x['sourceRow']==row)
            finance_guards.append(f'eq(count_if(eq({c("finance","Source Row")}, const_num({row}))), const_num(1))')
            finance_guards.append(f'eq(count_if('+_join([f'eq({c("finance","Source Row")}, const_num({row}))',
                f'eq({c("finance","Line Item")}, const_text({cell["lineItem"]!r}))',f'eq({c("finance","Section")}, const_text({cell["section"]!r}))'])+'), const_num(1))')
        for i,cell in enumerate(definition['cells']):
            v='v'+str(i);lines.append(f'{v} = as_num(lookup({c("finance",cell["period"])}, const_num({cell["sourceRow"]}), {c("finance","Source Row")}))')
            source_values[cell['operationFrozenColumn']]=v;finance_guards.append(finite(v))
        # All getters/metadata assertions precede the first mutation.
        lines += [f'op = add_text_control({op_page}, value="", variable_name="forecast-snapshot-operation")',
            'op.set_frame(1, technical_cursor, min(24, op.parent_grid_width()), 2)',f'button = {ui_page}.add_button("Prepare forecast snapshot")',
            'button.set_frame(1, ui_cursor, min(12, button.parent_grid_width()), 2)']
        values={c('operations',n):repr(v) for n,v in constants.items()}
        values[c('operations','Created At')]='now()';values.update({c('operations',n):v for n,v in source_values.items()})
        lines += [f'insert = insert_row({operation_table}, '+ '{'+', '.join(k+': '+v for k,v in values.items())+'}, name="Freeze complete forecast snapshot")',
            'returned = insert.returns()', 'sequence = button.on_click("Prepare complete snapshot", [insert, set_control_value(op, returned)])']
        pending=_join([f'eq({c("operations","Workflow")}, const_text("snapshot"))',f'eq({c("operations","State")}, const_text("PREPARED"))'])
        guard=_join([*primary_guards,*finance_guards,'eq(op, const_text(""))',f'eq(as_num(coalesce(count_if({pending}), const_num(0))), const_num(0))'])
        effect_count=1
    else:
        control=elements.get(operation_control_id,{})
        if (control.get('kind')!='control' or control.get('controlType')!='text'
                or control.get('controlId')!='forecast-snapshot-operation' or owners.get(operation_control_id)!=bindings['operations']['pageId']):
            raise ValueError('Prepare actual assigned native operation TextControl required')
        lines += [f'op = {op_page}.get_element({operation_control_id!r})',
            'if not isinstance(op, TextControl): raise RuntimeError("operation control")',
            'if op.variable_name() != "forecast-snapshot-operation": raise RuntimeError("operation control variable")']
        guard_parts=[*primary_guards,'neq(op, const_text(""))',f'eq(count_if(eq({rid}, op)), const_num(1))']
        # Authoring version/hash belong to the Prepare operation, whose source
        # contract must be supplied unchanged by the caller after readback.
        prepare_contract=prepared_contract
        if prepare_contract:
            if set(prepare_contract)!=set(constants):raise ValueError('exact Prepare metadata contract required')
            for name in constants:
                if name not in ('Authoring Baseline Document Version','Authoring Baseline Contents Hash') and prepare_contract[name]!=constants[name]:
                    raise ValueError('Prepare metadata contract changed')
            constants=prepare_contract
        else:
            raise ValueError('actual Prepare source metadata contract required for Confirm')
        for name,value in constants.items():
            literal='const_num('+str(value)+')' if type(value) in (int,float) else 'const_text('+repr(value)+')'
            guard_parts.append(f'eq({op_lookup(name)}, {literal})')
        lines += ['frozen_assumptions = ['+', '.join(c('operations',field['operationColumn']) for field in definition['inputFields'])+']',
            'current_assumptions = ['+', '.join(source_values[field['operationColumn']] for field in definition['inputFields'])+']']
        # A typed loop keeps all 108 identical guards/effects within Code
        # Mode's request size; data arrays contain identities, never values.
        lines += ['frozen_columns = ['+', '.join(columns[('operations',cell['operationFrozenColumn'])] for cell in definition['cells'])+']',
            f'snapshot_keys = {[cell["snapshotKey"] for cell in definition["cells"]]!r}',
            f'source_cells = {[cell["sourceCell"] for cell in definition["cells"]]!r}',
            f'source_rows = {[cell["sourceRow"] for cell in definition["cells"]]!r}',
            f'line_items = {[cell["lineItem"] for cell in definition["cells"]]!r}',
            f'sections = {[cell["section"] for cell in definition["cells"]]!r}',
            f'periods = {[cell["period"] for cell in definition["cells"]]!r}',
            f'capture_time = {op_lookup("Created At")}',
            f'confirm_guard_parts = [{_join([*guard_parts, "is_not_null(capture_time)"])}]',
            'for i in range(96):', f'    before = lookup(frozen_assumptions[i], op, {rid})',
            '    actual = current_assumptions[i]', f'    confirm_guard_parts.append({nullsame("actual","before")})',
            'allowed_parts = [eq(const_num(0), const_num(1))]', 'effects = []', 'for i in range(108):',
            f'    frozen = as_num(lookup(frozen_columns[i], op, {rid}))',
            f'    key = and_(eq({c("snapshots","Operation ID")}, op), eq({c("snapshots","Snapshot Key")}, const_text(snapshot_keys[i])))',
            f'    allowed_parts.append(eq({c("snapshots","Snapshot Key")}, const_text(snapshot_keys[i])))']
        payload={'Operation ID':'op','Snapshot Key':'const_text(snapshot_keys[i])','Source Sheet':'const_text("Monthly Model")',
            'Source Cell':'const_text(source_cells[i])','Source Row':'const_num(source_rows[i])',
            'Line Item':'const_text(line_items[i])','Section':'const_text(sections[i])',
            'Period':'const_text(periods[i])','Value':'frozen','Captured At':'capture_time',
            **{n:op_lookup(n) for n in ('Source Hash','Plan Hash','Source Procedure Hash','Replacement Decision ID')}}
        matching=_join([f'eq({c("snapshots",name)}, {value})' for name,value in payload.items()])
        lines += [f'    matching = and_(key, {matching})',
            f'    confirm_guard_parts.append(and_({finite("frozen")}, or_(eq(count_if(key), const_num(0)), and_(eq(count_if(key), const_num(1)), eq(count_if(matching), const_num(1))))))',
            '    effects.append(if_else([(eq(count_if(key), const_num(0)), [insert_row('+snapshot_table+', {'+', '.join(c('snapshots',n)+': '+v for n,v in payload.items())+'})])], []))']
        # Deep left folds were rejected by the live condition setter. Balance
        # only associative booleans; every nullable/conflict/foreign-key guard
        # and all frozen-value arithmetic remain unchanged.
        lines += _balance_lines('allowed_parts', 'or_', 109)
        lines += [f'confirm_guard_parts.append(eq(count_if(and_(eq({c("snapshots","Operation ID")}, op), not_(allowed_parts[0]))), const_num(0)))']
        lines += _balance_lines('confirm_guard_parts', 'and_', 206)
        lines += [
            f'button = {ui_page}.add_button("Commit prepared forecast snapshot")',
            'button.set_frame(1, ui_cursor, min(12, button.parent_grid_width()), 2)',
            'sequence = button.on_click("Fill missing frozen snapshot values", effects)']
        guard='confirm_guard_parts[0]'
        effect_count=108
    lines += ['sequence.set_disabled(True)', f'sequence.set_condition({guard})','sequence.remove_success_toast()',
        'sequence.set_disabled(True)', 'if not sequence.is_disabled(): raise RuntimeError("disabled")',
        'print({"snapshotPhase": '+repr(phase)+', "buttonId": button.id(), "sequenceId": sequence.id(), "operationControlId": op.id(), "enabled": False, "businessWrites": False})']
    program='\n'.join(lines)+'\n';compile(program,'snapshot-'+phase,'exec')
    result={'phase':phase,'program':program,'programHash':hashlib.sha256(program.encode()).hexdigest(),
        'controllerPrecondition':{'documentId':current['workbookId'],'version':current['documentVersion'],'contentsHash':digest(contents)},
        'preparedContract':constants,'captureContract':contract,'frozenValueCount':204,'insertEffectCount':effect_count,
        'definitionHash':definition['definitionHash'],'status':'DISABLED_COMPONENT_CANDIDATE','productionStatus':'BLOCKED',
        'limitations':['backend type admission pending','source/row preflight must be fresh before submit','no version CAS or multi-effect atomicity',
            'published roles/recovery/concurrency and full persisted readback pending','no optimistic success or automatic outcome update']}
    return result
