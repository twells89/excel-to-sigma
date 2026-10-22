"""Draft a ``build-table-report.py`` plan from a workbook of formal Excel Tables.

The workbook shape (inventory projection, sales pipeline, headcount plan) is:

* formal Tables - literal columns are data, calculated columns are row formulas;
* report sheets - grids whose cells are formulas over those tables.

``draft_table_report`` recognises each grid as one of three frames, lowers every
formula to the IR in ``table_report_model`` and then recomputes *every* recognised cell
from the table data, comparing it with Excel's cached value. Anything it does not
recognise, and any mismatch, raises ``PlanDraftError`` naming sheet!cell; nothing is
coerced and no cached value is ever frozen into the plan as a result.

Frames
------
``summary``    row labels x columns; each cell aggregates a table (SUMIFS/COUNTIFS/AVERAGEIFS
               filtered by the row label) or is arithmetic/lookup over the row's other cells.
``expansion``  ordinal rows (1..N) pulling table rows with INDEX, periods across: the table
               crossed with a period axis, the cell formula as a calculated column.
``keyed``      labels x periods whose cells pick one table row per (label, period), optionally
               as a running balance (``prev -/+ terms``) -> CumulativeSum partitioned by label.
"""
from __future__ import annotations

import re
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

try:
    from . import table_report_model as trm
    from .excel_expression import (Binary, Boolean, ColumnRange, ExpressionError, Function,
                                   Group, Number, Percent, Range, Reference, StructuredReference, Text,
                                   Unary, parse_expression)
    from .migration_gate import source_sha256
except ImportError:  # direct import when scripts/lib itself is on sys.path
    import table_report_model as trm
    from excel_expression import (Binary, Boolean, ColumnRange, ExpressionError, Function,
                                  Group, Number, Percent, Range, Reference, StructuredReference, Text,
                                  Unary, parse_expression)
    from migration_gate import source_sha256

INTENTS = ("read-only", "editable", "what-if")
CRIT_RE = re.compile(r"^(>=|<=|<>|>|<|=)?(.*)$", re.S)


class PlanDraftError(RuntimeError):
    """The workbook does not have a shape this drafter supports."""


class NoTableReportShape(PlanDraftError):
    """The workbook has no formula-driven table report at all (not a near miss)."""


# ---------------------------------------------------------------- small helpers
def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _text(value):
    return value.strip() if isinstance(value, str) and value.strip() else None


def _is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def _addr(sheet, row, col):
    return f"{sheet}!{get_column_letter(col)}{row}"


def _rc(ref):
    """(row, col) of a parsed single-cell reference."""
    match = re.fullmatch(r"([A-Z]+)(\d+)", ref.address)
    return int(match.group(2)), column_index_from_string(match.group(1))


def _fmt_kind(cell):
    fmt = cell.number_format or ""
    if "%" in fmt:
        return "percent"
    if any(sym in fmt for sym in "$€£¥"):
        return "currency"
    return None


def _title(text):
    """A frame/table title Sigma element references can hold (no [ ] /)."""
    return re.sub(r"\s+", " ", re.sub(r"[\[\]/]", "-", str(text))).strip()


def _plain(value):
    """JSON-safe cached value (dates -> ISO)."""
    if isinstance(value, datetime):
        return _date(value, "cached value").isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _clean_name(value, where):
    name = _text(value) if not isinstance(value, (datetime, date)) else None
    if name is None:
        name = str(value).strip() if value is not None else None
    if not name:
        raise PlanDraftError(f"{where}: empty name")
    if re.search(r"[\[\]/]", name):
        raise PlanDraftError(f"{where}: name {name!r} contains [ ] or /, which Sigma column refs cannot hold")
    return name


def _date(value, where):
    try:
        return trm.parse_date(value)
    except trm.PlanError as error:
        raise PlanDraftError(f"{where}: {error}") from error


def col_values(tab, name):
    """Raw values of a table column: literals as typed, calculated columns as cached."""
    column = tab.col(name)
    return column.formulas if column.kind == "literal" else column.cached


WILDCARD_RE = re.compile(r"[*?~]")


def _wildcard_free(value, where, what):
    """Excel reads ``* ? ~`` in a label/criteria cell as wildcards (and ``~`` as an escape); Sigma matches
    text exactly, so such a label would silently match different rows. Refuse it, naming the cell."""
    if isinstance(value, str) and WILDCARD_RE.search(value):
        raise PlanDraftError(f"{where}: {what} {value!r} contains * ? or ~, which Excel treats as wildcards in "
                             "criteria and MATCH but Sigma matches literally; rename the label")


# ---------------------------------------------------------------- workbook model
class Col:
    def __init__(self, name, kind, cells, formulas, cached, fmt):
        self.name, self.kind = name, kind            # kind: literal | calculated
        self.cells, self.formulas, self.cached, self.fmt = cells, formulas, cached, fmt
        self.type = None


class Tab:
    """One formal table."""

    def __init__(self, name, ws, min_col, min_row, max_row, columns):
        self.name, self.ws, self.min_col, self.min_row, self.max_row = name, ws, min_col, min_row, max_row
        self.columns = columns
        self.n = len(columns[0].cells) if columns else 0

    def col(self, name):
        for column in self.columns:
            if column.name == name:
                return column
        return None

    def cells_of(self, name):
        column = self.col(name)
        return column.cells if column is not None else []


def _infer_type(values, where, name):
    seen = {"number": None, "text": None, "datetime": None, "boolean": None}
    for coord, value in values:
        if value is None:
            continue
        if isinstance(value, bool):
            kind = "boolean"
        elif _is_number(value):
            kind = "number"
        elif isinstance(value, (datetime, date)):
            kind = "datetime"
        elif isinstance(value, str):
            kind = "text"
        else:
            raise PlanDraftError(f"{coord}: unsupported value {value!r} in column {name!r}")
        seen[kind] = seen[kind] or coord
    kinds = [k for k, v in seen.items() if v]
    if len(kinds) > 1:
        raise PlanDraftError(f"{seen[kinds[1]]}: column {name!r} mixes {kinds[0]} and {kinds[1]} values "
                             f"(first {kinds[0]} at {seen[kinds[0]]})")
    return kinds[0] if kinds else None


def _guess_empty_type(name):
    lowered = name.lower()
    if re.search(r"\bdate\b|\bat$|\btime\b|\bstart\b|\bend\b|\bdue\b|\bweek\b", lowered):
        return "datetime"
    if re.search(r"qty|units|amount|count|cost|price|total|quantity", lowered):
        return "number"
    return "text"


def load_tables(wf, wv):
    tabs = OrderedDict()
    for ws in wf.worksheets:
        for table in ws.tables.values():
            from openpyxl.utils import range_boundaries
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            where = f"{ws.title}!{table.ref} ({table.name})"
            if (table.totalsRowCount or 0) > 0 or table.totalsRowShown:
                raise PlanDraftError(f"{where}: tables with a totals row are not supported")
            if table.name in tabs:
                raise PlanDraftError(f"duplicate table name {table.name}")
            names = []
            for col in range(min_col, max_col + 1):
                names.append(_clean_name(ws.cell(min_row, col).value, f"{_addr(ws.title, min_row, col)}"))
            if len(set(names)) != len(names):
                raise PlanDraftError(f"{where}: duplicate column names")
            data_rows = list(range(min_row + 1, max_row + 1))
            empty = all(ws.cell(r, c).value is None for r in data_rows for c in range(min_col, max_col + 1))
            if empty:
                data_rows = []
            columns = []
            for index, col in enumerate(range(min_col, max_col + 1)):
                cells = [_addr(ws.title, r, col) for r in data_rows]
                raw = [ws.cell(r, col).value for r in data_rows]
                formulas = [v for v in raw if _is_formula(v)]
                if raw and len(formulas) == len(raw):
                    kind = "calculated"
                elif not formulas:
                    kind = "literal"
                else:
                    first_literal = next(c for c, v in zip(cells, raw) if not _is_formula(v))
                    raise PlanDraftError(f"{first_literal}: column {names[index]!r} of {table.name} mixes "
                                         "formulas and literal values; a table column is either data or "
                                         "a calculated column")
                for coord, value in zip(cells, raw):
                    if value is not None and not isinstance(value, (str, int, float, bool, datetime, date)):
                        raise PlanDraftError(f"{coord}: unsupported cell content {type(value).__name__}")
                cached = [wv[ws.title].cell(r, col).value for r in data_rows]
                fmt = _fmt_kind(ws.cell(data_rows[0], col)) if data_rows else None
                columns.append(Col(names[index], kind, cells, raw, cached, fmt))
            tabs[table.name] = Tab(table.name, ws, min_col, min_row, max_row, columns)
            tabs[table.name].data_rows = data_rows
            tabs[table.name].ref = table.ref
    return tabs


def _covered(tabs):
    covered = {}
    for tab in tabs.values():
        cells = covered.setdefault(tab.ws.title, set())
        for r in range(tab.min_row, tab.max_row + 1):
            for c in range(tab.min_col, tab.min_col + len(tab.columns)):
                cells.add((r, c))
    return covered


# ---------------------------------------------------------------- AST helpers
def _norm(node, row, col, sheet):
    """Canonical text of a formula with references relative to ``(row, col)``."""
    if isinstance(node, Number):
        return f"n{node.value!r}"
    if isinstance(node, Text):
        return f"s{node.value!r}"
    if isinstance(node, Boolean):
        return f"b{node.value}"
    if isinstance(node, Reference):
        ref = node.value
        if ref.name:
            return f"name:{ref.name}"
        r, c = _rc(ref)
        rr = f"${r}" if ref.row_absolute else f"[{r - row}]"
        cc = f"${c}" if ref.column_absolute else f"[{c - col}]"
        return f"ref({ref.sheet or sheet}!R{rr}C{cc})"
    if isinstance(node, Range):
        return f"({_norm(Reference(node.start), row, col, sheet)}:{_norm(Reference(node.end), row, col, sheet)})"
    if isinstance(node, ColumnRange):
        return f"cols({node.sheet or sheet}!{node.start_column}:{node.end_column})"
    if isinstance(node, StructuredReference):
        return f"struct({node.table}{node.specification})"
    if isinstance(node, Group):
        return f"({_norm(node.expression, row, col, sheet)})"
    if isinstance(node, Unary):
        return f"u{node.operator}({_norm(node.operand, row, col, sheet)})"
    if isinstance(node, Percent):
        return f"pct({_norm(node.operand, row, col, sheet)})"
    if isinstance(node, Binary):
        return f"({_norm(node.left, row, col, sheet)}{node.operator}{_norm(node.right, row, col, sheet)})"
    if isinstance(node, Function):
        return f"{node.name}(" + ",".join(_norm(a, row, col, sheet) for a in node.arguments) + ")"
    raise PlanDraftError(f"unsupported expression node {type(node).__name__}")


def _struct(node, where):
    """``tbl[...]`` -> ("this", col) | ("col", col) | ("range", first, last)."""
    spec = node.specification
    if "'" in spec:
        raise PlanDraftError(f"{where}: escaped structured reference {node.table}{spec} is not supported")
    this = re.fullmatch(r"\[\[#This Row\],\[([^\]]+)\]\]", spec) or re.fullmatch(r"\[@\[?([^\]]+)\]?\]", spec)
    if this:
        return "this", this.group(1)
    span = re.fullmatch(r"\[\[([^\]]+)\]:\[([^\]]+)\]\]", spec)
    if span:
        return "range", span.group(1), span.group(2)
    column = re.fullmatch(r"\[([^\[\]#@]+)\]", spec)
    if column:
        return "col", column.group(1)
    raise PlanDraftError(f"{where}: structured reference {node.table}{spec} is not supported")


class Registry:
    """Shared drafting state: parameters, synthetic tables, review notes."""

    def __init__(self, src, wf, wv, tabs):
        self.src, self.wf, self.wv, self.tabs = src, wf, wv, tabs
        self.covered = _covered(tabs)
        self.params = OrderedDict()      # (sheet, address) -> {name,value,cell}
        self.param_names = {}
        self.synthetic = OrderedDict()   # table name -> plan table dict
        self.added_columns = []          # (table, column) calculated columns the drafter synthesised

    def param(self, sheet, address, where, allow=("number",)):
        """Name of the parameter for a literal cell. ``allow`` widens the kinds a context accepts
        (the period-ledger recogniser takes text and date parameters; every other grid numbers only)."""
        key = (sheet, address)
        if key in self.params:
            if self.params[key].get("kind", "number") not in allow:
                raise PlanDraftError(f"{where}: parameter {sheet}!{address} is "
                                     f"{self.params[key].get('kind')}, not a number")
            return self.params[key]["name"]
        if sheet not in self.wf.sheetnames:
            raise PlanDraftError(f"{where}: sheet {sheet!r} not found")
        ws = self.wf[sheet]
        match = re.fullmatch(r"([A-Z]+)(\d+)", address)
        row, col = int(match.group(2)), column_index_from_string(match.group(1))
        if (row, col) in self.covered.get(sheet, ()):
            raise PlanDraftError(f"{where}: {sheet}!{address} is inside a table; reference it with structured "
                                 "table syntax")
        value = ws.cell(row, col).value
        if _is_formula(value):
            raise PlanDraftError(f"{where}: {sheet}!{address} is a formula cell, not a literal parameter")
        kind = "number"
        if _is_number(value):
            pass
        elif "date" in allow and isinstance(value, (datetime, date)):
            kind, value = "date", _date(value, f"{sheet}!{address}").isoformat()
        elif "text" in allow and isinstance(value, str) and value.strip():
            kind = "text"
        else:
            raise PlanDraftError(f"{where}: parameter {sheet}!{address} is {value!r}; only numeric literal "
                                 "parameters are supported")
        label = next((_text(ws.cell(row, c).value) for c in range(col - 1, 0, -1)
                      if _text(ws.cell(row, c).value) and not _is_formula(ws.cell(row, c).value)), None)
        name = label or f"{sheet}!{address}"
        if name in self.param_names:
            name = f"{name} ({sheet}!{address})"
        self.param_names[name] = key
        self.params[key] = {"name": name, "value": value, "cell": f"{sheet}!{address}"}
        if kind != "number":
            self.params[key]["kind"] = kind
        return name


# ---------------------------------------------------------------- lowering
class Ctx:
    """Context hooks; the defaults refuse (and say what was seen)."""

    def ref(self, low, ref):
        raise low.fail(f"cell reference {ref.address or ref.name} is not supported in this formula context")

    def struct(self, low, node):
        raise low.fail(f"structured reference {node.table}{node.specification} is not supported here")

    def func(self, low, node):
        return None

    def named(self, low, name):
        raise low.fail(f"named range {name!r} is not supported")

    def range_(self, low, node):
        raise low.fail("a cell range outside SUM/MIN/MAX/aggregate arguments is not supported")


class Lowerer:
    def __init__(self, reg, sheet, where, ctx):
        self.reg, self.sheet, self.where, self.ctx = reg, sheet, where, ctx

    def fail(self, message):
        return PlanDraftError(f"{self.where}: {message}")

    # ---- generic expression
    def lower(self, node):
        if isinstance(node, Number):
            return ["n", node.value]
        if isinstance(node, Text):
            return ["s", node.value]
        if isinstance(node, Boolean):
            return ["b", node.value]
        if isinstance(node, Group):
            return self.lower(node.expression)
        if isinstance(node, Unary):
            inner = self.lower(node.operand)
            return ["neg", inner] if node.operator == "-" else inner
        if isinstance(node, Percent):
            return ["op", "/", self.lower(node.operand), ["n", 100]]
        if isinstance(node, Binary):
            if node.operator == "&":
                raise self.fail("text concatenation (&) is only supported inside aggregate criteria")
            return ["op", node.operator, self.lower(node.left), self.lower(node.right)]
        if isinstance(node, Reference):
            if node.value.name:
                return self.ctx.named(self, node.value.name)
            return self.ctx.ref(self, node.value)
        if isinstance(node, StructuredReference):
            return self.ctx.struct(self, node)
        if isinstance(node, (Range, ColumnRange)):
            return self.ctx.range_(self, node)
        if isinstance(node, Function):
            handled = self.ctx.func(self, node)
            return handled if handled is not None else self.function(node)
        raise self.fail(f"unsupported expression {type(node).__name__}")

    def literal_int(self, node, what):
        if isinstance(node, Group):
            return self.literal_int(node.expression, what)
        if isinstance(node, Unary) and node.operator == "-":
            return -self.literal_int(node.operand, what)
        if isinstance(node, Number) and float(node.value) == int(node.value):
            return int(node.value)
        raise self.fail(f"{what} must be a literal integer")

    def function(self, node):
        name, args = node.name, node.arguments
        n = len(args)
        if name == "IF":
            if n != 3:
                raise self.fail("IF needs three arguments (no implicit FALSE branch)")
            return ["f", "IF"] + [self.lower(a) for a in args]
        if name in ("AND", "OR"):
            return ["f", name] + [self.lower(a) for a in args]
        if name == "NOT" and n == 1:
            return ["f", "NOT", self.lower(args[0])]
        if name in ("MIN", "MAX") and n >= 2:
            return ["f", name] + [self.lower(a) for a in args]
        if name == "SUM" and n >= 1:
            return ["f", "SUM"] + [self.lower(a) for a in args]
        if name == "ABS" and n == 1:
            return ["f", "ABS", self.lower(args[0])]
        if name == "MOD" and n == 2:
            return ["f", "MOD", self.lower(args[0]), self.lower(args[1])]
        if name == "ROUND" and n == 2:
            return ["f", "ROUND", self.lower(args[0]), self.lower(args[1])]
        if name == "CEILING" and n == 2:
            return ["f", "CEILING", self.lower(args[0]), self.lower(args[1])]
        if name == "IFERROR" and n == 2:
            return ["f", "IFERROR", self.lower(args[0]), self.lower(args[1])]
        if name in ("MONTH", "YEAR") and n == 1:
            return ["f", name, self.lower(args[0])]
        if name == "DATE" and n == 3:
            y, m, d = (self.literal_int(a, "DATE argument") for a in args)
            try:
                return ["d", date(y, m, d).isoformat()]
            except ValueError as error:
                raise self.fail(f"invalid DATE({y},{m},{d})") from error
        if name == "EOMONTH" and n == 2:
            if self.literal_int(args[1], "EOMONTH months") != 0:
                raise self.fail("EOMONTH with a non-zero month offset is not supported")
            inner = self.lower(args[0])
            if inner[0] == "d":
                return ["d", trm.eomonth0(date.fromisoformat(inner[1])).isoformat()]
            return ["f", "EOMONTH0", inner]
        if name == "XLOOKUP":
            return self.xlookup(node)
        if name == "INDEX":
            return self.index_match(node)
        raise self.fail(f"unsupported function {name}")

    # ---- lookups into another table
    def table_ref(self, node, what):
        if not isinstance(node, StructuredReference):
            raise self.fail(f"{what} must be a table column reference such as tblX[Col]")
        kind = _struct(node, self.where)
        if kind[0] != "col":
            raise self.fail(f"{what} must be a whole table column, got {node.table}{node.specification}")
        table = self.reg.tabs.get(node.table)
        if table is None:
            raise self.fail(f"unknown table {node.table!r}")
        if table.col(kind[1]) is None:
            raise self.fail(f"table {node.table} has no column {kind[1]!r}")
        return table, kind[1]

    def lookup_node(self, key_ir, key_table, key_col, value_table, value_col, default, wildcards=False):
        if key_table is not value_table:
            raise self.fail("lookup key and result columns belong to different tables")
        keys = [k for k in col_values(key_table, key_col) if k is not None]
        folded = [k.lower() if isinstance(k, str) else k for k in keys]
        if len(set(folded)) != len(folded):
            raise self.fail(f"lookup key {key_table.name}[{key_col}] has duplicate values (text keys are "
                            "compared ignoring case); Sigma's Lookup would not match Excel's first-match rule")
        if wildcards:     # MATCH(.., 0) reads * ? ~ in its lookup value as wildcards
            for coord, key in zip(key_table.cells_of(key_col), col_values(key_table, key_col)):
                _wildcard_free(key, coord, f"{key_table.name}[{key_col}] key")
        return ["lk", key_table.name, key_ir, key_col, value_col, default]

    def xlookup(self, node):
        args = node.arguments
        if not 3 <= len(args) <= 6:
            raise self.fail("XLOOKUP argument count")
        if len(args) >= 5 and not (isinstance(args[4], Number) and args[4].value == 0):
            raise self.fail("XLOOKUP match mode must be exact (0)")
        if len(args) >= 6 and not (isinstance(args[5], Number) and args[5].value == 1):
            raise self.fail("XLOOKUP search mode must be first-to-last (1)")
        key = self.lower(args[0])
        key_table, key_col = self.table_ref(args[1], "XLOOKUP lookup column")
        value_table, value_col = self.table_ref(args[2], "XLOOKUP return column")
        default = self.lower(args[3]) if len(args) >= 4 else None
        return self.lookup_node(key, key_table, key_col, value_table, value_col, default)

    def index_match(self, node):
        args = node.arguments
        if len(args) == 2 and isinstance(args[1], Function) and args[1].name == "MATCH":
            match = args[1].arguments
            if len(match) == 3 and isinstance(match[2], Number) and match[2].value == 0:
                value_table, value_col = self.table_ref(args[0], "INDEX return column")
                key_table, key_col = self.table_ref(match[1], "MATCH lookup column")
                return self.lookup_node(self.lower(match[0]), key_table, key_col, value_table, value_col, None,
                                        wildcards=True)
        raise self.fail("INDEX is only supported as INDEX(tblX[Col], MATCH(key, tblX[Key], 0)) here")

    def param_ref(self, ref):
        sheet = ref.sheet or self.sheet
        name = self.reg.param(sheet, ref.address, self.where)
        return ["p", name]


# ---- table calculated columns
class TableCtx(Ctx):
    def __init__(self, table, row=None):
        self.table, self.row = table, row

    def ref(self, low, ref):
        """Same-row A1 references (``=I3*J3``) read the table's own columns; other cells are parameters."""
        table = self.table
        if self.row is not None and ref.address and (ref.sheet or low.sheet) == table.ws.title:
            row, col = _rc(ref)
            if table.min_row < row <= table.max_row and table.min_col <= col < table.min_col + len(table.columns):
                if row != self.row:
                    raise low.fail(f"{ref.address} is on another row of {table.name}; only same-row references "
                                   "are supported in a calculated column")
                return ["c", table.columns[col - table.min_col].name]
        return low.param_ref(ref)

    def struct(self, low, node):
        kind = _struct(node, low.where)
        if kind[0] != "this" or node.table != self.table.name:
            raise low.fail(f"only same-row references to {self.table.name} are supported here, "
                           f"got {node.table}{node.specification}")
        if self.table.col(kind[1]) is None:
            raise low.fail(f"table {self.table.name} has no column {kind[1]!r}")
        return ["c", kind[1]]


# ---- criteria
def _crit(low, node, label_ok):
    """Aggregate criterion operand -> ``("label",)`` | ``(op, valueIR)``."""
    if isinstance(node, Group):
        return _crit(low, node.expression, label_ok)
    if isinstance(node, Text):
        match = CRIT_RE.match(node.value)
        op, rest = (match.group(1) or "="), match.group(2)
        if re.search(r"[*?~]", rest):
            raise low.fail(f"wildcard criterion {node.value!r} is not supported")
        try:
            return op, ["n", float(rest) if "." in rest or "e" in rest.lower() else int(rest)]
        except ValueError:
            return op, ["s", rest]
    if isinstance(node, Binary) and node.operator == "&" and isinstance(node.left, Text) \
            and node.left.value in (">=", "<=", "<>", ">", "<", "="):
        return node.left.value, low.lower(node.right)
    if isinstance(node, Reference) and label_ok(node.value):
        return ("label",)
    return "=", low.lower(node)


def _and(terms):
    if not terms:
        return None
    if len(terms) == 1:
        return terms[0]
    return ["f", "AND"] + terms


# ---- summary grids
class Block:
    """A header row, label column and rows of one report grid."""

    def __init__(self, ws, header_row, label_col, last_col, first_row, last_row):
        self.ws, self.sheet = ws, ws.title
        self.hr, self.lc, self.last_col = header_row, label_col, last_col
        self.first_row, self.last_row = first_row, last_row
        self.rows = list(range(first_row, last_row + 1))
        self.title = None
        self.roles = {}   # col -> dict(role, name, ...)

    def header(self, col):
        return self.ws.cell(self.hr, col).value

    def label(self, row):
        return self.ws.cell(row, self.lc).value

    @property
    def cols(self):
        return list(range(self.lc + 1, self.last_col + 1))


def find_blocks(wf, covered, tabs):
    blocks = []
    for ws in wf.worksheets:
        cov = covered.get(ws.title, set())

        def cell(r, c):
            return None if (r, c) in cov else ws.cell(r, c).value

        row = 1
        while row <= ws.max_row:
            present = [c for c in range(1, ws.max_column + 1) if cell(row, c) is not None]
            if len(present) >= 2:
                first = present[0]
                run = [first]
                while cell(row, run[-1] + 1) is not None:
                    run.append(run[-1] + 1)
                headers = [cell(row, c) for c in run]
                header_like = all((isinstance(h, (str, datetime, date)) or _is_number(h))
                                  and not _is_formula(h) for h in headers)
                if len(run) >= 2 and header_like and cell(row + 1, first) is not None:
                    last_row = row + 1
                    while last_row + 1 <= ws.max_row and cell(last_row + 1, first) is not None:
                        last_row += 1
                    block = Block(ws, row, first, run[-1], row + 1, last_row)
                    above = [r for r in range(max(1, row - 3), row)
                             if len([c for c in range(1, ws.max_column + 1) if cell(r, c) is not None]) == 1]
                    if above:
                        text = _text(cell(above[-1], next(c for c in range(1, ws.max_column + 1)
                                                          if cell(above[-1], c) is not None)))
                        block.title = text
                    blocks.append(block)
                    row = last_row + 1
                    continue
            row += 1
    return blocks


def _formula_cells_outside(wf, covered, blocks):
    inside = {}
    for b in blocks:
        for r in b.rows:
            for c in range(b.lc, b.last_col + 1):
                inside.setdefault(b.sheet, set()).add((r, c))
    stray = []
    for ws in wf.worksheets:
        cov, ins = covered.get(ws.title, set()), inside.get(ws.title, set())
        for row in ws.iter_rows():
            for c in row:
                if (c.row, c.column) in cov or (c.row, c.column) in ins:
                    continue
                if c.data_type == "f" or _is_formula(c.value):
                    stray.append(f"{ws.title}!{c.coordinate}")
    return stray


class SummaryCtx(Ctx):
    """Lowers one summary column's cell. Records the aggregate sources/keys it meets."""

    def __init__(self, block, name_of, kinds_of, grids, tabs, reg, row):
        self.b, self.name_of, self.kinds_of, self.grids = block, name_of, kinds_of, grids
        self.tabs, self.reg, self.row = tabs, reg, row
        self.sources, self.by_cols, self.deps = set(), set(), set()
        self.calc_columns = []

    # --- single cells
    def ref(self, low, ref):
        sheet = ref.sheet or low.sheet
        row, col = _rc(ref)
        if sheet == self.b.sheet and row == self.row and self.b.lc <= col <= self.b.last_col:
            if col == self.b.lc:
                return ["key"]
            kind = self.kinds_of.get(col)
            if kind in ("summary", "literal"):
                self.deps.add(self.name_of[col])
                return ["c", self.name_of[col]]
            raise low.fail(f"{ref.address} is a {kind or 'matrix'} column; reference ranges over it with "
                           "SUM/MIN/MAX instead")
        if sheet == self.b.sheet and self.b.first_row <= row <= self.b.last_row:
            raise low.fail(f"{ref.address} is on another row of this grid; only same-row references are supported")
        return low.param_ref(ref)

    def is_label(self, low, ref):
        sheet = ref.sheet or low.sheet
        if not ref.address:
            return False
        row, col = _rc(ref)
        return sheet == self.b.sheet and col == self.b.lc and row == self.row

    # --- ranges inside SUM/MIN/MAX/AVERAGE
    def same_row_range(self, low, node):
        if not isinstance(node, Range):
            return None
        sheet = node.start.sheet or low.sheet
        r1, c1 = _rc(node.start)
        r2, c2 = _rc(node.end)
        if sheet != self.b.sheet or r1 != self.row or r2 != self.row:
            return None
        return list(range(min(c1, c2), max(c1, c2) + 1))

    def func(self, low, node):
        name, args = node.name, node.arguments
        if name in ("SUMIFS", "COUNTIFS", "AVERAGEIFS", "SUMIF", "COUNTIF"):
            return self.aggregate(low, node)
        if name in ("SUM", "MIN", "MAX", "AVERAGE") and len(args) == 1 and isinstance(args[0], Range):
            cols = self.same_row_range(low, args[0])
            if cols is None:
                raise low.fail(f"{name} over a range outside this row is not supported")
            roles = {self.kinds_of.get(c, "matrix" if c in self.grids["matrix_cols"] else None) for c in cols}
            if roles == {"summary"} or roles == {"summary", "literal"} or roles == {"literal"}:
                for c in cols:
                    self.deps.add(self.name_of[c])
                items = [["c", self.name_of[c]] for c in cols]
                if name == "SUM":
                    return ["f", "SUM"] + items if len(items) > 1 else items[0]
                if name in ("MIN", "MAX") and len(items) > 1:
                    return ["f", name] + items
                raise low.fail(f"{name} over {len(items)} column(s) is not supported")
            if roles == {"matrix"} and set(cols) == set(self.grids["matrix_cols"]):
                fn = {"SUM": "sum", "MIN": "min", "MAX": "max", "AVERAGE": "avg"}[name]
                self.sources.add(("frame", self.grids["keyed_id"]))
                self.by_cols.add(self.grids["keyed_row_key"])
                return ["agg", fn, self.grids["keyed_measure"], None]
            raise low.fail(f"{name} range spans columns of different kinds or only part of the period axis")
        if name == "SUM" and len(args) == 1 and isinstance(args[0], Function) and args[0].name == "INDEX":
            return self.row_sum_lookup(low, args[0])
        return None

    def row_sum_lookup(self, low, node):
        args = node.arguments
        if len(args) == 3 and isinstance(args[0], StructuredReference) and isinstance(args[1], Function) \
                and args[1].name == "MATCH" and isinstance(args[2], Number) and args[2].value == 0:
            kind = _struct(args[0], low.where)
            match = args[1].arguments
            if kind[0] == "range" and len(match) == 3 and isinstance(match[2], Number) and match[2].value == 0:
                table = self.tabs.get(args[0].table)
                if table is None:
                    raise low.fail(f"unknown table {args[0].table!r}")
                key_table, key_col = low.table_ref(match[1], "MATCH lookup column")
                if key_table is not table:
                    raise low.fail("MATCH column and INDEX range belong to different tables")
                names = [c.name for c in table.columns]
                if kind[1] not in names or kind[2] not in names:
                    raise low.fail(f"table {table.name} has no columns {kind[1]!r}..{kind[2]!r}")
                span = names[names.index(kind[1]):names.index(kind[2]) + 1]
                if any(table.col(n).kind != "literal" or table.col(n).type != "number" for n in span
                       if table.col(n).type):
                    raise low.fail(f"{table.name}[{kind[1]}:{kind[2]}] holds non-numeric columns")
                total_name = f"Total ({kind[1]} to {kind[2]})"
                self.reg.added_columns.append((table.name, total_name, span))
                key = low.lower(match[0])
                low.lookup_node(key, key_table, key_col, key_table, key_col, None, wildcards=True)
                return ["lk", table.name, key, key_col, total_name, None]
        raise low.fail("INDEX with a column range is only supported as "
                       "SUM(INDEX(tblX[[A]:[B]], MATCH(label, tblX[Key], 0), 0))")

    # --- SUMIFS / COUNTIFS / AVERAGEIFS
    def aggregate(self, low, node):
        name, args = node.name, node.arguments
        fn = {"SUMIFS": "sum", "SUMIF": "sum", "COUNTIFS": "count", "COUNTIF": "count",
              "AVERAGEIFS": "avg"}[name]
        if name in ("SUMIF", "COUNTIF"):
            if name == "SUMIF":
                if len(args) not in (2, 3):
                    raise low.fail("SUMIF argument count")
                value, pairs = (args[2] if len(args) == 3 else args[0]), [(args[0], args[1])]
            else:
                if len(args) != 2:
                    raise low.fail("COUNTIF argument count")
                value, pairs = None, [(args[0], args[1])]
        elif fn == "count":
            if len(args) < 2 or len(args) % 2:
                raise low.fail("COUNTIFS needs range/criterion pairs")
            value, pairs = None, list(zip(args[0::2], args[1::2]))
        else:
            if len(args) < 3 or len(args) % 2 == 0:
                raise low.fail(f"{name} needs a value range then range/criterion pairs")
            value, pairs = args[0], list(zip(args[1::2], args[2::2]))
        probe = [value] if value is not None else []
        probe += [r for r, _ in pairs]
        if all(isinstance(p, StructuredReference) for p in probe):
            return self.over_table(low, fn, value, pairs)
        if all(isinstance(p, (Range, ColumnRange)) for p in probe):
            return self.over_grid(low, fn, value, pairs)
        raise low.fail("aggregate ranges must all be table columns or all be grid columns")

    def over_table(self, low, fn, value, pairs):
        table = None
        value_col = None
        if value is not None:
            table, value_col = low.table_ref(value, "aggregate value range")
        terms, label_col = [], None
        for rng, crit in pairs:
            t, col = low.table_ref(rng, "aggregate criteria range")
            if table is not None and t is not table:
                raise low.fail("aggregate ranges belong to different tables")
            table = t
            result = _crit(low, crit, lambda ref: self.is_label(low, ref))
            if result == ("label",):
                if label_col is not None:
                    raise low.fail("more than one criterion is bound to the row label")
                label_col = col
            else:
                terms.append(["op", result[0], ["c", col], result[1]])
        if label_col is None:
            raise low.fail("aggregate is not filtered by the row label, so it is not a per-row summary")
        self.sources.add(("table", table.name))
        self.by_cols.add(label_col)
        return ["agg", fn, value_col, _and(terms)]

    # --- aggregates over another grid's columns (expansion frames)
    def grid_column(self, low, node):
        """(grid, role dict) for a whole-column / bounded range into an earlier grid."""
        if isinstance(node, ColumnRange):
            sheet = node.sheet or low.sheet
            if node.start_column != node.end_column:
                raise low.fail("multi-column ranges are not supported in aggregates")
            col = column_index_from_string(node.start_column)
            rows = None
        else:
            sheet = node.start.sheet or low.sheet
            r1, c1 = _rc(node.start)
            r2, c2 = _rc(node.end)
            if c1 != c2:
                raise low.fail("multi-column ranges are not supported in aggregates")
            col, rows = c1, (min(r1, r2), max(r1, r2))
        grid = self.grids["expansions"].get(sheet)
        if grid is None:
            raise low.fail(f"range on {sheet!r} is not a recognised expansion grid (ordinal rows x periods)")
        role = grid["cols"].get(col)
        if role is None:
            raise low.fail(f"column {get_column_letter(col)} of {sheet!r} is not part of the recognised grid")
        if rows is not None and rows != (grid["first_row"], grid["last_row"]):
            raise low.fail(f"range rows {rows} do not match the grid rows {grid['first_row']}..{grid['last_row']}")
        if rows is None:
            ws = self.reg.wf[sheet]
            for r in range(1, ws.max_row + 1):
                if grid["first_row"] <= r <= grid["last_row"] or r == grid["header_row"]:
                    continue
                if ws.cell(r, col).value is not None and not (
                        isinstance(ws.cell(r, col).value, str) and r < grid["header_row"]):
                    raise low.fail(f"whole-column range {sheet}!{get_column_letter(col)} also holds "
                                   f"{sheet}!{get_column_letter(col)}{r}, outside the grid")
        return grid, role

    def over_grid(self, low, fn, value, pairs):
        grid = None
        value_role = None
        if value is not None:
            grid, value_role = self.grid_column(low, value)
        label_col, filters, period_value = None, [], None
        for rng, crit in pairs:
            g, role = self.grid_column(low, rng)
            if grid is not None and g is not grid:
                raise low.fail("aggregate ranges belong to different grids")
            grid = g
            result = _crit(low, crit, lambda ref: self.is_label(low, ref))
            if role["role"] == "attr":
                if result == ("label",):
                    if label_col is not None:
                        raise low.fail("more than one criterion is bound to the row label")
                    label_col = role["name"]
                else:
                    filters.append(["op", result[0], ["c", role["name"]], result[1]])
            elif role["role"] == "period":
                if result == ("label",):
                    raise low.fail("a period column cannot be matched to the row label")
                if period_value is not None and period_value != role["value"]:
                    raise low.fail("criteria on two different period columns cannot be combined")
                period_value = role["value"]
                filters.append(["op", "=", ["c", grid["axis_name"]], ["d", role["value"]]])
                filters.append(["op", result[0], ["c", grid["measure"]], result[1]])
            else:
                raise low.fail("a row-total column cannot be used as a criteria range")
        if label_col is None:
            raise low.fail("aggregate is not filtered by the row label, so it is not a per-row summary")
        if value_role is not None:
            if value_role["role"] == "rowtotal":
                if period_value is not None:
                    raise low.fail("a row-total value range cannot be combined with a period-column criterion")
                column = grid["measure"]
            elif value_role["role"] == "period":
                if period_value is not None and period_value != value_role["value"]:
                    raise low.fail("value and criteria refer to different period columns")
                filters.append(["op", "=", ["c", grid["axis_name"]], ["d", value_role["value"]]])
                column = grid["measure"]
            else:
                raise low.fail("the value range must be a period or row-total column")
        else:
            if period_value is None:
                raise low.fail("a COUNTIFS over an expansion grid needs one period-column criterion "
                               "(otherwise it would count rows once per period)")
            column = None
        self.sources.add(("frame", grid["frame_id"]))
        self.by_cols.add(label_col)
        # duplicate period filters collapse (value and criteria on the same period column)
        unique = []
        for term in filters:
            if term not in unique:
                unique.append(term)
        return ["agg", fn, column, _and(unique)]


# ---- expansion grids (ordinal rows x periods)
class ExpansionCtx(Ctx):
    def __init__(self, block, base, own_col, attr_cols, reg):
        self.b, self.base, self.own_col, self.attr_cols, self.reg = block, base, own_col, attr_cols, reg
        self.row = None

    def ref(self, low, ref):
        sheet = ref.sheet or low.sheet
        if sheet == self.b.sheet and ref.address:
            row, col = _rc(ref)
            if row == self.b.hr and col == self.own_col:
                return ["c", "__axis__"]
        return low.param_ref(ref)

    def func(self, low, node):
        if node.name == "INDEX" and len(node.arguments) == 2 and isinstance(node.arguments[1], Reference) \
                and isinstance(node.arguments[0], StructuredReference):
            ref = node.arguments[1].value
            if ref.address:
                row, col = _rc(ref)
                sheet = ref.sheet or low.sheet
                if sheet == self.b.sheet and col == self.b.lc and row == self.row:
                    table, column = low.table_ref(node.arguments[0], "INDEX range")
                    if table is not self.base:
                        raise low.fail(f"INDEX reads {table.name}, but this grid expands {self.base.name}")
                    return ["c", column]
        return None


def _index_pattern(node, label_col, row):
    """If ``node`` is INDEX(tbl[col], $A<row>) return (table, col) else None."""
    if isinstance(node, Function) and node.name == "INDEX" and len(node.arguments) == 2 \
            and isinstance(node.arguments[0], StructuredReference) and isinstance(node.arguments[1], Reference):
        ref = node.arguments[1].value
        if ref.address and _rc(ref) == (row, label_col):
            kind = _struct(node.arguments[0], "")
            if kind[0] == "col":
                return node.arguments[0].table, kind[1]
    return None


def _parse(formula, sheet, where):
    try:
        return parse_expression(formula, sheet=sheet, workbook_mode=True)
    except ExpressionError as error:
        raise PlanDraftError(f"{where}: {error}") from error


def _contains_ref_to(node, pred):
    """True when any single-cell Reference in ``node`` satisfies ``pred(ref)``."""
    if isinstance(node, Reference):
        return bool(node.value.address) and pred(node.value)
    if isinstance(node, Range):
        return pred(node.start) or pred(node.end)
    if isinstance(node, (Group,)):
        return _contains_ref_to(node.expression, pred)
    if isinstance(node, (Unary, Percent)):
        return _contains_ref_to(node.operand, pred)
    if isinstance(node, Binary):
        return _contains_ref_to(node.left, pred) or _contains_ref_to(node.right, pred)
    if isinstance(node, Function):
        return any(_contains_ref_to(a, pred) for a in node.arguments)
    return False


def _walk_functions(node):
    if isinstance(node, Function):
        yield node
        for a in node.arguments:
            yield from _walk_functions(a)
    elif isinstance(node, Group):
        yield from _walk_functions(node.expression)
    elif isinstance(node, (Unary, Percent)):
        yield from _walk_functions(node.operand)
    elif isinstance(node, Binary):
        yield from _walk_functions(node.left)
        yield from _walk_functions(node.right)


class Drafter:
    def __init__(self, src, wf, wv, tabs, intent, tolerance):
        self.src, self.wf, self.wv, self.tabs = src, wf, wv, tabs
        self.intent, self.tol = intent, tolerance
        self.reg = Registry(src, wf, wv, tabs)
        self.frames, self.checks, self.targets = [], [], []
        self.expansions = {}          # sheet -> grid info (for aggregates over expansion frames)
        self.issues = []
        self.ledger_scenarios = []    # period ledgers' headline targets, for the scenario recompute

    # ---- checks
    def add_checks(self, frame, column, kind, cells):
        if cells:
            self.checks.append({"frame": frame, "column": column, "kind": kind,
                                "cells": [[a, sel, _plain(v)] for a, sel, v in cells]})

    # ---- table calculated columns
    def table_columns(self):
        for tab in self.tabs.values():
            for column in tab.columns:
                if column.kind != "calculated":
                    continue
                first = column.formulas[0]
                where = column.cells[0]
                ast = _parse(first, tab.ws.title, where)
                first_row = tab.data_rows[0]
                norm = _norm(ast, first_row, tab.min_col, tab.ws.title)
                for row, coord, formula in zip(tab.data_rows, column.cells, column.formulas):
                    other = _norm(_parse(formula, tab.ws.title, coord), row, tab.min_col, tab.ws.title)
                    if other != norm:
                        raise PlanDraftError(f"{coord}: calculated column {column.name!r} of {tab.name} is "
                                             f"not uniform ({formula!r} vs {first!r} in {column.cells[0]})")
                low = Lowerer(self.reg, tab.ws.title, where, TableCtx(tab, first_row))
                column.expr = low.lower(ast)
                column.type = self.type_from_cached(column.cells, column.cached, column.expr, where, column.name)

    def type_from_cached(self, cells, cached, expr, where, name):
        found = _infer_type(list(zip(cells, cached)), where, name)
        if found:
            return found
        return "text" if self.guess_text(expr) else "number"

    @staticmethod
    def guess_text(expr):
        if expr[0] == "f" and expr[1] == "IF":
            return any(b[0] == "s" for b in expr[3:5])
        return expr[0] == "s"

    def finalize_literal_types(self):
        notes = []
        for tab in self.tabs.values():
            for column in tab.columns:
                if column.kind != "literal":
                    continue
                where = f"{tab.ws.title}!{tab.name}"
                column.type = _infer_type(list(zip(column.cells, column.formulas)), where, column.name)
                if column.type is None:
                    column.type = _guess_empty_type(column.name)
                    notes.append(f"{tab.name}[{column.name}]")
        if notes:
            self.issues.append({
                "code": "TABLE_COLUMN_TYPES_GUESSED", "severity": "review",
                "route": "table-report-column-type-review",
                "message": "These table columns hold no values, so their types were guessed from the column "
                           "name; confirm them before seeding.",
                "evidence": notes})

    # ---- blocks
    def run(self):
        self.finalize_literal_types()
        self.table_columns()
        self.draft_ledgers()
        blocks = find_blocks(self.wf, self.reg.covered, self.tabs)
        stray = _formula_cells_outside(self.wf, self.reg.covered, blocks)
        if stray:
            raise PlanDraftError(f"{stray[0]}: formula cell outside any recognised grid "
                                 f"({len(stray)} such cell(s); first {', '.join(stray[:5])})")
        formula_blocks = []
        for block in blocks:
            has_formula = any(_is_formula(block.ws.cell(r, c).value)
                              for r in block.rows for c in block.cols)
            if has_formula:
                formula_blocks.append(block)
        # expansions and keyed matrices first: summaries may aggregate over them
        ordered = sorted(formula_blocks, key=lambda b: 0 if self.is_ordinal(b) else 1)
        for block in ordered:
            if self.is_ordinal(block):
                self.expansion_block(block)
        for block in ordered:
            if not self.is_ordinal(block):
                self.label_block(block)
        if not self.frames:
            raise NoTableReportShape("no formula-driven grid over table data was found "
                                     "(a table-report needs at least one summary, expansion or keyed grid)")

    def draft_ledgers(self):
        """Period ledgers (cash-flow style: line items down, period ordinals across) come first: their
        cells are then covered, so the grid finders below do not mistake them for summaries."""
        try:
            from . import table_report_ledger as ledger
        except ImportError:  # direct import when scripts/lib itself is on sys.path
            import table_report_ledger as ledger
        ledger.draft_ledgers(self)

    @staticmethod
    def is_ordinal(block):
        labels = [block.label(r) for r in block.rows]
        return all(isinstance(v, int) and not isinstance(v, bool) for v in labels)

    # ---- helpers on columns of a block
    def column_formulas(self, block, col):
        values = [block.ws.cell(r, col).value for r in block.rows]
        formulas = [v for v in values if _is_formula(v)]
        if formulas and len(formulas) != len(values):
            first = next(_addr(block.sheet, r, col) for r, v in zip(block.rows, values) if not _is_formula(v))
            raise PlanDraftError(f"{first}: column {get_column_letter(col)} mixes formulas and literal cells")
        return values, bool(formulas)

    def uniform(self, block, col, values):
        first_row = block.first_row
        base = _parse(values[0], block.sheet, _addr(block.sheet, first_row, col))
        norm = _norm(base, first_row, col, block.sheet)
        for r, v in zip(block.rows, values):
            other = _norm(_parse(v, block.sheet, _addr(block.sheet, r, col)), r, col, block.sheet)
            if other != norm:
                raise PlanDraftError(f"{_addr(block.sheet, r, col)}: formula differs from "
                                     f"{_addr(block.sheet, first_row, col)} ({v!r}); grid columns must be "
                                     "one formula copied down")
        return base, norm

    def cached(self, block, row, col):
        return self.wv[block.sheet].cell(row, col).value

    def col_type(self, block, col, expr):
        cells = [_addr(block.sheet, r, col) for r in block.rows]
        cached = [self.cached(block, r, col) for r in block.rows]
        return self.type_from_cached(cells, cached, expr, cells[0], str(block.header(col)))

    # ---- expansion (ordinal rows x periods)
    def expansion_block(self, block):
        labels = [block.label(r) for r in block.rows]
        if labels != list(range(1, len(labels) + 1)):
            raise PlanDraftError(f"{_addr(block.sheet, block.first_row, block.lc)}: ordinal labels must be "
                                 f"1..{len(labels)} in order")
        attr, matrix, rowtotal, bases = {}, [], None, set()
        for col in block.cols:
            values, is_formula = self.column_formulas(block, col)
            if not is_formula:
                raise PlanDraftError(f"{_addr(block.sheet, block.first_row, col)}: literal column in an ordinal "
                                     "expansion grid is not supported")
            first_ast = _parse(values[0], block.sheet, _addr(block.sheet, block.first_row, col))
            hit = _index_pattern(first_ast, block.lc, block.first_row)
            if hit:
                self.uniform(block, col, values)
                table = self.tabs.get(hit[0])
                if table is None or table.col(hit[1]) is None:
                    raise PlanDraftError(f"{_addr(block.sheet, block.first_row, col)}: unknown {hit[0]}[{hit[1]}]")
                attr[col] = hit[1]
                bases.add(hit[0])
                continue
            header = block.header(col)
            if isinstance(header, (datetime, date)):
                matrix.append(col)
                continue
            if isinstance(first_ast, Function) and first_ast.name == "SUM" and len(first_ast.arguments) == 1 \
                    and isinstance(first_ast.arguments[0], Range):
                rowtotal = col
                continue
            raise PlanDraftError(f"{_addr(block.sheet, block.hr, col)}: column {header!r} is not a table "
                                 "attribute (INDEX), a period column or a row total")
        if not matrix:
            raise PlanDraftError(f"{block.sheet}: ordinal grid has no date-headed period columns")
        if not bases:
            for col in matrix:
                for fn in _walk_functions(_parse(block.ws.cell(block.first_row, col).value, block.sheet, "")):
                    if fn.name == "INDEX":
                        hit = _index_pattern(fn, block.lc, block.first_row)
                        if hit:
                            bases.add(hit[0])
        if len(bases) != 1:
            raise PlanDraftError(f"{block.sheet}: ordinal grid must expand exactly one table "
                                 f"(found {sorted(bases) or 'none'})")
        base = self.tabs[next(iter(bases))]
        row_key = self.expansion_row_key(block, base, attr)
        if base.n != len(labels):
            raise PlanDraftError(f"{block.sheet}: ordinal rows 1..{len(labels)} do not cover all "
                                 f"{base.n} rows of {base.name}; a full expansion would add rows Excel omits")
        if matrix != list(range(matrix[0], matrix[-1] + 1)):
            raise PlanDraftError(f"{block.sheet}: period columns are not contiguous")
        periods = [_date(block.header(c), _addr(block.sheet, block.hr, c)) for c in matrix]
        if periods != sorted(set(periods)):
            raise PlanDraftError(f"{block.sheet}: period headers must be unique and ascending")
        # lower the measure from the first period column, require one copied formula everywhere
        norms = set()
        for col in matrix:
            values, _ = self.column_formulas(block, col)
            _, norm = self.uniform(block, col, values)
            norms.add(norm)
        if len(norms) != 1:
            raise PlanDraftError(f"{_addr(block.sheet, block.first_row, matrix[1] if len(matrix) > 1 else matrix[0])}: "
                                 "period columns do not share one formula")
        first_col = matrix[0]
        ctx = ExpansionCtx(block, base, first_col, attr, self.reg)
        ctx.row = block.first_row
        where = _addr(block.sheet, block.first_row, first_col)
        low = Lowerer(self.reg, block.sheet, where, ctx)
        measure_ir = low.lower(_parse(block.ws.cell(block.first_row, first_col).value, block.sheet, where))
        axis_name = "Period"
        for existing in [c.name for c in base.columns]:
            if existing == axis_name:
                axis_name = "Period (grid)"
        measure_ir = self.rename_axis(measure_ir, axis_name)
        measure_name = self.measure_name(block)
        used = {c.name for c in base.columns} | {axis_name}
        if measure_name in used:
            measure_name = f"{measure_name} (calc)"
        if block.sheet in self.expansions:
            raise PlanDraftError(f"{block.sheet}: more than one ordinal expansion grid on a sheet is not supported")
        title = self.unique_title(_title(block.sheet))
        frame_id = f"x_{trm.slug(title)}"
        measure_type = "number"
        fmt = _fmt_kind(block.ws.cell(block.first_row, first_col))
        self.frames.append({
            "id": frame_id, "kind": "expansion", "title": title, "sheet": block.sheet,
            "base": base.name, "attrs": ([row_key] if row_key else []) + [attr[c] for c in sorted(attr)],
            "measure": measure_name,
            "axis": {"name": axis_name, "type": "datetime", "values": [p.isoformat() for p in periods]},
            "columns": [{"name": measure_name, "type": measure_type, "expr": measure_ir,
                         **({"format": fmt} if fmt else {})}]})
        grid_cols = {}
        for col, base_col in attr.items():
            grid_cols[col] = {"role": "attr", "name": base_col}
        for col, period in zip(matrix, periods):
            grid_cols[col] = {"role": "period", "value": period.isoformat()}
        if rowtotal:
            grid_cols[rowtotal] = {"role": "rowtotal"}
        self.expansions[block.sheet] = {"frame_id": frame_id, "cols": grid_cols, "measure": measure_name,
                                        "axis_name": axis_name, "first_row": block.first_row,
                                        "last_row": block.last_row, "header_row": block.hr}
        # checks
        for col, base_col in attr.items():
            self.add_checks(base.name, base_col, "table",
                            [[_addr(block.sheet, r, col), [i], self.cached(block, r, col)]
                             for i, r in enumerate(block.rows)])
        cells = []
        for col, period in zip(matrix, periods):
            for i, r in enumerate(block.rows):
                cells.append([_addr(block.sheet, r, col), [i, period.isoformat()], self.cached(block, r, col)])
        self.add_checks(frame_id, measure_name, "expansion", cells)
        if rowtotal:
            first_ast = _parse(block.ws.cell(block.first_row, rowtotal).value, block.sheet, "")
            span = first_ast.arguments[0]
            self.uniform(block, rowtotal, [block.ws.cell(r, rowtotal).value for r in block.rows])
            r1, c1 = _rc(span.start)
            r2, c2 = _rc(span.end)
            if (r1, r2) != (block.first_row, block.first_row) or (c1, c2) != (matrix[0], matrix[-1]):
                raise PlanDraftError(f"{_addr(block.sheet, block.first_row, rowtotal)}: row total must be "
                                     "SUM over exactly the period columns of its row")
            self.add_checks(frame_id, measure_name, "rowtotal",
                            [[_addr(block.sheet, r, rowtotal), [i], self.cached(block, r, rowtotal)]
                             for i, r in enumerate(block.rows)])

    def expansion_row_key(self, block, base, attr):
        """The pivot over an expansion groups by the attribute columns the grid shows, so those must
        identify ONE base row each, or distinct rows collapse into one (a roster grid showing only
        ``Dept`` becomes one row per department). Returns the base column to lead the row dimensions
        with when the shown attributes are not unique, else ``None``; raises when no column can."""
        shown = [attr[c] for c in sorted(attr)]
        keys = [tuple(col_values(base, name)[i] for name in shown) for i in range(base.n)]
        if shown and len(set(keys)) == len(keys):
            return None
        for column in base.columns:
            if column.kind != "literal" or column.name in shown:
                continue
            values = col_values(base, column.name)
            if all(v is not None for v in values) and len(set(values)) == len(values):
                self.issues.append({
                    "code": "EXPANSION_ROW_KEY_ADDED", "severity": "review",
                    "route": "expansion-row-key-review",
                    "message": f"{block.sheet}: the grid shows {shown or 'no attribute columns'}, which do not "
                               f"identify one {base.name} row, so the Sigma pivot also groups by "
                               f"{base.name}[{column.name}] (a unique column) to keep one pivot row per "
                               f"{base.name} row, as in Excel; the pivot therefore shows that extra column.",
                    "evidence": [f"{block.sheet}!{get_column_letter(block.lc)}{block.first_row}"]})
                return column.name
        where = _addr(block.sheet, block.first_row, min(attr) if attr else block.lc)
        raise PlanDraftError(f"{where}: the columns this grid shows from {base.name} ({shown or 'none'}) do not "
                             f"identify one row each and {base.name} has no unique, non-blank column to add, so "
                             "a Sigma pivot would merge distinct rows; add a unique key column to the table")

    @staticmethod
    def rename_axis(node, axis_name):
        if isinstance(node, list):
            if node and node[0] == "c" and node[1] == "__axis__":
                return ["c", axis_name]
            return [Drafter.rename_axis(n, axis_name) if isinstance(n, list) else n for n in node]
        return node

    @staticmethod
    def measure_name(block):
        title = block.title
        if title and len(title) <= 42 and not re.search(r"[\[\]/]", title):
            return title
        return f"{block.sheet} Value"

    # ---- label blocks (summary / keyed)
    def label_block(self, block):
        total = None
        last_label = block.label(block.last_row)
        if len(block.rows) > 2 and isinstance(last_label, str) and re.match(r"^\s*(grand\s+)?total\b", last_label, re.I):
            total = {"row": block.last_row, "label": last_label.strip()}
            block.rows = block.rows[:-1]   # the total row is verified, not treated as a data row
            block.last_row -= 1
        labels = [block.label(r) for r in block.rows]
        if len(set(labels)) != len(labels):
            raise PlanDraftError(f"{_addr(block.sheet, block.first_row, block.lc)}: duplicate row labels in "
                                 f"{block.sheet}")
        for r, v in zip(block.rows, labels):
            if _is_formula(v) or v is None:
                raise PlanDraftError(f"{_addr(block.sheet, r, block.lc)}: row labels must be literal values")
            _wildcard_free(v, _addr(block.sheet, r, block.lc), "row label")
        header_names = {}
        matrix_cols, summary_cols, literal_cols = [], [], []
        for col in block.cols:
            values, is_formula = self.column_formulas(block, col)
            header = block.header(col)
            if not is_formula:
                literal_cols.append(col)
                continue
            if isinstance(header, (datetime, date)):
                first_ast = _parse(values[0], block.sheet, _addr(block.sheet, block.first_row, col))
                hdr = lambda ref, hr=block.hr, c=col: bool(ref.address) and _rc(ref) == (hr, c)  # noqa: E731
                prev = lambda ref, r=block.first_row, c=col: bool(ref.address) \
                    and (ref.sheet or block.sheet) == block.sheet and _rc(ref) == (r, c - 1)  # noqa: E731
                if _contains_ref_to(first_ast, hdr) or (
                        matrix_cols and matrix_cols[-1] == col - 1 and _contains_ref_to(first_ast, prev)):
                    matrix_cols.append(col)
                    continue
            summary_cols.append(col)
        for col in summary_cols + literal_cols:
            name = _clean_name(block.header(col), _addr(block.sheet, block.hr, col))
            if isinstance(block.header(col), (datetime, date)):
                name = block.header(col).strftime("%Y-%m-%d")
            if name in header_names.values():
                raise PlanDraftError(f"{_addr(block.sheet, block.hr, col)}: duplicate column header {name!r}")
            header_names[col] = name
        grids = {"matrix_cols": matrix_cols, "expansions": self.expansions}
        keyed_id = None
        if total and matrix_cols:
            raise PlanDraftError(f"{_addr(block.sheet, total['row'], block.lc)}: a total row under period columns "
                                 "is not supported")
        if matrix_cols:
            if matrix_cols != list(range(matrix_cols[0], matrix_cols[-1] + 1)):
                raise PlanDraftError(f"{block.sheet}: period columns are not contiguous")
            keyed = self.keyed_block(block, labels, matrix_cols)
            keyed_id = keyed["id"]
            grids.update(keyed_id=keyed_id, keyed_row_key=keyed["rowKey"], keyed_measure=keyed["measure"])
        if not summary_cols and not literal_cols:
            return
        self.summary_block(block, labels, summary_cols, literal_cols, header_names, grids, matrix_cols, total)

    # ---- keyed matrix
    def keyed_block(self, block, labels, matrix_cols):
        first_col = matrix_cols[0]
        periods = [_date(block.header(c), _addr(block.sheet, block.hr, c)) for c in matrix_cols]
        if periods != sorted(set(periods)):
            raise PlanDraftError(f"{block.sheet}: period headers must be unique and ascending")
        irs, norms = [], []
        state = {"table": None, "row_key": None, "period_key": None}
        for position, col in enumerate(matrix_cols[:2]):
            values, _ = self.column_formulas(block, col)
            base_ast, norm = self.uniform(block, col, values)
            norms.append(norm)
            where = _addr(block.sheet, block.first_row, col)
            ctx = KeyedCtx(block, col, col > first_col, state, self.tabs, self.reg)
            ctx.row = block.first_row
            irs.append(Lowerer(self.reg, block.sheet, where, ctx).lower(base_ast))
        for col in matrix_cols[2:]:
            values, _ = self.column_formulas(block, col)
            _, norm = self.uniform(block, col, values)
            if len(norms) > 1 and norm != norms[1]:
                raise PlanDraftError(f"{_addr(block.sheet, block.first_row, col)}: period column formula differs "
                                     f"from {_addr(block.sheet, block.first_row, matrix_cols[1])}")
        table = self.tabs[state["table"]]
        row_key, period_key = state["row_key"], state["period_key"]
        if row_key is None or period_key is None:
            raise PlanDraftError(f"{_addr(block.sheet, block.first_row, first_col)}: a keyed grid needs "
                                 "SUMIFS(tbl[V], tbl[RowKey], $label, tbl[PeriodKey], header)")
        running = len(matrix_cols) > 1 and any(n[0] == "prev" for n in trm.walk(irs[1]))
        measure = self.measure_name(block)
        if measure in {c.name for c in table.columns}:
            measure = f"{measure} (calc)"
        columns = []
        if running:
            opening, terms = self.peel(irs[0])
            tail_leaf, tail_terms = self.peel(irs[1])
            if tail_leaf != ["prev"]:
                raise PlanDraftError(f"{_addr(block.sheet, block.first_row, matrix_cols[1])}: a running balance "
                                     "must start from the previous period's cell (prev +/- terms)")
            if terms != tail_terms:
                raise PlanDraftError(f"{_addr(block.sheet, block.first_row, first_col)}: the first period's "
                                     "formula differs from the running formula beyond its opening term")
            if not terms:
                first_value = ["f", "IF", ["op", "=", ["period"], ["d", periods[0].isoformat()]],
                               opening, ["n", 0]]
                columns.append({"name": measure, "type": "number",
                                "expr": ["cumsum", first_value, row_key, period_key]})
            else:
                delta_name = "Net Change"
                if delta_name in {c.name for c in table.columns} | {measure}:
                    delta_name = f"{measure} Change"
                columns.append({"name": delta_name, "type": "number", "expr": self.rebuild(terms)})
                running_expr = ["op", "+", opening, ["cumsum", ["c", delta_name], row_key, period_key]]
                columns.append({"name": measure, "type": "number", "expr": running_expr})
                self.issues.append({
                    "code": "RUNNING_BALANCE_PARTITION_REVIEW", "severity": "review",
                    "route": "cumulative-sum-live-validation",
                    "message": f"{block.sheet}: the running balance is a CumulativeSum partitioned by {row_key} "
                               f"and ordered by {period_key}; confirm the live query matches before relying on it.",
                    "evidence": [f"{block.sheet}!{get_column_letter(matrix_cols[0])}{block.first_row}"]})
        else:
            if len(matrix_cols) > 1 and norms[0] != norms[1]:
                raise PlanDraftError(f"{_addr(block.sheet, block.first_row, first_col)}: first period formula "
                                     "differs from the others")
            columns.append({"name": measure, "type": "number", "expr": irs[0]})
        fmt = _fmt_kind(block.ws.cell(block.first_row, first_col))
        if fmt:
            columns[-1]["format"] = fmt
        # the grid must be exactly the table's (row key, period key) pairs
        pairs = {}
        for i, row in enumerate(table.data_rows):
            key = (col_values(table, row_key)[i], _date(col_values(table, period_key)[i], f"{table.name}"))
            if key in pairs:
                raise PlanDraftError(f"{table.name}: more than one row for {row_key}={key[0]!r}, "
                                     f"{period_key}={key[1]}; the grid sums them, a keyed frame cannot")
            pairs[key] = i
        wanted = {(lab, p) for lab in labels for p in periods}
        missing = sorted(wanted - set(pairs), key=str)
        extra = sorted(set(pairs) - wanted, key=str)
        if missing or extra:
            sample = (missing or extra)[0]
            raise PlanDraftError(
                f"{block.sheet}: the grid and {table.name} do not pair one-to-one on ({row_key}, {period_key}): "
                f"{len(missing)} grid cell(s) have no table row (e.g. {sample}) and {len(extra)} table row(s) "
                "fall outside the grid; this builder has no spine for sparse grids")
        title = self.unique_title(_title(block.sheet))
        frame_id = f"k_{trm.slug(title)}"
        keyed = {"id": frame_id, "kind": "keyed", "title": title, "sheet": block.sheet,
                 "table": table.name, "rowKey": row_key, "periodKey": period_key, "measure": measure,
                 "periods": [p.isoformat() for p in periods], "columns": columns}
        self.frames.append(keyed)
        cells = []
        for col, period in zip(matrix_cols, periods):
            for r, lab in zip(block.rows, labels):
                cells.append([_addr(block.sheet, r, col), [pairs[(lab, period)]], self.cached(block, r, col)])
        self.add_checks(frame_id, measure, "keyed", cells)
        return keyed

    @staticmethod
    def peel(node):
        terms = []
        while node[0] == "op" and node[1] in ("+", "-"):
            terms.append((node[1], node[3]))
            node = node[2]
        terms.reverse()
        return node, [list(t) for t in terms]

    @staticmethod
    def rebuild(terms):
        sign, first = terms[0]
        out = first if sign == "+" else ["neg", first]
        for sign, term in terms[1:]:
            out = ["op", sign, out, term]
        return out

    # ---- summary
    def summary_block(self, block, labels, summary_cols, literal_cols, header_names, grids, matrix_cols,
                      total=None):
        cols = sorted(summary_cols + literal_cols)
        kinds = {c: ("literal" if c in literal_cols else "summary") for c in cols}
        sources, by_cols, columns, deps = set(), set(), {}, {}
        for col in cols:
            if kinds[col] == "literal":
                continue
            values, _ = self.column_formulas(block, col)
            ast, _ = self.uniform(block, col, values)
            where = _addr(block.sheet, block.first_row, col)
            ctx = SummaryCtx(block, header_names, kinds, grids, self.tabs, self.reg, block.first_row)
            expr = Lowerer(self.reg, block.sheet, where, ctx).lower(ast)
            sources |= ctx.sources
            by_cols |= ctx.by_cols
            columns[col] = expr
            deps[col] = {c for c in cols if header_names[c] in ctx.deps}
        if not sources:
            raise PlanDraftError(f"{block.sheet}: summary grid at row {block.hr} has no aggregate over a table "
                                 "or grid, so its rows cannot be derived from data")
        if len(sources) != 1:
            raise PlanDraftError(f"{block.sheet}: summary columns aggregate more than one source "
                                 f"({sorted(sources)}); one grid must read one table or frame")
        if len(by_cols) != 1:
            raise PlanDraftError(f"{block.sheet}: summary columns group by different columns {sorted(by_cols)}")
        source_kind, source_name = next(iter(sources))
        by = next(iter(by_cols))
        # row labels must exist in the source (a grouped table cannot show a row with no data)
        key_values = self.source_key_values(source_kind, source_name, by)
        absent = [lab for lab in labels if lab not in key_values]
        if absent:
            raise PlanDraftError(f"{_addr(block.sheet, block.first_row + labels.index(absent[0]), block.lc)}: "
                                 f"row label {absent[0]!r} has no rows in {source_name}[{by}]; a grouped table "
                                 "would not show that row")
        extra = sorted((k for k in key_values if k not in labels and k is not None), key=str)
        if extra:
            self.issues.append({
                "code": "EXTRA_GROUP_KEYS", "severity": "review", "route": "grouped-rows-from-data-review",
                "message": f"{block.sheet}: {source_name}[{by}] has values the grid does not list; the Sigma "
                           "grouped table will show them as additional rows.",
                "evidence": [str(k) for k in extra[:10]]})
        # literal columns become an editable synthetic table keyed by the label
        title = _title(block.title or f"{block.sheet} summary")
        if source_kind == "frame" and any(f["id"] == source_name and f["kind"] == "keyed"
                                          for f in self.frames):
            title = f"{_title(block.sheet)} Summary"
        title = self.unique_title(title)
        frame_id = f"s_{trm.slug(title)}"
        frame_cols = []
        if literal_cols:
            lit_name = self.unique_title(f"{title} Inputs")
            lit_columns = [{"name": by, "type": self.key_type(labels)}]
            lit_rows = [[lab] for lab in labels]
            for col in literal_cols:
                values = [block.ws.cell(r, col).value for r in block.rows]
                kind = _infer_type([(_addr(block.sheet, r, col), v) for r, v in zip(block.rows, values)],
                                   block.sheet, header_names[col]) or "number"
                lit_columns.append({"name": header_names[col], "type": kind,
                                    **({"format": _fmt_kind(block.ws.cell(block.first_row, col))}
                                       if _fmt_kind(block.ws.cell(block.first_row, col)) else {})})
                for row_values, v in zip(lit_rows, values):
                    row_values.append(trm.iso(v) if isinstance(v, (date, datetime)) else v)
            self.reg.synthetic[lit_name] = {
                "name": lit_name, "sheet": block.sheet, "ref": f"{get_column_letter(block.lc)}{block.first_row}:"
                f"{get_column_letter(block.last_col)}{block.last_row}",
                "editable": True, "synthetic": True, "columns": lit_columns, "rows": lit_rows,
                "note": f"Typed values from {block.sheet} (not a formula) - editable inputs keyed by {by}."}
        for col in self.order_columns(cols, deps):
            name = header_names[col]
            fmt = _fmt_kind(block.ws.cell(block.first_row, col))
            if kinds[col] == "literal":
                expr = ["lk", lit_name, ["key"], by, name, None]
                ctype = next(c["type"] for c in lit_columns if c["name"] == name)
            else:
                expr = columns[col]
                ctype = self.col_type(block, col, expr)
            entry = {"name": name, "type": ctype, "expr": expr}
            if fmt:
                entry["format"] = fmt
            frame_cols.append(entry)
        frame = {"id": frame_id, "kind": "summary", "title": title, "sheet": block.sheet,
                 "source": {"kind": source_kind, "name": source_name}, "by": by,
                 "labels": [trm.iso(v) if isinstance(v, (date, datetime)) else v for v in labels],
                 "columns": frame_cols}
        self.frames.append(frame)
        for col in cols:
            name = header_names[col]
            self.add_checks(frame_id, name, "summary",
                            [[_addr(block.sheet, r, col), [lab], self.cached(block, r, col)]
                             for r, lab in zip(block.rows, labels)])
        if total:
            self.total_row(block, frame_id, title, cols, header_names, total)
        # headline parity targets: every cell of the summary (row x column)
        for col in cols:
            name = header_names[col]
            for r, lab in zip(block.rows, labels):
                value = self.cached(block, r, col)
                self.targets.append({"label": f"{title} / {lab} / {name}", "frame": frame_id, "column": name,
                                     "kind": "summary", "sel": [lab],
                                     "value": _plain(value) if value is not None else ""})

    def total_row(self, block, frame_id, title, cols, header_names, total):
        """A ``Total`` row of ``SUM(<column above>)`` cells: verified against the frame, not emitted."""
        for col in block.cols:
            value = block.ws.cell(total["row"], col).value
            if value is None:
                continue
            where = _addr(block.sheet, total["row"], col)
            ast = _parse(value, block.sheet, where) if _is_formula(value) else None
            ok = (col in cols and isinstance(ast, Function) and ast.name == "SUM" and len(ast.arguments) == 1
                  and isinstance(ast.arguments[0], Range))
            if ok:
                start, end = _rc(ast.arguments[0].start), _rc(ast.arguments[0].end)
                ok = start == (block.first_row, col) and end == (block.last_row, col)
            if not ok:
                raise PlanDraftError(f"{where}: total row cell must be SUM of the grid rows above it "
                                     f"({get_column_letter(col)}{block.first_row}:{get_column_letter(col)}"
                                     f"{block.last_row}), got {value!r}")
            self.add_checks(frame_id, header_names[col], "summary-total",
                            [[where, [], self.wv[block.sheet].cell(total["row"], col).value]])
        self.issues.append({
            "code": "TOTAL_ROW_VERIFIED_ONLY", "severity": "review", "route": "grouped-table-totals-review",
            "message": f"{block.sheet}: the {total['label']!r} row was verified as the sum of the grid rows "
                       "but is not emitted as a row; show totals on the grouped table if it is needed.",
            "evidence": [f"{block.sheet}!{get_column_letter(block.lc)}{total['row']}"]})

    @staticmethod
    def key_type(labels):
        return "number" if all(_is_number(v) for v in labels) else "text"

    def unique_title(self, title):
        taken = set(self.tabs) | {f["title"] for f in self.frames} | set(self.reg.synthetic) | {trm.PARAM_TABLE}
        candidate, n = title, 2
        while candidate in taken:
            candidate, n = f"{title} {n}", n + 1
        return candidate

    @staticmethod
    def order_columns(cols, deps):
        ordered, done = [], set()
        remaining = list(cols)
        while remaining:
            progressed = False
            for col in list(remaining):
                if deps.get(col, set()) <= done:
                    ordered.append(col)
                    done.add(col)
                    remaining.remove(col)
                    progressed = True
            if not progressed:
                raise PlanDraftError("summary columns reference each other in a cycle")
        return ordered

    def source_key_values(self, kind, name, by):
        if kind == "table":
            return {v for v in col_values(self.tabs[name], by) if v is not None}
        frame = next(f for f in self.frames if f["id"] == name)
        base = frame["base"] if frame["kind"] == "expansion" else frame["table"]
        return {v for v in col_values(self.tabs[base], by) if v is not None}


class KeyedCtx(Ctx):
    """Lowers one keyed-matrix cell: SUMIFS picks, label/header references, previous period."""

    def __init__(self, block, own_col, allow_prev, state, tabs, reg):
        self.b, self.own_col, self.allow_prev = block, own_col, allow_prev
        self.state, self.tabs, self.reg = state, tabs, reg
        self.row = None

    def _is(self, low, ref, row, col):
        return bool(ref.address) and (ref.sheet or low.sheet) == self.b.sheet and _rc(ref) == (row, col)

    def ref(self, low, ref):
        if self._is(low, ref, self.row, self.b.lc):
            return ["key"]
        if self._is(low, ref, self.b.hr, self.own_col):
            return ["period"]
        if self._is(low, ref, self.row, self.own_col - 1):
            if not self.allow_prev:
                raise low.fail("the first period column cannot reference a previous period")
            return ["prev"]
        if (ref.sheet or low.sheet) == self.b.sheet and ref.address:
            row, col = _rc(ref)
            if self.b.first_row <= row <= self.b.last_row or row == self.b.hr:
                raise low.fail(f"{ref.address} is a grid cell this recogniser cannot map (only $label, the "
                               "column header, and the previous period are supported)")
        return low.param_ref(ref)

    def func(self, low, node):
        if node.name != "SUMIFS":
            return None
        args = node.arguments
        if len(args) != 5:
            raise low.fail("keyed grids support SUMIFS(tbl[V], tbl[RowKey], $label, tbl[PeriodKey], header)")
        table, value_col = low.table_ref(args[0], "SUMIFS value range")
        row_key = period_key = None
        for rng, crit in ((args[1], args[2]), (args[3], args[4])):
            t, col = low.table_ref(rng, "SUMIFS criteria range")
            if t is not table:
                raise low.fail("SUMIFS ranges belong to different tables")
            if not isinstance(crit, Reference):
                raise low.fail("keyed SUMIFS criteria must be the row label and the column header")
            if self._is(low, crit.value, self.row, self.b.lc):
                row_key = col
            elif self._is(low, crit.value, self.b.hr, self.own_col):
                period_key = col
            else:
                raise low.fail("keyed SUMIFS criteria must be the row label ($A#) and the column header (B$#)")
        if row_key is None or period_key is None:
            raise low.fail("keyed SUMIFS needs one criterion on the row label and one on the column header")
        state = self.state
        for key, value in (("table", table.name), ("row_key", row_key), ("period_key", period_key)):
            if state[key] is None:
                state[key] = value
            elif state[key] != value:
                raise low.fail(f"SUMIFS reads {value!r} but earlier cells read {state[key]!r}")
        return ["c", value_col]


# ---------------------------------------------------------------- charts
def _range_cells(ref):
    """Addresses of a chart range such as ``'Forecast'!$B$4:$B$7`` -> ("Forecast", ["B4", ...])."""
    from openpyxl.utils import range_boundaries
    found = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^!]+))!(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)", ref or "")
    if not found:
        return None
    sheet = (found.group(1) or found.group(2)).replace("''", "'")
    c1, r1, c2, r2 = range_boundaries(found.group(3).replace("$", ""))
    return sheet, [f"{get_column_letter(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def list_charts(wf, checks=()):
    """Source charts are not emitted; say which verified frame column each series reads."""
    owner = {}
    for group in checks:
        for cell in group["cells"]:
            owner[cell[0]] = (group["frame"], group["column"])
    out = []
    for ws in wf.worksheets:
        for chart in getattr(ws, "_charts", []):
            title = None
            try:
                title = "".join(run.t for p in chart.title.tx.rich.p for run in (p.r or [])).strip() or None
            except AttributeError:
                pass
            kind = type(chart).__name__ + (" stacked" if getattr(chart, "grouping", None) == "stacked" else "")
            reads = []
            for series in chart.series:
                ref = series.val.numRef.f if series.val is not None and series.val.numRef is not None else None
                cells = _range_cells(ref)
                if not cells:
                    continue
                frames = {owner.get(f"{cells[0]}!{c}") for c in cells[1]}
                if len(frames) == 1 and None not in frames:
                    frame, column = next(iter(frames))
                    reads.append(f"{frame}.{column}")
            detail = f"; series read verified {', '.join(dict.fromkeys(reads))}" if reads else ""
            out.append(f"{title or type(chart).__name__} ({kind}) on {ws.title!r}{detail}: not emitted - "
                       "rebuild from the presentation plan (refs/presentation-migration.md)")
    return out


# ---------------------------------------------------------------- text-case hazards
def _frame_tables(frame, frames):
    """Names of the plan tables whose columns a frame's expressions read."""
    kind = frame["kind"]
    if kind == "expansion":
        return [frame["base"]]
    if kind == "keyed":
        return [frame["table"]]
    if kind == "ledger":
        return [frame["spine"]]
    source = frame["source"]
    if source["kind"] == "table":
        return [source["name"]]
    parent = next((f for f in frames if f["id"] == source["name"]), None)
    return _frame_tables(parent, frames) if parent else []


def _text_literal_compares(expr):
    """``(column, literal)`` for every comparison of a column with a non-empty text literal."""
    for node in trm.walk(expr):
        if node[0] == "op" and node[1] in trm.COMPARE:
            for column, literal in ((node[2], node[3]), (node[3], node[2])):
                if column[0] == "c" and literal[0] == "s" and literal[1] != "":
                    yield column[1], literal[1]


def reject_case_hazards(plan, tabs):
    """Excel compares text case-insensitively everywhere (``=``, SUMIFS criteria, MATCH, grouping into
    one row per label); Sigma compares exactly. Rather than guess, refuse a workbook whose data would
    behave differently: case variants of one value in a column the report groups by, or a text literal
    that the data holds in another case. Raises ``PlanDraftError`` naming the cell."""
    def text_column(table, name):
        tab = tabs.get(table)
        column = tab.col(name) if tab is not None else None
        if column is None or column.kind != "literal" or column.type != "text":
            return None
        return column

    def collisions(table, name, why):
        column = text_column(table, name)
        if column is None:
            return
        first = {}
        for coord, value in zip(column.cells, column.formulas):
            if isinstance(value, str):
                other = first.setdefault(value.lower(), (coord, value))
                if other[1] != value:
                    raise PlanDraftError(f"{coord}: {table}[{name}] holds {value!r} and {other[1]!r} "
                                         f"({other[0]}), the same text in different case; {why}")

    def against_literal(table, name, literal, where):
        column = text_column(table, name)
        if column is None:
            return
        for coord, value in zip(column.cells, column.formulas):
            if isinstance(value, str) and value != literal and value.lower() == literal.lower():
                raise PlanDraftError(f"{coord}: {table}[{name}] holds {value!r} but {where} compares it with "
                                     f"{literal!r}; Excel matches text ignoring case, Sigma exactly")

    for table in plan["tables"]:
        for column in table["columns"]:
            if "expr" in column:
                for name, literal in _text_literal_compares(column["expr"]):
                    against_literal(table["name"], name, literal, f"{table['name']}[{column['name']}]")
                for node in trm.walk(column["expr"]):
                    if node[0] == "lk":
                        collisions(node[1], node[3], "Sigma's Lookup would not match Excel's case-insensitive match")
    for frame in plan["frames"]:
        for name in _frame_tables(frame, plan["frames"]):
            for column in frame["columns"]:
                for col, literal in _text_literal_compares(column["expr"]):
                    against_literal(name, col, literal, f"{frame['title']}[{column['name']}]")
                for node in trm.walk(column["expr"]):
                    if node[0] == "lk":
                        collisions(node[1], node[3], "Sigma's Lookup would not match Excel's case-insensitive match")
            if frame["kind"] == "summary":
                collisions(name, frame["by"], "Excel merges them into one row, a Sigma grouped table keeps two")
            if frame["kind"] == "keyed":
                collisions(name, frame["rowKey"], "Excel merges them into one row, a Sigma grouped table keeps two")


# ---------------------------------------------------------------- entry point
def draft_table_report(source, *, intent, connection, folder, name=None, currency_format="$,.0f",
                       pct_format=".1%", tolerance=None):
    """Return ``(plan, summary)``; raise ``PlanDraftError`` on any shape or value it cannot prove."""
    if intent not in INTENTS:
        raise PlanDraftError(f"intent must be one of {', '.join(INTENTS)}")
    src = Path(source).expanduser().resolve()
    if not src.is_file():
        raise PlanDraftError(f"source workbook not found: {src}")
    wf = load_workbook(src)
    wv = load_workbook(src, data_only=True)
    tabs = load_tables(wf, wv)
    if not tabs:
        raise NoTableReportShape("no formal Excel Tables found; a table-report is built on formal tables")
    drafter = Drafter(src, wf, wv, tabs, intent, tolerance)
    drafter.run()
    reg = drafter.reg

    # tables -> plan dicts
    for table_name, column_name, span in reg.added_columns:
        tab = tabs[table_name]
        if tab.col(column_name) is None:
            col = Col(column_name, "calculated", [], [], [], None)
            col.expr = ["f", "SUM"] + [["c", n] for n in span]
            col.type = "number"
            col.synthetic = True
            tab.columns.append(col)
    editable_default = intent != "read-only"
    referenced = set()
    for frame in drafter.frames:
        if frame["kind"] == "expansion":
            referenced.add(frame["base"])
        if frame["kind"] == "keyed":
            referenced.add(frame["table"])
        if frame["kind"] == "summary" and frame["source"]["kind"] == "table":
            referenced.add(frame["source"]["name"])
        for column in frame["columns"]:
            for node in trm.walk(column["expr"]):
                if node[0] == "lk":
                    referenced.add(node[1])
    for tab in tabs.values():
        for column in tab.columns:
            if column.kind == "calculated" and hasattr(column, "expr"):
                for node in trm.walk(column.expr):
                    if node[0] == "lk":
                        referenced.add(node[1])
    plan_tables = []
    for tab in tabs.values():
        columns, rows = [], []
        literal = [c for c in tab.columns if c.kind == "literal"]
        for column in tab.columns:
            entry = {"name": column.name, "type": column.type}
            if column.fmt:
                entry["format"] = column.fmt
            if column.kind == "calculated":
                entry["expr"] = column.expr
                if getattr(column, "synthetic", False):
                    entry["synthetic"] = True
            columns.append(entry)
        for i in range(tab.n):
            row = []
            for column in literal:
                value = column.formulas[i]
                if column.type == "datetime" and value is not None:
                    value = _date(value, column.cells[i]).isoformat()
                row.append(value)
            rows.append(row)
        plan_tables.append({
            "name": tab.name, "sheet": tab.ws.title,
            "ref": tab.ref,
            "editable": bool(editable_default and (tab.name in referenced or not rows)),
            "columns": columns, "rows": rows})
    for synthetic in reg.synthetic.values():
        synthetic["editable"] = editable_default
        plan_tables.append(synthetic)

    # parameters
    params = list(reg.params.values())
    # checks of table calculated columns
    for tab in tabs.values():
        for column in tab.columns:
            if column.kind == "calculated" and column.cells:
                drafter.add_checks(tab.name, column.name, "table",
                                   [[cell, [i], value] for i, (cell, value)
                                    in enumerate(zip(column.cells, column.cached))])

    issues = list(drafter.issues)
    if any(f["kind"] == "summary" and not f.get("helper") for f in drafter.frames):
        issues.append({
            "code": "GROUPED_ROWS_FROM_DATA", "severity": "review", "route": "grouped-rows-from-data-review",
            "message": "Excel's summary row labels are typed cells; in Sigma each summary is a grouped table "
                       "whose rows come from the data, so new keys appear and missing ones vanish. Every "
                       "Excel row label was verified to exist in the data.",
            "evidence": [f["title"] for f in drafter.frames
                         if f["kind"] == "summary" and not f.get("helper")]})
    unmapped = list_charts(wf, drafter.checks)

    plan = {
        "archetype": trm.ARCHETYPE, "planVersion": trm.PLAN_VERSION,
        "name": name or (wf.properties.title or "").strip() or src.stem,
        "intent": intent, "connectionId": connection, "folderId": folder,
        "currencyFormat": currency_format, "pctFormat": pct_format,
        "sourcePath": str(src), "sourceSha256": source_sha256(src),
        "parametersEditable": editable_default,
        "params": params, "tables": plan_tables, "frames": drafter.frames,
        "checks": drafter.checks, "parityTargets": drafter.targets,
        "unmappedCharts": unmapped,
        "askUser": [
            "Which tables do people edit? Editable tables become empty input tables seeded afterwards "
            "(plan table.editable; the default is every table the formulas read, so a typed value can "
            "change a result)",
            "Confirm period/row grouping and the editable parameters "
            f"({', '.join(p['name'] for p in params) or 'none'})",
        ],
        "draftReviewIssues": issues,
    }
    reject_case_hazards(plan, tabs)
    try:
        trm.validate_plan(plan)
        model = trm.compute(plan)
    except (trm.PlanError, trm.EvalError) as error:
        raise PlanDraftError(f"the drafted plan does not evaluate: {error}") from error
    checked, problems = trm.check_cells(plan, model, tolerance)
    if problems:
        shown = "\n  ".join(problems[:10])
        more = f"\n  ... and {len(problems) - 10} more" if len(problems) > 10 else ""
        raise PlanDraftError("drafted plan does not reproduce the workbook's cached values "
                             f"({len(problems)} mismatch(es) of {checked} cells):\n  {shown}{more}")
    try:
        from . import table_report_ledger as ledger
    except ImportError:
        import table_report_ledger as ledger
    scenarios = ledger.scenario_results(plan, drafter)
    if scenarios:
        plan["scenarioResults"] = scenarios
    # completeness: every formula cell in the workbook must be one of the verified cells
    verified = {cell[0] for group in plan["checks"] for cell in group["cells"]}
    unverified = []
    for ws in wf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if (_is_formula(c.value) or c.data_type == "f") and f"{ws.title}!{c.coordinate}" not in verified:
                    unverified.append(f"{ws.title}!{c.coordinate}")
    if unverified:
        raise PlanDraftError(f"{unverified[0]}: formula cell is not covered by any recognised grid or table "
                             f"column ({len(unverified)} such cell(s))")
    formula_cells = sum(1 for ws in wf.worksheets for row in ws.iter_rows() for c in row
                        if _is_formula(c.value))
    summary = {
        "tables": [t["name"] for t in plan_tables], "frames": [(f["kind"], f["title"]) for f in drafter.frames],
        "verifiedCells": checked, "formulaCells": formula_cells, "params": len(params),
        "editable": [t["name"] for t in plan_tables if t["editable"]],
        "readOnly": [t["name"] for t in plan_tables if not t["editable"]],
        "reviewIssues": [i["code"] for i in issues], "unmappedCharts": unmapped,
        "parityTargets": len(plan["parityTargets"]),
    }
    return plan, summary
