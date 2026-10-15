"""Fail-loud capability inventory for OOXML Excel workbooks.

The audit does not claim that every inventoried feature is converted. It
classifies each feature into an automatic path or an explicit review route so a
migration cannot silently discard workbook behavior.
"""
from __future__ import annotations

import json
import re
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl

SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}
VOLATILE_FUNCTIONS = {
    "CELL", "INDIRECT", "NOW", "OFFSET", "RAND", "RANDBETWEEN", "TODAY",
}
KNOWN_FUNCTIONS = {
    "ABS", "AND", "AVERAGE", "AVERAGEIF", "AVERAGEIFS", "CEILING", "CHOOSE",
    "COLUMN", "COLUMNS", "CONCAT", "CONCATENATE", "COUNT", "COUNTA", "COUNTBLANK",
    "COUNTIF", "COUNTIFS", "DATE", "DATEVALUE", "DAY", "EDATE", "EOMONTH", "ERROR.TYPE",
    "FILTER", "FIND", "FLOOR", "HLOOKUP", "IF", "IFERROR", "IFS", "INDEX", "INT",
    "ISBLANK", "ISERROR", "ISNA", "LARGE", "LEFT", "LEN", "LET", "LOOKUP", "LOWER",
    "MATCH", "MAX", "MID", "MIN", "MOD", "MONTH", "NOT", "OR", "PROPER", "RANK",
    "REPLACE", "RIGHT", "ROUND", "ROUNDDOWN", "ROUNDUP", "ROW", "ROWS", "SEARCH",
    "SEQUENCE", "SMALL", "SORT", "SORTBY", "SUBSTITUTE", "SUM", "SUMIF", "SUMIFS",
    "SUMPRODUCT", "TEXT", "TEXTJOIN", "TIME", "TRIM", "UNIQUE", "UPPER", "VALUE",
    "VLOOKUP", "WEEKDAY", "XLOOKUP", "XMATCH", "YEAR",
}
FUNCTION_RE = re.compile(r"(?<![A-Z0-9_.])(?:_xlfn\.)?([A-Z][A-Z0-9_.]*)\s*\(", re.I)
EXTERNAL_REF_RE = re.compile(r"\[[^\]]+\.(?:xlsx|xlsm|xlsb|xls)\]", re.I)


def _issue(code, message, route, evidence=None, severity="review"):
    item = {
        "code": code,
        "severity": severity,
        "route": route,
        "message": message,
    }
    if evidence:
        item["evidence"] = evidence
    return item


def _zip_inventory(path):
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    lowered = [name.lower() for name in names]

    def any_path(*needles):
        return any(any(needle in name for needle in needles) for name in lowered)

    return {
        "vbaProject": any(name.endswith("vbaproject.bin") for name in lowered),
        "externalLinks": sum("/externallinks/" in name for name in lowered),
        "connections": any(name.endswith("/connections.xml") for name in lowered),
        "queryTables": sum("/querytables/" in name for name in lowered),
        "powerQuery": any_path("/queries/", "datamashup"),
        "pivotParts": sum("/pivottables/" in name for name in lowered),
        "slicerParts": sum(
            "/slicers/" in name or "/slicercaches/" in name for name in lowered
        ),
        "timelineParts": sum(
            "/timelines/" in name or "/timelinecaches/" in name for name in lowered
        ),
        "activeXParts": sum("/activex/" in name for name in lowered),
        "embeddedObjects": sum("/embeddings/" in name for name in lowered),
        "customXmlParts": sum(name.startswith("customxml/") and name.endswith(".xml")
                              for name in lowered),
        "webExtensions": sum("/webextensions/" in name for name in lowered),
    }


def _workbook_inventory(path):
    suffix = Path(path).suffix.lower()
    wb = openpyxl.load_workbook(
        path,
        data_only=False,
        keep_vba=suffix == ".xlsm",
        read_only=False,
    )
    functions = Counter()
    formulas = []
    external_formula_cells = []
    dynamic_formula_cells = []
    tables = []
    charts = []
    validations = 0
    conditional_formats = 0
    pivots = 0
    protected_sheets = []
    comments = 0
    images = 0

    for ws in wb.worksheets:
        tables.extend({
            "sheet": ws.title,
            "name": table.name,
            "ref": table.ref,
        } for table in ws.tables.values())
        charts.extend({
            "sheet": ws.title,
            "type": type(chart).__name__,
            "title": str(getattr(chart, "title", "") or ""),
        } for chart in getattr(ws, "_charts", []))
        validations += len(getattr(getattr(ws, "data_validations", None), "dataValidation", []) or [])
        conditional_formats += len(getattr(ws, "conditional_formatting", []) or [])
        pivots += len(getattr(ws, "_pivots", []) or [])
        if getattr(ws.protection, "sheet", False):
            protected_sheets.append(ws.title)
        images += len(getattr(ws, "_images", []) or [])

        for row in ws.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    comments += 1
                value = cell.value
                if not isinstance(value, str) or not value.startswith("="):
                    continue
                formulas.append(f"{ws.title}!{cell.coordinate}")
                functions.update(match.group(1).upper() for match in FUNCTION_RE.finditer(value))
                if EXTERNAL_REF_RE.search(value):
                    external_formula_cells.append(f"{ws.title}!{cell.coordinate}")
                if "#" in value or any(token in value.upper() for token in
                                       ("_XLFN.FILTER(", "_XLFN.UNIQUE(", "_XLFN.SORT(",
                                        "_XLFN.SEQUENCE(", "FILTER(", "UNIQUE(", "SORT(")):
                    dynamic_formula_cells.append(f"{ws.title}!{cell.coordinate}")

    defined_names = []
    for item in wb.defined_names.values():
        defined_names.append({
            "name": item.name,
            "type": getattr(item, "type", None),
            "hidden": bool(getattr(item, "hidden", False)),
        })

    workbook_protected = bool(
        getattr(wb.security, "lockStructure", False)
        or getattr(wb.security, "workbookPassword", None)
        or getattr(wb.security, "revisionsPassword", None)
    )
    return {
        "sheets": [
            {"name": ws.title, "state": ws.sheet_state, "rows": ws.max_row, "columns": ws.max_column}
            for ws in wb.worksheets
        ],
        "tables": tables,
        "formulaCells": len(formulas),
        "functions": dict(sorted(functions.items())),
        "charts": charts,
        "pivotTables": pivots,
        "dataValidations": validations,
        "conditionalFormattingRules": conditional_formats,
        "definedNames": defined_names,
        "protectedSheets": protected_sheets,
        "workbookProtected": workbook_protected,
        "comments": comments,
        "images": images,
        "externalFormulaCells": external_formula_cells[:25],
        "dynamicFormulaCells": dynamic_formula_cells[:25],
    }


def audit_workbook(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        return {
            "source": str(path),
            "status": "BLOCKED",
            "inventory": {},
            "issues": [_issue(
                "LEGACY_OR_UNSUPPORTED_FORMAT",
                f"{suffix or 'extensionless'} files are not OOXML workbooks",
                "convert-to-xlsx-or-xlsm",
                severity="blocker",
            )],
            "routing": {"automatic": [], "review": [], "blocked": ["convert-to-xlsx-or-xlsm"]},
        }
    if not path.exists():
        raise FileNotFoundError(path)

    package = _zip_inventory(path)
    workbook = _workbook_inventory(path)
    inventory = {**workbook, "package": package}
    issues = []

    function_names = set(workbook["functions"])
    unknown = sorted(function_names - KNOWN_FUNCTIONS - VOLATILE_FUNCTIONS)
    volatile = sorted(function_names & VOLATILE_FUNCTIONS)
    if unknown:
        issues.append(_issue(
            "UNKNOWN_FORMULAS",
            "Formula functions outside the translator allowlist require a disposition.",
            "formula-review-or-python",
            unknown,
        ))
    if volatile:
        issues.append(_issue(
            "VOLATILE_FORMULAS",
            "Volatile/location-dependent formulas need an explicit Sigma replacement.",
            "formula-review-or-python",
            volatile,
        ))
    if workbook["externalFormulaCells"] or package["externalLinks"]:
        issues.append(_issue(
            "EXTERNAL_WORKBOOK_LINKS",
            "External workbook references must be landed as governed Sigma sources.",
            "source-mapping-review",
            workbook["externalFormulaCells"] or [f"{package['externalLinks']} OOXML link part(s)"],
        ))
    if workbook["dynamicFormulaCells"]:
        issues.append(_issue(
            "DYNAMIC_ARRAY_FORMULAS",
            "Spill formulas require a grain-aware table or Python translation.",
            "formula-review-or-python",
            workbook["dynamicFormulaCells"],
        ))
    if package["connections"] or package["queryTables"] or package["powerQuery"]:
        issues.append(_issue(
            "POWER_QUERY_OR_CONNECTIONS",
            "Power Query and workbook connections are source pipelines, not cell formulas.",
            "warehouse-or-api-source-review",
        ))
    if package["vbaProject"]:
        issues.append(_issue(
            "VBA_PROJECT",
            "VBA must pass macro-classify.py; no procedure may be silently dropped.",
            "macro-classifier",
        ))
    if workbook["charts"]:
        issues.append(_issue(
            "CHARTS_PRESENT",
            "Chart types, encodings, filters, and aggregation must be mapped and rendered.",
            "chart-migration-and-render-check",
            sorted({chart["type"] for chart in workbook["charts"]}),
        ))
    if workbook["pivotTables"] or package["pivotParts"]:
        issues.append(_issue(
            "PIVOT_TABLES_PRESENT",
            "Pivot rows, columns, values, filters, subtotals, and sorts require mapping.",
            "pivot-migration-and-parity-check",
        ))
    if package["slicerParts"] or package["timelineParts"]:
        issues.append(_issue(
            "SLICERS_OR_TIMELINES",
            "Slicers and timelines must become Sigma controls with equivalent targets.",
            "control-migration-review",
        ))
    if workbook["dataValidations"]:
        issues.append(_issue(
            "DATA_VALIDATION_PRESENT",
            "Translate supported list/range/file constraints and review custom formulas.",
            "input-validation-migration",
        ))
    if workbook["conditionalFormattingRules"]:
        issues.append(_issue(
            "CONDITIONAL_FORMATTING_PRESENT",
            "Conditional formatting semantics need a visual parity check.",
            "formatting-migration-review",
        ))
    if workbook["protectedSheets"] or workbook["workbookProtected"]:
        issues.append(_issue(
            "WORKBOOK_PROTECTION",
            "Excel protection is not RLS/CLS; capture its intent and map permissions explicitly.",
            "security-review",
        ))
    if package["activeXParts"] or package["webExtensions"]:
        issues.append(_issue(
            "ADDINS_OR_ACTIVEX",
            "ActiveX and web add-ins cannot be inferred as workbook formulas.",
            "external-application-review",
            severity="blocker",
        ))
    if package["embeddedObjects"]:
        issues.append(_issue(
            "EMBEDDED_OBJECTS",
            "Embedded OLE objects require extraction and an explicit destination.",
            "content-extraction-review",
        ))
    if package["customXmlParts"]:
        issues.append(_issue(
            "CUSTOM_XML",
            "Custom XML may contain add-in or business metadata and must be dispositioned.",
            "package-metadata-review",
        ))
    if workbook["comments"] or workbook["images"]:
        issues.append(_issue(
            "ANNOTATIONS_OR_IMAGES",
            "Comments, notes, and images require content-level migration review.",
            "content-migration-review",
            [f"comments={workbook['comments']}", f"images={workbook['images']}"],
        ))

    blockers = [item for item in issues if item["severity"] == "blocker"]
    status = "BLOCKED" if blockers else ("REVIEW_REQUIRED" if issues else "AUTO")
    automatic = [
        "formal-table-to-input-or-data-model" if workbook["tables"] else
        "cell-model-or-research-template-routing"
    ]
    routes = sorted({item["route"] for item in issues})
    return {
        "source": str(path),
        "status": status,
        "inventory": inventory,
        "issues": issues,
        "routing": {
            "automatic": automatic,
            "review": routes if not blockers else [
                item["route"] for item in issues if item["severity"] != "blocker"
            ],
            "blocked": sorted({item["route"] for item in blockers}),
        },
    }


def dumps(report):
    return json.dumps(report, indent=2, sort_keys=False)
