"""Versioned, source-bound migration manifests and evidence gates."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

try:
    from .excel_expression import UnsupportedExpression, extract_references, parse_expression
    from .formula_graph import FormulaGraph
except ImportError:  # standalone tests/scripts that place lib directly on sys.path
    from excel_expression import UnsupportedExpression, extract_references, parse_expression
    from formula_graph import FormulaGraph


MANIFEST_VERSION = 1


class ManifestError(ValueError):
    pass


_RUNTIME_CONTAINERS = {
    "evidence", "runtimeEvidence", "execution", "executionState", "verification",
    "targetVersions", "targetVersion", "targetBinding", "actualTarget", "runtime",
}
_ASSIGNED_TARGET_FIELDS = {"workbookId", "elementId", "documentId", "actualColumnId", "version"}
_ROOT_SOURCE_METADATA = {"source", "sourcePath", "sourceSha256", "migrationManifest"}


def _sha256(path):
    with Path(path).open("rb") as stream:
        return hashlib.sha256(stream.read()).hexdigest()


def _logical(value, *, parent=None):
    if isinstance(value, dict):
        return {
            key: _logical(item, parent=key)
            for key, item in sorted(value.items())
            if key not in _RUNTIME_CONTAINERS
            and not (parent == "targetMapping" and key in _ASSIGNED_TARGET_FIELDS)
        }
    if isinstance(value, list):
        return [_logical(item, parent=parent) for item in value]
    return value


def _plan_hash(source_hash, plan):
    plan = plan or {}
    projected = {key: value for key, value in plan.items() if key not in _ROOT_SOURCE_METADATA}
    logical = {"sourceHash": source_hash, "plan": _logical(projected)}
    encoded = json.dumps(logical, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def build_manifest(source_path, plan=None) -> dict:
    source = Path(source_path).expanduser().resolve()
    if not source.is_file():
        raise ManifestError(f"source workbook not found: {source}")
    source_hash = _sha256(source)
    workbook = load_workbook(source, data_only=False, read_only=False)
    cached = load_workbook(source, data_only=True, read_only=False)
    graph = FormulaGraph.from_workbook(workbook)
    cycle_members = {cell for component in graph.cycles() for cell in component}
    plan_cells = _plan_cell_index(plan or {})
    formulas = []
    for identity in sorted(graph.formulas):
        formula = graph.formulas[identity]
        record = plan_cells.get(identity.id, {})
        mapping_status = record.get("mappingStatus") or ("verification" if identity.sheet == "Audit" else "unresolved")
        try:
            expression = parse_expression(formula, sheet=identity.sheet)
            named_refs = [_reference_id(ref) for ref in extract_references(expression) if ref.name]
            refs = sorted({cell.id for cell in graph.dependencies.get(identity, ())} | set(named_refs))
            parse_error = None
        except UnsupportedExpression as error:
            refs, parse_error = [], str(error)
        item = {
            "id": identity.id,
            "sheet": identity.sheet,
            "address": identity.address,
            "originalFormula": formula,
            "dependencies": refs,
            "mappingStatus": mapping_status,
            "targetMapping": record.get("targetMapping"),
            "evidence": list(record.get("evidence") or []),
            "verification": record.get("verification"),
        }
        cached_value = cached[identity.sheet][identity.address].value
        if record.get("mappingStatus") == "cached":
            item["cachedFallback"] = cached_value
        if parse_error:
            item["mappingStatus"] = "unresolved"
            item["unresolvedReason"] = parse_error
        if identity in cycle_members:
            item["mappingStatus"] = "unresolved"
            item["unresolvedReason"] = "true cell dependency cycle"
        elif identity in graph.unresolved_references:
            item["mappingStatus"] = "unresolved"
            item["unresolvedReason"] = "missing source reference"
        formulas.append(item)
    literals = []
    for identity in sorted(graph.literals):
        value = graph.literals[identity]
        if isinstance(value, (str, int, float, bool)):
            literals.append({"id": identity.id, "sheet": identity.sheet,
                             "address": identity.address, "value": value})
    workbook.close(); cached.close()
    return {
        "manifestVersion": MANIFEST_VERSION,
        "source": {"name": source.name, "path": str(source), "sha256": source_hash},
        "sourceHash": source_hash,
        "planHash": _plan_hash(source_hash, plan),
        "formulaCells": formulas,
        "literalInputs": literals,
        "cycles": [[cell.id for cell in component] for component in graph.cycles()],
        "unsupported": {cell.id: reason for cell, reason in sorted(graph.unsupported_formulas.items())},
        "stagingApproval": None,
        "target": None,
    }


def _plan_cell_index(plan):
    result = {}
    for key in ("formulaCells", "cells", "provenance"):
        items = plan.get(key) or []
        if isinstance(items, dict):
            items = items.values()
        for item in items:
            if isinstance(item, dict) and item.get("id"):
                result[item["id"]] = item
    return result


def _reference_id(ref):
    if ref.name:
        return f"name:{ref.name}"
    return f"{ref.sheet}!{ref.address}"


def validate_manifest(manifest, source_path=None):
    if not isinstance(manifest, dict) or manifest.get("manifestVersion") != MANIFEST_VERSION:
        raise ManifestError("unsupported or malformed manifest version")
    source_hash = manifest.get("sourceHash")
    plan_hash = manifest.get("planHash")
    if not _digest(source_hash) or not _digest(plan_hash):
        raise ManifestError("manifest requires valid sourceHash and planHash")
    formulas = manifest.get("formulaCells")
    literals = manifest.get("literalInputs")
    if not isinstance(formulas, list) or not isinstance(literals, list):
        raise ManifestError("manifest requires formulaCells and literalInputs arrays")
    formula_ids = [item.get("id") for item in formulas if isinstance(item, dict)]
    literal_ids = [item.get("id") for item in literals if isinstance(item, dict)]
    if len(formula_ids) != len(formulas) or len(set(formula_ids)) != len(formula_ids):
        raise ManifestError("each source formula must appear exactly once")
    if set(formula_ids) & set(literal_ids):
        raise ManifestError("literal inputs and formula cells must be distinct")
    if any(item.get("mappingStatus") not in {"live", "cached", "unresolved", "verification"}
           for item in formulas):
        raise ManifestError("formula mappingStatus is invalid")
    if source_path is not None and _sha256(source_path) != source_hash:
        raise ManifestError("source hash does not match manifest")
    return True


def validate_staging_approval(manifest, approval):
    validate_manifest(manifest)
    if not isinstance(approval, dict):
        raise ManifestError("staging approval must be an object")
    if approval.get("sourceHash") != manifest["sourceHash"] or approval.get("planHash") != manifest["planHash"]:
        raise ManifestError("staging approval is stale")
    if not str(approval.get("owner") or "").strip():
        raise ManifestError("staging approval requires a named owner")
    mapped = {item["id"] for item in manifest["formulaCells"]
              if item["mappingStatus"] != "unresolved"}
    reviewed = approval.get("reviewItems")
    if not isinstance(reviewed, list) or set(reviewed) != mapped:
        raise ManifestError("staging approval must explicitly list every mapped review item")
    unresolved = [item["id"] for item in manifest["formulaCells"]
                  if item["mappingStatus"] == "unresolved"]
    if unresolved:
        raise ManifestError("staging approval cannot include unresolved formula mappings")
    return True


def validate_completion(manifest):
    validate_manifest(manifest)
    target = manifest.get("target") or {}
    if not target.get("documentId") or target.get("version") is None:
        raise ManifestError("completion requires an actual target document ID and version")
    for item in manifest["formulaCells"]:
        if item.get("mappingStatus") in {"cached", "unresolved"}:
            raise ManifestError(f"formula is not completion-ready: {item['id']}")
        evidence = item.get("verification") or {}
        mapping = item.get("targetMapping") or {}
        if evidence.get("status") != "passed" or not evidence.get("evidenceId"):
            raise ManifestError(f"verification evidence missing for {item['id']}")
        if evidence.get("sourceCellId") != item["id"]:
            raise ManifestError(f"verification evidence source binding is stale for {item['id']}")
        if not evidence.get("targetId"):
            raise ManifestError(f"verification evidence target binding is missing for {item['id']}")
        expected_target = mapping.get("columnId") or mapping.get("lineId") or mapping.get("checkId")
        if expected_target and evidence.get("targetId") != expected_target:
            raise ManifestError(f"verification evidence target binding is stale for {item['id']}")
        if evidence.get("sourceHash") != manifest["sourceHash"]:
            raise ManifestError(f"verification evidence source hash is stale for {item['id']}")
        if evidence.get("planHash") != manifest["planHash"]:
            raise ManifestError(f"verification evidence plan hash is stale for {item['id']}")
        if target and (evidence.get("targetVersion") != target.get("version") or
                       evidence.get("targetDocumentId") != target.get("documentId")):
            raise ManifestError(f"verification evidence target version is stale for {item['id']}")
    return True


def _digest(value):
    return isinstance(value, str) and len(value) == 64 and all(ch in "0123456789abcdef" for ch in value)
