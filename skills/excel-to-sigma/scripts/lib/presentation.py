"""Source-bound chart, validation, formatting, and layout migration.

Extraction is generic OOXML/openpyxl work. Emission requires explicit live target
bindings and never uses cached formula results as chart data. Cached values are
retained only as independent source-side acceptance expectations.
"""
from __future__ import annotations

import copy
import hashlib
import json
import posixpath
import re
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

try:
    from .input_sync import _sdk_literal, sigma_input_type
except ImportError:
    from input_sync import _sdk_literal, sigma_input_type

try:
    from .input_sync import build_operational_input_elements, build_shared_input_architecture
except ImportError:
    from input_sync import build_operational_input_elements, build_shared_input_architecture


MAIN_COLUMNS = {
    "forecast": ["Line Item", "Method", "Monthly Growth", "Jul-26", "Aug-26", "Sep-26", "Oct-26", "Nov-26", "Dec-26"],
    "research": ["Metric", "Unit", "Forecast Value", "Description"],
    "sales": ["Opportunity ID", "Account", "Region", "Segment", "Rep", "Stage", "Probability", "ACV", "Weighted ACV", "Expected Close", "Timing Status"],
    "inventory": ["SKU", "Product", "Category", "Supplier", "Lead Days", "On Hand", "Allocated", "Available", "Reorder Point", "Reorder Qty", "Status"],
    "expense": ["Expense ID", "Date", "Employee", "Department", "Category", "Amount", "Receipt Status", "Approval Status", "Exception"],
}
EDITABLE_COLUMNS = {
    "forecast": ["Method", "Monthly Growth", "Jul-26", "Aug-26", "Sep-26", "Oct-26", "Nov-26", "Dec-26"],
    "research": ["Forecast Value"],
    "sales": ["Region", "Segment", "Stage", "Probability", "ACV", "Expected Close"],
    "inventory": ["Category", "Supplier", "Lead Days", "On Hand", "Allocated", "Unit Cost", "Avg Daily Use", "Safety Stock"],
    "expense": ["Department", "Category", "Amount", "Receipt Status", "Approval Status", "Card Amount", "Ledger Amount"],
}
CALCULATED_COLUMNS = {
    "forecast": ["Forecast Amount", "FY Total"],
    "research": ["Revenue", "EBITDA", "EPS", "Net debt / EBITDA"],
    "sales": ["Weighted ACV", "Cycle Days", "Timing Status"],
    "inventory": ["Available", "Reorder Point", "Days Cover", "Reorder Qty", "Status"],
    "expense": ["Policy Limit", "Recon Variance", "Policy Excess", "Exception"],
}
DETAIL_COLUMNS = {
    "forecast": MAIN_COLUMNS["forecast"] + ["Section", "Section Order", "Line Order", "Owner"],
    "research": [f"Source Cell {i}" for i in range(1, 146)],
    "sales": ["Opportunity ID", "Created Date", "Expected Close", "Account", "Region", "Segment", "Rep", "Stage", "Probability", "ACV", "Forecast Category", "Weighted ACV", "Cycle Days", "Timing Status"],
    "inventory": ["SKU", "Product", "Category", "Supplier", "Lead Days", "On Hand", "Allocated", "Unit Cost", "Avg Daily Use", "Safety Stock", "Available", "Reorder Point", "Days Cover", "Reorder Qty", "Status"],
    "expense": ["Expense ID", "Date", "Employee", "Department", "Category", "Vendor", "Amount", "Receipt Status", "Approval Status", "Card Amount", "Ledger Amount", "Policy Limit", "Recon Variance", "Policy Excess", "Exception"],
}
BUSINESS_TABLES = {
    "forecast": ("Forecast Inputs", "tblFixtureForecastInputs", ["Line Item"]),
    "sales": ("Pipeline", "tblSalesPipeline", ["Opportunity ID"]),
    "inventory": ("Inventory", "tblInventory", ["SKU"]),
    "expense": ("Expense Register", "tblExpenseRegister", ["Expense ID"]),
}

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def _digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"), default=str).encode()).hexdigest()


def _workflow(path):
    name = Path(path).name.lower()
    for workflow in ("forecast", "research", "sales", "inventory", "expense"):
        if workflow in name:
            return workflow
    return re.sub(r"[^a-z0-9]+", "-", Path(path).stem.lower()).strip("-")


def _resolve(base, target):
    return posixpath.normpath(posixpath.join(posixpath.dirname(base), target)).lstrip("/")


def _rels(archive, part):
    rel_path = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if rel_path not in archive.namelist():
        return {}
    root = ET.fromstring(archive.read(rel_path))
    return {node.attrib["Id"]: _resolve(part, node.attrib["Target"])
            for node in root.findall("pr:Relationship", NS)}


def _chart_locations(path):
    locations = {}
    with ZipFile(path) as archive:
        workbook_part = "xl/workbook.xml"
        workbook = ET.fromstring(archive.read(workbook_part))
        workbook_rels = _rels(archive, workbook_part)
        for sheet in workbook.findall("m:sheets/m:sheet", NS):
            sheet_name = sheet.attrib["name"]
            sheet_part = workbook_rels[sheet.attrib[f"{{{NS['r']}}}id"]]
            sheet_xml = ET.fromstring(archive.read(sheet_part))
            sheet_rels = _rels(archive, sheet_part)
            for drawing_node in sheet_xml.findall("m:drawing", NS):
                drawing_part = sheet_rels[drawing_node.attrib[f"{{{NS['r']}}}id"]]
                drawing = ET.fromstring(archive.read(drawing_part))
                drawing_rels = _rels(archive, drawing_part)
                for anchor in list(drawing):
                    chart_node = anchor.find(".//c:chart", NS)
                    if chart_node is None:
                        continue
                    chart_part = drawing_rels[chart_node.attrib[f"{{{NS['r']}}}id"]]
                    marker = anchor.find("xdr:from", NS)
                    anchor_cell = None
                    if marker is not None:
                        col = int(marker.findtext("xdr:col", default="0", namespaces=NS)) + 1
                        row = int(marker.findtext("xdr:row", default="0", namespaces=NS)) + 1
                        anchor_cell = f"{get_column_letter(col)}{row}"
                    locations[chart_part] = {
                        "worksheet": sheet_name, "worksheetPart": sheet_part,
                        "drawing": drawing_part, "chartPart": chart_part,
                        "anchor": anchor_cell,
                    }
    return locations


def _formula_range(reference):
    match = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^!]+))!(\$?[A-Z]+\$?\d+):(\$?[A-Z]+\$?\d+)", reference or "")
    if not match:
        return None
    sheet = (match.group(1) or match.group(2)).replace("''", "'")
    return sheet, match.group(3).replace("$", ""), match.group(4).replace("$", "")


def _range_values(workbook, reference):
    parsed = _formula_range(reference)
    if not parsed:
        return []
    sheet, start, end = parsed
    return [cell.value for row in workbook[sheet][f"{start}:{end}"] for cell in row]


def _range_cells(reference):
    parsed = _formula_range(reference)
    if not parsed:
        return []
    sheet, start, end = parsed
    min_col, min_row, max_col, max_row = range_boundaries(f"{start}:{end}")
    return [f"{sheet}!{get_column_letter(column)}{row}"
            for row in range(min_row, max_row + 1)
            for column in range(min_col, max_col + 1)]


def _unit(format_code, title=""):
    value = f"{format_code or ''} {title}".lower()
    if "$" in value or "usd" in value:
        return "currency"
    if "%" in value:
        return "percentage"
    if "count" in value:
        return "count"
    if "eps" in value or "leverage" in value or "/ ebitda" in value:
        return "ratio"
    return "number"


def _aggregation(formulas):
    kinds = set()
    for formula in formulas:
        if not isinstance(formula, str) or not formula.startswith("="):
            kinds.add("literal")
        elif re.match(r"=SUM(?:IFS?|PRODUCT)?\(", formula, re.I):
            kinds.add("sum")
        elif re.match(r"=COUNT(?:A|IF|IFS)?\(", formula, re.I):
            kinds.add("count")
        elif re.fullmatch(r"='?[^']+'?!\$?[A-Z]+\$?\d+", formula):
            kinds.add("direct")
        else:
            kinds.add("expression")
    return kinds.pop() if len(kinds) == 1 else "mixed:" + ",".join(sorted(kinds))


def _extract_charts(path, formulas, cached):
    locations = _chart_locations(path)
    charts = []
    with ZipFile(path) as archive:
        for index, chart_part in enumerate(sorted(locations, key=lambda p: int(re.search(r"(\d+)", Path(p).stem).group(1))), 1):
            root = ET.fromstring(archive.read(chart_part))
            title = "".join(node.text or "" for node in root.findall(".//c:title//a:t", NS)) or f"Chart {index}"
            chart_type = "line" if root.find(".//c:lineChart", NS) is not None else "bar" if root.find(".//c:barChart", NS) is not None else "unsupported"
            series_nodes = root.findall(f".//c:{chart_type}Chart/c:ser", NS) if chart_type != "unsupported" else []
            series = []
            category_ref = None
            for order, node in enumerate(series_nodes):
                name = "".join(item.text or "" for item in node.findall(".//c:tx//c:v", NS)) or f"Series {order + 1}"
                cat = node.find(".//c:cat//c:f", NS)
                val = node.find(".//c:val//c:f", NS)
                fmt = node.find(".//c:val//c:numCache/c:formatCode", NS)
                category_ref = cat.text if cat is not None else category_ref
                value_ref = val.text if val is not None else None
                source_formulas = _range_values(formulas, value_ref) if value_ref else []
                expected = _range_values(cached, value_ref) if value_ref else []
                series.append({
                    "name": name, "order": order, "valueRef": value_ref,
                    "numberFormat": fmt.text if fmt is not None else None,
                    "unit": _unit(fmt.text if fmt is not None else None, title),
                    "aggregation": _aggregation(source_formulas),
                    "sourceFormulas": source_formulas,
                    "sourceFormulaCells": _range_cells(value_ref),
                    "expectedValues": expected,
                })
            charts.append({
                "id": f"chart-{index}", "title": title, "chartType": chart_type,
                "categoryRef": category_ref, "categories": _range_values(cached, category_ref),
                "categoryOrder": "source", "series": series,
                "unit": series[0]["unit"] if series and len({s["unit"] for s in series}) == 1 else "mixed",
                "status": "supported" if chart_type in {"bar", "line"} else "unresolved",
                "sourceLocation": locations[chart_part],
            })
    return charts


def _field_for_range(ws, cell_range):
    min_col, min_row, max_col, max_row = range_boundaries(str(cell_range).split()[0])
    if max_col > min_col and max_row == min_row:
        for column in range(min_col - 1, 0, -1):
            value = ws.cell(min_row, column).value
            if value not in (None, "") and not (isinstance(value, str) and value.startswith("=")):
                return str(value).strip()
    for row in range(min_row - 1, 0, -1):
        value = ws.cell(row, min_col).value
        if value not in (None, "") and not (isinstance(value, str) and value.startswith("=")):
            return str(value)
    return get_column_letter(min_col)


def _fields_for_range(ws, cell_range):
    min_col, _min_row, max_col, _max_row = range_boundaries(str(cell_range).split()[0])
    if max_col == min_col:
        return [_field_for_range(ws, cell_range)]
    return [_field_for_range(ws, f"{get_column_letter(column)}{_min_row}:{get_column_letter(column)}{_max_row}")
            for column in range(min_col, max_col + 1)]


def _list_values(formula):
    if isinstance(formula, str) and len(formula) >= 2 and formula[0] == formula[-1] == '"':
        return formula[1:-1].split(",")
    return None


def validation_accepts(rule, value):
    if value in (None, ""):
        return bool(rule.get("allowBlank"))
    emitted = rule.get("emission") or {}
    if "values" in emitted:
        return value in emitted["values"] and type(value) in {type(item) for item in emitted["values"]}
    limits = emitted.get("range") or {}
    try:
        number = float(value)
    except (TypeError, ValueError):
        return False
    return (limits.get("min") is None or number >= limits["min"]) and (limits.get("max") is None or number <= limits["max"])


def _extract_validations(workbook):
    result = []
    for ws in workbook.worksheets:
        for dv in ws.data_validations.dataValidation:
            for sqref in str(dv.sqref).split():
                values = _list_values(dv.formula1)
                status = "supported"
                if dv.type == "list" and values is not None:
                    emission = {"values": values}
                elif dv.type == "decimal" and dv.operator == "between":
                    emission = {"range": {"min": float(dv.formula1), "max": float(dv.formula2)}}
                elif dv.type == "whole" and dv.operator == "between":
                    low, high = int(dv.formula1), int(dv.formula2)
                    emission = {"values": list(range(low, high + 1))}
                else:
                    emission, status = None, "unresolved"
                result.append({
                    "id": f"validation-{len(result) + 1}", "sheet": ws.title,
                    "range": sqref, "field": _field_for_range(ws, sqref),
                    "fields": _fields_for_range(ws, sqref),
                    "sourceType": dv.type, "operator": dv.operator,
                    "allowBlank": bool(dv.allow_blank), "formula1": dv.formula1,
                    "formula2": dv.formula2, "inclusive": dv.operator == "between",
                    "status": status, "emission": emission,
                })
    return result


def _rgb(color):
    if color is None:
        return None
    value = getattr(color, "rgb", None)
    if isinstance(value, str):
        return "#" + value[-6:]
    return None


def _dxf_style(rule):
    dxf = getattr(rule, "dxf", None)
    if dxf is None:
        return {}
    style = {}
    if getattr(dxf, "fill", None) is not None:
        color = _rgb(getattr(dxf.fill, "fgColor", None)) or _rgb(getattr(dxf.fill, "start_color", None))
        if color:
            style["backgroundColor"] = color
    if getattr(dxf, "font", None) is not None:
        color = _rgb(dxf.font.color)
        if color:
            style["color"] = color
        for key in ("bold", "italic", "underline"):
            source = {"bold": "b", "italic": "i", "underline": "u"}[key]
            value = getattr(dxf.font, source, None)
            if value:
                style[key] = True
    return style


def _relative_reference(ws, sqref, formula):
    min_col, min_row, _max_col, _max_row = range_boundaries(str(sqref).split()[0])
    refs = []
    for col_abs, col, row_abs, row in re.findall(r"(\$?)([A-Z]{1,3})(\$?)(\d+)", formula or ""):
        refs.append({
            "reference": f"{col_abs}{col}{row_abs}{row}",
            "columnOffset": 0 if col_abs else column_index_from_string(col) - min_col,
            "rowOffset": 0 if row_abs else int(row) - min_row,
            "absoluteColumn": bool(col_abs), "absoluteRow": bool(row_abs),
        })
    return {"origin": f"{get_column_letter(min_col)}{min_row}", "references": refs}


def _extract_conditional_formats(workbook):
    result = []
    condition_map = {"lessThan": "<", "lessThanOrEqual": "<=", "greaterThan": ">",
                     "greaterThanOrEqual": ">=", "equal": "=", "notEqual": "!=",
                     "between": "Between", "notBetween": "NotBetween",
                     "containsText": "Contains", "notContainsText": "NotContains"}
    for ws in workbook.worksheets:
        for key, rules in ws.conditional_formatting._cf_rules.items():
            sqref = str(key.sqref)
            for rule in rules:
                source_type = rule.type
                emitted = None
                status = "supported"
                formula = (rule.formula or [None])[0]
                if source_type == "colorScale":
                    scale = rule.colorScale
                    colors = [_rgb(color) for color in scale.color]
                    cfvos = list(scale.cfvo)
                    domain = {"min": "Min([Column])", "max": "Max([Column])"}
                    if len(cfvos) == 3:
                        mid = cfvos[1]
                        if mid.type == "percentile":
                            domain["mid"] = f"Percentile([Column], {float(mid.val) / 100:g})"
                        elif mid.type == "num":
                            domain["mid"] = float(mid.val)
                        elif mid.type == "formula":
                            domain["mid"] = str(mid.val)
                        else:
                            status = "unresolved"
                    emitted = {"type": "backgroundScale", "columnIds": [], "scheme": colors,
                               "domain": domain, "order": "ascending", "nullBehavior": "asNull"}
                else:
                    condition = condition_map.get(rule.operator or source_type)
                    value = rule.text if rule.text is not None else formula
                    if condition in {"<", "<=", ">", ">=", "=", "!="}:
                        try:
                            value = float(value)
                            if value.is_integer(): value = int(value)
                        except (TypeError, ValueError):
                            pass
                    if condition:
                        emitted = {"type": "single", "columnIds": [], "condition": condition,
                                   "value": value, "style": _dxf_style(rule)}
                    else:
                        status = "unresolved"
                result.append({
                    "id": f"conditional-{len(result) + 1}", "sheet": ws.title,
                    "range": sqref, "field": _field_for_range(ws, sqref),
                    "sourceType": source_type, "operator": rule.operator,
                    "formula": formula, "text": rule.text, "priority": rule.priority,
                    "stopIfTrue": bool(rule.stopIfTrue), "status": status,
                    "relativeReference": _relative_reference(ws, sqref, formula),
                    "emission": emitted,
                })
    return sorted(result, key=lambda item: (item["sheet"], item["range"], item["priority"]))


def _precision(number_format):
    section = str(number_format or "").split(";")[0]
    match = re.search(r"[0#],?([0#]+)?\.([0#]+)", section)
    return len(match.group(2)) if match else 0


def _semantic_formats(workbook):
    result = {}
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                fmt = cell.number_format or "General"
                if cell.is_date:
                    semantic = "date"
                    emission = {"kind": "datetime", "formatString": "%Y-%m-%d"}
                    grain = "day" if getattr(cell.value, "hour", 0) == getattr(cell.value, "minute", 0) == getattr(cell.value, "second", 0) == 0 else "timestamp"
                elif "%" in fmt:
                    semantic, emission, grain = "percentage", {"kind": "number", "formatString": f".{_precision(fmt)}%"}, None
                elif "$" in fmt:
                    semantic, emission, grain = "currency", {"kind": "number", "formatString": f",.{_precision(fmt)}f", "currencySymbol": "$"}, None
                elif isinstance(cell.value, (int, float)) or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    semantic, emission, grain = "number", {"kind": "number", "formatString": f",.{_precision(fmt)}f"}, None
                    if str(fmt).rstrip().lower().endswith("x"):
                        emission["suffix"] = "x"
                else:
                    continue
                result[f"{ws.title}!{cell.coordinate}"] = {
                    "sourceFormat": fmt, "semantic": semantic, "precision": _precision(fmt),
                    "dateGrain": grain, "emission": emission,
                }
    return result


def extract_presentation(source_path):
    """Extract presentation features with source identities and no target guesses."""
    source = Path(source_path).resolve()
    formulas = load_workbook(source, data_only=False, read_only=False)
    cached = load_workbook(source, data_only=True, read_only=False)
    formula_provenance = []
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_provenance.append({"cell": f"{ws.title}!{cell.coordinate}",
                                               "formula": cell.value, "status": "source-accounted"})
    result = {
        "version": 1, "source": {"path": str(source), "name": source.name,
                                    "sha256": hashlib.sha256(source.read_bytes()).hexdigest()},
        "workflow": _workflow(source),
        "charts": _extract_charts(source, formulas, cached),
        "validations": _extract_validations(formulas),
        "conditionalFormats": _extract_conditional_formats(formulas),
        "semanticFormats": _semantic_formats(formulas),
        "formulaProvenance": formula_provenance,
    }
    result["extractionHash"] = _digest({key: value for key, value in result.items() if key != "extractionHash"})
    formulas.close(); cached.close()
    return result


def _chart_binding(title, workflow):
    exact = {
        "Monthly revenue, EBITDA, and net income": ("pnlCombined", 'DateTrunc("month", [PnL Combined/Month End])', "Min([PnL Combined/Month Num])", ["Sum([PnL Combined/Revenue Amt])", "Sum([PnL Combined/Revenue Amt]) - Sum([PnL Combined/COGS Amt]) - Sum([PnL Combined/OpEx Amt])", "Sum([PnL Combined/Revenue Amt]) - Sum([PnL Combined/COGS Amt]) - Sum([PnL Combined/OpEx Amt]) - Sum([PnL Combined/Below Amt])"]),
        "FY total by statement section": ("pnlCombined", "[PnL Combined/Section]", "Min([PnL Combined/Section Order])", ["Sum([PnL Combined/Amount])"]),
        "Revenue and EBITDA (USD m)": ("researchCalc", "[Calculation/Year]", "Min([Calculation/Year])", ["Sum([Calculation/c8])", "Sum([Calculation/c18])"]),
        "EPS and leverage": ("researchCalc", "[Calculation/Year]", "Min([Calculation/Year])", ["Sum([Calculation/c37])", "Sum([Calculation/c55]) / Sum([Calculation/c18])"]),
        "Weighted pipeline by stage": ("salesPipeline", "[Pipeline/Stage]", 'Min(Switch([Pipeline/Stage], "Prospecting", 1, "Qualification", 2, "Proposal", 3, "Negotiation", 4, "Closed Won", 5, "Closed Lost", 6))', ["Sum([Pipeline/Weighted ACV])"]),
        "Weighted pipeline by expected close month": ("salesPipeline", 'DateTrunc("month", [Pipeline/Expected Close])', 'Min(DateTrunc("month", [Pipeline/Expected Close]))', ["Sum([Pipeline/Weighted ACV])"]),
        "Recommended reorder quantity by category": ("inventory", "[Inventory/Category]", 'Min(Switch([Inventory/Category], "Electronics", 1, "Components", 2, "Packaging", 3, "Service Parts", 4))', ["Sum([Inventory/Reorder Qty])"]),
        "Available inventory value by supplier": ("inventory", "[Inventory/Supplier]", 'Min(Switch([Inventory/Supplier], "Apex Supply", 1, "Beacon Components", 2, "CorePack", 3, "Delta Service Co", 4))', ["Sum([Inventory/Available] * [Inventory/Unit Cost])"]),
        "Spend by department": ("expenseRegister", "[Expense Register/Department]", 'Min(Switch([Expense Register/Department], "Sales", 1, "Engineering", 2, "Operations", 3, "Corporate", 4))', ["Sum([Expense Register/Amount])"]),
        "Expenses by review status": ("expenseRegister", "[Expense Register/Exception]", 'Min(Switch([Expense Register/Exception], "RECON", 1, "OVER LIMIT", 2, "MISSING RECEIPT", 3, "PENDING", 4, "CLEAR", 5))', ["Count([Expense Register/Expense ID])"]),
    }
    return exact.get(title)


def _source_columns(plan):
    columns = defaultdict(dict)
    workbook = load_workbook(plan["source"]["path"], data_only=False, read_only=False)
    for ws in workbook.worksheets:
        for table in ws.tables.values():
            min_col, min_row, max_col, _max_row = range_boundaries(table.ref)
            for col in range(min_col, max_col + 1):
                name = str(ws.cell(min_row, col).value)
                columns[ws.title][name] = {
                    "columnId": re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_"),
                    "sourceCell": f"{ws.title}!{get_column_letter(col)}{min_row + 1}",
                }
    workbook.close()
    return columns


def _research_view_columns(plan):
    workbook = load_workbook(plan["source"]["path"], data_only=False, read_only=False)
    ws = workbook["FY results"]
    metrics = []
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if label in (None, "") or not any(ws.cell(row, column).value is not None for column in range(2, 14)):
            continue
        semantic = next((plan["semanticFormats"].get(f"FY results!{get_column_letter(column)}{row}")
                         for column in range(2, 14)
                         if plan["semanticFormats"].get(f"FY results!{get_column_letter(column)}{row}")), None)
        column = {"id": f"research_c{row}", "name": str(label).strip(),
                  "formula": f"[Calculation/c{row}]", "sourceCell": f"FY results!B{row}"}
        if semantic:
            column["format"] = copy.deepcopy(semantic["emission"])
        metrics.append(column)
    workbook.close()
    year = {"id": "research_year", "name": "Year", "formula": "[Calculation/Year]",
            "format": {"kind": "number", "formatString": ".0f"}}
    by_row = {int(column["id"].split("c")[-1]): column for column in metrics}
    main = [year] + [copy.deepcopy(by_row[row]) for row in (8, 18, 37, 57)]
    return main, [year] + metrics


def fixture_bindings(plan):
    """Return logical target bindings for the five synthetic acceptance fixtures.

    Runtime callers replace IDs with GET/readback IDs before merge/PUT.
    """
    workflow = plan["workflow"]
    charts = {}
    for chart in plan["charts"]:
        binding = _chart_binding(chart["title"], workflow)
        if binding:
            element_id, category, order_formula, series = binding
            charts[chart["id"]] = {"sourceElementId": element_id,
                                    "categoryFormula": category, "orderFormula": order_formula,
                                    "seriesFormulas": series}
    sheets = _source_columns(plan)
    element_for_sheet = {sheet: re.sub(r"[^a-z0-9]+", "-", sheet.lower()).strip("-") for sheet in sheets}
    if workflow in BUSINESS_TABLES:
        element_for_sheet[BUSINESS_TABLES[workflow][0]] = f"{workflow}-input"
    if workflow == "research":
        element_for_sheet["Model Inputs"] = "research-input-existing"
    columns = {f"{sheet}!{name}": {"elementId": element_for_sheet[sheet], **info}
               for sheet, fields in sheets.items() for name, info in fields.items()}
    detail = list(DETAIL_COLUMNS.get(workflow, []))
    edit = list(EDITABLE_COLUMNS.get(workflow, []))
    if workflow == "research":
        edit = detail[:6]
    source_names = {"forecast": "PnL Combined", "research": "Calculation", "sales": "Pipeline",
                    "inventory": "Inventory", "expense": "Expense Register"}
    source_ids = {"forecast": "pnlCombined", "research": "researchCalc", "sales": "salesPipeline",
                  "inventory": "inventory", "expense": "expenseRegister"}
    source_name, source_id = source_names[workflow], source_ids[workflow]
    default_view = lambda names: [
        {"id": re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_"),
         "name": name, "formula": f"[{source_name}/{name}]"}
        for name in names
    ]
    main_view_columns = default_view(MAIN_COLUMNS.get(workflow, []))
    detail_view_columns = default_view(detail)
    research_formats_exact = None
    if workflow == "research":
        main_view_columns, detail_view_columns = _research_view_columns(plan)
        research_formats_exact = all(column.get("format") for column in detail_view_columns)
    main_ids = {column["name"]: column["id"] for column in main_view_columns}
    detail_ids = {column["name"]: column["id"] for column in detail_view_columns}
    conditional_targets = {}
    for rule in plan["conditionalFormats"]:
        field = rule["field"].strip()
        if field in main_ids:
            conditional_targets[rule["id"]] = {"elementId": f"{workflow}-main-view",
                                                "columnId": main_ids[field]}
        elif field in detail_ids:
            conditional_targets[rule["id"]] = {"elementId": f"{workflow}-detail-view",
                                                "columnId": detail_ids[field]}
        else:
            conditional_targets[rule["id"]] = {
                "elementId": source_id,
                "columnId": re.sub(r"[^a-z0-9]+", "_", field.lower()).strip("_") or "value",
            }
    validation_targets = {}
    for rule in plan["validations"]:
        if workflow == "research" and rule["sheet"] == "Model Inputs" and rule["range"] == "C2:C7":
            validation_targets[rule["id"]] = [
                {"elementId": "research-input-existing", "columnId": f"i_model_inputs_c{row}"}
                for row in range(2, 8)
            ]
            continue
        targets = []
        for field in rule["fields"]:
            target = columns.get(f"{rule['sheet']}!{field}")
            if target:
                targets.append(target)
        validation_targets[rule["id"]] = targets
    return {
        "workflow": workflow, "charts": charts, "columns": columns,
        "mainElementId": f"{workflow}-main-view", "editElementId": f"{workflow}-input",
        "detailElementId": f"{workflow}-detail-view", "viewSourceElementId": source_id,
        "mainViewColumns": main_view_columns, "detailViewColumns": detail_view_columns,
        "conditionalTargets": conditional_targets,
        "validationTargets": validation_targets,
        "sharedInputScope": "shared" if workflow == "research" else None,
        "sharedInputColumns": detail, "prominentEditableColumns": edit,
        "researchMetricFormatsExact": research_formats_exact,
        "businessInput": {"pageId": f"{workflow}-edit", "connectionId": "ASSIGN_FROM_RUNTIME",
                          "folderId": "ASSIGN_FROM_RUNTIME",
                          "name": f"{workflow.title()} Business Inputs", "editMode": "published-explore",
                          **({"elementId": "research-input-existing", "rowId": "ASSIGN_FROM_RUNTIME",
                              "scopeColumnId": "ASSIGN_FROM_RUNTIME"} if workflow == "research" else {})},
        "sourceHash": plan["source"]["sha256"], "planHash": "ASSIGN_HASH_BOUND_PLAN",
        "stagingReceiptHash": "ASSIGN_APPROVED_STAGING_RECEIPT_HASH",
        "documentId": "ASSIGN_FROM_RUNTIME", "documentVersion": "ASSIGN_FROM_RUNTIME",
    }


def extract_business_input(plan):
    """Extract an ordinary input table without copying formula/cache columns."""
    workflow = plan["workflow"]
    if workflow == "research":
        workbook = load_workbook(plan["source"]["path"], data_only=False, read_only=False)
        ws = workbook["FY results"]
        additional = []
        for column in range(column_index_from_string("B"), column_index_from_string("M") + 1):
            cell = ws.cell(56, column)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                raise ValueError(f"research nondependency input unexpectedly became a formula: {cell.coordinate}")
            additional.append({
                "id": f"i_fy_results_{cell.coordinate.lower()}",
                "name": f"i_fy_results_{cell.coordinate.lower()}", "label": f"FY results!{cell.coordinate}",
                "sourceCellId": f"FY results!{cell.coordinate}", "value": cell.value,
                "type": sigma_input_type(cell.value), "role": "nondependency_line_input_literal",
                "mappingStatus": "live",
            })
        workbook.close()
        return {
            "workflow": workflow, "mode": "extend-existing", "expectedScope": "shared",
            "existingSourceInputColumns": 133, "requiredSourceInputColumns": 145,
            "expectedExistingEditableColumns": 134, "expectedEditableColumns": 146,
            "additionalColumns": additional, "sourceHash": plan["source"]["sha256"],
            "extractionHash": plan["extractionHash"],
            "reason": "preserve the Task2 physical table and add unreferenced source-authored line inputs",
        }
    sheet_name, table_name, key_columns = BUSINESS_TABLES[workflow]
    workbook = load_workbook(plan["source"]["path"], data_only=False, read_only=False)
    ws = workbook[sheet_name]
    table = ws.tables[table_name]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    headers = [str(ws.cell(min_row, column).value) for column in range(min_col, max_col + 1)]
    literal_headers, derived = [], []
    for offset, header in enumerate(headers):
        column = min_col + offset
        formulas = [ws.cell(row, column).value for row in range(min_row + 1, max_row + 1)
                    if isinstance(ws.cell(row, column).value, str)
                    and ws.cell(row, column).value.startswith("=")]
        if formulas:
            derived.append({"name": header, "sourceCells": [f"{sheet_name}!{get_column_letter(column)}{row}"
                                                              for row in range(min_row + 1, max_row + 1)],
                            "sourceFormulas": formulas, "status": "live-calculation-required"})
        else:
            literal_headers.append((header, column))
    rows = []
    for row in range(min_row + 1, max_row + 1):
        record = {header: ws.cell(row, column).value for header, column in literal_headers}
        if any(value is not None for value in record.values()):
            rows.append(record)
    columns = []
    for header, column in literal_headers:
        sample = next((row[header] for row in rows if row.get(header) is not None), None)
        columns.append({"name": header, "type": sigma_input_type(sample),
                        "sourceColumn": f"{sheet_name}!{get_column_letter(column)}",
                        "provenance": "source-literal"})
    workbook.close()
    result = {
        "workflow": workflow, "mode": "construct-ordinary-input", "sheet": sheet_name,
        "table": table_name, "sourceRange": table.ref, "keyColumns": key_columns,
        "columns": columns, "rows": rows, "derivedColumns": derived,
        "sourceHash": plan["source"]["sha256"], "extractionHash": plan["extractionHash"],
    }
    result["inputArtifactHash"] = _digest(result)
    return result


def build_business_operational_wiring(plan, binding):
    """Build the reviewed PR2 sales/expense joins from the source workbook."""
    workflow = plan["workflow"]
    if workflow not in {"sales", "expense"}:
        return None
    for field in ("connectionId", "folderId", "elementId", "name"):
        if not binding.get(field):
            raise ValueError(f"{workflow} operational wiring requires {field}")
    workbook = load_workbook(plan["source"]["path"], data_only=False, read_only=False)
    if workflow == "sales":
        source_plan = {"sharedInputSources": {"salesSettings": {
            "sourceBinding": "Assumptions!F2",
            "rows": [{"Setting": "AsOfDate", "AsOfDate": workbook["Assumptions"]["F2"].value}],
            "sourceElementId": binding["elementId"], "sourceElementName": binding["name"],
        }}}
    else:
        ws = workbook["Policy"]
        min_col, min_row, max_col, max_row = range_boundaries("A1:D7")
        headers = [str(ws.cell(min_row, column).value) for column in range(min_col, max_col + 1)]
        rows = [dict(zip(headers, [ws.cell(row, column).value
                                  for column in range(min_col, max_col + 1)]))
                for row in range(min_row + 1, max_row + 1)]
        source_plan = {"sharedInputSources": {"expensePolicy": {
            "sourceRange": "Policy!A1:D7", "rows": rows,
            "sourceElementId": binding["elementId"], "sourceElementName": binding["name"],
        }}}
    workbook.close()
    architecture = build_shared_input_architecture(
        source_plan, connection_id=binding["connectionId"], folder_id=binding["folderId"],
        workflows=[workflow])
    return {"architecture": architecture, "fragments": build_operational_input_elements(architecture),
            "requires": ["construct shared input from hash-bound PR2 manifest/receipt",
                         "replace logical shared-input element ID with assigned readback ID",
                         "merge native calculation tables with a guarded full-contents PUT"]}


def emit_business_input_program(spec, binding):
    """Emit an initial ordinary-input construction program; no formula cache rows."""
    if spec.get("mode") == "extend-existing":
        for field in ("pageId", "elementId", "rowId", "scopeColumnId"):
            if not (binding or {}).get(field):
                raise ValueError(f"research input extension requires {field}")
        additional = spec["additionalColumns"]
        extension_hash = _digest({"sourceHash": spec["sourceHash"], "elementId": binding["elementId"],
                                  "rowId": binding["rowId"], "additionalColumns": additional,
                                  "planHash": binding.get("planHash"),
                                  "stagingReceiptHash": binding.get("stagingReceiptHash")})
        lines = [f"# sourceHash={spec['sourceHash']}", f"# extractionHash={spec['extractionHash']}",
                 f"# planHash={binding.get('planHash')}",
                 f"# stagingReceiptHash={binding.get('stagingReceiptHash')}",
                 f"# inputExtensionHash={extension_hash}",
                 f"host = sigma.workbook.get_page({binding['pageId']!r})",
                 "if not isinstance(host, Page):", "    raise RuntimeError('expected current page host')",
                 f"table = host.get_element({binding['elementId']!r})",
                 "if not isinstance(table, InputTable):", "    raise RuntimeError('expected existing research input table')"]
        for index, column in enumerate(additional):
            lines.append(f"n{index} = table.add_editable_column({column['name']!r}, {column['type']!r})")
        fields = ", ".join(f"n{index}: {_sdk_literal(column['value'], column['type'])}"
                           for index, column in enumerate(additional))
        lines += [f"table.update_rows([({binding['rowId']!r}, {{{fields}}})])",
                  "print('inputTableId=' + table.id())"]
        for index, column in enumerate(additional):
            lines.append(f"print({('column:' + column['name'] + '=')!r} + n{index}.id())")
        return {"mode": "extend-existing", "elementId": binding["elementId"],
                "expectedExistingEditableColumns": spec["expectedExistingEditableColumns"],
                "expectedEditableColumns": spec["expectedEditableColumns"],
                "extensionHash": extension_hash, "additionalColumns": additional,
                "program": "\n".join(lines) + "\n"}
    if binding.get("elementId"):
        return {"mode": "reconcile-existing", "elementId": binding["elementId"],
                "sourceHash": spec["sourceHash"], "artifactHash": spec["inputArtifactHash"],
                "keyColumns": spec["keyColumns"], "sourceRows": spec["rows"],
                "nextStep": "query current rows, call input_sync.plan_seed, then emit hash-bound writes; do not recreate"}
    for field in ("pageId", "connectionId", "name"):
        if not (binding or {}).get(field):
            raise ValueError(f"business input binding requires {field}")
    columns = spec["columns"]
    lines = [f"# sourceHash={spec['sourceHash']}", f"# extractionHash={spec['extractionHash']}",
             f"# planHash={binding.get('planHash')}",
             f"# stagingReceiptHash={binding.get('stagingReceiptHash')}",
             f"# inputArtifactHash={spec['inputArtifactHash']}",
             f"host = sigma.workbook.get_page({binding['pageId']!r})",
             "if not isinstance(host, Page):", "    raise RuntimeError('expected current page host')"]
    create = f"created = host.add_input_table({binding['connectionId']!r}"
    if binding.get("writebackSchemaId"):
        create += f", writeback_schema_id={binding['writebackSchemaId']!r}"
    edit_mode = binding.get("editMode", "published-explore")
    if edit_mode not in {"draft", "published-explore", "published-all-users"}:
        raise ValueError("business input editMode is not a supported explicit publish policy")
    lines += [create + ")", f"created.set_name({binding['name']!r})",
              f"created.set_edit_mode({edit_mode!r})", "created.set_frame(1, 4, 24, 24)"]
    for index, column in enumerate(columns):
        lines.append(f"c{index} = created.add_editable_column({column['name']!r}, {column['type']!r})")
    by_name = {column["name"]: f"c{index}" for index, column in enumerate(columns)}
    live_formulas = []
    if spec["workflow"] == "sales":
        live_formulas = [
            ("Weighted ACV", f"sigma.formula.multiply({by_name['Probability']}, {by_name['ACV']})"),
            ("Cycle Days", f"sigma.formula.date_diff('day', {by_name['Created Date']}, {by_name['Expected Close']})"),
        ]
    elif spec["workflow"] == "inventory":
        available = f"sigma.formula.subtract({by_name['On Hand']}, {by_name['Allocated']})"
        reorder_point = (f"sigma.formula.add(sigma.formula.multiply({by_name['Avg Daily Use']}, "
                         f"{by_name['Lead Days']}), {by_name['Safety Stock']})")
        reorder_raw = ("sigma.formula.subtract(sigma.formula.add(d1, "
                       f"sigma.formula.multiply({by_name['Avg Daily Use']}, sigma.formula.const_num(30))), d0)")
        live_formulas = [
            ("Available", available),
            ("Reorder Point", reorder_point),
            ("Days Cover", f"sigma.formula.divide(d0, {by_name['Avg Daily Use']})"),
            ("Reorder Qty", "sigma.formula.if_(sigma.formula.gt(" + reorder_raw
             + ", sigma.formula.const_num(0)), sigma.formula.round_up(" + reorder_raw
             + ", 0), sigma.formula.const_num(0))"),
            ("Status", "sigma.formula.if_(sigma.formula.lte(d0, d1), "
             "sigma.formula.const_text('REORDER'), sigma.formula.const_text('OK'))"),
        ]
    for index, (name, expression) in enumerate(live_formulas):
        lines.append(f"d{index} = created.add_formula_column({name!r}, {expression})")
    lines += ["created_column_ids = [" + ", ".join(f"c{i}.id()" for i in range(len(columns))) + "]",
              "created.add_system_columns(['row_id'])", "table = host.get_element(created.id())",
              "if not isinstance(table, InputTable):", "    raise RuntimeError('expected ordinary input table')"]
    if spec["rows"]:
        lines.append("row_ids = table.upsert_rows([")
        for row in spec["rows"]:
            fields = ", ".join(f"c{i}: {_sdk_literal(row.get(column['name']), column['type'])}"
                               for i, column in enumerate(columns))
            lines.append("    {" + fields + "},")
        lines.append("])")
    lines += ["print('inputTableId=' + table.id())"]
    for index, column in enumerate(columns):
        lines.append(f"print({('column:' + column['name'] + '=')!r} + c{index}.id())")
    target_hash = _digest({key: binding.get(key) for key in
                           ("pageId", "connectionId", "writebackSchemaId", "name", "editMode",
                            "planHash", "stagingReceiptHash")})
    return {"mode": "construct-ordinary-input", "feature": f"{spec['workflow']}BusinessInput",
            "sourceHash": spec["sourceHash"], "artifactHash": spec["inputArtifactHash"],
            "targetBindingHash": target_hash,
            "keyColumns": spec["keyColumns"], "sourceRows": spec["rows"],
            "derivedColumns": spec["derivedColumns"],
            "liveFormulaColumns": [name for name, _expression in live_formulas],
            "downstreamWiring": (
                "existing forecast builder" if spec["workflow"] == "forecast"
                else "PR2 sales settings join supplies TimingStatus" if spec["workflow"] == "sales"
                else "native input formula columns" if spec["workflow"] == "inventory"
                else "PR2 expense policy join supplies PolicyExcess and Exception"),
            "program": "\n".join(lines) + "\n"}


def _layout(workflow, charts, bindings):
    overview, edit, detail = (f"{workflow}-overview", f"{workflow}-edit", f"{workflow}-detail")
    chart_ids = [chart["id"] for chart in charts]
    pages = []
    overview_items = [(f"{workflow}-overview-title", 1, 25, 1, 4)]
    if len(chart_ids) >= 1: overview_items.append((chart_ids[0], 1, 13, 4, 20))
    if len(chart_ids) >= 2: overview_items.append((chart_ids[1], 13, 25, 4, 20))
    overview_items.append((bindings["mainElementId"], 1, 25, 20, 38))
    page_items = {
        overview: overview_items,
        edit: [(f"{workflow}-edit-title", 1, 25, 1, 4), (bindings["editElementId"], 1, 25, 4, 28)],
        detail: [(f"{workflow}-detail-title", 1, 25, 1, 4), (bindings["detailElementId"], 1, 25, 4, 42)],
    }
    for page_id, items in page_items.items():
        lines = [f'<Page type="grid" gridTemplateColumns="repeat(24, 1fr)" gridTemplateRows="auto" id="{page_id}">']
        for eid, c1, c2, r1, r2 in items:
            lines.append(f'  <Element elementId="{eid}" gridColumn="{c1} / {c2}" gridRow="{r1} / {r2}"/>')
        lines.append("</Page>")
        pages.append("\n".join(lines))
    return '<?xml version="1.0" encoding="utf-8"?>\n' + "\n".join(pages)


def emit_presentation(plan, bindings):
    """Emit mergeable current-code-representation fragments from explicit bindings."""
    if bindings.get("workflow") != plan.get("workflow"):
        raise ValueError("presentation bindings target a different workflow")
    if bindings.get("sourceHash") != plan["source"]["sha256"]:
        raise ValueError("presentation bindings sourceHash is stale")
    for field in ("planHash", "stagingReceiptHash", "documentId", "documentVersion"):
        if not bindings.get(field):
            raise ValueError(f"presentation bindings require {field}")
    workflow = plan["workflow"]
    ownership = f"excel-to-sigma:{workflow}:{plan['source']['sha256']}"
    elements, unresolved, aggregate_checks = [], [] , []
    for role in ("overview", "edit", "detail"):
        elements.append({"id": f"{workflow}-{role}-title", "kind": "text",
                         "body": f"<!-- {ownership} -->\n# {workflow.title()} {role.title()}",
                         "verticalAlign": "middle"})
    for role, element_id, columns in (
            ("main", bindings["mainElementId"], bindings.get("mainViewColumns") or []),
            ("detail", bindings["detailElementId"], bindings.get("detailViewColumns") or [])):
        if not columns:
            raise ValueError(f"presentation bindings require {role}ViewColumns")
        elements.append({"id": element_id, "kind": "table", "name": f"{workflow.title()} {role.title()}",
                         "description": ownership,
                         "source": {"kind": "table", "elementId": bindings["viewSourceElementId"]},
                         "columns": copy.deepcopy(columns),
                         "style": {"backgroundColor": "#FFFFFF", "borderColor": {"kind": "theme", "ref": "colors-border"},
                                   "borderWidth": 1, "borderRadius": "round", "padding": "medium"}})
    emitted_charts = []
    for chart_index, source in enumerate(plan["charts"], 1):
        binding = (bindings.get("charts") or {}).get(source["id"])
        if not binding:
            unresolved.append({"kind": "chart", "sourceId": source["id"],
                               "title": source["title"], "reason": "missing live target binding"})
            continue
        formulas = binding.get("seriesFormulas") or []
        if len(formulas) != len(source["series"]):
            unresolved.append({"kind": "chart", "sourceId": source["id"],
                               "title": source["title"], "reason": "series binding count mismatch"})
            continue
        chart_id = f"{workflow}-chart-{chart_index}"
        columns = [{"id": f"chart_{chart_index}_category", "name": "Category",
                    "formula": binding["categoryFormula"]}]
        y_ids = []
        for series_index, (series, formula) in enumerate(zip(source["series"], formulas), 1):
            column_id = f"chart_{chart_index}_series_{series_index}"
            columns.append({"id": column_id, "name": series["name"], "formula": formula,
                            "format": _series_format(series)})
            y_ids.append(column_id)
            for point_index, (category, expected) in enumerate(zip(source["categories"], series["expectedValues"])):
                aggregate_checks.append({"id": f"{source['id']}-s{series_index}-p{point_index + 1}",
                                         "chartTitle": source["title"], "series": series["name"],
                                         "category": category, "expected": expected,
                                         "sourceFormula": series["sourceFormulas"][point_index],
                                         "liveFormula": formula, "unit": series["unit"],
                                         "order": point_index, "status": "pending-runtime-query"})
        order_id = f"chart_{chart_index}_order"
        columns.append({"id": order_id, "name": "Source Order", "formula": binding["orderFormula"]})
        chart = {"id": chart_id, "kind": f"{source['chartType']}-chart", "name": source["title"],
                 "description": ownership,
                 "source": {"kind": "table", "elementId": binding["sourceElementId"]},
                 "columns": columns,
                 "xAxis": {"columnId": columns[0]["id"], "sort": {"direction": "ascending", "by": order_id, "aggregation": "min"}},
                 "yAxis": {"columnIds": y_ids},
                 "style": {"backgroundColor": "#FFFFFF", "borderColor": {"kind": "theme", "ref": "colors-border"},
                           "borderWidth": 1, "borderRadius": "round", "padding": "medium"}}
        elements.append(chart); emitted_charts.append(chart)

    patches = defaultdict(lambda: {"columnUpdates": {}, "conditionalFormats": []})
    for validation in plan["validations"]:
        targets = (bindings.get("validationTargets") or {}).get(validation["id"])
        if targets is None:
            targets = [target for field in validation.get("fields") or [validation["field"]]
                       if (target := (bindings.get("columns") or {}).get(f"{validation['sheet']}!{field}"))]
        if validation["status"] != "supported" or not targets:
            unresolved.append({"kind": "validation", "sourceId": validation["id"],
                               "reason": "unsupported source rule" if validation["status"] != "supported" else "missing editable column binding"})
            continue
        for target in targets:
            patches[target["elementId"]]["columnUpdates"].setdefault(target["columnId"], {}).update(
                copy.deepcopy(validation["emission"]))
        if validation["allowBlank"] is False:
            unresolved.append({"kind": "validation-null-behavior", "sourceId": validation["id"],
                               "reason": "current input-column schema has no allowBlank field; runtime UI rejection remains required"})
    for rule in plan["conditionalFormats"]:
        target = ((bindings.get("conditionalTargets") or {}).get(rule["id"])
                  or (bindings.get("columns") or {}).get(f"{rule['sheet']}!{rule['field']}"))
        if rule["status"] != "supported" or not target:
            unresolved.append({"kind": "conditionalFormat", "sourceId": rule["id"],
                               "reason": "unsupported source rule" if rule["status"] != "supported" else "missing target column binding"})
            continue
        emitted = copy.deepcopy(rule["emission"])
        emitted["columnIds"] = [target["columnId"]]
        if emitted["type"] == "backgroundScale":
            for key, value in list(emitted["domain"].items()):
                if isinstance(value, str): emitted["domain"][key] = value.replace("[Column]", f"[{target['columnId']}]")
        patches[target["elementId"]]["conditionalFormats"].append((rule["priority"], emitted))
    for patch in patches.values():
        # Sigma says later rules take precedence. Excel priority 1 wins, so emit
        # larger priorities first and priority 1 last.
        patch["conditionalFormats"] = [item for _priority, item in sorted(
            patch["conditionalFormats"], key=lambda pair: pair[0], reverse=True)]

    # Table/input-table column display formats are applied from the first body
    # cell of each source table column. The source format record remains in the
    # extraction for exact precision/date-grain audit.
    for _source_field, target in (bindings.get("columns") or {}).items():
        semantic = plan["semanticFormats"].get(target.get("sourceCell"))
        if semantic:
            patches[target["elementId"]]["columnUpdates"].setdefault(target["columnId"], {}).setdefault(
                "format", copy.deepcopy(semantic["emission"]))

    main = list(MAIN_COLUMNS.get(workflow, []))
    detail = list(bindings.get("sharedInputColumns") or DETAIL_COLUMNS.get(workflow, main))
    edit = list(bindings.get("prominentEditableColumns") or EDITABLE_COLUMNS.get(workflow, []))
    pages = [{"id": f"{workflow}-{role}", "name": role.title()} for role in ("overview", "edit", "detail")]
    business_spec = extract_business_input(plan)
    business_binding = {**(bindings.get("businessInput") or {}),
                        "planHash": bindings["planHash"],
                        "stagingReceiptHash": bindings["stagingReceiptHash"]}
    business_program = emit_business_input_program(business_spec, business_binding)
    operational_binding = {**business_binding,
                           "elementId": business_binding.get("elementId") or bindings["editElementId"]}
    operational_wiring = build_business_operational_wiring(plan, operational_binding)
    chart_formula_cells = {cell for chart in plan["charts"] for series in chart["series"]
                           for cell in series.get("sourceFormulaCells") or []}
    formula_accounting = [
        {**item,
         "targetStatus": ("mapped-live-chart-check" if item["cell"] in chart_formula_cells
                           else "pending-runtime-parity"),
         "verification": ("aggregateChecks" if item["cell"] in chart_formula_cells
                          else "source formula retained; target parity evidence required")}
        for item in plan["formulaProvenance"]
    ]
    if workflow == "research" and not bindings.get("researchMetricFormatsExact"):
        unresolved.append({"kind": "mixed-unit-row-format", "sourceId": "FY results-values",
                           "reason": "current column format is uniform; currency, percentage, ratio, and count rows require separate bound views or UI verification"})
    artifact = {
        "version": 1, "sourceHash": plan["source"]["sha256"], "extractionHash": plan["extractionHash"],
        "planHash": bindings["planHash"], "stagingReceiptHash": bindings["stagingReceiptHash"],
        "workflow": workflow, "pages": pages, "elements": elements,
        "ownershipMarker": ownership,
        "layout": _layout(workflow, emitted_charts, bindings),
        "pageRoles": {role: f"{workflow}-{role}" for role in ("overview", "edit", "detail")},
        "referencedElementIds": sorted(set(
            [bindings["editElementId"], bindings["viewSourceElementId"]]
            + [binding["sourceElementId"] for binding in (bindings.get("charts") or {}).values()])),
        "elementPatches": dict(patches), "aggregateChecks": aggregate_checks,
        "ruleProvenance": copy.deepcopy(plan["conditionalFormats"]),
        "formulaAccounting": formula_accounting,
        "editableColumns": list(EDITABLE_COLUMNS.get(workflow, [])),
        "principalCalculatedColumns": list(CALCULATED_COLUMNS.get(workflow, [])),
        "mainGridColumns": main, "editColumns": edit, "detailColumns": detail,
        "sharedInputScope": bindings.get("sharedInputScope"), "unresolved": unresolved,
        "businessInput": business_spec, "businessInputProgram": business_program,
        "businessOperationalWiring": operational_wiring,
        "restPlan": {"method": "PUT", "path": "/v2/workbooks/{workbookId}/contents",
                     "requires": ["complete current contents", "documentVersion", "runtime element/column IDs"],
                     "documentId": bindings.get("documentId"), "documentVersion": bindings.get("documentVersion")},
    }
    artifact["artifactHash"] = _digest({key: value for key, value in artifact.items() if key != "artifactHash"})
    return artifact


def _series_format(series):
    unit = series.get("unit")
    if unit == "currency":
        return {"kind": "number", "formatString": ",.0f", "currencySymbol": "$"}
    if unit == "percentage":
        return {"kind": "number", "formatString": ".1%"}
    if unit == "count":
        return {"kind": "number", "formatString": ",.0f"}
    return {"kind": "number", "formatString": ",.2f" if unit == "ratio" else ",.0f"}


def integrate_presentation(contents, artifact):
    """Merge an emitted artifact into a full GET-back contents document.

    The caller must use a version-guarded full PUT. Missing referenced runtime
    elements/columns fail closed; this function performs no network calls.
    """
    result = copy.deepcopy(contents)
    existing_elements = {element.get("id"): element for element in result.get("elements", [])}
    missing = [eid for eid in artifact["referencedElementIds"] if eid not in existing_elements]
    if missing:
        raise ValueError("presentation references missing runtime elements: " + ", ".join(missing))
    marker = artifact["ownershipMarker"]

    def owned(element):
        return element.get("description") == marker or marker in str(element.get("body") or "")

    new_ids = [element["id"] for element in artifact["elements"]]
    collisions = set(new_ids) & set(existing_elements)
    foreign = [element_id for element_id in collisions if not owned(existing_elements[element_id])]
    if foreign:
        raise ValueError("presentation element IDs collide with foreign contents: " + ", ".join(sorted(foreign)))
    result["elements"] = [element for element in result.get("elements", [])
                          if element.get("id") not in collisions]
    result["elements"].extend(copy.deepcopy(artifact["elements"]))
    page_ids = {page["id"] for page in artifact["pages"]}
    existing_pages = {page.get("id") for page in result.get("pages", [])}
    if page_ids & existing_pages and not collisions:
        raise ValueError("presentation page IDs collide without owned elements")
    result["pages"] = [page for page in result.get("pages", []) if page.get("id") not in page_ids]
    result["pages"].extend(copy.deepcopy(artifact["pages"]))
    layout = result.get("layout", "")
    for page_id in page_ids:
        layout = re.sub(rf'<Page\b(?=[^>]*\bid="{re.escape(page_id)}")[^>]*>.*?</Page>\s*', "", layout,
                        flags=re.S)
    presentation_layout = re.sub(r'^<\?xml[^>]*>\s*', "", artifact["layout"])
    result["layout"] = layout.rstrip() + "\n" + presentation_layout
    existing_elements = {element.get("id"): element for element in result.get("elements", [])}
    for element_id, patch in artifact["elementPatches"].items():
        element = existing_elements.get(element_id)
        if element is None:
            raise ValueError(f"presentation patch target is missing: {element_id}")
        columns = {column.get("id"): column for column in element.get("columns", [])}
        for column_id, update in patch["columnUpdates"].items():
            if column_id not in columns:
                raise ValueError(f"presentation patch column is missing: {element_id}/{column_id}")
            columns[column_id].update(copy.deepcopy(update))
        prior = element.setdefault("conditionalFormats", [])
        for rule in patch["conditionalFormats"]:
            if rule not in prior:
                prior.append(copy.deepcopy(rule))
    return result


def audit_presentation_openapi(document):
    """Audit the exact current fields used by presentation artifacts."""
    property_sets, literals = [], set()

    def walk(value):
        if isinstance(value, dict):
            properties = value.get("properties")
            if isinstance(properties, dict):
                property_sets.append(set(properties))
            enum = value.get("enum")
            if isinstance(enum, list):
                literals.update(item for item in enum if isinstance(item, str))
            for child in value.values():
                walk(child)
        elif isinstance(value, list):
            for child in value:
                walk(child)

    walk((document.get("components") or {}).get("schemas") or {})
    has = lambda required: any(set(required).issubset(props) for props in property_sets)
    capabilities = {
        "nativeCharts": {"bar-chart", "line-chart"}.issubset(literals)
                        and all(any(name in props for props in property_sets)
                                for name in ("xAxis", "yAxis", "columns", "columnIds")),
        "inputAllowedValuesAndRanges": has({"values", "valuesFrom", "range"}),
        "conditionalRules": {"single", "backgroundScale", "fontScale"}.issubset(literals)
                            and has({"columnIds"}) and has({"scheme", "domain", "order", "nullBehavior"}),
        "semanticFormats": {"number", "datetime"}.issubset(literals)
                           and has({"formatString"}) and has({"currencySymbol"}),
        "fullContentsUpdate": "/v2/workbooks/{workbookId}/contents" in (document.get("paths") or {}),
    }
    missing = [name for name, value in capabilities.items() if not value]
    return {"status": "PASS" if not missing else "DRIFT", "capabilities": capabilities,
            "missing": missing, "openapiVersion": (document.get("info") or {}).get("version")}
